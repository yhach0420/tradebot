"""Write report.json / report.md / audit.xlsx only. No CSV."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.full_causal_mechanism_discovery_v1.analyze import public_row
from research.post_open_native_discontinuous_up_repricing_full_strategy_v1 import ANALYSIS_ID, STRATEGY_ID
from research.post_open_native_discontinuous_up_repricing_full_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Parent_Pin",
    "Data_Boundary",
    "Status_Semantics",
    "Status_Distribution",
    "Raw_Support",
    "Discontinuous_Events",
    "ISQ_Overlap",
    "Duplicate_Check",
    "Precommit",
    "Synthetic_Tests",
    "Canary",
    "Episodes",
    "Signals",
    "Execution",
    "Trades",
    "Daily",
    "Blocks",
    "Coverage_Gates",
    "Economic_Gates",
    "Concentration",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    evals = [public_row(e) for e in list(report.get("candidate_evals") or [])]
    ev = (report.get("candidate_evals") or [{}])[0] if report.get("candidate_evals") else {}
    extra = dict(ev.get("extra") or {})
    g = dict(ev.get("g_table") or {})
    trades = []
    daily_rows = []
    block_rows = []
    for e in list(report.get("candidate_evals") or []):
        for t in list(e.get("_trades") or []):
            trades.append(
                {
                    "STRATEGY_ID": e.get("STRATEGY_ID"),
                    "date": t.get("date"),
                    "symbol": t.get("symbol"),
                    "fill_t": t.get("fill_time") or t.get("fill_t"),
                    "exit_t": t.get("exit_time") or t.get("exit_t"),
                    "pnl": t.get("pnl_yen_100"),
                    "exit_reason": t.get("exit_reason"),
                    "episode_id": (t.get("src") or {}).get("episode_id"),
                }
            )
        for day, pnl in dict(e.get("daily") or {}).items():
            daily_rows.append({"STRATEGY_ID": e.get("STRATEGY_ID"), "date": day, "pnl": pnl})
        for b in list((e.get("blocks") or {}).get("blocks") or []):
            block_rows.append({"STRATEGY_ID": e.get("STRATEGY_ID"), **b})
    return {
        "Summary": _kv(dict(report.get("answers") or {})),
        "Parent_Pin": _kv(dict(report.get("parent") or {})),
        "Data_Boundary": _kv(dict(report.get("data_boundary") or {})),
        "Status_Semantics": _kv(dict(report.get("status_semantics") or {})),
        "Status_Distribution": _kv(dict((report.get("raw_support") or {}).get("status_distribution") or {}))
        or [{"empty": True, "reason": "not_run"}],
        "Raw_Support": _kv(dict(report.get("raw_support") or {})) or [{"empty": True, "reason": "not_run"}],
        "Discontinuous_Events": _kv(dict(report.get("structural") or {})) or [{"empty": True, "reason": "not_run"}],
        "ISQ_Overlap": _kv(dict(report.get("isq_overlap") or {})) or [{"empty": True, "reason": "not_run"}],
        "Duplicate_Check": list((report.get("duplicate_check") or {}).get("FAMILY_ROWS") or [])
        or _kv(dict(report.get("duplicate_check") or {}))
        or [{"empty": True, "reason": "not_run"}],
        "Precommit": _kv(dict(report.get("precommit") or {})),
        "Synthetic_Tests": list((report.get("unit_tests") or {}).get("rows") or [])
        or _kv(dict(report.get("unit_tests") or {})),
        "Canary": _kv(dict(report.get("canary") or {})),
        "Episodes": list(report.get("episodes") or [])[:5000] or [{"empty": True}],
        "Signals": list(report.get("signals") or [])[:5000] or [{"empty": True}],
        "Execution": [
            {
                "STRATEGY_ID": STRATEGY_ID,
                "X1_fill_n": extra.get("X1_fill_n") or ev.get("X1_fill_n"),
                "X1_nonfill_n": extra.get("X1_nonfill_n"),
                "technical_exit_fire_n": extra.get("technical_exit_fire_n"),
                "session_exit_unfilled_n": ev.get("session_exit_unfilled_n"),
            }
        ],
        "Trades": trades or [{"empty": True}],
        "Daily": daily_rows or [{"empty": True}],
        "Blocks": block_rows or [{"empty": True}],
        "Coverage_Gates": [
            {
                "C1": ev.get("C1"),
                "C2": ev.get("C2"),
                "C3": ev.get("C3"),
                "C4": ev.get("C4"),
                "coverage_ok": ev.get("coverage_ok"),
            }
        ],
        "Economic_Gates": evals or _kv({"economics": "not_run"}),
        "Concentration": [
            {
                "CAUSAL_EX_TOP1_PNL": ev.get("CAUSAL_EX_TOP1_PNL"),
                "G6": g.get("G6"),
            }
        ],
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        f"CASE: {d.get('CASE')}",
        f"LOGIC_COMPLETE: {d.get('LOGIC_COMPLETE')}",
        f"ROBUST_DEV_QUALIFIED: {d.get('ROBUST_DEV_QUALIFIED')}",
        "TRUE_OOS: false",
        "CERTIFIED: false",
        "",
        str(d.get("INTERPRETATION") or ""),
        "",
    ]
    for i in range(1, 73):
        matches = [k for k in a if k.split("_", 1)[0] == str(i)]
        if matches:
            lines.append(f"{i}. {matches[0]}: {a.get(matches[0])}")
    lines.extend(
        [
            "",
            "No MBO. No futures. No Sizing. No Holdout/Stress/future. Opening CLOSED. ISQ CLOSED. AOP CLOSED. CalcPrice CLOSED (no retune).",
            "",
            "STOP.",
            "",
        ]
    )
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
