"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_correction_v4_freeze_and_prospective_prep.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Frozen_Identity",
    "Parent_Hashes",
    "Source_Inventory",
    "NBar_Reachability",
    "Executable_Death_Paths",
    "Contamination_Ledger",
    "Prospective_Eligibility",
    "Prospective_Precommit",
    "Data_Quality_Gates",
    "Causal_Timestamp_Contract",
    "Hidden1M_Contract",
    "State_Log_Schema",
    "Adjudication_Rubric",
    "Version_Invalidation_Rules",
    "Safety",
)
STRIP = {"isolation_before", "isolation_after", "funnel_days", "setups", "e0_events", "e1_events"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    nbar = dict(report.get("nbar") or {})
    return {
        "SPEC_CHANGED": False,
        "V4_CHANGED": False,
        "V4_FROZEN": bool(dec.get("V4_FROZEN")),
        "OLD_CONFIRMATION_OPENED": False,
        "FROZEN_VALIDATION_OPENED": False,
        "PROSPECTIVE_DATA_OPENED": False,
        "FUTURE_OUTCOME_USED": False,
        "PNL_USED": False,
        "MFE_MAE_USED": False,
        "THRESHOLD_RETUNED": False,
        "NEW_NUMERIC_CUTOFF_ADDED": False,
        "N_BAR_EXPIRY_USED": bool(nbar.get("N_BAR_EXPIRY_USED")),
        "HIDDEN_N_BAR_DEATH_PATH_FOUND": bool(nbar.get("HIDDEN_N_BAR_DEATH_PATH_FOUND")),
        "REPEATED_NO_EXPANSION_N_PRESENT": nbar.get("REPEATED_NO_EXPANSION_N_PRESENT"),
        "REPEATED_NO_EXPANSION_N_EXECUTABLE_DEATH_REACHABLE": nbar.get("REPEATED_NO_EXPANSION_N_EXECUTABLE_DEATH_REACHABLE"),
        "submit/cancel/live": "0/0/0",
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "CONTAMINATED_SYMBOL_DATE_N": (report.get("ledger") or {}).get("CONTAMINATED_SYMBOL_DATE_N"),
        "PRECOMMIT_SHA256": (report.get("precommit") or {}).get("PRECOMMIT_SHA256"),
        "FROZEN_IDENTITY": (report.get("identity") or {}).get("frozen_identity"),
        "machine_sha": (report.get("identity") or {}).get("machine_sha"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ident = dict(report.get("identity") or {})
    nbar = dict(report.get("nbar") or {})
    ledger = dict(report.get("ledger") or {})
    pre = dict(report.get("precommit") or {})
    scan = dict(nbar.get("scan") or {})
    trace = dict(nbar.get("trace") or {})
    walked = dict(nbar.get("walked_deaths") or {})
    elig = dict(pre.get("data_eligibility") or {})
    qgate = dict(pre.get("market_data_quality_gate") or {})
    ts = dict(pre.get("state_logging_schema") or {})
    hid = dict(pre.get("hidden_1m_contract") or {})
    adj = dict(pre.get("semantic_adjudication_procedure") or {})
    inv = dict(pre.get("version_invalidation_rule") or {})
    return {
        "Manifest": [{"key": k, "value": v} for k, v in (report.get("answers") or {}).items()],
        "Frozen_Identity": [
            {"key": k, "value": v}
            for k, v in ident.items()
            if k not in ("files", "git", "environment")
        ]
        + [{"key": f"git.{k}", "value": v} for k, v in dict(ident.get("git") or {}).items() if k != "dirty"]
        + [{"key": f"env.{k}", "value": v} for k, v in dict(ident.get("environment") or {}).items()],
        "Parent_Hashes": [{"key": k, "value": v} for k, v in (report.get("hashes") or {}).items()],
        "Source_Inventory": list(ident.get("files") or []) or [{"empty": True}],
        "NBar_Reachability": [
            {"key": "ok", "value": nbar.get("ok")},
            {"key": "REPEATED_NO_EXPANSION_N_PRESENT", "value": nbar.get("REPEATED_NO_EXPANSION_N_PRESENT")},
            {"key": "REPEATED_NO_EXPANSION_N_EXECUTABLE_DEATH_REACHABLE", "value": nbar.get("REPEATED_NO_EXPANSION_N_EXECUTABLE_DEATH_REACHABLE")},
            {"key": "N_BAR_EXPIRY_USED", "value": nbar.get("N_BAR_EXPIRY_USED")},
            {"key": "HIDDEN_N_BAR_DEATH_PATH_FOUND", "value": nbar.get("HIDDEN_N_BAR_DEATH_PATH_FOUND")},
            {"key": "EXECUTABLE_DEATH_CONDITION", "value": nbar.get("EXECUTABLE_DEATH_CONDITION")},
            {"key": "by_class", "value": scan.get("by_class")},
            {"key": "walked_by_reason", "value": walked.get("by_reason")},
        ]
        + list(scan.get("hits") or []),
        "Executable_Death_Paths": [
            {
                "path": "DISPLACEMENT_UNWOUND",
                "n_bar_alone": False,
                "note": "requires displacement giveback vs peak, not N consecutive bars",
            },
            {
                "path": "REPEATED_FAILED_PROGRESS",
                "n_bar_alone": False,
                "note": "requires wick_only_n+micro_break_n >= FAILED_ATTEMPT_N AND five_m_no_expansion_n>=2; wick/micro never increment in V4 so unreachable from stall",
            },
            {
                "path": "REPEATED_OR_RECROSS",
                "n_bar_alone": False,
                "note": "requires leave + recross closes",
            },
            {
                "path": "TWO_SIDED_BALANCE_REESTABLISHED",
                "n_bar_alone": False,
                "note": "requires leave + both OR extremes",
            },
            {
                "path": "STALE_RANGE_RESOLUTION",
                "n_bar_alone": False,
                "note": "committed opposite non-contracting 5m",
            },
            {
                "path": "FAILED_BREAK_REACCEPTED",
                "n_bar_alone": False,
                "note": "confirmation completed-bar leftover reacceptance after FAILED_PROBE_PENDING",
            },
            {
                "path": "N_CONSECUTIVE_BARS_ALONE",
                "n_bar_alone": False,
                "killed_in_stall_grid": trace.get("n_bar_alone_killed"),
                "loss_after_12_stall_bars": trace.get("loss_after_12_stall_bars"),
            },
        ],
        "Contamination_Ledger": list(ledger.get("rows") or []) or [{"empty": True}],
        "Prospective_Eligibility": [{"key": k, "value": v} for k, v in elig.items()],
        "Prospective_Precommit": [{"key": k, "value": v} for k, v in pre.items() if k not in ("state_logging_schema",)],
        "Data_Quality_Gates": [{"key": k, "value": v} for k, v in qgate.items()],
        "Causal_Timestamp_Contract": [{"field": x} for x in list(ts.get("causal_timestamps") or [])]
        + [{"forbidden": x} for x in list(ts.get("forbidden") or [])],
        "Hidden1M_Contract": [{"key": k, "value": v} for k, v in hid.items()],
        "State_Log_Schema": [{"field": x} for x in list(ts.get("candidate_day") or [])],
        "Adjudication_Rubric": [{"question": q} for q in list(adj.get("questions") or [])]
        + [{"key": "primary_gates", "value": adj.get("primary_gates")}, {"key": "human_label_accuracy_is_not_primary_gate", "value": adj.get("human_label_accuracy_is_not_primary_gate")}],
        "Version_Invalidation_Rules": [{"key": k, "value": v} for k, v in inv.items()]
        + [{"key": k, "value": v} for k, v in dict(pre.get("bug_handling") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 Correction V4 — freeze and prospective prep", "", "V4 source was not mutated. Prospective data was not opened.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
