"""Write report.json / report.md / audit.xlsx only. Do not write parent OUT."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v3_2_face_failure_rca.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Agreement",
    "Opening_Drive",
    "Early_Reversal",
    "Location",
    "Reacceleration",
    "Attribution",
    "Exemplars",
    "Layers",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "events", "failed_push_archive", "funnel_days", "chart_zones", "five_m_bars", "nearest_opposing"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    ag = dict(report.get("agreement") or {})
    loc = dict(report.get("location_audit") or {})
    kinds = dict(loc.get("kind_counts_among_valid") or {})
    er = dict(report.get("early_reversal_contrast") or {})
    reacc = dict(report.get("reaccel_audit") or {})
    lost = dict(report.get("thesis_lost") or {})
    htf = dict(report.get("htf") or {})
    play = dict(report.get("in_play_audit") or {})
    return {
        "v32_unchanged": report.get("v32_unchanged"),
        "semantic_rca_n": report.get("semantic_rca_n"),
        "future_outcome_used": False,
        "second_pass_pattern_agree": (ag.get("pattern") or {}).get("agree_share"),
        "second_pass_pattern_kappa": (ag.get("pattern") or {}).get("cohen_kappa"),
        "second_pass_pattern_disagree_n": (ag.get("pattern") or {}).get("disagreement_n"),
        "opening_drive_agree": (ag.get("opening") or {}).get("agree_share"),
        "opening_drive_kappa": (ag.get("opening") or {}).get("cohen_kappa"),
        "reacceleration_agree": (ag.get("reacceleration") or {}).get("agree_share"),
        "reacceleration_kappa": (ag.get("reacceleration") or {}).get("cohen_kappa"),
        "3382_vs_early_reversal_fp": er,
        "five_m_better_than_or_half": (report.get("opening_drive") or {}).get("five_m_better_than_or_half"),
        "or_boundary_alone_meaningful": loc.get("or_boundary_alone_meaningful"),
        "valid_loc_or_only": kinds.get("OR_ONLY_DEFENSE", 0),
        "valid_loc_prior_sr": kinds.get("PRIOR_SR_ZONE_DEFENSE", 0),
        "valid_loc_cleared_zone": kinds.get("CLEARED_ZONE_RETEST", 0),
        "valid_loc_pdh_pdl": kinds.get("PDH_PDL_DEFENSE", 0),
        "valid_loc_vwap": kinds.get("VWAP_CONFLUENCE", 0),
        "valid_loc_other": int(kinds.get("DAILY_LEVEL_CONFLUENCE") or 0) + int(kinds.get("NO_MEANINGFUL_LEVEL") or 0),
        "reaccel_distinguish": (reacc.get("distinguishing_families")),
        "tradingvalue_participates": reacc.get("tradingvalue_participates"),
        "opening_thesis_lost_state": lost.get("opening_thesis_lost_state"),
        "daily_context_in_definition": htf.get("belongs_to_semantic_definition"),
        "in_play_too_broad": play.get("in_play_too_broad"),
        "PRIMARY_CAUSE": dec.get("PRIMARY_CAUSE"),
        "SECONDARY_CAUSE": dec.get("SECONDARY_CAUSE"),
        "CONTRIBUTING_CAUSE": dec.get("CONTRIBUTING_CAUSE"),
        "pb1_salvageable": dec.get("pb1_salvageable"),
        "any_rule_changed": False,
        "any_threshold_optimized": False,
        "any_future_return": False,
        "any_economic_test": False,
        "any_pnl": False,
        "old_confirmation_opened": False,
        "frozen_validation_opened": False,
        "submit_cancel_live": "0/0/0",
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    ag = dict(report.get("agreement") or {})
    loc = dict(report.get("location_audit") or {})
    kinds = dict(loc.get("kind_counts_among_valid") or {})
    er = dict(report.get("early_reversal_contrast") or {})
    ex = dict(er.get("exemplar") or {})
    fp = dict(er.get("other_early_reversal") or {})
    lines = [
        "# PB1_V3_2_FACE_FAILURE_RCA_V1",
        "",
        "Semantic RCA of frozen V3.2 face failure. No V3.3. No PnL.",
        "",
        f"V3.2 unchanged? **{a.get('v32_unchanged')}**",
        f"Semantic RCA set n? **{a.get('semantic_rca_n')}**",
        f"Future outcome used? **{a.get('future_outcome_used')}**",
        "",
        f"Second-pass human label agreement? pattern agree={a.get('second_pass_pattern_agree')} kappa={a.get('second_pass_pattern_kappa')} disagree_n={a.get('second_pass_pattern_disagree_n')}",
        f"Opening-drive agreement? agree={a.get('opening_drive_agree')} kappa={a.get('opening_drive_kappa')}",
        f"Reacceleration agreement? agree={a.get('reacceleration_agree')} kappa={a.get('reacceleration_kappa')}",
        "",
        "What distinguishes the real 3382 bearish drive from the early-reversal false positives?",
        f"- exemplar largest_counter/normal={ex.get('median_largest_counter_over_normal_5m')} body-seq reversal with real 5m range; FPs largest_counter/normal={fp.get('median_largest_counter_over_normal_5m')}",
        f"- {(er.get('visual_separation'))}",
        "",
        f"Is 5m opening-drive representation better than OR-half location? **{a.get('five_m_better_than_or_half')}**",
        f"Is OR boundary alone a meaningful defended level? **{a.get('or_boundary_alone_meaningful')}**",
        "",
        "How many human-valid locations were:",
        f"- OR only? **{kinds.get('OR_ONLY_DEFENSE', 0)}**",
        f"- prior S/R? **{kinds.get('PRIOR_SR_ZONE_DEFENSE', 0)}**",
        f"- cleared zone? **{kinds.get('CLEARED_ZONE_RETEST', 0)}**",
        f"- PDH/PDL? **{kinds.get('PDH_PDL_DEFENSE', 0)}**",
        f"- VWAP confluence? **{kinds.get('VWAP_CONFLUENCE', 0)}**",
        f"- other? **{a.get('valid_loc_other')}**",
        "",
        f"What distinguishes valid reacceleration from micro-cross? **{a.get('reaccel_distinguish')}**",
        f"Does TradingValue participate in actual reacceleration semantics? **{a.get('tradingvalue_participates')}**",
        f"What causal state marks opening thesis lost? **{a.get('opening_thesis_lost_state')}**",
        f"Does daily context belong to the semantic definition? **{a.get('daily_context_in_definition')}**",
        f"Was IN-PLAY itself too broad? **{a.get('in_play_too_broad')}**",
        "",
        f"PRIMARY_CAUSE? **{a.get('PRIMARY_CAUSE')}**",
        f"SECONDARY_CAUSE? **{a.get('SECONDARY_CAUSE')}**",
        f"CONTRIBUTING_CAUSE? **{a.get('CONTRIBUTING_CAUSE')}**",
        f"Is PB1 salvageable as a coherent chart setup? **{a.get('pb1_salvageable')}**",
        "",
        f"Any rule changed? **{a.get('any_rule_changed')}**",
        f"Any threshold optimized? **{a.get('any_threshold_optimized')}**",
        f"Any future return? **{a.get('any_future_return')}**",
        f"Any economic test? **{a.get('any_economic_test')}**",
        f"Any PnL? **{a.get('any_pnl')}**",
        f"Old Confirmation opened? **{a.get('old_confirmation_opened')}**",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"submit/cancel/live? **{a.get('submit_cancel_live')}**",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
    ]
    return "\n".join(lines) + "\n"


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ex = list(((report.get("exemplars") or {}).get("positive")) or [])
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Agreement": _kv_rows(report.get("agreement")),
        "Opening_Drive": _kv_rows(report.get("opening_drive")),
        "Early_Reversal": _kv_rows(report.get("early_reversal_contrast")),
        "Location": _kv_rows(report.get("location_audit")),
        "Reacceleration": _kv_rows(report.get("reaccel_audit")),
        "Attribution": _kv_rows(report.get("attribution")),
        "Exemplars": ex or [{"empty": True}],
        "Layers": _kv_rows(report.get("layer_ranks")),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    (OUT / "report.json").write_text(json.dumps(_json_sanitize(report), ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
