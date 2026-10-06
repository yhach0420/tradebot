"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.bg_cont_vwap_largest_causal_deficiency_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Frozen_Parent",
    "Exit_Specs",
    "State_Machine",
    "Replay_Parity",
    "E0_Baseline",
    "E1_Persist2",
    "E2_MFE8_Grace",
    "D2_D3_D4",
    "D2_D3_Core",
    "Giveback",
    "Tail_Dependency",
    "Symbol_Contribution",
    "Occupancy",
    "Exit_Reason",
    "Event_Order",
    "Candidate_Comparison",
    "Safety",
)
ECON_KEYS = (
    "exit_id",
    "trade_n",
    "day_n",
    "symbol_n",
    "mean_x0_bps",
    "mean_x1_bps",
    "profit_factor",
    "hit_rate",
    "median_x0_bps",
    "daily_mean_bps",
    "daily_median_bps",
    "max_dd_daily_mean_bps",
    "mean_hold_min",
    "median_hold_min",
    "occupancy_skips",
    "same_symbol_skips",
    "immediate_failure_path_n",
    "warning_cancel_n_total",
    "mfe8_grace_used_n",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _econ(pack: dict[str, Any]) -> dict[str, Any]:
    return {k: pack.get(k) for k in ECON_KEYS}


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    e0 = dict(report.get("E0") or {})
    e1 = dict(report.get("E1") or {})
    e2 = dict(report.get("E2") or {})
    packs = [("E0", e0), ("E1", e1), ("E2", e2)]
    sym_rows = []
    for eid, pack in packs:
        for r in list(pack.get("symbols") or []):
            sym_rows.append({"exit_id": eid, **r})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "frozen_mechanism_hash": report.get("frozen_mechanism_hash"),
                "entry_unchanged": True,
                "promoted": False,
                "frozen_validation_opened": False,
                "design_evidence_not_certification": True,
            }
        ),
        "Frozen_Parent": _kv(dict(report.get("parent_spec") or {})),
        "Exit_Specs": list(report.get("exit_specs") or [{"empty": True}]),
        "State_Machine": [
            {"exit_id": "E0", "path": "OPEN → EXIT_PENDING on first VWAP loss → CLOSED"},
            {"exit_id": "E1", "path": "OPEN → VWAP_WARNING on first loss → OPEN if reclaim else EXIT_PENDING on 2nd loss"},
            {"exit_id": "E2", "path": "if MFE8 not activated: E0; if MFE8 activated: E1 warning logic"},
            {"slot": "occupied until EXIT fill", "states": "OPEN / VWAP_WARNING / EXIT_PENDING / CLOSED"},
        ],
        "Replay_Parity": _kv(dict(report.get("replay_parity") or {})),
        "E0_Baseline": [_econ(e0)],
        "E1_Persist2": [_econ(e1)],
        "E2_MFE8_Grace": [_econ(e2)],
        "D2_D3_D4": [
            {"exit_id": eid, "D2": (pack.get("block_mean_x0") or {}).get("D2"), "D3": (pack.get("block_mean_x0") or {}).get("D3"), "D4": (pack.get("block_mean_x0") or {}).get("D4")}
            for eid, pack in packs
        ],
        "D2_D3_Core": [{"exit_id": eid, **dict(pack.get("d2_d3") or {})} for eid, pack in packs],
        "Giveback": [{"exit_id": eid, **dict(pack.get("giveback") or {})} for eid, pack in packs],
        "Tail_Dependency": [{"exit_id": eid, **{k: v for k, v in dict(pack.get("tail") or {}).items() if k != "n" or True}} for eid, pack in packs],
        "Symbol_Contribution": sym_rows or [{"empty": True}],
        "Occupancy": [
            {
                "exit_id": eid,
                "occupancy_skips": pack.get("occupancy_skips"),
                "same_symbol_skips": pack.get("same_symbol_skips"),
                "mean_hold_min": pack.get("mean_hold_min"),
                "mean_slots_at_entry": pack.get("mean_slots_at_entry"),
                "CAP": 3,
            }
            for eid, pack in packs
        ],
        "Exit_Reason": [{"exit_id": eid, **dict(pack.get("exit_reasons") or {})} for eid, pack in packs],
        "Event_Order": [{"scope": "E0", "rule": r} for r in list(report.get("event_order_e0") or [])]
        + [{"scope": "E1_E2", "rule": r} for r in list(report.get("event_order_e1_e2") or [])],
        "Candidate_Comparison": list((report.get("comparisons") or {}).values()) or [{"empty": True}],
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0", "v27_bolted": False, "promoted": False}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        "# BG_CONT_VWAP_LARGEST_CAUSAL_DEFICIENCY_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        f"E0 replay parity exact? **{a.get('E0_replay_parity_exact')}** n=`{a.get('E0_trade_n')}` X0=`{a.get('E0_X0')}`",
        f"E1 X0/X1/PF: **{a.get('E1_X0')}** / **{a.get('E1_X1')}** / **{a.get('E1_PF')}**",
        f"E2 X0/X1/PF: **{a.get('E2_X0')}** / **{a.get('E2_X1')}** / **{a.get('E2_PF')}**",
        f"Reduced giveback most: **{a.get('Which_reduced_giveback_most')}** `{a.get('giveback_means')}`",
        f"Preserved large winners: **{a.get('Which_preserved_large_winners')}**",
        f"PTL delta E1/E2: `{a.get('Did_profitable_then_loss_decrease_E1')}` / `{a.get('Did_profitable_then_loss_decrease_E2')}`",
        f"Immediate-failure delta E1/E2: `{a.get('Did_true_immediate_failures_worsen_E1')}` / `{a.get('Did_true_immediate_failures_worsen_E2')}`",
        f"Occupancy-skip delta E1/E2: `{a.get('Did_delayed_exit_increase_capacity_blocking_E1')}` / `{a.get('Did_delayed_exit_increase_capacity_blocking_E2')}`",
        f"Symbols improved E1: `{a.get('Which_symbols_improved_E1')}` ({a.get('Was_improvement_broad_or_8136_dominated_E1')})",
        f"Symbols improved E2: `{a.get('Which_symbols_improved_E2')}` ({a.get('Was_improvement_broad_or_8136_dominated_E2')})",
        f"E1 D2/D3/D4 / D2+D3: `{a.get('E1_D2')}` / `{a.get('E1_D3')}` / `{a.get('E1_D4')}` / `{a.get('E1_D2_D3')}`",
        f"E2 D2/D3/D4 / D2+D3: `{a.get('E2_D2')}` / `{a.get('E2_D3')}` / `{a.get('E2_D4')}` / `{a.get('E2_D2_D3')}`",
        f"Any X1 > 0? **{a.get('Any_X1_gt_0')}**",
        f"Any candidate materially better than E0? **{a.get('Any_candidate_materially_better_than_E0')}**",
        f"E2 outperform E1 (one-bar noise after continuation)? **{a.get('Does_E2_outperform_E1')}**",
        f"Winning EXIT: **{a.get('Winning_EXIT')}**",
        "",
        f"Complete strategy promoted? **{a.get('Complete_strategy_promoted')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
        f"New paid data? **{a.get('New_paid_data')}**",
        f"Kabu50? **{a.get('Kabu50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
