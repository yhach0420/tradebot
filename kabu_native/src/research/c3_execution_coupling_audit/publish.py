"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.c3_execution_coupling_audit import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "c3_execution_coupling_audit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "OOF_Exact",
    "Funnel",
    "Replacement",
    "TargetFill",
    "FillTo600",
    "Adverse",
    "Deciles",
    "Lineage",
    "Occupancy",
    "FirstEntry",
    "EntryExit",
    "Questions",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    q = report.get("questions") or {}
    return "\n".join(
        [
            "# C3 OOF → EXACT EXECUTION COUPLING AUDIT V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE RESEARCH ONLY. Diagnosis only. No new model / feature / threshold / Clock / EXIT / C4.",
            "C3_TOP_EDGE_SUPPORTED_PORTFOLIO_FAIL maintained.",
            "C2_OOF_RANKING_EDGE_NOT_STABLE maintained. B_NO_ROBUST_IMPROVEMENT maintained.",
            "REENTRY_ORDINAL_IS_CLOCK_PROXY maintained.",
            "OOF_EXACT is OOF_EXACT_DIAGNOSTIC_ONLY. Not a deployable final model.",
            "Fill: Corrected Passive Fill SoT. WAIT_SEC=1.0. CLOCK=CURRENT IRREGULAR. TARGET V4 M4_PERSISTENT 600s.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A2_PARITY: {g.get('A2_PARITY')}",
            f"C3_FINAL_PARITY: {g.get('C3_FINAL_PARITY')}",
            "",
            f"OOF_EXACT_TRADES: {g.get('OOF_EXACT_TRADES')}",
            f"OOF_EXACT_PNL: {g.get('OOF_EXACT_PNL')}",
            f"OOF_EXACT_PF: {g.get('OOF_EXACT_PF')}",
            f"OOF_EXACT_DD: {g.get('OOF_EXACT_DD')}",
            f"OOF_EXACT_FIRST_N: {g.get('OOF_EXACT_FIRST_N')}",
            f"OOF_EXACT_FIRST_PNL: {g.get('OOF_EXACT_FIRST_PNL')}",
            f"OOF_EXACT_REENTRY_N: {g.get('OOF_EXACT_REENTRY_N')}",
            f"OOF_EXACT_REENTRY_PNL: {g.get('OOF_EXACT_REENTRY_PNL')}",
            "",
            f"CURRENT_TOP3_WOULD_FILL_RATE: {g.get('CURRENT_TOP3_WOULD_FILL_RATE')}",
            f"C3_OOF_TOP3_WOULD_FILL_RATE: {g.get('C3_OOF_TOP3_WOULD_FILL_RATE')}",
            f"C3_FINAL_TOP3_WOULD_FILL_RATE: {g.get('C3_FINAL_TOP3_WOULD_FILL_RATE')}",
            "",
            f"CURRENT_ONLY_TOP3_TARGET: {g.get('CURRENT_ONLY_TOP3_TARGET')}",
            f"CURRENT_ONLY_TOP3_FILL_RATE: {g.get('CURRENT_ONLY_TOP3_FILL_RATE')}",
            f"C3_ONLY_TOP3_TARGET: {g.get('C3_ONLY_TOP3_TARGET')}",
            f"C3_ONLY_TOP3_FILL_RATE: {g.get('C3_ONLY_TOP3_FILL_RATE')}",
            "",
            f"CURRENT_FILLABLE_TOP3_TARGET: {g.get('CURRENT_FILLABLE_TOP3_TARGET')}",
            f"C3_OOF_FILLABLE_TOP3_TARGET: {g.get('C3_OOF_FILLABLE_TOP3_TARGET')}",
            f"C3_FINAL_FILLABLE_TOP3_TARGET: {g.get('C3_FINAL_FILLABLE_TOP3_TARGET')}",
            "",
            f"CURRENT_FILL_TO_600: {g.get('CURRENT_FILL_TO_600')}",
            f"C3_OOF_FILL_TO_600: {g.get('C3_OOF_FILL_TO_600')}",
            f"C3_FINAL_FILL_TO_600: {g.get('C3_FINAL_FILL_TO_600')}",
            "",
            f"UPWARD_TARGET_LOW_FILLABILITY: {g.get('UPWARD_TARGET_LOW_FILLABILITY')}",
            f"PASSIVE_FILL_ADVERSE_SELECTION: {g.get('PASSIVE_FILL_ADVERSE_SELECTION')}",
            "",
            f"C3_LOST_VS_A2_N: {g.get('C3_LOST_VS_A2_N')}",
            f"PRIMARY_LOST_TRADE_CAUSE: {g.get('PRIMARY_LOST_TRADE_CAUSE')}",
            f"ENTRY_OR_EXIT_FAILURE: {g.get('ENTRY_OR_EXIT_FAILURE')}",
            f"FINAL_SPEC_INSTABILITY: {g.get('FINAL_SPEC_INSTABILITY')}",
            f"PRIMARY_CAUSE: {g.get('PRIMARY_CAUSE')}",
            f"RECOMMENDED_NEXT_RESEARCH: {g.get('RECOMMENDED_NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## QUESTIONS",
            "",
            f"Q1: {q.get('Q1')}",
            f"Q2: {q.get('Q2')}",
            f"Q3: {q.get('Q3')}",
            f"Q4: {q.get('Q4')}",
            f"Q5: {q.get('Q5')}",
            f"Q6: {q.get('Q6')}",
            f"Q7: {q.get('Q7')}",
            f"Q8: {q.get('Q8')}",
            "",
            "## STOP",
            "",
            "Audit complete. No C4. Runtime not changed. C3 not implemented. C14 unchanged.",
            "No execution-aware model / fill-probability model / new target / new EXIT / new Clock.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
