"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.full_causal_mechanism_discovery_v1 import ANALYSIS_ID
from research.full_causal_mechanism_discovery_v1.analyze import public_row
from research.full_causal_mechanism_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "mechanism_library",
    "exit_mapping",
    "candidate_set",
    "closed_references",
    "integrity",
    "canary",
    "coverage",
    "economics",
    "daily",
    "symbols",
    "causal_ex_top",
    "blocks",
    "base_qualification",
    "fold_selection",
    "fold_transfer",
    "selection_stability",
    "final_selection",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"CASE: {_fmt(d.get('CASE'))}",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT'))}",
        "",
        "Classification: DEV_CANDIDATE only if CASE A. Never validated / true OOS / certified / Formal.",
        "",
        "## Answers",
        "",
    ]
    for k, v in a.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    freeze = dict(report.get("candidate_set") or {})
    cands = list(freeze.get("candidates") or [])
    evaluated = [public_row(r) for r in list(report.get("evaluated") or [])]
    closed_e = [public_row(r) for r in list(report.get("closed_evaluated") or [])]
    daily_rows = []
    for r in list(report.get("evaluated") or []) + list(report.get("closed_evaluated") or []):
        for day, pnl in (r.get("daily") or {}).items():
            daily_rows.append({"STRATEGY_ID": r.get("STRATEGY_ID"), "date": day, "pnl": pnl})
    block_rows = []
    for r in list(report.get("evaluated") or []) + list(report.get("closed_evaluated") or []):
        for b in ((r.get("blocks") or {}).get("blocks") or []):
            block_rows.append({"STRATEGY_ID": r.get("STRATEGY_ID"), **b})
    return {
        "answers": kv_rows(report.get("answers") or {}),
        "objective_alignment": kv_rows(report.get("objective_alignment") or {}),
        "parent": kv_rows(report.get("parent") or {}),
        "mechanism_library": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "OPERATOR": r.get("OPERATOR"),
                "PREDICATES": r.get("PREDICATES"),
                "SELECTABLE": r.get("SELECTABLE"),
                "EXACT_CLOSED_ENTRY_IDENTITY": r.get("EXACT_CLOSED_ENTRY_IDENTITY"),
            }
            for r in cands
        ]
        or [{"empty": True}],
        "exit_mapping": kv_rows(report.get("exit_mapping") or {}),
        "candidate_set": [
            {
                "STRATEGY_ID": r.get("STRATEGY_ID"),
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "OPERATOR": r.get("OPERATOR"),
                "EXECUTION_ID": r.get("EXECUTION_ID"),
                "TECHNICAL_EXIT_ID": r.get("TECHNICAL_EXIT_ID"),
                "CAP": r.get("CAP"),
                "SELECTABLE": r.get("SELECTABLE"),
                "EXACT_CLOSED_ENTRY_IDENTITY": r.get("EXACT_CLOSED_ENTRY_IDENTITY"),
                "CLOSED_ENTRY_ID": r.get("CLOSED_ENTRY_ID"),
            }
            for r in cands
        ]
        or [{"empty": True}],
        "closed_references": [r for r in cands if r.get("EXACT_CLOSED_ENTRY_IDENTITY")] or [{"empty": True}],
        "integrity": list((report.get("integrity") or {}).get("gates") or [{"empty": True}]),
        "canary": kv_rows(report.get("canary") or {}),
        "coverage": [
            {
                "STRATEGY_ID": r.get("STRATEGY_ID"),
                "SELECTABLE": r.get("SELECTABLE"),
                "C1": r.get("C1"),
                "C2": r.get("C2"),
                "C3": r.get("C3"),
                "C4": r.get("C4"),
                "trade_n": r.get("trade_n"),
                "fill_day_n": r.get("fill_day_n"),
                "trades_per_day": r.get("trades_per_day"),
                "session_exit_unfilled_n": r.get("session_exit_unfilled_n"),
                "coverage_ok": r.get("coverage_ok"),
            }
            for r in evaluated + closed_e
        ]
        or [{"empty": True}],
        "economics": evaluated + closed_e or [{"empty": True}],
        "daily": daily_rows or [{"empty": True}],
        "symbols": [
            {"STRATEGY_ID": r.get("STRATEGY_ID"), "top_symbol": r.get("top_symbol"), "top_symbol_pnl": r.get("top_symbol_pnl")}
            for r in evaluated + closed_e
        ]
        or [{"empty": True}],
        "causal_ex_top": [
            {
                "STRATEGY_ID": r.get("STRATEGY_ID"),
                "top_symbol": r.get("top_symbol"),
                "CAUSAL_EX_TOP1_PNL": r.get("CAUSAL_EX_TOP1_PNL"),
                "METHOD": r.get("CAUSAL_EX_TOP1_METHOD"),
                "G6": (r.get("g_table") or {}).get("G6"),
            }
            for r in evaluated + closed_e
        ]
        or [{"empty": True}],
        "blocks": block_rows or [{"empty": True}],
        "base_qualification": [
            {
                "STRATEGY_ID": r.get("STRATEGY_ID"),
                "SELECTABLE": r.get("SELECTABLE"),
                "BASE_QUALIFIED": r.get("BASE_QUALIFIED"),
                "coverage_ok": r.get("coverage_ok"),
                "g_table": r.get("g_table"),
                "S1": (r.get("blocks") or {}).get("S1"),
                "S2": (r.get("blocks") or {}).get("S2"),
            }
            for r in evaluated + closed_e
        ]
        or [{"empty": True}],
        "fold_selection": list((report.get("selection_transfer") or {}).get("folds") or [{"empty": True}]),
        "fold_transfer": kv_rows(
            {
                k: v
                for k, v in dict(report.get("selection_transfer") or {}).items()
                if k != "folds"
            }
        ),
        "selection_stability": kv_rows(
            {
                "FULL_DEV_SELECTED_ID": report.get("FULL_DEV_SELECTED_ID"),
                "FULL_DEV_WINNER_TRAIN_TOP3_N": (report.get("selection_transfer") or {}).get("FULL_DEV_WINNER_TRAIN_TOP3_N"),
                "ST1": (report.get("selection_transfer") or {}).get("ST1"),
                "ST2": (report.get("selection_transfer") or {}).get("ST2"),
                "ST3": (report.get("selection_transfer") or {}).get("ST3"),
                "ST4": (report.get("selection_transfer") or {}).get("ST4"),
            }
        ),
        "final_selection": kv_rows(public_row(report.get("selected") or {}) if report.get("selected") else {"DEV_CANDIDATE": None}),
        "decision": kv_rows(report.get("decision") or {}),
        "safety": kv_rows(report.get("safety") or {}),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k not in {"_markdown", "evaluated", "closed_evaluated", "qualified", "closed_qualified"}}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))
