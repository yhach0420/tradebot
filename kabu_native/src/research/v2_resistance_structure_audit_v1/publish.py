"""Publish the resistance audit. Charts are drawn from recorded prices and embedded."""
from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

from research.event_time_impulse_complete_strategy_v2.identity import bind
from research.v2_resistance_structure_audit_v1.isolation import assert_isolated
from research.v2_resistance_structure_audit_v1.metrics import build
from research.v2_resistance_structure_audit_v1.scan import scan

OUT = Path("results/research/v2_resistance_structure_audit_v1")


def _rows(sheet, header, data):
    sheet.append(header)
    for row in data:
        sheet.append(row)


def _chart(item: dict, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.plot(item["t"], item["px"], color="black", lw=0.8, label="CurrentPrice")
    if item["bt"]:
        ax.plot(item["bt"], item["bid"], color="tab:blue", lw=0.6, label="Bid")
        ax.plot(item["bt"], item["ask"], color="tab:orange", lw=0.6, label="Ask")
    ax.axhline(item["entry_px"], color="green", ls="--", lw=0.7, label="ENTRY_ASK")
    if item.get("pre_break") is not None:
        ax.axhline(float(item["pre_break"]), color="red", ls=":", lw=0.8, label="PRE_BREAK_HIGH")
    if item.get("center") is not None and item.get("tick"):
        ax.axhspan(float(item["center"]) - float(item["tick"]), float(item["center"]) + float(item["tick"]), color="red", alpha=0.12, label="CURRENT_ZONE")
    ax.axvline(item["entry_t"], color="green", lw=0.7)
    ax.axvline(item["exit_t"], color="purple", lw=0.7)
    ax.set_title(f"{item['kind']} {item['symbol']} {item['date']} {item['reason']} pnl={item['pnl_yen']:.0f}")
    ax.legend(fontsize=7, loc="best")
    fig.tight_layout()
    fig.savefig(path, dpi=100)
    plt.close(fig)


def publish() -> dict:
    assert_isolated(OUT)
    bound = bind()
    raw = scan()
    report = build(raw)
    report["signal_sha256"] = bound["signal_sha256"]
    report["strategy_sha256"] = bound["strategy_sha256"]
    OUT.mkdir(parents=True, exist_ok=True)
    public = {k: v for k, v in report.items() if k not in ("rows", "charts")}
    (OUT / "report.json").write_text(json.dumps(public, indent=2, default=str), encoding="utf-8")
    lines = [
        f"# {report['verdict']}",
        "",
        f"NEXT: {report['next']}",
        "",
        f"parity={report['parity']} n={report['trade_n']} pnl={report['pnl_yen']} pf={report['pf']}",
        "",
        "ENTRY changed: false",
        "EXIT changed: false",
        "Complete Strategy changed: false",
        "submit/cancel/live: 0/0/0",
    ]
    (OUT / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    _rows(summary, ["key", "value"], [[k, public[k]] for k in ("verdict", "next", "parity", "trade_n", "pnl_yen", "pf", "quality_explanatory", "headroom_explanatory", "support_conversion_mechanism_supported", "v3_justified")])
    anatomy = wb.create_sheet("v2_trade_anatomy")
    _rows(anatomy, ["symbol", "date", "pnl_yen", "bps", "ratchets", "touch_class", "touch_n", "rejection_n", "headroom_ticks", "conversion", "reason"], [
        [r["symbol"], r["date"], r["pnl_yen"], r["bps"], r["ratchets"], r.get("touch_class"), r.get("touch_n"), r.get("rejection_n"), r.get("headroom_ticks"), r.get("conversion"), r["reason"]]
        for r in report["rows"]
    ])
    for name, payload in (
        ("resistance_touch_anatomy", report["classes"]),
        ("support_conversion", report["conversion"]),
        ("ratchet_groups", report["ratchet_groups"]),
    ):
        sheet = wb.create_sheet(name)
        _rows(sheet, ["group", "json"], [[key, json.dumps(value, default=str)] for key, value in payload.items()])
    head = wb.create_sheet("next_resistance_headroom")
    _rows(head, ["decile", "n", "mean_bps", "mfe", "ratchet_ge2"], [[d["decile"], d["trade_n"], d["mean_bps"], d["mfe_mean"], d["ratchet_ge2"]] for d in report["deciles"]])
    wb.create_sheet("rejection_anatomy").append(["see spearman in summary report.json"])
    wb.create_sheet("ratchet_structure").append([json.dumps(report["ratchet_ge2"], default=str)])
    wb.create_sheet("hold_groups").append(["hold groups are outcome labels; immediate sheet is <100ms"])
    imm = wb.create_sheet("immediate_failures")
    imm.append([json.dumps(report["immediate"], default=str)])
    wb.create_sheet("original_extension").append([json.dumps({"original": report["original18"], "extension": report["extension17"]}, default=str)])
    wb.create_sheet("chronological_folds").append([json.dumps(report["folds"], default=str)])
    visual = wb.create_sheet("visual_examples")
    _rows(visual, ["kind", "symbol", "date", "entry_t", "exit_t", "pnl", "ratchet_n", "pre_break", "structured", "touch_n", "rejection_n", "next", "headroom", "conversion", "explanation"], [
        [c["kind"], c["symbol"], c["date"], c["entry_t"], c["exit_t"], c["pnl_yen"], c["ratchets"], c.get("pre_break"), c.get("status"), c.get("touch_n"), c.get("rejection_n"), c.get("next_status"), c.get("headroom_ticks"), c.get("conversion"),
         f"{c['symbol']} {c['date']}: pre-break {c.get('pre_break')} had {c.get('touch_n')} prior zone entries and {c.get('rejection_n')} downward leaves. Next resistance {c.get('next_status')}. After entry the zone label is {c.get('conversion')}. Exit {c['reason']}."]
        for c in report["charts"]
    ])
    with TemporaryDirectory() as tmp:
        for i, item in enumerate(report["charts"]):
            png = Path(tmp) / f"chart_{i}.png"
            _chart(item, png)
            image = XLImage(str(png))
            image.anchor = f"A{20 + i * 22}"
            visual.add_image(image)
        wb.save(OUT / "audit.xlsx")
    return {"verdict": report["verdict"], "next": report["next"], "parity": report["parity"], "trade_n": report["trade_n"]}
