"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_implementation_correction_rca.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "OpeningBars",
    "Sequence",
    "Baseline",
    "S1_FP",
    "S1_FN",
    "MICRO",
    "Dimensions",
    "Constants",
    "FailedOpen",
    "CLEAR19",
    "Conditioned",
    "Timeframe",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def finite_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): finite_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [finite_sanitize(v) for v in out]
    return out


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    s1 = dict(report.get("s1") or {})
    micro = list(report.get("micro") or [])
    cons = dict(report.get("constants") or {})
    loc = dict(report.get("3382_location") or {})
    h7011 = dict(report.get("7011_20250523") or {})
    d4063 = dict(report.get("4063") or {})
    clear = dict(report.get("clear19") or {})
    tf = dict(report.get("timeframe") or {})
    return {
        "Corrected SHA unchanged?": bool(report.get("corrected_source_unchanged")) and bool(dec.get("corrected_sha_unchanged")),
        "RCA set n?": report.get("rca_set_n"),
        "Future outcome used?": False,
        "Prospective event consumed?": False,
        "S1 FP total?": s1.get("s1_fp_n"),
        "S1 FN total?": s1.get("s1_fn_n"),
        "Human-state composition of S1 FPs?": s1.get("human_state_composition_of_s1_fps"),
        "For each 3 MICRO FP?": micro,
        "Is NORMAL_OPENING_5M_RANGE part of the problem?": s1.get("NORMAL_OPENING_5M_RANGE_part_of_problem"),
        "Does current machine encode continued directional intent?": s1.get("continued_intent_encoded"),
        "Is path efficiency necessary to represent the frozen concept?": s1.get("path_efficiency_necessary"),
        "FAIL_EXTEND_MAX_BARS=6 explicitly supported by frozen spec?": cons.get("FAIL_EXTEND_MAX_BARS_explicitly_supported_by_frozen_spec"),
        "Is completed-5m leave explicitly required by frozen spec?": cons.get("completed_5m_leave_explicitly_required_by_frozen_spec"),
        "Is completed-5m retest/hold explicitly required by frozen spec?": cons.get("completed_5m_retest_hold_explicitly_required_by_frozen_spec"),
        "3382 OR_TOUCH_ONLY cause?": {"death": loc.get("death"), "visibility": loc.get("visibility"), "detail": loc.get("detail")},
        "7011 20250523 opening-horizon interpretation?": h7011,
        "4063 mismatch decomposition?": d4063,
        "CLEAR19 death distribution?": clear.get("death_distribution"),
        "How many CLEAR misses are caused by layers?": clear.get("miss_class_n"),
        "Which timeframe interpretation best matches the frozen intent?": tf.get("best_match_frozen_intent"),
        "Would that interpretation turn PB1 into a 1m strategy?": tf.get("would_that_turn_pb1_into_1m_strategy"),
        "Is S2 currently conflating location identification with retest confirmation?": tf.get("s2_currently_conflates_location_with_retest"),
        "Conditioned parity by layer?": report.get("conditioned_parity"),
        "Is the remaining problem?": dec.get("remaining_problem"),
        "combination?": dec.get("combination"),
        "Any rule changed?": False,
        "Any PnL?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    bind = dict(report.get("bind") or {})
    s1 = dict(report.get("s1") or {})
    dims = dict(s1.get("dimension_medians_by_group") or {})
    cons = dict(report.get("constants") or {})
    material = list(report.get("material_rows") or [])
    opening_bars: list[dict[str, Any]] = []
    sequences: list[dict[str, Any]] = []
    baselines: list[dict[str, Any]] = []
    for r in material:
        for i, b in enumerate(list(r.get("bars") or [])):
            opening_bars.append(
                {
                    "rca_id": r.get("rca_id"),
                    "symbol": r.get("symbol"),
                    "date": r.get("date"),
                    "bar_i": i + 1,
                    "human_opening_state": r.get("human_opening_state"),
                    "machine_opening_state": r.get("machine_opening_state"),
                    **dict(b),
                }
            )
        sequences.append(
            {
                "rca_id": r.get("rca_id"),
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "human_opening_state": r.get("human_opening_state"),
                "machine_S1": r.get("machine_S1"),
                "s1_fp": r.get("s1_fp"),
                "s1_fn": r.get("s1_fn"),
                **dict(r.get("sequence") or {}),
                **{f"clause_{k}": v for k, v in dict(r.get("true_clauses") or {}).items() if k != "passing_clause_ids"},
            }
        )
        baselines.append(
            {
                "rca_id": r.get("rca_id"),
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "human_opening_state": r.get("human_opening_state"),
                **dict(r.get("baseline") or {}),
            }
        )
    return {
        "Binding": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in bind.items()],
        "OpeningBars": opening_bars,
        "Sequence": sequences,
        "Baseline": baselines,
        "S1_FP": list(report.get("s1_fp_rows") or []),
        "S1_FN": list(report.get("s1_fn_rows") or []),
        "MICRO": list(report.get("micro") or []),
        "Dimensions": [{"group": k, **(v if isinstance(v, dict) else {"n": v})} for k, v in dims.items()],
        "Constants": list(cons.get("rows") or []),
        "FailedOpen": list((report.get("failed_open") or {}).get("rows") or []),
        "CLEAR19": list(report.get("clear19_rows") or []),
        "Conditioned": [{"layer": k, **(v if isinstance(v, dict) else {"v": v})} for k, v in dict(report.get("conditioned_parity") or {}).items() if isinstance(v, dict)],
        "Timeframe": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in (report.get("timeframe") or {}).items()],
        "Decision": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in (report.get("answers") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 implementation correction RCA", "", "Diagnosis only. Corrected SHA unchanged. DEVELOPMENT 88 only.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
