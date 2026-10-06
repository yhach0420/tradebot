"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_opening_range_continuation_face_valid_v2.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Parent_V1",
    "State_Machine",
    "In_Play",
    "Opening_Impulse",
    "OR15",
    "Break_Leave_Retest",
    "Triggers",
    "Reward_Space",
    "Matching_Roles",
    "Events",
    "Sample",
    "Human_Classification",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "day_rows", "events"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[chr(64 + i) if i < 27 else "A"].width = min(48, max(12, len(str(c)) + 2))


def build_markdown(report: dict[str, Any]) -> str:
    d = dict(report.get("decision") or {})
    h = dict(report.get("human") or {})
    ev = dict(report.get("events_summary") or {})
    ip = dict(report.get("in_play") or {})
    oi = dict(report.get("opening_impulse") or {})
    lines = [
        "# PB1_OPENING_RANGE_CONTINUATION_FACE_VALID_V2",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        "Semantic rebuild from blinded V1 face-validity evidence only.",
        "Thesis unchanged. False-break reversal is not this machine.",
        "No matching. No PnL. No Confirmation. No Frozen Validation. V1 not mutated.",
        "",
        f"V1 in-play share: {ip.get('v1_in_play_share')}",
        f"V1 too broad: {ip.get('v1_too_broad_because')}",
        f"V2 in-play share: {ip.get('in_play_share')}",
        f"V1-style failed-open setups removed: {report.get('v1_failed_open_reversals_removed')}",
        f"V1-style stale retests removed: {report.get('v1_stale_retest_removed')}",
        f"Retest freshness limit: {d.get('retest_freshness_minutes')} trading minutes",
        "",
        f"setup_n={ev.get('setup_n')} in_play_setup_n={ev.get('in_play_setup_n')} "
        f"bull/bear={ev.get('by_direction')}",
        f"opening states CLEAN/FAILED/RANGE/AMBIGUOUS="
        f"{oi.get('open_CLEAN_OPENING_IMPULSE')}/"
        f"{oi.get('open_FAILED_OPEN_REVERSAL')}/"
        f"{oi.get('open_RANGE_OPEN')}/"
        f"{oi.get('open_AMBIGUOUS_OPEN')}",
        f"triggers={ev.get('trigger_primary')}",
        "",
        f"Independent VERIFY sample_n={h.get('sample_n')} reviewed_n={h.get('reviewed_n')}",
        f"PATTERN CLEAR={h.get('CLEAR_OPENING_CONTINUATION')} "
        f"QUESTIONABLE={h.get('QUESTIONABLE')} NOT={h.get('NOT_OPENING_CONTINUATION')}",
        f"Pattern-face-valid n/share={h.get('pattern_face_valid_n')}/{h.get('pattern_face_valid_share')}",
        f"Trade-face-valid n/share={h.get('trade_face_valid_n')}/{h.get('trade_face_valid_share')}",
        f"RANGE_NOISE share={h.get('RANGE_NOISE_share')}",
        f"FAILED_OPEN_REVERSAL share={h.get('FAILED_OPEN_REVERSAL_share')}",
        f"STALE_RETEST share={h.get('STALE_RETEST_share')}",
        f"Most common semantic failure: {h.get('most_common_semantic_failure')}",
        "",
        f"future_outcome_used={d.get('future_outcome_used')} "
        f"outcome_based_threshold_choice={d.get('outcome_based_threshold_choice')} "
        f"pnl_test={d.get('pnl_test')}",
        f"Old Confirmation opened? {d.get('old_confirmation_opened')} "
        f"Frozen Validation opened? {d.get('frozen_validation_opened')}",
        f"submit/cancel/live? {d.get('submit_cancel_live')}",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ev = dict(report.get("events_summary") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Parent_V1": _kv_rows(report.get("parent_preserved")),
        "State_Machine": [{"text": report.get("state_machine")}],
        "In_Play": _kv_rows(report.get("in_play")),
        "Opening_Impulse": _kv_rows(report.get("opening_impulse")),
        "OR15": _kv_rows(report.get("or15_freeze")),
        "Break_Leave_Retest": _kv_rows(
            {
                "valid_break": report.get("valid_break"),
                "leave": report.get("leave"),
                "first_retest": report.get("first_retest"),
                "hold": report.get("hold"),
            }
        ),
        "Triggers": _kv_rows(ev.get("trigger_primary") or {}),
        "Reward_Space": _kv_rows(
            {
                "target_space": report.get("target_space"),
                "room_class": ev.get("room_class"),
            }
        ),
        "Matching_Roles": _kv_rows(report.get("matching_variable_roles") or {}),
        "Events": _kv_rows(ev),
        "Sample": list(report.get("sample") or [{"empty": True}]),
        "Human_Classification": list((report.get("human") or {}).get("rows") or [_kv_rows(report.get("human"))[0]]),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
