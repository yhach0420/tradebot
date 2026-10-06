"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.sector_symbol_driver_response_mapping_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Symbol_Sector_Map",
    "Driver_Catalog",
    "Source_Availability",
    "Market_Common_Mode",
    "Sector_Response",
    "Sector_LeadLag",
    "Sector_Asymmetry",
    "Sector_Regime",
    "Sector_TimeOfDay",
    "Stock_Response",
    "Stock_Residual_Response",
    "Stock_LeadLag",
    "Response_Vectors",
    "Response_Clusters",
    "Driver_Sector_Matrix",
    "Driver_Stock_Matrix",
    "Unknown_Driver",
    "True_Futures_Subtrack",
    "HM1_PostMap_Diagnostic",
    "D1_D4",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    cat = dict(report.get("catalog") or {})
    cells = list(report.get("sector_driver_cells") or [])
    stocks = list(report.get("stock_response") or [])
    syms = list(report.get("symbol_sector_map") or [])
    clus = list((report.get("clusters") or {}).get("clusters") or [])
    hm1 = dict(report.get("hm1_postmap") or {})
    live = dict(report.get("live_20260914") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "architecture": report.get("architecture"),
                "hm1_role": report.get("hm1_role"),
                "kabu_50": report.get("kabu_50_applied"),
                "frozen_validation_opened": False,
                "old_confirmation_used_to_design": report.get("old_confirmation_used_to_design"),
            }
        ),
        "Symbol_Sector_Map": [
            {k: r.get(k) for k in ("symbol", "name_ja", "sector", "industry_topix17", "first", "last", "n_days", "primary_driver", "cluster_id", "stock_specific_rule_justified")}
            for r in syms
        ]
        or [{"empty": True}],
        "Driver_Catalog": list(cat.get("drivers") or [{"empty": True}]),
        "Source_Availability": _kv(dict(cat.get("by_class") or {})) + _kv({"testable": cat.get("testable_ids"), "did_not_purchase": True}),
        "Market_Common_Mode": [r for r in cells if r.get("driver") == "MKT_MEDIAN_RET_1M"] or [{"note": "no_cells"}],
        "Sector_Response": list(report.get("sector_response") or [{"empty": True}]),
        "Sector_LeadLag": [{k: r.get(k) for k in ("sector", "driver", "lead_lag", "fwd_rho", "reverse_rho_sector_now_vs_driver", "usable_before_stock_move", "only_contemporaneous")} for r in cells] or [{"empty": True}],
        "Sector_Asymmetry": list(report.get("sector_response") or [{"empty": True}]),
        "Sector_Regime": list(report.get("sector_response") or [{"empty": True}]),
        "Sector_TimeOfDay": cells or [{"empty": True}],
        "Stock_Response": stocks or [{"empty": True}],
        "Stock_Residual_Response": [{k: r.get(k) for k in ("symbol", "driver", "rho_after_market", "rho_after_sector", "direct_beyond_sector", "status")} for r in stocks] or [{"empty": True}],
        "Stock_LeadLag": [{"note": "stock_primary_horizon_predeclared_+5m_labels_only", "primary_horizon_min": 5}],
        "Response_Vectors": [{k: r.get(k) for k in ("symbol", "sector", "primary_driver", "secondary_driver", "named", "cluster_id")} for r in syms] or [{"empty": True}],
        "Response_Clusters": clus or [{"empty": True}],
        "Driver_Sector_Matrix": cells or [{"empty": True}],
        "Driver_Stock_Matrix": stocks or [{"empty": True}],
        "Unknown_Driver": [r for r in syms if str(r.get("primary_driver") or "").startswith("UNKNOWN")] or [{"none": True}],
        "True_Futures_Subtrack": _kv(dict(report.get("true_futures_subtrack") or {})),
        "HM1_PostMap_Diagnostic": _kv(hm1),
        "D1_D4": [{k: r.get(k) for k in ("sector", "driver", "status", "block_rho")} for r in cells] or [{"empty": True}],
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    clus = list(a.get("clusters") or [])
    lines = [
        "# SECTOR_SYMBOL_DRIVER_RESPONSE_MAPPING_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted parent: `HISTORICAL_FUTURES_SOURCE_NOT_AVAILABLE_V1` as a source-availability / global-HM1 test, not as a mandate to wait on futures before mapping other families.",
        "Architecture: DRIVER → SECTOR → STOCK → local setup. Multiple playbooks. Not one global 105-name strategy.",
        "",
        f"Sectors: **{a.get('sector_n')}**  Stocks: **{a.get('stock_n')}**",
        f"Historically testable drivers: `{a.get('historically_testable_drivers')}`",
        f"PROXY_ONLY: `{a.get('PROXY_ONLY')}`",
        f"PROSPECTIVE_ONLY: `{a.get('PROSPECTIVE_ONLY')}`",
        f"UNAVAILABLE_NO_PURCHASE: `{a.get('UNAVAILABLE_NO_PURCHASE')}`",
        f"Stable sector-driver relations: **{a.get('stable_sector_driver_n')}**",
        f"Usable causal lead (before stock move): **{a.get('usable_lead_n')}**",
        f"Contemporaneous-only cells: **{a.get('relations_only_contemporaneous_n')}**",
        f"Stocks with extra direct sensitivity: **{a.get('stocks_additional_direct_sensitivity')}**",
        f"Stock-specific rules justified: **{a.get('stock_specific_rules_justified')}**",
        f"Fallback to sector/group: **{a.get('stocks_fallback_to_sector_or_group')}**",
        f"Stable clusters found? **{a.get('stable_response_clusters_found')}**",
        f"Electrical machinery one common driver? **{a.get('electrical_machinery_one_common_driver')}** split? **{a.get('electrical_machinery_splits')}**",
        f"True NK/TOPIX historical futures still unavailable? **{a.get('true_nk_topix_historical_futures_still_unavailable')}**",
        f"Prospective futures track intact? **{a.get('prospective_futures_track_intact')}**",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('old_confirmation_used_to_design')}**",
        f"Kabu 50? **{a.get('kabu_50_applied')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "Clusters:",
    ]
    for g in clus:
        lines.append(f"- `{g.get('id')}` n={g.get('n')} {g.get('interpretation')} symbols={g.get('symbols')}")
    lines += ["", "STOP. Do not return to generic global strategy mining.", ""]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
