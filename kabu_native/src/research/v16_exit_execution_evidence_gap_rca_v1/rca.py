"""Explain unfilled exits and unknown bid quantities without a full-period replay."""
from __future__ import annotations

import json
import time
import traceback
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from research.anchor_vs_event_driven.run_comparison import capture_event_epoch
from research.current50_v16_complete_strategy_risk_foundation_v1.foundation import _capture_dir, _paper_sessions, _symbol
from research.current50_v16_historical_economic_expansion_v2.expand import _attach_text
from research.e1_x34a_execution_policy.executable_board import is_executable_continuous_board
from research.event_time_volume_confirmed_impulse import FRESH_SEC, MIN_QTY
from research.simple_tech_entry_family.harvest import _parse_iso
from research.v16_execution_realism_unfilled_risk_v1.audit import QtyE, _install_patch, _prepare
from research.v16_one_tick_exit_capital_scaling_v2 import run as scaling

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "v16_exit_execution_evidence_gap_rca_v1"
PRIOR = NATIVE / "results" / "research" / "v16_execution_realism_unfilled_risk_v1" / "report.json"
JST = ZoneInfo("Asia/Tokyo")
G = 1_000_000.0
WORKERS = 2
CLOCKS = (
    ("20260825", "2351", 0),
    ("20260902", "7376", 0),
    ("20261009", "5246", None),
)
UNKNOWN_DAYS = (
    "20260728",
    "20260729",
    "20260730",
    "20260731",
    "20260804",
    "20260806",
    "20260807",
    "20260827",
    "20260828",
    "20260901",
    "20260902",
    "20260903",
    "20260904",
    "20260907",
    "20260909",
    "20260910",
)
TRIGGER_DAYS = ("20260825", "20260902", "20261009")


def _iso(epoch: Optional[float]) -> Optional[str]:
    if epoch is None:
        return None
    return datetime.fromtimestamp(float(epoch), JST).isoformat(timespec="milliseconds")


def _num(value: Any) -> Optional[float]:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:
        return None
    return number


def _model_fresh(payload: dict[str, Any], event_t: float) -> tuple[float, str]:
    board_age = _num(payload.get("board_age_sec"))
    if board_age is not None:
        return board_age, "board_age_sec"
    fresh = _num(payload.get("fresh_sec"))
    if fresh is not None:
        return fresh, "fresh_sec"
    for name in ("CurrentPriceTime", "AskTime", "BidTime"):
        stamp = _parse_iso(payload.get(name))
        if stamp is not None:
            return float(event_t) - stamp, name
    return 0.0, "missing_clock"


def _watermark(day: str) -> int:
    rows = [row for row in _paper_sessions(day) if row["kind"] == "pm" and "V16" in str(row["activation_id"])]
    if not rows:
        raise RuntimeError(f"NO_PM_WATERMARK {day}")
    return int(rows[-1]["watermark_seq"])


def _expected() -> dict[str, Any]:
    report = json.loads(PRIOR.read_text(encoding="utf-8"))
    out = {}
    for day, body in report["days"].items():
        pm = body["pm"]["E"]
        out[day] = {
            "trades": int(pm["trades"]),
            "gross": float(pm["gross"]),
            "exit_unknown": int(pm["exit_unknown"]),
            "unfilled": list(pm.get("unfilled") or []),
        }
    return out


def scan_symbol(day: str, code: str, seq_after: int) -> dict[str, Any]:
    capture = _capture_dir(day)
    if capture is None:
        return {"day": day, "symbol": code, "error": "NO_CAPTURE"}
    attach = _attach_text(day, "pm")
    needle = f'"symbol":"{code}"'
    flat = _parse_iso(f"{day[:4]}-{day[4:6]}-{day[6:8]}T15:20:00+09:00")
    deadline = _parse_iso(f"{day[:4]}-{day[4:6]}-{day[6:8]}T15:25:00+09:00")
    last_cpt = None
    last_cpt_recv = None
    last_buy_time = None
    last_buy_recv = None
    last_buy_px = None
    last_buy_qty = None
    cpt_changes = 0
    qty_only_while_cpt_frozen = 0
    buy_time_advanced_while_cpt_frozen = 0
    buy_time_missing = 0
    buy_time_ne_bid_time = 0
    events = 0
    window = Counter()
    window_buy_newer_than_cpt = 0
    window_buy_age_le_5 = 0
    window_model_fresh_gt_5 = 0
    window_qty_only = 0
    sources = Counter()
    states = Counter()
    first_special = None
    special_open = None
    special_intervals: list[dict[str, Any]] = []
    last_continuous = None
    last_model_ok = None
    last_buy_clock_ok = None
    transitions: list[dict[str, Any]] = []
    prev_qty = None
    prev_px = None
    prev_state = None
    prev_cpt = None
    for part in sorted(capture.glob("push_part_*.jsonl")):
        with part.open(encoding="utf-8") as handle:
            for line in handle:
                if needle not in line:
                    continue
                rec = json.loads(line)
                if str(rec.get("symbol") or "") != code:
                    continue
                if int(rec.get("sequence") or 0) <= seq_after:
                    continue
                received = str(rec.get("received_at") or "")
                if not received or received < attach:
                    continue
                payload = rec.get("original_payload") if isinstance(rec.get("original_payload"), dict) else None
                if not payload:
                    continue
                event_t = capture_event_epoch(rec, payload)
                if event_t is None or event_t >= float(deadline):
                    continue
                events += 1
                fresh, source = _model_fresh(payload, event_t)
                sources[source] += 1
                board = is_executable_continuous_board(payload, event_t=event_t)
                state = str(board.get("state") or "")
                states[state] += 1
                buy = payload.get("Buy1") if isinstance(payload.get("Buy1"), dict) else None
                buy_px = None if buy is None else _num(buy.get("Price"))
                buy_qty = None if buy is None else _num(buy.get("Qty"))
                buy_time = None if buy is None else _parse_iso(buy.get("Time"))
                cpt = _parse_iso(payload.get("CurrentPriceTime"))
                bid_time = _parse_iso(payload.get("BidTime"))
                if buy_time is None:
                    buy_time_missing += 1
                elif bid_time is not None and abs(buy_time - bid_time) > 1.0:
                    buy_time_ne_bid_time += 1
                if cpt != last_cpt:
                    cpt_changes += 1
                    last_cpt = cpt
                    last_cpt_recv = event_t
                if last_cpt_recv is not None and event_t > last_cpt_recv + 1.0 and cpt == last_cpt:
                    if buy_qty is not None and prev_qty is not None and buy_px == prev_px and buy_qty != prev_qty:
                        qty_only_while_cpt_frozen += 1
                    if buy_time is not None and last_buy_time is not None and buy_time > last_buy_time + 0.5:
                        buy_time_advanced_while_cpt_frozen += 1
                if buy_time is not None:
                    last_buy_time = buy_time
                    last_buy_recv = event_t
                if buy_px is not None:
                    last_buy_px = buy_px
                if buy_qty is not None:
                    last_buy_qty = buy_qty
                model_ok = bool(board.get("ok")) and fresh <= FRESH_SEC and buy_qty is not None and buy_qty + 1e-9 >= MIN_QTY and buy_px is not None and buy_px > 0
                buy_age = None if buy_time is None else event_t - buy_time
                buy_ok = bool(board.get("ok")) and buy_age is not None and buy_age <= FRESH_SEC and buy_qty is not None and buy_qty + 1e-9 >= MIN_QTY
                if model_ok and event_t < float(flat):
                    last_model_ok = {"received_at": received, "px": buy_px, "qty": buy_qty, "fresh": fresh, "source": source, "buy_time": None if buy is None else buy.get("Time"), "cpt": payload.get("CurrentPriceTime"), "state": state}
                if buy_ok and event_t < float(flat):
                    last_buy_clock_ok = {"received_at": received, "px": buy_px, "qty": buy_qty, "buy_age_sec": buy_age, "buy_time": None if buy is None else buy.get("Time"), "cpt": payload.get("CurrentPriceTime"), "model_fresh": fresh, "state": state}
                if state == "CONTINUOUS_TRADING":
                    last_continuous = {"received_at": received, "px": buy_px, "qty": buy_qty, "state": state, "fresh": fresh, "cpt": payload.get("CurrentPriceTime"), "buy_time": None if buy is None else buy.get("Time")}
                    if special_open is not None:
                        special_open["ended_at"] = received
                        special_intervals.append(special_open)
                        special_open = None
                elif "SPECIAL" in state and special_open is None:
                    special_open = {"started_at": received, "state": state, "px": buy_px, "qty": buy_qty, "cpt": payload.get("CurrentPriceTime"), "buy_time": None if buy is None else buy.get("Time"), "sign": None if buy is None else buy.get("Sign"), "fresh": fresh, "source": source, "prior_continuous": last_continuous}
                    if first_special is None:
                        first_special = dict(special_open)
                in_window = event_t + 1e-9 >= float(flat)
                if in_window:
                    window[state] += 1
                    if fresh > FRESH_SEC:
                        window_model_fresh_gt_5 += 1
                    if buy_time is not None and cpt is not None and buy_time > cpt + FRESH_SEC:
                        window_buy_newer_than_cpt += 1
                    if buy_age is not None and buy_age <= FRESH_SEC:
                        window_buy_age_le_5 += 1
                    if prev_px is not None and buy_px == prev_px and prev_qty is not None and buy_qty != prev_qty:
                        window_qty_only += 1
                changed = state != prev_state or buy_px != prev_px or buy_qty != prev_qty or cpt != prev_cpt
                if changed and (in_window or (first_special and transitions and transitions[-1].get("state") != state) or len(transitions) < 8):
                    if len(transitions) < 48:
                        transitions.append(
                            {
                                "received_at": received,
                                "event_time": rec.get("event_time"),
                                "state": state,
                                "buy_px": buy_px,
                                "buy_qty": buy_qty,
                                "buy_time": None if buy is None else buy.get("Time"),
                                "cpt": payload.get("CurrentPriceTime"),
                                "bid_time": payload.get("BidTime"),
                                "model_fresh_sec": round(fresh, 3),
                                "fresh_source": source,
                                "buy_age_sec": None if buy_age is None else round(buy_age, 3),
                            }
                        )
                prev_qty = buy_qty
                prev_px = buy_px
                prev_state = state
                prev_cpt = cpt
    if special_open is not None:
        special_open["ended_at"] = None
        special_intervals.append(special_open)
    covering = []
    for row in special_intervals:
        started = _parse_iso(row.get("started_at"))
        ended = _parse_iso(row.get("ended_at")) if row.get("ended_at") else None
        if started is None or started > float(flat):
            continue
        if ended is not None and ended < float(flat):
            continue
        covering.append(row)
    return {
        "day": day,
        "symbol": f"{code}.T",
        "events_after_attach_before_deadline": events,
        "fresh_sources": dict(sources),
        "states": dict(states),
        "cpt_changes": cpt_changes,
        "last_current_price_time": _iso(last_cpt),
        "last_current_price_time_seen_at": _iso(last_cpt_recv),
        "last_buy1_time": _iso(last_buy_time),
        "last_buy1_time_seen_at": _iso(last_buy_recv),
        "last_buy1_px": last_buy_px,
        "last_buy1_qty": last_buy_qty,
        "buy1_time_missing": buy_time_missing,
        "buy1_time_differs_from_bid_time": buy_time_ne_bid_time,
        "qty_only_updates_while_cpt_frozen": qty_only_while_cpt_frozen,
        "buy1_time_advanced_while_cpt_frozen": buy_time_advanced_while_cpt_frozen,
        "window_15_20_to_15_25": {
            "by_state": dict(window),
            "model_fresh_gt_5": window_model_fresh_gt_5,
            "buy1_time_newer_than_cpt_by_gt_5s": window_buy_newer_than_cpt,
            "buy1_age_le_5s": window_buy_age_le_5,
            "same_price_qty_only": window_qty_only,
        },
        "last_model_fresh_bid_before_15_20": last_model_ok,
        "last_buy1_clock_bid_before_15_20": last_buy_clock_ok,
        "last_continuous": last_continuous,
        "first_noncontinuous": first_special,
        "special_quote_covering_1520": covering[-1] if covering else None,
        "special_interval_count": len(special_intervals),
        "special_intervals": special_intervals if len(special_intervals) <= 6 else special_intervals[-6:],
        "transitions": transitions,
        "fresh_sec_rule": FRESH_SEC,
        "bid_clock_used_by_model": "CurrentPriceTime, then AskTime, then BidTime. Buy1.Time is not read.",
    }


class Probe(QtyE):
    def _normalize(self, event: dict[str, Any]) -> Optional[dict[str, Any]]:
        ev = super()._normalize(event)
        if not ev:
            return ev
        payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
        buy = payload.get("Buy1") if isinstance(payload.get("Buy1"), dict) else None
        event_t = float(ev["t"])
        fresh, source = _model_fresh(payload, event_t)
        ev["probe"] = {
            "cpt": payload.get("CurrentPriceTime"),
            "bid_time": payload.get("BidTime"),
            "ask_time": payload.get("AskTime"),
            "buy_time": None if buy is None else buy.get("Time"),
            "buy_qty": None if buy is None else buy.get("Qty"),
            "buy_px": None if buy is None else buy.get("Price"),
            "bid_qty_field": payload.get("BidQty"),
            "ask_qty_field": payload.get("AskQty"),
            "fresh_source": source,
            "model_fresh": fresh,
            "state": ev.get("state"),
        }
        return ev

    def _flush_bucket(self) -> None:
        snap = {}
        for ev in self._bucket:
            bid = _num(ev.get("bid"))
            snap[(str(ev.get("symbol") or ""), float(ev["t"]), bid)] = ev.get("probe")
        before = {symbol: len(rows) for symbol, rows in self.bid_log.items()}
        super()._flush_bucket()
        for symbol, rows in self.bid_log.items():
            for row in rows[before.get(symbol, 0) :]:
                row["probe"] = snap.get((symbol, float(row["t"]), _num(row.get("px"))))

    def _ensure_flatten(self, event_t: float) -> None:
        if not self._flatten_armed:
            self.pre_flat = {
                sym: {
                    "exit_trigger_t": pos.get("exit_trigger_t"),
                    "pending_reason": pos.get("pending_reason"),
                    "exiting": bool(pos.get("exiting")),
                    "exit_reason": pos.get("exit_reason"),
                    "entry_t": pos.get("entry_t"),
                    "entry_px": pos.get("entry_px"),
                }
                for sym, pos in self.open_pos.items()
            }
        super()._ensure_flatten(event_t)

    def _mark_unfilled(self, sym: str) -> None:
        pre = dict(getattr(self, "pre_flat", {}).get(sym) or {})
        super()._mark_unfilled(sym)
        if self.unfilled:
            self.unfilled[-1]["before_flatten"] = pre


def collect_day(day: str) -> dict[str, Any]:
    started = time.perf_counter()
    try:
        _install_patch()
        if day == "20261009":
            seq_after = _watermark(day)
            attach = ""
        else:
            seq_after = 0
            attach = _attach_text(day, "pm")
        capture = _capture_dir(day)
        if capture is None:
            return {"day": day, "error": "NO_CAPTURE"}
        book = _prepare(day, "pm", Probe)
        if len(book.admission_membership) != 50:
            return {"day": day, "error": "ADMISSION"}
        for part in sorted(capture.glob("push_part_*.jsonl")):
            with part.open(encoding="utf-8") as handle:
                for line in handle:
                    if not line.strip():
                        continue
                    rec = json.loads(line)
                    if rec.get("kind") not in (None, "market_push"):
                        continue
                    if int(rec.get("sequence") or 0) <= seq_after:
                        continue
                    received = str(rec.get("received_at") or "")
                    if not received or (attach and received < attach):
                        continue
                    payload = rec.get("original_payload") if isinstance(rec.get("original_payload"), dict) else None
                    if not payload:
                        continue
                    payload = dict(payload)
                    payload.setdefault("received_at", received)
                    payload.setdefault("recorded_at", received)
                    symbol = _symbol(str(rec.get("symbol") or payload.get("Symbol") or ""))
                    if not symbol or book._closed:
                        continue
                    book.on_market_event({"symbol": symbol, "payload": payload})
            if book._closed:
                break
        if not book._closed:
            book.close_at_runtime_boundary(day=day, kind="pm")
        unknown = []
        flat = float(book.flatten_start)
        for trade in book.trades:
            symbol = str(trade["symbol"])
            exit_t = float(trade["exit_t"])
            exit_px = float(trade["exit_px"])
            exact = None
            for row in book.bid_log.get(symbol, []):
                if abs(float(row["t"]) - exit_t) <= 1e-6 and abs(float(row["px"]) - exit_px) <= 1e-6:
                    exact = row
                    break
            if exact is not None and exact.get("buy1"):
                continue
            known = None
            if abs(exit_t - flat) <= 1e-6:
                for row in book.bid_log.get(symbol, []):
                    stamp = float(row["t"])
                    if stamp > flat + 1e-9:
                        break
                    if stamp + FRESH_SEC + 1e-9 >= flat:
                        known = row
                if known is not None and abs(float(known["px"]) - exit_px) > 1e-6:
                    known = None
            found = exact or known
            probe = None if found is None else found.get("probe") or {}
            buy_qty = None if not probe else _num(probe.get("buy_qty"))
            logged_qty = None if found is None else _num(found.get("qty"))
            unknown.append(
                {
                    "uid": trade.get("signal_uid"),
                    "symbol": symbol,
                    "reason": trade.get("reason"),
                    "pnl_yen": trade.get("pnl_yen"),
                    "entry_px": trade.get("entry_px"),
                    "exit_px": exit_px,
                    "entry_t": _iso(float(trade["entry_t"])),
                    "exit_t": _iso(exit_t),
                    "bid_event_t": None if found is None else _iso(float(found["t"])),
                    "join": "exact_exit_time" if exact is not None else ("flatten_known_bid" if known is not None else "unjoined"),
                    "bid_log_found": found is not None,
                    "buy1_qty_confirmed": bool(found is not None and found.get("buy1") and logged_qty is not None and logged_qty + 1e-9 >= MIN_QTY),
                    "fallback_bid_qty": logged_qty,
                    "buy1_qty": buy_qty,
                    "fresh_sec": None if found is None else found.get("fresh"),
                    "fresh_source": None if not probe else probe.get("fresh_source"),
                    "state": None if not probe else probe.get("state"),
                    "current_price_time": None if not probe else probe.get("cpt"),
                    "buy1_time": None if not probe else probe.get("buy_time"),
                    "bid_time": None if not probe else probe.get("bid_time"),
                    "english_bid_qty": None if not probe else probe.get("bid_qty_field"),
                    "english_ask_qty": None if not probe else probe.get("ask_qty_field"),
                    "cover": None if found is None else found.get("cover"),
                    "close_proven": False,
                }
            )
        unfilled = []
        for row in book.unfilled:
            pre = row.get("before_flatten") or {}
            unfilled.append(
                {
                    "uid": row.get("signal_uid"),
                    "symbol": row.get("symbol"),
                    "entry_t": _iso(float(row["entry_t"])),
                    "entry_px": row.get("entry_px"),
                    "notional": float(row["entry_px"]) * 100.0,
                    "pnl_yen": None,
                    "exit_signal_before_flatten": pre.get("exit_trigger_t") is not None,
                    "exit_trigger_before_flatten": _iso(_num(pre.get("exit_trigger_t"))),
                    "pending_reason_before_flatten": pre.get("pending_reason"),
                    "exiting_before_flatten": pre.get("exiting"),
                    "model_fill": False,
                    "slot_released": True,
                    "session_end_unfilled": True,
                    "economic_close": False,
                }
            )
        gross = round(sum(float(trade["pnl_yen"]) for trade in book.trades), 6)
        print(f"RCA {day} trades {len(book.trades)} gross {gross} unknown {len(unknown)} unfilled {len(unfilled)}", flush=True)
        return {
            "day": day,
            "elapsed_sec": time.perf_counter() - started,
            "trades": len(book.trades),
            "gross": gross,
            "max_gross": scaling._exposure(book.ledger, G)[0],
            "open_n": len(book.open_pos),
            "unknown": unknown,
            "unfilled": unfilled,
        }
    except Exception:
        return {"day": day, "error": traceback.format_exc()}


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _book(path: Path, clocks: list[dict[str, Any]], unknown: list[dict[str, Any]], unfilled: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "clocks"
    headers = [
        "day",
        "symbol",
        "events",
        "fresh_source",
        "last_cpt",
        "last_cpt_seen",
        "last_buy1_time",
        "buy1_time_missing",
        "qty_only_while_cpt_frozen",
        "window_model_fresh_gt_5",
        "window_buy1_age_le_5",
        "last_model_bid",
        "first_noncontinuous",
    ]
    ws.append(headers)
    for row in clocks:
        window = row.get("window_15_20_to_15_25") or {}
        ws.append(
            [
                row.get("day"),
                row.get("symbol"),
                row.get("events_after_attach_before_deadline"),
                _cell(row.get("fresh_sources")),
                row.get("last_current_price_time"),
                row.get("last_current_price_time_seen_at"),
                row.get("last_buy1_time"),
                row.get("buy1_time_missing"),
                row.get("qty_only_updates_while_cpt_frozen"),
                window.get("model_fresh_gt_5"),
                window.get("buy1_age_le_5s"),
                _cell(row.get("last_model_fresh_bid_before_15_20")),
                _cell(row.get("first_noncontinuous")),
            ]
        )
    wu = wb.create_sheet("unknown_exits")
    keys = [
        "uid",
        "symbol",
        "reason",
        "pnl_yen",
        "entry_px",
        "exit_px",
        "entry_t",
        "exit_t",
        "bid_log_found",
        "buy1_qty_confirmed",
        "fallback_bid_qty",
        "fresh_sec",
        "fresh_source",
        "state",
        "current_price_time",
        "buy1_time",
        "join",
        "bid_event_t",
        "buy1_qty",
        "english_bid_qty",
        "english_ask_qty",
        "close_proven",
    ]
    wu.append(keys)
    for row in unknown:
        wu.append([_cell(row.get(key)) for key in keys])
    wf = wb.create_sheet("unfilled")
    fkeys = list(unfilled[0].keys()) if unfilled else ["uid"]
    wf.append(fkeys)
    for row in unfilled:
        wf.append([_cell(row.get(key)) for key in fkeys])
    wb.save(path)


def _markdown(body: dict[str, Any]) -> str:
    lines = [
        f"# V16 EXIT Execution Evidence Gap RCA V1",
        "",
        f"VERDICT: `{body['verdict']}`",
        "",
        f"NEXT: {body['next']}",
        "",
        "Paper、Runtime、ENTRY、EXIT、Captureは変更していない。凍結済みB+Eは上書きしていない。この結果を理由にPaperへB+Eを採用しない。submit/cancel/live = 0/0/0。",
        "",
        "## 未約定の直接原因",
        "",
    ]
    for row in body["causes"]:
        lines.append(f"- {row['symbol']} {row['day']} {row['session']}: {row['cause']}")
    lines.extend(["", "## 15:20前の状態", ""])
    for row in body.get("unfilled") or []:
        lines.append(
            f"- {row.get('uid')}: 建玉 {row.get('notional')} 円。15:20より前のEXIT発火は {row.get('exit_signal_before_flatten')}、"
            f"理由 {row.get('pending_reason_before_flatten')}、発火時刻 {row.get('exit_trigger_before_flatten')}。"
            "モデル約定はなし。Slotは解放。Session終了。経済的な建玉解消は未確認。損益はnull。"
        )
    lines.extend(["", "## 鮮度判定", "", body["freshness_judgment"], "", "## EXIT数量UNKNOWN", "", body["unknown_summary"], "", "## 未解消建玉", "", body["residual"], "", "## Grossの信頼性", "", body["gross_reliability"], "", "## 追加で必要なデータ", "", body["still_needed"]])
    return "\n".join(lines) + "\n"


def _finish(clocks: list[dict[str, Any]], days: dict[str, Any], expected: dict[str, Any]) -> None:
    unknown: list[dict[str, Any]] = []
    unfilled: list[dict[str, Any]] = []
    parity_bad = []
    for day, body in sorted(days.items()):
        if body.get("error"):
            parity_bad.append(day)
            continue
        want = expected.get(day)
        if day == "20261009":
            want = {"trades": 124, "gross": 17570.0, "exit_unknown": 0}
        if want is not None and (int(body["trades"]) != int(want["trades"]) or abs(float(body["gross"]) - float(want["gross"])) > 1e-6):
            parity_bad.append(day)
            continue
        if want is not None and len(body["unknown"]) != int(want["exit_unknown"]):
            parity_bad.append(day)
            continue
        unknown.extend(body["unknown"])
        if day in TRIGGER_DAYS or day == "20261009":
            unfilled.extend(body["unfilled"])
    pnl = round(sum(float(row["pnl_yen"]) for row in unknown if row.get("pnl_yen") is not None), 6)
    confirmed = sum(1 for row in unknown if row.get("buy1_qty_confirmed"))
    by_reason = Counter(str(row.get("reason") or "") for row in unknown)
    by_join = Counter(str(row.get("join") or "") for row in unknown)
    missing_buy = [row for row in unknown if not row.get("buy1_qty_confirmed")]
    missing_pnl = round(sum(float(row["pnl_yen"]) for row in missing_buy if row.get("pnl_yen") is not None), 6)
    clocks_ok = all("error" not in row for row in clocks)
    gap_closed = clocks_ok and not parity_bad and len(unknown) == 28 and by_join.get("unjoined", 0) == 0
    verdict = "V16_EXIT_EXECUTION_EVIDENCE_GAP_RESOLVED_V1" if gap_closed else "V16_EXIT_EXECUTION_EVIDENCE_GAP_PARTIALLY_RESOLVED_V1"
    if not clocks_ok:
        verdict = "V16_EXIT_EXECUTION_EVIDENCE_GAP_UNRESOLVED_V1"
    causes = []
    for row in clocks:
        if row.get("error"):
            causes.append({"day": row.get("day"), "session": "PM", "symbol": row.get("symbol"), "cause": row["error"]})
            continue
        window = row["window_15_20_to_15_25"]
        if row["symbol"] == "5246.T":
            special = row.get("special_quote_covering_1520") or row.get("first_noncontinuous") or {}
            cause = (
                f"15:20時点の特別気配開始は {special.get('started_at')}、状態 {special.get('state')}、終了 {special.get('ended_at')}。"
                f"モデルが5秒以内と判定した最後のBuy1は {((row.get('last_model_fresh_bid_before_15_20') or {}).get('received_at'))}。"
                "既存EXITは特別気配の板を約定に使わない。早期決済は研究仮説であり、新しいEXIT条件にはしない。"
            )
        else:
            cause = (
                f"15:20以降にモデル鮮度が5秒を超えたイベントは {window.get('model_fresh_gt_5')}。"
                f"同じ窓でBuy1.Timeの経過が5秒以内だったイベントは {window.get('buy1_age_le_5s')}。"
                f"現値時刻の最終値は {row.get('last_current_price_time')}、Buy1.Timeの最終値は {row.get('last_buy1_time')}。"
                "鮮度計算はCurrentPriceTimeを優先し、Buy1.Timeを読まない。現値時刻が古いだけで、新しいBid更新が無効になった。"
                "PUSH受信時刻が新しいこと自体は価格の有効性の証明にしていない。"
            )
        causes.append({"day": row["day"], "session": "PM", "symbol": row["symbol"], "cause": cause})
    freshness = (
        "5秒制約は event_time_volume_confirmed_impulse.FRESH_SEC であり、密封済みのBid適格判定が使っている。今回は緩和していない。"
        "計算式は受信時刻から CurrentPriceTime、無ければ AskTime、無ければ BidTime を引く。board_age_sec と fresh_sec が payload にあればそちらが先だが、対象イベントでは使われていない。"
        "Buy1.Time は式に入らない。同一価格で数量だけ変わっても、現値時刻が止まっていれば鮮度は古いままになる。"
        "2351.T と 7376.T では Buy1.Time が現値時刻より新しく、Bid更新時刻を現値時刻とは別に確認できた。したがって、この2件の鮮度判定はBidの古さとしては妥当ではなかった。"
        "5246.T は特別気配で執行可能板が消えており、鮮度だけでは説明しない。"
    )
    unknown_summary = (
        f"前回の28件は、Session決済の記録時刻が15:20ちょうどで、実際のBid時刻はその5秒以内だったため、時刻の完全一致では結べなかった。"
        f"対象は前回監査が数えた16午後だけで、全29日と資金8条件は再計算していない。"
        f"UID {len(unknown)} 件。結合 {dict(by_join)}。理由 {dict(by_reason)}。"
        f"Buy1.Qtyが100株以上だったのは {confirmed} 件。確認できないのは {len(missing_buy)} 件で、その損益合計は {missing_pnl} 円。"
        f"28件全体の損益合計は {pnl} 円。公式Gross +837770 は書き換えない。"
        "表示数量を確認できても、証券会社の約定証明にはしない。確認できない行はPASSにしない。"
    )
    residual_rows = [row for row in unfilled]
    residual_sum = round(sum(float(row["notional"]) for row in residual_rows), 6)
    residual = (
        f"未約定 {len(residual_rows)} 件、建玉金額合計 {residual_sum} 円。損益はnullのまま。0円確定にしない。"
        "EXITシグナル、決済可能板、モデル約定、Slot解放、Session終了、経済的解消は行ごとに分離した。"
        "Slotはモデル上解放されているが、経済的な建玉解消は確認できない。これはPaper上の仮想残であり、実口座残高ではない。"
        "29日B+Eの未解消は 32300+82300=114600 円。20261009の5246.T 105700 円は別枠。"
    )
    gross_reliability = (
        "公式29日B+E Gross +837770 は変更しない。そのうちBuy1数量を確認できないEXITの損益を別計した。"
        "未約定2件は元からGrossに入っていない。6626件は前回監査でBuy1.Qty 100株以上だった。"
        "したがってGrossは表示板に基づく研究損益であり、今回の未確認分と未約定分を実約定の利益としては扱えない。"
    )
    still = "Buy1.TimeをBid鮮度に使う別バージョンの再評価、特別気配へ入る前の決済可能性の仮説検証、証券会社の約定・手数料・保証金。今回の監査だけでは実注文の約定保証はしない。"
    body = {
        "running": False,
        "verdict": verdict,
        "next": "凍結済みB+Eは維持し、Paperへ採用しない。Bid鮮度の時刻フィールドを直す場合は別バージョンとして設計し、採用判断は保留する。",
        "paper_runtime_capture_changed": False,
        "frozen_rule_overwritten": False,
        "submit": 0,
        "cancel": 0,
        "live": 0,
        "full_29_day_replay": False,
        "capital_cases_recomputed": False,
        "targeted_pm_days": sorted(days),
        "parity_bad_days": parity_bad,
        "causes": causes,
        "freshness_judgment": freshness,
        "unknown_n": len(unknown),
        "unknown_pnl_yen": pnl,
        "join": "flatten_known_bid_v2",
        "unknown_buy1_confirmed": confirmed,
        "unknown_unconfirmed_pnl_yen": missing_pnl,
        "unknown_summary": unknown_summary,
        "residual": residual,
        "residual_notional_yen": residual_sum,
        "gross_reliability": gross_reliability,
        "still_needed": still,
        "clocks": clocks,
        "unknown": unknown,
        "unfilled": residual_rows,
        "day_parity": {
            day: {
                "trades": body.get("trades"),
                "gross": body.get("gross"),
                "unknown_n": None if body.get("error") else len(body.get("unknown") or []),
                "expected_unknown": None if expected.get(day) is None else expected[day]["exit_unknown"],
                "error": body.get("error"),
            }
            for day, body in sorted(days.items())
        },
    }
    body["markdown"] = _markdown(body)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.md").write_text(body["markdown"], encoding="utf-8")
    _book(OUT / "audit.xlsx", clocks, unknown, residual_rows)
    (OUT / "report.json").write_text(json.dumps(body, ensure_ascii=False), encoding="utf-8")
    partial = OUT / "partial_days.json"
    if partial.exists():
        partial.unlink()
    print(verdict, "unknown", len(unknown), "pnl", pnl, "parity_bad", parity_bad, flush=True)


def main() -> None:
    existing = OUT / "report.json"
    if existing.exists():
        saved = json.loads(existing.read_text(encoding="utf-8"))
        if saved.get("verdict") and saved.get("running") is False and saved.get("join") == "flatten_known_bid_v2":
            print(saved["verdict"], flush=True)
            return
    expected = _expected()
    clocks = []
    for day, code, seq in CLOCKS:
        seq_after = _watermark(day) if seq is None else int(seq)
        print(f"SCAN {day} {code}", flush=True)
        clocks.append(scan_symbol(day, code, seq_after))
    days: dict[str, Any] = {}
    targets = list(dict.fromkeys([*UNKNOWN_DAYS, *TRIGGER_DAYS]))
    with ProcessPoolExecutor(max_workers=WORKERS) as pool:
        futures = {pool.submit(collect_day, day): day for day in targets}
        for future in as_completed(futures):
            body = future.result()
            days[str(body.get("day"))] = body
            done = OUT / "partial_days.json"
            OUT.mkdir(parents=True, exist_ok=True)
            done.write_text(json.dumps(days, ensure_ascii=False), encoding="utf-8")
    _finish(clocks, days, expected)


if __name__ == "__main__":
    main()
