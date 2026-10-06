"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_c0_exit_ptl_guard import ANALYSIS_ID, ARCHITECTURE_ID
from research.am_entry_profit_improvement.publish import json_sanitize

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_c0_exit_ptl_guard"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "summary",
    "precommit",
    "fixed_trades",
    "e1_augment",
    "daily_pnl",
    "current_preservation",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    lines = [
        "# AM C0 EXIT ARCHITECTURE V1",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        f"EXIT_ARCHITECTURE: {ARCHITECTURE_ID}",
        "C0 ENTRY frozen. CURRENT keeps C14. PTL guard on C0 augment only.",
        "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "",
        "## REQUIRED FINAL OUTPUT",
        "",
        f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
        "",
        f"EXIT_ARCHITECTURE: {_fmt(g.get('EXIT_ARCHITECTURE'))}",
        "",
        f"EXIT_PRECOMMIT_SPEC_SHA256: {_fmt(g.get('EXIT_PRECOMMIT_SPEC_SHA256'))}",
        "",
        f"FIXED_TRADE_N: {_fmt(g.get('FIXED_TRADE_N'))}",
        "",
        f"FIXED_DELTA_NET: {_fmt(g.get('FIXED_DELTA_NET'))}",
        "",
        f"L2_BASE_N: {_fmt(g.get('L2_BASE_N'))}",
        "",
        f"PTL_TRIGGER_N: {_fmt(g.get('PTL_TRIGGER_N'))}",
        "",
        f"L2_PTL_TRIGGER_N: {_fmt(g.get('L2_PTL_TRIGGER_N'))}",
        "",
        f"L2_CONVERTED_TO_NONLOSS_N: {_fmt(g.get('L2_CONVERTED_TO_NONLOSS_N'))}",
        "",
        f"WINNER_EARLY_CUT_N: {_fmt(g.get('WINNER_EARLY_CUT_N'))}",
        "",
        f"LOSER_PNL_SAVED_BY_GUARD: {_fmt(g.get('LOSER_PNL_SAVED_BY_GUARD'))}",
        "",
        f"WINNER_PNL_LOST_BY_GUARD: {_fmt(g.get('WINNER_PNL_LOST_BY_GUARD'))}",
        "",
        f"E1_AUGMENT_TRADE_N: {_fmt(g.get('E1_AUGMENT_TRADE_N'))}",
        "",
        f"E1_AUGMENT_NET: {_fmt(g.get('E1_AUGMENT_NET'))}",
        "",
        f"E1_AUGMENT_PF: {_fmt(g.get('E1_AUGMENT_PF'))}",
        "",
        f"E1_OVERLAY_NET: {_fmt(g.get('E1_OVERLAY_NET'))}",
        "",
        f"E1_OVERLAY_PF: {_fmt(g.get('E1_OVERLAY_PF'))}",
        "",
        f"E1_OVERLAY_DD: {_fmt(g.get('E1_OVERLAY_DD'))}",
        "",
        f"PAIRED_POS_DAYS: {_fmt(g.get('PAIRED_POS_DAYS'))}",
        "",
        f"PAIRED_NEG_DAYS: {_fmt(g.get('PAIRED_NEG_DAYS'))}",
        "",
        f"PAIRED_ZERO_DAYS: {_fmt(g.get('PAIRED_ZERO_DAYS'))}",
        "",
        f"PAIRED_MEDIAN: {_fmt(g.get('PAIRED_MEDIAN'))}",
        "",
        f"EX_BEST: {_fmt(g.get('EX_BEST'))}",
        "",
        f"EX_TOP3: {_fmt(g.get('EX_TOP3'))}",
        "",
        f"DELTA_NET_VS_C0: {_fmt(g.get('DELTA_NET_VS_C0'))}",
        "",
        f"DELTA_PF_VS_C0: {_fmt(g.get('DELTA_PF_VS_C0'))}",
        "",
        f"DELTA_DD_VS_C0: {_fmt(g.get('DELTA_DD_VS_C0'))}",
        "",
        f"DELTA_PAIRED_MEDIAN_VS_C0: {_fmt(g.get('DELTA_PAIRED_MEDIAN_VS_C0'))}",
        "",
        f"CURRENT_PRESERVATION_PASS: {_tf(g.get('CURRENT_PRESERVATION_PASS'))}",
        "",
        f"E1_FULL_PASS: {_tf(g.get('E1_FULL_PASS'))}",
        "",
        f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
        "",
        f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
        "",
        f"VERDICT: {g.get('VERDICT')}",
        "",
        f"NEXT: {g.get('NEXT')}",
        "",
        "## STOP",
        "",
        "Original C0 remains FROZEN_FOR_FUTURE_OOS_ONLY with C14.",
        "This architecture is C0_EXIT_PTL_GUARD_120_600_V1. Runtime adoption forbidden.",
        "submit/cancel/live=0/0/0.",
        "",
    ]
    return "\n".join(lines)
