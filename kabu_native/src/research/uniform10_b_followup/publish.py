"""Write report.json / report.md / audit.xlsx only. Isolated from C2."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.uniform10_b_followup import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "uniform10_b_followup"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Rank_Audit",
    "Flow",
    "B0",
    "B1_Folds",
    "B1_Search",
    "B2",
    "Comparison",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                obj = float(obj)
            elif isinstance(obj, np.integer):
                return int(obj)
    except Exception:
        pass
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# UNIFORM10 B FOLLOW-UP",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE Exact Dual-Lane. Occupancy approximation was not used.",
            "C2 was not modified. C14 / Runtime / CLOCK / ENTRY / EXIT / Fill SoT unchanged.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "Per-feature thresholds were not started.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"B0_PARITY: {g.get('B0_PARITY')}",
            f"NONEXEC_ALL_RATE: {g.get('NONEXEC_ALL_RATE')}",
            f"NONEXEC_TOP1_RATE: {g.get('NONEXEC_TOP1_RATE')}",
            f"NONEXEC_TOP3_RATE: {g.get('NONEXEC_TOP3_RATE')}",
            f"NONEXEC_TOP5_RATE: {g.get('NONEXEC_TOP5_RATE')}",
            f"NONEXEC_TOP3_ENRICHMENT: {g.get('NONEXEC_TOP3_ENRICHMENT')}",
            f"PENDING_NONEXEC_RATE: {g.get('PENDING_NONEXEC_RATE')}",
            f"EXPIRED_FROM_NONEXEC_RATE: {g.get('EXPIRED_FROM_NONEXEC_RATE')}",
            f"FILL_FROM_NONEXEC_T0_N: {g.get('FILL_FROM_NONEXEC_T0_N')}",
            f"CURRENT_ENTRY_NONEXEC_RANK_BIAS: {g.get('CURRENT_ENTRY_NONEXEC_RANK_BIAS')}",
            f"B1_SELECTED_THRESHOLD: {g.get('B1_SELECTED_THRESHOLD')}",
            f"B1_EXACT_TRADES: {g.get('B1_EXACT_TRADES')}",
            f"B1_EXACT_PNL: {g.get('B1_EXACT_PNL')}",
            f"B1_EXACT_PF: {g.get('B1_EXACT_PF')}",
            f"B1_EXACT_DD: {g.get('B1_EXACT_DD')}",
            f"B1_MEDIAN_DAILY_PNL: {g.get('B1_MEDIAN_DAILY_PNL')}",
            f"B1_POSITIVE_DAY_RATE: {g.get('B1_POSITIVE_DAY_RATE')}",
            f"B1_ROBUST_IMPROVEMENT: {g.get('B1_ROBUST_IMPROVEMENT')}",
            f"B2_TESTED: {g.get('B2_TESTED')}",
            f"B2_PNL: {g.get('B2_PNL')}",
            f"B2_PF: {g.get('B2_PF')}",
            f"B2_DD: {g.get('B2_DD')}",
            f"B2_ROBUST_IMPROVEMENT: {g.get('B2_ROBUST_IMPROVEMENT')}",
            f"RECOMMENDED_B_VARIANT: {g.get('RECOMMENDED_B_VARIANT')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## STOP",
            "",
            "B results are not written to Runtime. C2 unchanged. Paper/OPVAL not started.",
            "",
        ]
    )
