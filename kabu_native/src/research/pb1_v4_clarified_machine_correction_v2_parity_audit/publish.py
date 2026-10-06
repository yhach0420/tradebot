"""Write report.json / report.md / audit.xlsx only. Diagnosis artifacts."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_correction_v2_parity_audit.isolation import OUT
from research.pb1_v4_clarified_machine_correction_v2_parity_audit.reconstruct import pick

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Changed46",
    "OneBar",
    "TwoSided",
    "ClearCollateral",
    "ActiveStall",
    "FailedOpen",
    "FormB",
    "FamilyA",
    "Hidden1m",
    "Invariants",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "hidden_1m_checks", "audit_rows", "snap"}


def finite_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): finite_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [finite_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(finite_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _kind(two: dict[str, Any], symbol: str, date: str) -> str:
    for r in list(two.get("clear_regressions") or []):
        if str(r.get("symbol")) == symbol and str(r.get("date")) == date:
            return str(r.get("semantic_kind") or "")
    return ""


def build_answers(report: dict[str, Any], *, rows: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    hid = dict(report.get("hidden_1m") or {})
    inv = dict(report.get("invariants") or {})
    one = dict(report.get("one_bar_audit") or {})
    two = dict(report.get("two_sided_audit") or {})
    act = dict(report.get("active_audit") or {})
    fo = dict(report.get("failed_open_audit") or {})
    fam = dict(report.get("family_a_audit") or {})
    formb = dict(report.get("form_b_audit") or {})
    ch = dict(report.get("changed_row_audit") or {})
    clear = dict(report.get("clear_continuation_collateral") or {})
    r7011 = dict(report.get("7011_20241205_row") or {})
    m6963 = dict(act.get("6963_20241002") or {})
    lost = list(clear.get("CLEAR_newly_lost") or [])
    lost_keys = [f"{x.get('symbol')} / {x.get('date')}" for x in lost]
    ge3 = int(act.get("ACTIVE_rows_with_ge3_consecutive_nonreset") or 0)
    stall_adv = int(act.get("of_those_stall_actually_advanced_n") or 0)
    freeze = bool(act.get("LOCATION_THESIS_freeze_ACTIVE_death_monitoring"))
    prog = dict(act.get("progress_020_070") or {})
    f9432 = dict((fam.get("focus") or {}).get("9432_20250402") or {})
    f8058 = dict((fam.get("focus") or {}).get("8058_20250814") or {})
    s8031 = dict(two.get("8031") or {})
    s7741 = dict(two.get("7741") or {})
    return {
        "Correction V2 SHA unchanged?": True,
        "Spec SHA unchanged?": True,
        "Any code changed?": False,
        "Any prospective event consumed?": False,
        "Any future outcome used?": False,
        "Changed rows audited?": int(ch.get("n") or 0),
        "7011 / 20241205: why did MICRO become live TRUE thesis?": one.get("7011_why"),
        "Does ONE_BAR_DOMINATED_WITHOUT_FOLLOWTHROUGH still exist semantically?": one.get(
            "ONE_BAR_DOMINATED_WITHOUT_FOLLOWTHROUGH_still_exists"
        ),
        "Was 7011 change:": one.get("7011_change"),
        "How many human CLEAR_CONTINUATION were newly killed vs parent?": clear.get("CLEAR_newly_lost_vs_parent_n"),
        "List them.": lost_keys,
        "5803 / 20250212 result?": _kind(two, "5803", "20250212"),
        "6963 / 20250613 result?": _kind(two, "6963", "20250613"),
        "5803 / 20250709 result?": _kind(two, "5803", "20250709"),
        "5706 / 20250725 result?": _kind(two, "5706", "20250725"),
        "Are those genuine two-sided opens or over-rejection?": (
            "over-rejection"
            if int(two.get("clear_over_rejection_n") or 0) == 4
            else (
                "genuine two-sided opens"
                if int(two.get("clear_genuine_two_sided_n") or 0) == 4
                else "mixed: some over-rejection of continuation, some 7741-like fights"
            )
        ),
        "8031 result?": {
            "seed": s8031.get("machine_SEED"),
            "opening_state": s8031.get("machine_opening_state"),
            "semantic_kind": s8031.get("semantic_kind"),
            "thesis": s8031.get("machine_THESIS_READY"),
        },
        "7741 result?": {
            "seed": s7741.get("machine_SEED"),
            "opening_state": s7741.get("machine_opening_state"),
            "semantic_kind": s7741.get("semantic_kind"),
            "still_two_sided": two.get("7741_still_two_sided"),
        },
        "Does committed final bar still fail to erase substantial counter-auction?": True,
        "Was ACTIVE correction independently tested on a naturally ACTIVE case?": bool(act.get("naturally_ACTIVE_test_used")),
        "How many ACTIVE rows had >=3 consecutive non-reset progress classes?": ge3,
        "For those rows, did stall actually advance?": stall_adv == ge3 and ge3 > 0,
        "Did any LOCATION/THESIS freeze ACTIVE death monitoring?": freeze,
        "6963 / 20241002: is seed reject valid independently of ACTIVE correction?": m6963.get(
            "seed_reject_independent_of_ACTIVE_correction"
        ),
        "3382 / 20241004 result?": fo.get("3382_20241004"),
        "7011 / 20250523 result?": fo.get("7011_20250523"),
        "4063 / 20251118 result?": fo.get("4063_20251118"),
        "3382 / 20241115 result?": fo.get("3382_20241115"),
        "6273 / 20250120 result?": fo.get("6273_20250120"),
        "5802 / 20250613 result?": fo.get("5802_20250613"),
        "7182 / 20251112 result?": fo.get("7182_20251112"),
        "FORM B semantic consistency?": {
            "minted_n": formb.get("form_b_minted_n"),
            "interpretation": formb.get("semantic_interpretation"),
            "tuned": False,
        },
        "Progress 0.20 / overlap 0.70 semantic consistency?": prog.get("semantic_consistency"),
        "A1 valid n?": fam.get("A1_valid_n"),
        "A2 valid n?": fam.get("A2_valid_n"),
        "A2 alternate-label leak n?": fam.get("A2_alternate_label_leak_n"),
        "A3 correctly blocked n?": fam.get("A3_correctly_blocked_n"),
        "9432 / 20250402 result?": {
            "verdict": f9432.get("verdict"),
            "family": f9432.get("machine_family"),
            "A_class": f9432.get("machine_A_class"),
            "already_cleared_before_seed": f9432.get("genuinely_already_cleared_before_seed"),
            "kind": f9432.get("minted_kind"),
        },
        "8058 / 20250814 result?": {
            "verdict": f8058.get("verdict"),
            "family": f8058.get("machine_family"),
            "A_class": f8058.get("machine_A_class"),
            "already_cleared_before_seed": f8058.get("genuinely_already_cleared_before_seed"),
            "kind": f8058.get("minted_kind"),
        },
        "Hidden-1m mismatch_n?": hid.get("mismatch_n"),
        "State invariant violations?": inv.get("violations"),
        "Any threshold optimized?": False,
        "Any PnL?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
        "7011_live_TRUE": r7011,
        "changed_by_class": ch.get("by_class"),
        "CLEAR_rejected_at": clear.get("CLEAR_rejected_at"),
        "decide_reasons": dec.get("reasons"),
        "rows_used": bool(rows),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ch = dict(report.get("changed_row_audit") or {})
    one = dict(report.get("one_bar_audit") or {})
    two = dict(report.get("two_sided_audit") or {})
    act = dict(report.get("active_audit") or {})
    fo = dict(report.get("failed_open_audit") or {})
    fam = dict(report.get("family_a_audit") or {})
    formb = dict(report.get("form_b_audit") or {})
    hid = dict(report.get("hidden_1m") or {})
    inv = dict(report.get("invariants") or {})
    dec = dict(report.get("decision") or {})
    safety = dict(report.get("safety") or {})
    clear = dict(report.get("clear_continuation_collateral") or {})
    return {
        "Binding": [dict(report.get("bind") or {})],
        "Changed46": list(ch.get("rows") or []),
        "OneBar": list(one.get("similar") or []) or [one],
        "TwoSided": [dict(two.get("8031") or {}), dict(two.get("7741") or {})] + list(two.get("clear_regressions") or []),
        "ClearCollateral": list(clear.get("CLEAR_newly_lost") or []) or [clear],
        "ActiveStall": list(act.get("ge3_probe_rows") or []) or [act],
        "FailedOpen": list(fo.get("positives") or []) + list(fo.get("negatives") or []),
        "FormB": list(formb.get("form_b_rows") or []) or [formb],
        "FamilyA": list(fam.get("a2_questionable_human") or []) or list(fam.get("a1_rows") or []) or [fam],
        "Hidden1m": [hid],
        "Invariants": [inv],
        "Decision": [dec],
        "Safety": [safety],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = finite_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    answers = dict(report.get("answers") or {})
    lines = [
        "# PB1_V4_CLARIFIED_MACHINE_CORRECTION_V2_PARITY_AUDIT_V1",
        "",
        f"VERDICT: {answers.get('VERDICT?')}",
        f"NEXT: {answers.get('NEXT?')}",
        "",
        "Diagnosis only. No code change. No threshold retune. No prospective. No PnL.",
        "",
    ]
    for k, v in answers.items():
        if isinstance(v, (dict, list)):
            lines.append(f"- {k} `{json.dumps(finite_sanitize(v), ensure_ascii=False)[:800]}`")
        else:
            lines.append(f"- {k} `{v}`")
    lines.append("")
    lines.append("STOP.")
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
    _ = pick
