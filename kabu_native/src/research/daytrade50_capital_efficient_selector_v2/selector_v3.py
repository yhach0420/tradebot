"""DAYTRADE50_SELECTOR_V3. One frozen correction: max of the two opportunity ranks.

The V2 selector is not modified. The 87-day RCA sample is not a confirmation set.
"""
from __future__ import annotations

import hashlib
import json
import statistics
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

from openpyxl import Workbook

from research.auto50_research_capture.mode import tick_yen
from research.daytrade50_capital_efficient_selector_v2.baselines import load_current50_symbols
from research.daytrade50_capital_efficient_selector_v2.features import (
    MARKETS,
    percentile_high_is_good,
    snapshot_days,
)
from research.daytrade50_capital_efficient_selector_v2.lineage import (
    download_daily_bars,
    jpx_sessions,
    previous_jpx_session,
    universe_symbols,
)
from research.daytrade50_capital_efficient_selector_v2.policy import (
    AWAIT_USER_G,
    BUDGET_INELIGIBLE_PRESELECT,
    SCORE_PRECISION,
    budget_share,
    one_lot_notional,
    preselect_status,
)
from research.daytrade50_capital_efficient_selector_v2.selector import (
    FAIL_CLOSED,
    RANK_FIELDS,
    RANK_NAMES,
    selection_sort_key,
)
from research.daytrade50_capital_efficient_selector_v2.target_blind_g1m import (
    PINNED_V16_SHA,
    REPORTS,
    RUNTIME,
    V16,
    _auto50,
    _book_metrics,
    _feature,
    _market_map,
    _series,
    _target_metrics,
    read_identity,
)
from research.daytrade50_capital_efficient_selector_v2.validate import _quantile
from small_paper.v1r_activation_binding import manifest_content_sha

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "results" / "research" / "daytrade50_selector_v3_long_horizon_correction"
SAVED = (
    ROOT
    / "results"
    / "research"
    / "daytrade50_capital_efficient_selector_v2"
    / "target_blind_validation_cap500k_lev2x_g1m"
    / "report.json"
)
RCA = (
    ROOT
    / "results"
    / "research"
    / "daytrade50_capital_efficient_selector_v2"
    / "range_deficit_mechanism_v1"
    / "report.json"
)

SELECTOR_ID = "DAYTRADE50_SELECTOR_V3"
PARENT_ID = "DAYTRADE50_SELECTOR_V2"
VALIDATED = "DAYTRADE50_SELECTOR_V3_BLIND_VALIDATED_V1"
MIXED = "DAYTRADE50_SELECTOR_V3_BLIND_MIXED_V1"
NOT_SUPPORTED = "DAYTRADE50_SELECTOR_V3_BLIND_NOT_SUPPORTED_V1"
NO_HOLDOUT = "DAYTRADE50_SELECTOR_V3_FROZEN_NO_INDEPENDENT_HISTORICAL_HOLDOUT"
MISMATCH = "CAPITAL_MODEL_IDENTITY_MISMATCH"


def spec_document() -> dict[str, Any]:
    """Canonical V3 spec. No sample result enters this document."""
    return {
        "selector_id": SELECTOR_ID,
        "parent_selector_id": PARENT_ID,
        "correction_count": 1,
        "correction": "OPPORTUNITY_5D_RANK and OPPORTUNITY_20D_RANK leave the primary median; OPPORTUNITY_ANY_HORIZON_RANK is their max",
        "feature_definitions": {
            "TV_5D_MEDIAN": "median of volume*close over the last 5 completed sessions",
            "TV_20D_MEDIAN": "median of volume*close over the last 20 completed sessions",
            "liquidity_feature": "min(TV_5D_MEDIAN, TV_20D_MEDIAN)",
            "RANGE_PCT_5D_MEDIAN": "median of (high-low)/close*100 over the last 5 completed sessions",
            "RANGE_PCT_20D_MEDIAN": "median of (high-low)/close*100 over the last 20 completed sessions",
            "execution_suitability": "-tick_yen(close)/close",
        },
        "rank_definitions": {
            "method": "percentile_high_is_good inside the capital-eligible cross-section",
            "ties": "average rank, symbol orders equal values",
            "LIQUIDITY_RANK": "percentile of liquidity_feature",
            "OPPORTUNITY_5D_RANK": "percentile of RANGE_PCT_5D_MEDIAN",
            "OPPORTUNITY_20D_RANK": "percentile of RANGE_PCT_20D_MEDIAN",
            "EXECUTION_SUITABILITY_RANK": "percentile of execution_suitability",
            "OPPORTUNITY_ANY_HORIZON_RANK": "max(OPPORTUNITY_5D_RANK, OPPORTUNITY_20D_RANK)",
        },
        "PRIMARY_QUALITY_SCORE_V3": "median(LIQUIDITY_RANK, OPPORTUNITY_ANY_HORIZON_RANK, EXECUTION_SUITABILITY_RANK)",
        "weights": None,
        "gap_in_score": False,
        "gap_threshold": None,
        "five_day_only": False,
        "twenty_day_deleted": False,
        "capital_policy": {
            "BASE_CAPITAL_YEN": 500000,
            "MAX_LEVERAGE_MULTIPLIER": "2",
            "SYSTEM_GROSS_EXPOSURE_BUDGET_YEN": 1000000,
            "POLICY_INSTANCE_SHA": "8ff7da05e5e5991050048dce24c14882c7ff14e1814429efb07a0e246589c598",
            "qty": 100,
            "max_concurrent": 5,
            "per_position_cap": None,
            "g_over_5": False,
            "eligibility": "ONE_LOT_NOTIONAL = prior_close * 100 <= G",
        },
        "tie_break": ["-round(PRIMARY_QUALITY_SCORE_V3, 6)", "BUDGET_SHARE ascending", "symbol ascending"],
        "top50_rule": "fewer than 50 eligible names publishes no symbols",
        "freshness_rule": "window end equals the last JPX session before the target; lag_sessions = 0; future_used = false",
        "target_missing_rule": "a missing target outcome does not refill or rebuild membership",
        "core10": False,
        "rca87_is_confirmation": False,
        "decision_rule": {
            "range_pair": "median of daily (V3 range median - V2 range median) > 0",
            "range_overall": "median of daily V3 range medians > median of daily V2 range medians",
            "top50_success": "every holdout day returns 50 names",
            "active": "V3 active rate >= V2 active rate",
            "dead": "V3 dead rate <= V2 dead rate",
            "eligibility": "V3 one-lot eligibility = 1",
            "consistent": "V3 win days > V2 win days",
            "major_deterioration": "V3 trading-value median < half of V2, or V3 tick suitability median worse than V2 by more than 0.001",
            "VALIDATED": "range_pair and range_overall and top50_success and active and dead and eligibility and consistent, and not major_deterioration",
            "NOT_SUPPORTED": "range_pair is false or range_overall is false or major_deterioration",
            "MIXED": "range improves without major deterioration, and at least one preservation or consistency condition fails",
            "auto50_required_for_validated": False,
        },
    }


def spec_sha() -> str:
    raw = json.dumps(spec_document(), sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def score_v3(rows: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Same four percentiles as V2. The primary median then uses the max opportunity rank."""
    scored: list[dict[str, Any]] = [dict(row) for row in rows]
    for field, name in zip(RANK_FIELDS, RANK_NAMES):
        pairs = [(str(row["symbol"]), float(row[field])) for row in scored]
        ranks = percentile_high_is_good(pairs)
        for row in scored:
            row[name] = ranks[str(row["symbol"])]
    for row in scored:
        any_horizon = max(float(row["OPPORTUNITY_5D_RANK"]), float(row["OPPORTUNITY_20D_RANK"]))
        row["OPPORTUNITY_ANY_HORIZON_RANK"] = any_horizon
        primary = float(statistics.median([
            float(row["LIQUIDITY_RANK"]),
            any_horizon,
            float(row["EXECUTION_SUITABILITY_RANK"]),
        ]))
        row["PRIMARY_QUALITY_SCORE_V3"] = primary
        row["PRIMARY_QUALITY_SCORE"] = primary
        row["score_rounded"] = round(primary, SCORE_PRECISION)
    return scored


def select_top50_v3(feature_rows: Sequence[Mapping[str, Any]], *, gross_budget_yen: Optional[float]) -> dict[str, Any]:
    """V2 eligibility and tie break. Only the primary score differs."""
    if gross_budget_yen is None:
        return {"selector_id": SELECTOR_ID, "status": AWAIT_USER_G, "symbols": [], "rows": []}
    eligible: list[dict[str, Any]] = []
    rejected = {"history": 0, "budget": 0, "other": 0}
    for raw in feature_rows:
        row = dict(raw)
        if not row.get("eligible"):
            rejected["history" if row.get("reject") == "INSUFFICIENT_HISTORY" else "other"] += 1
            continue
        status = preselect_status(float(row["prior_close"]), float(gross_budget_yen))
        row["preselect_status"] = status
        row["ONE_LOT_NOTIONAL"] = one_lot_notional(float(row["prior_close"]))
        row["BUDGET_SHARE"] = budget_share(float(row["prior_close"]), float(gross_budget_yen))
        if status == BUDGET_INELIGIBLE_PRESELECT:
            rejected["budget"] += 1
            continue
        eligible.append(row)
    ordered = sorted(score_v3(eligible), key=selection_sort_key)
    if len(ordered) < 50:
        return {
            "selector_id": SELECTOR_ID,
            "status": FAIL_CLOSED,
            "symbols": [],
            "rows": [],
            "eligible_n": len(ordered),
            "rejected": rejected,
        }
    chosen = ordered[:50]
    for index, row in enumerate(chosen, start=1):
        row["rank"] = index
        row["origin"] = SELECTOR_ID
    return {
        "selector_id": SELECTOR_ID,
        "status": "TOP50",
        "symbols": [str(row["symbol"]) for row in chosen],
        "rows": chosen,
        "eligible_n": len(ordered),
        "rejected": rejected,
    }


def _formula_self_check() -> None:
    """Locks the algebra. It does not read a market day."""
    rows = []
    for symbol, liquidity, range5, range20, suitability in (
        ("A", 3.0, 9.0, 1.0, -0.001),
        ("B", 2.0, 1.0, 9.0, -0.001),
        ("C", 1.0, 2.0, 2.0, -0.01),
    ):
        rows.append({
            "symbol": symbol,
            "eligible": True,
            "prior_close": 1000.0,
            "liquidity_feature": liquidity,
            "range_pct_5d_median": range5,
            "range_pct_20d_median": range20,
            "suitability_feature": suitability,
        })
    scored = {row["symbol"]: row for row in score_v3(rows)}
    for row in scored.values():
        if row["OPPORTUNITY_ANY_HORIZON_RANK"] != max(row["OPPORTUNITY_5D_RANK"], row["OPPORTUNITY_20D_RANK"]):
            raise RuntimeError("V3 opportunity rank is not the max")
        expected = statistics.median([
            row["LIQUIDITY_RANK"],
            row["OPPORTUNITY_ANY_HORIZON_RANK"],
            row["EXECUTION_SUITABILITY_RANK"],
        ])
        if row["PRIMARY_QUALITY_SCORE_V3"] != expected:
            raise RuntimeError("V3 primary score is not the three-rank median")
    if scored["A"]["OPPORTUNITY_ANY_HORIZON_RANK"] <= scored["A"]["OPPORTUNITY_20D_RANK"]:
        raise RuntimeError("a fresh 5D name did not keep its 5D rank")
    if scored["B"]["OPPORTUNITY_ANY_HORIZON_RANK"] <= scored["B"]["OPPORTUNITY_5D_RANK"]:
        raise RuntimeError("a persistent 20D name did not keep its 20D rank")


def classify_holdout(summary: Mapping[str, Any]) -> str:
    """The rule in spec_document. It is not edited after the holdout is read."""
    if summary["major_deterioration"] or not summary["range_pair"] or not summary["range_overall"]:
        return NOT_SUPPORTED
    preserved = (
        summary["top50_success"]
        and summary["active_ok"]
        and summary["dead_ok"]
        and summary["eligibility_ok"]
        and summary["consistent"]
    )
    if preserved:
        return VALIDATED
    return MIXED


def _median(values: Sequence[Optional[float]]) -> Optional[float]:
    usable = [float(value) for value in values if value is not None]
    if not usable:
        return None
    return float(statistics.median(usable))


def _mean(values: Sequence[Optional[float]]) -> Optional[float]:
    usable = [float(value) for value in values if value is not None]
    if not usable:
        return None
    return sum(usable) / len(usable)


def _level(rows: Sequence[Mapping[str, Any]], key: str) -> Optional[float]:
    return _median([row.get(key) for row in rows])


def _rate(rows: Sequence[Mapping[str, Any]], key: str) -> Optional[float]:
    return _mean([row.get(key) for row in rows])


def _used_targets() -> list[str]:
    days = list(snapshot_days(REPORTS))
    if SAVED.is_file():
        saved = json.loads(SAVED.read_text(encoding="utf-8"))
        days.extend(str(row.get("target_trading_date")) for row in saved.get("days") or [])
    if RCA.is_file():
        rca = json.loads(RCA.read_text(encoding="utf-8"))
        days.extend(str(row.get("target_trading_date")) for row in rca.get("days") or [])
    return sorted({day for day in days if day and day.isdigit()})


def _holdout_calendar(sessions: Sequence[str], used_targets: Sequence[str]) -> dict[str, Any]:
    """Dates only. No next-day range is read here."""
    used_bars: set[str] = set()
    for target in used_targets:
        used_bars.add(target)
        required = previous_jpx_session(sessions, target)
        if required is None or required not in sessions:
            continue
        index = sessions.index(required)
        used_bars.update(sessions[max(0, index - 19): index + 1])
    earliest_used = min(used_bars) if used_bars else ""
    holdout = []
    for index, day in enumerate(sessions):
        if index < 20 or day >= "20260519" or (earliest_used and day >= earliest_used):
            continue
        holdout.append(day)
    return {
        "used_target_count": len(set(used_targets)),
        "earliest_used_session": earliest_used,
        "latest_used_target": max(used_targets) if used_targets else "",
        "holdout_dates": holdout,
        "holdout_start": holdout[0] if holdout else None,
        "holdout_end": holdout[-1] if holdout else None,
        "holdout_day_count": len(holdout),
        "prior_use": {
            "v2_design_targets": "feature artifacts and V2 selection targets begin 20260519",
            "v2_validation_and_rca": "saved target dates and their 20-session input windows are excluded",
            "v3_design": "the formula was fixed from the RCA mechanism, not from a pre-20260519 sample",
            "selector_tuning_artifacts_before_20260519": "none in results/reports or results/daily",
            "other_strategy_studies": "PB1 and entry studies are not this universe selector and do not admit these dates",
        },
    }


def _next(verdict: str) -> str:
    if verdict == VALIDATED:
        return "FREEZE_V3_AND_PREPARE_HISTORICAL_COMPLETE_STRATEGY_PATH"
    if verdict in (MIXED, NOT_SUPPORTED):
        return "CLOSE_V3_WITHOUT_ADOPTION"
    if verdict == NO_HOLDOUT:
        return "DO_NOT_CLAIM_VALIDATION"
    return "STOP_BEFORE_CONFIRMATION"


def _confirm(holdout: Sequence[str], sessions: Sequence[str], gross: float) -> dict[str, Any]:
    start = datetime.strptime(sessions[0], "%Y%m%d").date()
    end = datetime.strptime(holdout[-1], "%Y%m%d").date() + timedelta(days=1)
    meta = _market_map(REPORTS)
    symbols = [f"{symbol}.T" for symbol in universe_symbols(REPORTS)]
    panel = download_daily_bars(symbols, start, end)
    series = {symbol: _series(rows) for symbol, rows in panel.items() if rows}
    days: list[dict[str, Any]] = []
    v2_rows: list[dict[str, Any]] = []
    v3_rows: list[dict[str, Any]] = []
    auto_rows: list[dict[str, Any]] = []
    current_rows: list[dict[str, Any]] = []
    future_conditioned = 0
    for target in holdout:
        required = previous_jpx_session(sessions, target)
        if required is None:
            continue
        features: list[dict[str, Any]] = []
        prior_for_auto: list[dict[str, Any]] = []
        target_by_symbol: dict[str, dict[str, float]] = {}
        notion_by_symbol: dict[str, float] = {}
        suit_by_symbol: dict[str, float] = {}
        for symbol, item in series.items():
            exchange, market = meta.get(symbol, ("1", "prime"))
            built = _feature(symbol, item, required, exchange, market)
            if built is not None and exchange == "1" and market in MARKETS:
                if built["SOURCE_BAR_DATE"] != required:
                    future_conditioned += 1
                features.append(built)
                notion_by_symbol[symbol] = one_lot_notional(float(built["prior_close"]))
                suit_by_symbol[symbol] = float(built["suitability_feature"])
                price = float(built["prior_close"])
                index = item["index"][required]
                prior_for_auto.append({
                    "symbol": symbol,
                    "close": price,
                    "range_pct": (item["high"][index] - item["low"][index]) / price * 100.0,
                    "trading_value": item["volume"][index] * price,
                    "tick_yen": tick_yen(price),
                })
            metrics = _target_metrics(item, target)
            if metrics is not None:
                target_by_symbol[symbol] = metrics
        v2 = select_top50_imported(features, gross)
        v3 = select_top50_v3(features, gross_budget_yen=gross)
        v2_book = _book_metrics(v2.get("symbols") or [], notion_by_symbol, suit_by_symbol, target_by_symbol, gross)
        v3_book = _book_metrics(v3.get("symbols") or [], notion_by_symbol, suit_by_symbol, target_by_symbol, gross)
        auto_symbols = _auto50(prior_for_auto)
        auto_book = _book_metrics(auto_symbols, notion_by_symbol, suit_by_symbol, target_by_symbol, gross)
        current_path = RUNTIME / f"same_day_am_frozen_universe_{target}.json"
        current_symbols = load_current50_symbols(current_path)[:50] if current_path.is_file() else []
        current_book = _book_metrics(current_symbols, notion_by_symbol, suit_by_symbol, target_by_symbol, gross) if current_symbols else None
        if v2.get("status") == "TOP50":
            v2_rows.append(v2_book)
        if v3.get("status") == "TOP50":
            v3_rows.append(v3_book)
        auto_rows.append(auto_book)
        if current_book is not None:
            current_rows.append(current_book)
        delta = None
        if v2_book.get("range_median") is not None and v3_book.get("range_median") is not None:
            delta = float(v3_book["range_median"]) - float(v2_book["range_median"])
        days.append({
            "target_trading_date": target,
            "required_previous_jpx_session": required,
            "v2_status": v2.get("status"),
            "v3_status": v3.get("status"),
            "v2_range_median": v2_book.get("range_median"),
            "v3_range_median": v3_book.get("range_median"),
            "v3_minus_v2": delta,
            "auto_range_median": auto_book.get("range_median"),
            "current_range_median": None if current_book is None else current_book.get("range_median"),
        })
    deltas = [float(row["v3_minus_v2"]) for row in days if row.get("v3_minus_v2") is not None]
    wins = sum(1 for value in deltas if value > 1e-9)
    losses = sum(1 for value in deltas if value < -1e-9)
    ties = len(deltas) - wins - losses
    v2_range = _level(v2_rows, "range_median")
    v3_range = _level(v3_rows, "range_median")
    v2_active = _rate(v2_rows, "active_rate")
    v3_active = _rate(v3_rows, "active_rate")
    v2_dead = _rate(v2_rows, "dead_rate")
    v3_dead = _rate(v3_rows, "dead_rate")
    v2_tv = _level(v2_rows, "trading_value_median")
    v3_tv = _level(v3_rows, "trading_value_median")
    v2_tick = _level(v2_rows, "tick_suitability_median")
    v3_tick = _level(v3_rows, "tick_suitability_median")
    v3_elig = _rate(v3_rows, "one_lot_eligibility_rate")
    pair_median = _quantile(deltas, 0.50)
    major = bool(
        v2_tv is not None and v3_tv is not None and v3_tv < 0.5 * v2_tv
    ) or bool(
        v2_tick is not None and v3_tick is not None and v3_tick < v2_tick - 0.001
    )
    summary = {
        "range_pair": pair_median is not None and pair_median > 0,
        "range_overall": v2_range is not None and v3_range is not None and v3_range > v2_range,
        "top50_success": len(days) > 0 and all(row.get("v3_status") == "TOP50" for row in days),
        "active_ok": v2_active is not None and v3_active is not None and v3_active + 1e-9 >= v2_active,
        "dead_ok": v2_dead is not None and v3_dead is not None and v3_dead <= v2_dead + 1e-9,
        "eligibility_ok": v3_elig is not None and v3_elig >= 1 - 1e-12,
        "consistent": wins > losses,
        "major_deterioration": major,
    }
    books = {
        "V2": _book_level(v2_rows),
        "V3": _book_level(v3_rows),
        "AUTO50_V1": _book_level(auto_rows),
        "CURRENT50": _book_level(current_rows),
    }
    return {
        "days": days,
        "books": books,
        "paired": {
            "days": len(deltas),
            "v3_win_days": wins,
            "v2_win_days": losses,
            "ties": ties,
            "median_delta": pair_median,
            "p25_delta": _quantile(deltas, 0.25),
            "p75_delta": _quantile(deltas, 0.75),
        },
        "decision_inputs": summary,
        "future_conditioned_selection": future_conditioned,
        "verdict": classify_holdout(summary),
    }


def select_top50_imported(features: Sequence[Mapping[str, Any]], gross: float) -> dict[str, Any]:
    from research.daytrade50_capital_efficient_selector_v2.selector import select_top50

    return select_top50(features, gross_budget_yen=gross)


def _book_level(rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    return {
        "days": len(rows),
        "range_median": _level(rows, "range_median"),
        "range_p25": _level(rows, "range_p25"),
        "range_p75": _level(rows, "range_p75"),
        "trading_value_median": _level(rows, "trading_value_median"),
        "active_rate": _rate(rows, "active_rate"),
        "dead_rate": _rate(rows, "dead_rate"),
        "tick_suitability_median": _level(rows, "tick_suitability_median"),
        "one_lot_eligibility_rate": _rate(rows, "one_lot_eligibility_rate"),
    }


def run() -> dict[str, Any]:
    _formula_self_check()
    frozen_sha = spec_sha()
    identity = read_identity()
    body = json.loads(V16.read_text(encoding="utf-8")) if V16.is_file() else {}
    v16_sha = manifest_content_sha(body) if body else ""
    base = {
        "identity": identity,
        "identity_parity": bool(identity.get("match")),
        "V3_SELECTOR_SPEC_SHA": frozen_sha,
        "exact_v3_score": spec_document()["PRIMARY_QUALITY_SCORE_V3"],
        "changed_items": ["PRIMARY score now uses max(OPPORTUNITY_5D_RANK, OPPORTUNITY_20D_RANK) instead of both ranks"],
        "unchanged_items": [
            "liquidity feature",
            "RANGE_PCT_5D_MEDIAN",
            "RANGE_PCT_20D_MEDIAN",
            "execution suitability",
            "capital eligibility",
            "tie break",
            "freshness",
            "target missing",
            "qty",
            "MAX_CONCURRENT",
        ],
        "rca87_reused_as_confirmation": False,
        "in_sample_label": "NOT_EXECUTED_NOT_A_CONFIRMATION",
        "score_changed_after_holdout": spec_sha() != frozen_sha,
        "capital_changed": False,
        "leverage_changed": False,
        "v16_changed": v16_sha != PINNED_V16_SHA,
        "v16_sha": v16_sha,
        "runtime_applied": False,
        "selector_adopted": False,
        "submit_cancel_live": "0/0/0",
    }
    if not identity["match"]:
        base["verdict"] = MISMATCH
        base["next"] = _next(MISMATCH)
        return base
    witness = download_daily_bars(["1301.T", "7203.T"], date(2026, 1, 1), date(2026, 5, 20))
    sessions = jpx_sessions({
        symbol: [(row["source_bar_date"], float(row["close"]), float(row["volume"])) for row in rows]
        for symbol, rows in witness.items()
    })
    provenance = _holdout_calendar(sessions, _used_targets())
    base["holdout_provenance"] = provenance
    base["untouched_historical_holdout_exists"] = bool(provenance["holdout_dates"])
    base["holdout_date_range"] = [provenance["holdout_start"], provenance["holdout_end"]]
    base["holdout_day_count"] = provenance["holdout_day_count"]
    if not provenance["holdout_dates"]:
        base["verdict"] = NO_HOLDOUT
        base["next"] = _next(NO_HOLDOUT)
        base["score_changed_after_holdout"] = spec_sha() != frozen_sha
        return base
    confirmation = _confirm(provenance["holdout_dates"], sessions, float(identity["SYSTEM_GROSS_EXPOSURE_BUDGET_YEN"]))
    base.update(confirmation)
    base["verdict"] = confirmation["verdict"]
    base["next"] = _next(confirmation["verdict"])
    base["score_changed_after_holdout"] = spec_sha() != frozen_sha
    return base


def _fmt(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.6g}"
    return str(value)


def _markdown(report: Mapping[str, Any]) -> str:
    books = report.get("books") or {}
    paired = report.get("paired") or {}
    provenance = report.get("holdout_provenance") or {}
    lines = [
        f"VERDICT: {report.get('verdict')}",
        "",
        f"NEXT: {report.get('next')}",
        "",
        f"V3_SELECTOR_SPEC_SHA: {report.get('V3_SELECTOR_SPEC_SHA')}",
        "",
        f"PRIMARY_QUALITY_SCORE_V3 = {report.get('exact_v3_score')}",
        "",
        "changed items: OPPORTUNITY_ANY_HORIZON_RANK = max(OPPORTUNITY_5D_RANK, OPPORTUNITY_20D_RANK). それ以外のprimary入力は変えていません。",
        "",
        f"RCA87 reused as confirmation: {report.get('rca87_reused_as_confirmation')}",
        f"untouched historical holdout exists: {report.get('untouched_historical_holdout_exists')}",
        f"holdout date range: {provenance.get('holdout_start')} .. {provenance.get('holdout_end')}",
        f"holdout day count: {report.get('holdout_day_count')}",
        f"earliest session already used by V2/RCA inputs or targets: {provenance.get('earliest_used_session')}",
        "",
    ]
    for name in ("V2", "V3", "AUTO50_V1", "CURRENT50"):
        row = books.get(name) or {}
        if not row:
            continue
        lines.append(
            f"{name}: days={row.get('days')} range_median={_fmt(row.get('range_median'))} "
            f"p25={_fmt(row.get('range_p25'))} p75={_fmt(row.get('range_p75'))} "
            f"tv={_fmt(row.get('trading_value_median'))} active={_fmt(row.get('active_rate'))} "
            f"dead={_fmt(row.get('dead_rate'))} tick={_fmt(row.get('tick_suitability_median'))} "
            f"eligibility={_fmt(row.get('one_lot_eligibility_rate'))}"
        )
    if paired:
        lines.extend([
            "",
            f"paired V3-V2 days={paired.get('days')} wins={paired.get('v3_win_days')} losses={paired.get('v2_win_days')} ties={paired.get('ties')}",
            f"delta median/p25/p75: {_fmt(paired.get('median_delta'))} / {_fmt(paired.get('p25_delta'))} / {_fmt(paired.get('p75_delta'))}",
        ])
    lines.extend([
        "",
        f"score changed after holdout: {report.get('score_changed_after_holdout')}",
        f"capital changed: {report.get('capital_changed')}",
        f"leverage changed: {report.get('leverage_changed')}",
        f"CURRENT V16 changed: {report.get('v16_changed')}",
        "runtime applied: false",
        "submit/cancel/live: 0/0/0",
        "",
    ])
    return "\n".join(lines)


def _xlsx(report: Mapping[str, Any], path: Path) -> None:
    book = Workbook()
    spec = book.active
    spec.title = "spec"
    spec.append(["field", "value"])
    for key in (
        "verdict", "next", "V3_SELECTOR_SPEC_SHA", "exact_v3_score",
        "rca87_reused_as_confirmation", "untouched_historical_holdout_exists",
        "holdout_day_count", "score_changed_after_holdout", "v16_changed",
    ):
        spec.append([key, report.get(key)])
    days = book.create_sheet("holdout_days")
    day_rows = report.get("days") or []
    if day_rows:
        days.append(list(day_rows[0].keys()))
        for row in day_rows:
            days.append([row[key] for key in day_rows[0].keys()])
    compare = book.create_sheet("compare")
    compare.append(["book", "days", "range_median", "range_p25", "range_p75", "trading_value_median", "active_rate", "dead_rate", "tick_suitability_median", "one_lot_eligibility_rate"])
    for name, row in (report.get("books") or {}).items():
        compare.append([
            name, row.get("days"), row.get("range_median"), row.get("range_p25"), row.get("range_p75"),
            row.get("trading_value_median"), row.get("active_rate"), row.get("dead_rate"),
            row.get("tick_suitability_median"), row.get("one_lot_eligibility_rate"),
        ])
    book.save(path)


def main() -> None:
    report = run()
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(_markdown(report), encoding="utf-8")
    _xlsx(report, OUT / "audit.xlsx")
    print(report["verdict"])
    print(report["next"])
    print(report["V3_SELECTOR_SPEC_SHA"])


if __name__ == "__main__":
    main()
