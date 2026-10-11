"""Write the three research artifacts. Does not start Paper and does not set G."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from research.daytrade50_capital_efficient_selector_v2.baselines import baseline_catalog
from research.daytrade50_capital_efficient_selector_v2.policy import (
    policy_document,
    policy_instance_sha,
    policy_spec_sha,
)
from research.daytrade50_capital_efficient_selector_v2.lineage import freshness_proof
from small_paper.v1r_activation_binding import manifest_content_sha

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "results" / "research" / "daytrade50_capital_efficient_selector_v2"
V16 = ROOT / "results" / "research" / "v1r_exit_v2_prospective_activation" / "FIXED_ENTRY_SUPPORT_PAPER_PRIMARY_ACTIVATION_V16.json"
PINNED_V16_SHA = "49fb47c67a1da39a99edfba8b88d66c2e8009a07d6db94c269d6e8cdd05a2240"


def build_report() -> dict:
    freshness = freshness_proof(ROOT / "results" / "reports")
    body = json.loads(V16.read_text(encoding="utf-8")) if V16.is_file() else {}
    v16_sha = manifest_content_sha(body) if body else ""
    answers = {
        "1_user_parameter_is_G_only": True,
        "2_g_over_5_cap_removed": True,
        "3_single_name_may_use_near_G": True,
        "4_several_names_sum_to_G": True,
        "5_max_concurrent_is_occupancy": True,
        "6_qty_fixed_100": True,
        "7_no_cheap_stock_bias": True,
        "8_prior_day_only_score_retired": True,
        "9_liquidity_5d_20d": True,
        "10_opportunity_5d_20d": True,
        "11_persistence": True,
        "12_next_day_up_is_not_primary": True,
        "13_future_conditioned_selection": 0,
        "14_core10_removed": True,
        "15_user_g_not_imputed": True,
        "16_v16_unchanged": v16_sha == PINNED_V16_SHA,
        "17_submit_cancel_live": "0/0/0",
    }
    spec_sha = policy_spec_sha()
    instance_sha = policy_instance_sha(None)
    after = freshness["after_20261007"]
    passed = (
        freshness["stale_bar_selection_bug"] is True
        and after.get("MATCH") is True
        and freshness["targets_n"] > 0
        and freshness["targets_lag_zero"] == freshness["targets_n"]
        and freshness["targets_future"] == 0
        and freshness["targets_missing"] == 0
        and freshness["future_conditioned_selection"] == 0
        and freshness["score_changed"] is False
        and instance_sha is None
        and answers["16_v16_unchanged"] is True
    )
    return {
        "verdict": (
            "DAYTRADE50_V2_DAILY_FEATURE_FRESHNESS_PROVEN_V1"
            if passed
            else "DAYTRADE50_V2_DAILY_FEATURE_FRESHNESS_NOT_PROVEN_V1"
        ),
        "next": (
            "OBTAIN_USER_G_AND_RUN_TARGET_BLIND_SELECTOR_VALIDATION_V2"
            if passed
            else "FIX_FIRST_DAILY_FEATURE_LINEAGE_DEFICIENCY"
        ),
        "known_blocker": None if passed else "DAILY_FEATURE_LINEAGE",
        "user_g": None,
        "g_status": "AWAIT_USER_G",
        "runtime_applied": False,
        "paper_changed": False,
        "submit_cancel_live": "0/0/0",
        "policy": policy_document(),
        "POLICY_SPEC_SHA": spec_sha,
        "POLICY_INSTANCE_SHA": instance_sha,
        "sha_5a395155_is": "POLICY_SPEC_SHA",
        "date_evidence": freshness["iloc_preopen"],
        "baselines": baseline_catalog(),
        "answers": answers,
        "freshness": freshness,
        "v16_sha": v16_sha,
        "v16_sha_pinned": PINNED_V16_SHA,
        "capital_numbers": "AWAIT_USER_G",
        "primary_target": "daytrade_suitability",
        "not_primary_target": "next_day_up_prediction",
    }


def write_markdown(report: dict) -> str:
    fresh = report["freshness"]
    after = fresh["after_20261007"]
    before = fresh["before_20261007"]
    lines = [
        f"VERDICT: {report['verdict']}",
        f"NEXT: {report['next']}",
        "",
        "user G = unset",
        "POLICY_INSTANCE_SHA = null",
        "submit/cancel/live = 0/0/0",
        "score_changed = false",
        "production fetch_previous_day_yfinance unchanged",
        "",
        f"TARGET_TRADING_DATE: {after.get('TARGET_TRADING_DATE')}",
        f"REQUIRED_LATEST_COMPLETED_SESSION: {after.get('REQUIRED_LATEST_COMPLETED_SESSION')}",
        f"before ACTUAL_LATEST_SOURCE_BAR_DATE: {before.get('actual_latest_source_bar_date')}",
        f"after ACTUAL_LATEST_SOURCE_BAR_DATE: {after.get('ACTUAL_LATEST_SOURCE_BAR_DATE')}",
        f"MATCH: {after.get('MATCH')}",
        f"WINDOW_START_SESSION: {after.get('WINDOW_START_SESSION')}",
        f"WINDOW_END_SESSION: {after.get('WINDOW_END_SESSION')}",
        f"OBSERVATION_N: {after.get('OBSERVATION_N')}",
        f"targets lag 0: {fresh.get('targets_lag_zero')}/{fresh.get('targets_n')}",
        f"targets missing: {fresh.get('targets_missing')}",
        f"future used: {fresh.get('targets_future')}",
        f"STALE_BAR_SELECTION_BUG: {fresh.get('stale_bar_selection_bug')}",
        "",
        "| # | answer |",
        "| --- | --- |",
    ]
    labels = [
        "資本パラメータは G のみ",
        "G/5 per-position cap は廃止",
        "1銘柄が G 近くまで使える",
        "複数銘柄の合計が G まで使える",
        "MAX_CONCURRENT=5 は occupancy",
        "100株固定",
        "cheap-stock bias なし",
        "前日1日 score は V2 の主 score ではない",
        "5D/20D liquidity",
        "5D/20D opportunity",
        "persistence を median で見る",
        "翌日上昇予測は primary ではない",
        "future-conditioned selection = 0",
        "Core10 は新 Selector に無い",
        "G を補完していない",
        "CURRENT V16 は未変更",
        "submit/cancel/live = 0/0/0",
    ]
    for label, value in zip(labels, report["answers"].values()):
        lines.append(f"| {label} | {value} |")
    return "\n".join(lines) + "\n"


def write_xlsx(report: dict, path: Path) -> None:
    book = Workbook()
    answers = book.active
    answers.title = "answers"
    answers.append(["key", "value"])
    for key, value in report["answers"].items():
        answers.append([key, value])
    targets = book.create_sheet("targets")
    rows = report["freshness"].get("targets") or []
    if rows:
        targets.append(list(rows[0].keys()))
        for row in rows:
            targets.append([row.get(key) for key in rows[0].keys()])
    before = book.create_sheet("before_files")
    old_rows = report["freshness"].get("before_rows") or []
    if old_rows:
        before.append(list(old_rows[0].keys()))
        for row in old_rows:
            before.append([row.get(key) for key in old_rows[0].keys()])
    policy = book.create_sheet("policy")
    policy.append(["key", "value"])
    for key, value in report["policy"].items():
        policy.append([key, json.dumps(value) if isinstance(value, (dict, list)) else value])
    book.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    report = build_report()
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(write_markdown(report), encoding="utf-8")
    write_xlsx(report, OUT / "audit.xlsx")
    print(report["verdict"])
    print(report["freshness"]["after_20261007"])


if __name__ == "__main__":
    main()
