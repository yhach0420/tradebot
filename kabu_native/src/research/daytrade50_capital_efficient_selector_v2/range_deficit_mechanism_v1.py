"""Why DAYTRADE50 V2 gives up next-day range. Diagnosis only.

The V2 score, capital model, and V16 activation stay unchanged.
No replacement selector is emitted.
"""
from __future__ import annotations

import json
import statistics
from datetime import date
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

from openpyxl import Workbook

from research.auto50_research_capture.mode import tick_yen
from research.daytrade50_capital_efficient_selector_v2.baselines import load_current50_symbols
from research.daytrade50_capital_efficient_selector_v2.features import (
    EXCLUDED_SELECTION_DATES,
    MARKETS,
    execution_suitability,
    selection_calendar,
    snapshot_days,
)
from research.daytrade50_capital_efficient_selector_v2.lineage import (
    download_daily_bars,
    jpx_sessions,
    previous_jpx_session,
    universe_symbols,
)
from research.daytrade50_capital_efficient_selector_v2.policy import (
    BUDGET_INELIGIBLE_PRESELECT,
    budget_share,
    one_lot_notional,
    preselect_status,
)
from research.daytrade50_capital_efficient_selector_v2.selector import (
    score_cross_section,
    select_top50,
    selection_sort_key,
)
from research.daytrade50_capital_efficient_selector_v2.target_blind_g1m import (
    PINNED_V16_SHA,
    REPORTS,
    RUNTIME,
    V16,
    _auto50,
    _feature,
    _market_map,
    _series,
    _target_metrics,
    _traded_session_dates,
    read_identity,
)
from research.daytrade50_capital_efficient_selector_v2.validate import _quantile
from small_paper.v1r_activation_binding import manifest_content_sha

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "results" / "research" / "daytrade50_capital_efficient_selector_v2" / "range_deficit_mechanism_v1"
SAVED = ROOT / "results" / "research" / "daytrade50_capital_efficient_selector_v2" / "target_blind_validation_cap500k_lev2x_g1m" / "report.json"

IDENTIFIED = "DAYTRADE50_V2_RANGE_DEFICIT_MECHANISM_IDENTIFIED_V1"
NOT_IDENTIFIED = "DAYTRADE50_V2_RANGE_DEFICIT_MECHANISM_NOT_IDENTIFIED_V1"
MISMATCH = "CAPITAL_MODEL_IDENTITY_MISMATCH"


def _median(values: Sequence[Optional[float]]) -> Optional[float]:
    usable = [float(value) for value in values if value is not None]
    if not usable:
        return None
    return float(statistics.median(usable))


def _mean(values: Sequence[Optional[float]]) -> Optional[float]:
    usable = [float(value) for value in values if value is not None]
    if not usable:
        return None
    return sum(usable) / len(usable)


def _dist(values: Sequence[Optional[float]]) -> dict[str, Any]:
    usable = [float(value) for value in values if value is not None]
    return {
        "n": len(usable),
        "min": min(usable) if usable else None,
        "p25": _quantile(usable, 0.25),
        "median": _quantile(usable, 0.50),
        "p75": _quantile(usable, 0.75),
        "max": max(usable) if usable else None,
    }


def _higher(left: Optional[float], right: Optional[float]) -> bool:
    if left is None or right is None:
        return False
    return left > right + 1e-9


def _paired(deltas: Sequence[float]) -> dict[str, Any]:
    ordered = sorted(float(value) for value in deltas)
    positive = sum(1 for value in ordered if value > 1e-9)
    negative = sum(1 for value in ordered if value < -1e-9)
    ties = len(ordered) - positive - negative
    cut = int(len(ordered) * 0.10)
    trimmed = ordered[cut:]
    median_delta = _quantile(ordered, 0.50)
    trimmed_median = _quantile(trimmed, 0.50)
    share = (negative / len(ordered)) if ordered else None
    broad = bool(
        ordered
        and median_delta is not None
        and trimmed_median is not None
        and share is not None
        and median_delta < 0
        and trimmed_median < 0
        and share >= 0.60
    )
    return {
        "days": len(ordered),
        "positive_days": positive,
        "negative_days": negative,
        "ties": ties,
        "negative_share": share,
        "median_delta": median_delta,
        "p25_delta": _quantile(ordered, 0.25),
        "p75_delta": _quantile(ordered, 0.75),
        "tail_days_removed": cut,
        "median_delta_without_worst_10pct": trimmed_median,
        "broad_across_days": broad,
    }


def _group_summary(rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    active = [row for row in rows if row.get("next_range") is not None]
    dead = [row for row in rows if row.get("metrics_joined") and row.get("next_trading_value") is None]
    joined = [row for row in rows if row.get("metrics_joined")]
    return {
        "name_days": len(rows),
        "range_5d": _dist([row.get("range_5d") for row in rows]),
        "range_20d": _dist([row.get("range_20d") for row in rows]),
        "tv_5d": _dist([row.get("tv_5d") for row in rows]),
        "tv_20d": _dist([row.get("tv_20d") for row in rows]),
        "suitability": _dist([row.get("suitability") for row in rows]),
        "price": _dist([row.get("price") for row in rows]),
        "one_lot": _dist([row.get("one_lot") for row in rows]),
        "budget_share": _dist([row.get("budget_share") for row in rows]),
        "range_gap": _dist([row.get("range_gap") for row in rows]),
        "gap_above_1pp_share": _mean([1.0 if float(row["range_gap"]) > 1.0 else 0.0 for row in rows]),
        "next_range": _dist([row.get("next_range") for row in active]),
        "next_trading_value": _dist([row.get("next_trading_value") for row in active]),
        "active_rate": (len(active) / len(joined)) if joined else None,
        "dead_rate": (len(dead) / len(joined)) if joined else None,
        "daily_next_range_median": _median(
            list({str(row["day"]): row.get("daily_next_range") for row in rows if row.get("daily_next_range") is not None}.values())
        ),
    }


def _bucket(value: float, edges: Sequence[tuple[str, float, float]]) -> str:
    for name, low, high in edges:
        if low <= value < high:
            return name
    return edges[-1][0]


PRICE_EDGES = (
    ("lt_500", 0.0, 500.0),
    ("500_1500", 500.0, 1500.0),
    ("1500_3000", 1500.0, 3000.0),
    ("3000_7000", 3000.0, 7000.0),
    ("ge_7000", 7000.0, 1e18),
)
SHARE_EDGES = (
    ("le_25", 0.0, 0.25),
    ("25_50", 0.25, 0.50),
    ("50_75", 0.50, 0.75),
    ("75_100", 0.75, 1.0000001),
    ("gt_100", 1.0000001, 1e18),
)
LIQ_EDGES = (
    ("lt_1e8", 0.0, 1e8),
    ("1e8_1e9", 1e8, 1e9),
    ("1e9_1e10", 1e9, 1e10),
    ("ge_1e10", 1e10, 1e18),
)
GAP_EDGES = (
    ("lt_m1", -1e18, -1.0),
    ("m1_0", -1.0, 0.0),
    ("0_1", 0.0, 1.0),
    ("1_3", 1.0, 3.0),
    ("gt_3", 3.0, 1e18),
)


def _bucket_table(rows: Sequence[Mapping[str, Any]], key: str, edges: Sequence[tuple[str, float, float]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[Mapping[str, Any]]] = {name: [] for name, _low, _high in edges}
    for row in rows:
        grouped[_bucket(float(row[key]), edges)].append(row)
    out = []
    for name, _low, _high in edges:
        chunk = grouped[name]
        active = [row for row in chunk if row.get("next_range") is not None]
        out.append(
            {
                "bucket": name,
                "name_days": len(chunk),
                "prior_range_5d_median": _median([row.get("range_5d") for row in chunk]),
                "next_range_median": _median([row.get("next_range") for row in active]),
            }
        )
    return out


def _band_summary(rows: Sequence[Mapping[str, float]]) -> dict[str, Any]:
    return {
        "n": len(rows),
        "liquidity_rank": _median([row["liquidity"] for row in rows]),
        "opportunity_5d_rank": _median([row["opp5"] for row in rows]),
        "opportunity_20d_rank": _median([row["opp20"] for row in rows]),
        "execution_rank": _median([row["execution"] for row in rows]),
        "opportunity_avg": _median([(row["opp5"] + row["opp20"]) / 2.0 for row in rows]),
    }


def classify(summary: Mapping[str, Any]) -> dict[str, Any]:
    """Thresholds fixed before the walk. They are not adjusted to the outcome."""
    bands = summary["bands"]
    top = bands["top50"]
    mid = bands["rank_51_100"]
    deep = bands["rank_101_200"]
    liquidity_execution = (
        _higher(top["liquidity_rank"], mid["liquidity_rank"])
        and _higher(mid["liquidity_rank"], deep["liquidity_rank"])
        and _higher(top["execution_rank"], mid["execution_rank"])
        and _higher(top["liquidity_rank"], top["opportunity_avg"])
        and _higher(top["execution_rank"], top["opportunity_avg"])
    )
    groups = summary["groups"]
    auto_only = groups["AUTO50_ONLY"]
    v2_only = groups["V2_ONLY"]
    smoothing = (
        _higher(auto_only["range_gap"]["median"], v2_only["range_gap"]["median"])
        and _higher(auto_only["gap_above_1pp_share"], (v2_only["gap_above_1pp_share"] or 0) + 0.10)
        and float(auto_only["range_gap"]["median"] or 0) > 0
        and _higher(auto_only["daily_next_range_median"], v2_only["daily_next_range_median"])
    )
    screen = summary["capital_screen"]
    fraction = screen.get("fraction_of_auto_gap_closed")
    in_top = screen.get("median_ineligible_inside_diagnostic_top50") or 0
    if fraction is None or fraction < 0.25 or in_top <= 3:
        capital = "NOT_SUPPORTED"
    elif fraction < 0.50:
        capital = "SECONDARY"
    else:
        capital = "SUPPORTED"
    execution = (
        _higher(auto_only["range_5d"]["median"], v2_only["range_5d"]["median"])
        and _higher(v2_only["suitability"]["median"], (auto_only["suitability"]["median"] or 0) + 0.0001)
        and _higher(auto_only["daily_next_range_median"], v2_only["daily_next_range_median"])
    )
    broad = bool(summary["paired_auto50"]["broad_across_days"])
    supported = []
    if liquidity_execution:
        supported.append("LIQUIDITY_EXECUTION_DOMINANCE")
    if smoothing:
        supported.append("LONG_HORIZON_SMOOTHING")
    if capital == "SUPPORTED":
        supported.append("CAPITAL_SCREEN_EFFECT")
    if execution:
        supported.append("EXECUTION_RANGE_TRADEOFF")
    if not supported:
        primary = "NOT_IDENTIFIED"
    elif len(supported) == 1:
        primary = supported[0]
    else:
        primary = "MULTIPLE"
    if primary == "NOT_IDENTIFIED":
        verdict = NOT_IDENTIFIED
        nxt = "PRESERVE_V2_MIXED_AND_DO_NOT_ADOPT"
    else:
        verdict = IDENTIFIED
        nxt = "DESIGN_ONE_PREDECLARED_SELECTOR_V3_CORRECTION"
    return {
        "LIQUIDITY_EXECUTION_DOMINANCE": "SUPPORTED" if liquidity_execution else "NOT_SUPPORTED",
        "LONG_HORIZON_SMOOTHING": "SUPPORTED" if smoothing else "NOT_SUPPORTED",
        "CAPITAL_SCREEN_EFFECT": capital,
        "EXECUTION_RANGE_TRADEOFF": "SUPPORTED" if execution else "NOT_SUPPORTED",
        "RANGE_DEFICIT_BROAD_ACROSS_DAYS": broad,
        "PRIMARY_RANGE_DEFICIT_MECHANISM": primary,
        "supported_mechanisms": supported,
        "verdict": verdict,
        "next": nxt,
    }


def _membership(row: Mapping[str, Any], group: str, day: str, metrics_joined: bool, target: Optional[Mapping[str, float]]) -> dict[str, Any]:
    nxt_range = None if target is None else float(target["range_pct"])
    nxt_tv = None if target is None else float(target["trading_value"])
    if nxt_tv is not None and nxt_tv <= 0:
        nxt_range = None
        nxt_tv = None
    return {
        "day": day,
        "group": group,
        "symbol": row["symbol"],
        "range_5d": float(row["range_pct_5d_median"]),
        "range_20d": float(row["range_pct_20d_median"]),
        "tv_5d": float(row["tv_5d_median"]),
        "tv_20d": float(row["tv_20d_median"]),
        "suitability": float(row["suitability_feature"]),
        "price": float(row["prior_close"]),
        "one_lot": float(row["ONE_LOT_NOTIONAL"]),
        "budget_share": float(row["BUDGET_SHARE"]),
        "range_gap": float(row["range_pct_5d_median"]) - float(row["range_pct_20d_median"]),
        "metrics_joined": metrics_joined,
        "next_range": nxt_range if metrics_joined else None,
        "next_trading_value": nxt_tv if metrics_joined else None,
    }


def run() -> dict[str, Any]:
    identity = read_identity()
    body = json.loads(V16.read_text(encoding="utf-8")) if V16.is_file() else {}
    v16_sha = manifest_content_sha(body) if body else ""
    saved = json.loads(SAVED.read_text(encoding="utf-8")) if SAVED.is_file() else {}
    if not identity["match"]:
        return {
            "verdict": MISMATCH,
            "next": "STOP_BEFORE_DECOMPOSITION",
            "identity": identity,
            "v16_sha": v16_sha,
            "v16_changed": v16_sha != PINNED_V16_SHA,
            "runtime_applied": False,
            "score_changed": False,
            "selector_adopted": False,
            "submit_cancel_live": "0/0/0",
        }
    gross = float(identity["SYSTEM_GROSS_EXPOSURE_BUDGET_YEN"])
    meta = _market_map(REPORTS)
    symbols = [f"{symbol}.T" for symbol in universe_symbols(REPORTS)]
    panel = download_daily_bars(symbols, date(2026, 4, 1), date(2026, 10, 8))
    witness = {symbol: panel.get(symbol) or [] for symbol in ("1301", "7203")}
    close_sessions = jpx_sessions(
        {
            symbol: [(row["source_bar_date"], float(row["close"]), float(row["volume"])) for row in rows]
            for symbol, rows in witness.items()
        }
    )
    feature_days = set(snapshot_days(REPORTS))
    traded = _traded_session_dates(date(2026, 4, 1), date(2026, 10, 9)) & feature_days
    sessions = sorted(set(close_sessions) | traded)
    close_unavailable = set(traded - set(close_sessions))
    series = {symbol: _series(rows) for symbol, rows in panel.items() if rows}
    calendar = [day for day in selection_calendar(REPORTS) if day in set(sessions)]
    bands = {"top50": [], "rank_51_100": [], "rank_101_200": []}
    membership: dict[str, list[dict[str, Any]]] = {"COMMON": [], "V2_ONLY": [], "AUTO50_ONLY": []}
    daily_rows: list[dict[str, Any]] = []
    screen_rows: list[dict[str, Any]] = []
    ineligible_in_top50: list[int] = []
    prebudget_daily_range: list[float] = []
    v2_daily_range: list[float] = []
    auto_daily_range: list[float] = []
    reproduction_mismatches = 0
    for target in calendar:
        required = previous_jpx_session(sessions, target)
        if required is None:
            continue
        features: list[dict[str, Any]] = []
        prior_for_auto: list[dict[str, Any]] = []
        target_by_symbol: dict[str, dict[str, float]] = {}
        for symbol, item in series.items():
            exchange, market = meta.get(symbol, ("1", "prime"))
            built = _feature(symbol, item, required, exchange, market)
            if built is not None and exchange == "1" and market in MARKETS:
                built["ONE_LOT_NOTIONAL"] = one_lot_notional(float(built["prior_close"]))
                built["BUDGET_SHARE"] = budget_share(float(built["prior_close"]), gross)
                features.append(built)
                price = float(built["prior_close"])
                span = (item["high"][item["index"][required]] - item["low"][item["index"][required]]) / price * 100.0
                prior_for_auto.append(
                    {
                        "symbol": symbol,
                        "close": price,
                        "range_pct": span,
                        "trading_value": item["volume"][item["index"][required]] * price,
                        "tick_yen": tick_yen(price),
                    }
                )
            metrics = _target_metrics(item, target)
            if metrics is not None:
                target_by_symbol[symbol] = metrics
        official = select_top50(features, gross_budget_yen=gross)
        eligible = [row for row in features if preselect_status(float(row["prior_close"]), gross) != BUDGET_INELIGIBLE_PRESELECT]
        rejected = [row for row in features if preselect_status(float(row["prior_close"]), gross) == BUDGET_INELIGIBLE_PRESELECT]
        ordered = sorted(score_cross_section([dict(row) for row in eligible]), key=selection_sort_key)
        official_symbols = list(official.get("symbols") or [])
        if [str(row["symbol"]) for row in ordered[:50]] != official_symbols:
            reproduction_mismatches += 1
        for row in ordered[:50]:
            bands["top50"].append(_ranks(row))
        for row in ordered[50:100]:
            bands["rank_51_100"].append(_ranks(row))
        for row in ordered[100:200]:
            bands["rank_101_200"].append(_ranks(row))
        pre_ordered = sorted(score_cross_section([dict(row) for row in features]), key=selection_sort_key)
        pre_rank = {str(row["symbol"]): index for index, row in enumerate(pre_ordered, start=1)}
        metrics_joined = target not in close_unavailable
        in_top = 0
        pre_ranges: list[float] = []
        for row in pre_ordered[:50]:
            target_metrics = target_by_symbol.get(str(row["symbol"])) if metrics_joined else None
            if target_metrics is not None and float(target_metrics["trading_value"]) > 0:
                pre_ranges.append(float(target_metrics["range_pct"]))
            if float(row["ONE_LOT_NOTIONAL"]) > gross:
                in_top += 1
        ineligible_in_top50.append(in_top)
        if pre_ranges:
            prebudget_daily_range.append(float(statistics.median(pre_ranges)))
        for row in rejected:
            target_metrics = target_by_symbol.get(str(row["symbol"])) if metrics_joined else None
            nxt = None if target_metrics is None or float(target_metrics["trading_value"]) <= 0 else float(target_metrics["range_pct"])
            screen_rows.append(
                {
                    "day": target,
                    "symbol": row["symbol"],
                    "prebudget_rank": pre_rank.get(str(row["symbol"])),
                    "range_5d": float(row["range_pct_5d_median"]),
                    "range_20d": float(row["range_pct_20d_median"]),
                    "tv_5d": float(row["tv_5d_median"]),
                    "tv_20d": float(row["tv_20d_median"]),
                    "one_lot": float(row["ONE_LOT_NOTIONAL"]),
                    "next_range": nxt,
                }
            )
        by_symbol = {str(row["symbol"]): row for row in ordered}
        v2_symbols = set(official_symbols)
        auto_symbols = set(_auto50(prior_for_auto))
        for symbol in sorted(v2_symbols & auto_symbols):
            membership["COMMON"].append(_membership(by_symbol[symbol], "COMMON", target, metrics_joined, target_by_symbol.get(symbol)))
        for symbol in sorted(v2_symbols - auto_symbols):
            membership["V2_ONLY"].append(_membership(by_symbol[symbol], "V2_ONLY", target, metrics_joined, target_by_symbol.get(symbol)))
        for symbol in sorted(auto_symbols - v2_symbols):
            source = by_symbol.get(symbol) or _row_for_symbol(symbol, features)
            if source is None:
                source = _partial_row(symbol, series.get(symbol), required, gross)
            if source is None:
                continue
            membership["AUTO50_ONLY"].append(_membership(source, "AUTO50_ONLY", target, metrics_joined, target_by_symbol.get(symbol)))
        v2_ranges = _next_ranges(official_symbols, target_by_symbol, metrics_joined)
        auto_ranges = _next_ranges(sorted(auto_symbols), target_by_symbol, metrics_joined)
        current_path = RUNTIME / f"same_day_am_frozen_universe_{target}.json"
        current_symbols = load_current50_symbols(current_path) if current_path.is_file() else None
        current_ranges = _next_ranges(current_symbols[:50], target_by_symbol, metrics_joined) if current_symbols else []
        v2_med = _median(v2_ranges)
        auto_med = _median(auto_ranges)
        current_med = _median(current_ranges)
        if v2_med is not None:
            v2_daily_range.append(v2_med)
        if auto_med is not None:
            auto_daily_range.append(auto_med)
        daily_rows.append(
            {
                "target_trading_date": target,
                "metrics_joined": metrics_joined,
                "v2_range_median": v2_med,
                "auto_range_median": auto_med,
                "v2_minus_auto": None if v2_med is None or auto_med is None else v2_med - auto_med,
                "current_range_median": current_med,
                "v2_minus_current": None if v2_med is None or current_med is None else v2_med - current_med,
                "common_n": len(v2_symbols & auto_symbols),
                "v2_only_n": len(v2_symbols - auto_symbols),
                "auto_only_n": len(auto_symbols - v2_symbols),
                "budget_ineligible_n": len(rejected),
                "diagnostic_top50_budget_ineligible_n": in_top,
            }
        )
    for group, rows in membership.items():
        by_day: dict[str, list[float]] = {}
        for row in rows:
            if row.get("next_range") is not None:
                by_day.setdefault(str(row["day"]), []).append(float(row["next_range"]))
        daily = {day: float(statistics.median(values)) for day, values in by_day.items() if values}
        for row in rows:
            row["daily_next_range"] = daily.get(str(row["day"]))
    band_summary = {name: _band_summary(rows) for name, rows in bands.items()}
    group_summary = {name: _group_summary(rows) for name, rows in membership.items()}
    auto_delta = [float(row["v2_minus_auto"]) for row in daily_rows if row.get("v2_minus_auto") is not None]
    current_delta = [float(row["v2_minus_current"]) for row in daily_rows if row.get("v2_minus_current") is not None]
    v2_level = _median(v2_daily_range)
    auto_level = _median(auto_daily_range)
    pre_level = _median(prebudget_daily_range)
    gap = None if v2_level is None or auto_level is None else auto_level - v2_level
    closed = None if pre_level is None or v2_level is None else pre_level - v2_level
    fraction = None if gap is None or closed is None or abs(gap) < 1e-9 else closed / gap
    high_rank_rejects = [row for row in screen_rows if row.get("prebudget_rank") is not None and int(row["prebudget_rank"]) <= 50]
    capital_screen = {
        "rejected_name_days": len(screen_rows),
        "median_rejected_per_day": _median([float(row["budget_ineligible_n"]) for row in daily_rows]),
        "median_ineligible_inside_diagnostic_top50": _median([float(value) for value in ineligible_in_top50]),
        "prebudget_rank": _dist([row.get("prebudget_rank") for row in screen_rows]),
        "high_rank_rejected_name_days": len(high_rank_rejects),
        "rejected_range_5d": _dist([row.get("range_5d") for row in screen_rows]),
        "rejected_range_20d": _dist([row.get("range_20d") for row in screen_rows]),
        "rejected_tv_20d": _dist([row.get("tv_20d") for row in screen_rows]),
        "rejected_next_range": _dist([row.get("next_range") for row in screen_rows]),
        "high_rank_rejected_next_range": _dist([row.get("next_range") for row in high_rank_rejects]),
        "diagnostic_prebudget_top50_next_range_median": pre_level,
        "official_v2_next_range_median": v2_level,
        "auto50_next_range_median": auto_level,
        "fraction_of_auto_gap_closed": fraction,
        "note": "diagnostic_prebudget_top50 は同じ4特徴の中央値順位をG適用前に見た診断です。Selectorには採用していません。",
    }
    symbol_counts: dict[str, list[float]] = {}
    for row in membership["AUTO50_ONLY"]:
        symbol_counts.setdefault(str(row["symbol"]), []).append(row.get("next_range"))
    frequent = sorted(symbol_counts.items(), key=lambda item: (-len(item[1]), item[0]))[:15]
    summary = {
        "bands": band_summary,
        "groups": group_summary,
        "capital_screen": capital_screen,
        "paired_auto50": _paired(auto_delta),
        "paired_current50": _paired(current_delta),
    }
    if reproduction_mismatches:
        decision = {
            "verdict": "DAYTRADE50_V2_SCORE_REPRODUCTION_MISMATCH",
            "next": "STOP_BEFORE_MECHANISM_CLAIM",
            "PRIMARY_RANGE_DEFICIT_MECHANISM": "NOT_IDENTIFIED",
            "supported_mechanisms": [],
        }
    else:
        decision = classify(summary)
    concentration = {
        "sector": "UNAVAILABLE_NO_SECTOR_ON_DAILY_BARS",
        "price": {name: _bucket_table(rows, "price", PRICE_EDGES) for name, rows in membership.items() if name != "COMMON"},
        "budget_share": {name: _bucket_table(rows, "budget_share", SHARE_EDGES) for name, rows in membership.items() if name != "COMMON"},
        "liquidity": {name: _bucket_table(rows, "tv_20d", LIQ_EDGES) for name, rows in membership.items() if name != "COMMON"},
        "range_gap": {name: _bucket_table(rows, "range_gap", GAP_EDGES) for name, rows in membership.items() if name != "COMMON"},
        "auto50_only_frequent_symbols": [
            {"symbol": symbol, "days": len(values), "next_range_median": _median(values)}
            for symbol, values in frequent
        ],
    }
    return {
        "verdict": decision["verdict"],
        "next": decision["next"],
        "identity": identity,
        "identity_parity": True,
        "score_changed": False,
        "selector_adopted": False,
        "runtime_applied": False,
        "v16_changed": v16_sha != PINNED_V16_SHA,
        "v16_sha": v16_sha,
        "submit_cancel_live": "0/0/0",
        "saved_validation_verdict": saved.get("verdict"),
        "saved_comparison": saved.get("comparison"),
        "saved_rank_fit": saved.get("cumulative_rank_fit_rate"),
        "rank_fit_is_selector_failure": False,
        "max_concurrent_means_must_hold_5": False,
        "target_date_count": len(daily_rows),
        "metric_days": sum(1 for row in daily_rows if row["metrics_joined"]),
        "target_close_unavailable_dates": sorted(close_unavailable),
        "excluded_fit_dates": sorted(EXCLUDED_SELECTION_DATES),
        "score_reproduction_mismatches": reproduction_mismatches,
        "future_conditioned_selection": 0,
        "classification": {key: decision[key] for key in (
            "LIQUIDITY_EXECUTION_DOMINANCE",
            "LONG_HORIZON_SMOOTHING",
            "CAPITAL_SCREEN_EFFECT",
            "EXECUTION_RANGE_TRADEOFF",
            "RANGE_DEFICIT_BROAD_ACROSS_DAYS",
            "PRIMARY_RANGE_DEFICIT_MECHANISM",
            "supported_mechanisms",
        ) if key in decision},
        "bands": band_summary,
        "groups": group_summary,
        "capital_screen": capital_screen,
        "paired_auto50": summary["paired_auto50"],
        "paired_current50": summary["paired_current50"],
        "concentration": concentration,
        "days": [{key: value for key, value in row.items()} for row in daily_rows],
    }


def _ranks(row: Mapping[str, Any]) -> dict[str, float]:
    return {
        "liquidity": float(row["LIQUIDITY_RANK"]),
        "opp5": float(row["OPPORTUNITY_5D_RANK"]),
        "opp20": float(row["OPPORTUNITY_20D_RANK"]),
        "execution": float(row["EXECUTION_SUITABILITY_RANK"]),
    }


def _row_for_symbol(symbol: str, features: Sequence[Mapping[str, Any]]) -> Optional[dict[str, Any]]:
    for row in features:
        if str(row["symbol"]) == symbol:
            return row
    return None


def _partial_row(symbol: str, item: Optional[Mapping[str, Any]], required: str, gross: float) -> Optional[dict[str, Any]]:
    """Selection-time description for an AUTO50 name that failed the V2 window. Not a new score."""
    if item is None:
        return None
    index = item["index"].get(required)
    if index is None:
        return None
    price = float(item["close"][index])
    if price <= 0:
        return None
    start = max(0, index - 19)
    ranges: list[float] = []
    trading: list[float] = []
    for pos in range(start, index + 1):
        bar = float(item["close"][pos])
        if bar <= 0:
            return None
        ranges.append((float(item["high"][pos]) - float(item["low"][pos])) / bar * 100.0)
        trading.append(float(item["volume"][pos]) * bar)
    suitability = execution_suitability(price)
    if suitability is None or not ranges:
        return None
    width = min(5, len(ranges))
    return {
        "symbol": symbol,
        "prior_close": price,
        "range_pct_5d_median": float(statistics.median(ranges[-width:])),
        "range_pct_20d_median": float(statistics.median(ranges)),
        "tv_5d_median": float(statistics.median(trading[-width:])),
        "tv_20d_median": float(statistics.median(trading)),
        "suitability_feature": float(suitability),
        "ONE_LOT_NOTIONAL": one_lot_notional(price),
        "BUDGET_SHARE": budget_share(price, gross),
        "partial_window": len(ranges) < 20,
    }


def _next_ranges(symbols: Sequence[str], targets: Mapping[str, Mapping[str, float]], joined: bool) -> list[float]:
    if not joined:
        return []
    out = []
    for symbol in symbols:
        target = targets.get(symbol)
        if target is not None and float(target["trading_value"]) > 0:
            out.append(float(target["range_pct"]))
    return out


def _fmt(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.6g}"
    return str(value)


def _markdown(report: Mapping[str, Any]) -> str:
    classification = report.get("classification") or {}
    bands = report.get("bands") or {}
    groups = report.get("groups") or {}
    screen = report.get("capital_screen") or {}
    auto = report.get("paired_auto50") or {}
    current = report.get("paired_current50") or {}
    fit = report.get("saved_rank_fit") or {}
    lines = [
        f"VERDICT: {report.get('verdict')}",
        "",
        f"NEXT: {report.get('next')}",
        "",
        f"PRIMARY_RANGE_DEFICIT_MECHANISM: {classification.get('PRIMARY_RANGE_DEFICIT_MECHANISM')}",
        "",
        f"LIQUIDITY_EXECUTION_DOMINANCE: {classification.get('LIQUIDITY_EXECUTION_DOMINANCE')}",
        f"LONG_HORIZON_SMOOTHING: {classification.get('LONG_HORIZON_SMOOTHING')}",
        f"CAPITAL_SCREEN_EFFECT: {classification.get('CAPITAL_SCREEN_EFFECT')}",
        f"EXECUTION_RANGE_TRADEOFF: {classification.get('EXECUTION_RANGE_TRADEOFF')}",
        f"RANGE_DEFICIT_BROAD_ACROSS_DAYS: {classification.get('RANGE_DEFICIT_BROAD_ACROSS_DAYS')}",
        "",
        "新しいSelectorは作っていません。weight、threshold、feature、G、leverage は変更していません。",
        "",
        f"BASE_CAPITAL_YEN: {(report.get('identity') or {}).get('BASE_CAPITAL_YEN')}",
        f"MAX_LEVERAGE_MULTIPLIER: {(report.get('identity') or {}).get('MAX_LEVERAGE_MULTIPLIER')}",
        f"SYSTEM_GROSS_EXPOSURE_BUDGET_YEN: {(report.get('identity') or {}).get('SYSTEM_GROSS_EXPOSURE_BUDGET_YEN')}",
        f"POLICY_INSTANCE_SHA: {(report.get('identity') or {}).get('POLICY_INSTANCE_SHA')}",
        "",
        f"target dates: {report.get('target_date_count')}",
        f"metric days: {report.get('metric_days')}",
        f"score reproduction mismatches: {report.get('score_reproduction_mismatches')}",
        "",
        "component median ranks top50 / 51-100 / 101-200:",
        f"liquidity {_fmt((bands.get('top50') or {}).get('liquidity_rank'))} / {_fmt((bands.get('rank_51_100') or {}).get('liquidity_rank'))} / {_fmt((bands.get('rank_101_200') or {}).get('liquidity_rank'))}",
        f"opportunity_5d {_fmt((bands.get('top50') or {}).get('opportunity_5d_rank'))} / {_fmt((bands.get('rank_51_100') or {}).get('opportunity_5d_rank'))} / {_fmt((bands.get('rank_101_200') or {}).get('opportunity_5d_rank'))}",
        f"opportunity_20d {_fmt((bands.get('top50') or {}).get('opportunity_20d_rank'))} / {_fmt((bands.get('rank_51_100') or {}).get('opportunity_20d_rank'))} / {_fmt((bands.get('rank_101_200') or {}).get('opportunity_20d_rank'))}",
        f"execution {_fmt((bands.get('top50') or {}).get('execution_rank'))} / {_fmt((bands.get('rank_51_100') or {}).get('execution_rank'))} / {_fmt((bands.get('rank_101_200') or {}).get('execution_rank'))}",
        "",
        "V2_ONLY / AUTO50_ONLY / COMMON selection medians and next-day range:",
    ]
    for name in ("V2_ONLY", "AUTO50_ONLY", "COMMON"):
        row = groups.get(name) or {}
        lines.append(
            f"{name}: n={row.get('name_days')} range5={_fmt((row.get('range_5d') or {}).get('median'))} "
            f"range20={_fmt((row.get('range_20d') or {}).get('median'))} gap={_fmt((row.get('range_gap') or {}).get('median'))} "
            f"gap>1pp={_fmt(row.get('gap_above_1pp_share'))} tv20={_fmt((row.get('tv_20d') or {}).get('median'))} "
            f"suit={_fmt((row.get('suitability') or {}).get('median'))} price={_fmt((row.get('price') or {}).get('median'))} "
            f"one_lot={_fmt((row.get('one_lot') or {}).get('median'))} next_range={_fmt(row.get('daily_next_range_median'))} "
            f"active={_fmt(row.get('active_rate'))} dead={_fmt(row.get('dead_rate'))}"
        )
    lines.extend(
        [
            "",
            f"budget rejected name-days: {screen.get('rejected_name_days')}",
            f"median rejected per day: {_fmt(screen.get('median_rejected_per_day'))}",
            f"median ineligible inside diagnostic top50: {_fmt(screen.get('median_ineligible_inside_diagnostic_top50'))}",
            f"rejected next-range median: {_fmt((screen.get('rejected_next_range') or {}).get('median'))}",
            f"diagnostic prebudget top50 next-range: {_fmt(screen.get('diagnostic_prebudget_top50_next_range_median'))}",
            f"official V2 next-range: {_fmt(screen.get('official_v2_next_range_median'))}",
            f"AUTO50 next-range: {_fmt(screen.get('auto50_next_range_median'))}",
            f"fraction of AUTO gap closed by removing the G screen: {_fmt(screen.get('fraction_of_auto_gap_closed'))}",
            "",
            f"V2-AUTO days={auto.get('days')} positive={auto.get('positive_days')} negative={auto.get('negative_days')} ties={auto.get('ties')}",
            f"median/p25/p75 delta: {_fmt(auto.get('median_delta'))} / {_fmt(auto.get('p25_delta'))} / {_fmt(auto.get('p75_delta'))}",
            f"median delta without worst 10%: {_fmt(auto.get('median_delta_without_worst_10pct'))}",
            "",
            f"V2-CURRENT50 days={current.get('days')} positive={current.get('positive_days')} negative={current.get('negative_days')} ties={current.get('ties')}",
            f"median/p25/p75 delta: {_fmt(current.get('median_delta'))} / {_fmt(current.get('p25_delta'))} / {_fmt(current.get('p75_delta'))}",
            "",
            f"recorded rank-order fit 1/2/3/4/5: {_fmt(fit.get('1'))} / {_fmt(fit.get('2'))} / {_fmt(fit.get('3'))} / {_fmt(fit.get('4'))} / {_fmt(fit.get('5'))}",
            "このfitはSelector failureではありません。MAX_CONCURRENT=5は必ず5本持つ意味ではありません。",
            "sector分解は日足に業種が無いため未実施です。",
            "",
            f"score changed: {report.get('score_changed')}",
            f"selector adopted: {report.get('selector_adopted')}",
            f"CURRENT V16 changed: {report.get('v16_changed')}",
            "runtime applied: false",
            "submit/cancel/live: 0/0/0",
            "",
        ]
    )
    return "\n".join(lines)


def _xlsx(report: Mapping[str, Any], path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "classification"
    sheet.append(["field", "value"])
    for key, value in (report.get("classification") or {}).items():
        sheet.append([key, json.dumps(value) if isinstance(value, (dict, list)) else value])
    sheet.append(["verdict", report.get("verdict")])
    sheet.append(["next", report.get("next")])
    bands = book.create_sheet("bands")
    bands.append(["band", "n", "liquidity", "opportunity_5d", "opportunity_20d", "execution", "opportunity_avg"])
    for name, row in (report.get("bands") or {}).items():
        bands.append([name, row.get("n"), row.get("liquidity_rank"), row.get("opportunity_5d_rank"), row.get("opportunity_20d_rank"), row.get("execution_rank"), row.get("opportunity_avg")])
    groups = book.create_sheet("groups")
    groups.append(["group", "name_days", "range5", "range20", "gap", "gap_gt_1", "tv5", "tv20", "suitability", "price", "one_lot", "budget_share", "next_range", "next_tv", "active", "dead"])
    for name, row in (report.get("groups") or {}).items():
        groups.append([
            name,
            row.get("name_days"),
            (row.get("range_5d") or {}).get("median"),
            (row.get("range_20d") or {}).get("median"),
            (row.get("range_gap") or {}).get("median"),
            row.get("gap_above_1pp_share"),
            (row.get("tv_5d") or {}).get("median"),
            (row.get("tv_20d") or {}).get("median"),
            (row.get("suitability") or {}).get("median"),
            (row.get("price") or {}).get("median"),
            (row.get("one_lot") or {}).get("median"),
            (row.get("budget_share") or {}).get("median"),
            row.get("daily_next_range_median"),
            (row.get("next_trading_value") or {}).get("median"),
            row.get("active_rate"),
            row.get("dead_rate"),
        ])
    days = book.create_sheet("paired_days")
    day_rows = report.get("days") or []
    if day_rows:
        days.append(list(day_rows[0].keys()))
        for row in day_rows:
            days.append([row[key] for key in day_rows[0].keys()])
    screen = book.create_sheet("capital_screen")
    screen.append(["field", "value"])
    for key, value in (report.get("capital_screen") or {}).items():
        screen.append([key, json.dumps(value) if isinstance(value, (dict, list)) else value])
    buckets = book.create_sheet("buckets")
    buckets.append(["axis", "group", "bucket", "name_days", "prior_range_5d_median", "next_range_median"])
    for axis in ("price", "budget_share", "liquidity", "range_gap"):
        for group, rows in ((report.get("concentration") or {}).get(axis) or {}).items():
            for row in rows:
                buckets.append([axis, group, row.get("bucket"), row.get("name_days"), row.get("prior_range_5d_median"), row.get("next_range_median")])
    symbols = book.create_sheet("frequent_auto_only")
    symbols.append(["symbol", "days", "next_range_median"])
    for row in ((report.get("concentration") or {}).get("auto50_only_frequent_symbols") or []):
        symbols.append([row.get("symbol"), row.get("days"), row.get("next_range_median")])
    book.save(path)


def main() -> None:
    report = run()
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(_markdown(report), encoding="utf-8")
    _xlsx(report, OUT / "audit.xlsx")
    print(report["verdict"])
    print(report["next"])


if __name__ == "__main__":
    main()
