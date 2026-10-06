"""report.json / report.md / audit.xlsx / optional precursor freeze. Day2 primary untouched."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_reversal_precursor_day1_v1 import ANALYSIS_ID
from research.futures_reversal_precursor_day1_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "MinuteStats",
    "EpisodeStats",
    "OriginalFive",
    "Episodes",
    "Minutes",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    ms = dict(body.get("minute_stats") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"CASE: {d.get('CASE')} strongest={d.get('strongest_precursor')}",
        "",
        "Question: can a later futures reversal be foreshadowed with T-time causal",
        "information inside AGREEMENT_180S=BOTH_DOWN? T+1m/T+3m BOTH_UP is a",
        "NONCAUSAL label only. Day2 primary freeze is unchanged. No ENTRY/EXIT.",
        "Minute-snapshot and episode views are reported separately; episode is primary.",
        "",
        "## Required answers",
        "",
        f"- 1 minute BOTH_DOWN N: `{a.get('1_minute_BOTH_DOWN_N')}`",
        f"- 2 BOTH_DOWN episode N: `{a.get('2_BOTH_DOWN_episode_N')}`",
        f"- 3 reversal_by_1m N/rate: `{a.get('3_reversal_by_1m_N_rate')}`",
        f"- 4 reversal_by_3m N/rate: `{a.get('4_reversal_by_3m_N_rate')}`",
        f"- 5 P1 prevalence: `{_fmt(a.get('5_P1_prevalence'))}`",
        f"- 6 P1 reversal3m TRUE/FALSE: `{a.get('6_P1_reversal3m_TRUE_FALSE')}`",
        f"- 7 P2 prevalence: `{_fmt(a.get('7_P2_prevalence'))}`",
        f"- 8 P2 reversal3m TRUE/FALSE: `{a.get('8_P2_reversal3m_TRUE_FALSE')}`",
        f"- 9 P3 result: TRUE/FALSE rev3={_fmt((ms.get('P3') or {}).get('rev3_true'))}/"
        f"{_fmt((ms.get('P3') or {}).get('rev3_false'))} lift={_fmt((ms.get('P3') or {}).get('executable_lift'))}",
        f"- 10 P4 result: TRUE/FALSE rev3={_fmt((ms.get('P4') or {}).get('rev3_true'))}/"
        f"{_fmt((ms.get('P4') or {}).get('rev3_false'))} lift={_fmt((ms.get('P4') or {}).get('executable_lift'))}",
        f"- 11 C1 result: TRUE/FALSE rev3={_fmt((ms.get('C1') or {}).get('rev3_true'))}/"
        f"{_fmt((ms.get('C1') or {}).get('rev3_false'))} lift={_fmt((ms.get('C1') or {}).get('executable_lift'))}",
        f"- 12 C2 result: TRUE/FALSE rev3={_fmt((ms.get('C2') or {}).get('rev3_true'))}/"
        f"{_fmt((ms.get('C2') or {}).get('rev3_false'))} lift={_fmt((ms.get('C2') or {}).get('executable_lift'))}",
        f"- 13 strongest precursor: `{a.get('13_strongest_precursor')}`",
        f"- 14 episode-level same direction: `{_fmt(a.get('14_episode_level_same_direction'))}`",
        f"- 15 precursor TRUE TOP LONG 10m mean/median: `{a.get('15_precursor_TRUE_TOP_LONG_10m_mean_median')}`",
        f"- 16 precursor FALSE TOP LONG: `{a.get('16_precursor_FALSE_TOP_LONG')}`",
        f"- 17 executable lift: `{_fmt(a.get('17_executable_lift'))}` bps",
        f"- 18 09:20: `{a.get('18_09:20_precursor_states')}`",
        f"- 19 09:40: `{a.get('19_09:40')}`",
        f"- 20 10:20: `{a.get('20_10:20')}`",
        f"- 21 10:30: `{a.get('21_10:30')}`",
        f"- 22 10:40: `{a.get('22_10:40')}`",
        f"- 23 separates 09:20/10:20 from continued: `{_fmt(a.get('23_separates_0920_1020_from_continued'))}`",
        f"- 24 distinguishes 10:30: `{_fmt(a.get('24_distinguishes_1030'))}`",
        f"- 25 exact causal sequence: {a.get('25_exact_causal_sequence')}",
        f"- 26 usable ENTRY thesis now plausible: `{_fmt(a.get('26_usable_ENTRY_thesis_now_plausible'))}`",
        f"- 27 ENTRY built: `{_fmt(a.get('27_ENTRY_built'))}`",
        f"- 28 EXIT built: `{_fmt(a.get('28_EXIT_built'))}`",
        f"- 29 original Day2 freeze changed: `{_fmt(a.get('29_original_Day2_freeze_changed'))}`",
        f"- 30 Runtime changed: `{_fmt(a.get('30_Runtime_changed'))}`",
        f"- 31 Paper changed: `{_fmt(a.get('31_Paper_changed'))}`",
        f"- 32 submit/cancel/live: `{a.get('32_submit_cancel_live')}`",
        f"- 33 VERDICT: `{a.get('33_VERDICT')}`",
        f"- 34 NEXT: `{a.get('34_NEXT')}`",
        "",
        "No significance claims on small N. 5m/10m reversal windows were not searched.",
        "",
    ]
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(body), encoding="utf-8")
    paths = {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
    }
    freeze = body.get("freeze_manifest")
    if freeze:
        fp = OUT / "reversal_precursor_freeze_manifest.json"
        fp.write_text(json.dumps(json_sanitize(freeze), ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
        paths["reversal_precursor_freeze_manifest"] = str(fp)
    sheets = {
        "Answers": _kv(dict(body.get("answers") or {})),
        "MinuteStats": list((body.get("minute_stats") or {}).values()),
        "EpisodeStats": list((body.get("episode_stats") or {}).values()),
        "OriginalFive": list(body.get("original_five") or []),
        "Episodes": list(body.get("episodes") or []),
        "Minutes": list(body.get("minutes_slim") or []),
        "Decision": _kv(dict(body.get("decision") or {})),
        "Safety": [
            {"key": "ENTRY_built", "value": False},
            {"key": "EXIT_built", "value": False},
            {"key": "Day2_primary_changed", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "submit_cancel_live", "value": (body.get("answers") or {}).get("32_submit_cancel_live")},
            {"key": "future_leakage_n", "value": body.get("future_leakage_n")},
            {"key": "reversal_label_is_noncausal", "value": True},
        ],
    }
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    paths["audit_xlsx"] = str(xlsx)
    return paths
