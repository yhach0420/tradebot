"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.usd_jpy_sector_symbol_response_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "USDJPY_Source",
    "Timestamp_Semantics",
    "Preopen_Response",
    "Intraday_Impulse",
    "Reverse_Causality",
    "Sector_Raw_Response",
    "Sector_Residual_Response",
    "Sector_LeadLag",
    "Sector_Asymmetry",
    "Stock_Raw_Response",
    "Stock_MarketAdjusted",
    "Stock_SectorAdjusted",
    "Stock_Response_Profile",
    "D1_D4",
    "Prior_Residual_Flags",
    "Unexplained_Groups",
    "Next_Driver",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    secs = list(report.get("sector_raw_and_residual") or [])
    stocks = list(report.get("stock_response") or [])
    flags = dict(report.get("prior_residual_flags") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "architecture": report.get("architecture"),
                "kabu_50": False,
                "frozen_validation_opened": False,
                "old_confirmation_used_to_design": False,
                "japan_internal_not_mined": True,
            }
        ),
        "USDJPY_Source": _kv(dict(report.get("source_probe") or {})) + _kv(dict(report.get("acquisition") or {})) + _kv(dict(report.get("fx_meta") or {})),
        "Timestamp_Semantics": _kv(dict(report.get("timestamp_semantics") or {})),
        "Preopen_Response": list(report.get("preopen_sector") or [{"empty": True}]),
        "Intraday_Impulse": [
            {"sector": r.get("sector"), "lead_lag": r.get("lead_lag"), "fwd_rho_residual": r.get("fwd_rho_residual"), "usable": r.get("usable_before_stock_move")}
            for r in secs
        ]
        or [{"empty": True}],
        "Reverse_Causality": [{"sector": r.get("sector"), "reverse_rho": r.get("reverse_rho"), "lead_lag": r.get("lead_lag")} for r in secs] or [{"empty": True}],
        "Sector_Raw_Response": [{"sector": r.get("sector"), "fwd_rho_raw": r.get("fwd_rho_raw"), "n": r.get("n")} for r in secs] or [{"empty": True}],
        "Sector_Residual_Response": [
            {"sector": r.get("sector"), "residual_1m": r.get("residual_1m"), "curve": r.get("residual_curve"), "persist": r.get("response_persistence_1m_to_5m")}
            for r in secs
        ]
        or [{"empty": True}],
        "Sector_LeadLag": [{"sector": r.get("sector"), "lead_lag": r.get("lead_lag"), "usable_before_stock_move": r.get("usable_before_stock_move"), "only_simultaneous": r.get("only_simultaneous")} for r in secs]
        or [{"empty": True}],
        "Sector_Asymmetry": [
            {"sector": r.get("sector"), "jpy_weakening_sector_up": r.get("jpy_weakening_sector_up"), "jpy_strengthening_sector_up": r.get("jpy_strengthening_sector_up"), "asymmetric": r.get("asymmetric")}
            for r in secs
        ]
        or [{"empty": True}],
        "Stock_Raw_Response": [{"symbol": r.get("symbol"), "sector": r.get("sector"), "raw_1m": r.get("raw_1m")} for r in stocks] or [{"empty": True}],
        "Stock_MarketAdjusted": [{"symbol": r.get("symbol"), "market_adj_1m": r.get("market_adj_1m"), "class": r.get("class")} for r in stocks] or [{"empty": True}],
        "Stock_SectorAdjusted": [{"symbol": r.get("symbol"), "sector_adj_1m": r.get("sector_adj_1m"), "class": r.get("class")} for r in stocks] or [{"empty": True}],
        "Stock_Response_Profile": [
            {
                "symbol": r.get("symbol"),
                "sector": r.get("sector"),
                "class": r.get("class"),
                "lead_lag": r.get("lead_lag"),
                "lead_time_min": r.get("lead_time_min"),
                "direct_sensitivity": r.get("direct_sensitivity"),
                "causal_lead": r.get("causal_lead"),
                "usable_before_stock_move": r.get("usable_before_stock_move"),
                "simultaneous_common_news": r.get("simultaneous_common_news"),
                "asymmetric": r.get("asymmetric"),
                "d1_d4_agree": r.get("d1_d4_agree"),
                "up_mean_sec": r.get("up_mean_sec"),
                "down_mean_sec": r.get("down_mean_sec"),
            }
            for r in stocks
        ]
        or [{"empty": True}],
        "D1_D4": [{"sector": r.get("sector"), "block_rho_1m": r.get("block_rho_1m"), "stable_residual_lead": r.get("stable_residual_lead")} for r in secs]
        + [{"symbol": r.get("symbol"), "d1_d4_sec": r.get("d1_d4_sec"), "agree": r.get("d1_d4_agree")} for r in stocks],
        "Prior_Residual_Flags": list(flags.get("explained_by_usdjpy") or []) + list(flags.get("remain_unexplained") or []) or [{"empty": True}],
        "Unexplained_Groups": _kv(dict(report.get("unexplained") or {})),
        "Next_Driver": _kv(dict(report.get("next_driver") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        "# USDJPY_SECTOR_SYMBOL_CAUSAL_RESPONSE_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted parent: `CURRENT_DRIVER_SET_INSUFFICIENT_V1`. Japan-internal drivers were not mined further.",
        "This does **not** mean no external driver exists. Native resolution is **1 minute**, not the previous 5-minute grid.",
        "",
        f"USDJPY source obtained? **{a.get('USDJPY_source_obtained')}**  Free or paid? **{a.get('free_or_paid')}**  Purchase? **{a.get('any_purchase')}**",
        f"History range: `{a.get('history_range')}`",
        f"Bid/Ask or OHLC: {a.get('bid_ask_or_ohlc')}",
        f"Timestamp semantics proven? **{a.get('timestamp_semantics_proven')}**",
        f"Sectors with USDJPY causal lead: **{a.get('sectors_with_USDJPY_causal_lead')}**",
        f"Sectors with only simultaneous association: **{a.get('sectors_with_only_simultaneous_association')}**",
        f"Preopen/overnight stable sectors: **{a.get('preopen_stable_n')}** `{a.get('preopen_stable_sectors')}`",
        f"JPY weakening (USDJPY up) sectors: `{a.get('sectors_respond_JPY_weakening')}`",
        f"JPY strengthening sectors: `{a.get('sectors_respond_JPY_strengthening')}`",
        f"Direct sensitivity after market+sector: **{a.get('stocks_direct_after_market_and_sector_n')}**",
        f"USABLE_DIRECT_CAUSAL_LEAD_N: **{a.get('USABLE_DIRECT_CAUSAL_LEAD_N')}** `{a.get('usable_direct_causal_lead_symbols')}`",
        f"Direct but not usable: `{a.get('direct_but_not_usable')}`",
        f"Prior 12 explained by USDJPY: `{a.get('prior_12_explained_by_USDJPY')}`",
        f"Groups change? **{a.get('response_based_groups_change')}** (C1–C5 not reused)",
        f"Electrical-machinery heterogeneity explained? **{a.get('USDJPY_explains_electrical_machinery_heterogeneity')}**",
        f"Usable before stock move? **{a.get('any_relation_usable_before_stock_move')}**",
        f"Next driver family: **{a.get('next_driver_family')}** — {a.get('next_driver_why')}",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('old_confirmation_used_to_design')}**",
        f"Kabu 50? **{a.get('kabu_50')}**",
        f"ENTRY rules optimized? **{a.get('entry_rules_optimized')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "Do not build playbooks yet.",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
