"""Write report.json / report.md / audit.xlsx only. No CSV dump."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.edge_decay_rca import (
    ANALYSIS_ID,
    DEVELOPMENT_DAYS,
    DOCUMENT_ID,
    POST_DAYS,
    PRIMARY_BASELINE,
    TASK_LABEL,
    UNIFORM10_ROLE,
)
from small_paper.v1r_native_entry_live import FEATURE_ORDER

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "edge_decay_rca"
JST = timezone(timedelta(hours=9))
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)

REQUIRED_KEYS = [
    "DEV_DAYS",
    "POST_DAYS",
    "DEV_PNL",
    "POST_PNL",
    "DEV_PF",
    "POST_PF",
    "DEV_TRADES",
    "POST_TRADES",
    "PNL_DECAY",
    "TRADE_COUNT_COMPONENT",
    "WIN_RATE_COMPONENT",
    "WIN_SIZE_COMPONENT",
    "LOSS_SIZE_COMPONENT",
    "FILL_COMPONENT",
    "SELECTION_COMPONENT",
    "EXIT_COMPONENT",
    "REENTRY_EFFECT",
    "TAIL_COMPONENT",
    "TOP1_TRADE_SHARE_DEV",
    "TOP1_TRADE_SHARE_POST",
    "TOP3_TRADE_SHARE_DEV",
    "TOP3_TRADE_SHARE_POST",
    "REMOVE_TOP3_TRADES_DEV_PNL",
    "REMOVE_TOP3_TRADES_POST_PNL",
    "COMMON_SYMBOL_EDGE",
    "UNIVERSE_COMPOSITION_EFFECT",
    "SELECTION_EDGE_DEV",
    "SELECTION_EDGE_POST",
    "FEATURE_DISTRIBUTION_SHIFT",
    "FEATURE_OUTCOME_RELATIONSHIP_SHIFT",
    "FILL_EXECUTION_SHIFT",
    "EXIT_EDGE_SHIFT",
    "AM_EDGE_DECAY",
    "PM_EDGE_DECAY",
    "PRIMARY_ROOT_CAUSE",
    "SECONDARY_ROOT_CAUSE",
    "EDGE_DECAY_CONFIDENCE",
    "ANCHOR_GRID_NOT_PRIMARY_DRIVER",
    "verdict",
]


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            if v == float("inf"):
                v = "Infinity"
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _required(analysis: dict[str, Any]) -> dict[str, Any]:
    br = analysis.get("bridge") or {}
    ov = analysis.get("overlay") or {}
    eco_d = analysis.get("eco_dev") or {}
    eco_p = analysis.get("eco_post") or {}
    tail_d = analysis.get("tail_dev") or {}
    tail_p = analysis.get("tail_post") or {}
    coh = analysis.get("cohort") or {}
    sel_d = analysis.get("sel_dev") or {}
    sel_p = analysis.get("sel_post") or {}
    feat = analysis.get("feat_dist") or {}
    rel = analysis.get("feat_rel") or {}
    fill = analysis.get("fill") or {}
    ex = analysis.get("exit") or {}
    tod = analysis.get("tod") or {}
    clf = analysis.get("classify") or {}
    u10 = analysis.get("uniform10") or {}
    return {
        "DEV_DAYS": 10,
        "POST_DAYS": 7,
        "DEV_PNL": br.get("DEV_PNL"),
        "POST_PNL": br.get("POST_PNL"),
        "DEV_PF": eco_d.get("PF"),
        "POST_PF": eco_p.get("PF"),
        "DEV_TRADES": eco_d.get("n") or eco_d.get("trades"),
        "POST_TRADES": eco_p.get("n") or eco_p.get("trades"),
        "PNL_DECAY": br.get("observed"),
        "TRADE_COUNT_COMPONENT": br.get("TRADE_COUNT_COMPONENT"),
        "WIN_RATE_COMPONENT": br.get("WIN_RATE_COMPONENT"),
        "WIN_SIZE_COMPONENT": br.get("WIN_SIZE_COMPONENT"),
        "LOSS_SIZE_COMPONENT": br.get("LOSS_SIZE_COMPONENT"),
        "FILL_COMPONENT": ov.get("FILL_COMPONENT"),
        "SELECTION_COMPONENT": ov.get("SELECTION_COMPONENT"),
        "EXIT_COMPONENT": ov.get("EXIT_COMPONENT"),
        "REENTRY_EFFECT": ov.get("REENTRY_EFFECT"),
        "TAIL_COMPONENT": ov.get("TAIL_COMPONENT"),
        "TOP1_TRADE_SHARE_DEV": tail_d.get("top1_share"),
        "TOP1_TRADE_SHARE_POST": tail_p.get("top1_share"),
        "TOP3_TRADE_SHARE_DEV": tail_d.get("top3_share"),
        "TOP3_TRADE_SHARE_POST": tail_p.get("top3_share"),
        "REMOVE_TOP3_TRADES_DEV_PNL": (tail_d.get("remove_top3_trades") or {}).get("pnl"),
        "REMOVE_TOP3_TRADES_POST_PNL": (tail_p.get("remove_top3_trades") or {}).get("pnl"),
        "COMMON_SYMBOL_EDGE": coh.get("COMMON_SYMBOL_EDGE"),
        "UNIVERSE_COMPOSITION_EFFECT": coh.get("UNIVERSE_COMPOSITION_EFFECT"),
        "SELECTION_EDGE_DEV": sel_d.get("rank1_minus_rank5_avg"),
        "SELECTION_EDGE_POST": sel_p.get("rank1_minus_rank5_avg"),
        "FEATURE_DISTRIBUTION_SHIFT": feat.get("FEATURE_DISTRIBUTION_SHIFT"),
        "FEATURE_OUTCOME_RELATIONSHIP_SHIFT": rel.get("FEATURE_OUTCOME_RELATIONSHIP_SHIFT"),
        "FILL_EXECUTION_SHIFT": fill.get("FILL_EXECUTION_SHIFT"),
        "EXIT_EDGE_SHIFT": ex.get("EXIT_EDGE_SHIFT"),
        "AM_EDGE_DECAY": tod.get("AM_EDGE_DECAY"),
        "PM_EDGE_DECAY": tod.get("PM_EDGE_DECAY"),
        "PRIMARY_ROOT_CAUSE": clf.get("PRIMARY_ROOT_CAUSE"),
        "SECONDARY_ROOT_CAUSE": clf.get("SECONDARY_ROOT_CAUSE"),
        "EDGE_DECAY_CONFIDENCE": clf.get("EDGE_DECAY_CONFIDENCE"),
        "ANCHOR_GRID_NOT_PRIMARY_DRIVER": (u10 or {}).get("ANCHOR_GRID_NOT_PRIMARY_DRIVER"),
        "verdict": clf.get("verdict"),
        "bridge_residual": br.get("residual"),
        "bridge_recon": br.get("recon"),
    }


def build_report(*, provenance: dict[str, Any], inventory_check: dict[str, Any], analysis: dict[str, Any]) -> dict[str, Any]:
    req = _required(analysis)
    report = {
        "ANALYSIS_ID": ANALYSIS_ID,
        "DOCUMENT_ID": DOCUMENT_ID,
        "label": TASK_LABEL,
        "PRIMARY_BASELINE": PRIMARY_BASELINE,
        "UNIFORM10_ROLE": UNIFORM10_ROLE,
        "generated_at_jst": datetime.now(JST).isoformat(timespec="seconds"),
        "runtime_changed": False,
        "clock_grid_changed": False,
        "strategy_changed": False,
        "optimization_attempted": False,
        "safety": {"submit": 0, "cancel": 0, "live": 0, "offline_research_only": True},
        "SOURCE_FILE": provenance.get("SOURCE_FILE"),
        "SOURCE_SHA": provenance.get("SOURCE_SHA"),
        "CANONICAL_CLOCK_GRID": provenance.get("CANONICAL_CLOCK_GRID"),
        "DEVELOPMENT_DAYS": list(DEVELOPMENT_DAYS),
        "POST_FREEZE_DAYS": list(POST_DAYS),
        "day_contract": inventory_check,
        "note": (
            "Canonical CLOCK_GRID REFERENCE Dual Lane is PRIMARY. UNIFORM10 is secondary diagnostic only. "
            "Exact 4-way identity: TRADE_COUNT + WIN_RATE + WIN_SIZE + LOSS_SIZE + residual = observed. "
            "Overlays (FILL/SELECTION/EXIT/REENTRY/TAIL) are not forced to sum to decay. "
            "No threshold / feature / TopK / CLOCK_GRID / fill / EXIT / re-entry change."
        ),
        **req,
        "eco_dev": analysis.get("eco_dev"),
        "eco_post": analysis.get("eco_post"),
        "bridge": analysis.get("bridge"),
        "overlay": analysis.get("overlay"),
        "tail_dev": analysis.get("tail_dev"),
        "tail_post": analysis.get("tail_post"),
        "symbols_dev": analysis.get("symbols_dev"),
        "symbols_post": analysis.get("symbols_post"),
        "TOP_SYMBOL_PNL_SHARE_DEV": analysis.get("TOP_SYMBOL_PNL_SHARE_DEV"),
        "TOP_SYMBOL_PNL_SHARE_POST": analysis.get("TOP_SYMBOL_PNL_SHARE_POST"),
        "sel_dev": analysis.get("sel_dev"),
        "sel_post": analysis.get("sel_post"),
        "feat_dist": analysis.get("feat_dist"),
        "feat_dist_selected": analysis.get("feat_dist_selected"),
        "feat_rel": analysis.get("feat_rel"),
        "fill": analysis.get("fill"),
        "exit": analysis.get("exit"),
        "tod": analysis.get("tod"),
        "reentry": analysis.get("reentry"),
        "cohort": analysis.get("cohort"),
        "market": analysis.get("market"),
        "daily": analysis.get("daily"),
        "uniform10": analysis.get("uniform10"),
        "classify": analysis.get("classify"),
        "FEATURE_ORDER": list(FEATURE_ORDER),
    }
    return json_sanitize(report)


def write_md(report: dict[str, Any]) -> str:
    clf = report.get("classify") or {}
    br = report.get("bridge") or {}
    lines = [
        "# TRADEBOT — Edge decay root cause",
        "",
        "Offline RCA. CLOCK_GRID / ENTRY / EXIT / Runtime unchanged. No optimization.",
        "PRIMARY = canonical CLOCK_GRID REFERENCE. UNIFORM10 = secondary diagnostic only.",
        "",
        f"SOURCE_FILE: {report.get('SOURCE_FILE')}",
        f"SOURCE_SHA: {report.get('SOURCE_SHA')}",
        "",
        "## REQUIRED",
    ]
    for k in REQUIRED_KEYS:
        lines.append(f"{k}: {report.get(k)}")
    lines += [
        "",
        "## Exact 4-way identity",
        f"TRADE_COUNT + WIN_RATE + WIN_SIZE + LOSS_SIZE = {br.get('recon')}",
        f"observed POST-DEV = {br.get('observed')}",
        f"residual = {br.get('residual')}",
        "",
        "## Classification",
        f"PRIMARY_ROOT_CAUSE: {clf.get('PRIMARY_ROOT_CAUSE')}",
        f"SECONDARY_ROOT_CAUSE: {clf.get('SECONDARY_ROOT_CAUSE')}",
        f"drivers: {clf.get('drivers')}",
        f"driver_magnitudes: {clf.get('driver_magnitudes')}",
        f"EDGE_DECAY_CONFIDENCE: {clf.get('EDGE_DECAY_CONFIDENCE')}",
        f"verdict: {report.get('verdict')}",
        "",
        "## Uniform10 secondary",
        f"ANCHOR_GRID_NOT_PRIMARY_DRIVER: {report.get('ANCHOR_GRID_NOT_PRIMARY_DRIVER')}",
        f"U10 POST/DEV ratio: {(report.get('uniform10') or {}).get('POST_OVER_DEV_RATIO_U10')}",
        f"REF POST/DEV ratio: {(report.get('uniform10') or {}).get('POST_OVER_DEV_RATIO_REF')}",
        "",
        "STOP. Do not change CLOCK_GRID / Runtime / Strategy.",
        "submit/cancel/live=0/0/0",
    ]
    return "\n".join(lines) + "\n"


def _eco_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for lab, st in (("DEV", report.get("eco_dev")), ("POST", report.get("eco_post"))):
        rec = {"period": lab}
        rec.update(st or {})
        rows.append(rec)
    return rows or [{"status": "empty"}]


def _tail_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for lab, st in (("DEV", report.get("tail_dev")), ("POST", report.get("tail_post"))):
        body = dict(st or {})
        tops = body.pop("top_trades", None)
        rec = {"period": lab, **body}
        rows.append(rec)
        for i, t in enumerate(tops or [], start=1):
            rows.append({"period": lab, "top_i": i, **(t or {})})
    return rows or [{"status": "empty"}]


def _symbol_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for lab, lst in (("DEV", report.get("symbols_dev")), ("POST", report.get("symbols_post"))):
        for r in lst or []:
            rows.append({"period": lab, **r})
    return rows or [{"status": "empty"}]


def _sel_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for lab, st in (("DEV", report.get("sel_dev")), ("POST", report.get("sel_post"))):
        body = st or {}
        cut = body.get("cutoff") or {}
        rec = {
            "period": lab,
            "rank1_minus_rank5_avg": body.get("rank1_minus_rank5_avg"),
            "selected_score_mean": body.get("selected_score_mean"),
            "selected_score_median": body.get("selected_score_median"),
            **{f"cutoff_{k}": v for k, v in cut.items()},
        }
        rows.append(rec)
        for rk, b in (body.get("by_rank") or {}).items():
            rows.append({"period": lab, "rank": rk, **(b or {})})
    return rows or [{"status": "empty"}]


def _feat_shift_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for pop_key in ("feat_dist", "feat_dist_selected"):
        body = report.get(pop_key) or {}
        pop = body.get("population") or pop_key
        for fname, st in (body.get("features") or {}).items():
            rec = {
                "population": pop,
                "feature": fname,
                "psi": st.get("psi"),
                "ks": st.get("ks"),
                "smd": st.get("smd"),
                "n_dev": st.get("n_dev"),
                "n_post": st.get("n_post"),
            }
            for side in ("dev", "post"):
                q = st.get(side) or {}
                for k, v in q.items():
                    rec[f"{side}_{k}"] = v
            rows.append(rec)
    return rows or [{"status": "empty"}]


def _feat_out_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    rel = report.get("feat_rel") or {}
    for fname, st in (rel.get("by_feature") or {}).items():
        for per in ("DEV", "POST"):
            b = (st or {}).get(per) or {}
            rows.append(
                {
                    "feature": fname,
                    "period": per,
                    "n": b.get("n"),
                    "spearman": b.get("spearman"),
                    "spearman_drop": st.get("spearman_drop"),
                }
            )
            for q in b.get("quintiles") or []:
                rows.append({"feature": fname, "period": per, **q})
    return rows or [{"status": "empty"}]


def _fill_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    fill = report.get("fill") or {}
    rows = []
    for lab in ("DEV", "POST"):
        rec = {"period": lab, **(fill.get(lab) or {})}
        rec["FILL_EXECUTION_SHIFT"] = fill.get("FILL_EXECUTION_SHIFT")
        rec["FILL_COMPONENT"] = fill.get("FILL_COMPONENT")
        rec["fill_rate_relative_drop"] = fill.get("fill_rate_relative_drop")
        rows.append(rec)
    spr = fill.get("spread_shift") or {}
    if spr:
        rows.append({"period": "SHIFT", **spr})
    return rows or [{"status": "empty"}]


def _exit_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    ex = report.get("exit") or {}
    rows = []
    for lab in ("DEV", "POST"):
        for fam, st in (ex.get(lab) or {}).items():
            rows.append({"period": lab, "cohort": "ALL", "exit_family": fam, **(st or {})})
        fvr = (ex.get("first_vs_reentry_exit") or {}).get(lab) or {}
        for kind, block in fvr.items():
            for fam, st in (block or {}).items():
                rows.append({"period": lab, "cohort": kind, "exit_family": fam, **(st or {})})
    re = report.get("reentry") or {}
    for lab, st in re.items():
        rec = {"period": lab, "cohort": "REENTRY_BLOCK"}
        rec.update({k: v for k, v in (st or {}).items() if k != "all"})
        if isinstance(st.get("all"), dict):
            rec.update({f"all_{k}": v for k, v in st["all"].items()})
        rows.append(rec)
    return rows or [{"status": "empty"}]


def _tod_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    tod = report.get("tod") or {}
    rows = []
    for lab in ("DEV", "POST"):
        for bucket, st in (tod.get(lab) or {}).items():
            rows.append({"period": lab, "bucket": bucket, **(st or {})})
    rows.append(
        {
            "period": "DECAY",
            "AM_EDGE_DECAY": tod.get("AM_EDGE_DECAY"),
            "PM_EDGE_DECAY": tod.get("PM_EDGE_DECAY"),
        }
    )
    return rows or [{"status": "empty"}]


def _cohort_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    coh = dict(report.get("cohort") or {})
    mkt = report.get("market") or {}
    return [{**coh, **{f"market_{k}": v for k, v in mkt.items() if not isinstance(v, dict)} , **{
        "market_universe_DEV": json.dumps(mkt.get("universe_DEV") or {}, default=str),
        "market_universe_POST": json.dumps(mkt.get("universe_POST") or {}, default=str),
    }}]


def _u10_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    u = report.get("uniform10") or {}
    rows = [{"section": "summary", **{k: v for k, v in u.items() if not isinstance(v, dict)}}]
    for key in ("bridge", "eco_dev", "eco_post", "sel_dev", "sel_post", "cohort"):
        body = u.get(key)
        if isinstance(body, dict):
            rec = {"section": key}
            for k, v in body.items():
                if isinstance(v, (dict, list)):
                    rec[k] = json.dumps(v, ensure_ascii=False, default=str)
                else:
                    rec[k] = v
            rows.append(rec)
    for lab in ("DEV", "POST"):
        tod = (u.get("tod") or {}).get(lab) or {}
        for bucket, st in tod.items():
            rows.append({"section": "tod", "period": lab, "bucket": bucket, **(st or {})})
    return rows or [{"status": "empty"}]


def _daily_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for lab, lst in (report.get("daily") or {}).items():
        for r in lst or []:
            rows.append({"series": lab, **r})
    return rows or [{"status": "empty"}]


def write_xlsx(report: dict[str, Any], path: Path) -> None:
    req = {k: report.get(k) for k in REQUIRED_KEYS}
    req.update(
        {
            "PRIMARY_BASELINE": report.get("PRIMARY_BASELINE"),
            "bridge_residual": report.get("bridge_residual"),
            "SOURCE_SHA": report.get("SOURCE_SHA"),
        }
    )
    sheets = {
        "Summary": [req],
        "Daily": _daily_rows(report),
        "TradeEconomics": _eco_rows(report),
        "Tail": _tail_rows(report),
        "Symbols": _symbol_rows(report),
        "Selection": _sel_rows(report),
        "FeatureShift": _feat_shift_rows(report),
        "FeatureOutcome": _feat_out_rows(report),
        "Fill": _fill_rows(report),
        "Exit": _exit_rows(report),
        "TimeOfDay": _tod_rows(report),
        "CommonCohort": _cohort_rows(report),
        "Uniform10_Check": _u10_rows(report),
    }
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(str(name)[:31])
        _sheet(ws, rows if isinstance(rows, list) else [rows])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_artifacts(report: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    jp, mp, xp = OUT / "report.json", OUT / "report.md", OUT / "audit.xlsx"
    slim = dict(report)
    jp.write_text(json.dumps(json_sanitize(slim), ensure_ascii=False, indent=2), encoding="utf-8")
    mp.write_text(write_md(report), encoding="utf-8")
    write_xlsx(report, xp)
    return {"report_json": str(jp), "report_md": str(mp), "audit_xlsx": str(xp)}
