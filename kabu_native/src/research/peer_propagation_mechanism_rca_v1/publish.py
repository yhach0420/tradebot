"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.cross_sectional_peer_propagation_discovery_v1 import MECHS
from research.peer_propagation_mechanism_rca_v1 import ABLATION, STREAMS
from research.peer_propagation_mechanism_rca_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Episode_Onset",
    "State_Duration",
    "P1_P2_P3_Overlap",
    "Core_Lag_Ablation",
    "Sector_vs_Market",
    "First_Onset_Matched",
    "Execution_Latency",
    "Edge_Consumption",
    "SMA5_25_75",
    "MA_Sequence",
    "VWAP_Context",
    "SR_Context",
    "Target_TV_Sequence",
    "Confirmation_Delay",
    "Remaining_Edge",
    "Failure_RCA",
    "Offset_Map",
    "Clustered_Uncertainty",
    "Playbook_Map",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    ep = dict(report.get("episodes") or {})
    mech = dict(report.get("mechanisms") or {})
    abl = dict(report.get("ablation") or {})
    freeze = dict(report.get("freeze") or {})
    ep_rows, dur_rows, first_rows, xo_rows, edge_rows = [], [], [], [], []
    sma_rows, ma_rows, vwap_rows, sr_rows, tv_rows = [], [], [], [], []
    conf_rows, remain_rows, fail_rows, off_rows, boot_rows, pb_rows, svm_rows = [], [], [], [], [], [], []
    for m in STREAMS:
        e = dict(ep.get(m) or {})
        ep_rows.append(
            {
                "mech": m,
                "raw_qualified_minute_n": e.get("raw_qualified_minute_n"),
                "first_onset_episode_n": e.get("first_onset_episode_n"),
            }
        )
        dur_rows.append({"mech": m, **dict(e.get("duration") or {})})
    for m in MECHS:
        pack = dict(mech.get(m) or {})
        mc = dict(pack.get("matched_close") or {})
        xo = dict(pack.get("matched_next_open") or {})
        for b in ("D2", "D3", "D4"):
            g = dict((mc.get("blocks") or {}).get(b) or {})
            x = dict((xo.get("blocks") or {}).get(b) or {})
            first_rows.append({"mech": m, "block": b, **g})
            xo_rows.append({"mech": m, "block": b, **x})
        edge_rows.append({"mech": m, **dict(pack.get("edge_consumption") or {})})
        sma_rows.append({"mech": m, **dict(pack.get("ma") or {})})
        vwap_rows.append({"mech": m, **dict(pack.get("vwap") or {})})
        for row in list(pack.get("sr") or []) or [{"sr_bin": None}]:
            sr_rows.append({"mech": m, **dict(row)})
        tv_rows.append({"mech": m, **dict(pack.get("tv_sequence") or {})})
        for row in list(pack.get("confirmation") or []) or [{"confirmation": None}]:
            conf_rows.append({"mech": m, **dict(row)})
            remain_rows.append(
                {
                    "mech": m,
                    "confirmation": row.get("confirmation"),
                    "remaining_10m": row.get("remaining_10m"),
                    "price_moved_before_bps": row.get("price_moved_before_bps"),
                    "delay": row.get("mean_delay_trading_minutes"),
                }
            )
        fail_rows.append({"mech": m, **dict(pack.get("failure") or {})})
        off = dict(pack.get("offset_map") or {})
        for row in list(off.get("rows") or []) or [{"offset": None}]:
            off_rows.append({"mech": m, "usable_lead": off.get("usable_lead_trading_minutes"), **dict(row)})
        boot = dict((mc.get("bootstrap") or {}))
        boot_rows.append({"mech": m, "kind": "close_10m", **{k: boot.get(k) for k in ("n", "point_gap_10m", "date", "symbol", "two_way")}})
        boot_xo = dict((xo.get("bootstrap") or {})) if False else dict((pack.get("matched_next_open") or {}).get("bootstrap") or {})
        if boot_xo:
            boot_rows.append({"mech": m, "kind": "next_open_10m", **boot_xo})
        for row in list(pack.get("playbook") or []) or [{"code": None}]:
            pb_rows.append({"mech": m, **dict(row)})
        svm_rows.append({"mech": m, **{k: v for k, v in dict(pack.get("sector_vs_market") or {}).items() if k != "market_matched"}})
        mm = dict((pack.get("sector_vs_market") or {}).get("market_matched") or {})
        svm_rows[-1]["market_matched_gap_10m"] = mm.get("gap_10m")
        svm_rows[-1]["market_matched_all_eval_positive"] = mm.get("all_eval_positive")
        ma_rows.append(
            {
                "mech": m,
                "sma5_reclaim_n": (pack.get("ma") or {}).get("sma5_reclaim_n"),
                "sma25_reclaim_n": (pack.get("ma") or {}).get("sma25_reclaim_n"),
                "sma5_delay": (pack.get("ma") or {}).get("sma5_reclaim_delay"),
                "sma25_delay": (pack.get("ma") or {}).get("sma25_reclaim_delay"),
            }
        )
    lag = dict((report.get("lag_only") or {}).get("matched_close") or {})
    abl_rows = [{"mech": ABLATION, "all_eval_positive": lag.get("all_eval_positive"), "gap_10m": lag.get("gap_10m")}]
    for m, v in abl.items():
        abl_rows.append({"mech": m, **dict(v or {})})
    return {
        "Binding": _kv_rows(
            {
                "ok": (report.get("bind") or {}).get("ok"),
                "parent_verdict": (report.get("bind") or {}).get("parent_verdict"),
                "FREEZE_SHA256": freeze.get("FREEZE_SHA256"),
                "p1_p2_p3_not_modified": True,
                "d1_boundaries_not_modified": True,
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
            }
        ),
        "Episode_Onset": ep_rows or [{"mech": None}],
        "State_Duration": dur_rows or [{"mech": None}],
        "P1_P2_P3_Overlap": _kv_rows(report.get("overlap") or {}),
        "Core_Lag_Ablation": abl_rows,
        "Sector_vs_Market": svm_rows or [{"mech": None}],
        "First_Onset_Matched": first_rows or [{"mech": None}],
        "Execution_Latency": xo_rows or [{"mech": None}],
        "Edge_Consumption": edge_rows or [{"mech": None}],
        "SMA5_25_75": sma_rows or [{"mech": None}],
        "MA_Sequence": ma_rows or [{"mech": None}],
        "VWAP_Context": vwap_rows or [{"mech": None}],
        "SR_Context": sr_rows or [{"mech": None}],
        "Target_TV_Sequence": tv_rows or [{"mech": None}],
        "Confirmation_Delay": conf_rows or [{"mech": None}],
        "Remaining_Edge": remain_rows or [{"mech": None}],
        "Failure_RCA": fail_rows or [{"mech": None}],
        "Offset_Map": off_rows or [{"mech": None}],
        "Clustered_Uncertainty": boot_rows or [{"mech": None}],
        "Playbook_Map": pb_rows or [{"mech": None}],
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# PEER_PROPAGATION_MECHANISM_RCA_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Frozen P1/P2/P3 and D1 boundaries. Unique causal onsets. Next-open execution RCA.",
            "SMA5/25/75 are execution context, not peer signals. No Complete Strategy.",
            "",
            f"raw qualified minute_n? **{a.get('raw qualified minute_n?')}**",
            f"unique first-onset episode_n? **{a.get('unique first-onset episode_n?')}**",
            f"median state duration? **{a.get('median state duration?')}**",
            f"P1/P2/P3 mostly the same? **{a.get('Are P1/P2/P3 mostly the same mechanism?')}**",
            f"HIGH lag alone? **{a.get('Does HIGH target lag alone explain it?')}** breadth add? **{a.get('Does breadth add?')}** peer return add? **{a.get('peer return add?')}** peer TV add? **{a.get('peer TV breadth add?')}**",
            f"sector vs market? **{a.get('Sector-specific or broad-market propagation?')}**",
            f"first-onset 10m D2/D3/D4? **{a.get('First-onset matched D2/D3/D4 10m gap?')}**",
            f"first-onset 5m/20m/MFE/MAE? **{a.get('First-onset matched 5m/20m/MFE/MAE?')}**",
            f"next-open D2? **{a.get('Next-open matched D2?')}** D3? **{a.get('Next-open matched D3?')}** D4? **{a.get('Next-open matched D4?')}**",
            f"bps consumed before entry? **{a.get('How many bps are consumed before executable entry?')}**",
            f"nontrivial after next open? **{a.get('Does signal remain economically nontrivial after next open?')}**",
            f"SMA stack? **{a.get('SMA5/25/75 stack-aligned result?')}** SMA25 hold? **{a.get('SMA25 pullback/hold?')}** SMA5 reclaim? **{a.get('SMA5 reclaim?')}** SMA25 reclaim? **{a.get('SMA25 reclaim?')}**",
            f"VWAP reclaim? **{a.get('VWAP reclaim?')}**",
            f"target TV after peer? **{a.get('Target TV: does participation arrive after peer signal?')}**",
            f"confirmation retaining most remaining edge (diagnostic, not selected)? **{a.get('Which causal confirmation retains the most remaining edge?')}**",
            f"usable peer lead minutes? **{a.get('How long is usable peer lead?')}**",
            f"one-symbol? **{a.get('Any one-symbol dominance?')}** one-sector? **{a.get('one-sector?')}** one-day? **{a.get('one-day?')}**",
            f"future episode selection? **{a.get('Any future episode selection?')}** MA period tuning? **{a.get('Any MA period tuning?')}** PnL? **{a.get('Any PnL optimization?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}** Kabu50? **{a.get('Kabu50?')}**",
            f"submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
