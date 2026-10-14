"""Audit the 11:25 SESSION_FLAT clock against the 11:30 morning close.

The saved 268-trade ledger is read only. No runtime file is changed.
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

NATIVE = Path(__file__).resolve().parents[3]
SRC = NATIVE / "src"
ROOT = NATIVE.parent
for item in (str(ROOT), str(NATIVE), str(SRC)):
    if item not in sys.path:
        sys.path.insert(0, item)

OUT = NATIVE / "results" / "research" / "session_flat_execution_boundary_audit_v1"
os.environ["TRADEBOT_C3_RESIDUAL_PATH"] = str(OUT / "residual_unused.json")

from research.anchor_timing_robustness.grid import hm_epoch
from research.current50_v16_complete_strategy_risk_foundation_v1.foundation import _capture_dir, _symbol
from research.execution_realism_layer_v1.model import ExecutionRealismBook, proven_sell
from small_paper.am_pm_session_policy import AmPmSessionPolicy
from small_paper.fixed_support_x1_session import production_session_bounds

DAY = "20261013"
JST = timezone(timedelta(hours=9))
REALISM = NATIVE / "results" / "research" / "execution_realism_layer_v1" / "report.json"
LEDGER = (
    NATIVE / "results" / "small_paper" / DAY / "live_session_085018" / "fixed_support_x1_paper_ledger.jsonl",
)
PROTECTED = (
    REALISM,
    NATIVE / "results" / "research" / "c3_execution_loss_mechanism_rca_v1" / "report.json",
    NATIVE / "results" / "research" / "v19_actual_paper_e2e_parity_20261013" / "report.json",
    NATIVE / "src" / "small_paper" / "am_pm_session_policy.py",
    NATIVE / "src" / "small_paper" / "fixed_support_x1_session.py",
    NATIVE / "src" / "small_paper" / "c3_paper_session.py",
)


def _fmt(stamp) -> str:
    return datetime.fromtimestamp(float(stamp), JST).strftime("%H:%M:%S.%f")[:-3]


def _sha(path: Path) -> str:
    if not path.is_file():
        return ""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _loads():
    try:
        import orjson

        return orjson.loads
    except Exception:
        return json.loads


def load_am_exits() -> list[dict]:
    rows = []
    path = LEDGER[0]
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        if row.get("event") != "EXIT":
            continue
        rows.append(row)
    return rows


def scan(symbols: set[str], capture: Path, t0: float, t1: float) -> dict[str, list[dict]]:
    probe = ExecutionRealismBook(0.0)
    probe.notifier = None
    loads = _loads()
    found: dict[str, list[dict]] = defaultdict(list)
    for part in sorted(capture.glob("push_part_*.jsonl")):
        with part.open("rb") as handle:
            for raw in handle:
                if not raw.strip():
                    continue
                rec = loads(raw)
                if rec.get("kind") not in (None, "market_push"):
                    continue
                payload = rec.get("original_payload") if isinstance(rec.get("original_payload"), dict) else None
                if not payload:
                    continue
                symbol = _symbol(str(rec.get("symbol") or payload.get("Symbol") or ""))
                if symbol not in symbols:
                    continue
                ev = probe._normalize({"symbol": symbol, "payload": payload, "received_at": rec.get("received_at")})
                if not ev:
                    continue
                stamp = float(ev["t"])
                if stamp < t0 - 1.0 or stamp > t1 + 1.0:
                    continue
                decision = proven_sell(ev)
                found[symbol].append(
                    {
                        "t": stamp,
                        "clock": _fmt(stamp),
                        "px": ev.get("px"),
                        "bid": ev.get("bid"),
                        "buy1_px": decision.get("px"),
                        "qty": decision.get("qty"),
                        "age": decision.get("age"),
                        "ok": bool(decision.get("ok")),
                        "reason": decision.get("reason"),
                        "state": ev.get("state"),
                        "executable": ev.get("executable"),
                        "special": ev.get("special"),
                    }
                )
    for rows in found.values():
        rows.sort(key=lambda row: float(row["t"]))
    return found


def causal_sell(events: list[dict], start: float, end: float, qty: float) -> dict:
    """Sell only on a proven bid strictly after the 11:25 request and before 11:30."""
    remaining = float(qty)
    fills = []
    for ev in events:
        stamp = float(ev["t"])
        if stamp <= start or stamp >= end or not ev["ok"]:
            continue
        size = ev.get("qty")
        if size is None or float(size) <= 0:
            continue
        take = min(remaining, float(size))
        fills.append(
            {
                "t": stamp,
                "clock": ev["clock"],
                "px": ev["buy1_px"],
                "qty": take,
                "displayed_qty": float(size),
                "age": ev["age"],
                "state": ev["state"],
            }
        )
        remaining -= take
        if remaining <= 1e-9:
            break
    filled = float(qty) - remaining
    return {
        "filled_qty": filled,
        "residual_qty": remaining,
        "fills": fills,
        "vwap": None
        if filled <= 0
        else round(sum(float(row["px"]) * float(row["qty"]) for row in fills) / filled, 6),
    }


def write_xlsx(path: Path, report: dict) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "判定"
    for row in (
        ("項目", "内容"),
        ("11:20", report["clocks"]["11:20"]),
        ("11:25", report["clocks"]["11:25"]),
        ("11:30", report["clocks"]["11:30"]),
        ("因果モデル", report["causal_model"]["valid"]),
        ("発注遅延", "実測なし。仮定0秒"),
        ("268件", "上書きしない"),
        ("submit/cancel/live", "0/0/0"),
    ):
        sheet.append(row)
    positions = book.create_sheet("5件")
    positions.append(("銘柄", "数量", "11:25の枠", "PaperのBid時刻", "Paper価格", "11:25後の最初の有効Bid", "数量", "状態", "11:30残数量", "研究価格"))
    for row in report["positions"]:
        first = row["causal"]["fills"][0] if row["causal"]["fills"] else {}
        positions.append(
            (
                row["symbol"],
                row["qty"],
                row["cap_at_1125"],
                row["paper_bid_clock"],
                row["paper_px"],
                first.get("clock"),
                first.get("displayed_qty"),
                first.get("state"),
                row["causal"]["residual_qty"],
                row["causal"]["vwap"],
            )
        )
    bids = book.create_sheet("板1125から1130")
    bids.append(("銘柄", "時刻", "Bid", "数量", "鮮度秒", "状態", "採用可", "理由"))
    for row in report["window_bids"]:
        bids.append((row["symbol"], row["clock"], row["buy1_px"], row["qty"], row["age"], row["state"], row["ok"], row["reason"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def write_md(path: Path, report: dict) -> None:
    lines = [
        f"# {report['verdict']}",
        "",
        f"NEXT: {report['next']}",
        "",
        "これはSESSION_FLATの時刻の監査である。C3の損失原因の評価ではない。268件 / −29,500円は変更していない。",
        "",
        "## 11:25の意味",
        "",
        f"- 11:20: {report['clocks']['11:20']}",
        f"- 11:25: {report['clocks']['11:25']}",
        f"- 11:30: {report['clocks']['11:30']}",
        "",
        f"前場の最終ENTRYは {report['last_am_entry_clock']}。11:20より後のENTRYは {report['entries_after_1120']} 件。",
        "",
        "## 因果モデル",
        "",
        report["causal_model"]["text"],
        "",
        "## 5件",
        "",
    ]
    for row in report["positions"]:
        first = row["causal"]["fills"][0] if row["causal"]["fills"] else None
        lines.append(
            f"- {row['symbol']}: 11:25時点で数量 {row['qty']}、枠を1つ占有。Paperの記帳Bidは {row['paper_bid_clock']} の {row['paper_px']}。"
            f"11:25より後かつ11:30より前の有効Bidは {row['proven_in_window_n']} 件。"
            + (
                f"最初の因果約定は {first['clock']}、表示数量 {first['displayed_qty']}、価格 {first['px']}、状態 {first['state']}。"
                if first
                else "因果約定に使えるBidはない。"
            )
            + f"11:30の残数量は {row['causal']['residual_qty']}。"
        )
    lines.extend(
        [
            "",
            f"11:25の未決済は {report['open_at_1125_n']} 銘柄、合計数量 {report['open_qty_at_1125']}、枠 {report['slots_at_1125']}。",
            f"研究モデルが11:30まで待ったあとの残数量合計は {report['residual_qty_at_1130']}。",
            "この研究価格は268件のPnLに加算していない。発注遅延は実測ではない。",
            "",
            "Runtimeは変更していない。submit/cancel/live = 0/0/0。",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    before = {str(path): _sha(path) for path in PROTECTED}
    realism = json.loads(REALISM.read_text(encoding="utf-8"))
    residuals = list(realism["primary_residuals"])
    if len(residuals) != 5:
        raise SystemExit(f"RESIDUAL_COUNT {len(residuals)}")
    exits = load_am_exits()
    by_uid = {row["signal_uid"]: row for row in exits}
    t_1120 = float(hm_epoch(DAY, 11, 20))
    t_1125 = float(hm_epoch(DAY, 11, 25))
    t_1130 = float(hm_epoch(DAY, 11, 30))
    _session, book_start, book_end = production_session_bounds(DAY, "am")
    morning = AmPmSessionPolicy.morning()
    pm_note = AmPmSessionPolicy.safe_pm_flatten()
    am_entries = [row for row in exits]
    last_entry = max(float(row["signal_timestamp"]) for row in am_entries)
    entries_after_1120 = sum(1 for row in am_entries if float(row["signal_timestamp"]) >= t_1120)
    capture = _capture_dir(DAY)
    if capture is None:
        raise SystemExit("NO_CAPTURE")
    symbols = {row["symbol"] for row in residuals}
    events = scan(symbols, capture, t_1125, t_1130)
    positions = []
    window_rows = []
    reject_reasons = Counter()
    for row in residuals:
        base = by_uid[row["signal_uid"]]
        symbol_events = events.get(row["symbol"], [])
        window = [ev for ev in symbol_events if t_1125 <= float(ev["t"]) <= t_1130]
        proven = [ev for ev in window if t_1125 < float(ev["t"]) < t_1130 and ev["ok"]]
        for ev in window:
            if not ev["ok"]:
                reject_reasons[str(ev["reason"])] += 1
            if t_1125 <= float(ev["t"]) < t_1130:
                window_rows.append({"symbol": row["symbol"], **ev})
        sale = causal_sell(symbol_events, t_1125, t_1130, 100.0)
        positions.append(
            {
                "signal_uid": row["signal_uid"],
                "symbol": row["symbol"],
                "qty": 100.0,
                "notional": row["notional"],
                "cap_at_1125": 1,
                "paper_bid_clock": _fmt(base["exit_fill_timestamp"]),
                "paper_trigger_clock": _fmt(base["exit_trigger_timestamp"]),
                "paper_px": base["exit_price"],
                "paper_reason": base["exit_reason"],
                "paper_pnl_yen": base["pnl_yen"],
                "paper_pnl_kept_out_of_this_audit": True,
                "proven_in_window_n": len(proven),
                "states_in_window": sorted({str(ev["state"]) for ev in proven}),
                "causal": sale,
            }
        )
    open_qty = sum(row["qty"] for row in positions)
    residual_qty = sum(float(row["causal"]["residual_qty"]) for row in positions)
    clocks = {
        "11:20": (
            f"Pilotの entry_stop と allowed_entry_end。新規のPilot側ENTRYを止める時刻。"
            f"C3のX1台帳は sess_end までシグナルを受け、11:20では止めない。実装の entry_stop は {morning.entry_stop}。"
        ),
        "11:25": (
            f"Pilotの session_end と force_close。ランナーはここで morning_session_close として停止し、"
            f"C3台帳は close_at_runtime_boundary を呼ぶ。X1の sess_end も { _fmt(book_end) } で、"
            f"これ以降の板はパスに入らない。EXIT要求を出してから後のBidを待つ窓口は無い。"
            f"SESSION_FLATの価格は、その時点の最後のBidであり、時刻は境界より前になる。"
            f"売却約定期限ではない。"
        ),
        "11:30": (
            f"前場の取引終了として、PM flatten のコメントが「11:25のフラットは11:30のセッション終了の前」と書いている。"
            f"AMポリシーのフィールドには11:30が無い。X1の境界コメントは research 11:30 を使わないと明記している。"
            f"PMの対応関係は flatten {pm_note.flatten_start}、連続約定の期限 {pm_note.continuous_deadline}、場の終了の言及は15:30。"
        ),
    }
    all_filled = all(float(row["causal"]["residual_qty"]) <= 1e-9 for row in positions)
    causal_text = (
        "11:25ちょうどにEXIT要求を置き、約定は11:25より後かつ11:30より前の、C3と同じBuy1条件を満たすBidだけを使う。"
        "要求と同じ板は使わない。11:24のBidは使わない。発注遅延は実測がなく、仮定は0秒。"
        + (
            "5件ともその窓のBidで100株が埋まる。これは境界監査の観察であり、268件のPnLではない。"
            if all_filled
            else "窓内で埋まらない数量は11:30の残とする。価格は付けない。"
        )
    )
    report = {
        "verdict": "SESSION_FLAT_BOUNDARY_AUDIT_COMPLETE",
        "next": "NO_RUNTIME_CHANGE",
        "scope": "SESSION_FLAT clock audit. Not a C3 loss-mechanism result.",
        "baseline_pnl_unchanged": {"n": 268, "pnl_yen": -29500.0},
        "clocks": clocks,
        "book_start": _fmt(book_start),
        "book_end": _fmt(book_end),
        "book_end_matches_1125": abs(book_end - t_1125) <= 1e-6,
        "policy": morning.to_dict(),
        "last_am_entry_clock": _fmt(last_entry),
        "entries_after_1120": entries_after_1120,
        "entries_after_1125": sum(1 for row in am_entries if float(row["signal_timestamp"]) >= t_1125),
        "causal_model": {
            "valid": True,
            "delay_class": "ASSUMED",
            "assumed_order_delay_sec": 0.0,
            "measured_order_latency_sec": None,
            "exit_request": "11:25:00",
            "fill_window": "(11:25:00, 11:30:00)",
            "text": causal_text,
            "replaces_268_pnl": False,
        },
        "positions": positions,
        "open_at_1125_n": len(positions),
        "open_qty_at_1125": open_qty,
        "slots_at_1125": len(positions),
        "residual_qty_at_1130": residual_qty,
        "window_bid_n": len(window_rows),
        "window_bids": window_rows,
        "reject_reasons": dict(reject_reasons),
        "submit": 0,
        "cancel": 0,
        "live": 0,
        "runtime_changed": False,
        "protected_sha_before": before,
    }
    after = {str(path): _sha(path) for path in PROTECTED}
    report["protected_sha_after"] = after
    report["protected_unchanged"] = before == after
    if not report["protected_unchanged"]:
        report["verdict"] = "BOUNDARY_AUDIT_PROTECTED_FILE_CHANGED"
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    write_md(OUT / "report.md", report)
    write_xlsx(OUT / "audit.xlsx", report)
    print(report["verdict"], report["entries_after_1120"], report["residual_qty_at_1130"], report["protected_unchanged"], flush=True)
    for row in positions:
        first = row["causal"]["fills"][:1]
        print(row["symbol"], row["proven_in_window_n"], first, row["causal"]["residual_qty"], flush=True)


if __name__ == "__main__":
    main()
