"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.entry_panel_exact_reconciliation import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "entry_panel_exact_reconciliation"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "RawExact3",
    "Causality",
    "PanelExactCauses",
    "PanelExactRows",
    "A2_155",
    "C3_1349",
    "FeatureParity",
    "ScoreParity",
    "Population",
    "RankParity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _feat_block(title: str, d: dict[str, Any] | None) -> list[str]:
    lines = [title, ""]
    for name, st in (d or {}).items():
        lines.append(
            f"  {name}: N={st.get('N')} max_abs_diff={st.get('max_abs_diff')} "
            f"mean_abs_diff={st.get('mean_abs_diff')} p95_abs_diff={st.get('p95_abs_diff')} "
            f"exact_match_N={st.get('exact_match_N')}"
        )
    lines.append("")
    return lines


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    raw = report.get("raw_exact") or {}
    cas = report.get("causality") or {}
    pe = report.get("panel_exact") or {}
    c3 = report.get("c3_1349") or {}
    fp = report.get("feature_parity") or {}
    dec = report.get("decision") or {}
    proof = report.get("code_proof") or {}
    return "\n".join(
        [
            "# CANONICAL ENTRY PANEL / EXACT REBASE V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE RESEARCH INFRASTRUCTURE REPAIR. No new ENTRY model. No C4.",
            "No execution-aware objective. No A0/A2/C3 PnL rebase this run.",
            "Frozen: C3_TOP_EDGE_SUPPORTED_PORTFOLIO_FAIL / C3_MULTIFACTOR_EXACT_FAILURE /",
            "C3_EXECUTION_COUPLING_NOT_YET_RESOLVED / ENTRY_DECISION_CONTRACT_MISMATCH.",
            "ENTRY origin = PENDING.anchor. fill_time / nearest-anchor / exit_time forbidden.",
            "Canonical panel = Exact Dual-Lane CLOCK snapshot (one implementation, two consumers).",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"RAW_EXACT_MISMATCH_N: {g.get('RAW_EXACT_MISMATCH_N')}",
            f"RAW_EXACT_EXPLAINED_N: {g.get('RAW_EXACT_EXPLAINED_N')}",
            f"RAW_EXACT_UNEXPLAINED_N: {g.get('RAW_EXACT_UNEXPLAINED_N')}",
            "",
            f"FUTURE_EVENT_USE_N: {g.get('FUTURE_EVENT_USE_N')}",
            "",
            f"PANEL_EXACT_MISMATCH_N_BEFORE: {g.get('PANEL_EXACT_MISMATCH_N_BEFORE')}",
            f"PRIMARY_MISMATCH_CAUSE: {g.get('PRIMARY_MISMATCH_CAUSE')}",
            "",
            f"A2_155_EXPLAINED_N: {g.get('A2_155_EXPLAINED_N')}",
            f"A2_155_UNEXPLAINED_N: {g.get('A2_155_UNEXPLAINED_N')}",
            "",
            f"C3_1349_PRIMARY_CAUSE: {g.get('C3_1349_PRIMARY_CAUSE')}",
            "",
            f"POSTFIX_PANEL_EXACT_EXEC_MISMATCH_N: {g.get('POSTFIX_PANEL_EXACT_EXEC_MISMATCH_N')}",
            "",
            f"CURRENT_SCORE_AVAILABILITY_MISMATCH_N: {g.get('CURRENT_SCORE_AVAILABILITY_MISMATCH_N')}",
            f"C3_SCORE_AVAILABILITY_MISMATCH_N: {g.get('C3_SCORE_AVAILABILITY_MISMATCH_N')}",
            "",
            f"POSTFIX_CURRENT_SCORE_MAX_DIFF: {g.get('POSTFIX_CURRENT_SCORE_MAX_DIFF')}",
            f"POSTFIX_C3_OOF_SCORE_MAX_DIFF: {g.get('POSTFIX_C3_OOF_SCORE_MAX_DIFF')}",
            f"POSTFIX_C3_FINAL_SCORE_MAX_DIFF: {g.get('POSTFIX_C3_FINAL_SCORE_MAX_DIFF')}",
            "",
            f"POSTFIX_DECISION_POPULATION_JACCARD: {g.get('POSTFIX_DECISION_POPULATION_JACCARD')}",
            f"POSTFIX_RESEARCH_ONLY_N: {g.get('POSTFIX_RESEARCH_ONLY_N')}",
            f"POSTFIX_EXACT_ONLY_N: {g.get('POSTFIX_EXACT_ONLY_N')}",
            "",
            f"CURRENT_TOP1_AGREEMENT: {g.get('CURRENT_TOP1_AGREEMENT')}",
            f"CURRENT_TOP3_OVERLAP: {g.get('CURRENT_TOP3_OVERLAP')}",
            f"CURRENT_TOP5_OVERLAP: {g.get('CURRENT_TOP5_OVERLAP')}",
            "",
            f"C3_TOP1_AGREEMENT: {g.get('C3_TOP1_AGREEMENT')}",
            f"C3_TOP3_OVERLAP: {g.get('C3_TOP3_OVERLAP')}",
            f"C3_TOP5_OVERLAP: {g.get('C3_TOP5_OVERLAP')}",
            "",
            f"CANONICAL_RESEARCH_CONTRACT_PROVEN: {_tf(g.get('CANONICAL_RESEARCH_CONTRACT_PROVEN'))}",
            f"PERFORMANCE_REBASE_ALLOWED: {_tf(g.get('PERFORMANCE_REBASE_ALLOWED'))}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## 1. RAW vs EXACT (3)",
            "",
            str(raw.get("note") or ""),
            f"causes={raw.get('cause_counts')} explained={raw.get('explained_n')} unexplained={raw.get('unexplained_n')}.",
            "Same source event id/time. Qty<=0 sets board.special after ingest gate; RAW reconstruct treats it as SpecialQuote.",
            "Not future-event use. Not a reason to STOP as EXACT_CAUSALITY_DEFECT_FOUND.",
            "",
            "## 2. CAUSALITY",
            "",
            f"FUTURE_EVENT_USE_N={cas.get('FUTURE_EVENT_USE_N')} BOARD_EVENT_AFTER_T0_N={cas.get('BOARD_EVENT_AFTER_T0_N')} "
            f"SERIES_FUTURE_UNUSED_TAIL_N={cas.get('SERIES_FUTURE_UNUSED_TAIL_N')}.",
            "Code proof:",
            f"- CLOCK: {proof.get('CLOCK_FIRE')}",
            f"- CURRENT: {proof.get('CURRENT_FEATURES')}",
            f"- C3 lookback: {proof.get('C3_LOOKBACK')}",
            f"- C3 VWAP: {proof.get('C3_VWAP')}",
            f"- C3 cache: {proof.get('C3_CACHE')}",
            "",
            "## 3. PANEL vs EXACT BEFORE",
            "",
            f"N={pe.get('n')} counts={pe.get('cause_counts')} PRIMARY={pe.get('primary')}.",
            str(pe.get("note") or ""),
            "Old C3 extract.py calls patch_board_buf_for_marks() (KEEPALL compact_tail no-op) while ingest still truncates boards[sym] to last 20000.",
            "last_idx on KEEPALL finds the same last quote as Exact CLOCK (lag_close=all mismatches).",
            "classify_t0_row(src_rows[i0]) then reads the truncated list at the KEEPALL index: empty src (MISSING_EVENT) or a later-session row (CACHE).",
            "",
            "## 4. A2 155",
            "",
            f"EXPLAINED={g.get('A2_155_EXPLAINED_N')} UNEXPLAINED={g.get('A2_155_UNEXPLAINED_N')}.",
            "Set = A2 CLOSED with RAW=true Exact=true Panel=false at PENDING.anchor. Same KEEPALL/boards desync. Not fill_time join.",
            "",
            "## 5. C3 1349",
            "",
            f"N={c3.get('n')} counts={c3.get('cause_counts')} PRIMARY={c3.get('primary')}.",
            str(c3.get("note") or ""),
            "Exact C3 scores exist because first-CLOCK source_series still contains morning volume increments; volume_features_at cuts at t0.",
            "Old panel series is truncated afternoon history, so morning t0 searchsorted yields i<0 and vwap_dist_bps is missing.",
            "",
            "## 6. FEATURE VECTOR PARITY (canonical vs Exact CLOCK)",
            "",
            *_feat_block("CURRENT six", fp.get("CURRENT")),
            *_feat_block("C3 six", fp.get("C3")),
            "## 7. CANONICAL BUILDER",
            "",
            "CanonicalEngine = BFollowEngine executable_t0_only CLOCK fire + classify_t0_row + preentry_from_board + causal_features + C3LookupEngine _xs_imbalance/apply_norm/_predict_row.",
            "No MarkBoardBuf patch. No second last_idx-after-full-day implementation.",
            "Runtime not changed. Research reconstruction aligned to runtime-equivalent Exact semantics.",
            "",
            "## 8. DECISION",
            "",
            f"VERDICT={g.get('VERDICT')}",
            f"CANONICAL_RESEARCH_CONTRACT_PROVEN={_tf(g.get('CANONICAL_RESEARCH_CONTRACT_PROVEN'))}",
            f"PERFORMANCE_REBASE_ALLOWED={_tf(g.get('PERFORMANCE_REBASE_ALLOWED'))}",
            str(dec.get("note") or ""),
            "NEXT if PASS: CANONICAL ENTRY PERFORMANCE REBASE. Not started this run.",
            "",
            "## STOP",
            "",
            "No Performance Rebase. No C4. No execution-aware model. Runtime/C14 unchanged.",
            "CLOCK/EXIT/CAP/Fill unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
