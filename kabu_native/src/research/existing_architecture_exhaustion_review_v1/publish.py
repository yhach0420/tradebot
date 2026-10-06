"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.existing_architecture_exhaustion_review_v1 import ANALYSIS_ID
from research.existing_architecture_exhaustion_review_v1.isolation import OUT
from research.existing_architecture_exhaustion_review_v1.inventory import ROW_FIELDS

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "architecture_inventory",
    "duplicate_map",
    "open_strength_audit",
    "x6r3_audit",
    "closure_evidence",
    "remaining_candidates",
    "decision",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def inventory_table_md(rows: list[dict[str, Any]]) -> list[str]:
    lines = [
        "| ARCHITECTURE_ID | FIRST_KNOWN | STATUS | PAPER | FULL_CAUSAL | REMAINING | VERDICT |",
        "|---|---|---|---|---|---|---|",
    ]
    for r in rows:
        remaining = r.get("REMAINING_ELIGIBLE")
        if remaining is None:
            remaining = False
        lines.append(
            "| "
            + " | ".join(
                [
                    str(r["ARCHITECTURE_ID"]),
                    str(r["FIRST_KNOWN_DATE"]),
                    str(r["CURRENT_STATUS"]),
                    "true" if r.get("ACTUAL_PAPER_RUN") else "false",
                    "true" if r.get("FULL_CAUSAL_RUN") else "false",
                    "true" if remaining else "false",
                    str(r.get("FINAL_VERDICT") or "").replace("|", "/"),
                ]
            )
            + " |"
        )
    return lines


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    inv = list(report.get("inventory") or [])
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"CASE: {_fmt(d.get('CASE_NAME') or d.get('CASE'))}",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT'))}",
        "",
        "PNL_USED_TO_SELECT_NEXT: false",
        "",
        "NEW_ARCHITECTURE_CREATED: false",
        "",
        "NEW_ECONOMICS_RUN: false",
        "",
        "STRESS_RAW_READ: false",
        "",
        "FUTURE_USED: false",
        "",
        "## Answers 1-40",
        "",
    ]
    for k, v in a.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("## Architecture inventory")
    lines.append("")
    compact = []
    remaining_ids = set(d.get("remaining_eligible_ids") or [])
    for r in inv:
        compact.append(
            {
                **{k: r.get(k) for k in ("ARCHITECTURE_ID", "FIRST_KNOWN_DATE", "CURRENT_STATUS", "FINAL_VERDICT", "ACTUAL_PAPER_RUN", "FULL_CAUSAL_RUN")},
                "REMAINING_ELIGIBLE": r["ARCHITECTURE_ID"] in remaining_ids,
            }
        )
    lines.extend(inventory_table_md(compact))
    lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def closure_rows(inv: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in inv:
        if str(r["CURRENT_STATUS"]).startswith("CLOSED_") or r["CURRENT_STATUS"] in {
            "ALREADY_EXECUTED_FULL_STRATEGY",
            "NOT_FULL_STRATEGY",
            "NEW_DERIVATIVE_NOT_PREEXISTING",
        }:
            out.append(
                {
                    "ARCHITECTURE_ID": r["ARCHITECTURE_ID"],
                    "CURRENT_STATUS": r["CURRENT_STATUS"],
                    "FINAL_VERDICT": r["FINAL_VERDICT"],
                    "WHY_CLOSED_OR_OPEN": r["WHY_CLOSED_OR_OPEN"],
                    "EVIDENCE_PATHS": r["EVIDENCE_PATHS"],
                }
            )
    return out


def duplicate_rows(inv: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for r in inv:
        if r.get("IS_DUPLICATE") or r.get("EXACT_DUPLICATE_OF"):
            rows.append(
                {
                    "ARCHITECTURE_ID": r["ARCHITECTURE_ID"],
                    "EXACT_DUPLICATE_OF": r.get("EXACT_DUPLICATE_OF") or "",
                    "CURRENT_STATUS": r["CURRENT_STATUS"],
                    "WHY": r["WHY_CLOSED_OR_OPEN"],
                }
            )
    if not rows:
        rows = [{"ARCHITECTURE_ID": "NONE", "EXACT_DUPLICATE_OF": "", "CURRENT_STATUS": "", "WHY": "no exact Full Strategy identity duplicates among remaining"}]
    return rows


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))


__all__ = [
    "SHEET_ORDER",
    "ROW_FIELDS",
    "build_markdown",
    "closure_rows",
    "duplicate_rows",
    "json_sanitize",
    "kv_rows",
    "write_artifacts",
]
