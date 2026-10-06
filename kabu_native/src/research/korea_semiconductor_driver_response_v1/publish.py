"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.korea_semiconductor_driver_response_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Preserved_Driver_Map",
    "Korea_Source_Audit",
    "Timestamp_Semantics",
    "Session_Semantics",
    "KOSPI200_Response",
    "Samsung_Response",
    "SKHynix_Response",
    "Samsung_After_KOSPI",
    "SKHynix_After_KOSPI",
    "Leader_Incremental",
    "Reverse_Causality",
    "Sector_Response",
    "Stock_Response",
    "Japan_Tech_Targets",
    "US_Risk_Comparison",
    "FX_Comparison",
    "D1_D4",
    "Driver_Coverage_105",
    "Unexplained_Groups",
    "Next_Driver",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    probe = dict(report.get("source_probe") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "parent_verdict": (report.get("bind") or {}).get("parent_nqes_verdict"),
                "architecture": report.get("architecture"),
                "fx_preserved": True,
                "us_risk_preserved": True,
                "nq_not_reinterpreted_as_true_cme": True,
                "kabu_50": False,
                "frozen_validation_opened": False,
                "old_confirmation_used_to_design": False,
                "playbooks_built": False,
                "purchase": False,
            }
        ),
        "Preserved_Driver_Map": _kv(dict(report.get("preserved_driver_map") or {})),
        "Korea_Source_Audit": _kv(
            {
                "status": probe.get("status"),
                "native_krx_available": probe.get("native_krx_available"),
                "proxy_available": probe.get("proxy_available"),
                "did_not_fabricate_korea_proxy": probe.get("did_not_fabricate_korea_proxy"),
                "did_not_use_hkg_or_chi_as_korea_proxy": probe.get("did_not_use_hkg_or_chi_as_korea_proxy"),
                "purchase": probe.get("purchase"),
                "krx": probe.get("krx"),
                "jetta": probe.get("jetta"),
                "required_history": probe.get("required_history"),
                "required_instruments": probe.get("required_instruments"),
                "rows": probe.get("rows"),
            }
        ),
        "Timestamp_Semantics": _kv(dict(report.get("timestamp_semantics") or {})),
        "Session_Semantics": _kv(dict(report.get("session_semantics") or {})),
        "KOSPI200_Response": list(report.get("kospi200_response") or [{"empty": True}]),
        "Samsung_Response": list(report.get("samsung_response") or [{"empty": True}]),
        "SKHynix_Response": list(report.get("skhynix_response") or [{"empty": True}]),
        "Samsung_After_KOSPI": list(report.get("samsung_after_kospi") or [{"empty": True}]),
        "SKHynix_After_KOSPI": list(report.get("skhynix_after_kospi") or [{"empty": True}]),
        "Leader_Incremental": _kv(dict(report.get("leader_incremental") or {})),
        "Reverse_Causality": _kv(dict(report.get("reverse_causality") or {})),
        "Sector_Response": list(report.get("sector_response") or [{"empty": True}]),
        "Stock_Response": list(report.get("stock_response") or [{"empty": True}]),
        "Japan_Tech_Targets": list(report.get("japan_tech_targets") or [{"empty": True}]),
        "US_Risk_Comparison": _kv(dict(report.get("us_risk_comparison") or {})),
        "FX_Comparison": _kv(dict(report.get("fx_comparison") or {})),
        "D1_D4": list(report.get("d1_d4") or [{"empty": True}]),
        "Driver_Coverage_105": list(report.get("coverage_rows") or [{"empty": True}]),
        "Unexplained_Groups": _kv(dict(report.get("unexplained") or {})),
        "Next_Driver": _kv(dict(report.get("next_driver") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    q = dict((report.get("coverage") or {}).get("quality") or report.get("coverage_quality") or {})
    lines = [
        "# KOREA_SEMICONDUCTOR_DRIVER_RESPONSE_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted parent: `US_TECH_PROXY_RESPONSE_FOUND_TRUE_FUTURES_PROOF_REQUIRED_V1`.",
        "NQ proxy is **not** true CME futures proof. FX_RESPONSE_MAP_V1 and US-risk proxy evidence are preserved.",
        "Do not build playbooks.",
        "",
        f"KOSPI200 minute source found? **{a.get('KOSPI200_minute_source_found')}** `{a.get('KOSPI200_native_or_proxy')}` `{a.get('KOSPI200_free_or_paid')}`",
        f"Purchase? **{a.get('purchase')}**",
        f"Samsung 005930 minute source found? **{a.get('Samsung_005930_minute_source_found')}**",
        f"SK hynix 000660 minute source found? **{a.get('SK_hynix_000660_minute_source_found')}**",
        f"KRX products listed: IX1002=`{a.get('krx_ix1002')}` ST1002=`{a.get('krx_st1002')}` exact_price_krw=`{a.get('krx_exact_price_krw')}`",
        f"Jetta catalog n=`{a.get('jetta_n_instruments')}` KR country n=`{a.get('jetta_kr_country_n')}`",
        f"Timestamp semantics proven? **{a.get('timestamp_semantics_proven')}**",
        f"Session semantics proven from bars? **{a.get('session_semantics_proven')}**",
        "",
        f"Korea broad market lead Japanese tech? **{a.get('Korea_broad_market_leads_Japanese_tech')}**",
        f"Samsung lead? **{a.get('Samsung_leads_Japanese_tech')}**",
        f"SK hynix lead? **{a.get('SK_hynix_leads_Japanese_tech')}**",
        f"Samsung incremental after KOSPI200? **{a.get('Samsung_incremental_after_KOSPI200')}**",
        f"SK hynix incremental after KOSPI200? **{a.get('SK_hynix_incremental_after_KOSPI200')}**",
        f"Strongest Korea semiconductor driver? **{a.get('strongest_Korea_semiconductor_driver')}**",
        f"STRICT_CAUSAL_LEAD: `{a.get('STRICT_CAUSAL_LEAD')}`",
        f"SIMULTANEOUS_COMMON_NEWS only? **{a.get('SIMULTANEOUS_COMMON_NEWS_only')}**",
        "",
        f"6857 explained? **{a.get('explained_6857')}**",
        f"8035 explained? **{a.get('explained_8035')}**",
        f"6920 explained? **{a.get('explained_6920')}**",
        f"285A0 explained? **{a.get('explained_285A0')}**",
        f"Tech 13 with usable causal drivers: **{a.get('tech_13_with_usable_causal_drivers_n')}**",
        f"Korea adds beyond US broad risk? **{a.get('Korea_adds_beyond_US_broad_risk')}**",
        f"Korea adds beyond USDJPY where relevant? **{a.get('Korea_adds_beyond_USDJPY_where_relevant')}**",
        "",
        f"Stocks with ≥1 usable external driver: **{a.get('stocks_with_at_least_one_usable_external_driver_n')}** / 105",
        f"Remain unexplained: **{a.get('remain_unexplained_n')}**",
        f"Evidence quality: TRUE_DRIVER_PROVEN n=**{a.get('TRUE_DRIVER_PROVEN_n')}** PROXY_SUPPORTED_CAUSAL n=**{a.get('PROXY_SUPPORTED_CAUSAL_n')}** OVERNIGHT_CAUSAL n=**{a.get('OVERNIGHT_CAUSAL_n')}**",
        f"Not all 16 equally proven? **{a.get('not_all_16_equally_proven')}** `{q}`",
        "",
        f"Unresolved economic group that determines NEXT: **{a.get('unresolved_economic_group_that_determines_NEXT')}**",
        f"Minimum missing for tech: **{a.get('minimum_missing_for_tech')}**",
        f"Next executable free family: **{a.get('next_executable_free_family')}** `{a.get('next_driver_family')}`",
        f"Why: {a.get('next_driver_why')}",
        "",
        f"ENTRY playbooks built? **{a.get('ENTRY_playbooks_built')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
        f"Kabu 50? **{a.get('Kabu_50')}**",
        f"purchase? **{a.get('purchase')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "Preserve confirmed FX and US-risk maps. Do not build playbooks yet.",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
