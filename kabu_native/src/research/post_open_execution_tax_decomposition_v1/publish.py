"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.post_open_execution_tax_decomposition_v1 import ANALYSIS_ID
from research.post_open_execution_tax_decomposition_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Parent_Parity",
    "Execution_Pin",
    "Anchor_Population",
    "W5_Fills",
    "Common_Endpoint",
    "Selection_Effect",
    "Population_Comparison",
    "Block_Comparison",
    "Univariate_W5",
    "Interactions_W5",
    "Tree_W5",
    "LOBO",
    "Interpretation",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    pops = dict(report.get("populations") or {})
    tree = dict(report.get("tree") or {})
    blocks = dict(report.get("blocks") or {})
    return {
        "Summary": _kv(dict(report.get("answers") or {})),
        "Parent_Parity": _kv(dict(report.get("parent_parity") or {})),
        "Execution_Pin": _kv(dict(report.get("w5_pin") or {})),
        "Anchor_Population": _kv(dict(pops.get("P0") or {})),
        "W5_Fills": _kv(
            {
                "W5_fill_n": pops.get("W5_fill_n"),
                "W5_fill_rate": pops.get("W5_fill_rate"),
                "W5_fill_day_n": pops.get("W5_fill_day_n"),
                "W5_fill_symbol_n": pops.get("W5_fill_symbol_n"),
                "support_ok": pops.get("support_ok"),
                "per_day_fill_rate": pops.get("per_day_fill_rate"),
                "per_clock_fill_rate": pops.get("per_clock_fill_rate"),
            }
        ),
        "Common_Endpoint": _kv(dict(report.get("common_endpoint") or {})),
        "Selection_Effect": _kv(dict(report.get("selection") or {})),
        "Population_Comparison": [
            {"pop": "P0", **dict(pops.get("P0") or {})},
            {"pop": "P1", **dict(pops.get("P1") or {})},
            {"pop": "P2", **dict(pops.get("P2") or {})},
            {"pop": "P2_minus_P1", "mean": pops.get("P2_minus_P1_mean"), "median": pops.get("P2_minus_P1_median")},
        ],
        "Block_Comparison": [{"block": k, **v} for k, v in dict(blocks.get("blocks") or {}).items()] or [{"empty": True}],
        "Univariate_W5": [
            {
                "feature": u.get("feature"),
                "spearman": u.get("spearman_w5_markout"),
                "stable": u.get("stable"),
                "agree": u.get("block_agree"),
            }
            for u in list(report.get("univariate") or [])
        ]
        or [{"empty": True}],
        "Interactions_W5": list(report.get("interactions") or []) or [{"empty": True}],
        "Tree_W5": list(tree.get("leaves") or []) or _kv({k: v for k, v in tree.items() if k != "leaves"}),
        "LOBO": list((report.get("lobo") or {}).get("folds") or []) or [{"empty": True}],
        "Interpretation": _kv(dict(report.get("interpretation") or {})),
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        f"CASE: {d.get('CASE')}",
        "TRUE_OOS: false",
        "CERTIFIED: false",
        "CANDIDATE_STRATEGY_N: 0",
        "",
        str((report.get("interpretation") or {}).get("MARKET_MECHANISM") or ""),
        "",
    ]
    for i in range(1, 39):
        matches = [k for k in a if k.split("_", 1)[0] == str(i)]
        if matches:
            lines.append(f"{i}. {matches[0]}: {a.get(matches[0])}")
    lines.extend(["", "No Full Causal. No CAP. Nonfill is not PnL. No V5 rescue.", "", "STOP.", ""])
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
