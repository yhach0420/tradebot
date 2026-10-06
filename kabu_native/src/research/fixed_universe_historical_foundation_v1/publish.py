"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.fixed_universe_historical_foundation_v1 import ANALYSIS_ID
from research.fixed_universe_historical_foundation_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "UNIVERSE_CANDIDATES",
    "SECTOR_COVERAGE",
    "DATA_SOURCE_MATRIX",
    "TIMESTAMP_SEMANTICS",
    "CORPORATE_ACTIONS",
    "SESSION_CALENDAR",
    "RESEARCH_SPLIT_PLAN",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        "UNIVERSE_CANDIDATES": list(report.get("universe_candidates") or []),
        "SECTOR_COVERAGE": list(report.get("sector_coverage") or []),
        "DATA_SOURCE_MATRIX": list(report.get("data_sources") or []),
        "TIMESTAMP_SEMANTICS": list(report.get("timestamp_semantics") or []),
        "CORPORATE_ACTIONS": _kv(dict(report.get("corporate_actions") or {})),
        "SESSION_CALENDAR": list(report.get("session_calendar") or []),
        "RESEARCH_SPLIT_PLAN": _kv(dict(report.get("research_split") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    u = dict(report.get("universe_methodology") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        f"THEN: {d.get('THEN')}",
        "",
        str(d.get("INTERPRETATION") or ""),
        "",
        "## Required answers",
        "",
        f"1. Fixed Universe approach viable? **{a.get('1_Fixed_Universe_approach_viable')}**",
        f"2. recommended stock N? **{a.get('2_recommended_stock_N')}**",
        f"3. selection methodology? {a.get('3_selection_methodology')}",
        f"4. survivorship/panel conditioning documented? **{a.get('4_survivorship_panel_conditioning_documented')}**",
        f"5. J-Quants minute data usable? **{a.get('5_J_Quants_minute_data_usable')}**",
        f"6. history available? **{a.get('6_history_available')} years**",
        f"7. futures historical source? {a.get('7_futures_historical_source')}",
        f"8. USDJPY source candidate? {a.get('8_USDJPY_source_candidate')}",
        f"9. US futures source candidate? {a.get('9_US_futures_source_candidate')}",
        f"10. oil source candidate? {a.get('10_oil_source_candidate')}",
        f"11. Asia source candidates? {a.get('11_Asia_source_candidates')}",
        f"12. rates source candidates? {a.get('12_rates_source_candidates')}",
        f"13. timestamp semantics fully known? **{a.get('13_timestamp_semantics_fully_known')}**",
        f"14. DST handling plan? **{a.get('14_DST_handling_plan')}**",
        f"15. corporate action plan? **{a.get('15_corporate_action_plan')}**",
        f"16. minimum common historical period? {a.get('16_minimum_common_historical_period')}",
        f"17. data gaps? {a.get('17_data_gaps')}",
        f"18. cost estimate if known? {a.get('18_cost_estimate_if_known')}",
        f"19. blockers before research? {a.get('19_blockers_before_research')}",
        f"20. same-bar leakage protection ready? **{a.get('20_same_bar_leakage_protection_ready')}** (spec yes, panel not built)",
        f"21. chronological split plan? **{a.get('21_chronological_split_plan')}**",
        f"22. simple technical library definition ready? **{a.get('22_simple_technical_library_definition_ready')}**",
        f"23. 20260914 live plan changed? **{a.get('23_20260914_live_plan_changed')}**",
        f"24. Day1 mining reopened? **{a.get('24_Day1_mining_reopened')}**",
        f"25. Runtime changed? **{a.get('25_Runtime_changed')}**",
        f"26. Paper changed? **{a.get('26_Paper_changed')}**",
        f"27. submit/cancel/live? **{a.get('27_submit_cancel_live')}**",
        f"28. VERDICT? **{a.get('28_VERDICT')}**",
        f"29. NEXT? **{a.get('29_NEXT')}**",
        "",
        "## Universe",
        "",
        f"- candidate N: {u.get('CANDIDATE_N')} (not frozen)",
        f"- panel-conditioned: {u.get('PANEL_CONDITIONED_HISTORICAL_RESEARCH')}",
        f"- claim 1-year-ago same universe: {u.get('CLAIM_ONE_YEAR_AGO_SAME_UNIVERSE')}",
        f"- required sectors covered: {u.get('REQUIRED_SECTORS_ALL_COVERED')}",
        "",
        "## Safety",
        "",
        "- No strategy search.",
        "- No profit optimization.",
        "- No data purchase in this run.",
        "- 20260914 futures / breadth / Day2 frozen confirmation unchanged.",
        "- Day1 mining remains closed.",
        "- Historical 1-minute is not fill/spread proof.",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
