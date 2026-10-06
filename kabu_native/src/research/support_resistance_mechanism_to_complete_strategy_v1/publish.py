"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import math

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out
from research.support_resistance_mechanism_to_complete_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Strategy_Definitions",
    "A2_Signals",
    "C1_Signals",
    "Entry_Lineage",
    "Structural_Targets",
    "Exit_Lineage",
    "A2_Trades",
    "C1_Trades",
    "Combined_Trades",
    "CAP_Blocked",
    "SameSymbol_Blocked",
    "Occupancy",
    "Slot_Release",
    "Reentry",
    "Target_Invalidation_Ambiguity",
    "A1_Passive_Diagnostic",
    "Cost_Stress",
    "Daily",
    "D1_D4",
    "Long_Short",
    "Support_Resistance",
    "MFE_MAE",
    "Giveback",
    "Concentration",
    "Tail_Robustness",
    "Tie_Order_Sensitivity",
    "Freeze_Candidate",
    "Safety",
)

SIG_COLS = (
    "signal_id",
    "mechanism",
    "symbol",
    "date",
    "block",
    "side",
    "zone_id",
    "zone_role",
    "decision_bar",
    "entry_time",
    "entry_price",
    "target_price",
    "target_zone_id",
    "exit_reason",
    "exit_t",
    "exit_price",
    "gross_yen",
    "stress_yen",
    "same_bar_entry",
    "target_future_leakage",
    "target_invalidation_ambiguous",
)

TRADE_COLS = SIG_COLS + (
    "strategy_id",
    "hold_min",
    "mfe_bps",
    "mae_bps",
    "attribution",
    "occupancy_at_entry",
    "reentry",
    "slot_release_t",
)

STRIP = {
    "_markdown",
    "_a2_signals",
    "_c1_signals",
    "_a1_rows",
    "_trades",
}


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in d.items()]
    return [{"key": "value", "value": d}]


def _rows(items: list[dict[str, Any]], cols: tuple[str, ...]) -> list[dict[str, Any]]:
    out = []
    for r in items:
        out.append({c: r.get(c) for c in cols})
    return out


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def _econ_block(name: str, e: dict[str, Any]) -> dict[str, Any]:
    return {
        "strategy": name,
        "signal_n": e.get("signal_n"),
        "trades": e.get("trades"),
        "gross_yen": e.get("gross_yen"),
        "stress_yen": e.get("stress_yen"),
        "pf_gross": e.get("pf_gross"),
        "pf_stress": e.get("pf_stress"),
        "median_trade": e.get("median_trade"),
        "maxDD_yen": e.get("maxDD_yen"),
        "positive_day_rate": e.get("positive_day_rate"),
        "break_even_cost_bps": e.get("break_even_cost_bps"),
        "target_exit_n": e.get("target_exit_n"),
        "invalidation_exit_n": e.get("invalidation_exit_n"),
        "session_close_n": e.get("session_close_n"),
        "coherent": e.get("coherent"),
        "occupancy_mean_at_entry": e.get("occupancy_mean_at_entry"),
        "slot_utilization": e.get("slot_utilization"),
        "reentry_n": e.get("reentry_n"),
        "cap_blocked_n": e.get("cap_blocked_n"),
        "same_symbol_blocked_n": e.get("same_symbol_blocked_n"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a2s = list(report.get("_a2_signals") or [])
    c1s = list(report.get("_c1_signals") or [])
    a2t = list((report.get("A2") or {}).get("_trades") or [])
    c1t = list((report.get("C1") or {}).get("_trades") or [])
    ct = list((report.get("COMBINED") or {}).get("_trades") or [])
    a2 = dict(report.get("A2") or {})
    c1 = dict(report.get("C1") or {})
    comb = dict(report.get("COMBINED") or {})
    ea = dict(a2.get("economics") or {})
    ec = dict(c1.get("economics") or {})
    ex = dict(comb.get("economics") or {})
    d = dict(report.get("decision") or {})
    lineage = []
    for r in a2s + c1s:
        lineage.append(
            {
                "signal_id": r.get("signal_id"),
                "decision_bar": r.get("decision_bar"),
                "decision_available_at": r.get("decision_available_at"),
                "entry_eligible_bar": r.get("entry_eligible_bar"),
                "entry_time": r.get("entry_time"),
                "entry_price": r.get("entry_price"),
                "same_bar_entry": r.get("same_bar_entry"),
            }
        )
    targets = []
    for r in a2s + c1s:
        targets.append(
            {
                "signal_id": r.get("signal_id"),
                "target_zone_id": r.get("target_zone_id"),
                "target_price": r.get("target_price"),
                "target_known_at_entry": r.get("target_known_at_entry"),
                "target_activated_at": r.get("target_activated_at"),
                "target_future_leakage": r.get("target_future_leakage"),
            }
        )
    exits = []
    for r in a2t + c1t:
        exits.append(
            {
                "signal_id": r.get("signal_id"),
                "entry_thesis": r.get("entry_thesis"),
                "entry_zone_id": r.get("entry_zone_id"),
                "entry_zone_bounds": r.get("entry_zone_bounds"),
                "target_zone_id": r.get("target_zone_id"),
                "target_known_at_entry": r.get("target_known_at_entry"),
                "invalidation_reason": r.get("invalidation_reason"),
                "exit_reason": r.get("exit_reason"),
                "exit_decision_time": r.get("exit_decision_time"),
                "exit_execution_time": r.get("exit_execution_time"),
                "exit_retroactive": r.get("exit_retroactive"),
            }
        )
    amb = [r for r in a2s + c1s if r.get("target_invalidation_ambiguous")]
    cap_b = list(a2.get("cap_blocked") or []) + list(c1.get("cap_blocked") or []) + list(comb.get("cap_blocked") or [])
    sym_b = list(a2.get("same_symbol_blocked") or []) + list(c1.get("same_symbol_blocked") or []) + list(comb.get("same_symbol_blocked") or [])
    occ = [
        _econ_block("A2", ea),
        _econ_block("C1", ec),
        _econ_block("COMBINED", ex),
    ]
    slot = [{"signal_id": t.get("signal_id"), "exit_reason": t.get("exit_reason"), "exit_t": t.get("exit_t"), "slot_release_t": t.get("slot_release_t")} for t in a2t + c1t + ct]
    reentry = [t for t in a2t + c1t + ct if t.get("reentry")]
    daily = []
    for name, e in (("A2", ea), ("C1", ec), ("COMBINED", ex)):
        for row in list(e.get("daily") or []):
            daily.append({"strategy": name, **row})
    d14 = []
    for name, e in (("A2", ea), ("C1", ec), ("COMBINED", ex)):
        for b, row in dict(e.get("blocks") or {}).items():
            d14.append({"strategy": name, "block": b, **row})
    ls = []
    for name, e in (("A2", ea), ("C1", ec), ("COMBINED", ex)):
        ls.append({"strategy": name, "side": "LONG", **dict(e.get("LONG") or {})})
        ls.append({"strategy": name, "side": "SHORT", **dict(e.get("SHORT") or {})})
    sr = []
    for name, e in (("A2", ea), ("C1", ec), ("COMBINED", ex)):
        sr.append({"strategy": name, "role": "SUPPORT", **dict(e.get("SUPPORT") or {})})
        sr.append({"strategy": name, "role": "RESISTANCE", **dict(e.get("RESISTANCE") or {})})
    mfe = [{"strategy": n, "mfe_mean": e.get("mfe_mean"), "mae_mean": e.get("mae_mean"), "attribution": e.get("attribution")} for n, e in (("A2", ea), ("C1", ec), ("COMBINED", ex))]
    gb = [{"strategy": n, "giveback_mean": e.get("giveback_mean")} for n, e in (("A2", ea), ("C1", ec), ("COMBINED", ex))]
    conc = []
    for name, e in (("A2", ea), ("C1", ec), ("COMBINED", ex)):
        conc.append({"strategy": name, **{k: v for k, v in dict(e.get("concentration") or {}).items() if not isinstance(v, (dict, list)) or k in {"top3_symbols", "top3_days"}}})
    tail = []
    for name, e in (("A2", ea), ("C1", ec), ("COMBINED", ex)):
        c = dict(e.get("concentration") or {})
        for k in ("top1_trade_removed", "top5_trades_removed", "top5pct_winners_removed", "best_day_removed", "top_symbol_removed"):
            tail.append({"strategy": name, "cut": k, **dict(c.get(k) or {})})
    ties = []
    for name, pack in (("A2", a2), ("C1", c1), ("COMBINED", comb)):
        t = dict(pack.get("tie") or {})
        ties.append({"strategy": name, "CAP_TIE_ORDER_FRAGILE": t.get("CAP_TIE_ORDER_FRAGILE"), "primary_stress_yen": t.get("primary_stress_yen")})
        for row in list(t.get("comparisons") or []):
            ties.append({"strategy": name, **row})
    return {
        "Binding": _kv_rows(
            {
                "analysis_id": report.get("analysis_id"),
                "parent_verdict": report.get("parent_verdict"),
                "bind_ok": (report.get("bind") or {}).get("ok"),
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
                "identity": report.get("identity"),
            }
        ),
        "Strategy_Definitions": _kv_rows(report.get("strategy_spec") or {}),
        "A2_Signals": _rows(a2s, SIG_COLS),
        "C1_Signals": _rows(c1s, SIG_COLS),
        "Entry_Lineage": lineage,
        "Structural_Targets": targets,
        "Exit_Lineage": exits,
        "A2_Trades": _rows(a2t, TRADE_COLS),
        "C1_Trades": _rows(c1t, TRADE_COLS),
        "Combined_Trades": _rows(ct, TRADE_COLS),
        "CAP_Blocked": _rows(cap_b, ("signal_id", "mechanism", "symbol", "date", "entry_time", "block_reason", "strategy_id")),
        "SameSymbol_Blocked": _rows(sym_b, ("signal_id", "mechanism", "symbol", "date", "entry_time", "block_reason")),
        "Occupancy": occ,
        "Slot_Release": slot,
        "Reentry": _rows(reentry, ("signal_id", "strategy_id", "symbol", "date", "entry_time", "reentry")),
        "Target_Invalidation_Ambiguity": _rows(amb, SIG_COLS),
        "A1_Passive_Diagnostic": list(report.get("_a1_rows") or [])[:5000] or _kv_rows(report.get("A1") or {}),
        "Cost_Stress": [
            {"strategy": "A2", "X0_GROSS": ea.get("gross_yen"), "X1_8BPS_EXECUTION_STRESS": ea.get("stress_yen"), "BREAK_EVEN_COST_BPS": ea.get("break_even_cost_bps")},
            {"strategy": "C1", "X0_GROSS": ec.get("gross_yen"), "X1_8BPS_EXECUTION_STRESS": ec.get("stress_yen"), "BREAK_EVEN_COST_BPS": ec.get("break_even_cost_bps")},
            {"strategy": "COMBINED", "X0_GROSS": ex.get("gross_yen"), "X1_8BPS_EXECUTION_STRESS": ex.get("stress_yen"), "BREAK_EVEN_COST_BPS": ex.get("break_even_cost_bps")},
        ],
        "Daily": daily,
        "D1_D4": d14,
        "Long_Short": ls,
        "Support_Resistance": sr,
        "MFE_MAE": mfe,
        "Giveback": gb,
        "Concentration": conc,
        "Tail_Robustness": tail,
        "Tie_Order_Sensitivity": ties,
        "Freeze_Candidate": _kv_rows(report.get("freeze_candidate") or {}),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# SUPPORT_RESISTANCE_MECHANISM_TO_COMPLETE_STRATEGY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "A2 is the stronger Discovery mechanism. C1 is supported with Holm p=0.07 caveat.",
            "Evaluation unit is ENTRY + EXECUTION + EXIT + CAP + occupancy, not ENTRY alone.",
            "",
            f"A2 signal_n? **{a.get('A2 signal_n?')}** trade_n? **{a.get('A2 trade_n?')}**",
            f"A2 gross? **{a.get('A2 gross PnL?')}** 8bps-stress? **{a.get('A2 8bps-stress PnL?')}**",
            f"A2 PF gross/stress? **{a.get('A2 PF gross/stress?')}** median? **{a.get('A2 median trade?')}** maxDD? **{a.get('A2 maxDD?')}** pos-day? **{a.get('A2 positive-day rate?')}**",
            "",
            f"C1 signal_n? **{a.get('C1 signal_n?')}** trade_n? **{a.get('C1 trade_n?')}**",
            f"C1 gross? **{a.get('C1 gross PnL?')}** 8bps-stress? **{a.get('C1 8bps-stress PnL?')}**",
            f"C1 PF gross/stress? **{a.get('C1 PF gross/stress?')}** median? **{a.get('C1 median trade?')}** maxDD? **{a.get('C1 maxDD?')}** pos-day? **{a.get('C1 positive-day rate?')}**",
            "",
            f"COMBINED signal_n? **{a.get('COMBINED signal_n?')}** trade_n? **{a.get('COMBINED trade_n?')}**",
            f"COMBINED gross? **{a.get('COMBINED gross PnL?')}** 8bps-stress? **{a.get('COMBINED 8bps-stress PnL?')}**",
            f"COMBINED PF? **{a.get('COMBINED PF gross/stress?')}**",
            "",
            f"target / invalidation / session-close? **{a.get('How many target exits?')}** / **{a.get('How many invalidation exits?')}** / **{a.get('How many session-close exits?')}**",
            f"target future leakage? **{a.get('Any target future leakage?')}** same-bar entry? **{a.get('Any same-bar entry?')}**",
            f"delayed CAP? **{a.get('Any delayed CAP entry?')}** delayed same-symbol? **{a.get('Any delayed same-symbol entry?')}**",
            f"CAP tie-order fragile? **{a.get('CAP tie-order fragile?')}**",
            f"A2 candidate? **{a.get('Does A2 remain a Complete Strategy candidate?')}** C1? **{a.get('Does C1?')}** combined value? **{a.get('Does combined add genuine portfolio value?')}**",
            f"PnL-based rule change? **{a.get('Any PnL-based rule change?')}**",
            f"Old Confirmation opened? **{a.get('Old Confirmation opened?')}** Frozen Validation opened? **{a.get('Frozen Validation opened?')}**",
            f"Kabu50 applied? **{a.get('Kabu50 applied?')}** submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    if isinstance(payload.get("A2"), dict):
        payload["A2"] = {k: v for k, v in payload["A2"].items() if k != "_trades"}
    if isinstance(payload.get("C1"), dict):
        payload["C1"] = {k: v for k, v in payload["C1"].items() if k != "_trades"}
    if isinstance(payload.get("COMBINED"), dict):
        payload["COMBINED"] = {k: v for k, v in payload["COMBINED"].items() if k != "_trades"}
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
