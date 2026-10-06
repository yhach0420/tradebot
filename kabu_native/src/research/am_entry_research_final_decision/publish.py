"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize
from research.am_entry_research_final_decision import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_entry_research_final_decision"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "required",
    "c0_spec",
    "closed_lines",
    "evidence",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    lines = [
        "# AM ENTRY RESEARCH FINAL DECISION V1",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "Development close. C0 frozen as prospective challenger only.",
        "Not a DEVELOPMENT_PASS. Not a RUNTIME_CANDIDATE. Not a CERTIFIED_POLICY.",
        "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "",
        "## REQUIRED FINAL OUTPUT",
        "",
        f"AM_ENTRY_FULL_GATE_POLICY_FOUND: {_tf(g.get('AM_ENTRY_FULL_GATE_POLICY_FOUND'))}",
        "",
        f"AM_ENTRY_DEVELOPMENT_ECONOMIC_EDGE_FOUND: {_tf(g.get('AM_ENTRY_DEVELOPMENT_ECONOMIC_EDGE_FOUND'))}",
        "",
        f"BEST_DEVELOPMENT_ARCHITECTURE: {_fmt(g.get('BEST_DEVELOPMENT_ARCHITECTURE'))}",
        "",
        f"C0_NET: {_fmt(g.get('C0_NET'))}",
        "",
        f"C0_PF: {_fmt(g.get('C0_PF'))}",
        "",
        f"C0_MAX_DD: {_fmt(g.get('C0_MAX_DD'))}",
        "",
        f"C0_PAIRED_MEDIAN: {_fmt(g.get('C0_PAIRED_MEDIAN'))}",
        "",
        f"C0_POS_DAYS: {_fmt(g.get('C0_POS_DAYS'))}",
        "",
        f"C0_NEG_DAYS: {_fmt(g.get('C0_NEG_DAYS'))}",
        "",
        f"C0_EX_BEST: {_fmt(g.get('C0_EX_BEST'))}",
        "",
        f"C0_EX_TOP3: {_fmt(g.get('C0_EX_TOP3'))}",
        "",
        f"C0_PROSPECTIVE_SPEC_SHA256: {_fmt(g.get('C0_PROSPECTIVE_SPEC_SHA256'))}",
        "",
        f"AM_EXISTING_18D_ARCHITECTURE_SEARCH_CLOSED: {_tf(g.get('AM_EXISTING_18D_ARCHITECTURE_SEARCH_CLOSED'))}",
        "",
        f"RUNTIME_ADOPTION_ALLOWED: {_tf(g.get('RUNTIME_ADOPTION_ALLOWED'))}",
        "",
        f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
        "",
        f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
        "",
        f"VERDICT: {g.get('VERDICT')}",
        "",
        f"NEXT: {g.get('NEXT')}",
        "",
        "## CLOSED RESEARCH LINES",
        "",
    ]
    for line in list(report.get("closed_lines") or []):
        lines.append(str(line))
        lines.append("")
    lines.extend(
        [
            "## STOP",
            "",
            "Existing 18-day AM ENTRY architecture search is closed.",
            "C0 is FROZEN_FOR_FUTURE_OOS_ONLY. No Runtime change.",
            "submit/cancel/live=0/0/0.",
            "",
        ]
    )
    return "\n".join(lines)
