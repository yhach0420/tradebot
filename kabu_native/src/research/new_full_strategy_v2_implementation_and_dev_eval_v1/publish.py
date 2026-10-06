"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.new_full_strategy_v2_implementation_and_dev_eval_v1 import ANALYSIS_ID
from research.new_full_strategy_v2_implementation_and_dev_eval_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "integrity",
    "coverage",
    "economics",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT') or a.get('51_VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT') or a.get('52_NEXT'))}",
        "",
        f"FULL_STRATEGY_SPEC_SHA256_V4: {_fmt(d.get('FULL_STRATEGY_SPEC_SHA256_V4'))}",
        "",
        f"V4_HASH_UNCHANGED: {_fmt(d.get('V4_HASH_UNCHANGED'))}",
        "",
        f"ECONOMICS_VISIBLE_BEFORE_INTEGRITY_PASS: {_fmt(d.get('ECONOMICS_VISIBLE_BEFORE_INTEGRITY_PASS'))}",
        "",
        f"BASE_ECONOMIC_RUN_N: {_fmt(d.get('BASE_ECONOMIC_RUN_N'))}",
        "",
        f"G6_CAUSAL_RERUN_N: {_fmt(d.get('G6_CAUSAL_RERUN_N'))}",
        "",
        f"FAILED_STAGE: {_fmt(d.get('FAILED_STAGE'))}",
        "",
        "ANOTHER_PRECOMMIT_RUN: false",
        "",
        "POST_RESULT_RETUNE: false",
        "",
        "## Answers",
        "",
    ]
    for k, v in a.items():
        if v is None:
            continue
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    dumped = json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n"
    (OUT / "report.json").write_text(dumped, encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
