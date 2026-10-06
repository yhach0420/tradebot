"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.wait5_execution_aware_rebase import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "wait5_execution_aware_rebase"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Capacity",
    "Sessions",
    "YFill5",
    "FillAligned",
    "Pareto",
    "CurrentFill",
    "CurrentQuality",
    "Delay",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# WAIT5 EXECUTION-AWARE ENTRY REBASE V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "Development WAIT=5s. Runtime WAIT_SEC remains 1.0.",
            "Corrected Passive Fill frozen. Standalone. No POSITION_CAP / occupancy / reentry.",
            "No model. No feature search. Old t0 joint is diagnostic only.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"W5_ANY_RATE: {_fmt(g.get('W5_ANY_RATE'))}",
            f"W5_MULTI_RATE: {_fmt(g.get('W5_MULTI_RATE'))}",
            f"W5_THREEPLUS_RATE: {_fmt(g.get('W5_THREEPLUS_RATE'))}",
            "",
            f"AM_ANY_RATE: {_fmt(g.get('AM_ANY_RATE'))}",
            f"AM_MULTI_RATE: {_fmt(g.get('AM_MULTI_RATE'))}",
            f"AM_THREEPLUS_RATE: {_fmt(g.get('AM_THREEPLUS_RATE'))}",
            "",
            f"PM_ANY_RATE: {_fmt(g.get('PM_ANY_RATE'))}",
            f"PM_MULTI_RATE: {_fmt(g.get('PM_MULTI_RATE'))}",
            f"PM_THREEPLUS_RATE: {_fmt(g.get('PM_THREEPLUS_RATE'))}",
            "",
            f"Y_FILL5_POS_N: {g.get('Y_FILL5_POS_N')}",
            "",
            f"CURRENT_TOP3_FILL5_RATE: {_fmt(g.get('CURRENT_TOP3_FILL5_RATE'))}",
            f"CURRENT_FILL5_ENRICHMENT: {_fmt(g.get('CURRENT_FILL5_ENRICHMENT'))}",
            "",
            f"CURRENT_SCORE_U_FILL_SPEARMAN: {_fmt(g.get('CURRENT_SCORE_U_FILL_SPEARMAN'))}",
            f"CURRENT_SCORE_D_FILL_SPEARMAN: {_fmt(g.get('CURRENT_SCORE_D_FILL_SPEARMAN'))}",
            "",
            f"AM_CURRENT_SCORE_U_FILL_SPEARMAN: {_fmt(g.get('AM_CURRENT_SCORE_U_FILL_SPEARMAN'))}",
            f"AM_CURRENT_SCORE_D_FILL_SPEARMAN: {_fmt(g.get('AM_CURRENT_SCORE_D_FILL_SPEARMAN'))}",
            "",
            f"PM_CURRENT_SCORE_U_FILL_SPEARMAN: {_fmt(g.get('PM_CURRENT_SCORE_U_FILL_SPEARMAN'))}",
            f"PM_CURRENT_SCORE_D_FILL_SPEARMAN: {_fmt(g.get('PM_CURRENT_SCORE_D_FILL_SPEARMAN'))}",
            "",
            f"FILLABLE_PARETO_SIZE_MEAN: {_fmt(g.get('FILLABLE_PARETO_SIZE_MEAN'))}",
            f"MULTI_FILLABLE_WITH_DOMINANCE_RATE: {_fmt(g.get('MULTI_FILLABLE_WITH_DOMINANCE_RATE'))}",
            f"THREEPLUS_WITH_DOMINANCE_RATE: {_fmt(g.get('THREEPLUS_WITH_DOMINANCE_RATE'))}",
            "",
            f"FILL_DELAY_U_SPEARMAN: {_fmt(g.get('FILL_DELAY_U_SPEARMAN'))}",
            f"FILL_DELAY_D_SPEARMAN: {_fmt(g.get('FILL_DELAY_D_SPEARMAN'))}",
            "",
            f"TARGET_ARCHITECTURE: {g.get('TARGET_ARCHITECTURE')}",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No ENTRY model. No W5 runtime adoption. Runtime WAIT_SEC remains 1.0.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
