"""Write report.json / report.md / audit.xlsx only. No CSV dump."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.anchor_timing_robustness import (
    ANALYSIS_ID,
    DOCUMENT_ID,
    SHIFT_KEYS,
    SHIFTS_MIN,
    TASK_LABEL,
)
from research.anchor_timing_robustness.grid import shift_validity
from research.anchor_timing_robustness.metrics import (
    classify_tod_status,
    interpret,
    maxdd,
    mean_finite,
    trade_stats,
)

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "anchor_timing_robustness"
JST = timezone(timedelta(hours=9))
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _mean_field(rows: list[dict[str, Any]], key: str, *, shift: str | None = None) -> Any:
    xs = []
    for r in rows:
        if not r.get("valid"):
            continue
        if shift is not None and str(r.get("shift_key")) != shift:
            continue
        v = r.get(key)
        if v is None:
            continue
        xs.append(v)
    return mean_finite(xs)


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            if v == float("inf"):
                v = "Infinity"
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _iso_trades(days: list[dict[str, Any]], shift_key: str) -> list[dict[str, Any]]:
    out = []
    for d in days:
        for t in (d.get("primary") or {}).get("isolated_trades") or []:
            if str(t.get("shift_key")) == shift_key and t.get("independent_filled"):
                out.append(t)
    return out


def _port_trades(days: list[dict[str, Any]], shift_key: str) -> list[dict[str, Any]]:
    out = []
    for d in days:
        pack = (d.get("portfolios") or {}).get(shift_key) or {}
        out.extend(pack.get("trades") or [])
    return out


def build_report(
    *,
    provenance: dict[str, Any],
    inventory: list[dict[str, Any]],
    days: list[dict[str, Any]],
    failed: list[str],
) -> dict[str, Any]:
    elig = [r for r in inventory if r.get("replay_eligible")]
    used = [d for d in days if d.get("ok")]
    used_dates = [str(d["date"]) for d in used]
    comparisons: list[dict[str, Any]] = []
    for d in used:
        for row in (d.get("primary") or {}).get("comparisons") or []:
            slim = {k: v for k, v in row.items() if k not in {"ranked_ref", "ranked_shift"}}
            comparisons.append(slim)

    spearman_pm = {k: _mean_field(comparisons, "spearman", shift=k) for k in ("M5", "P5", "M10", "P10")}
    top5_pm = {k: _mean_field(comparisons, "top5_overlap", shift=k) for k in ("M5", "P5", "M10", "P10")}
    entry_pm = {k: _mean_field(comparisons, "entry_jaccard", shift=k) for k in ("M5", "P5", "M10", "P10")}

    tod_status = classify_tod_status(comparisons)
    isolated_stats = {
        k: trade_stats(_iso_trades(used, k)) for k in ("REFERENCE", "M5", "P5", "M10", "P10")
    }
    portfolio_stats = {}
    for k in ("REFERENCE", "M5", "P5", "M10", "P10"):
        tr = _port_trades(used, k)
        st = trade_stats(tr)
        st["maxDD"] = maxdd(tr)
        st["cap_blocked"] = sum(int(((d.get("portfolios") or {}).get(k) or {}).get("cap_blocked") or 0) for d in used)
        st["same_symbol_blocked"] = sum(
            int(((d.get("portfolios") or {}).get(k) or {}).get("same_symbol_blocked") or 0) for d in used
        )
        st["reentry_extra_fills"] = sum(
            int(((d.get("portfolios") or {}).get(k) or {}).get("reentry_extra_fills") or 0) for d in used
        )
        st["anchor_fires"] = sum(
            int(((d.get("portfolios") or {}).get(k) or {}).get("anchor_fires") or 0) for d in used
        )
        portfolio_stats[k] = st

    interp = interpret(
        spearman_pm=spearman_pm,
        top5_pm=top5_pm,
        entry_pm=entry_pm,
        isolated_stats=isolated_stats,
        tod_status=tod_status,
    )

    def _split_stats(period: str) -> dict[str, Any]:
        sub = [d for d in used if str(d.get("period")) == period]
        cmp_ = []
        for d in sub:
            for row in (d.get("primary") or {}).get("comparisons") or []:
                if row.get("valid"):
                    cmp_.append(row)
        iso = {k: trade_stats(_iso_trades(sub, k)) for k in ("REFERENCE", "M5", "P5", "M10", "P10")}
        return {
            "days": [str(d["date"]) for d in sub],
            "n_days": len(sub),
            "spearman_M5": _mean_field(cmp_, "spearman", shift="M5"),
            "spearman_P5": _mean_field(cmp_, "spearman", shift="P5"),
            "top5_M5": _mean_field(cmp_, "top5_overlap", shift="M5"),
            "top5_P5": _mean_field(cmp_, "top5_overlap", shift="P5"),
            "entry_M5": _mean_field(cmp_, "entry_jaccard", shift="M5"),
            "entry_P5": _mean_field(cmp_, "entry_jaccard", shift="P5"),
            "isolated": iso,
        }

    leak_days = [str(d["date"]) for d in used if d.get("snapshot_future_leak")]
    validity = shift_validity()

    ref = isolated_stats["REFERENCE"]
    report = {
        "ANALYSIS_ID": ANALYSIS_ID,
        "DOCUMENT_ID": DOCUMENT_ID,
        "label": TASK_LABEL,
        "generated_at_jst": datetime.now(JST).isoformat(timespec="seconds"),
        "runtime_changed": False,
        "strategy_changed": False,
        "paper_opval_changed": False,
        "new_grid_created": False,
        "best_shift_adopted": False,
        "safety": {"submit": 0, "cancel": 0, "live": 0, "offline_research_only": True},
        "CANONICAL_CLOCK_GRID": provenance.get("CANONICAL_CLOCK_GRID"),
        "CANONICAL_AM_ANCHORS": provenance.get("CANONICAL_AM_ANCHORS"),
        "CANONICAL_PM_ANCHORS": provenance.get("CANONICAL_PM_ANCHORS"),
        "SOURCE_FILE": provenance.get("SOURCE_FILE"),
        "SOURCE_SHA": provenance.get("SOURCE_SHA"),
        "CLOCK_GRID_PROVENANCE": provenance,
        "DATA_DAYS": {
            "inventoried": len(inventory),
            "eligible": len(elig),
            "completed": len(used),
            "failed": failed,
            "used": used_dates,
            "development": [r["date"] for r in elig if r.get("period") == "DEVELOPMENT"],
            "post_freeze_holdout": [r["date"] for r in elig if r.get("period") == "POST_FREEZE_HOLDOUT"],
            "excluded": [
                {"date": r["date"], "reason": r.get("exclusion_reason"), "class": r.get("capture_class")}
                for r in inventory
                if not r.get("replay_eligible")
            ],
        },
        "TRUE_HOLDOUT_AVAILABLE": bool(provenance.get("TRUE_HOLDOUT_AVAILABLE"))
        and any(r.get("period") == "POST_FREEZE_HOLDOUT" and r.get("replay_eligible") for r in inventory),
        "tod_status": interp.get("tod_status"),
        **{k: interp[k] for k in interp if k != "tod_status"},
        "REFERENCE_TRADES": ref.get("trades"),
        "REFERENCE_PNL": ref.get("pnl"),
        "REFERENCE_PF": ref.get("PF"),
        "M5_TRADES": isolated_stats["M5"].get("trades"),
        "M5_PNL": isolated_stats["M5"].get("pnl"),
        "M5_PF": isolated_stats["M5"].get("PF"),
        "P5_TRADES": isolated_stats["P5"].get("trades"),
        "P5_PNL": isolated_stats["P5"].get("pnl"),
        "P5_PF": isolated_stats["P5"].get("PF"),
        "M10_TRADES": isolated_stats["M10"].get("trades"),
        "M10_PNL": isolated_stats["M10"].get("pnl"),
        "M10_PF": isolated_stats["M10"].get("PF"),
        "P10_TRADES": isolated_stats["P10"].get("trades"),
        "P10_PNL": isolated_stats["P10"].get("pnl"),
        "P10_PF": isolated_stats["P10"].get("PF"),
        "ECONOMICS_BASIS": "ISOLATED_ANCHOR_NO_PORTFOLIO",
        "ISOLATED_ECONOMICS": isolated_stats,
        "PORTFOLIO_ECONOMICS": portfolio_stats,
        "PERIOD_SPLIT": {
            "DEVELOPMENT": _split_stats("DEVELOPMENT"),
            "POST_FREEZE_HOLDOUT": _split_stats("POST_FREEZE_HOLDOUT"),
        },
        "SHIFT_VALIDITY": validity,
        "LEAKAGE": {
            "snapshot_future_leak_days": leak_days,
            "any": bool(leak_days),
        },
        "inventory": [
            {k: v for k, v in r.items() if k != "universe_symbols"}
            for r in inventory
        ],
        "daily": [
            {
                "date": d.get("date"),
                "period": d.get("period"),
                "elapsed_sec": d.get("elapsed_sec"),
                "shift_errors": d.get("shift_errors"),
                "leak": d.get("snapshot_future_leak"),
                "portfolio": {
                    k: {
                        "trades": (p or {}).get("trade_n"),
                        "pnl": (p or {}).get("pnl"),
                        "PF": (p or {}).get("PF"),
                        "fires": (p or {}).get("anchor_fires"),
                        "cap_blocked": (p or {}).get("cap_blocked"),
                        "same_symbol_blocked": (p or {}).get("same_symbol_blocked"),
                    }
                    for k, p in (d.get("portfolios") or {}).items()
                },
            }
            for d in used
        ],
        "comparisons": comparisons,
        "regular_grid": [row for d in used for row in ((d.get("primary") or {}).get("regular") or [])],
        "nearest_regular10": [row for d in used for row in ((d.get("primary") or {}).get("nearest_regular10") or [])],
        "isolated_trades": [t for d in used for t in ((d.get("primary") or {}).get("isolated_trades") or [])],
        "portfolio_trades": [
            {**t, "shift_key": k}
            for d in used
            for k, p in (d.get("portfolios") or {}).items()
            for t in (p or {}).get("trades") or []
        ],
        "entry_rank_cross": [
            {
                "date": r.get("date"),
                "anchor": r.get("anchor"),
                "shift_key": r.get("shift_key"),
                "direction": "ref_entry_at_shift",
                **item,
            }
            for r in comparisons
            if r.get("valid")
            for item in (r.get("ref_entry_rank_at_shift") or [])
        ]
        + [
            {
                "date": r.get("date"),
                "anchor": r.get("anchor"),
                "shift_key": r.get("shift_key"),
                "direction": "shift_entry_at_ref",
                **item,
            }
            for r in comparisons
            if r.get("valid")
            for item in (r.get("shift_entry_rank_at_ref") or [])
        ],
        "note": (
            "Headline TRADES/PNL/PF are isolated (no cap / same-symbol / re-entry). "
            "Portfolio figures are secondary and must not be used to pick a new CLOCK_GRID. "
            "P3-0 PnL-beats-all-shifts is not reused as this verdict."
        ),
    }
    return json_sanitize(report)


def write_md(report: dict[str, Any]) -> str:
    prov = report.get("CLOCK_GRID_PROVENANCE") or {}
    data = report.get("DATA_DAYS") or {}
    iso = report.get("ISOLATED_ECONOMICS") or {}
    port = report.get("PORTFOLIO_ECONOMICS") or {}
    lines = [
        "# TRADEBOT — Anchor Timing Robustness",
        "",
        "Offline causal study. Runtime / CLOCK_GRID / ENTRY / EXIT / Paper / OPVAL unchanged.",
        "No new grid. No best-shift adoption. Headline economics = isolated anchors (no portfolio).",
        "",
        "## CANONICAL_CLOCK_GRID",
        f"CANONICAL_AM_ANCHORS: {report.get('CANONICAL_AM_ANCHORS')}",
        f"CANONICAL_PM_ANCHORS: {report.get('CANONICAL_PM_ANCHORS')}",
        f"SOURCE_FILE: {report.get('SOURCE_FILE')}",
        f"SOURCE_SHA: {report.get('SOURCE_SHA')}",
        "",
        "## CLOCK_GRID_PROVENANCE",
        f"introduced_as: {prov.get('introduced_as')}",
        f"x32_precommitted_not_outcome_derived: {prov.get('x32_precommitted_not_outcome_derived')}",
        f"freeze_matches_runtime: {prov.get('freeze_matches_runtime')}",
        f"x32_matches_runtime: {prov.get('x32_matches_runtime')}",
        f"DEVELOPMENT_PERIOD: {prov.get('DEVELOPMENT_PERIOD')}",
        f"POST_FREEZE_FROM: {prov.get('POST_FREEZE_FROM')}",
        f"pnl_selected_clock_evidence: {prov.get('pnl_selected_clock_evidence')}",
        "",
        "## DATA_DAYS",
        f"inventoried={data.get('inventoried')} eligible={data.get('eligible')} completed={data.get('completed')}",
        f"used: {data.get('used')}",
        f"development: {data.get('development')}",
        f"post_freeze_holdout: {data.get('post_freeze_holdout')}",
        f"failed: {data.get('failed')}",
        "excluded:",
    ]
    for ex in data.get("excluded") or []:
        lines.append(f"- {ex.get('date')}: {ex.get('class')} {ex.get('reason')}")
    lines += [
        "",
        f"TRUE_HOLDOUT_AVAILABLE: {report.get('TRUE_HOLDOUT_AVAILABLE')}",
        "",
        "## CROSS-SECTION (mean over valid day×anchor pairs)",
        f"RANK_SPEARMAN_M5: {report.get('RANK_SPEARMAN_M5')}",
        f"RANK_SPEARMAN_P5: {report.get('RANK_SPEARMAN_P5')}",
        f"RANK_SPEARMAN_M10: {report.get('RANK_SPEARMAN_M10')}",
        f"RANK_SPEARMAN_P10: {report.get('RANK_SPEARMAN_P10')}",
        f"TOP5_OVERLAP_M5: {report.get('TOP5_OVERLAP_M5')}",
        f"TOP5_OVERLAP_P5: {report.get('TOP5_OVERLAP_P5')}",
        f"TOP5_OVERLAP_M10: {report.get('TOP5_OVERLAP_M10')}",
        f"TOP5_OVERLAP_P10: {report.get('TOP5_OVERLAP_P10')}",
        f"ENTRY_OVERLAP_M5: {report.get('ENTRY_OVERLAP_M5')}",
        f"ENTRY_OVERLAP_P5: {report.get('ENTRY_OVERLAP_P5')}",
        f"ENTRY_OVERLAP_M10: {report.get('ENTRY_OVERLAP_M10')}",
        f"ENTRY_OVERLAP_P10: {report.get('ENTRY_OVERLAP_P10')}",
        "",
        "## ISOLATED ANCHOR ECONOMICS (FILL + Early Guard + 600/750, no portfolio)",
        f"REFERENCE_TRADES: {report.get('REFERENCE_TRADES')}",
        f"REFERENCE_PNL: {report.get('REFERENCE_PNL')}",
        f"REFERENCE_PF: {report.get('REFERENCE_PF')}",
        f"M5_TRADES: {report.get('M5_TRADES')}",
        f"M5_PNL: {report.get('M5_PNL')}",
        f"M5_PF: {report.get('M5_PF')}",
        f"P5_TRADES: {report.get('P5_TRADES')}",
        f"P5_PNL: {report.get('P5_PNL')}",
        f"P5_PF: {report.get('P5_PF')}",
        f"M10_TRADES: {report.get('M10_TRADES')}",
        f"M10_PNL: {report.get('M10_PNL')}",
        f"M10_PF: {report.get('M10_PF')}",
        f"P10_TRADES: {report.get('P10_TRADES')}",
        f"P10_PNL: {report.get('P10_PNL')}",
        f"P10_PF: {report.get('P10_PF')}",
        "",
        "Isolated extra:",
    ]
    for k in ("REFERENCE", "M5", "P5", "M10", "P10"):
        st = iso.get(k) or {}
        lines.append(
            f"- {k}: trades={st.get('trades')} pnl={st.get('pnl')} PF={st.get('PF')} "
            f"win_rate={st.get('win_rate')} avg={st.get('avg_trade')} median={st.get('median_trade')} "
            f"MFE={st.get('MFE')} MAE={st.get('MAE')}"
        )
    lines += ["", "## FULL-GRID PORTFOLIO (secondary; not used to pick a grid)"]
    for k in ("REFERENCE", "M5", "P5", "M10", "P10"):
        st = port.get(k) or {}
        lines.append(
            f"- {k}: trades={st.get('trades')} pnl={st.get('pnl')} PF={st.get('PF')} "
            f"maxDD={st.get('maxDD')} fires={st.get('anchor_fires')} cap_blocked={st.get('cap_blocked')} "
            f"same_symbol_blocked={st.get('same_symbol_blocked')} reentry_extra={st.get('reentry_extra_fills')}"
        )
    lines += [
        "",
        "## TIME-OF-DAY",
        f"OPEN_EARLY_RESULT: {report.get('OPEN_EARLY_RESULT')}",
        f"NORMAL_SESSION_RESULT: {report.get('NORMAL_SESSION_RESULT')}",
        f"PM_OPEN_RESULT: {report.get('PM_OPEN_RESULT')}",
        "",
        "## VERDICT",
        f"ECONOMICS_STABLE: {report.get('ECONOMICS_STABLE')}",
        f"SELECTION_STABLE: {report.get('SELECTION_STABLE')}",
        f"EXACT_TIME_DEPENDENCE: {report.get('EXACT_TIME_DEPENDENCE')}",
        f"OVERFIT_CONCERN: {report.get('OVERFIT_CONCERN')}",
        f"verdict: {report.get('verdict')}",
        "",
        f"thresholds: {json.dumps(report.get('thresholds'), ensure_ascii=False)}",
        f"economics_notes: {report.get('economics_notes')}",
        f"leakage: {report.get('LEAKAGE')}",
        "",
        "STOP. Do not change CLOCK_GRID from this result.",
        "submit/cancel/live=0/0/0",
    ]
    return "\n".join(lines) + "\n"


def write_xlsx(report: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {
        "README": [
            {
                "analysis": ANALYSIS_ID,
                "verdict": report.get("verdict"),
                "economics_basis": report.get("ECONOMICS_BASIS"),
                "safety": "submit/cancel/live=0/0/0",
                "note": report.get("note"),
            }
        ],
        "CANONICAL_CLOCK_GRID": [
            {"role": "AM", "anchors": json.dumps(report.get("CANONICAL_AM_ANCHORS"))},
            {"role": "PM", "anchors": json.dumps(report.get("CANONICAL_PM_ANCHORS"))},
            {"role": "SOURCE_FILE", "anchors": report.get("SOURCE_FILE")},
            {"role": "SOURCE_SHA", "anchors": report.get("SOURCE_SHA")},
        ],
        "PROVENANCE": [{"key": k, "value": json.dumps(v, ensure_ascii=False, default=str) if isinstance(v, (dict, list)) else v} for k, v in (report.get("CLOCK_GRID_PROVENANCE") or {}).items()],
        "DATA_DAYS": report.get("inventory") or [{"status": "empty"}],
        "EXCLUDED": (report.get("DATA_DAYS") or {}).get("excluded") or [{"status": "none"}],
        "SHIFT_VALIDITY": report.get("SHIFT_VALIDITY") or [{"status": "empty"}],
        "RANK_COMPARE": report.get("comparisons") or [{"status": "empty"}],
        "ENTRY_RANK_CROSS": report.get("entry_rank_cross") or [{"status": "empty"}],
        "ISOLATED_TRADES": report.get("isolated_trades") or [{"status": "empty"}],
        "ISOLATED_ECONOMICS": [{"shift_key": k, **v} for k, v in (report.get("ISOLATED_ECONOMICS") or {}).items()] or [{"status": "empty"}],
        "PORTFOLIO_TRADES": report.get("portfolio_trades") or [{"status": "empty"}],
        "PORTFOLIO_ECONOMICS": [{"shift_key": k, **v} for k, v in (report.get("PORTFOLIO_ECONOMICS") or {}).items()] or [{"status": "empty"}],
        "DAILY_PORTFOLIO": report.get("daily") or [{"status": "empty"}],
        "TOD_VERDICT": [{"bucket": k, "status": v} for k, v in ((report.get("tod_status") or {}) | {"OPEN_EARLY": report.get("OPEN_EARLY_RESULT"), "NORMAL_SESSION": report.get("NORMAL_SESSION_RESULT"), "PM_OPEN": report.get("PM_OPEN_RESULT")}).items()],
        "REGULAR_GRID": report.get("regular_grid") or [{"status": "empty"}],
        "NEAREST_REGULAR10": report.get("nearest_regular10") or [{"status": "empty"}],
        "PERIOD_SPLIT": [
            {"period": k, **{kk: json.dumps(vv, ensure_ascii=False, default=str) if isinstance(vv, (dict, list)) else vv for kk, vv in v.items()}}
            for k, v in (report.get("PERIOD_SPLIT") or {}).items()
        ],
        "PRECOMMITTED_RULES": [{"key": k, "value": v} for k, v in (report.get("thresholds") or {}).items()],
        "VERDICT": [
            {
                "verdict": report.get("verdict"),
                "SELECTION_STABLE": report.get("SELECTION_STABLE"),
                "ECONOMICS_STABLE": report.get("ECONOMICS_STABLE"),
                "EXACT_TIME_DEPENDENCE": report.get("EXACT_TIME_DEPENDENCE"),
                "OVERFIT_CONCERN": report.get("OVERFIT_CONCERN"),
                "TRUE_HOLDOUT_AVAILABLE": report.get("TRUE_HOLDOUT_AVAILABLE"),
            }
        ],
        "LEAKAGE": [report.get("LEAKAGE") or {"any": False}],
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(str(name)[:31])
        _sheet(ws, rows if isinstance(rows, list) else [rows])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_artifacts(report: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    jp = OUT / "report.json"
    mp = OUT / "report.md"
    xp = OUT / "audit.xlsx"
    # keep json smaller: drop bulky trade arrays from the json copy written to disk? User asked report.json as deliverable with required fields. Keep trades in xlsx; json keeps summary + comparisons without isolated/portfolio full dumps if huge.
    json_body = dict(report)
    json_body.pop("isolated_trades", None)
    json_body.pop("portfolio_trades", None)
    json_body.pop("entry_rank_cross", None)
    jp.write_text(json.dumps(json_sanitize(json_body), ensure_ascii=False, indent=2), encoding="utf-8")
    mp.write_text(write_md(report), encoding="utf-8")
    write_xlsx(report, xp)
    return {"report_json": str(jp), "report_md": str(mp), "audit_xlsx": str(xp)}
