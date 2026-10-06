"""Write report.json / report.md / audit.xlsx only under v1/."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize
from research.simple_tech_entry_family import ANALYSIS_ID, STRATEGY_ID
from research.simple_tech_entry_family.isolation import V1_OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Precommit",
    "Data_Manifest",
    "Bar_Integrity",
    "Opportunity_Funnel",
    "Stage_Contribution",
    "Signal_State",
    "Winner_Path",
    "Loser_Path",
    "Failure_Taxonomy",
    "Trend_RCA",
    "RCI_RCA",
    "Volume_RCA",
    "Trigger_RCA",
    "Board_RCA",
    "Execution_RCA",
    "Exit_Attribution",
    "Trades",
    "Daily",
    "Robustness",
    "Baseline",
    "Non_Interference",
    "Integrity",
)
REQUIRED_KEYS = (
    "STRATEGY_ID",
    "RAW_OPPORTUNITY_N",
    "TREND_PASS_N",
    "PULLBACK_PASS_N",
    "RCI_PASS_N",
    "PRICE_ACTION_PASS_N",
    "VOLUME_PASS_N",
    "BOARD_PASS_N",
    "PENDING_N",
    "FILL_N",
    "TRADE_N",
    "NET",
    "PF",
    "MAX_DD",
    "POSITIVE_DAY_N",
    "NEGATIVE_DAY_N",
    "DAILY_MEDIAN",
    "EX_BEST",
    "EX_TOP3",
    "F1_TREND_FALSE_N",
    "F2_PULLBACK_FALSE_N",
    "F3_REVERSAL_FALSE_N",
    "F4_VOLUME_FALSE_CONFIRM_N",
    "F5_PRICE_TRIGGER_LATE_N",
    "F6_BOARD_FALSE_SUPPORT_N",
    "F7_EXECUTION_COST_N",
    "F8_EXIT_GIVEBACK_N",
    "PRIMARY_DEFICIENCY",
    "PRIMARY_DEFICIENCY_GROSS_LOSS",
    "PRIMARY_DEFICIENCY_DAY_COVERAGE",
    "V2_RECOMMENDED_COMPONENT",
    "TRUE_OOS",
    "NON_INTERFERENCE_PASS",
    "VERDICT",
    "NEXT",
)


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": d.get(k)} for k in d.keys()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    V1_OUT.mkdir(parents=True, exist_ok=True)
    extra = [p for p in V1_OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    for p in extra:
        p.unlink()
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (V1_OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (V1_OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(V1_OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    req = dict(report.get("required") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"STRATEGY_ID: `{STRATEGY_ID}`",
        f"SPEC_SHA256: `{req.get('SPEC_SHA256') or report.get('SPEC_SHA256')}`",
        f"VERDICT: **{req.get('VERDICT')}**",
        f"PRIMARY_DEFICIENCY: `{req.get('PRIMARY_DEFICIENCY')}`",
        f"V2_RECOMMENDED_COMPONENT: `{req.get('V2_RECOMMENDED_COMPONENT')}`",
        f"TRUE_OOS: `{req.get('TRUE_OOS')}`",
        "",
        "## Required output",
        "",
    ]
    for k in REQUIRED_KEYS:
        lines.append(f"{k}: {req.get(k)}")
    lines.extend(
        [
            "",
            "## Notes",
            "",
            "- V1 is a reference instrument. Family remains open. V2 is not implemented here.",
            "- Existing 18-day set is burned development data. Do not treat this as OOS.",
            "- Board is support/veto only. C14 EXIT is unchanged.",
            "- Runtime/Capture were not stopped, restarted, or written.",
            "",
            f"NEXT: {req.get('NEXT')}",
            "",
        ]
    )
    return "\n".join(lines) + "\n"
