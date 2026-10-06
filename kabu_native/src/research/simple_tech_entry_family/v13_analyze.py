"""V13 E4 identity, causal fill audit, numeric parity, reporting semantics. No ranking. No EXIT."""
from __future__ import annotations

import hashlib
import json
import os
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from research.simple_tech_entry_family.v12_analyze import policy_metrics
from research.simple_tech_entry_family.v12_harvest import improve_1tick_limit
from research.simple_tech_entry_family.v13_spec import (
    B1_EXECUTABLE_N_EXPECTED,
    DEVELOPMENT_CHALLENGER,
    E4_BID_EX_BEST_180_EXPECTED,
    E4_BID_EX_BEST_300_EXPECTED,
    E4_BID_NEG_DAY_180_EXPECTED,
    E4_BID_NEG_DAY_300_EXPECTED,
    E4_BID_POS_DAY_180_EXPECTED,
    E4_BID_POS_DAY_300_EXPECTED,
    E4_FILL_RATE_EXPECTED,
    E4_FILLED_GROSS_EXPECTED,
    E4_FILLED_N_EXPECTED,
    E4_INSIDE_COLLAPSE_N_EXPECTED,
    E4_UNCOND_BID_EXPECTED,
    E4_UNCOND_MID_EXPECTED,
    E4_UNFILLED_GROSS_EXPECTED,
    E4_UNFILLED_N_EXPECTED,
    E4_WAIT_BUDGET_SEC,
    FILL_EVIDENCE,
    PARITY_ABS_TOL,
    V12_CASE_EXPECTED,
    V12_SELECTED_POLICY,
    V12_SPEC_SHA256_EXPECTED,
    V12_VERDICT_EXPECTED,
)


def _finite(v: Any) -> bool:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return False
    return x == x


def _close(a: Any, b: Any, tol: float = PARITY_ABS_TOL) -> bool:
    try:
        return abs(float(a) - float(b)) <= float(tol)
    except (TypeError, ValueError):
        return False


def _canon_num(v: Any) -> Any:
    if isinstance(v, bool):
        return bool(v)
    if isinstance(v, int) and not isinstance(v, bool):
        return int(v)
    if _finite(v):
        return float(v)
    if v is None:
        return None
    return str(v)


def set_hash(rows: list[tuple[Any, ...]]) -> str:
    blob = json.dumps(rows, separators=(",", ":"), ensure_ascii=True, default=str).encode("utf-8")
    return hashlib.sha256(blob).hexdigest()


def _e4(row: dict[str, Any]) -> dict[str, Any]:
    return dict((row.get("exec") or {}).get(DEVELOPMENT_CHALLENGER) or {})


def identity_keys(row: dict[str, Any]) -> tuple[str, str, float]:
    return (str(row.get("date") or ""), str(row.get("symbol") or "").replace(".T", ""), float(row.get("t0") or 0.0))


def signal_tuples(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    return sorted(identity_keys(r) for r in rows)


def eligible_tuples(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    return sorted(identity_keys(r) for r in rows if r.get("executable_signal"))


def fill_tuples(exe: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    out = []
    for r in exe:
        ex = _e4(r)
        if not ex.get("filled"):
            continue
        d, s, t0 = identity_keys(r)
        out.append(
            (
                d,
                s,
                t0,
                _canon_num(ex.get("fill_t")),
                _canon_num(ex.get("fill_price")),
                _canon_num(ex.get("limit_price")),
                bool(ex.get("collapsed_to_bid") or r.get("inside_collapsed_to_bid")),
            )
        )
    return sorted(out)


def fill_rows(exe: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in exe:
        ex = _e4(r)
        if not ex.get("filled"):
            continue
        t0 = float(r["t0"])
        ft = ex.get("fill_t")
        wait = (float(ft) - t0) if _finite(ft) else None
        out.append(
            {
                "date": r.get("date"),
                "symbol": str(r.get("symbol") or "").replace(".T", ""),
                "signal_t0": t0,
                "Bid1_t0": r.get("bid0"),
                "Ask1_t0": r.get("ask0"),
                "limit_price": ex.get("limit_price"),
                "tick_size": r.get("tick"),
                "inside_collapsed": bool(ex.get("collapsed_to_bid") or r.get("inside_collapsed_to_bid")),
                "fill_time": ft,
                "wait_sec": wait if wait is not None else ex.get("waited_sec"),
                "fill_evidence_event": f"{FILL_EVIDENCE}:{ft}" if _finite(ft) else None,
                "fill_price": ex.get("fill_price"),
                "evidence": ex.get("evidence"),
            }
        )
    out.sort(key=lambda rec: (str(rec.get("date") or ""), str(rec.get("symbol") or ""), float(rec.get("signal_t0") or 0.0)))
    return out


def placement_ok(row: dict[str, Any]) -> tuple[bool, str]:
    bid = row.get("bid0")
    ask = row.get("ask0")
    tick = row.get("tick")
    if not _finite(bid) or not _finite(ask):
        return False, "NO_T0_QUOTES"
    lim_calc, collapsed, tick_calc = improve_1tick_limit(float(bid), float(ask))
    ex = _e4(row)
    stored = ex.get("limit_price")
    if not _finite(stored):
        return False, "NO_STORED_LIMIT"
    if abs(float(stored) - float(lim_calc)) > 1e-12:
        return False, "LIMIT_NE_RECOMPUTED"
    if _finite(tick) and abs(float(tick) - float(tick_calc)) > 1e-12:
        return False, "TICK_NE_RECOMPUTED"
    stored_c = bool(ex.get("collapsed_to_bid") or row.get("inside_collapsed_to_bid"))
    if bool(collapsed) != stored_c:
        return False, "COLLAPSE_MISMATCH"
    if float(stored) >= float(ask) - 1e-12:
        return False, "ASK_CROSS_AT_PLACEMENT"
    if abs(float(ex.get("wait_budget_sec") or 0.0) - float(E4_WAIT_BUDGET_SEC)) > 1e-12:
        return False, "WAIT_BUDGET_NE_5"
    return True, ""


def causal_audit(exe: list[dict[str, Any]]) -> dict[str, Any]:
    future = 0
    touch = 0
    queue = 0
    window_fail = 0
    evidence_fail = 0
    price_ne_limit = 0
    placement_fail = 0
    details: list[dict[str, Any]] = []
    for r in exe:
        ok_p, why = placement_ok(r)
        if not ok_p:
            placement_fail += 1
            details.append({"key": identity_keys(r), "placement": why})
        ex = _e4(r)
        if not ex.get("filled"):
            continue
        t0 = float(r["t0"])
        ft = ex.get("fill_t")
        if not _finite(ft):
            future += 1
            continue
        if float(ft) + 1e-12 < t0:
            future += 1
        if float(ft) > t0 + float(E4_WAIT_BUDGET_SEC) + 1e-12:
            window_fail += 1
        ev = str(ex.get("evidence") or "")
        if ev != FILL_EVIDENCE:
            evidence_fail += 1
            if "TOUCH" in ev:
                touch += 1
            if "QUEUE" in ev:
                queue += 1
        if _finite(ex.get("fill_price")) and _finite(ex.get("limit_price")):
            if abs(float(ex["fill_price"]) - float(ex["limit_price"])) > 1e-12:
                price_ne_limit += 1
    n_fill = sum(1 for r in exe if _e4(r).get("filled"))
    pass_ok = (
        n_fill == int(E4_FILLED_N_EXPECTED)
        and future == 0
        and touch == 0
        and queue == 0
        and window_fail == 0
        and evidence_fail == 0
        and price_ne_limit == 0
        and placement_fail == 0
    )
    return {
        "CAUSAL_FILL_AUDIT_PASS": bool(pass_ok),
        "FILLED_AUDITED_N": n_fill,
        "FUTURE_FILL_USE_N": future,
        "TOUCH_ONLY_FILL_N": touch,
        "QUEUE_ASSUMED_FILL_N": queue,
        "REPRICE_N": 0,
        "CHASE_N": 0,
        "WINDOW_FAIL_N": window_fail,
        "EVIDENCE_FAIL_N": evidence_fail,
        "FILL_PRICE_NE_LIMIT_N": price_ne_limit,
        "PLACEMENT_FAIL_N": placement_fail,
        "details": details[:20],
    }


def horizon_unfilled_audit(e4: dict[str, Any]) -> dict[str, Any]:
    uncond = e4.get("UNCONDITIONAL_POLICY_MARKOUT") or {}
    denom_ok = int(uncond.get("DENOMINATOR") or 0) == int(B1_EXECUTABLE_N_EXPECTED)
    unfilled_zero = bool(uncond.get("UNFILLED_AS_ZERO"))
    primary_kind = str(uncond.get("KIND") or "") == "FILL_TO_MID_SIGNAL_ANCHORED"
    ft = e4.get("FILL_TIME_MARKOUT_DIAG") or {}
    ft_not_primary = bool(ft.get("NOT_FOR_SELECTION"))
    filled_n = int(e4.get("FILLED_N") or 0)
    unfilled_n = int(e4.get("UNFILLED_N") or 0)
    n_ok = filled_n + unfilled_n == int(B1_EXECUTABLE_N_EXPECTED)
    return {
        "HORIZON_SEMANTICS_PASS": bool(denom_ok and primary_kind and ft_not_primary and n_ok),
        "UNFILLED_ZERO_PASS": bool(unfilled_zero and n_ok),
        "DENOMINATOR": uncond.get("DENOMINATOR"),
        "PRIMARY_KIND": uncond.get("KIND"),
        "FILL_TIME_NOT_FOR_SELECTION": ft_not_primary,
        "FILLED_PLUS_UNFILLED": filled_n + unfilled_n,
    }


def numeric_parity(e4: dict[str, Any]) -> dict[str, Any]:
    mid = e4.get("UNCONDITIONAL_POLICY_MARKOUT") or {}
    bid = e4.get("UNCONDITIONAL_FILL_TO_BID") or {}
    fg = e4.get("FILLED_GROSS_MID") or {}
    ug = e4.get("UNFILLED_GROSS_MID") or {}
    checks = {
        "FILLED_N": int(e4.get("FILLED_N") or -1) == int(E4_FILLED_N_EXPECTED),
        "UNFILLED_N": int(e4.get("UNFILLED_N") or -1) == int(E4_UNFILLED_N_EXPECTED),
        "ELIGIBLE_N": int(e4.get("ELIGIBLE_N") or -1) == int(B1_EXECUTABLE_N_EXPECTED),
        "FILL_RATE": _close(e4.get("FILL_RATE"), E4_FILL_RATE_EXPECTED),
        "INSIDE_COLLAPSE_N": int(e4.get("INSIDE_COLLAPSE_N") or -1) == int(E4_INSIDE_COLLAPSE_N_EXPECTED),
        "UNCOND_MID_60": _close(mid.get("60"), E4_UNCOND_MID_EXPECTED["60"]),
        "UNCOND_MID_180": _close(mid.get("180"), E4_UNCOND_MID_EXPECTED["180"]),
        "UNCOND_MID_300": _close(mid.get("300"), E4_UNCOND_MID_EXPECTED["300"]),
        "UNCOND_BID_60": _close(bid.get("60"), E4_UNCOND_BID_EXPECTED["60"]),
        "UNCOND_BID_180": _close(bid.get("180"), E4_UNCOND_BID_EXPECTED["180"]),
        "UNCOND_BID_300": _close(bid.get("300"), E4_UNCOND_BID_EXPECTED["300"]),
        "FILLED_GROSS_60": _close(fg.get("60"), E4_FILLED_GROSS_EXPECTED["60"]),
        "FILLED_GROSS_180": _close(fg.get("180"), E4_FILLED_GROSS_EXPECTED["180"]),
        "FILLED_GROSS_300": _close(fg.get("300"), E4_FILLED_GROSS_EXPECTED["300"]),
        "UNFILLED_GROSS_60": _close(ug.get("60"), E4_UNFILLED_GROSS_EXPECTED["60"]),
        "UNFILLED_GROSS_180": _close(ug.get("180"), E4_UNFILLED_GROSS_EXPECTED["180"]),
        "UNFILLED_GROSS_300": _close(ug.get("300"), E4_UNFILLED_GROSS_EXPECTED["300"]),
        "BID_POS_DAY_180": int(e4.get("BID_POS_DAY_N_180") or -1) == int(E4_BID_POS_DAY_180_EXPECTED),
        "BID_NEG_DAY_180": int(e4.get("BID_NEG_DAY_N_180") or -1) == int(E4_BID_NEG_DAY_180_EXPECTED),
        "BID_POS_DAY_300": int(e4.get("BID_POS_DAY_N_300") or -1) == int(E4_BID_POS_DAY_300_EXPECTED),
        "BID_NEG_DAY_300": int(e4.get("BID_NEG_DAY_N_300") or -1) == int(E4_BID_NEG_DAY_300_EXPECTED),
        "BID_EX_BEST_180": _close(e4.get("BID_EX_BEST_180"), E4_BID_EX_BEST_180_EXPECTED),
        "BID_EX_BEST_300": _close(e4.get("BID_EX_BEST_300"), E4_BID_EX_BEST_300_EXPECTED),
        "ADVERSE_FALSE": e4.get("EXTREME_ADVERSE_SELECTION") is False,
    }
    return {
        "E4_NUMERIC_PARITY": all(bool(v) for v in checks.values()),
        "checks": checks,
        "observed": {
            "FILLED_N": e4.get("FILLED_N"),
            "UNCOND_MID": mid,
            "UNCOND_BID": bid,
            "FILLED_GROSS": fg,
            "UNFILLED_GROSS": ug,
            "BID_DAYS_180": [e4.get("BID_POS_DAY_N_180"), e4.get("BID_NEG_DAY_N_180")],
            "BID_DAYS_300": [e4.get("BID_POS_DAY_N_300"), e4.get("BID_NEG_DAY_N_300")],
            "BID_EX_BEST": [e4.get("BID_EX_BEST_180"), e4.get("BID_EX_BEST_300")],
            "ADVERSE": e4.get("EXTREME_ADVERSE_SELECTION"),
        },
    }


def v12_official_preserved(v12_req: dict[str, Any], v12_decision: dict[str, Any]) -> dict[str, Any]:
    ok = (
        str(v12_req.get("VERDICT") or "") == V12_VERDICT_EXPECTED
        and str(v12_req.get("SELECTED_EXECUTION_POLICY") or "") == V12_SELECTED_POLICY
        and v12_req.get("ENTRY_EXECUTION_EDGE_REPAIRED") is False
        and str(v12_req.get("V12_SPEC_SHA256") or "") == V12_SPEC_SHA256_EXPECTED
        and str(v12_decision.get("CASE") or "") == V12_CASE_EXPECTED
    )
    return {
        "V12_OFFICIAL_VERDICT_PRESERVED": bool(ok),
        "V12_CASE": v12_decision.get("CASE"),
        "V12_VERDICT": v12_req.get("VERDICT"),
        "V12_SELECTED_EXECUTION_POLICY": v12_req.get("SELECTED_EXECUTION_POLICY"),
        "V12_ENTRY_EXECUTION_EDGE_REPAIRED": v12_req.get("ENTRY_EXECUTION_EDGE_REPAIRED"),
        "V12_SPEC_SHA256": v12_req.get("V12_SPEC_SHA256"),
    }


def _pid_in(snap: dict[str, Any], pid: Any) -> bool:
    if pid is None:
        return False
    try:
        want = int(pid)
    except (TypeError, ValueError):
        return False
    ids = []
    for key in ("LIVE_PIDS", "RESEARCH_PIDS"):
        for x in snap.get(key) or []:
            try:
                ids.append(int(x))
            except (TypeError, ValueError):
                continue
    if want in set(ids):
        return True
    return _process_exists(want)


def _process_exists(pid: int) -> bool:
    if os.name != "nt":
        return False
    try:
        import ctypes

        PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
        k = ctypes.windll.kernel32
        h = k.OpenProcess(PROCESS_QUERY_LIMITED_INFORMATION, False, int(pid))
        if h:
            k.CloseHandle(h)
            return True
        err = int(k.GetLastError() or 0)
        return err == 5
    except Exception:
        return False


def _ts_gt(a: Any, b: Any) -> bool:
    if a is None or b is None:
        return False
    try:
        return float(a) > float(b) + 1e-15
    except (TypeError, ValueError):
        pass
    sa, sb = str(a), str(b)
    if not sa or not sb:
        return False
    try:
        ta = datetime.fromisoformat(sa.replace("Z", "+00:00"))
        tb = datetime.fromisoformat(sb.replace("Z", "+00:00"))
        return ta > tb
    except Exception:
        return sa > sb


def reporting_semantics(pre: dict[str, Any], post: dict[str, Any]) -> dict[str, Any]:
    rt_pid = post.get("RUNTIME_PID")
    cap_pid = post.get("CAPTURE_PID")
    rt_alive = bool(rt_pid is not None and _pid_in(post, rt_pid))
    cap_alive = bool(cap_pid is not None and _pid_in(post, cap_pid))
    if rt_pid is None:
        rt_alive = False
    if cap_pid is None:
        cap_alive = False
    hb_adv = _ts_gt(post.get("RUNTIME_HEARTBEAT"), pre.get("RUNTIME_HEARTBEAT"))
    cap_adv = _ts_gt(post.get("CAPTURE_LAST_EVENT"), pre.get("CAPTURE_LAST_EVENT"))
    pack = {
        "RUNTIME_PID_UNCHANGED": pre.get("RUNTIME_PID") == post.get("RUNTIME_PID"),
        "CAPTURE_PID_UNCHANGED": pre.get("CAPTURE_PID") == post.get("CAPTURE_PID"),
        "RUNTIME_HEARTBEAT_ADVANCED": bool(hb_adv),
        "CAPTURE_ADVANCED": bool(cap_adv),
        "RUNTIME_STILL_ALIVE": bool(rt_alive),
        "CAPTURE_STILL_ALIVE": bool(cap_alive),
        "RUNTIME_PID": rt_pid,
        "CAPTURE_PID": cap_pid,
        "PRE_RUNTIME_HEARTBEAT": pre.get("RUNTIME_HEARTBEAT"),
        "POST_RUNTIME_HEARTBEAT": post.get("RUNTIME_HEARTBEAT"),
        "PRE_CAPTURE_LAST_EVENT": pre.get("CAPTURE_LAST_EVENT"),
        "POST_CAPTURE_LAST_EVENT": post.get("CAPTURE_LAST_EVENT"),
    }
    cons = True
    if pack["CAPTURE_PID"] is None and pack["CAPTURE_STILL_ALIVE"]:
        cons = False
    if pack["RUNTIME_PID"] is None and pack["RUNTIME_STILL_ALIVE"]:
        cons = False
    if pre.get("RUNTIME_HEARTBEAT") == post.get("RUNTIME_HEARTBEAT") and pack["RUNTIME_HEARTBEAT_ADVANCED"]:
        cons = False
    if pre.get("CAPTURE_LAST_EVENT") == post.get("CAPTURE_LAST_EVENT") and pack["CAPTURE_ADVANCED"]:
        cons = False
    pack["REPORTING_SEMANTICS_PASS"] = bool(cons)
    return pack


def read_v12_audit_e4(path: Any) -> dict[str, Any]:
    if not path or not path.is_file():
        return {"ok": False, "reason": "MISSING"}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Policies" not in wb.sheetnames:
            return {"ok": False, "reason": "NO_POLICIES_SHEET", "sheets": list(wb.sheetnames)}
        ws = wb["Policies"]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {"ok": False, "reason": "EMPTY"}
        headers = [str(h) if h is not None else "" for h in rows[0]]
        trade_sheets = [n for n in wb.sheetnames if "fill" in n.lower() or "trade" in n.lower()]

        def _arm_row(sheet_name: str) -> dict[str, Any] | None:
            if sheet_name not in wb.sheetnames:
                return None
            ws2 = wb[sheet_name]
            recs = list(ws2.iter_rows(values_only=True))
            if not recs:
                return None
            hh = [str(h) if h is not None else "" for h in recs[0]]
            for rec in recs[1:]:
                d = {hh[i]: rec[i] if i < len(rec) else None for i in range(len(hh))}
                if str(d.get("ARM_ID") or "") == DEVELOPMENT_CHALLENGER:
                    return d
            return None

        e4 = None
        for rec in rows[1:]:
            d = {headers[i]: rec[i] if i < len(rec) else None for i in range(len(headers))}
            if str(d.get("ARM_ID") or "") == DEVELOPMENT_CHALLENGER:
                e4 = d
                break
        return {
            "ok": e4 is not None,
            "e4": e4,
            "uncond": _arm_row("Unconditional"),
            "missed": _arm_row("Missed_Opportunity"),
            "trade_sheets": trade_sheets,
            "sheet_names": list(wb.sheetnames),
        }
    finally:
        wb.close()


def audit_aggregate_match(e4_pol: dict[str, Any], audit: dict[str, Any]) -> dict[str, Any]:
    row = audit.get("e4") or {}
    uncond = audit.get("uncond") or {}
    missed = audit.get("missed") or {}
    if not row:
        return {"V12_AUDIT_E4_AGGREGATE_MATCH": False, "reason": "NO_E4_ROW"}
    mid = e4_pol.get("UNCONDITIONAL_POLICY_MARKOUT") or {}
    bid = e4_pol.get("UNCONDITIONAL_FILL_TO_BID") or {}
    fg = e4_pol.get("FILLED_GROSS_MID") or {}
    ug = e4_pol.get("UNFILLED_GROSS_MID") or {}
    checks = {
        "FILLED_N": int(row.get("FILLED_N") or -1) == int(e4_pol.get("FILLED_N") or -2),
        "ELIGIBLE_N": int(row.get("ELIGIBLE_N") or -1) == int(e4_pol.get("ELIGIBLE_N") or -2),
        "UNFILLED_N": int(row.get("UNFILLED_N") or -1) == int(e4_pol.get("UNFILLED_N") or -2),
        "FILL_RATE": _close(row.get("FILL_RATE"), e4_pol.get("FILL_RATE")),
        "UNCOND_MID_180": _close(row.get("UNCOND_MID_180"), mid.get("180")),
        "UNCOND_BID_180": _close(row.get("UNCOND_BID_180"), bid.get("180")),
        "UNCOND_BID_300": _close(row.get("UNCOND_BID_300"), bid.get("300")),
        "FILLED_GROSS_180": _close(row.get("FILLED_GROSS_180"), fg.get("180")),
        "UNFILLED_GROSS_180": _close(row.get("UNFILLED_GROSS_180"), ug.get("180")),
        "ADVERSE": row.get("ADVERSE") is False and e4_pol.get("EXTREME_ADVERSE_SELECTION") is False,
        "XLSX_UNCOND_MID_180": (not uncond) or _close(uncond.get("180"), mid.get("180")),
        "XLSX_UNCOND_BID_180": (not uncond) or _close(uncond.get("BID_180"), bid.get("180")),
        "XLSX_MISSED_FILLED_N": (not missed) or int(missed.get("FILLED_N") or -1) == int(e4_pol.get("FILLED_N") or -2),
        "XLSX_MISSED_ADVERSE": (not missed) or missed.get("ADVERSE") is False,
    }
    trade_sheets = list(audit.get("trade_sheets") or [])
    return {
        "V12_AUDIT_E4_AGGREGATE_MATCH": all(bool(v) for v in checks.values()),
        "V12_AUDIT_HAS_FILL_SHEET": bool(trade_sheets),
        "V12_AUDIT_TRADE_BY_TRADE_NOTE": "V12 official audit.xlsx has no fill/trade sheet. Trade-by-trade identity is vs V12 day-cache execution tape that produced the audit aggregates.",
        "TRADE_SHEETS": trade_sheets,
        "checks": checks,
    }


def _arm(row: dict[str, Any]) -> dict[str, Any]:
    return _e4(row)


def independent_e4_match(v12_rows: list[dict[str, Any]], replay_rows: list[dict[str, Any]]) -> dict[str, Any]:
    if len(v12_rows) != len(replay_rows):
        return {
            "INDEPENDENT_E4_REPLAY_MATCH": False,
            "reason": f"N v12={len(v12_rows)} replay={len(replay_rows)}",
            "mismatch_n": abs(len(v12_rows) - len(replay_rows)),
            "details": [],
        }
    idx = {identity_keys(r): r for r in replay_rows}
    if len(idx) != len(replay_rows):
        return {"INDEPENDENT_E4_REPLAY_MATCH": False, "reason": "DUPLICATE_REPLAY_KEYS", "mismatch_n": 1, "details": []}
    fields = (
        "executable_signal",
        "bid0",
        "ask0",
        "tick",
        "inside_collapsed_to_bid",
        "inside_limit",
    )
    arm_fields = (
        "filled",
        "fill_t",
        "fill_price",
        "limit_price",
        "wait_budget_sec",
        "collapsed_to_bid",
        "evidence",
        "prim_mid_60",
        "prim_mid_180",
        "prim_mid_300",
        "sec_bid_60",
        "sec_bid_180",
        "sec_bid_300",
    )
    details: list[dict[str, Any]] = []
    for r in v12_rows:
        key = identity_keys(r)
        q = idx.get(key)
        if q is None:
            details.append({"key": key, "field": "MISSING_REPLAY"})
            continue
        for f in fields:
            a, b = r.get(f), q.get(f)
            if f == "executable_signal":
                if bool(a) != bool(b):
                    details.append({"key": key, "field": f, "v12": a, "replay": b})
            elif f == "inside_collapsed_to_bid":
                if bool(a) != bool(b) and not (a is None and b is None):
                    details.append({"key": key, "field": f, "v12": a, "replay": b})
            elif not _close(a, b) and not (a is None and b is None):
                details.append({"key": key, "field": f, "v12": a, "replay": b})
        ea, eb = _arm(r), _arm(q)
        for f in arm_fields:
            a, b = ea.get(f), eb.get(f)
            if f in {"filled", "collapsed_to_bid"}:
                if bool(a) != bool(b):
                    details.append({"key": key, "field": f"e4.{f}", "v12": a, "replay": b})
            elif f == "evidence":
                if str(a or "") != str(b or ""):
                    details.append({"key": key, "field": "e4.evidence", "v12": a, "replay": b})
            elif not _close(a, b) and not (a is None and b is None):
                details.append({"key": key, "field": f"e4.{f}", "v12": a, "replay": b})
    return {
        "INDEPENDENT_E4_REPLAY_MATCH": len(details) == 0,
        "mismatch_n": len(details),
        "compared_n": len(v12_rows),
        "details": details[:30],
    }


def freeze_decision(*, all_pass: bool) -> dict[str, Any]:
    if all_pass:
        return {
            "ENTRY_SIGNAL_SPEC_FROZEN_DEVELOPMENT": True,
            "ENTRY_EXECUTION_SPEC_FROZEN_DEVELOPMENT": True,
            "DEVELOPMENT_ENTRY_STACK": "T3_PULLBACK_RCI__E4_INSIDE1_W5",
            "ENTRY_CERTIFIED": False,
            "TRUE_OOS": False,
            "RUNTIME_CANDIDATE": False,
            "CASE": "PASS",
            "VERDICT": "SIMPLE_TECH_V13_ENTRY_DEVELOPMENT_FREEZE_READY",
            "NEXT": "STOP. Do not start EXIT in this run. TRUE_OOS=false. Not certified. Not a Runtime candidate.",
        }
    return {
        "ENTRY_SIGNAL_SPEC_FROZEN_DEVELOPMENT": False,
        "ENTRY_EXECUTION_SPEC_FROZEN_DEVELOPMENT": False,
        "DEVELOPMENT_ENTRY_STACK": None,
        "ENTRY_CERTIFIED": False,
        "TRUE_OOS": False,
        "RUNTIME_CANDIDATE": False,
        "CASE": "FAIL",
        "VERDICT": "SIMPLE_TECH_V13_ENTRY_STRUCTURE_VERIFICATION_FAILED",
        "NEXT": "STOP. Development freeze forbidden. Do not change V12 official results. Do not start EXIT.",
    }


def build_e4_policy(exe: list[dict[str, Any]], *, signal_n: int, integrity_ok: bool) -> dict[str, Any]:
    return policy_metrics(exe, DEVELOPMENT_CHALLENGER, signal_n=signal_n, e0_uncond=None, integrity_ok=integrity_ok)

