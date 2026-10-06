"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_opening_drive_location_reaccel_spec.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Role_Map",
    "States",
    "Invariants",
    "State_Diagram",
    "Positive_Exemplars",
    "Negative_Exemplars",
    "Descriptors",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(52, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    spec = dict(report.get("specification") or {})
    return {
        "v4_is_1m_strategy": spec.get("v4_is_1m_strategy"),
        "primary_setup_timeframe": spec.get("primary_setup_timeframe"),
        "one_m_role": spec.get("one_m_role"),
        "one_m_can_create_eligibility_without_5m": (spec.get("one_m_execution") or {}).get("can_create_eligibility_without_5m"),
        "daily_role": spec.get("daily_role"),
        "or_role": spec.get("or_role"),
        "valid_opening_states": spec.get("valid_opening_states"),
        "invalid_opening_states": spec.get("invalid_opening_states"),
        "TRUE_OPENING_DRIVE": spec.get("TRUE_OPENING_DRIVE"),
        "FAILED_OPEN_THEN_REAL_DRIVE": spec.get("FAILED_OPEN_THEN_REAL_DRIVE"),
        "meaningful_location": spec.get("location"),
        "or_alone_automatically_valid": (spec.get("location") or {}).get("or_alone_automatically_valid"),
        "opening_thesis_lost": spec.get("opening_thesis_lost"),
        "five_m_state_before_1m": "FIVE_M_CONTINUATION_STATE",
        "allowed_1m_role": spec.get("one_m_role"),
        "allowed_1m_descriptor_families": (spec.get("one_m_execution") or {}).get("state_change_families_allowed"),
        "numeric_1m_thresholds_frozen": (spec.get("one_m_execution") or {}).get("numeric_thresholds_frozen"),
        "tradingvalue_is_1m_trigger_gate": (spec.get("one_m_execution") or {}).get("tradingvalue_is_1m_gate"),
        "daily_bias_is_gate": (spec.get("daily_context") or {}).get("gate"),
        "economic_outcomes_used": False,
        "v4_machine_implemented": False,
        "VERDICT": (report.get("decision") or {}).get("VERDICT"),
        "NEXT": (report.get("decision") or {}).get("NEXT"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    spec = dict(report.get("specification") or {})
    tod = dict(spec.get("TRUE_OPENING_DRIVE") or {})
    fod = dict(spec.get("FAILED_OPEN_THEN_REAL_DRIVE") or {})
    loc = dict(spec.get("location") or {})
    lost = dict(spec.get("opening_thesis_lost") or {})
    one = dict(spec.get("one_m_execution") or {})
    lines = [
        "# PB1_V4_OPENING_DRIVE_LOCATION_REACCEL_SPEC_V1",
        "",
        "Semantic specification only. No V4 machine. No event_n. No PnL.",
        "",
        "Is V4 a 1m strategy? **false**",
        f"Primary setup timeframe? **{a.get('primary_setup_timeframe')}**",
        f"1m role? **{a.get('one_m_role')}**",
        "Can 1m create eligibility without valid 5m setup? **false**",
        f"Daily role? **{a.get('daily_role')}**",
        f"OR role? **{a.get('or_role')}**",
        "",
        f"Valid opening states? **{a.get('valid_opening_states')}**",
        f"Invalid opening states? **{a.get('invalid_opening_states')}**",
        "",
        "Definition of TRUE_OPENING_DRIVE?",
        f"- {tod.get('chart_test')}",
        f"- required: {tod.get('required_concepts')}",
        f"- not: {tod.get('not')}",
        "",
        "Definition of FAILED_OPEN_THEN_REAL_DRIVE?",
        f"- {fod.get('3382_must_match')}",
        f"- required: {fod.get('required_concepts')}",
        f"- not: {fod.get('not')}",
        "",
        "What constitutes meaningful location?",
        f"- A: { (loc.get('valid_forms') or {}).get('A') }",
        f"- B: { (loc.get('valid_forms') or {}).get('B') }",
        f"- C: { (loc.get('valid_forms') or {}).get('C') }",
        "Is OR alone automatically valid? **false**",
        "",
        f"What constitutes opening thesis lost? **{lost.get('lost_when')}** (state, not clock; cutoffs forbidden={lost.get('clock_cutoffs_forbidden')})",
        "",
        "What 5m state must exist before 1m is consulted? **FIVE_M_CONTINUATION_STATE**",
        f"What is the allowed 1m role? **{a.get('allowed_1m_role')}**",
        f"What normalized 1m descriptor families are allowed? **{a.get('allowed_1m_descriptor_families')}**",
        "Are numeric 1m thresholds frozen? **false**",
        "Is TradingValue a 1m trigger gate? **false**",
        "Is daily bias a gate? **false**",
        "Were any economic outcomes used? **false**",
        "Was any V4 machine implemented? **false**",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
    ]
    _ = one
    return "\n".join(lines) + "\n"


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    spec = dict(report.get("specification") or {})
    ex = dict(report.get("exemplars") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Role_Map": _kv_rows(spec.get("role_map")),
        "States": _kv_rows(
            {
                "valid": spec.get("valid_opening_states"),
                "invalid": spec.get("invalid_opening_states"),
                "TRUE_OPENING_DRIVE": spec.get("TRUE_OPENING_DRIVE"),
                "FAILED_OPEN_THEN_REAL_DRIVE": spec.get("FAILED_OPEN_THEN_REAL_DRIVE"),
            }
        ),
        "Invariants": [{"invariant": x} for x in list(spec.get("invariants") or [])] or [{"empty": True}],
        "State_Diagram": list((spec.get("state_diagram") or {}).get("transitions") or []) or [{"empty": True}],
        "Positive_Exemplars": list(ex.get("positive") or []) or [{"empty": True}],
        "Negative_Exemplars": list(ex.get("negative") or []) or [{"empty": True}],
        "Descriptors": _kv_rows(spec.get("candidate_descriptors")),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    (OUT / "report.json").write_text(json.dumps(_json_sanitize(report), ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
