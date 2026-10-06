"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.uniform10_entry_validity_audit import ANALYSIS_ID, NEW_FORWARD_N

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "uniform10_entry_validity_audit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(32, max(12, len(str(_k)) + 2))


def _kv(ws, data: dict[str, Any]) -> None:
    ws.append(["key", "value"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for k, v in data.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        ws.append([k, v])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 90


def _flat(prefix: str, obj: Any, out: dict[str, Any]) -> None:
    if isinstance(obj, dict):
        for k, v in obj.items():
            _flat(f"{prefix}.{k}" if prefix else str(k), v, out)
    else:
        out[prefix] = obj


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    order = [
        "Summary",
        "Tradability_State",
        "Target_Base_Audit",
        "Ranking_Executable",
        "Executable_Flag",
        "LODO_Leakage",
        "LODO_Recheck",
        "MFE_Audit",
        "C_Tail",
        "C_Session",
        "B_Summary",
        "B_Daily",
        "B_Period",
        "B_Concentration",
        "A_vs_B",
        "Decision",
        "Safety",
    ]
    first = True
    for name in order:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        if len(rows) == 1 and set(rows[0].keys()) <= {"key", "value"} or (
            rows and "key" in rows[0] and "value" in rows[0] and len(rows[0]) == 2
        ):
            _sheet(ws, rows)
        elif name in {"Summary", "Decision", "Safety", "MFE_Audit"} and rows and list(rows[0].keys()) == ["key", "value"]:
            _sheet(ws, rows)
        else:
            _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}
    _flat("", data, flat)
    return [{"key": k, "value": v} for k, v in flat.items()]


def build_markdown(report: dict[str, Any]) -> str:
    c = report.get("C_final") or {}
    b = report.get("B_final") or {}
    flags = report.get("flags") or {}
    rec = report.get("ranking_sets") or {}
    s1 = rec.get("SET1_ORIGINAL_ALL_ROWS") or {}
    s2 = rec.get("SET2_EXECUTABLE_AT_ANCHOR_ONLY") or {}
    lines = [
        "# UNIFORM10 ENTRY rebuild validity audit V2",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "OFFLINE AUDIT ONLY. Corrected Passive Fill is the execution SoT.",
        "C14 / Runtime / CLOCK / ENTRY / EXIT unchanged. UNIFORM10 not activated.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "",
        "## C required fields",
        "",
        f"TARGET_ITAYOSE_PRICE_CONTAMINATION: {flags.get('TARGET_ITAYOSE_PRICE_CONTAMINATION')}",
        f"NON_EXECUTABLE_RANKING_CONTAMINATION: {flags.get('NON_EXECUTABLE_RANKING_CONTAMINATION')}",
        f"MODEL_DESIGN_INVALID_TRADABILITY_FEATURE: {flags.get('MODEL_DESIGN_INVALID_TRADABILITY_FEATURE')}",
        f"LODO_INFORMATION_LEAK: {flags.get('LODO_INFORMATION_LEAK')}",
        f"MFE_METRIC_VALID: {flags.get('MFE_METRIC_VALID')}",
        f"ORIGINAL_TOP1_TARGET: {(s1.get('Top1') or {}).get('mean_target')}",
        f"EXECUTABLE_ONLY_TOP1_TARGET: {(s2.get('Top1') or {}).get('mean_target')}",
        f"ORIGINAL_TOP3_TARGET: {(s1.get('Top3') or {}).get('mean_target')}",
        f"EXECUTABLE_ONLY_TOP3_TARGET: {(s2.get('Top3') or {}).get('mean_target')}",
        f"C_PNL: {c.get('C_PNL')}",
        f"C_PNL_EX_TOP1: {c.get('C_PNL_EX_TOP1')}",
        f"C_PNL_EX_TOP3: {c.get('C_PNL_EX_TOP3')}",
        f"C_PNL_EX_285A: {c.get('C_PNL_EX_285A')}",
        f"C_VERDICT: {c.get('C_VERDICT')}",
        "",
        "## B required fields",
        "",
        f"trades: {b.get('trades')}",
        f"PnL: {b.get('PnL')}",
        f"PF: {b.get('PF')}",
        f"maxDD: {b.get('maxDD')}",
        f"positive_day_rate: {b.get('positive_day_rate')}",
        f"median_daily_pnl: {b.get('median_daily_pnl')}",
        f"AM_PnL: {b.get('AM_PnL')}",
        f"PM_PnL: {b.get('PM_PnL')}",
        f"first_entry_PnL: {b.get('first_entry_PnL')}",
        f"re_entry_PnL: {b.get('re_entry_PnL')}",
        f"PnL_ex_top1_trade: {b.get('PnL_ex_top1_trade')}",
        f"PnL_ex_top3_trades: {b.get('PnL_ex_top3_trades')}",
        f"PnL_ex_best_day: {b.get('PnL_ex_best_day')}",
        f"PnL_ex_top3_days: {b.get('PnL_ex_top3_days')}",
        f"PnL_ex_top_symbol: {b.get('PnL_ex_top_symbol')}",
        f"PnL_ex_top3_symbols: {b.get('PnL_ex_top3_symbols')}",
        f"PnL_ex_285A: {b.get('PnL_ex_285A')}",
        f"DEV10_PnL/PF: {b.get('DEV10_PnL_PF')}",
        f"POST7_PnL/PF: {b.get('POST7_PnL_PF')}",
        f"B_VERDICT: {b.get('B_VERDICT')}",
        "",
        f"RECOMMENDED_NEXT_STEP: {report.get('RECOMMENDED_NEXT_STEP')}",
        f"NEW_FORWARD_N: {NEW_FORWARD_N}",
        "TRUE_OOS: false",
        "",
        "STOP. B/C not written to Runtime. No new Strategy candidate. Paper not started.",
        "",
    ]
    return "\n".join(lines)
