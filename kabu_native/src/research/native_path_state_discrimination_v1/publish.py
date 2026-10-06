"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.native_path_state_discrimination_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Episode_Set",
    "Outcome_Class",
    "Feature_Set",
    "Availability_Audit",
    "D1_Training",
    "D2_Prediction",
    "D3_Prediction",
    "D4_Prediction",
    "Tree_Rules",
    "Logistic_Diagnostic",
    "Rule_Stability",
    "Symbol_Residual",
    "Sector_Residual",
    "Path_Lift",
    "Complete_Strategy",
    "Execution",
    "Tail_Dependency",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _block_row(step: dict[str, Any]) -> dict[str, Any]:
    tm = dict(step.get("tree_metrics") or {})
    lm = dict(step.get("logistic_metrics") or {})
    return {
        "block": step.get("block"),
        "role": step.get("role"),
        "train_n": step.get("train_n"),
        "eval_n": step.get("eval_n"),
        "tree_auc": tm.get("auc"),
        "tree_lift": tm.get("lift_at_cutoff"),
        "tree_precision": tm.get("precision_at_cutoff"),
        "base_rate": tm.get("base_favorable_rate"),
        "logistic_auc": lm.get("auc"),
        "top_family": step.get("top_family"),
        "rules_text": (step.get("tree") or {}).get("rules_text"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    pre = dict(report.get("prequential") or {})
    ext = dict(report.get("extracted_rules") or {})
    replays = list(report.get("replays") or [])
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "bg_cont_vwap_closed": True,
                "promoted": False,
                "frozen_validation_opened": False,
                "random_cv": False,
            }
        ),
        "Episode_Set": _kv(dict(report.get("episode_set") or {})),
        "Outcome_Class": [{"class": k, "n": v} for k, v in dict(report.get("outcome_counts") or {}).items()],
        "Feature_Set": _kv({"feature_spec_sha256": report.get("feature_spec_sha256"), "primary": (pre.get("spec") or {}).get("primary_features")}),
        "Availability_Audit": _kv(dict(report.get("availability_audit") or {})),
        "D1_Training": [_block_row(dict(pre.get("D1") or {}))],
        "D2_Prediction": [_block_row(dict(pre.get("D2") or {}))],
        "D3_Prediction": [_block_row(dict(pre.get("D3") or {}))],
        "D4_Prediction": [_block_row(dict(pre.get("D4") or {}))],
        "Tree_Rules": [
            {"block": b, "rules_text": ((pre.get(b) or {}).get("tree") or {}).get("rules_text"), "top_family": (pre.get(b) or {}).get("top_family")}
            for b in ("D1", "D2", "D3", "D4")
        ],
        "Logistic_Diagnostic": [
            {"block": b, "nested": json.dumps((pre.get(b) or {}).get("nested_logistic") or [], default=str), "coef_top": json.dumps((pre.get(b) or {}).get("logistic_coef_top") or [], default=str)}
            for b in ("D1", "D2", "D3", "D4")
        ],
        "Rule_Stability": list(ext.get("rules") or [{"empty": True}]),
        "Symbol_Residual": list(report.get("symbol_residual") or [{"empty": True}]),
        "Sector_Residual": list(report.get("sector_residual") or [{"empty": True}]),
        "Path_Lift": [
            {"block": b, **dict((pre.get(b) or {}).get("tree_metrics") or {})}
            for b in ("D1", "D2", "D3", "D4")
        ]
        + [{"block": "D2_D3", **dict(pre.get("D2_D3") or {})}],
        "Complete_Strategy": [
            {
                "rule_id": p.get("rule_id"),
                "trade_n": p.get("trade_n"),
                "day_n": p.get("day_n"),
                "symbol_n": p.get("symbol_n"),
                "mean_x0_bps": p.get("mean_x0_bps"),
                "mean_x1_bps": p.get("mean_x1_bps"),
                "profit_factor": p.get("profit_factor"),
                "block_mean_x0": p.get("block_mean_x0"),
                "survives": (p.get("survival") or {}).get("survives"),
                "exit_kind": p.get("exit_kind"),
            }
            for p in replays
        ]
        or [{"empty": True}],
        "Execution": [
            {"rule_id": p.get("rule_id"), "occupancy": p.get("occupancy"), "skipped": p.get("skipped"), "x1_tax_bps": p.get("x1_tax_bps"), "full_event_time_replay": True}
            for p in replays
        ]
        or [{"CAP": 3, "x1_tax_bps": 8.0}],
        "Tail_Dependency": [{"rule_id": p.get("rule_id"), **{k: v for k, v in dict(p.get("tail") or {}).items() if k != "n" or True}} for p in replays] or [{"empty": True}],
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0", "v27_bolted": False, "promoted": False}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# NATIVE_PATH_STATE_DISCRIMINATION_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            d.get("INTERPRETATION") or "",
            "",
            f"Episodes: **{a.get('How_many_episodes')}**",
            f"Outcome counts: `{a.get('Outcome_class_counts')}`",
            f"Feature leakage? **{a.get('Any_feature_leakage')}**",
            f"Random CV? **{a.get('Random_CV_used')}**",
            f"Tree spec: `{a.get('Primary_shallow_tree_spec')}`",
            f"D2/D3/D4/D2+D3 AUC: **{a.get('D2_auc')}** / **{a.get('D3_auc')}** / **{a.get('D4_auc')}** / **{a.get('D2_D3_auc')}**",
            f"Top families: `{a.get('Which_causal_states_distinguish_favorable_paths')}`",
            f"Stable across blocks? **{a.get('Are_the_same_interactions_stable_across_blocks')}**",
            f"Group adds? **{a.get('Does_behavior_group_add_incremental_information')}** Sector-rel? **{a.get('Does_sector_relative_state_add')}** Activity? **{a.get('Does_volume_activity_add')}** VWAP after others? **{a.get('Does_VWAP_state_add_after_the_other_information')}**",
            f"Extracted rules: **{a.get('How_many_extracted_rules')}** stable=`{a.get('stable_rule_ids')}`",
            f"Any complete-strategy survive? **{a.get('Any_rule_survives_Complete_Strategy_replay')}**",
            f"Best X0/X1/PF n/day/symbol: **{a.get('X0')}** / **{a.get('X1')}** / **{a.get('PF')}** / **{a.get('trade_N')}** / **{a.get('day_N')}** / **{a.get('symbol_N')}**",
            f"D2/D3/D4: `{a.get('D2_D3_D4')}`",
            f"Any X1>0? **{a.get('Does_any_candidate_have_X1_gt_0')}** 8136-dependent? **{a.get('Is_result_dependent_on_8136')}** top5%-dependent? **{a.get('Is_result_dependent_on_top_5pct_winners')}**",
            "",
            f"Complete strategy promoted? **{a.get('Any_Complete_Strategy_promoted')}**",
            f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
            f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
            f"New paid data? **{a.get('New_paid_data')}**",
            f"Kabu50? **{a.get('Kabu50')}**",
            f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
