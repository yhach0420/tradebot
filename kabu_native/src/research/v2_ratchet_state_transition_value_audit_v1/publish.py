"""Publish the ratchet state-transition value audit. Diagnostic only."""
from __future__ import annotations

import json
import pickle
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

from research.event_time_impulse_complete_strategy_v2.identity import bind
from research.v2_ratchet_state_transition_value_audit_v1.metrics import build
from research.v2_ratchet_state_transition_value_audit_v1.scan import scan
from research.v2_ratchet_state_transition_value_audit_v1.states import self_check

OUT = Path("results/research/v2_ratchet_state_transition_value_audit_v1")
BASE_N, BASE_PNL, BASE_PF = 11902, 1189150.0, 1.409303686366296
STRATEGY_SHA = "1d5ac587d4a67a7da4de9def64a6471428a136e4e30e1d26ca904983a8bcf505"
SIGNAL_SHA = "06345e9f7ddd2fcdf21a7beac49241495d3a4b6607715c879cf2ef0a307b4f68"


def _pf(rows: list[dict[str, Any]]) -> float | None:
    gain = sum(float(r["pnl_yen"]) for r in rows if float(r["pnl_yen"]) > 0)
    loss = sum(-float(r["pnl_yen"]) for r in rows if float(r["pnl_yen"]) < 0)
    return gain / loss if loss else None


def _chart(item: dict[str, Any], path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 4.4))
    ax.plot(item["t"], item["px"], color="black", lw=0.7, label="CurrentPrice")
    if item["bt"]:
        ax.plot(item["bt"], item["bid"], color="tab:blue", lw=0.5, label="Bid")
        ax.plot(item["bt"], item["ask"], color="tab:orange", lw=0.5, label="Ask")
    states = item["states"]
    for i, state in enumerate(states):
        t_end = states[i + 1]["t"] if i + 1 < len(states) else item["exit_t"]
        ax.hlines(state["support"], state["t"], t_end, colors="red", lw=0.9)
        ax.axvline(state["t"], color="green" if state["state"] == 0 else "magenta", lw=0.6)
    ax.axvline(item["exit_t"], color="purple", lw=0.8, label="EXIT")
    ax.plot([], [], color="red", lw=0.9, label="ACTIVE_SUPPORT")
    notes = []
    for state in states:
        now = state["exit_now_yen"]
        cont = state["continue_yen"]
        notes.append(f"S{state['state']} now={None if now is None else round(now)} final={round(state['final_yen'])} cont={None if cont is None else round(cont)}")
    ax.set_title(f"{item['kind']} {item['symbol']} {item['date']} r={item['ratchets']} {item['reason']}\n" + " | ".join(notes), fontsize=7)
    ax.legend(fontsize=6, loc="best")
    fig.tight_layout()
    fig.savefig(path, dpi=100)
    plt.close(fig)


def publish() -> dict[str, Any]:
    self_check()
    OUT.mkdir(parents=True, exist_ok=True)
    bound = bind()
    if bound["strategy_sha256"] != STRATEGY_SHA or bound["signal_sha256"] != SIGNAL_SHA:
        raise RuntimeError("frozen_sha_mismatch")
    cache = OUT / "_scan.pkl"
    if cache.exists():
        got = pickle.loads(cache.read_bytes())
    else:
        got = scan()
        cache.write_bytes(pickle.dumps(got, protocol=pickle.HIGHEST_PROTOCOL))
    base = got["baseline"]
    pnl = sum(float(r["pnl_yen"]) for r in base)
    pf = _pf(base)
    parity = len(base) == BASE_N and abs(pnl - BASE_PNL) < 1e-6 and pf is not None and abs(pf - BASE_PF) < 1e-12 and got["identity_miss"] == 0
    uid_ok = len(got["uids"]) == 11930 and len(set(got["uids"])) == 11930 and got["uids"] == got["uids2"]
    mapped = len({r["signal_uid"] for r in base})
    multi = len(base) - mapped
    integrity = bool(parity and uid_ok and mapped == BASE_N and multi == 0 and got["observer_mismatch"] == 0 and got["walk_mismatch"] == 0)
    stats = build(got["records"], got["dates"], got["original"], got["extension"])
    if not integrity:
        stats["verdict"] = "INTEGRITY_BLOCKED_V1"
        stats["next"] = "REPAIR_RECONSTRUCTION_INTEGRITY_ONLY"
        stats["confirmation_ladder_supported"] = False
        stats["stagewise_risk_allocation_justified"] = False
        stats["stationary_repeated_hazard"] = False
    report = {
        "study": "V2_RATCHET_STATE_TRANSITION_VALUE_AUDIT_V1",
        "research_only": True,
        "v2_changed": False,
        "new_data_acquired": False,
        "submit_cancel_live": "0/0/0",
        "strategy_sha256": bound["strategy_sha256"],
        "signal_sha256": bound["signal_sha256"],
        "baseline_parity_exact": parity,
        "trade_n": len(base),
        "pnl_yen": pnl,
        "pf": pf,
        "signal_uid_integrity": integrity,
        "full_signal_n": len(got["uids"]),
        "unique_signal_uid_n": len(set(got["uids"])),
        "mapped_trade_n": mapped,
        "unmapped_trade_n": len(base) - mapped,
        "multi_mapped_trade_n": multi,
        "observer_mismatch": got["observer_mismatch"],
        "walk_mismatch": got["walk_mismatch"],
        **stats,
    }
    (OUT / "report.json").write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    lines = [
        f"# {report['verdict']}",
        "",
        f"NEXT: {report['next']}",
        "",
        report["survivorship_warning"],
        "",
        "This is conditional-state analysis. It is not an exit rule and not a sizing rule.",
        "",
        f"parity={parity} integrity={integrity} n={len(base)} pnl={pnl} pf={pf}",
        "",
        "| STATE | AT_RISK_N | NEXT_RATCHET_N | EXIT_N | P_NEXT | P_EXIT | MEDIAN_TIME_TO_NEXT | MEDIAN_TIME_TO_EXIT | MEAN_CONTINUE | MEDIAN_CONTINUE | POSITIVE_FRACTION |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for block in report["states"].values():
        lines.append(
            "| {label} | {at_risk_n} | {next_ratchet_n} | {exit_n} | {p_next} | {p_exit} | {t_next} | {t_exit} | {mean_c} | {med_c} | {pos} |".format(
                label=block["label"],
                at_risk_n=block["at_risk_n"],
                next_ratchet_n=block["next_ratchet_n"],
                exit_n=block["exit_n"],
                p_next=block["p_next"],
                p_exit=block["p_exit"],
                t_next=block["time_to_next"]["median"],
                t_exit=block["time_to_exit"]["median"],
                mean_c=block["mean_continue_yen"],
                med_c=block["median_continue_yen"],
                pos=block["positive_fraction"],
            )
        )
    lines.extend(
        [
            "",
            f"transition rises={report['transition_probability_rises']}",
            f"continuation rises={report['continuation_value_rises']}",
            f"reward/risk improves={report['state_local_reward_risk_improves']}",
            f"stationary={report['stationary_repeated_hazard']}",
            f"ladder={report['confirmation_ladder_supported']}",
            "",
            "ENTRY changed: false",
            "EXIT changed: false",
            "submit/cancel/live: 0/0/0",
            "",
        ]
    )
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    for key in ("verdict", "next", "baseline_parity_exact", "signal_uid_integrity", "trade_n", "pnl_yen", "pf", "transition_probability_rises", "continuation_value_rises", "state_local_reward_risk_improves", "stationary_repeated_hazard", "confirmation_ladder_supported", "stagewise_risk_allocation_justified"):
        ws.append([key, report[key]])
    table = wb.create_sheet("transition")
    table.append(["state", "at_risk_n", "next_ratchet_n", "exit_n", "p_next", "p_exit", "median_time_to_next", "median_time_to_exit", "mean_continue_yen", "median_continue_yen", "positive_fraction", "median_mfe_bps", "median_mae_bps", "reward_risk", "median_bid_minus_support"])
    for block in report["states"].values():
        table.append([
            block["label"], block["at_risk_n"], block["next_ratchet_n"], block["exit_n"], block["p_next"], block["p_exit"],
            block["time_to_next"]["median"], block["time_to_exit"]["median"], block["mean_continue_yen"], block["median_continue_yen"],
            block["positive_fraction"], block["median_mfe_bps"], block["median_mae_bps"], block["reward_risk"], block["median_bid_minus_support"],
        ])
    visual = wb.create_sheet("visual_examples")
    visual.append(["kind", "symbol", "date", "ratchets", "reason", "pnl"])
    for item in got["charts"]:
        visual.append([item["kind"], item["symbol"], item["date"], item["ratchets"], item["reason"], item["pnl_yen"]])
    with TemporaryDirectory() as tmp:
        for i, item in enumerate(got["charts"]):
            png = Path(tmp) / f"c{i}.png"
            _chart(item, png)
            image = XLImage(str(png))
            image.width = 640
            image.height = 280
            image.anchor = f"A{8 + i * 18}"
            visual.add_image(image)
        wb.save(OUT / "audit.xlsx")
    cache.unlink(missing_ok=True)
    print(json.dumps({"verdict": report["verdict"], "next": report["next"], "parity": parity, "integrity": integrity}, sort_keys=True), flush=True)
    return report


if __name__ == "__main__":
    publish()
