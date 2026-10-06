"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_complete_strategy_economic_confirmation1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Precommit",
    "Identity",
    "Trades",
    "Sessions",
    "PnL",
    "Drawdown",
    "E0_E1",
    "Long_Short",
    "Exit_Reasons",
    "Execution_Tax",
    "CAP",
    "Same_Symbol",
    "Occupancy",
    "Reentry",
    "Concentration",
    "Invariants",
    "Decision",
    "Safety",
)
STRIP = {
    "isolation_before",
    "isolation_after",
    "minutes",
    "funnel_days",
    "e0_events",
    "e1_events",
    "walked_raw",
    "blocked_rows",
    "candidates",
    "drawdown",
}
TRADE_COLS = (
    "date",
    "symbol",
    "entry_type",
    "exec_variant",
    "side",
    "signal_t",
    "entry_t",
    "entry_px",
    "exit_t",
    "exit_px",
    "exit_reason",
    "reentry_n",
    "gross_pnl_yen",
    "execution_cost_yen",
    "net_pnl_yen",
    "holding_min",
    "thesis_death",
    "ops_flatten",
    "THESIS_LOST_AT",
    "THESIS_LOST_REASON",
    "execution_id",
)


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _kv(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    ident = dict(report.get("identity") or {})
    metrics = dict(report.get("metrics") or {})
    pre = dict(report.get("precommit") or {})
    conc = dict(metrics.get("concentration") or {})
    return {
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "FROZEN_ENTRY_IDENTITY": ident.get("frozen_identity"),
        "machine_sha": ident.get("machine_sha"),
        "source_inventory_sha": ident.get("source_inventory_sha"),
        "complete_strategy_identity": ident.get("complete_strategy_identity"),
        "COMPLETE_STRATEGY_SHA256": ident.get("COMPLETE_STRATEGY_SHA256"),
        "PRECOMMIT_ID": pre.get("precommit_id"),
        "PRECOMMIT_SHA256": pre.get("PRECOMMIT_SHA256"),
        "ECONOMIC_EDGE_CONCENTRATED": conc.get("ECONOMIC_EDGE_CONCENTRATED"),
        "net_pnl_yen": metrics.get("net_pnl_yen"),
        "profit_factor": metrics.get("profit_factor"),
        "mean_net_pnl_per_trade": metrics.get("mean_net_pnl_per_trade"),
        "trade_n": metrics.get("trade_n"),
        "primary_gate": metrics.get("primary_gate"),
        "V4_CHANGED": False,
        "COMPLETE_STRATEGY_CHANGED": False,
        "THRESHOLD_RETUNED": False,
        "FROZEN_VALIDATION_ECONOMIC_OPENED": False,
        "PROSPECTIVE_DATA_OPENED": False,
        "submit/cancel/live": "0/0/0",
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ident = dict(report.get("identity") or {})
    pre = dict(report.get("precommit") or {})
    metrics = dict(report.get("metrics") or {})
    inv = dict(report.get("invariants") or {})
    dec = dict(report.get("decision") or {})
    replay = dict(report.get("replay") or {})
    trades = list(report.get("trades") or [])
    compact = [{k: t.get(k) for k in TRADE_COLS} for t in trades]
    blocked = dict(metrics.get("blocked_hypothetical_not_used_for_verdict") or {})
    lost = dict(metrics.get("thesis_lost_reasons") or {})
    return {
        "Manifest": _kv(dict(report.get("answers") or {})),
        "Precommit": _kv({k: v for k, v in pre.items() if k not in ("pass_fail_rules",)}),
        "Identity": _kv({k: v for k, v in ident.items() if k != "files"}),
        "Trades": compact or [{"empty": True}],
        "Sessions": list(metrics.get("sessions") or []) or [{"empty": True}],
        "PnL": _kv(
            {
                k: metrics.get(k)
                for k in (
                    "signal_n",
                    "fill_n",
                    "trade_n",
                    "gross_pnl_yen",
                    "execution_cost_yen",
                    "net_pnl_yen",
                    "mean_net_pnl_per_trade",
                    "median_net_pnl_per_trade",
                    "profit_factor",
                    "win_rate",
                    "average_win",
                    "average_loss",
                    "max_drawdown",
                    "positive_session_n",
                    "negative_session_n",
                    "session_n",
                    "holding_time_mean",
                    "holding_time_median",
                    "max_concurrent",
                    "CAP_blocked_n",
                    "same_symbol_blocked_n",
                    "reentry_n",
                )
            }
        )
        + list(metrics.get("monthly_net_pnl") or []),
        "Drawdown": list(metrics.get("drawdown") or []) or [{"empty": True}],
        "E0_E1": [{"entry_type": "E0", **dict(metrics.get("E0") or {})}, {"entry_type": "E1", **dict(metrics.get("E1") or {})}],
        "Long_Short": [{"side": "LONG", **dict(metrics.get("long") or {})}, {"side": "SHORT", **dict(metrics.get("short") or {})}],
        "Exit_Reasons": (
            [{"exit_reason": "PB1_V4_THESIS_LOST_NEXT_OPEN", **dict(metrics.get("exit_technical") or {})}]
            + [{"exit_reason": "SESSION_FLAT_1520", **dict(metrics.get("exit_session_flat") or {})}]
            + [{"exit_reason": k, **v} for k, v in lost.items()]
        ),
        "Execution_Tax": _kv(dict(metrics.get("execution_tax") or {})),
        "CAP": _kv(
            {
                "cap": 5,
                "max_concurrent": replay.get("max_concurrent"),
                "CAP_blocked_n": replay.get("cap_blocked_n"),
                "CAP_blocked_hypothetical_net": blocked.get("CAP_blocked_net_pnl_yen"),
            }
        ),
        "Same_Symbol": _kv(
            {
                "same_symbol_blocked_n": replay.get("same_symbol_blocked_n"),
                "overlap_violation_n": replay.get("same_symbol_overlap_violation_n"),
                "hypothetical_net": blocked.get("same_symbol_blocked_net_pnl_yen"),
            }
        ),
        "Occupancy": _kv(dict(replay.get("occupancy") or {})),
        "Reentry": _kv({"reentry": metrics.get("reentry"), "first_entry": metrics.get("first_entry")}),
        "Concentration": _kv(dict(metrics.get("concentration") or {})),
        "Invariants": _kv(dict(inv.get("counts") or inv)),
        "Decision": _kv(dec),
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    metrics = dict(report.get("metrics") or {})
    conc = dict(metrics.get("concentration") or {})
    gate = dict(metrics.get("primary_gate") or {})
    md = [
        "# PB1 V4 Complete Strategy — Economic Confirmation 1",
        "",
        "Frozen Complete Strategy was not changed. Evaluation unit is Complete Strategy, not ENTRY alone.",
        "Window: 20251127–20260421. Frozen Validation economics and prospective remain sealed.",
        "",
        f"**VERDICT** `{ans.get('VERDICT')}`",
        f"**NEXT** `{ans.get('NEXT')}`",
        f"**ECONOMIC_EDGE_CONCENTRATED** `{conc.get('ECONOMIC_EDGE_CONCENTRATED')}`",
        "",
        "## Identity",
        "",
        f"- ENTRY `{ans.get('FROZEN_ENTRY_IDENTITY')}`",
        f"- machine_sha `{ans.get('machine_sha')}`",
        f"- source_inventory_sha `{ans.get('source_inventory_sha')}`",
        f"- Complete Strategy `{ans.get('complete_strategy_identity')}`",
        f"- COMPLETE_STRATEGY_SHA256 `{ans.get('COMPLETE_STRATEGY_SHA256')}`",
        f"- PRECOMMIT_SHA256 `{ans.get('PRECOMMIT_SHA256')}`",
        "",
        "## Primary economic gate (execution cost included)",
        "",
        f"- net_pnl_yen `{metrics.get('net_pnl_yen')}` gate `{gate.get('net_pnl_yen_gt_0')}`",
        f"- profit_factor `{metrics.get('profit_factor')}` gate `{gate.get('profit_factor_gt_1')}`",
        f"- mean_net_pnl_per_trade `{metrics.get('mean_net_pnl_per_trade')}` gate `{gate.get('mean_net_pnl_per_trade_gt_0')}`",
        f"- trade_n `{metrics.get('trade_n')}`",
        "",
        "## Concentration (diagnostic; no drop rule)",
        "",
        f"- net_pnl_ex_top1_trade `{conc.get('net_pnl_ex_top1_trade')}`",
        f"- net_pnl_ex_top1_day `{conc.get('net_pnl_ex_top1_day')}`",
        f"- net_pnl_ex_top1_symbol `{conc.get('net_pnl_ex_top1_symbol')}`",
        "",
        "Confirmation 1 is not rescored after seeing these results.",
        "",
        "STOP.",
        "",
    ]
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
