"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.causal_mechanism_representation_expansion_v1 import ANALYSIS_ID
from research.causal_mechanism_representation_expansion_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent_pin",
    "library_freeze",
    "p1_raw",
    "p2_raw",
    "gate",
    "candidate_library",
    "prior_use",
    "coverage",
    "markout_h5",
    "discovery_gates",
    "selection",
    "p1_signals",
    "p2_signals",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _sig_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in rows:
        out.append(
            {
                "date": r.get("date"),
                "symbol": r.get("symbol"),
                "i": r.get("i"),
                "signal_t0": r.get("signal_t0"),
                "entry_t": r.get("entry_t"),
                "entry_ask": r.get("entry_ask"),
                "markout_h3": r.get("markout_yen100_h3"),
                "markout_h5": r.get("markout_yen100_h5"),
                "markout_h10": r.get("markout_yen100_h10"),
                "excess_h5": r.get("excess_yen100_h5"),
                "complete_triple": r.get("complete_triple"),
            }
        )
    return out or [{"empty": True}]


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT') or a.get('29_VERDICT'))}",
        "",
        f"DISCOVERY_UNIT_VALID: {_fmt(d.get('DISCOVERY_UNIT_VALID') or a.get('DISCOVERY_UNIT_VALID'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT') or a.get('30_NEXT'))}",
        "",
        "NEW_STRATEGY_CREATED: false",
        "",
        "NEW_MECHANISM_OUTCOME_READ_BEFORE_UNIT_GATE: false",
        "",
        "## Answers",
        "",
    ]
    for k, v in a.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    dumped = json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n"
    (OUT / "report.json").write_text(dumped, encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_sheets(report: dict[str, Any], pack: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    scored = list(pack.get("scored") or [])
    return {
        "answers": kv_rows(report.get("answers")),
        "objective_alignment": kv_rows(pack.get("objective_alignment")),
        "parent_pin": kv_rows(pack.get("pin")),
        "library_freeze": kv_rows(pack.get("library_freeze")),
        "p1_raw": kv_rows(pack.get("P1")),
        "p2_raw": kv_rows(pack.get("P2")),
        "gate": kv_rows(pack.get("gate")),
        "candidate_library": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "TEMPLATE": r.get("TEMPLATE"),
                "PRIMITIVE_N": r.get("PRIMITIVE_N"),
                "DEFINITION": r.get("DEFINITION"),
                "RECLAIM_ID": r.get("RECLAIM_ID"),
            }
            for r in (pack.get("candidate_library") or [])
        ]
        or [{"empty": True}],
        "prior_use": list(pack.get("prior_use") or [{"empty": True}]),
        "coverage": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "EVENT_N": r.get("EVENT_N"),
                "EVENT_DAY_N": r.get("EVENT_DAY_N"),
                "COVERAGE_OK": bool(r.get("D1_EVENT_N") and r.get("D2_EVENT_DAY_N")),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "markout_h5": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "EVENT_N": r.get("EVENT_N"),
                "mean_markout_yen100": (r.get("h5") or {}).get("mean_markout_yen100"),
                "mean_excess_yen100": (r.get("h5") or {}).get("mean_excess_yen100"),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "discovery_gates": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "D1_EVENT_N": r.get("D1_EVENT_N"),
                "D2_EVENT_DAY_N": r.get("D2_EVENT_DAY_N"),
                "D3_MEAN_ABS_H5_GT0": r.get("D3_MEAN_ABS_H5_GT0"),
                "D4_MEAN_EXCESS_H5_GT0": r.get("D4_MEAN_EXCESS_H5_GT0"),
                "D5_MEAN_ABS_H3_GE0": r.get("D5_MEAN_ABS_H3_GE0"),
                "D6_MEAN_ABS_H10_GE0": r.get("D6_MEAN_ABS_H10_GE0"),
                "D7_MEAN_EXCESS_H3_GE0": r.get("D7_MEAN_EXCESS_H3_GE0"),
                "D8_MEAN_EXCESS_H10_GE0": r.get("D8_MEAN_EXCESS_H10_GE0"),
                "D9_POS_BLOCK_N": r.get("D9_POS_BLOCK_N"),
                "D10_EX_BEST_DAY": r.get("D10_EX_BEST_DAY"),
                "D11_BH_Q_H5": r.get("D11_BH_Q_H5"),
                "PASS_D1_D11": r.get("PASS_D1_D11"),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "selection": list(pack.get("selection_ranked") or [{"empty": True}]),
        "p1_signals": _sig_rows(list(pack.get("labeled_p1") or [])),
        "p2_signals": _sig_rows(list(pack.get("labeled_p2") or [])),
        "decision": kv_rows(pack.get("decision")),
        "safety": kv_rows(report.get("safety")),
    }
