"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v3_1_face_validity_fix.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Rule_Diff",
    "Funnel",
    "Deaths",
    "Old67_Causes",
    "Human_Sample",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "events", "failed_push_archive", "funnel_days", "chart_zones"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    deaths = dict(report.get("deaths") or {})
    human = dict(report.get("human") or {})
    dec = dict(report.get("decision") or {})
    att = dict(report.get("attribution_old67") or {})
    causes = dict(att.get("cause_counts") or {})
    meta = dict(report.get("sample_meta") or {})
    return {
        "parent_v3_unchanged": report.get("v3_unchanged"),
        "V3_1_MACHINE_SHA256": report.get("MACHINE_SHA256"),
        "PARENT_V3_SHA": report.get("PARENT_V3_SHA"),
        "existing_67_used_only_for_semantics": att.get("existing_67_used_only_for_semantics"),
        "any_future_outcome_used": att.get("future_outcome_used"),
        "old67_NON_DIRECTIONAL_OPEN": causes.get("NON_DIRECTIONAL_OPEN"),
        "old67_OPENING_IMPULSE_LOST": causes.get("OPENING_IMPULSE_LOST"),
        "old67_MICRO_OR_LEAK": causes.get("MICRO_OR_LEAK"),
        "old67_MICRO_STRUCTURE_NOT_TRADABLE": causes.get("MICRO_STRUCTURE_NOT_TRADABLE"),
        "old67_STRUCTURALLY_BLOCKED": causes.get("STRUCTURALLY_BLOCKED"),
        "old67_MOVE_ALREADY_REACHED_STRUCTURE": causes.get("MOVE_ALREADY_REACHED_STRUCTURE"),
        "old67_STALE_RETEST": causes.get("STALE_RETEST"),
        "old67_WEAK_RECLAIM": causes.get("WEAK_RECLAIM"),
        "NORMAL_1M_RANGE": report.get("NORMAL_1M_RANGE_defined"),
        "planned_R_must_exceed_one_local_noise_unit": report.get("planned_R_must_exceed_one_local_noise_unit"),
        "one_r_searched": report.get("one_r_noise_searched"),
        "meaningful_leave_must_exceed_one_local_noise_unit": report.get("meaningful_leave_must_exceed_one_local_noise_unit"),
        "leave_threshold_optimized": report.get("leave_threshold_optimized"),
        "CLEAN_OPENING_IMPULSE_rebuilt": report.get("CLEAN_OPENING_IMPULSE_rebuilt"),
        "retest_5min_gate": report.get("retest_5min_gate"),
        "clock_0930_cutoff": report.get("clock_0930_cutoff"),
        "new_indicator": report.get("new_indicator"),
        "V3_1_event_n": report.get("setup_n"),
        "unseen_candidate_n": report.get("unseen_candidate_n"),
        "independent_sample_n": report.get("independent_sample_n"),
        "independent_unavailable": dec.get("INDEPENDENT_FACE_VERIFY_NOT_AVAILABLE_IN_DISCOVERY"),
        "independent_unavailable_explicit": bool(dec.get("INDEPENDENT_FACE_VERIFY_NOT_AVAILABLE_IN_DISCOVERY")),
        "semantic_consistency_label": (report.get("semantic_consistency") or {}).get("label"),
        "old67_still_emitted_n": (report.get("semantic_consistency") or {}).get("still_emitted_n"),
        "weak_reclaims_mostly_disappeared": (report.get("semantic_consistency") or {}).get("weak_reclaims_mostly_disappeared"),
        "CLEAR_share": human.get("clear_share"),
        "QUESTIONABLE_share": human.get("questionable_share"),
        "NOT_share": human.get("not_share"),
        "opening_drive_human_agreement": human.get("opening_drive_human_agreement"),
        "structural_route_human_agreement": human.get("structural_route_human_agreement"),
        "valid_reclaim_share": human.get("valid_reclaim_share"),
        "any_return_test": report.get("pnl_test"),
        "any_pnl": report.get("pnl_test"),
        "old_confirmation_opened": False,
        "frozen_validation_opened": False,
        "submit_cancel_live": "0/0/0",
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "FACE_VALID": dec.get("FACE_VALID"),
        "NON_DIRECTIONAL_OPEN_funnel": (deaths.get("NON_DIRECTIONAL_OPEN") or {}).get("day_deaths"),
        "OPENING_IMPULSE_LOST_funnel": (deaths.get("OPENING_IMPULSE_LOST") or {}).get("day_deaths"),
        "MICRO_OR_LEAK_funnel": (deaths.get("MICRO_OR_LEAK") or {}).get("day_deaths"),
        "MICRO_STRUCTURE_funnel": (deaths.get("MICRO_STRUCTURE_NOT_TRADABLE") or {}).get("day_deaths"),
        "sampling": meta.get("sampling"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    lines = [
        "# FIX_PB1_V3_FACE_VALIDITY_V1 / PB1_V3_1_FACE_VALIDITY_FIX",
        "",
        "Semantic face-validity fix. Frozen V2 and V3 unchanged. No PnL.",
        "",
        f"Parent V3 unchanged? **{a.get('parent_v3_unchanged')}**",
        f"V3.1 SHA? `{a.get('V3_1_MACHINE_SHA256')}`",
        f"Existing 67 used only for semantics? **{a.get('existing_67_used_only_for_semantics')}**",
        f"Any future outcome used? **{a.get('any_future_outcome_used')}**",
        "",
        "How many of old 67 fail:",
        f"NON_DIRECTIONAL_OPEN? **{a.get('old67_NON_DIRECTIONAL_OPEN')}**",
        f"OPENING_IMPULSE_LOST? **{a.get('old67_OPENING_IMPULSE_LOST')}**",
        f"MICRO_OR_LEAK? **{a.get('old67_MICRO_OR_LEAK')}**",
        f"MICRO_STRUCTURE_NOT_TRADABLE? **{a.get('old67_MICRO_STRUCTURE_NOT_TRADABLE')}**",
        f"STRUCTURALLY_BLOCKED? **{a.get('old67_STRUCTURALLY_BLOCKED')}**",
        f"MOVE_ALREADY_REACHED_STRUCTURE? **{a.get('old67_MOVE_ALREADY_REACHED_STRUCTURE')}**",
        f"STALE_RETEST? **{a.get('old67_STALE_RETEST')}**",
        f"WEAK_RECLAIM? **{a.get('old67_WEAK_RECLAIM')}**",
        "",
        f"What is NORMAL_1M_RANGE? {a.get('NORMAL_1M_RANGE')}",
        f"Is planned_R required to exceed one local noise unit? **{a.get('planned_R_must_exceed_one_local_noise_unit')}**",
        f"Was 1.0 searched? **{a.get('one_r_searched')}**",
        f"Is meaningful leave required to exceed one local noise unit? **{a.get('meaningful_leave_must_exceed_one_local_noise_unit')}**",
        f"Was leave threshold optimized? **{a.get('leave_threshold_optimized')}**",
        f"Was CLEAN_OPENING_IMPULSE rebuilt? **{a.get('CLEAN_OPENING_IMPULSE_rebuilt')}**",
        f"Any hard 5m retest rule? **{a.get('retest_5min_gate')}**",
        f"Any 09:30/09:45 cutoff? **{a.get('clock_0930_cutoff')}**",
        f"Any new indicator? **{a.get('new_indicator')}**",
        "",
        f"V3.1 event_n? **{a.get('V3_1_event_n')}**",
        f"Unseen Discovery face-verify candidate_n? **{a.get('unseen_candidate_n')}**",
        f"New independent sample_n? **{a.get('independent_sample_n')}**",
        f"If unavailable, explicitly labeled? **{a.get('independent_unavailable_explicit')}**",
        f"Semantic consistency overlay? **{a.get('semantic_consistency_label')}**",
        f"Old 67 still emitted under V3.1? **{a.get('old67_still_emitted_n')}**",
        f"Weak reclaims mostly disappeared? **{a.get('weak_reclaims_mostly_disappeared')}**",
        "",
        f"CLEAR share? **{a.get('CLEAR_share')}**",
        f"QUESTIONABLE share? **{a.get('QUESTIONABLE_share')}**",
        f"NOT share? **{a.get('NOT_share')}**",
        f"Opening-drive human agreement? **{a.get('opening_drive_human_agreement')}**",
        f"Structural-route human agreement? **{a.get('structural_route_human_agreement')}**",
        f"Valid reclaim share? **{a.get('valid_reclaim_share')}**",
        "",
        f"Any return test? **{a.get('any_return_test')}**",
        f"Any PnL? **{a.get('any_pnl')}**",
        f"Old Confirmation opened? **{a.get('old_confirmation_opened')}**",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"submit/cancel/live? **{a.get('submit_cancel_live')}**",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    human = dict(report.get("human") or {})
    att = dict(report.get("attribution_old67") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Rule_Diff": _kv_rows(report.get("rule_diff")),
        "Funnel": _kv_rows(report.get("funnel")),
        "Deaths": _kv_rows(report.get("deaths")),
        "Old67_Causes": list(att.get("rows") or [_kv_rows(att.get("cause_counts") or {})[0]]),
        "Human_Sample": list(human.get("rows") or [_kv_rows(human)[0]]),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
