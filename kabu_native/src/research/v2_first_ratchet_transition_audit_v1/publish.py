"""Publish the first-ratchet transition audit."""
from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

from research.event_time_impulse_complete_strategy_v2.identity import bind
from research.v2_first_ratchet_transition_audit_v1.metrics import build
from research.v2_first_ratchet_transition_audit_v1.scan import scan
from research.v2_genuine_resistance_gate_complete_strategy_v1 import DETECTOR_SHA256, EXPECTED_PF, EXPECTED_PNL_YEN, EXPECTED_TRADE_N, SIGNAL_SHA256, STRATEGY_SHA256

OUT = Path("results/research/v2_first_ratchet_transition_audit_v1")


def _chart(item: dict[str, Any], path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.plot(item["t"], item["px"], color="black", lw=0.7, label="CurrentPrice")
    if item["bt"]:
        ax.plot(item["bt"], item["bid"], color="tab:blue", lw=0.5, label="Bid")
        ax.plot(item["bt"], item["ask"], color="tab:orange", lw=0.5, label="Ask")
    ax.axhline(item["pre_break"], color="red", ls=":", lw=0.8, label="PRE_BREAK_HIGH")
    ax.axhline(item["entry_px"], color="green", ls="--", lw=0.7, label="ENTRY_ASK")
    ax.axvline(item["entry_t"], color="green", lw=0.8)
    for ms, color in ((50, "gray"), (100, "gray"), (250, "olive"), (500, "olive"), (1000, "brown")):
        ax.axvline(item["entry_t"] + ms / 1000.0, color=color, lw=0.4, alpha=0.7)
    if item.get("first_ratchet_t"):
        ax.axvline(float(item["first_ratchet_t"]), color="magenta", lw=0.8, label="FIRST_RATCHET")
    ax.axvline(item["exit_t"], color="purple", lw=0.8, label="EXIT")
    ax.set_title(f"{item['kind']} {item['symbol']} {item['date']} r={item['ratchets']} {item['reason']}")
    ax.legend(fontsize=6, loc="best")
    fig.tight_layout()
    fig.savefig(path, dpi=100)
    plt.close(fig)


def publish() -> dict[str, Any]:
    OUT.mkdir(parents=True, exist_ok=True)
    bound = bind()
    raw = scan()
    stats = build(raw)
    pnl = sum(float(r["pnl_yen"]) for r in raw["rows"])
    n = len(raw["rows"])
    gain = sum(float(r["pnl_yen"]) for r in raw["rows"] if r["pnl_yen"] > 0)
    loss = -sum(float(r["pnl_yen"]) for r in raw["rows"] if r["pnl_yen"] < 0)
    pf = gain / loss if loss else None
    uids = raw["uids"]
    parity = n == EXPECTED_TRADE_N and abs(pnl - EXPECTED_PNL_YEN) < 1e-6 and pf is not None and abs(pf - EXPECTED_PF) < 1e-12
    uid_ok = len(uids) == 11930 and len(set(uids)) == 11930 and uids == raw["uids2"]
    mapped = sum(1 for r in raw["rows"] if r.get("signal_uid"))
    report = {
        "study": "V2_FIRST_RATCHET_TRANSITION_AUDIT_V1",
        "verdict": stats["verdict"],
        "next": stats["next"],
        "parity": parity,
        "trade_n": n,
        "pnl_yen": pnl,
        "pf": pf,
        "strategy_sha256": bound["strategy_sha256"],
        "signal_sha256": bound["signal_sha256"],
        "strategy_sha_ok": bound["strategy_sha256"] == STRATEGY_SHA256 and bound["signal_sha256"] == SIGNAL_SHA256,
        "signal_uid_integrity": bool(uid_ok and mapped == 11902 and raw["identity_miss"] == 0 and raw["observer_mismatch"] == 0),
        "full_signal_n": len(uids),
        "unique_signal_uid_n": len(set(uids)),
        "mapped_trade_n": mapped,
        "unmapped_trade_n": n - mapped,
        "multi_mapped_trade_n": mapped - len({r["signal_uid"] for r in raw["rows"]}),
        "observer_mismatch": raw["observer_mismatch"],
        "identity_miss": raw["identity_miss"],
        "detector_sha_unchanged": DETECTOR_SHA256,
        "v2_changed": False,
        "new_data_acquired": False,
        "submit_cancel_live": [0, 0, 0],
        **stats,
    }
    (OUT / "report.json").write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    lines = [
        f"# {report['verdict']}",
        "",
        f"NEXT: {report['next']}",
        "",
        f"parity={parity} n={n} pnl={pnl} pf={pf}",
        f"T0 separation={stats['t0_stable_separation']} post-signal={stats['post_signal_separation']} earliest_ms={stats['earliest_landmark_ms']}",
        "",
        "This is not an implemented gate.",
        "ENTRY changed: false",
        "EXIT changed: false",
        "submit/cancel/live: 0/0/0",
    ]
    (OUT / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    for key in ("verdict", "next", "parity", "trade_n", "pnl_yen", "pf", "signal_uid_integrity", "t0_stable_separation", "post_signal_separation", "earliest_landmark_ms", "predictable_at_entry", "information_before_first_ratchet", "later_test_justified", "concentration_problem"):
        summary.append([key, json.dumps(report[key], default=str) if isinstance(report[key], (dict, list)) else report[key]])
    t0 = wb.create_sheet("t0")
    t0.append(["feature", "n0", "n1", "mean0", "mean1", "median0", "median1", "d", "auc", "stable"])
    for key, block in report["t0"].items():
        t0.append([key, block.get("n0"), block.get("n1"), block.get("mean0"), block.get("mean1"), block.get("median0"), block.get("median1"), block.get("standardized_difference"), block.get("auc"), key in report["t0_stable_features"]])
    marks = wb.create_sheet("landmarks")
    marks.append(["ms", "unresolved", "ratcheted", "exited", "stable", "features"])
    for ms, block in report["landmarks"].items():
        marks.append([ms, block["at_risk_n"], block["already_ratcheted_n"], block["already_exited_n"], block["stable_precursor"], ",".join(block["stable_features"])])
    visual = wb.create_sheet("visual_examples")
    visual.append(["kind", "symbol", "date", "ratchets", "reason", "pnl"])
    for item in raw["charts"]:
        visual.append([item["kind"], item["symbol"], item["date"], item["ratchets"], item["reason"], item["pnl_yen"]])
    with TemporaryDirectory() as tmp:
        for i, item in enumerate(raw["charts"]):
            png = Path(tmp) / f"c{i}.png"
            _chart(item, png)
            image = XLImage(str(png))
            image.anchor = f"A{8 + i * 22}"
            visual.add_image(image)
        wb.save(OUT / "audit.xlsx")
    return {"verdict": report["verdict"], "next": report["next"], "parity": parity, "t0": report["t0_stable_separation"], "post": report["post_signal_separation"]}


if __name__ == "__main__":
    print(publish())
