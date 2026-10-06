"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.passive_fill_corrected_rebase import (
    ANALYSIS_ID,
    EXECUTION_SEMANTICS_DEFECT_FIXED,
    FILL_ELIGIBILITY_SOT,
    FILL_PRICE_RULE,
    STRATEGY_RETUNED,
    TASK_LABEL,
    WAIT_SEC,
)

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "passive_fill_corrected_rebase"
JST = timezone(timedelta(hours=9))
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            if v == float("inf"):
                v = "Infinity"
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _kv(ws, data: dict[str, Any]) -> None:
    ws.append(["key", "value"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for k, v in data.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        ws.append([k, v])
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 80


def write_artifacts(report: dict[str, Any], extra_sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(report), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    md = report.get("_markdown") or ""
    (OUT / "report.md").write_text(str(md), encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    _kv(ws, {k: v for k, v in report.items() if k not in {"_markdown", "baseline", "period_recheck"} and not isinstance(v, (dict, list))})
    for name, rows in extra_sheets.items():
        ws2 = wb.create_sheet(name[:31])
        _sheet(ws2, rows if isinstance(rows, list) else [rows])
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    hl = ((report.get("baseline") or {}).get("headline") or {})
    cand = report.get("NEW_RUNTIME_CANDIDATE") or {}
    q = report.get("capture_20260827") or {}
    pr = report.get("period_recheck") or {}
    lines = [
        "# Corrected Passive Fill historical rebase",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        f"TASK: {TASK_LABEL}",
        f"STRATEGY_RETUNED: {STRATEGY_RETUNED}",
        f"EXECUTION_SEMANTICS_DEFECT_FIXED: {EXECUTION_SEMANTICS_DEFECT_FIXED}",
        f"WAIT_SEC: {WAIT_SEC}",
        f"fill_price: {FILL_PRICE_RULE}",
        f"Fill SoT: {FILL_ELIGIBILITY_SOT}",
        "SAFETY: submit/cancel/live=0/0/0. Paper/OPVAL not started.",
        "",
        "## NEW_RUNTIME_CANDIDATE",
        "",
        f"- candidate_id: `{cand.get('candidate_id')}`",
        f"- candidate_sha: `{cand.get('candidate_sha')}`",
        f"- runtime_code_sha: `{cand.get('runtime_code_sha')}`",
        f"- git_commit: `{cand.get('runtime_code_git_commit')}`",
        f"- inventory_digest: `{cand.get('runtime_inventory_digest')}`",
        f"- config_digest: `{cand.get('config_sha256')}`",
        f"- STRATEGY_CHANGED: {cand.get('STRATEGY_CHANGED')}",
        f"- CLOCK_GRID_CHANGED: {cand.get('CLOCK_GRID_CHANGED')}",
        f"- ENTRY_CHANGED: {cand.get('ENTRY_CHANGED')}",
        f"- EXIT_CHANGED: {cand.get('EXIT_CHANGED')}",
        f"- EXECUTION_SEMANTICS_CHANGED: {cand.get('EXECUTION_SEMANTICS_CHANGED')}",
        "",
        "## CORRECTED_BASELINE",
        "",
        f"- trades: {hl.get('trades')}",
        f"- PnL: {hl.get('pnl')}",
        f"- PF: {hl.get('PF')}",
        f"- maxDD: {hl.get('maxDD')}",
        f"- win/loss/draw: {hl.get('win')}/{hl.get('loss')}/{hl.get('draw')}",
        "",
        "## OLD_RESULTS_STATUS",
        "",
        str((report.get("supersede") or {}).get("OLD_RESULTS_STATUS")),
        "",
        "## 20260827_CAPTURE",
        "",
        f"- DECISION: {q.get('DECISION')}",
        f"- PASS: {q.get('PASS')}",
        f"- fail: {q.get('fail')}",
        "",
        "## Period recheck (no new RCA)",
        "",
        f"- status: {pr.get('period_decay_status')}",
        f"- DEV10 PnL/PF/trades: {(pr.get('DEV10') or {}).get('pnl')} / {(pr.get('DEV10') or {}).get('PF')} / {(pr.get('DEV10') or {}).get('trades')}",
        f"- POST7 PnL/PF/trades: {(pr.get('POST7') or {}).get('pnl')} / {(pr.get('POST7') or {}).get('PF')} / {(pr.get('POST7') or {}).get('trades')}",
        "",
        f"verdict: **{report.get('verdict')}**",
        "",
        "STOP. UNIFORM10 Runtime activation is forbidden in this study.",
        "",
    ]
    return "\n".join(lines)
