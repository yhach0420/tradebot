"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.target_architecture_selection import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "target_architecture_selection"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "T3Components",
    "T1Cross",
    "T2Cross",
    "T3Gates",
    "T1Gates",
    "T2Gates",
    "ConsensusDays",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.8g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# TARGET ARCHITECTURE PRECOMMIT SELECTION V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine. Common 600s cohort. Existing 27 Ridge specs. No nested selector.",
            "T1=OPPORTUNITY_ONLY. T2=RISK_ONLY. T3=T1+T2 scalar. Not independent alpha families.",
            "Selection by cross-component alignment only. No raw-uplift cherry-pick.",
            "No Exact. No PnL. No new model. No T0/T4 rescue. No path-order gate.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"T3_MFE_MEDIAN_SPEC_DELTA: {_fmt(g.get('T3_MFE_MEDIAN_SPEC_DELTA'))}",
            f"T3_DOWNSIDE_MEDIAN_SPEC_DELTA: {_fmt(g.get('T3_DOWNSIDE_MEDIAN_SPEC_DELTA'))}",
            f"T3_MFE_POSITIVE_SPEC_N: {g.get('T3_MFE_POSITIVE_SPEC_N')}",
            f"T3_DOWNSIDE_POSITIVE_SPEC_N: {g.get('T3_DOWNSIDE_POSITIVE_SPEC_N')}",
            "",
            f"T3_MFE_CONSENSUS_POS_DAYS: {g.get('T3_MFE_CONSENSUS_POS_DAYS')}",
            f"T3_MFE_CONSENSUS_NEG_DAYS: {g.get('T3_MFE_CONSENSUS_NEG_DAYS')}",
            f"T3_DOWNSIDE_CONSENSUS_POS_DAYS: {g.get('T3_DOWNSIDE_CONSENSUS_POS_DAYS')}",
            f"T3_DOWNSIDE_CONSENSUS_NEG_DAYS: {g.get('T3_DOWNSIDE_CONSENSUS_NEG_DAYS')}",
            "",
            f"T3_MFE_EX_BEST_DAY: {_fmt(g.get('T3_MFE_EX_BEST_DAY'))}",
            f"T3_MFE_EX_TOP3_DAYS: {_fmt(g.get('T3_MFE_EX_TOP3_DAYS'))}",
            f"T3_DOWNSIDE_EX_BEST_DAY: {_fmt(g.get('T3_DOWNSIDE_EX_BEST_DAY'))}",
            f"T3_DOWNSIDE_EX_TOP3_DAYS: {_fmt(g.get('T3_DOWNSIDE_EX_TOP3_DAYS'))}",
            "",
            f"T3_COMPONENT_ALIGNMENT_PASS: {_tf(g.get('T3_COMPONENT_ALIGNMENT_PASS'))}",
            "",
            f"T1_MFE_MEDIAN_DELTA: {_fmt(g.get('T1_MFE_MEDIAN_DELTA'))}",
            f"T1_DOWNSIDE_MEDIAN_DELTA: {_fmt(g.get('T1_DOWNSIDE_MEDIAN_DELTA'))}",
            f"T1_CLEAN_OPPORTUNITY: {_tf(g.get('T1_CLEAN_OPPORTUNITY'))}",
            "",
            f"T2_DOWNSIDE_MEDIAN_DELTA: {_fmt(g.get('T2_DOWNSIDE_MEDIAN_DELTA'))}",
            f"T2_MFE_MEDIAN_DELTA: {_fmt(g.get('T2_MFE_MEDIAN_DELTA'))}",
            f"T2_CLEAN_RISK: {_tf(g.get('T2_CLEAN_RISK'))}",
            "",
            f"PRIMARY_TARGET: {g.get('PRIMARY_TARGET')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No target-specific model this run. No Exact. No new model. No T3 weight search.",
            "Runtime/C14 unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
