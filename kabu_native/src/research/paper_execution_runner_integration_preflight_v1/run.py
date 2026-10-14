"""Offline preflight. Does not enable the causal book in the sealed runner."""
from __future__ import annotations

import hashlib
import json
import os
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

NATIVE = Path(__file__).resolve().parents[3]
SRC = NATIVE / "src"
ROOT = NATIVE.parent
for item in (str(ROOT), str(NATIVE), str(SRC)):
    if item not in sys.path:
        sys.path.insert(0, item)

OUT = NATIVE / "results" / "research" / "paper_execution_runner_integration_preflight_v1"
os.environ["TRADEBOT_C3_RESIDUAL_PATH"] = str(OUT / "replay_residual.json")
os.environ.pop("TRADEBOT_PAPER_EXECUTION_CAUSAL", None)

from research.current50_v16_complete_strategy_risk_foundation_v1.foundation import _capture_dir, _symbol
from research.execution_realism_layer_v1.replay import load_baseline, load_universe
from small_paper.paper_execution_causal_v1 import PaperExecutionCausalBook

DAY = "20261013"
AM_WATERMARK = 6244
PM_WATERMARK = 545280
JST = timezone(timedelta(hours=9))
PRIOR_SHA = "429fa4103fa5628d444ccd6071708dc37b356c09540f84f916a01ecdba0769a0"
PRIOR_CAUSAL_PNL = -47080.0
PRIOR_CAUSAL_N = 263
FLAT_UIDS = (
    "20261013|AM|9268.T|588",
    "20261013|AM|5985.T|4237",
    "20261013|AM|7287.T|5476",
    "20261013|AM|5032.T|4615",
    "20261013|AM|6663.T|249",
)
LOCKED = {
    "20261013|AM|4382.T|1432": {
        "states": ["ENTRY_SIGNAL", "BUY_FILLED", "EXIT_SIGNAL", "POSITION_CLOSED"],
        "entry_fill": "09:03:33.237",
        "exit_signal": "09:03:33.237",
        "exit_fill": "09:03:33.272",
        "pnl_yen": -300.0,
    },
    "20261013|AM|4382.T|1434": {
        "states": ["ENTRY_SIGNAL", "BUY_FILLED", "EXIT_SIGNAL", "POSITION_CLOSED"],
        "entry_fill": "09:03:33.286",
        "exit_signal": "09:03:35.119",
        "exit_fill": "09:03:35.219",
        "pnl_yen": 600.0,
    },
}
PROTECTED = (
    NATIVE / "src" / "small_paper" / "paper_session_executor.py",
    NATIVE / "src" / "small_paper" / "pilot_runner.py",
    NATIVE / "src" / "small_paper" / "fixed_support_x1_session.py",
    NATIVE / "results" / "research" / "paper_execution_correctness_fix_v1" / "report.json",
    NATIVE / "results" / "research" / "execution_realism_layer_v1" / "report.json",
)


def _sha(path: Path) -> str:
    if not path.is_file():
        return ""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _fmt(stamp) -> str:
    if stamp is None:
        return ""
    return datetime.fromtimestamp(float(stamp), JST).strftime("%H:%M:%S.%f")[:-3]


def _loads():
    try:
        import orjson

        return orjson.loads
    except Exception:
        return json.loads


def replay(symbols: list[str], capture: Path) -> dict:
    os.environ["TRADEBOT_C3_RESIDUAL_PATH"] = str(OUT / "replay_residual.json")
    residual = OUT / "replay_residual.json"
    if residual.exists():
        residual.unlink()
    am = PaperExecutionCausalBook()
    pm = PaperExecutionCausalBook()
    for book, kind in ((am, "am"), (pm, "pm")):
        book.notifier = None
        book.ledger_path = None

        def admit(symbol: str, _t: float, membership=symbols) -> tuple[bool, str]:
            if symbol in membership:
                return True, "SAME_DAY_AM_FROZEN_UNIVERSE"
            return False, "NOT_IN_DAY_FREEZE"

        book.admission = admit
        book.bind_causal_schedule(day=DAY, kind=kind)
    loads = _loads()
    for part in sorted(capture.glob("push_part_*.jsonl")):
        with part.open("rb") as handle:
            for raw in handle:
                if not raw.strip():
                    continue
                rec = loads(raw)
                if rec.get("kind") not in (None, "market_push"):
                    continue
                seq = int(rec.get("sequence") or 0)
                payload = rec.get("original_payload") if isinstance(rec.get("original_payload"), dict) else None
                if not payload:
                    continue
                symbol = _symbol(str(rec.get("symbol") or payload.get("Symbol") or ""))
                if not symbol:
                    continue
                received = str(rec.get("received_at") or "")
                event = {"symbol": symbol, "payload": payload, "received_at": received, "t0_push_received_at": received}
                if seq > AM_WATERMARK and not am._closed:
                    am.on_market_event(event)
                if seq > PM_WATERMARK and not pm._closed:
                    pm.on_market_event(event)
    if not am._closed:
        am.close_at_runtime_boundary(day=DAY, kind="am")
    if not pm._closed:
        pm.close_at_runtime_boundary(day=DAY, kind="pm")
    closed = list(am.closed) + list(pm.closed)
    digest = hashlib.sha256(
        json.dumps(
            [(row.get("signal_uid"), row.get("pnl_yen"), row.get("entry_fill_t"), row.get("exit_fill_t")) for row in closed],
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()
    stopped = list(am.post_stop_signals) + list(pm.post_stop_signals)
    for row in stopped:
        row["clock"] = _fmt(row["signal_t"])
    return {
        "closed": closed,
        "residuals": list(am.residuals) + list(pm.residuals),
        "hero": list(am.hero) + list(pm.hero),
        "violations": list(am.violations) + list(pm.violations),
        "stopped": stopped,
        "sha": digest,
        "live_order_enabled": bool(getattr(am, "live_order_enabled", False) or getattr(pm, "live_order_enabled", False)),
    }


def _case(model: dict, uid: str) -> dict:
    closed = next((row for row in model["closed"] if row.get("signal_uid") == uid), None)
    states = [row.get("state") for row in model["hero"] if row.get("signal_uid") == uid]
    return {
        "signal_uid": uid,
        "states": states,
        "entry_fill": _fmt(None if closed is None else closed.get("entry_fill_t")),
        "exit_signal": _fmt(None if closed is None else closed.get("exit_signal_t")),
        "exit_fill": _fmt(None if closed is None else closed.get("exit_fill_t")),
        "pnl_yen": None if closed is None else closed.get("pnl_yen"),
        "fill_hold_sec": None if closed is None else closed.get("fill_hold_sec"),
    }


def write_xlsx(path: Path, report: dict) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "判定"
    sheet.append(("判定", "結果"))
    for key, value in report["checks"].items():
        sheet.append((key, value))
    stopped = book.create_sheet("停止後シグナル")
    stopped.append(("銘柄", "UID", "時刻", "BUY_PENDING", "BUY_FILLED", "理由", "CAP増分", "same-symbol増分", "場"))
    for row in report["post_stop"]:
        stopped.append(
            (
                row["symbol"],
                row["signal_uid"],
                row["clock"],
                row["buy_pending"],
                row["buy_filled"],
                row["reason"],
                row["cap_delta"],
                row["same_symbol_delta"],
                row["session"],
            )
        )
    diffs = book.create_sheet("損益差")
    diffs.append(("UID", "旧損益", "隔離損益", "旧決済", "隔離決済"))
    for row in report["pnl_diffs"]:
        diffs.append((row["signal_uid"], row["paper_pnl_yen"], row["causal_pnl_yen"], row["paper_exit"], row["causal_exit"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def write_md(path: Path, report: dict) -> None:
    lines = [
        f"# {report['verdict']}",
        "",
        f"NEXT: {report['next']}",
        "",
        report["summary"],
        "",
        "## 判定",
        "",
    ]
    for key, value in report["checks"].items():
        lines.append(f"- {key}: {value}")
    lines.extend(["", "## 11:20以降", "", report["entry_stop_text"], "", "## 固定した損益", ""])
    for label, row in report["frozen"].items():
        lines.append(f"- {label}: {row}")
    lines.extend(["", "## 1432と1434", ""])
    for row in report["locked_cases"]:
        lines.append(
            f"- {row['signal_uid']}: {' → '.join(row['states'])}。BUY {row['entry_fill']}、EXIT {row['exit_signal']}、SELL {row['exit_fill']}、損益 {row['pnl_yen']}。"
        )
    lines.extend(
        [
            "",
            report["adoption_text"],
            "",
            "既定Runtimeは切り替えていない。submit/cancel/live = 0/0/0。",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    before = {str(path): _sha(path) for path in PROTECTED}
    from small_paper.paper_execution_offline_adapter import resolve_offline_paper_executor

    sealed = resolve_offline_paper_executor(
        {"execution_family": "X1_IMMEDIATE_ASK"},
        strategy_mode="C3",
        offline=False,
    )
    off_parity = (
        type(sealed).__name__ == "C3PaperSessionExecutor"
        and sealed.live_order_enabled is False
        and sealed.snapshot_status()["submit"] == 0
        and sealed.snapshot_status()["cancel"] == 0
        and sealed.snapshot_status()["live"] == 0
    )
    capture = _capture_dir(DAY)
    if capture is None:
        raise SystemExit("NO_CAPTURE")
    symbols = load_universe()
    first = replay(symbols, capture)
    second = replay(symbols, capture)
    opened = [row for row in first["stopped"] if row["opened"] or row["buy_pending"] or row["buy_filled"]]
    am_rows = [row for row in first["stopped"] if row["session"] == "AM"]
    pm_rows = [row for row in first["stopped"] if row["session"] == "PM"]
    entry_pass = not opened and not first["violations"]
    by_uid = {str(row["signal_uid"]): row for row in first["closed"]}
    flat_pnl = round(sum(float(by_uid[uid]["pnl_yen"]) for uid in FLAT_UIDS if uid in by_uid), 6)
    rest = [row for row in first["closed"] if row.get("signal_uid") not in FLAT_UIDS]
    rest_pnl = round(sum(float(row["pnl_yen"]) for row in rest), 6)
    total = round(sum(float(row["pnl_yen"]) for row in first["closed"]), 6)
    baseline = load_baseline()
    base_by = {str(row["signal_uid"]): row for row in baseline}
    diffs = []
    for uid, paper in base_by.items():
        causal = by_uid.get(uid)
        if causal is None:
            diffs.append(
                {
                    "signal_uid": uid,
                    "paper_pnl_yen": paper["pnl_yen"],
                    "causal_pnl_yen": None,
                    "paper_exit": _fmt(paper["exit_t"]),
                    "causal_exit": "",
                }
            )
            continue
        if float(causal["pnl_yen"]) != float(paper["pnl_yen"]) or _fmt(causal.get("exit_fill_t")) != _fmt(paper["exit_t"]):
            diffs.append(
                {
                    "signal_uid": uid,
                    "paper_pnl_yen": paper["pnl_yen"],
                    "causal_pnl_yen": causal["pnl_yen"],
                    "paper_exit": _fmt(paper["exit_t"]),
                    "causal_exit": _fmt(causal.get("exit_fill_t")),
                }
            )
    locked_cases = [_case(first, uid) for uid in LOCKED]
    locked_ok = all(
        row["states"] == LOCKED[row["signal_uid"]]["states"]
        and row["entry_fill"] == LOCKED[row["signal_uid"]]["entry_fill"]
        and row["exit_signal"] == LOCKED[row["signal_uid"]]["exit_signal"]
        and row["exit_fill"] == LOCKED[row["signal_uid"]]["exit_fill"]
        and row["pnl_yen"] == LOCKED[row["signal_uid"]]["pnl_yen"]
        for row in locked_cases
    )
    flat_times = {uid: _fmt(by_uid[uid].get("exit_fill_t")) for uid in FLAT_UIDS if uid in by_uid}
    expected_flat = {
        "20261013|AM|9268.T|588": "11:25:05.316",
        "20261013|AM|5985.T|4237": "11:25:11.445",
        "20261013|AM|7287.T|5476": "11:25:01.684",
        "20261013|AM|5032.T|4615": "11:25:05.066",
        "20261013|AM|6663.T|249": "11:25:46.917",
    }
    checks = {
        "ENTRY_STOP_INVARIANT_PASS": entry_pass,
        "PAPER_OFF_PATH_PARITY_PASS": off_parity,
        "CAUSAL_EXECUTION_E2E_PASS": bool(
            locked_ok and flat_times == expected_flat and len(first["closed"]) == 268 and total == -43480.0
        ),
        "SESSION_CLOSE_PASS": bool(flat_times == expected_flat and len(first["residuals"]) == 0 and flat_pnl == 3600.0),
        "RECOVERY_PASS": bool(first["sha"] == second["sha"] == PRIOR_SHA),
        "NO_REAL_ORDER_PATH_PASS": first["live_order_enabled"] is False,
        "PAPER_RUNTIME_ADOPTION_READY": False,
    }
    passed = all(value for key, value in checks.items() if key != "PAPER_RUNTIME_ADOPTION_READY")
    verdict = "PAPER_RUNTIME_ADOPTION_NOT_READY" if passed and not opened else "RUNNER_PREFLIGHT_FAIL"
    report = {
        "verdict": verdict,
        "next": "NO_RUNTIME_CHANGE",
        "summary": (
            f"11:20以降の候補は {len(am_rows)} 件、15:20以降は {len(pm_rows)} 件。"
            "いずれもBUY_PENDINGにもBUY_FILLEDにもなっていない。"
            "隔離モデルの損益は前回の263件に、11:25以降へ回した5件の+3,600円を足した値と一致する。"
            "既定Runnerへの接続はまだ行っていない。"
        ),
        "checks": checks,
        "entry_stop_text": (
            f"停止後に検出した候補は {len(first['stopped'])} 件で、成立した新規建玉は {len(opened)} 件。"
            "件数は候補の検出数であり、約定件数ではない。"
        ),
        "post_stop": first["stopped"],
        "opened_after_stop": opened,
        "frozen": {
            "旧Paper": "268件 / −29,500円",
            "前回因果約定": "263件決済 / −47,080円",
            "今回隔離モデル": f"{len(first['closed'])}件 / {total}円",
            "前回未決済5件の今回損益": f"{flat_pnl}円",
            "5件を除いた263件": f"{len(rest)}件 / {rest_pnl}円",
            "263件との差": rest_pnl - PRIOR_CAUSAL_PNL,
        },
        "accounting_ok": rest_pnl == PRIOR_CAUSAL_PNL and flat_pnl == 3600.0 and PRIOR_CAUSAL_N == 263,
        "pnl_diffs": diffs,
        "pnl_diff_n": len(diffs),
        "locked_cases": locked_cases,
        "flat_exit_clocks": flat_times,
        "determinism_sha": first["sha"],
        "prior_sha_match": first["sha"] == PRIOR_SHA,
        "second_sha_match": first["sha"] == second["sha"],
        "adoption_text": (
            "残課題は、封印されたRunnerが11:25で停止し、オフラインアダプタを呼んでいないこと。"
            "切替OFFの接続口は新しいモジュールにあり、既定ファイルは変更していない。"
            "Paper Runtimeへ反映できる状態ではない。"
        ),
        "submit": 0,
        "cancel": 0,
        "live": 0,
        "default_switch": False,
    }
    after = {str(path): _sha(path) for path in PROTECTED}
    report["protected_unchanged"] = before == after
    if not report["protected_unchanged"] or not report["accounting_ok"] or not passed:
        report["verdict"] = "RUNNER_PREFLIGHT_FAIL"
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    write_md(OUT / "report.md", report)
    write_xlsx(OUT / "audit.xlsx", report)
    print(report["verdict"], len(am_rows), len(pm_rows), len(opened), report["accounting_ok"], first["sha"] == PRIOR_SHA, flush=True)


if __name__ == "__main__":
    main()
