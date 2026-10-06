"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.top1_objective_development import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "top1_objective_development"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Integrity",
    "OOFGate",
    "PairedDays",
    "Folds",
    "RankShape",
    "SpecStability",
    "Parity",
    "Exact",
    "ExecCheck",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# TOP1 ENTRY OBJECTIVE PRECOMMITTED DEVELOPMENT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "New development hypothesis. Not a C3 continuation candidate.",
            "CanonicalEngine. TARGET V4 M4_PERSISTENT 600s. CURRENT IRREGULAR CLOCK.",
            "Primary: mean daily Top1 TARGET V4 uplift. RANK1_ONLY admission.",
            "Existing 27 Ridge specs. No new feature/model. No RF/pairwise. No Top3 selection. No PnL selection.",
            "Audit best spec was not adopted.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"CANONICAL_TRAINING_ROWS: {g.get('CANONICAL_TRAINING_ROWS')}",
            f"MATCHED_RANKING_ROWS: {g.get('MATCHED_RANKING_ROWS')}",
            "",
            f"CURRENT_TOP1_UPLIFT: {g.get('CURRENT_TOP1_UPLIFT')}",
            f"TOP1_MODEL_UPLIFT: {g.get('TOP1_MODEL_UPLIFT')}",
            "",
            f"TOP1_DELTA_MEAN: {g.get('TOP1_DELTA_MEAN')}",
            f"TOP1_DELTA_MEDIAN: {g.get('TOP1_DELTA_MEDIAN')}",
            "",
            f"TOP1_POSITIVE_DAYS: {g.get('TOP1_POSITIVE_DAYS')}",
            f"TOP1_NEGATIVE_DAYS: {g.get('TOP1_NEGATIVE_DAYS')}",
            "",
            f"TOP1_EX_BEST_DAY: {g.get('TOP1_EX_BEST_DAY')}",
            f"TOP1_EX_TOP3_DAYS: {g.get('TOP1_EX_TOP3_DAYS')}",
            "",
            f"TOP1_MAX_DAY_CONTRIBUTION: {g.get('TOP1_MAX_DAY_CONTRIBUTION')}",
            "",
            f"TOP1_OOF_GATE_PASS: {_tf(g.get('TOP1_OOF_GATE_PASS'))}",
            "",
            f"RANK1_TARGET: {g.get('RANK1_TARGET')}",
            f"RANK2_TARGET: {g.get('RANK2_TARGET')}",
            f"RANK3_TARGET: {g.get('RANK3_TARGET')}",
            "",
            f"MOST_COMMON_SPEC: {g.get('MOST_COMMON_SPEC')}",
            f"MOST_COMMON_SPEC_SHARE: {g.get('MOST_COMMON_SPEC_SHARE')}",
            "",
            f"FINAL_SPEC: {g.get('FINAL_SPEC')}",
            "",
            f"OOF_EXACT_RAN: {_tf(g.get('OOF_EXACT_RAN'))}",
            f"OOF_TOP1_EXACT: {g.get('OOF_TOP1_EXACT')}",
            "",
            f"FINAL_EXACT_RAN: {_tf(g.get('FINAL_EXACT_RAN'))}",
            f"FINAL_TOP1_EXACT: {g.get('FINAL_TOP1_EXACT')}",
            "",
            f"CURRENT_RANK1_WOULD_FILL: {g.get('CURRENT_RANK1_WOULD_FILL')}",
            f"TOP1_MODEL_RANK1_WOULD_FILL: {g.get('TOP1_MODEL_RANK1_WOULD_FILL')}",
            "",
            f"CURRENT_RANK1_FILL_TO_600: {g.get('CURRENT_RANK1_FILL_TO_600')}",
            f"TOP1_MODEL_RANK1_FILL_TO_600: {g.get('TOP1_MODEL_RANK1_FILL_TO_600')}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Runtime/C14 change. No C4 auto-start. No Paper/OPVAL.",
            "submit/cancel/live=0/0/0.",
            "",
        ]
    )
