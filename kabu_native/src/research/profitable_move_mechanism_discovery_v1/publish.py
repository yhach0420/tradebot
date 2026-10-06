"""Write report.json / report.md / audit.xlsx only. No CSV proliferation."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.profitable_move_mechanism_discovery_v1 import ANALYSIS_ID
from research.profitable_move_mechanism_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "data",
    "outcome_semantics",
    "predicate_library",
    "prior_use",
    "candidate_library",
    "coverage",
    "markout_h3",
    "markout_h5",
    "markout_h10",
    "daily",
    "blocks",
    "statistics",
    "fdr",
    "discovery_gates",
    "closed_lineage",
    "selection",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _markout_rows(scored: list[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    out = []
    for r in scored:
        h = dict(r.get(key) or {})
        out.append(
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "EVENT_N": r.get("EVENT_N"),
                "mean_markout_yen100": h.get("mean_markout_yen100"),
                "median_markout_yen100": h.get("median_markout_yen100"),
                "mean_excess_yen100": h.get("mean_excess_yen100"),
                "median_excess_yen100": h.get("median_excess_yen100"),
                "mean_markout_bps": h.get("mean_markout_bps"),
                "median_markout_bps": h.get("median_markout_bps"),
            }
        )
    return out or [{"empty": True}]


def _daily_rows(scored: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in scored:
        daily = dict(r.get("daily_mean_excess_h5") or {})
        if not daily:
            out.append({"MECHANISM_ID": r.get("MECHANISM_ID"), "date": None, "mean_excess_h5": None})
            continue
        for day, val in sorted(daily.items()):
            out.append({"MECHANISM_ID": r.get("MECHANISM_ID"), "date": day, "mean_excess_h5": val})
    return out or [{"empty": True}]


def _block_rows(scored: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in scored:
        blocks = dict(r.get("blocks_mean_excess_h5") or {})
        rec = {"MECHANISM_ID": r.get("MECHANISM_ID"), "H5_POSITIVE_BLOCK_N": r.get("H5_POSITIVE_BLOCK_N")}
        rec.update(blocks)
        out.append(rec)
    return out or [{"empty": True}]


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        "TERMINOLOGY: PROFITABLE_MOVE_MECHANISM",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT') or a.get('35_VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT') or a.get('36_NEXT'))}",
        "",
        f"SELECTED_MECHANISM_ID: {_fmt(d.get('SELECTED_MECHANISM_ID') or a.get('29_selected_mechanism_id'))}",
        "",
        "FULL_STRATEGY_FROZEN: false",
        "",
        "FULL_CAUSAL_PORTFOLIO_RUN: false",
        "",
        "STRATEGY_PNL_CLAIMED: false",
        "",
        "INFORMATION_FAMILY_INVENTORY_AS_PRIMARY_NEXT: false",
        "",
        "## Answers",
        "",
    ]
    for k, v in a.items():
        if v is None:
            lines.append(f"{k}: null")
            lines.append("")
            continue
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    dumped = json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n"
    (OUT / "report.json").write_text(dumped, encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_sheets(report: dict[str, Any], pack: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    scored = list(pack.get("scored") or [])
    return {
        "answers": kv_rows(report.get("answers")),
        "objective_alignment": kv_rows(pack.get("objective_alignment")),
        "data": kv_rows(pack.get("data")),
        "outcome_semantics": kv_rows(pack.get("outcome_semantics")),
        "predicate_library": list(pack.get("predicate_library") or [{"empty": True}]),
        "prior_use": list(pack.get("prior_use") or [{"empty": True}]),
        "candidate_library": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "TEMPLATE": r.get("TEMPLATE"),
                "PRIMITIVE_N": r.get("PRIMITIVE_N"),
                "ONSET_OF": r.get("ONSET_OF"),
                "DEFINITION": r.get("DEFINITION"),
                "FAMILIES": r.get("FAMILIES"),
            }
            for r in (pack.get("candidate_library") or [])
        ]
        or [{"empty": True}],
        "coverage": list(pack.get("coverage_rows") or [{"empty": True}]),
        "markout_h3": _markout_rows(scored, "h3"),
        "markout_h5": _markout_rows(scored, "h5"),
        "markout_h10": _markout_rows(scored, "h10"),
        "daily": _daily_rows(scored),
        "blocks": _block_rows(scored),
        "statistics": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "EVENT_N": r.get("EVENT_N"),
                "EVENT_DAY_N": r.get("EVENT_DAY_N"),
                "H5_PERM_P": r.get("H5_PERM_P"),
                "Q_VALUE_H5": r.get("Q_VALUE_H5"),
                "H5_POSITIVE_BLOCK_N": r.get("H5_POSITIVE_BLOCK_N"),
                "H5_EX_BEST_DAY_TOTAL_EXCESS": r.get("H5_EX_BEST_DAY_TOTAL_EXCESS"),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "fdr": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "H5_PERM_P": r.get("H5_PERM_P"),
                "Q_VALUE_H5": r.get("Q_VALUE_H5"),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "discovery_gates": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "D1_EVENT_N": r.get("D1_EVENT_N"),
                "D2_EVENT_DAY_N": r.get("D2_EVENT_DAY_N"),
                "D3_MEAN_ABS_H5_GT0": r.get("D3_MEAN_ABS_H5_GT0"),
                "D4_MEAN_EXCESS_H5_GT0": r.get("D4_MEAN_EXCESS_H5_GT0"),
                "D5_MEAN_ABS_H3_GE0": r.get("D5_MEAN_ABS_H3_GE0"),
                "D6_MEAN_ABS_H10_GE0": r.get("D6_MEAN_ABS_H10_GE0"),
                "D7_MEAN_EXCESS_H3_GE0": r.get("D7_MEAN_EXCESS_H3_GE0"),
                "D8_MEAN_EXCESS_H10_GE0": r.get("D8_MEAN_EXCESS_H10_GE0"),
                "D9_POS_BLOCK_N": r.get("D9_POS_BLOCK_N"),
                "D10_EX_BEST_DAY": r.get("D10_EX_BEST_DAY"),
                "D11_BH_Q_H5": r.get("D11_BH_Q_H5"),
                "PASS_D1_D11": r.get("PASS_D1_D11"),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "closed_lineage": [
            {
                "MECHANISM_ID": r.get("MECHANISM_ID"),
                "PASS_D1_D11": r.get("PASS_D1_D11"),
                "LINEAGE_CLASS": r.get("LINEAGE_CLASS"),
                "EXACT_PRIOR_TEST_MATCH": r.get("EXACT_PRIOR_TEST_MATCH"),
                "EXACT_CLOSED_LINEAGE": r.get("EXACT_CLOSED_LINEAGE"),
                "NEAREST_CLOSED_LINEAGE": r.get("NEAREST_CLOSED_LINEAGE"),
                "ELIGIBLE_AS_NEW_ARCHITECTURE": r.get("ELIGIBLE_AS_NEW_ARCHITECTURE"),
            }
            for r in scored
        ]
        or [{"empty": True}],
        "selection": list(pack.get("selection_ranked") or [{"empty": True}]),
        "decision": kv_rows(pack.get("decision")),
        "safety": kv_rows(report.get("safety")),
    }
