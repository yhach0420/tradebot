"""Write report.json / report.md / audit.xlsx only. Raw capture stays in data/market_context_capture/."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.new_causal_information_acquisition_v1 import ANALYSIS_ID
from research.new_causal_information_acquisition_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "Parent",
    "Standard_Config",
    "Exclusive",
    "Universe",
    "Preflight",
    "Capture_Today",
    "Futures_NK",
    "Futures_TOPIX",
    "Tests",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(body: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    cap = dict(body.get("capture_today") or {})
    uni = dict((body.get("universe_pack") or {}).get("universe") or {})
    return {
        "Answers": _kv(dict(body.get("answers") or {})),
        "Parent": _kv(dict(body.get("parent") or {})),
        "Standard_Config": _kv(dict(body.get("standard_config") or {})),
        "Exclusive": _kv(dict(body.get("exclusive") or {})),
        "Universe": _kv(uni),
        "Preflight": _kv(dict(body.get("preflight") or {})),
        "Capture_Today": _kv({k: v for k, v in cap.items() if k not in {"nk225mini", "topix"}}),
        "Futures_NK": _kv(dict(cap.get("nk225mini") or {})),
        "Futures_TOPIX": _kv(dict(cap.get("topix") or {})),
        "Tests": _kv(dict(body.get("tests") or {})),
        "Decision": _kv(dict(body.get("decision") or {})),
        "Safety": [
            {
                "submit_cancel_live": (body.get("answers") or {}).get("31_submit_cancel_live"),
                "sendorder": (body.get("answers") or {}).get("32_sendorder_call_N"),
                "Runtime_changed": (body.get("answers") or {}).get("29_Runtime_changed"),
                "Paper_changed": (body.get("answers") or {}).get("30_Paper_changed"),
                "TRUE_OOS": False,
                "CERTIFIED": False,
            }
        ],
    }


def build_markdown(body: dict[str, Any]) -> str:
    a = dict(body.get("answers") or {})
    d = dict(body.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        "",
        "Parent: POST_OPEN_CAUSAL_DOWNSIDE_MECHANISM_DISCOVERY_V1 "
        "CURRENT_CAPTURE_BIDIRECTIONAL_INFORMATION_LIMIT_CONFIRMED_V1 → NEW_CAUSAL_INFORMATION_ACQUISITION_V1.",
        "",
        "Dedicated mode is Core10 + Dynamic38 + NK225mini + TOPIX = 50. "
        "Standard Paper stays Core10 + Dynamic40 = 50.",
        "",
        "## Required answers",
        "",
    ]
    for k, v in a.items():
        lines.append(f"- {k}: `{json.dumps(v, ensure_ascii=False, default=str)}`")
    lines.extend(
        [
            "",
            "## Rules held",
            "",
            "- No ENTRY / EXIT / threshold / ML / TopK / PnL.",
            "- No Yahoo / historical futures backfill onto 20260722-20260807.",
            "- Raw capture path: `data/market_context_capture/` (not `data/market_capture/`).",
            "- Bid1 = Buy1, Ask1 = Sell1 via canonical_board.",
            "- Availability clock = received_at / ingress.",
            "",
        ]
    )
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any], *, out_dir: Optional[Path] = None) -> dict[str, str]:
    dest = Path(out_dir) if out_dir else OUT
    dest.mkdir(parents=True, exist_ok=True)
    report = json_sanitize(body)
    (dest / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (dest / "report.md").write_text(build_markdown(report), encoding="utf-8")
    wb = Workbook()
    first = True
    sheets = build_sheets(report)
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    extra = [
        ("Gates", _kv(dict(report.get("gates") or {}))),
        ("RCA", _kv(dict(report.get("rca") or {}))),
        ("PID", _kv(dict(report.get("pid") or {}))),
    ]
    for name, rows in extra:
        if not rows:
            continue
        ws = wb.create_sheet(name)
        _sheet(ws, rows)
    xlsx = dest / "audit.xlsx"
    wb.save(xlsx)
    return {
        "report_json": str(dest / "report.json"),
        "report_md": str(dest / "report.md"),
        "audit_xlsx": str(xlsx),
    }
