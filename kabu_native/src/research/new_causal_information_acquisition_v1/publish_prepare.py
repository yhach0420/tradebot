"""Write prepare readiness report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OUT = Path(__file__).resolve().parents[3] / "results" / "research" / "new_info_first_capture_readiness_v1"
SHEET_ORDER = ("Answers", "Gates", "Prepared", "Prebuild", "Exclusive", "Decision")


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def write_prepare_artifacts(result: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    body = json_sanitize(
        {
            "ANALYSIS_ID": "NEW_INFO_FIRST_CAPTURE_READINESS_V1",
            "TRUE_OOS": False,
            "CERTIFIED": False,
            **result,
        }
    )
    (OUT / "report.json").write_text(json.dumps(body, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    a = dict(body.get("answers") or {})
    md = [
        "# NEW_INFO_FIRST_CAPTURE_READINESS_V1",
        "",
        f"VERDICT: **{body.get('VERDICT')}**",
        f"NEXT: `{body.get('NEXT')}`",
        "",
        "## Required answers",
        "",
    ]
    for k, v in a.items():
        md.append(f"- {k}: `{json.dumps(v, ensure_ascii=False, default=str)}`")
    md.append("")
    (OUT / "report.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    wb = Workbook()
    sheets = {
        "Answers": _kv(a),
        "Gates": _kv(dict((body.get("gates") or {}).get("checks") or {})),
        "Prepared": _kv(
            {
                k: v
                for k, v in dict(body.get("prepared_manifest") or {}).items()
                if k not in {"core10", "dynamic40", "dynamic38", "competitor_audit"}
            }
        ),
        "Prebuild": _kv(dict(body.get("prebuild") or {})),
        "Exclusive": _kv(dict((body.get("prepared_manifest") or {}).get("competitor_audit") or {})),
        "Decision": _kv({"VERDICT": body.get("VERDICT"), "NEXT": body.get("NEXT"), "ok": body.get("ok")}),
    }
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
    return {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(OUT / "audit.xlsx"),
    }
