"""Write report.json / report.md / audit.xlsx only. No extra CSV."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_complete_strategy_build_and_economic_validation.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Strategy_Identity",
    "Data_Roles",
    "Execution_Binding",
    "Exit_Binding",
    "State_Machine",
    "Trades",
    "Session_Summary",
    "PnL",
    "Drawdown",
    "E0_E1",
    "Exit_Reasons",
    "CAP",
    "Reentry",
    "Execution_Tax",
    "Economic_Confirmation",
    "Safety",
)
STRIP = {
    "isolation_before",
    "isolation_after",
    "minutes",
    "funnel_days",
    "e0_events",
    "e1_events",
    "walked_raw",
    "blocked_rows",
    "candidates",
}
TRADE_COLS = (
    "date",
    "symbol",
    "entry_type",
    "exec_variant",
    "side",
    "signal_t",
    "entry_t",
    "entry_px",
    "exit_t",
    "exit_px",
    "exit_reason",
    "reentry_n",
    "gross_pnl_yen",
    "execution_cost_yen",
    "net_pnl_yen",
    "holding_min",
    "thesis_death",
    "ops_flatten",
    "THESIS_LOST_AT",
    "THESIS_LOST_REASON",
    "execution_id",
)


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _kv(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    ident = dict(report.get("identity") or {})
    metrics = dict(report.get("development_metrics") or {})
    replay = dict(report.get("replay") or {})
    return {
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "FROZEN_ENTRY_IDENTITY": ident.get("frozen_identity"),
        "machine_sha": ident.get("machine_sha"),
        "source_inventory_sha": ident.get("source_inventory_sha"),
        "complete_strategy_identity": dec.get("complete_strategy_identity"),
        "COMPLETE_STRATEGY_SHA256": dec.get("COMPLETE_STRATEGY_SHA256"),
        "RESEARCH_EXECUTION_APPROXIMATION": True,
        "DEVELOPMENT_PNL_IS_CERTIFICATION": False,
        "ECONOMIC_CONFIRMATION1_OPENED": False,
        "ECONOMIC_CONFIRMATION2_OPENED": False,
        "PROSPECTIVE_DATA_OPENED": False,
        "V4_CHANGED": False,
        "ENTRY_RETUNED": False,
        "PNL_USED_FOR_FREEZE": False,
        "MFE_MAE_USED": False,
        "submit/cancel/live": "0/0/0",
        "development_trade_n": metrics.get("trade_n"),
        "development_net_pnl_yen_not_certification": metrics.get("net_pnl_yen"),
        "max_concurrent": replay.get("max_concurrent"),
        "cap_blocked_n": replay.get("cap_blocked_n"),
        "same_symbol_blocked_n": replay.get("same_symbol_blocked_n"),
        "freeze_basis": dec.get("freeze_basis"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ident = dict(report.get("identity") or {})
    roles = dict(report.get("data_roles") or {})
    exec_inv = dict(report.get("execution_inventory") or {})
    exit_inv = dict(report.get("exit_inventory") or {})
    freeze = dict(report.get("decision") or {})
    replay = dict(report.get("replay") or {})
    metrics = dict(report.get("development_metrics") or {})
    trades = list(report.get("trades") or [])
    compact = [{k: t.get(k) for k in TRADE_COLS} for t in trades]
    eq = 0.0
    peak = 0.0
    dd_rows = []
    for t in trades:
        eq += float(t.get("net_pnl_yen") or 0.0)
        peak = max(peak, eq)
        dd_rows.append(
            {
                "date": t.get("date"),
                "symbol": t.get("symbol"),
                "entry_t": t.get("entry_t"),
                "net_pnl_yen": t.get("net_pnl_yen"),
                "equity": eq,
                "drawdown": eq - peak,
            }
        )
    return {
        "Manifest": _kv(dict(report.get("answers") or {})),
        "Strategy_Identity": _kv(
            {
                "complete_strategy_identity": freeze.get("complete_strategy_identity"),
                "COMPLETE_STRATEGY_SHA256": freeze.get("COMPLETE_STRATEGY_SHA256"),
                "entry_identity": ident.get("frozen_identity"),
                "machine_sha": ident.get("machine_sha"),
                "source_inventory_sha": ident.get("source_inventory_sha"),
                "v4_source_sha": ident.get("v4_source_sha"),
                "contract": freeze.get("contract"),
                "freeze_basis": freeze.get("freeze_basis"),
            }
        ),
        "Data_Roles": _kv({k: v for k, v in roles.items() if k != "development_dates"}),
        "Execution_Binding": _kv(exec_inv),
        "Exit_Binding": _kv(exit_inv),
        "State_Machine": _kv(
            {
                "path": "FLAT→ENTRY_ALLOWED→FILL→OPEN_POSITION→EXIT_PENDING→EXIT_EXECUTED→SLOT_RELEASE",
                "event_priority": "EXIT before FILL before ADMIT",
                "occupancy": replay.get("occupancy"),
                "structural": freeze.get("structural"),
            }
        ),
        "Trades": compact or [{"empty": True}],
        "Session_Summary": list(metrics.get("sessions") or []) or [{"empty": True}],
        "PnL": _kv({k: v for k, v in metrics.items() if k not in {"sessions", "equity_tail"}}),
        "Drawdown": dd_rows or [{"empty": True}],
        "E0_E1": [
            {"entry_type": "E0", **dict(metrics.get("E0") or {})},
            {"entry_type": "E1", **dict(metrics.get("E1") or {})},
        ],
        "Exit_Reasons": [{"exit_reason": k, **v} for k, v in dict(metrics.get("exit_reasons") or {}).items()]
        or [{"empty": True}],
        "CAP": _kv(
            {
                "cap": 5,
                "max_concurrent": replay.get("max_concurrent"),
                "cap_blocked_n": replay.get("cap_blocked_n"),
                "cap_violation_n": replay.get("cap_violation_n"),
            }
        ),
        "Reentry": _kv(
            {
                "reentry_n": metrics.get("reentry_n"),
                "same_symbol_blocked_n": replay.get("same_symbol_blocked_n"),
                "winner_only_reentry": False,
                "loser_reentry_block": False,
                "rule": "after slot release if Frozen V4 still permits execution",
            }
        ),
        "Execution_Tax": _kv(
            {
                "cost_model": "X1_8BPS_EXECUTION_STRESS",
                "execution_cost_yen": metrics.get("execution_cost_yen"),
                "gross_pnl_yen": metrics.get("gross_pnl_yen"),
                "net_pnl_yen": metrics.get("net_pnl_yen"),
                "RESEARCH_EXECUTION_APPROXIMATION": True,
            }
        ),
        "Economic_Confirmation": _kv(
            {
                "confirmation_1": roles.get("economic_confirmation_1"),
                "confirmation_2": roles.get("economic_confirmation_2"),
                "status": "UNOPENED",
                "this_task": "FROZEN_READY_FOR_ECONOMIC_CONFIRMATION_ONLY",
            }
        ),
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = [
        "# PB1 V4 Complete Strategy — build and freeze",
        "",
        "Frozen PB1 ENTRY was not mutated. Evaluation unit is Complete Strategy.",
        "This task freezes structure, execution binding, technical EXIT, and the position state machine.",
        "Development PnL is binding validation only. It is not economic certification.",
        "Old Confirmation and Frozen Validation economic outcomes were not opened.",
        "",
        f"**VERDICT** `{ans.get('VERDICT')}`",
        f"**NEXT** `{ans.get('NEXT')}`",
        "",
        "## Identity",
        "",
        f"- Frozen ENTRY `{ans.get('FROZEN_ENTRY_IDENTITY')}`",
        f"- machine_sha `{ans.get('machine_sha')}`",
        f"- source_inventory_sha `{ans.get('source_inventory_sha')}`",
        f"- Complete Strategy `{ans.get('complete_strategy_identity')}`",
        f"- COMPLETE_STRATEGY_SHA256 `{ans.get('COMPLETE_STRATEGY_SHA256')}`",
        "",
        "## Execution",
        "",
        "- Live/paper source of truth remains `X1_IMMEDIATE_ASK` (first fresh Ask1 / causal Bid1).",
        "- Discovery parquet has no Bid/Ask, so development replay uses `HISTORICAL_NEXT_BAR_OPEN_EXECUTION`.",
        "- That path is labeled `RESEARCH_EXECUTION_APPROXIMATION`. X1 = 8bps stress tax, not observed spread.",
        "- Mid fill and same-bar fill are forbidden.",
        "",
        "## EXIT / state",
        "",
        "- Technical EXIT: Frozen V4 `THESIS_LOST`, then first causal next-open (lunch skipped, PM allowed).",
        "- Ops flatten: `15:20`. Not thesis death. Forced-close PnL is part of Complete Strategy.",
        "- CAP=5 from V1R runtime. Same-symbol prohibition. Occupancy +1 at FILL, slot release at EXIT fill.",
        "- Reentry only after slot release if Frozen V4 still permits execution. No winner/loser filter.",
        "",
        "## Freeze basis",
        "",
        f"- `{ans.get('freeze_basis')}`",
        "- Development PnL did not gate freeze.",
        f"- development_trade_n `{ans.get('development_trade_n')}` (not certification)",
        "",
        "STOP.",
        "",
    ]
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
