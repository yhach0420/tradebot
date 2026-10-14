"""Exact replay of the 2026-10-13 capture under the research execution model.

The saved V19 C3 ledger is the unchanged baseline. This process does not
rewrite it, and it does not send orders.
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
import time
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

NATIVE = Path(__file__).resolve().parents[3]
SRC = NATIVE / "src"
ROOT = NATIVE.parent
for item in (str(ROOT), str(NATIVE), str(SRC)):
    if item not in sys.path:
        sys.path.insert(0, item)

OUT = NATIVE / "results" / "research" / "execution_realism_layer_v1"
os.environ["TRADEBOT_C3_RESIDUAL_PATH"] = str(OUT / "replay_residual.json")

from research.current50_v16_complete_strategy_risk_foundation_v1.foundation import _capture_dir, _symbol
from research.execution_realism_layer_v1.model import ExecutionRealismBook

DAY = "20261013"
AM_WATERMARK = 6244
PM_WATERMARK = 545280
DELAYS = (0.0, 0.1, 1.0)
JST = timezone(timedelta(hours=9))
BASELINE_GROSS = -29500.0
BASELINE_N = 268


def _fmt(stamp) -> str:
    if stamp is None:
        return ""
    return datetime.fromtimestamp(float(stamp), JST).strftime("%H:%M:%S.%f")[:-3]


def _sha(path: Path) -> str:
    if not path.is_file():
        return ""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_baseline() -> list[dict]:
    rows = []
    for folder in ("live_session_085018", "live_session_122543"):
        path = NATIVE / "results" / "small_paper" / DAY / folder / "fixed_support_x1_paper_ledger.jsonl"
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            row = json.loads(line)
            if row.get("event") != "EXIT":
                continue
            hold = float(row["exit_fill_timestamp"]) - float(row["fill_timestamp"])
            rows.append(
                {
                    "signal_uid": row.get("signal_uid"),
                    "symbol": row.get("symbol"),
                    "reason": row.get("exit_reason"),
                    "pnl_yen": float(row["pnl_yen"] or 0.0),
                    "hold_sec": hold,
                    "signal_t": float(row["signal_timestamp"]),
                    "fill_t": float(row["fill_timestamp"]),
                    "exit_t": float(row["exit_fill_timestamp"]),
                    "entry_px": float(row["fill_price"]),
                    "exit_px": float(row["exit_price"]),
                    "reentry": bool(row.get("reentry")),
                    "same_event_fill": abs(float(row["signal_timestamp"]) - float(row["fill_timestamp"])) <= 1e-9,
                }
            )
    return rows


def _bucket(rows: list[dict], key: str) -> dict:
    out = {}
    for label, limit in (("lt_1s", 1.0), ("lt_5s", 5.0), ("lt_30s", 30.0), ("ge_30s", None)):
        chosen = []
        for row in rows:
            hold = row.get(key)
            if hold is None:
                continue
            if limit is None:
                if float(hold) >= 30.0:
                    chosen.append(row)
            elif float(hold) < limit:
                chosen.append(row)
        pnl = [float(item["pnl_yen"]) for item in chosen if item.get("pnl_yen") is not None]
        out[label] = {"n": len(chosen), "pnl_yen": round(sum(pnl), 6)}
    return out


def _economics(rows: list[dict]) -> dict:
    priced = [row for row in rows if row.get("pnl_yen") is not None]
    ordered = sorted(priced, key=lambda row: float(row.get("exit_fill_t") or row.get("exit_t") or 0.0))
    equity = 0.0
    peak = 0.0
    drawdown = 0.0
    for row in ordered:
        equity += float(row["pnl_yen"])
        peak = max(peak, equity)
        drawdown = min(drawdown, equity - peak)
    gains = sum(float(row["pnl_yen"]) for row in priced if float(row["pnl_yen"]) > 0)
    losses = sum(float(row["pnl_yen"]) for row in priced if float(row["pnl_yen"]) < 0)
    return {
        "closed_n": len(priced),
        "wins": sum(1 for row in priced if float(row["pnl_yen"]) > 0),
        "losses": sum(1 for row in priced if float(row["pnl_yen"]) < 0),
        "draws": sum(1 for row in priced if float(row["pnl_yen"]) == 0),
        "gross_pnl_yen": round(sum(float(row["pnl_yen"]) for row in priced), 6),
        "profit_factor": None if losses >= 0 else round(gains / abs(losses), 6),
        "exit_order_drawdown_yen": round(drawdown, 6),
        "reentry_n": sum(1 for row in priced if row.get("reentry")),
    }


def bind(kind: str, symbols: list[str], delay: float) -> ExecutionRealismBook:
    book = ExecutionRealismBook(delay)
    book.notifier = None
    membership = list(symbols)

    def admit(symbol: str, _t: float) -> tuple[bool, str]:
        if symbol in membership:
            return True, "SAME_DAY_AM_FROZEN_UNIVERSE"
        return False, "NOT_IN_DAY_FREEZE"

    book.admission = admit
    if kind == "pm":
        book.bind_safe_pm_flatten(day=DAY)
    else:
        book.bind_production_schedule(day=DAY, kind="am")
    book.ledger_path = None
    book.notifier = None
    return book


def load_universe() -> list[str]:
    path = NATIVE / "results" / "reports" / "same_day_am_frozen_universe_20261013.csv"
    symbols = []
    for line in path.read_text(encoding="utf-8").splitlines()[1:]:
        cell = line.split(",")[0].strip()
        if cell:
            symbols.append(_symbol(cell) or cell)
    return symbols


def pack(book: ExecutionRealismBook) -> dict:
    return {
        "closed": list(book.closed),
        "cancels": list(book.cancels),
        "residuals": list(book.residuals),
        "hero": list(book.hero),
        "intent_ids": list(book.intent_ids),
        "violations": list(book.violations),
        "quote_misaligned": int(book.quote_misaligned),
        "intent_n": int(book.intent_n),
        "exit_signal_n": int(book.exit_signal_n),
        "buy_full_n": int(book.buy_full_n),
        "buy_partial_n": int(book.buy_partial_n),
        "sell_full_n": int(book.sell_full_n),
        "sell_partial_n": int(book.sell_partial_n),
        "real_reentry_n": int(book.real_reentry_n),
        "peak_slots": int(book.peak_slots),
        "peak_reserved_yen": round(float(book.peak_reserved_yen), 6),
        "cap": int(book.counts.get("cap") or 0),
        "same_symbol": int(book.counts.get("same_symbol") or 0),
        "occupancy": int(book.counts.get("occupancy") or 0),
        "yen_deny": int(getattr(book, "gross_cap_denied_n", 0) or 0),
        "signal_seen": int(book.signal_seen_n),
        "admission": int(book.admission_count),
        "open_n": len(book.open_pos),
    }


def merge(left: dict, right: dict) -> dict:
    body = {}
    for key in ("closed", "cancels", "residuals", "hero", "intent_ids", "violations"):
        body[key] = list(left[key]) + list(right[key])
    for key in (
        "quote_misaligned",
        "intent_n",
        "exit_signal_n",
        "buy_full_n",
        "buy_partial_n",
        "sell_full_n",
        "sell_partial_n",
        "real_reentry_n",
        "cap",
        "same_symbol",
        "occupancy",
        "yen_deny",
        "signal_seen",
        "admission",
        "open_n",
    ):
        body[key] = int(left[key]) + int(right[key])
    body["peak_slots"] = max(int(left["peak_slots"]), int(right["peak_slots"]))
    body["peak_reserved_yen"] = round(max(float(left["peak_reserved_yen"]), float(right["peak_reserved_yen"])), 6)
    return body


def replay_delay(delay: float, symbols: list[str], capture: Path) -> dict:
    residual = OUT / "replay_residual.json"
    if residual.exists():
        residual.unlink()
    started = time.perf_counter()
    am = bind("am", symbols, delay)
    pm = bind("pm", symbols, delay)
    try:
        import orjson

        loads = orjson.loads
    except Exception:
        loads = json.loads
    lines = 0
    for part in sorted(capture.glob("push_part_*.jsonl")):
        with part.open("rb") as handle:
            for raw in handle:
                if not raw.strip():
                    continue
                lines += 1
                rec = loads(raw)
                if rec.get("kind") not in (None, "market_push"):
                    continue
                seq = int(rec.get("sequence") or 0)
                payload = rec.get("original_payload") if isinstance(rec.get("original_payload"), dict) else None
                if not payload:
                    continue
                symbol = _symbol(str(rec.get("symbol") or payload.get("Symbol") or ""))
                if not symbol:
                    continue
                received = str(rec.get("received_at") or "")
                event = {"symbol": symbol, "payload": payload, "received_at": received, "t0_push_received_at": received}
                if seq > AM_WATERMARK and not am._closed:
                    am.on_market_event(event)
                if seq > PM_WATERMARK and not pm._closed:
                    pm.on_market_event(event)
                if lines % 200000 == 0:
                    print(f"PROGRESS delay={delay} lines={lines}", flush=True)
    if not am._closed:
        am.close_at_runtime_boundary(day=DAY, kind="am")
    if not pm._closed:
        pm.close_at_runtime_boundary(day=DAY, kind="pm")
    body = merge(pack(am), pack(pm))
    body["elapsed_sec"] = round(time.perf_counter() - started, 3)
    body["assumed_order_delay_sec"] = delay
    body["delay_class"] = "ASSUMED"
    body["measured_order_latency_sec"] = None
    del am, pm
    return body


def outcome_for(uid: str, model: dict) -> str:
    if any(row.get("signal_uid") == uid for row in model["closed"]):
        return "CLOSED"
    if any(row.get("signal_uid") == uid for row in model["residuals"]):
        return "SELL_UNFILLED"
    if any(row.get("signal_uid") == uid for row in model["cancels"]):
        return "BUY_CANCELLED"
    if uid in set(model["intent_ids"]):
        return "INTENT_UNRESOLVED"
    return "NOT_ADMITTED"


def short_fate(baseline: list[dict], model: dict) -> dict:
    fate = {}
    for label, limit in (("lt_1s", 1.0), ("lt_5s", 5.0), ("lt_30s", 30.0)):
        chosen = [row for row in baseline if float(row["hold_sec"]) < limit]
        counts = Counter(outcome_for(str(row["signal_uid"]), model) for row in chosen)
        closed = [row for row in chosen if outcome_for(str(row["signal_uid"]), model) == "CLOSED"]
        model_pnl = 0.0
        for row in closed:
            match = [item for item in model["closed"] if item.get("signal_uid") == row["signal_uid"]]
            model_pnl += sum(float(item["pnl_yen"]) for item in match)
        fate[label] = {
            "baseline_n": len(chosen),
            "baseline_pnl_yen": round(sum(float(row["pnl_yen"]) for row in chosen), 6),
            "outcomes": dict(counts),
            "closed_model_pnl_yen": round(model_pnl, 6),
        }
    return fate


def uid_gap(baseline: list[dict], model: dict) -> dict:
    base = {str(row["signal_uid"]): row for row in baseline}
    modelled = {}
    for row in model["closed"]:
        modelled.setdefault(str(row["signal_uid"]), 0.0)
        modelled[str(row["signal_uid"])] += float(row["pnl_yen"])
    both = sorted(set(base) & set(modelled))
    price_gap = round(sum(modelled[uid] - float(base[uid]["pnl_yen"]) for uid in both), 6)
    missing = [uid for uid in base if uid not in set(model["intent_ids"])]
    cancelled = [row["signal_uid"] for row in model["cancels"]]
    return {
        "same_uid_n": len(both),
        "same_uid_pnl_gap_yen": price_gap,
        "baseline_not_admitted_n": len(missing),
        "baseline_not_admitted_pnl_yen": round(sum(float(base[uid]["pnl_yen"]) for uid in missing), 6),
        "buy_cancelled_n": len(model["cancels"]),
        "cancelled_uids_sample": cancelled[:12],
    }


def hero_rows(baseline: list[dict], model: dict) -> list[dict]:
    wanted = [row for row in baseline if str(row["symbol"]) == "4382.T" and float(row["hold_sec"]) < 5.0]
    traces = model["hero"]
    out = []
    for row in wanted:
        uid = str(row["signal_uid"])
        states = [item for item in traces if item.get("signal_uid") == uid]
        out.append(
            {
                "signal_uid": uid,
                "baseline_signal": _fmt(row["signal_t"]),
                "baseline_fill": _fmt(row["fill_t"]),
                "baseline_exit": _fmt(row["exit_t"]),
                "baseline_hold_sec": round(float(row["hold_sec"]), 6),
                "baseline_entry_px": row["entry_px"],
                "baseline_exit_px": row["exit_px"],
                "baseline_pnl_yen": row["pnl_yen"],
                "baseline_same_event_fill": row["same_event_fill"],
                "model_outcome": outcome_for(uid, model),
                "model_states": [
                    {"state": item["state"], "t": _fmt(item["event_t"]), "qty": item.get("filled_qty"), "reason": item.get("reason")}
                    for item in states
                ],
            }
        )
    return out


def write_xlsx(path: Path, baseline: list[dict], scenarios: list[dict], primary: dict) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "判定"
    for row in (
        ("項目", "値"),
        ("基準件数", len(baseline)),
        ("基準PnL", round(sum(float(row["pnl_yen"]) for row in baseline), 6)),
        ("研究PnL", _economics(primary["closed"])["gross_pnl_yen"]),
        ("研究決済件数", _economics(primary["closed"])["closed_n"]),
        ("買い取消", len(primary["cancels"])),
        ("売り未約定", len(primary["residuals"])),
        ("因果違反", len(primary["violations"])),
        ("板索引ずれ", primary["quote_misaligned"]),
        ("仮定遅延秒", primary["assumed_order_delay_sec"]),
        ("実測遅延", "なし"),
        ("submit/cancel/live", "0/0/0"),
    ):
        sheet.append(row)
    delays = book.create_sheet("遅延別")
    delays.append(("仮定遅延秒", "決済", "PnL", "PF", "DD", "買い全量", "買い部分", "売り全量", "売り部分", "買い取消", "売り未約定", "ピーク枠", "ピーク予約円", "同一銘柄拒否", "再ENTRY"))
    for item in scenarios:
        econ = _economics(item["closed"])
        delays.append(
            (
                item["assumed_order_delay_sec"],
                econ["closed_n"],
                econ["gross_pnl_yen"],
                econ["profit_factor"],
                econ["exit_order_drawdown_yen"],
                item["buy_full_n"],
                item["buy_partial_n"],
                item["sell_full_n"],
                item["sell_partial_n"],
                len(item["cancels"]),
                len(item["residuals"]),
                item["peak_slots"],
                item["peak_reserved_yen"],
                item["same_symbol"],
                econ["reentry_n"],
            )
        )
    shorts = book.create_sheet("短時間")
    shorts.append(("区分", "基準件数", "基準PnL", "研究で決済", "研究で取消", "研究で未約定", "研究で未採用", "研究決済PnL"))
    fate = short_fate(baseline, primary)
    for label in ("lt_1s", "lt_5s", "lt_30s"):
        row = fate[label]
        outcomes = row["outcomes"]
        shorts.append(
            (
                label,
                row["baseline_n"],
                row["baseline_pnl_yen"],
                outcomes.get("CLOSED", 0),
                outcomes.get("BUY_CANCELLED", 0),
                outcomes.get("SELL_UNFILLED", 0),
                outcomes.get("NOT_ADMITTED", 0),
                row["closed_model_pnl_yen"],
            )
        )
    hero = book.create_sheet("4382")
    hero.append(("uid", "基準シグナル", "基準約定", "基準EXIT", "保有秒", "基準PnL", "研究結果", "状態列"))
    for row in hero_rows(baseline, primary):
        hero.append(
            (
                row["signal_uid"],
                row["baseline_signal"],
                row["baseline_fill"],
                row["baseline_exit"],
                row["baseline_hold_sec"],
                row["baseline_pnl_yen"],
                row["model_outcome"],
                " | ".join(f"{item['state']}@{item['t']}" for item in row["model_states"]),
            )
        )
    fills = book.create_sheet("約定")
    fills.append(("uid", "銘柄", "理由", "数量", "エントリー", "エグジット", "保有秒", "PnL", "再ENTRY", "状態"))
    for row in primary["closed"]:
        fills.append(
            (
                row.get("signal_uid"),
                row.get("symbol"),
                row.get("reason"),
                row.get("qty"),
                row.get("entry_px"),
                row.get("exit_px"),
                None if row.get("fill_hold_sec") is None else round(float(row["fill_hold_sec"]), 6),
                row.get("pnl_yen"),
                row.get("reentry"),
                row.get("state"),
            )
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def write_md(path: Path, report: dict) -> None:
    base = report["baseline"]
    primary = report["primary"]
    gap = report["gap"]
    lines = [
        f"# {report['verdict']}",
        "",
        f"NEXT: {report['next']}",
        "",
        "## 基準と研究執行",
        "",
        f"- 現行Paper PnL: {base['gross_pnl_yen']} 円、{base['closed_n']} 件。シグナル時刻と買い約定時刻は同一。",
        f"- 執行現実性PnL: {primary['economics']['gross_pnl_yen']} 円、決済 {primary['economics']['closed_n']} 件。",
        f"- 差: {report['pnl_gap_yen']} 円。",
        f"- 同一UIDの価格差: {gap['same_uid_pnl_gap_yen']} 円（{gap['same_uid_n']} 件）。",
        f"- 基準にあって研究側が注文まで到達しなかった件数: {gap['baseline_not_admitted_n']}、その基準PnL {gap['baseline_not_admitted_pnl_yen']} 円。",
        f"- 買い取消: {len(report['primary_cancels'])}。売り未約定: {len(report['primary_residuals'])}。未約定に価格は付けていない。",
        "",
        "差は、約定を次の到達可能な板まで遅らせたこと、その間枠を占有したこと、板数量またはBuy1時刻が証明できない約定を成立させなかったことから出る。短時間の取引は削除していない。",
        "",
        "## 4382",
        "",
        "現行の1432はシグナルと買い約定が同一イベントで、9ms後の板でEXITが成立し、その解放のあと1434が再ENTRYしている。研究モデルでは、そのシグナル板では買えない。EXIT要求は保有成立後すぐに出せるが、売り約定はその後のBidが証明できたときだけである。",
        "",
    ]
    for row in report["hero_4382"]:
        if row["signal_uid"].endswith("|1432") or row["signal_uid"].endswith("|1434"):
            lines.append(
                f"- {row['signal_uid']}: 基準 {row['baseline_signal']} 約定 {row['baseline_fill']} EXIT {row['baseline_exit']} 保有 {row['baseline_hold_sec']} 秒 PnL {row['baseline_pnl_yen']}。研究結果 {row['model_outcome']}。状態 {row['model_states']}"
            )
    residual_notional = round(sum(float(row.get("notional") or 0.0) for row in report["primary_residuals"]), 6)
    lines.extend(
        [
            "",
            f"売り未約定の想定元本は {residual_notional} 円。価格が証明できないので、この分は研究PnLに入れていない。",
            "",
            "## 遅延",
            "",
            "発注往復の実測値はない。submitは0だった。下の遅延は仮定であり、0秒でもシグナルと同一の板では約定しない。",
            "",
        ]
    )
    for row in report["delay_table"]:
        lines.append(
            f"- 仮定 {row['assumed_order_delay_sec']} 秒: 決済 {row['closed_n']}、PnL {row['gross_pnl_yen']}、PF {row['profit_factor']}、DD {row['exit_order_drawdown_yen']}、取消 {row['buy_cancelled_n']}、未約定 {row['sell_unfilled_n']}"
        )
    lines.extend(
        [
            "",
            "## 短時間保有",
            "",
            "基準の保有時間はそのまま残し、各件が研究モデルでどう処理されたかを分けた。",
            "",
        ]
    )
    for label, row in report["short_fate"].items():
        lines.append(f"- {label}: 基準 {row['baseline_n']} 件 / {row['baseline_pnl_yen']} 円。内訳 {row['outcomes']}。研究側で決済できた分のPnL {row['closed_model_pnl_yen']} 円。")
    lines.extend(
        [
            "",
            "## 因果",
            "",
            f"- 違反 {report['violations_n']}。板索引のずれ {report['quote_misaligned']}。",
            f"- 決定性ハッシュ {report['determinism_sha256']}。乱数と壁時計は約定に使っていない。",
            "",
            "## 安全",
            "",
            "- V20、AUTO50、C3のENTRY/EXIT条件は変更していない。",
            "- Runtimeには載せていない。submit/cancel/live = 0/0/0。",
            "- 29日C3とV19突合の成果物は上書きしていない。",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    residual = NATIVE / "results" / "paper" / "c3_residual_ledger.json"
    mode = NATIVE / "src" / "small_paper" / "c3_paper_mode.json"
    selector = NATIVE / "src" / "small_paper" / "universe_selector_mode.json"
    before = {"residual": _sha(residual), "mode": _sha(mode), "selector": _sha(selector)}
    baseline = load_baseline()
    base_econ = _economics(baseline)
    if len(baseline) != BASELINE_N or abs(base_econ["gross_pnl_yen"] - BASELINE_GROSS) > 1e-6:
        raise SystemExit(f"BASELINE_MISMATCH n={len(baseline)} pnl={base_econ['gross_pnl_yen']}")
    capture = _capture_dir(DAY)
    if capture is None:
        raise SystemExit("NO_CAPTURE")
    symbols = load_universe()
    scenarios = []
    for delay in DELAYS:
        print(f"START delay={delay}", flush=True)
        scenarios.append(replay_delay(delay, symbols, capture))
        print(
            f"DONE delay={delay} closed={len(scenarios[-1]['closed'])} pnl={_economics(scenarios[-1]['closed'])['gross_pnl_yen']}",
            flush=True,
        )
    primary = scenarios[0]
    econ = _economics(primary["closed"])
    payload = json.dumps(primary["closed"], ensure_ascii=False, sort_keys=True, default=str)
    report = {
        "verdict": "EXECUTION_REALISM_RESEARCH_COMPLETE" if not primary["violations"] and primary["quote_misaligned"] == 0 else "EXECUTION_REPLAY_FAILED",
        "next": "RUNTIME_ADOPTION_NOT_AUTHORIZED",
        "baseline": {
            **base_econ,
            "same_event_entry_n": sum(1 for row in baseline if row["same_event_fill"]),
            "reentry_n": sum(1 for row in baseline if row["reentry"]),
            "hold_buckets": _bucket(baseline, "hold_sec"),
            "source": "saved V19 C3 ledger 20261013. parity replay was not rerun",
        },
        "primary": {"economics": econ, "assumed_order_delay_sec": 0.0, "delay_class": "ASSUMED", "measured_order_latency_sec": None},
        "primary_cancels": primary["cancels"],
        "primary_residuals": primary["residuals"],
        "pnl_gap_yen": round(econ["gross_pnl_yen"] - base_econ["gross_pnl_yen"], 6),
        "gap": uid_gap(baseline, primary),
        "short_fate": short_fate(baseline, primary),
        "fill_hold_buckets": _bucket(primary["closed"], "fill_hold_sec"),
        "delay_table": [
            {
                "assumed_order_delay_sec": item["assumed_order_delay_sec"],
                "delay_class": "ASSUMED",
                **_economics(item["closed"]),
                "buy_full_n": item["buy_full_n"],
                "buy_partial_n": item["buy_partial_n"],
                "sell_full_n": item["sell_full_n"],
                "sell_partial_n": item["sell_partial_n"],
                "buy_cancelled_n": len(item["cancels"]),
                "sell_unfilled_n": len(item["residuals"]),
                "peak_slots": item["peak_slots"],
                "peak_reserved_yen": item["peak_reserved_yen"],
                "same_symbol": item["same_symbol"],
                "occupancy": item["occupancy"],
                "cap": item["cap"],
                "yen_deny": item["yen_deny"],
                "intent_n": item["intent_n"],
                "exit_signal_n": item["exit_signal_n"],
                "violations_n": len(item["violations"]),
                "quote_misaligned": item["quote_misaligned"],
                "elapsed_sec": item["elapsed_sec"],
            }
            for item in scenarios
        ],
        "hero_4382": hero_rows(baseline, primary),
        "violations_n": len(primary["violations"]),
        "violations_sample": primary["violations"][:20],
        "quote_misaligned": primary["quote_misaligned"],
        "determinism_sha256": hashlib.sha256(payload.encode("utf-8")).hexdigest(),
        "event_time_causality": "fill only on a strictly later proven quote",
        "minimum_hold_added": False,
        "entry_exit_rules_changed": False,
        "runtime_adopted": False,
        "paper_started": False,
        "submit": 0,
        "cancel": 0,
        "live": 0,
        "protected_sha_before": before,
    }
    after = {"residual": _sha(residual), "mode": _sha(mode), "selector": _sha(selector)}
    report["protected_sha_after"] = after
    report["protected_unchanged"] = before == after
    if not report["protected_unchanged"]:
        report["verdict"] = "EXECUTION_REPLAY_FAILED"
    # The full cancel list can be large. Keep it in the json; trades too.
    report["closed_trades"] = primary["closed"]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    write_md(OUT / "report.md", report)
    write_xlsx(OUT / "audit.xlsx", baseline, scenarios, primary)
    print(report["verdict"], report["pnl_gap_yen"], report["protected_unchanged"], flush=True)


if __name__ == "__main__":
    main()
