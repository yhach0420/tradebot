"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_c0_exit_continuation_reassessment import ANALYSIS_ID
from research.am_entry_profit_improvement.publish import json_sanitize

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_c0_exit_continuation_reassessment"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "summary",
    "precommit",
    "fixed_trades",
    "extend_diag",
    "arms",
    "daily_pnl",
    "current_preservation",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    lines = [
        "# AM C0 EXIT CONTINUATION REASSESSMENT V1",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "C0 ENTRY frozen. CURRENT keeps C14. Force CONT_EXIT_600 into existing CONT_EXTEND_750 only.",
        "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "",
        "## REQUIRED FINAL OUTPUT",
        "",
        f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
        "",
        f"CONTINUATION_ARM_N: {_fmt(g.get('CONTINUATION_ARM_N'))}",
        "",
        f"E0_NET: {_fmt(g.get('E0_NET'))}",
        "",
        f"E0_PF: {_fmt(g.get('E0_PF'))}",
        "",
        f"E0_DD: {_fmt(g.get('E0_DD'))}",
        "",
        f"E1_FORCED_EXTEND_N: {_fmt(g.get('E1_FORCED_EXTEND_N'))}",
        "",
        f"E1_FIXED_DELTA_NET: {_fmt(g.get('E1_FIXED_DELTA_NET'))}",
        "",
        f"E1_OVERLAY_NET: {_fmt(g.get('E1_OVERLAY_NET'))}",
        "",
        f"E1_OVERLAY_PF: {_fmt(g.get('E1_OVERLAY_PF'))}",
        "",
        f"E1_OVERLAY_DD: {_fmt(g.get('E1_OVERLAY_DD'))}",
        "",
        f"E1_PAIRED_MEDIAN: {_fmt(g.get('E1_PAIRED_MEDIAN'))}",
        "",
        f"E1_POS_DAYS: {_fmt(g.get('E1_POS_DAYS'))}",
        "",
        f"E1_NEG_DAYS: {_fmt(g.get('E1_NEG_DAYS'))}",
        "",
        f"E1_EX_BEST: {_fmt(g.get('E1_EX_BEST'))}",
        "",
        f"E1_EX_TOP3: {_fmt(g.get('E1_EX_TOP3'))}",
        "",
        f"E2_FORCED_EXTEND_N: {_fmt(g.get('E2_FORCED_EXTEND_N'))}",
        "",
        f"E2_FIXED_DELTA_NET: {_fmt(g.get('E2_FIXED_DELTA_NET'))}",
        "",
        f"E2_OVERLAY_NET: {_fmt(g.get('E2_OVERLAY_NET'))}",
        "",
        f"E2_OVERLAY_PF: {_fmt(g.get('E2_OVERLAY_PF'))}",
        "",
        f"E2_OVERLAY_DD: {_fmt(g.get('E2_OVERLAY_DD'))}",
        "",
        f"E2_PAIRED_MEDIAN: {_fmt(g.get('E2_PAIRED_MEDIAN'))}",
        "",
        f"E2_POS_DAYS: {_fmt(g.get('E2_POS_DAYS'))}",
        "",
        f"E2_NEG_DAYS: {_fmt(g.get('E2_NEG_DAYS'))}",
        "",
        f"E2_EX_BEST: {_fmt(g.get('E2_EX_BEST'))}",
        "",
        f"E2_EX_TOP3: {_fmt(g.get('E2_EX_TOP3'))}",
        "",
        f"PASS_ARM_N: {_fmt(g.get('PASS_ARM_N'))}",
        "",
        f"BEST_PASS_ARM: {_fmt(g.get('BEST_PASS_ARM'))}",
        "",
        f"CURRENT_PRESERVATION_PASS_ALL: {_tf(g.get('CURRENT_PRESERVATION_PASS_ALL'))}",
        "",
        f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
        "",
        f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
        "",
        f"VERDICT: {g.get('VERDICT')}",
        "",
        f"NEXT: {g.get('NEXT')}",
        "",
        "## STOP",
        "",
        "Original C0 remains FROZEN_FOR_FUTURE_OOS_ONLY with C14.",
        "No PTL. No new 750 logic. Runtime adoption forbidden.",
        "submit/cancel/live=0/0/0.",
        "",
    ]
    return "\n".join(lines)
