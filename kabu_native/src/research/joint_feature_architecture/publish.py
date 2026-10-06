"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.joint_feature_architecture import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "joint_feature_architecture"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Architectures",
    "Gates",
    "Attribution",
    "Redundancy",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.8g}"
    return str(v)


def _arch_line(s: dict[str, Any] | None, title: str) -> str:
    if not s:
        return f"{title}: missing"
    g = "pass" if s.get("PASS") else "fail"
    return (
        f"{title}: MFE={_fmt(s.get('MFE_DELTA'))} / downside={_fmt(s.get('DOWNSIDE_DELTA'))} "
        f"/ joint={_fmt(s.get('JOINT_RATE'))} / gate={g}"
    )


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    by = {str(s.get("architecture_id")): s for s in (report.get("architectures") or [])}
    return "\n".join(
        [
            "# JOINT ENTRY FEATURE ARCHITECTURE REDESIGN V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine. Common 600s. Exact executable-at-decision.",
            "Direct joint label frozen. RandomForestClassifier frozen. Top3 frozen.",
            "Architectures: A0=B0, A1=B0+P, A2=B0+L, A3=B0+M, A4=B0+X, A5=B0+P+L+M+X.",
            "No individual feature search. No Exact. No PnL.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"ARCHITECTURE_N: {g.get('ARCHITECTURE_N')}",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"A0_MFE_DELTA: {_fmt(g.get('A0_MFE_DELTA'))}",
            f"A0_DOWNSIDE_DELTA: {_fmt(g.get('A0_DOWNSIDE_DELTA'))}",
            f"A0_JOINT_RATE: {_fmt(g.get('A0_JOINT_RATE'))}",
            "",
            _arch_line(by.get("A1"), "A1_PRICE_PATH"),
            _arch_line(by.get("A2"), "A2_LIQUIDITY"),
            _arch_line(by.get("A3"), "A3_MICROSTRUCTURE"),
            _arch_line(by.get("A4"), "A4_CROSS_SECTIONAL"),
            _arch_line(by.get("A5"), "A5_FULL"),
            "",
            f"PASSING_ARCHITECTURES: {g.get('PASSING_ARCHITECTURES')}",
            f"PASSING_ARCHITECTURE_N: {g.get('PASSING_ARCHITECTURE_N')}",
            "",
            f"PRIMARY_INFORMATION_SOURCE: {g.get('PRIMARY_INFORMATION_SOURCE')}",
            f"INTERACTION_DEPENDENT: {_tf(g.get('INTERACTION_DEPENDENT'))}",
            f"BEST_DIAGNOSTIC_ARCHITECTURE: {g.get('BEST_DIAGNOSTIC_ARCHITECTURE')}",
            "",
            f"MAX_FEATURE_CORRELATION: {_fmt(g.get('MAX_FEATURE_CORRELATION'))}",
            f"FUTURE_EVENT_USE_N: {g.get('FUTURE_EVENT_USE_N')}",
            f"TARGET_CONTAMINATION_N: {g.get('TARGET_CONTAMINATION_N')}",
            "",
            f"FEATURE_ARCHITECTURE_PASS: {_tf(g.get('FEATURE_ARCHITECTURE_PASS'))}",
            "",
            f"PRIMARY_FINDING: {g.get('PRIMARY_FINDING')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Exact. No runtime candidate. No sequence model. No feature selection rerun.",
            "Runtime/C14 unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
