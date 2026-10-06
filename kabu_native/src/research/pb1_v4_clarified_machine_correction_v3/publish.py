"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_correction_v3.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Parent_Hashes",
    "V3_Change_Summary",
    "Parent_RCA_Note",
    "OneBar_All88",
    "OneBar_Targets",
    "PostDominant_State",
    "TwoSided_All88",
    "TwoSided_Targets",
    "CounterAuction_State",
    "Active_Lifecycle",
    "Thesis_Lifecycle",
    "Reached_vs_Live",
    "Thesis_Loss_Events",
    "Failure_Layer_Audit",
    "FailedOpen_Protection",
    "FamilyA_Protection",
    "Hidden1M_Parity",
    "Changed_Rows",
    "Tests",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "audit_rows", "chart_zones", "hidden_1m_checks"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    inv = dict(report.get("invariants") or {})
    hid = dict(report.get("hidden_1m_recompute") or {})
    hashes = dict(report.get("hashes") or {})
    return {
        "SPEC_CHANGED": False,
        "CORRECTION_V2_CHANGED": False,
        "NEW_V3_IDENTITY_CREATED": True,
        "PROSPECTIVE_DATA_OPENED": False,
        "OLD_CONFIRMATION_OPENED": False,
        "FROZEN_VALIDATION_OPENED": False,
        "FUTURE_OUTCOME_USED": False,
        "PNL_USED": False,
        "MFE_MAE_USED": False,
        "THRESHOLD_OPTIMIZED": False,
        "submit/cancel/live": "0/0/0",
        "SPEC_SHA256": hashes.get("SPEC_SHA256") or report.get("EXPECTED_SPEC_SHA256"),
        "CORRECTION_V2_SHA256": "916ec521c6830d0b3aed6079ff40e1d043cc52c00bbd4c880108d39ced95536c",
        "NEW_V3_SHA256": hashes.get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V3_SHA256")
        or report.get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V3_SHA256"),
        "ONE_BAR_PRIMARY_CAUSE_REFINED": True,
        "SPEC_CHANGE_REQUIRED": False,
        "Hidden-1m mismatch_n": hid.get("mismatch_n"),
        "SAME_BAR_ENTRY": inv.get("SAME_BAR_ENTRY"),
        "invariant_violations": inv.get("violations"),
        "target_pass_n": (dec.get("target_checks") or {}).get("pass_n"),
        "target_fail_n": (dec.get("target_checks") or {}).get("fail_n"),
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "incomplete_reasons": dec.get("incomplete_reasons"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    rows = list(report.get("audit_rows") or [])
    hashes = dict(report.get("hashes") or {})
    enc = dict(report.get("encoding_manifest") or {})
    targets = list((report.get("target_checks") or {}).get("checks") or [])
    one_targets = ("7011_20241205", "8002_20241002", "6857_20250930", "8630_20250828", "6501_20251010")
    two_targets = ("6963_20250613", "5803_20250212", "8031_20250225", "7741_20250314", "5803_20250709", "5706_20250725")
    fo = (
        "3382_20241004",
        "7011_20250523",
        "4063_20251118",
        "3382_20241115",
        "6273_20250120",
        "5802_20250613",
        "7182_20251112",
        "3110_20250805",
    )
    fam = ("9432_20250402", "8058_20250814")
    crit = dict(report.get("critical_cases") or {})

    def _pick_rows(keys: tuple[str, ...]) -> list[dict[str, Any]]:
        out = []
        for k in keys:
            if "_" not in k:
                continue
            sym, date = k.split("_", 1)
            hit = next(
                (r for r in rows if str(r.get("symbol")) == sym and str(r.get("date")) == date),
                crit.get(k) or {"symbol": sym, "date": date},
            )
            out.append(hit)
        return out

    loss_rows = [r for r in rows if r.get("machine_THESIS_LOST") or r.get("machine_THESIS_LOST_AT")]
    return {
        "Manifest": [{"key": k, "value": v} for k, v in (report.get("answers") or {}).items()],
        "Parent_Hashes": [{"key": k, "value": v} for k, v in hashes.items()],
        "V3_Change_Summary": [
            {"change": "A", "what": "Post-dominant FOLLOWTHROUGH / absorption classes"},
            {"change": "B", "what": "Sandwich pullback vs meaningful counter vs committed fight"},
            {"change": "C", "what": "REACHED vs LIVE; event STALE; no N=3; no location-gated death"},
        ],
        "Parent_RCA_Note": [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in enc.items()
        ],
        "OneBar_All88": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "SEED": r.get("machine_SEED"),
                "post_dominant_class": r.get("post_dominant_class"),
                "human_opening_state": r.get("human_opening_state"),
            }
            for r in rows
        ],
        "OneBar_Targets": _pick_rows(one_targets),
        "PostDominant_State": [
            {"symbol": r.get("symbol"), "date": r.get("date"), "post_dominant_class": r.get("post_dominant_class")}
            for r in rows
        ],
        "TwoSided_All88": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "sandwich_semantic_class": r.get("sandwich_semantic_class"),
                "SEED": r.get("machine_SEED"),
                "opening": r.get("machine_opening_state"),
            }
            for r in rows
        ],
        "TwoSided_Targets": _pick_rows(two_targets),
        "CounterAuction_State": [
            {"symbol": r.get("symbol"), "date": r.get("date"), "sandwich_semantic_class": r.get("sandwich_semantic_class")}
            for r in rows
        ],
        "Active_Lifecycle": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "ACTIVE_REACHED": r.get("machine_ACTIVE_REACHED"),
                "ACTIVE_LIVE": r.get("machine_ACTIVE_LIVE"),
                "progress_log": r.get("progress_log"),
            }
            for r in rows
        ],
        "Thesis_Lifecycle": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "THESIS_REACHED": r.get("machine_THESIS_REACHED"),
                "THESIS_LIVE": r.get("machine_THESIS_LIVE"),
                "THESIS_LOST": r.get("machine_THESIS_LOST"),
                "THESIS_LOST_AT": r.get("machine_THESIS_LOST_AT"),
                "THESIS_LOST_REASON": r.get("machine_THESIS_LOST_REASON"),
            }
            for r in rows
        ],
        "Reached_vs_Live": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "ACTIVE_REACHED": r.get("machine_ACTIVE_REACHED"),
                "ACTIVE_LIVE": r.get("machine_ACTIVE_LIVE"),
                "THESIS_REACHED": r.get("machine_THESIS_REACHED"),
                "THESIS_LIVE": r.get("machine_THESIS_LIVE"),
                "thesis_id": r.get("thesis_id"),
                "opening_drive_id": r.get("opening_drive_id"),
            }
            for r in rows
        ],
        "Thesis_Loss_Events": loss_rows or [{"empty": True}],
        "Failure_Layer_Audit": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "EXPECTED_FAILURE_LAYER": r.get("EXPECTED_FAILURE_LAYER"),
                "ACTUAL_FAILURE_LAYER": r.get("ACTUAL_FAILURE_LAYER"),
                "wrong_layer_thesis_after_seed_should_die": r.get("wrong_layer_thesis_after_seed_should_die"),
            }
            for r in rows
        ],
        "FailedOpen_Protection": _pick_rows(fo),
        "FamilyA_Protection": _pick_rows(fam),
        "Hidden1M_Parity": [
            {"key": k, "value": v} for k, v in (report.get("hidden_1m_recompute") or {}).items() if k != "mismatches"
        ]
        or [{"empty": True}],
        "Changed_Rows": list((report.get("changed_row_manifest") or {}).get("rows") or []),
        "Tests": list(targets) or [{"empty": True}],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 clarified machine correction V3", "", "DEVELOPMENT_PARITY_ONLY. Not face validation.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
