"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.mtf_5min_sma5_25_75_with_1min_trigger_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "5M_Bar_Semantics",
    "Live_Causal_MA",
    "Confirmed_MA",
    "Trend_Structure",
    "SMA25_Pullback",
    "SMA5_Trigger",
    "SMA75_Failure",
    "A_Primary",
    "B_Control",
    "C_NoPullback",
    "Matched_AB",
    "Matched_AC",
    "D1",
    "D2",
    "D3",
    "D4",
    "MFE_MAE",
    "Structural_Risk",
    "VWAP_Diagnostic",
    "SR_Diagnostic",
    "TV_Diagnostic",
    "Timeframe_Comparison",
    "Causality_Audit",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def _block_rows(pack: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for b in ("D2", "D3", "D4"):
        rows.append({"block": b, **dict((pack.get("blocks") or {}).get(b) or {})})
    return rows or [{"block": None}]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    freeze = dict(report.get("freeze") or {})
    s = dict(report.get("setups") or {})
    a = dict(s.get("A") or {})
    ab = dict(report.get("a_vs_b") or {})
    ac = dict(report.get("a_vs_c") or {})
    audit = dict(report.get("bars5_audit") or {})
    block_rows = {}
    for b in ("D1", "D2", "D3", "D4"):
        block_rows[b] = _kv_rows(dict((s.get("by_block") or {}).get(b) or {}))
    return {
        "Binding": _kv_rows(
            {
                "ok": (report.get("bind") or {}).get("ok"),
                "parent_verdict": (report.get("bind") or {}).get("parent_verdict"),
                "one_min_sma_closed": True,
                "peer_rescue": False,
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
            }
        ),
        "5M_Bar_Semantics": _kv_rows(
            {
                "constructed_from_1m": True,
                "lunch_synthetic_n": audit.get("lunch_synthetic_n"),
                "overnight_synthetic_n": audit.get("overnight_synthetic_n"),
                "sma_reset_at_open": False,
                "ma_history_spans_sessions": True,
                "wait_for_5m_completion": False,
                "FUTURE_5M_CLOSE_USAGE_N": report.get("FUTURE_5M_CLOSE_USAGE_N"),
                "completed_5m_n": audit.get("completed_5m_n"),
            }
        ),
        "Live_Causal_MA": _kv_rows(
            {
                "definition": freeze.get("primary_ma"),
                "live_sma_defined_n": audit.get("live_sma_defined_n"),
                "provisional_close": "current completed 1m close",
                "chosen_by_performance": False,
            }
        ),
        "Confirmed_MA": _kv_rows(
            {
                "definition": freeze.get("confirmed_ma"),
                "confirmed_sma_defined_n": audit.get("confirmed_sma_defined_n"),
                "live_eq_confirmed_n": audit.get("live_eq_confirmed_n"),
                "live_ne_confirmed_n": audit.get("live_ne_confirmed_n"),
                "agree_rate_on_A": report.get("confirmed_agree_rate"),
                "confirmed_a_vs_b": report.get("confirmed_a_vs_b"),
                "chose_better": False,
            }
        ),
        "Trend_Structure": _kv_rows({"trend": freeze.get("trend"), "slope": freeze.get("slope"), "MTF_TREND_ALIGNED": True}),
        "SMA25_Pullback": _kv_rows(
            {
                "zone": freeze.get("sma25_zone"),
                "zone_atr_mult": freeze.get("zone_atr_mult"),
                "adds_all_eval": (report.get("sma25_pullback") or {}).get("adds_all_eval"),
                "D2": dict((ac.get("blocks") or {}).get("D2") or {}),
                "D3": dict((ac.get("blocks") or {}).get("D3") or {}),
                "D4": dict((ac.get("blocks") or {}).get("D4") or {}),
            }
        ),
        "SMA5_Trigger": _kv_rows({"trigger": freeze.get("trigger"), "reclaim_n": a.get("reclaim_n"), "swing_break_n": a.get("swing_break_n")}),
        "SMA75_Failure": _kv_rows({"role": freeze.get("sma75_role"), "sma25_failed_n": a.get("sma25_failed_n"), "time_to_sma25_fail": a.get("time_to_sma25_fail")}),
        "A_Primary": _kv_rows(s.get("A") or {}),
        "B_Control": _kv_rows(s.get("B") or {}),
        "C_NoPullback": _kv_rows(s.get("C") or {}),
        "Matched_AB": _block_rows(ab),
        "Matched_AC": _block_rows(ac),
        "D1": block_rows.get("D1") or [{"key": "empty"}],
        "D2": block_rows.get("D2") or [{"key": "empty"}],
        "D3": block_rows.get("D3") or [{"key": "empty"}],
        "D4": block_rows.get("D4") or [{"key": "empty"}],
        "MFE_MAE": _kv_rows({"mfe": a.get("mfe"), "mae": a.get("mae"), "time_to_mfe": a.get("time_to_mfe"), "time_to_invalidation": a.get("time_to_invalidation")}),
        "Structural_Risk": _kv_rows({"risk_bps": a.get("risk_bps"), "mfe_over_risk": a.get("mfe_over_risk"), "time_to_sma25_fail": a.get("time_to_sma25_fail")}),
        "VWAP_Diagnostic": _kv_rows(report.get("vwap") or {}),
        "SR_Diagnostic": _kv_rows(report.get("sr") or {}),
        "TV_Diagnostic": _kv_rows(report.get("tv") or {}),
        "Timeframe_Comparison": _kv_rows(report.get("timeframe_comparison") or {}),
        "Causality_Audit": _kv_rows(
            {
                **audit,
                "same_bar_entry_n": report.get("same_bar_entry_n"),
                "FUTURE_SETUP_SELECTION_N": report.get("FUTURE_SETUP_SELECTION_N"),
                "FUTURE_5M_CLOSE_USAGE_N": report.get("FUTURE_5M_CLOSE_USAGE_N"),
            }
        ),
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# MTF_5MIN_SMA5_25_75_WITH_1MIN_TRIGGER_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "5-minute live-causal SMA5/25/75 for trend structure. 1-minute price action for timing.",
            "Closes the 1-minute SMA architecture. Not a period search. Not a peer rescue. Not a Complete Strategy.",
            "",
            f"5m bars constructed causally? **{a.get('5m bars constructed causally?')}**",
            f"Any future 5m close used? **{a.get('Any future 5m close used?')}**",
            f"Does MA history span sessions? **{a.get('Does MA history span sessions?')}**",
            f"Any lunch synthetic bars? **{a.get('Any lunch synthetic bars?')}**",
            f"Live causal MA semantics? **{a.get('Live causal MA semantics?')}**",
            f"A setup_n? **{a.get('A setup_n?')}** bull/bear? **{a.get('bull/bear?')}**",
            f"D1/D2/D3/D4? **{a.get('D1/D2/D3/D4?')}**",
            f"A vs B matched 10m D2/D3/D4? **{a.get('A vs B matched 10m:')}**",
            f"A vs C D2/D3/D4? **{a.get('A vs C:')}**",
            f"Does 5m SMA structure add to same 1m trigger? **{a.get('Does 5m SMA structure add to same 1m trigger?')}**",
            f"Does SMA25 pullback add inside 5m aligned trend? **{a.get('Does SMA25 pullback add inside 5m aligned trend?')}**",
            f"Next-open MFE/MAE? **{a.get('Next-open MFE/MAE?')}**",
            f"Structural MFE/risk? **{a.get('Structural MFE/risk?')}**",
            f"How much movement before entry? **{a.get('How much movement before entry?')}**",
            f"Confirmed agrees with live-causal? **{a.get('Does confirmed-5m MA diagnostic agree with live-causal semantics?')}**",
            f"Differs from previous 1m SMA test? **{a.get('Does result differ materially from previous 1m SMA test?')}**",
            f"MA period opt? **{a.get('Any MA period optimization?')}** aux opt? **{a.get('Any auxiliary filter optimization?')}** PnL opt? **{a.get('Any PnL optimization?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}**",
            f"submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
