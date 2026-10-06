"""report.json / report.md / audit.xlsx only. No raw rewrite."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_context_day1_lead_lag_rca_v1 import ANALYSIS_ID
from research.futures_context_day1_lead_lag_rca_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "Offsets",
    "Clocks",
    "FourState",
    "Reversal",
    "Exec",
    "Preopen",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    primary = dict(body.get("primary_row") or {})
    p5 = next((r for r in (body.get("offset_rows") or []) if r.get("offset_sec") == 300), {})
    execu = dict(body.get("absolute_exec") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"TYPE: {d.get('TYPE')}",
        f"NEXT: `{d.get('NEXT')}`",
        f"reason: {d.get('reason')}",
        "",
        "Exploratory RCA. Day2 frozen effect-check specification is unchanged.",
        "Positive offsets are NONCAUSAL_DIAGNOSTIC and are not candidate-strategy features.",
        "Statistical unit is CLOCK. Sign split is feature>0 vs feature<0 only.",
        "",
        "## Primary row",
        "",
        f"- feature: `{primary.get('feature')}` horizon `{primary.get('horizon')}`",
        f"- Day1 causal 0m shift: {_fmt(primary.get('day1_causal_shift_bps'))} bps",
        f"- Day1 +5m placebo: {_fmt(primary.get('day1_placebo_p5m_bps'))} bps",
        f"- recomputed 0m: {_fmt(primary.get('recomputed_0m_shift_bps'))} bps",
        f"- recomputed +5m: {_fmt(primary.get('recomputed_p5m_shift_bps'))} bps",
        "",
        "## Required answers",
        "",
        f"- 1_offset_-5m_MID_shift: `{_fmt(a.get('1_offset_-5m_MID_shift'))}`",
        f"- 2_offset_-3m: `{_fmt(a.get('2_offset_-3m'))}`",
        f"- 3_offset_-1m: `{_fmt(a.get('3_offset_-1m'))}`",
        f"- 4_offset_0m: `{_fmt(a.get('4_offset_0m'))}`",
        f"- 5_offset_+1m: `{_fmt(a.get('5_offset_+1m'))}`  NONCAUSAL_DIAGNOSTIC",
        f"- 6_offset_+3m: `{_fmt(a.get('6_offset_+3m'))}`  NONCAUSAL_DIAGNOSTIC",
        f"- 7_offset_+5m: `{_fmt(a.get('7_offset_+5m'))}`  NONCAUSAL_DIAGNOSTIC",
        f"- 8_peak_offset: `{a.get('8_peak_offset')}`",
        f"- 9_causal_only_offsets_same_direction: `{_fmt(a.get('9_causal_only_offsets_same_direction'))}`",
        f"- 10_stock_PAST_RET_180S_relation: shift=`{_fmt((a.get('10_stock_PAST_RET_180S_relation') or {}).get('mid_shift_bps'))}` "
        f"pos={((a.get('10_stock_PAST_RET_180S_relation') or {}).get('pos_clock_n'))} "
        f"neg={((a.get('10_stock_PAST_RET_180S_relation') or {}).get('neg_clock_n'))}",
        f"- 11_NK180_adds_beyond_stock_prior: `{json.dumps(a.get('11_NK180_adds_beyond_stock_prior'), ensure_ascii=False, default=str)}`",
        f"- 12_NK_up_STK_down_future: `{json.dumps(a.get('12_NK_up_STK_down_future'), ensure_ascii=False, default=str)}`",
        f"- 13_NK_down_STK_up_future: `{json.dumps(a.get('13_NK_down_STK_up_future'), ensure_ascii=False, default=str)}`",
        f"- 14_30_60_inversion_explained_by_reversal: `{_fmt(a.get('14_30_60_inversion_explained_by_reversal'))}`",
        f"- 15_180s_explained_by_trend_persistence: `{_fmt(a.get('15_180s_explained_by_trend_persistence'))}`",
        f"- 16_placebo_superiority_due_temporal_overlap: `{_fmt(a.get('16_placebo_superiority_due_temporal_overlap'))}`",
        f"- 17_absolute_LONG_positive: `{_fmt(a.get('17_absolute_LONG_positive'))}`",
        f"- 18_absolute_SHORT_positive: `{_fmt(a.get('18_absolute_SHORT_positive'))}`",
        f"- 19_mechanism_classification: `{a.get('19_mechanism_classification')}`",
        f"- 20_exact_mechanism_sentence: {a.get('20_exact_mechanism_sentence')}",
        f"- 21_usable_ENTRY_thesis_exists: `{_fmt(a.get('21_usable_ENTRY_thesis_exists'))}`",
        f"- 22_Day2_frozen_test_changed: `{_fmt(a.get('22_Day2_frozen_test_changed'))}`",
        f"- 23_Runtime_changed: `{_fmt(a.get('23_Runtime_changed'))}`",
        f"- 24_Paper_changed: `{_fmt(a.get('24_Paper_changed'))}`",
        f"- 25_submit_cancel_live: `{a.get('25_submit_cancel_live')}`",
        "",
        "## Lead/lag map (clock-level 10m MID, NK_RET_180S at T+offset)",
        "",
        "| offset | causal? | pos/neg | MID shift bps | NK window vs T | overlap with stock 10m | feature fully inside outcome |",
        "|---|---|---|---:|---|---:|---|",
    ]
    for r in body.get("offset_rows") or []:
        win = r.get("feature_window_sec_from_T")
        lines.append(
            "| {lab} | {cau} | {pn} | {sh} | [{a},{b}]s | {ov}s ({frac}) | {inside} |".format(
                lab=r.get("label"),
                cau="causal" if not r.get("noncausal_diagnostic") else "NONCAUSAL_DIAGNOSTIC",
                pn=f"{r.get('pos_clock_n')}/{r.get('neg_clock_n')}",
                sh=_fmt(r.get("mid_shift_bps")),
                a=win[0] if isinstance(win, list) and win else "",
                b=win[1] if isinstance(win, list) and len(win) > 1 else "",
                ov=r.get("overlap_sec"),
                frac=_fmt(r.get("overlap_frac_of_outcome"), 2),
                inside=_fmt(r.get("feature_fully_inside_outcome")),
            )
        )
    long_p = dict(execu.get("NK180_pos_LONG") or {})
    short_n = dict(execu.get("NK180_neg_SHORT") or {})
    lines.extend(
        [
            "",
            "## +5m placebo overlap",
            "",
            f"- NK_RET_180S at T+5m uses futures in [{p5.get('feature_window_sec_from_T')}] seconds from T.",
            f"- Stock 10m MID outcome is [0, 600] seconds from T.",
            f"- overlap_sec={p5.get('overlap_sec')} feature_fully_inside_outcome={p5.get('feature_fully_inside_outcome')}.",
            "- If the feature window is inside the outcome window, placebo strength is contemporaneous overlap, not forecast skill.",
            "",
            "## Absolute executability (clock unit)",
            "",
            f"- NK180>0 clocks LONG mean/median: {_fmt(long_p.get('mean'))} / {_fmt(long_p.get('median'))} (n={long_p.get('n')})",
            f"- NK180<0 clocks SHORT mean/median: {_fmt(short_n.get('mean'))} / {_fmt(short_n.get('median'))} (n={short_n.get('n')})",
            f"- note: `{execu.get('note')}`",
            "",
            f"PREOPEN_DIRECTIONAL_FOLLOWING: `{_fmt(body.get('PREOPEN_DIRECTIONAL_FOLLOWING'))}` (Day1 denied if false).",
            "",
            "ENTRY: false",
            "EXIT: false",
            "Day2 frozen test changed: false",
            "submit/cancel/live: 0/0/0",
            "",
        ]
    )
    return "\n".join(lines)


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    report = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(report), encoding="utf-8")
    four = []
    for name, block in (report.get("four_state") or {}).items():
        four.append({"state": name, **(block if isinstance(block, dict) else {"value": block})})
    wb = Workbook()
    first = True
    sheets = {
        "Answers": _kv(dict(report.get("answers") or {})),
        "Offsets": list(report.get("offset_rows") or []),
        "Clocks": list(report.get("clocks") or []),
        "FourState": four,
        "Reversal": _kv(dict(report.get("reversal_vs_persistence") or {})),
        "Exec": _kv(dict(report.get("absolute_exec") or {})),
        "Preopen": _kv(dict(report.get("preopen") or {})),
        "Decision": _kv(dict(report.get("decision") or {})),
        "Safety": [
            {
                "submit_cancel_live": (report.get("answers") or {}).get("25_submit_cancel_live"),
                "ENTRY": False,
                "EXIT": False,
                "Day2_frozen_test_changed": False,
                "Runtime_changed": False,
                "Paper_changed": False,
            }
        ],
    }
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {"report_json": str(OUT / "report.json"), "report_md": str(OUT / "report.md"), "audit_xlsx": str(xlsx)}
