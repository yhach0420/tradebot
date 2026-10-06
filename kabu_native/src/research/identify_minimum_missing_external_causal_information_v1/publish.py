"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.identify_minimum_missing_external_causal_information_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "HM1_Anchor",
    "External_Source_Audit",
    "Futures_Semantics",
    "Historical_Availability",
    "Prospective_Availability",
    "NK_Features",
    "TOPIX_Features",
    "NK_TOPIX_Divergence",
    "HM1_FirstMove",
    "Incremental_Value",
    "Matched_Comparison",
    "D1_D4",
    "Proxy_Track",
    "Prospective_Track",
    "Next_External_Decision",
    "Live_20260914",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    bind = dict(report.get("bind") or {})
    freeze = dict(report.get("external_join_anchor") or {})
    hm1 = dict(freeze.get("hm1") or {})
    audit = dict(report.get("source_audit") or {})
    sem = dict(report.get("futures_semantics") or {})
    inc = dict(report.get("incremental") or {})
    live = dict(report.get("live_20260914") or {})
    proxy = dict(report.get("proxy_track") or {})
    nxt = dict(report.get("next_external_decision") or {})
    nk_sign = dict(inc.get("nk_sign") or {})
    tx_sign = dict(inc.get("topix_sign") or {})
    states = dict(inc.get("agreement_states") or {})
    base = dict(inc.get("base") or {})
    return {
        "Binding": _kv(
            {
                "parent_verdict": report.get("parent_verdict_accepted"),
                "bind_ok": bind.get("ok"),
                "hm1_tuned": report.get("hm1_tuned"),
                "old_confirmation_used_to_design": report.get("old_confirmation_used_to_design"),
                "kabu_50": report.get("kabu_50_applied"),
                "frozen_validation_opened": False,
                "interpretation_narrow": report.get("interpretation_narrow"),
            }
        ),
        "HM1_Anchor": _kv(hm1) + _kv({"hm2": freeze.get("hm2"), "identities_frozen_before_external_join": freeze.get("identities_frozen_before_external_join")}),
        "External_Source_Audit": list(audit.get("futures_endpoint_probes") or [{"empty": True}])
        + _kv({"datacube_local_file_n": audit.get("datacube_local_file_n"), "did_not_purchase": audit.get("did_not_purchase")}),
        "Futures_Semantics": _kv({k: v for k, v in sem.items() if k != "days"})
        + [{"day": d, **row} for d, row in dict(sem.get("days") or {}).items()],
        "Historical_Availability": _kv(
            {
                "nk": audit.get("true_historical_nk225mini_minute"),
                "topix": audit.get("true_historical_topix_futures_minute"),
                "additional_purchase_required": audit.get("additional_purchase_required"),
                "jquants_futures_minute_available": audit.get("jquants_futures_minute_available"),
                "jquants_futures_daily_available": audit.get("jquants_futures_daily_available"),
            }
        )
        + list(audit.get("datacube_catalog") or []),
        "Prospective_Availability": _kv(dict(report.get("prospective_track") or {})),
        "NK_Features": _kv(nk_sign) + list(inc.get("nk_ret_60s_quintiles") or [{"empty": True}]),
        "TOPIX_Features": _kv(tx_sign) + list(inc.get("topix_ret_60s_quintiles") or [{"empty": True}]),
        "NK_TOPIX_Divergence": [{"state": k, **(v if isinstance(v, dict) else {"value": v})} for k, v in states.items()]
        or [{"empty": True}],
        "HM1_FirstMove": _kv(
            {
                "base_p_mfe_before_mae": base.get("p_mfe_before_mae"),
                "nk_pos_p": (nk_sign.get("pos") or {}).get("p_mfe_before_mae"),
                "nk_neg_p": (nk_sign.get("neg") or {}).get("p_mfe_before_mae"),
                "delta_p": nk_sign.get("delta_p_first"),
            }
        ),
        "Incremental_Value": _kv(
            {
                "base_n": base.get("n"),
                "base_x0": base.get("mean_x0_bps"),
                "base_pf": base.get("profit_factor"),
                "nk_delta_x0": nk_sign.get("delta_x0"),
                "topix_delta_x0": tx_sign.get("delta_x0"),
                "structure": (inc.get("structure_precommitted_gate") or {}).get("structure_exists"),
                "role": inc.get("simple_causal_role"),
                "duplicate": (inc.get("duplicate_of_stock_momentum") or {}).get("flag"),
            }
        ),
        "Matched_Comparison": _kv(dict(inc.get("matched_comparison") or {})) + _kv(dict(inc.get("duplicate_of_stock_momentum") or {})),
        "D1_D4": _kv(
            {
                "base_blocks": base.get("block_mean_x0"),
                "nk_pos_blocks": (nk_sign.get("pos") or {}).get("block_mean_x0"),
                "nk_neg_blocks": (nk_sign.get("neg") or {}).get("block_mean_x0"),
                "stable": inc.get("d1_d4_stable"),
                "block_signs": (inc.get("structure_precommitted_gate") or {}).get("block_signs"),
            }
        ),
        "Proxy_Track": _kv(proxy),
        "Prospective_Track": _kv(dict(report.get("prospective_track") or {})),
        "Next_External_Decision": _kv(nxt),
        "Live_20260914": _kv(
            {
                "futures_pid": live.get("futures_pid"),
                "futures_alive": live.get("futures_alive"),
                "nk225mini_latest_timestamp": live.get("nk225mini_latest_timestamp"),
                "nk_row_count": live.get("nk225mini_row_count"),
                "topix_latest_timestamp": live.get("topix_latest_timestamp"),
                "topix_row_count": live.get("topix_row_count"),
                "breadth_pid": live.get("breadth_pid"),
                "breadth_alive": live.get("breadth_alive"),
                "breadth_latest_snapshot": live.get("breadth_latest_snapshot"),
                "breadth_snapshot_count": live.get("breadth_snapshot_count"),
                "submit_cancel_live": live.get("submit_cancel_live"),
                "recovery": live.get("recovery"),
            }
        ),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    live = dict(a.get("live_20260914") or {})
    return "\n".join(
        [
            "# IDENTIFY_MINIMUM_MISSING_EXTERNAL_CAUSAL_INFORMATION_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            d.get("INTERPRETATION") or "",
            "",
            "Accepted parent: `STOCK_PANEL_MAGNITUDE_LIMIT_CONFIRMED_V1` as CURRENT_STOCK_PANEL_DIRECTIONAL_EDGE_DOES_NOT_CLEAR_EXECUTION_SCALE, not zero stock value.",
            "",
            f"HM1 frozen? **{a.get('hm1_stock_side_anchor_frozen')}**",
            f"Anchor SHA: `{a.get('anchor_sha')}`",
            f"Historical NK225mini minute available? **{a.get('historical_nk225mini_minute_available')}** source `{a.get('nk_source')}`",
            f"Historical TOPIX futures minute available? **{a.get('historical_topix_futures_minute_available')}** source `{a.get('topix_source')}`",
            f"Additional purchase required? **{a.get('additional_purchase_required')}** (did not purchase: **{a.get('did_not_purchase')}**)",
            f"Proxies labeled PROXY_NOT_FUTURES? **{a.get('proxies_explicitly_labeled_proxy')}** `{a.get('historical_proxies_available')}`",
            f"Base P(MFE before MAE): **{a.get('base_mfe_before_mae_probability')}**",
            f"External-conditioned P(first): **{a.get('external_conditioned_probability')}**",
            f"Base HM1 Top1 X0: **{a.get('base_hm1_top1_x0')}**",
            f"External-conditioned X0: **{a.get('external_conditioned_x0')}**",
            f"Delta X0: **{a.get('delta_x0')}**",
            f"D1–D4 stable? **{a.get('improvement_stable_across_d1_d4')}**",
            f"NK incremental X0: `{a.get('nk_incremental_value')}` TOPIX `{a.get('topix_incremental_value')}`",
            f"Duplicate of stock/sector momentum? **{a.get('futures_merely_duplicated_stock_sector_momentum')}**",
            f"Role: `{a.get('simple_causal_role')}` GATE={a.get('DIRECTION_GATE')} VETO={a.get('VETO')} REGIME={a.get('REGIME')} TIE={a.get('TIE_BREAKER')}",
            f"NK-up vs NK-down X0 (descriptive, not a promoted filter): `{a.get('nk_up_x0')}` / `{a.get('nk_down_x0')}` first-move `{a.get('nk_up_p_mfe_before_mae')}` / `{a.get('nk_down_p_mfe_before_mae')}`",
            f"Enough for external-conditioned complete strategy? **{a.get('enough_evidence_for_external_conditioned_complete_strategy')}**",
            f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
            f"Old Confirmation used to design? **{a.get('old_confirmation_used_to_design')}**",
            f"Kabu 50? **{a.get('kabu_50_applied')}**",
            f"20260914 Futures PID `{live.get('futures_pid')}` alive **{live.get('futures_alive')}** NK `{live.get('nk225mini_latest_timestamp')}` n={live.get('nk_row_count')} TOPIX `{live.get('topix_latest_timestamp')}` n={live.get('topix_row_count')}",
            f"20260914 Breadth PID `{live.get('breadth_pid')}` alive **{live.get('breadth_alive')}** snapshot `{live.get('breadth_latest_snapshot')}` count={live.get('breadth_snapshot_count')}",
            f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
            "",
            "STOP. Do not return to technical threshold mining.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
