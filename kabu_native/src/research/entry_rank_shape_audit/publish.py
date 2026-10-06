"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.entry_rank_shape_audit import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "entry_rank_shape_audit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "NestedParity",
    "NestedFolds",
    "FixedSpecs",
    "NestedVsFixed",
    "SpecStability",
    "RankShape",
    "Top1Robust",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# CANONICAL C3 FAILURE MECHANISM AUDIT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine only. TARGET V4 M4_PERSISTENT 600s. CURRENT IRREGULAR CLOCK.",
            "Existing 27 Ridge specs. No new feature. No new model. No Top1 strategy.",
            "No Exact. No PnL. No execution-aware. Rank then join TARGET. NO TARGET BACKFILL.",
            "BEST_FIXED_DIAGNOSTIC_SPEC is diagnostic only and is not adopted.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"NESTED_OOF_PARITY: {_tf(g.get('NESTED_OOF_PARITY'))}",
            "",
            f"FIXED_SPECS_N: {g.get('FIXED_SPECS_N')}",
            f"FIXED_SPECS_PASSING_ALL_AK_N: {g.get('FIXED_SPECS_PASSING_ALL_AK_N')}",
            f"FIXED_SPECS_TOP3_POSITIVE_N: {g.get('FIXED_SPECS_TOP3_POSITIVE_N')}",
            f"FIXED_SPECS_TOP3_GT_CURRENT_N: {g.get('FIXED_SPECS_TOP3_GT_CURRENT_N')}",
            "",
            f"BEST_FIXED_DIAGNOSTIC_SPEC: {g.get('BEST_FIXED_DIAGNOSTIC_SPEC')}",
            f"BEST_FIXED_DIAGNOSTIC_TOP3_DELTA: {g.get('BEST_FIXED_DIAGNOSTIC_TOP3_DELTA')}",
            f"MEDIAN_FIXED_TOP3_DELTA: {g.get('MEDIAN_FIXED_TOP3_DELTA')}",
            "",
            f"INNER_OUTER_TOP3_CORRELATION: {g.get('INNER_OUTER_TOP3_CORRELATION')}",
            f"INNER_OUTER_DELTA_CORRELATION: {g.get('INNER_OUTER_DELTA_CORRELATION')}",
            "",
            f"MOST_COMMON_SPEC: {g.get('MOST_COMMON_SPEC')}",
            f"MOST_COMMON_SPEC_SHARE: {g.get('MOST_COMMON_SPEC_SHARE')}",
            f"SPEC_SELECTION_STABLE: {_tf(g.get('SPEC_SELECTION_STABLE'))}",
            "",
            f"C3_RANK1_TARGET: {g.get('C3_RANK1_TARGET')}",
            f"C3_RANK2_TARGET: {g.get('C3_RANK2_TARGET')}",
            f"C3_RANK3_TARGET: {g.get('C3_RANK3_TARGET')}",
            f"C3_RANK4_TARGET: {g.get('C3_RANK4_TARGET')}",
            f"C3_RANK5_TARGET: {g.get('C3_RANK5_TARGET')}",
            "",
            f"C3_RANK2_3_MEAN: {g.get('C3_RANK2_3_MEAN')}",
            f"C3_RANK4_5_MEAN: {g.get('C3_RANK4_5_MEAN')}",
            "",
            f"CURRENT_RANK1_TARGET: {g.get('CURRENT_RANK1_TARGET')}",
            f"CURRENT_RANK2_3_MEAN: {g.get('CURRENT_RANK2_3_MEAN')}",
            f"CURRENT_RANK4_5_MEAN: {g.get('CURRENT_RANK4_5_MEAN')}",
            "",
            f"TOP1_DELTA_MEAN: {g.get('TOP1_DELTA_MEAN')}",
            f"TOP1_DELTA_MEDIAN: {g.get('TOP1_DELTA_MEDIAN')}",
            f"TOP1_POSITIVE_DAYS: {g.get('TOP1_POSITIVE_DAYS')}",
            f"TOP1_NEGATIVE_DAYS: {g.get('TOP1_NEGATIVE_DAYS')}",
            f"TOP1_EX_BEST_DAY: {g.get('TOP1_EX_BEST_DAY')}",
            f"TOP1_EX_TOP3_DAYS: {g.get('TOP1_EX_TOP3_DAYS')}",
            "",
            f"TOP1_EDGE_ROBUST: {_tf(g.get('TOP1_EDGE_ROBUST'))}",
            "",
            f"PRIMARY_FAILURE_MECHANISM: {g.get('PRIMARY_FAILURE_MECHANISM')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Top1 strategy. No Exact. No C4. No execution-aware model.",
            "Runtime/C14 unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
