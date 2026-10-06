"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.wait5_session_target_learnability import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "wait5_session_target_learnability"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "AMFillability",
    "AMU",
    "AMD",
    "PMFillability",
    "Representations",
    "ConsensusDays",
    "Gates",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# WAIT5 SESSION TARGET LEARNABILITY V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM / PM are independent research tasks. 9 representations. Median/consensus only.",
            "No deployable strategy. Runtime WAIT_SEC remains 1.0. No common AM/PM model.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"REPRESENTATION_N: {g.get('REPRESENTATION_N')}",
            "",
            f"AM_FILLABILITY_TOP3_DELTA: {_fmt(g.get('AM_FILLABILITY_TOP3_DELTA'))}",
            f"AM_FILLABILITY_POSITIVE_REP_N: {g.get('AM_FILLABILITY_POSITIVE_REP_N')}",
            f"AM_FILLABILITY_LEARNABLE: {_tf(g.get('AM_FILLABILITY_LEARNABLE'))}",
            "",
            f"AM_U_FILL_SPEARMAN: {_fmt(g.get('AM_U_FILL_SPEARMAN'))}",
            f"AM_U_POSITIVE_REP_N: {g.get('AM_U_POSITIVE_REP_N')}",
            f"AM_U_FILL_LEARNABLE: {_tf(g.get('AM_U_FILL_LEARNABLE'))}",
            "",
            f"AM_D_FILL_SPEARMAN: {_fmt(g.get('AM_D_FILL_SPEARMAN'))}",
            f"AM_D_POSITIVE_REP_N: {g.get('AM_D_POSITIVE_REP_N')}",
            f"AM_D_FILL_LEARNABLE: {_tf(g.get('AM_D_FILL_LEARNABLE'))}",
            "",
            f"AM_TARGET_STATUS: {g.get('AM_TARGET_STATUS')}",
            f"AM_VERDICT: {g.get('AM_VERDICT')}",
            "",
            f"PM_FILLABILITY_TOP1_DELTA: {_fmt(g.get('PM_FILLABILITY_TOP1_DELTA'))}",
            f"PM_FILLABILITY_POSITIVE_REP_N: {g.get('PM_FILLABILITY_POSITIVE_REP_N')}",
            f"PM_FILLABILITY_LEARNABLE: {_tf(g.get('PM_FILLABILITY_LEARNABLE'))}",
            "",
            f"PM_TARGET_STATUS: {g.get('PM_TARGET_STATUS')}",
            f"PM_VERDICT: {g.get('PM_VERDICT')}",
            "",
            f"COMMON_AM_PM_MODEL_ALLOWED: {_tf(g.get('COMMON_AM_PM_MODEL_ALLOWED'))}",
            f"COMMON_AM_PM_TARGET_ALLOWED: {_tf(g.get('COMMON_AM_PM_TARGET_ALLOWED'))}",
            "",
            f"AM_TRAIN_PM_ROW_N: {g.get('AM_TRAIN_PM_ROW_N')}",
            f"PM_TRAIN_AM_ROW_N: {g.get('PM_TRAIN_AM_ROW_N')}",
            f"FUTURE_EVENT_USE_N: {g.get('FUTURE_EVENT_USE_N')}",
            f"TARGET_CONTAMINATION_N: {g.get('TARGET_CONTAMINATION_N')}",
            f"HELDOUT_FIT_LEAK_N: {g.get('HELDOUT_FIT_LEAK_N')}",
            "",
            f"AM_NEXT: {g.get('AM_NEXT')}",
            f"PM_NEXT: {g.get('PM_NEXT')}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"OVERALL_VERDICT: {g.get('OVERALL_VERDICT')}",
            "",
            "## DECISION",
            "",
            f"AM_CASE={dec.get('AM_CASE')} PM_CASE={dec.get('PM_CASE')}",
            str(dec.get("PRIMARY_FINDING") or ""),
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No next model. No AM/PM commonization. W5 not adopted into Runtime.",
            "Runtime WAIT_SEC remains 1.0. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
