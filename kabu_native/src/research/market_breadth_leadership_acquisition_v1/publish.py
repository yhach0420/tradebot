"""report.json / report.md / audit.xlsx. Probe dumps stay under probe_YYYYMMDD."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.market_breadth_leadership_acquisition_v1 import ANALYSIS_ID
from research.market_breadth_leadership_acquisition_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = ("Answers", "Schemas", "SoftLimit", "Decision", "Safety")


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _fmt(x: Any) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, (dict, list, tuple)):
        return json.dumps(x, ensure_ascii=False, default=str)
    return str(x)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"CASE: {d.get('CASE')}",
        f"NEXT: `{d.get('NEXT')}`",
        "",
        "kabu STATION GET /ranking leadership tape. Not monitor-48 recompute.",
        "Not ADVANCE_DECLINE_RATIO. Top-ranking list = LEADERSHIP / SURGE / SECTOR STATE.",
        "No /register, /unregister, sendorder. Standard Paper 50 unchanged.",
        "20260911 ranking backfill forbidden. Day2 futures frozen test unchanged.",
        "First live day is transport proof only. No alpha claim. No strategy.",
        "",
        "## Required answers",
        "",
        f"- 1_apisoftlimit_reachable: `{_fmt(a.get('1_apisoftlimit_reachable'))}`",
        f"- 2_soft_limit_response: `{_fmt(a.get('2_soft_limit_response'))}`",
        f"- 3_ranking_reachable: `{_fmt(a.get('3_ranking_reachable'))}`",
        f"- 4_Type1_schema: `{_fmt(a.get('4_Type1_schema'))}`",
        f"- 5_Type2_schema: `{_fmt(a.get('5_Type2_schema'))}`",
        f"- 6_Type5_schema: `{_fmt(a.get('6_Type5_schema'))}`",
        f"- 7_Type6_schema: `{_fmt(a.get('7_Type6_schema'))}`",
        f"- 8_Type7_schema: `{_fmt(a.get('8_Type7_schema'))}`",
        f"- 9_Type14_schema: `{_fmt(a.get('9_Type14_schema'))}`",
        f"- 10_Type15_schema: `{_fmt(a.get('10_Type15_schema'))}`",
        f"- 11_registration_mutation_n: `{_fmt(a.get('11_registration_mutation_n'))}`",
        f"- 12_unregister_n: `{_fmt(a.get('12_unregister_n'))}`",
        f"- 13_sendorder_n: `{_fmt(a.get('13_sendorder_n'))}`",
        f"- 14_standard_Paper_config_changed: `{_fmt(a.get('14_standard_Paper_config_changed'))}`",
        f"- 15_proposed_safe_cadence: `{_fmt(a.get('15_proposed_safe_cadence'))}`",
        f"- 16_collector_path: `{_fmt(a.get('16_collector_path'))}`",
        f"- 17_raw_root: `{_fmt(a.get('17_raw_root'))}`",
        f"- 18_tests_passed: `{_fmt(a.get('18_tests_passed'))}`",
        f"- 19_VERDICT: `{_fmt(a.get('19_VERDICT'))}`",
        f"- 20_NEXT: `{_fmt(a.get('20_NEXT'))}`",
        "",
        f"submit/cancel/live: `{body.get('submit_cancel_live')}`",
        "ENTRY: false",
        "EXIT: false",
        "Day2 frozen price-return test changed: false",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    report = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(report), encoding="utf-8")
    probe = dict(report.get("probe") or {})
    schema_rows = []
    for k, v in (probe.get("schemas") or {}).items():
        row = {"type": k}
        if isinstance(v, dict):
            row.update(v)
        schema_rows.append(row)
    wb = Workbook()
    first = True
    sheets = {
        "Answers": _kv(dict(report.get("answers") or {})),
        "Schemas": schema_rows or [{"empty": True}],
        "SoftLimit": _kv(dict((report.get("answers") or {}).get("2_soft_limit_response") or {})),
        "Decision": _kv(dict(report.get("decision") or {})),
        "Safety": [
            {
                "submit_cancel_live": report.get("submit_cancel_live"),
                "register_n": (report.get("answers") or {}).get("11_registration_mutation_n"),
                "unregister_n": (report.get("answers") or {}).get("12_unregister_n"),
                "sendorder_n": (report.get("answers") or {}).get("13_sendorder_n"),
                "Paper_changed": (report.get("answers") or {}).get("14_standard_Paper_config_changed"),
                "ENTRY": False,
                "EXIT": False,
                "Day2_frozen_price_return_test_changed": False,
                "Runtime_changed": False,
            }
        ],
    }
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {"report_json": str(OUT / "report.json"), "report_md": str(OUT / "report.md"), "audit_xlsx": str(xlsx)}
