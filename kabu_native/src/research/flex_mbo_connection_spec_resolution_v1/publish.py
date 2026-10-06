"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.flex_mbo_connection_spec_resolution_v1 import ANALYSIS_ID
from research.flex_mbo_connection_spec_resolution_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "access_boundary",
    "local_spec_search",
    "document_identity",
    "message_catalog",
    "event_order",
    "gap_recovery",
    "historical_pcap",
    "session_bootstrap",
    "order_lifecycle",
    "book_reconstruction",
    "replay_contract",
    "timestamp_boundary",
    "F8_F9",
    "legal_boundary",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    search = dict(report.get("local_spec_search") or {})
    search_sheet = _kv({k: v for k, v in search.items() if k not in {"CANDIDATES", "ELIGIBLE"}})
    search_sheet.extend(list(search.get("CANDIDATES") or []))
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "parent": _kv(dict(report.get("parent") or {})),
        "access_boundary": _kv(dict(report.get("access_boundary") or {})),
        "local_spec_search": search_sheet,
        "document_identity": _kv(dict(report.get("document_identity") or {})),
        "message_catalog": _kv(dict(report.get("message_catalog") or {})),
        "event_order": _kv(dict(report.get("event_order") or {})),
        "gap_recovery": _kv(dict(report.get("gap_recovery") or {})),
        "historical_pcap": _kv(dict(report.get("historical_pcap") or {})),
        "session_bootstrap": _kv(dict(report.get("session_bootstrap") or {})),
        "order_lifecycle": _kv(dict(report.get("order_lifecycle") or {})),
        "book_reconstruction": _kv(dict(report.get("book_reconstruction") or {})),
        "replay_contract": _kv(dict(report.get("replay_contract") or {})),
        "timestamp_boundary": _kv(dict(report.get("timestamp_boundary") or {})),
        "F8_F9": _kv(dict(report.get("F8_F9") or {})),
        "legal_boundary": _kv(dict(report.get("legal_boundary") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            "STRATEGY_RESEARCH_STATUS: PAUSED",
            "RESTART_ALLOWED: false",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- eligible official spec N: {a.get('6_eligible_official_spec_N')}",
            f"- connection spec available: {a.get('10_connection_spec_available')}",
            f"- protocol resolution possible: {a.get('11_protocol_resolution_possible')}",
            f"- F8: {a.get('34_F8')}",
            f"- F9: {a.get('35_F9')}",
            f"- F12 resolved: {a.get('39_F12_resolved')}",
            f"- F13 resolved: {a.get('40_F13_resolved')}",
            f"- account created: {a.get('12_account_created')}",
            f"- provider contacted: {a.get('13_provider_contacted')}",
            "",
            "No market data purchased. No collector. No strategy. Do not consume 20260907+.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
