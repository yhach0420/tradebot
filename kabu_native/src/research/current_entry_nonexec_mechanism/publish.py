"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.current_entry_nonexec_mechanism import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "current_entry_nonexec_mechanism"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Data_Manifest",
    "A0_Parity",
    "A_Rank_Audit",
    "A_Flow",
    "A0_vs_A2",
    "A2_Opportunity_Cost",
    "B0_vs_B2_Daily",
    "B2_Trade_Decomposition",
    "Mechanism",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                obj = float(obj)
            elif isinstance(obj, np.integer):
                return int(obj)
    except Exception:
        pass
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    q = report.get("answers") or {}
    return "\n".join(
        [
            "# CURRENT ENTRY NON-EXECUTABILITY MECHANISM AUDIT",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE Exact Dual-Lane. CURRENT IRREGULAR CLOCK. Occupancy approximation was not used.",
            "B2 formal verdict is UNCHANGED (POST_HOC_DIAGNOSTIC_ONLY).",
            "C2_OOF_RANKING_EDGE_NOT_STABLE maintained. C3 not started.",
            "C14 / Runtime / CLOCK / ENTRY / EXIT / Fill SoT unchanged.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A0_PARITY: {g.get('A0_PARITY')}",
            f"A_NONEXEC_ALL_RATE: {g.get('A_NONEXEC_ALL_RATE')}",
            f"A_NONEXEC_TOP1_RATE: {g.get('A_NONEXEC_TOP1_RATE')}",
            f"A_NONEXEC_TOP3_RATE: {g.get('A_NONEXEC_TOP3_RATE')}",
            f"A_NONEXEC_TOP3_ENRICHMENT: {g.get('A_NONEXEC_TOP3_ENRICHMENT')}",
            f"A_PENDING_NONEXEC_RATE: {g.get('A_PENDING_NONEXEC_RATE')}",
            f"A_FILL_FROM_NONEXEC_T0_N: {g.get('A_FILL_FROM_NONEXEC_T0_N')}",
            f"A_CURRENT_ENTRY_NONEXEC_RANK_BIAS: {g.get('A_CURRENT_ENTRY_NONEXEC_RANK_BIAS')}",
            f"A2_TESTED: {g.get('A2_TESTED')}",
            f"A2_TRADES: {g.get('A2_TRADES')}",
            f"A2_PNL: {g.get('A2_PNL')}",
            f"A2_PF: {g.get('A2_PF')}",
            f"A2_MAXDD: {g.get('A2_MAXDD')}",
            f"A2_POSITIVE_DAY_RATE: {g.get('A2_POSITIVE_DAY_RATE')}",
            f"A2_MEDIAN_DAILY_PNL: {g.get('A2_MEDIAN_DAILY_PNL')}",
            f"A2_WOULD_DROP_FILL_N: {g.get('A2_WOULD_DROP_FILL_N')}",
            f"A2_WOULD_DROP_FILL_PNL: {g.get('A2_WOULD_DROP_FILL_PNL')}",
            f"B2_ADDED_TRADES_N: {g.get('B2_ADDED_TRADES_N')}",
            f"B2_ADDED_FIRST_PNL: {g.get('B2_ADDED_FIRST_PNL')}",
            f"B2_ADDED_REENTRY_PNL: {g.get('B2_ADDED_REENTRY_PNL')}",
            f"B2_REMOVED_TRADES_N: {g.get('B2_REMOVED_TRADES_N')}",
            f"B2_DELTA_POSITIVE_DAYS: {g.get('B2_DELTA_POSITIVE_DAYS')}",
            f"B2_DELTA_NEGATIVE_DAYS: {g.get('B2_DELTA_NEGATIVE_DAYS')}",
            f"B2_MEDIAN_DAY_DELTA: {g.get('B2_MEDIAN_DAY_DELTA')}",
            f"B2_PAIRED_ANALYSIS_CLASSIFICATION: {g.get('B2_PAIRED_ANALYSIS_CLASSIFICATION')}",
            f"NONEXEC_BIAS_SCOPE: {g.get('NONEXEC_BIAS_SCOPE')}",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"RECOMMENDED_NEXT_RESEARCH: {g.get('RECOMMENDED_NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## Q1–Q5",
            "",
            f"Q1: {q.get('Q1')}",
            f"Q2: {q.get('Q2')}",
            f"Q3: {q.get('Q3')}",
            f"Q4: {q.get('Q4')}",
            f"Q5: {q.get('Q5')}",
            "",
            "## STOP",
            "",
            "Runtime not changed. C14 not changed. C2 not changed. B2 verdict not rewritten.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
