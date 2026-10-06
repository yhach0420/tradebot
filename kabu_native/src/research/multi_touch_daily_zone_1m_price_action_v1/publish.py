"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.multi_touch_daily_zone_1m_price_action_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Daily_Bars",
    "Reaction_Point_Definition",
    "Confirmed_Reactions",
    "Zone_Construction",
    "Zone_Availability",
    "Zone_Strength",
    "Zone_State",
    "Resistance_Interactions",
    "Support_Interactions",
    "Break_Accept",
    "Retest_Hold",
    "Support_Resistance_Flip",
    "Limit_Order_Simulation",
    "Limit_Adverse_Selection",
    "Participation",
    "Anchored_VWAP",
    "Volume_Profile_Approx",
    "PDH_Control",
    "Path_Outcomes",
    "D1_D4",
    "Symbol_Coverage",
    "Candidate_Playbooks",
    "Complete_Strategy",
    "Causality_Audit",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    zc = dict(report.get("zone_construction") or {})
    cmp = dict(report.get("comparisons") or {})
    lim = dict((report.get("replay") or {}).get("limit") or {})
    packs = list((report.get("replay") or {}).get("packs") or [])
    return {
        "Binding": _kv({"parent": report.get("parent_verdict_accepted"), "rca_deferred": report.get("rca_deferred"), "pdh_control": report.get("pdh_control")}),
        "Daily_Bars": _kv({"source": "native 1-minute aggregated to session OHLC", "ATR20": "completed prior sessions only"}),
        "Reaction_Point_Definition": _kv(
            {
                "resistance": "same-day close in lower 40% of range after tagging high; range >= 0.30*ATR20",
                "support": "same-day close in upper 40% of range after tagging low",
                "confirmed_at": "that session",
                "available_from": "next session",
                "future_pivot": False,
            }
        ),
        "Confirmed_Reactions": _kv({"reaction_n": zc.get("reaction_n")}),
        "Zone_Construction": _kv(zc),
        "Zone_Availability": _kv({"ZONE_ACTIVATED_AT": "available_from of 2nd distinct-day reaction", "lookback": 60, "half_width": "0.15*ATR20"}),
        "Zone_Strength": _kv({"touch_count_at_T": "only touches with available_from <= T", "future_touches_used": False}),
        "Zone_State": _kv({"roles": "RESISTANCE_ZONE / SUPPORT_ZONE", "flip": "tracked, not erased on first break"}),
        "Resistance_Interactions": _kv({k: v for k, v in dict(report.get("interaction_counts") or {}).items()}),
        "Support_Interactions": _kv(cmp.get("single_break") or {}),
        "Break_Accept": _kv(cmp.get("mt_break") or {}) + _kv({"_": "accept"}) + _kv(cmp.get("mt_accept") or {}),
        "Retest_Hold": _kv(cmp.get("mt_retest") or {}),
        "Support_Resistance_Flip": _kv(cmp.get("flip") or {}),
        "Limit_Order_Simulation": _kv({k: lim.get(k) for k in ("order_n", "fill_n", "fill_rate", "no_fill_rate", "X0_after_fill_d2d4", "X1_8BPS_STRESS_d2d4", "8bps_is_actual_passive_cost")}),
        "Limit_Adverse_Selection": _kv(dict(lim.get("adverse_selection") or {})),
        "Participation": _kv(dict(report.get("participation") or {})),
        "Anchored_VWAP": _kv({"SESSION_VWAP": "context on events", "ANCHORED_VWAP_FROM_ZONE_BREAK": "not used as S/R truth", "MINUTE_BAR_VOLUME_PROFILE_APPROXIMATION": "not claimed as VAP"}),
        "Volume_Profile_Approx": _kv({"label": "MINUTE_BAR_VOLUME_PROFILE_APPROXIMATION", "used_as_truth": False}),
        "PDH_Control": _kv(cmp.get("pdh") or {}),
        "Path_Outcomes": list(cmp.get("pairs") or []),
        "D1_D4": list(report.get("d1_d4") or []),
        "Symbol_Coverage": _kv({"pool": 105, "prefiltered_by_symbol_pnl": False, "symbol_zone_identity_integrity": True}),
        "Candidate_Playbooks": packs,
        "Complete_Strategy": _kv({"justified_x1_ids": (report.get("replay") or {}).get("justified_x1_ids"), "lunch": "HOLD_THROUGH_LUNCH_RESUME_PM", "cap": 3}),
        "Causality_Audit": _kv(dict(report.get("causality_audit") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0"}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# MULTI_TOUCH_DAILY_ZONE_1M_PRICE_ACTION_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            d.get("INTERPRETATION") or "",
            "",
            "REFERENCE_LEVEL_MECHANISM_RCA_V1 deferred. PB_PDH_BREAK preserved as control.",
            "",
            f"Confirmed reactions: **{a.get('How_many_confirmed_daily_reaction_points')}**",
            f"Active resistance zone-days: **{a.get('How_many_active_resistance_zones')}** support zone-days: **{a.get('support_zones')}**",
            f"2-touch / 3-touch / 4+: **{a.get('How_many_2_touch')}** / **{a.get('How_many_3_touch')}** / **{a.get('How_many_4plus_touch')}**",
            f"Zones active only after required reactions knowable? **{a.get('Were_zones_active_only_after_all_required_reactions_historically_knowable')}**",
            f"Retroactive zone creation? **{a.get('Any_retroactive_zone_creation')}** Future pivot leakage? **{a.get('Any_future_pivot_leakage')}**",
            f"Symbol-specific? **{a.get('Are_zones_symbol_specific')}**",
            f"Multi-touch break vs PDH: `{a.get('Does_multi_touch_resistance_break_beat_simple_PDH_break')}`",
            f"Retest hold vs break-only: `{a.get('Does_break_retest_hold_beat_break_only')}`",
            f"Touch-count monotonic? **{a.get('Does_touch_count_add_monotonic_information')}**",
            f"Passive limit: `{a.get('passive_retest_entry')}`",
            f"Complete-strategy X1>0 if justified? **{a.get('Any_Complete_Strategy_X1_gt_0_if_justified')}**",
            f"Old Confirmation opened? **{a.get('Old_Confirmation_opened')}** Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
            f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
