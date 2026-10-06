"""Write report.json / report.md / audit.xlsx only under branch_u_temporal_holdout_v1/."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook, load_workbook

from research.am_entry_profit_improvement.publish import json_sanitize
from research.simple_tech_entry_family.publish import _sheet, kv_rows
from research.simple_tech_redesign.branch_u_holdout_spec import ANALYSIS_ID
from research.simple_tech_redesign.isolation import BRANCH_U_HOLDOUT_OUT

SHEET_ORDER = (
    "summary",
    "integrity",
    "provenance",
    "trades",
    "path",
    "core_added",
    "day_robustness",
    "series",
    "decision",
)

REQUIRED_KEYS = (
    "SIGNAL_N",
    "HOLDOUT_DAYS",
    "TESTED_BRANCH_U_EXIT_MECHANISMS",
    "CONTROL_COUNTS",
    "TREATMENT_COUNTS",
    "incremental_trade_n",
    "control_only_trade_n",
    "CONTROL_ECONOMICS",
    "TREATMENT_ECONOMICS",
    "ATTRIBUTION",
    "WINNER_SAFETY_OCCUPANCY",
    "GOOD_BRANCH_U_EXIT_N",
    "DIP_BRANCH_U_EXIT_N",
    "DAY_ROBUSTNESS",
    "PROVENANCE",
    "DELTA_TOTAL_PNL",
    "DELTA_PF",
    "DELTA_MAX_DD",
    "OCCUPANCY_SOT_PARITY",
    "PRE_DIVERGENCE_PARITY",
    "FORCE_TREATMENT_FILL_SET",
    "TRUE_OOS",
    "HOLDOUT_CLASS",
    "NON_INTERFERENCE_PASS",
    "VERDICT",
    "NEXT",
)


def _trade_sheet_row(arm: str, t: dict[str, Any]) -> dict[str, Any]:
    return {
        "arm": arm,
        "trade_id": t.get("trade_id"),
        "date": t.get("date"),
        "symbol": t.get("symbol"),
        "t0": t.get("t0"),
        "fill_role": t.get("fill_role"),
        "path_type": t.get("path_type"),
        "fill_time": t.get("fill_time"),
        "fill_price": t.get("fill_price"),
        "exit_time": t.get("exit_time"),
        "exit_price": t.get("exit_price"),
        "exit_reason": t.get("exit_reason"),
        "pnl_yen_100": t.get("pnl_yen_100"),
        "exit_miss": t.get("exit_miss"),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    BRANCH_U_HOLDOUT_OUT.mkdir(parents=True, exist_ok=True)
    extra = [p for p in BRANCH_U_HOLDOUT_OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    for p in extra:
        p.unlink()
    body = {k: v for k, v in report.items() if k != "_markdown"}
    summary = dict(body.get("summary") or {})
    summary.pop("trades_control", None)
    summary.pop("trades_treatment", None)
    body["summary"] = summary
    (BRANCH_U_HOLDOUT_OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (BRANCH_U_HOLDOUT_OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(BRANCH_U_HOLDOUT_OUT / "audit.xlsx")


def build_sheets(report: dict[str, Any], leak: dict[str, Any], reporting: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    req = dict(report.get("required") or {})
    summary = dict(report.get("summary") or {})
    attr = dict(summary.get("ATTRIBUTION") or {})
    occ = dict(summary.get("WINNER_SAFETY_OCCUPANCY") or {})
    un = dict(summary.get("WINNER_SAFETY_UNCONSTRAINED") or {})
    path_rows = [
        {"scope": "DIRECT_EXIT_DELTA", "value": attr.get("DIRECT_EXIT_DELTA")},
        {"scope": "SLOT_RELEASE_DOWNSTREAM_DELTA", "value": attr.get("SLOT_RELEASE_DOWNSTREAM_DELTA")},
        {"scope": "DISPLACED_TRADE_DELTA", "value": attr.get("DISPLACED_TRADE_DELTA")},
        {"scope": "TOTAL_CAUSAL_DELTA", "value": attr.get("TOTAL_CAUSAL_DELTA")},
        {"scope": "OCC_GOOD_U", "value": occ.get("GOOD_BRANCH_U_EXIT_N")},
        {"scope": "OCC_DIP_U", "value": occ.get("DIP_BRANCH_U_EXIT_N")},
        {"scope": "OCC_EARLY_U", "value": occ.get("EARLY_BRANCH_U_EXIT_N")},
        {"scope": "OCC_PTF_U", "value": occ.get("PTF_BRANCH_U_EXIT_N")},
        {"scope": "OCC_OTHER_U", "value": occ.get("OTHER_BRANCH_U_EXIT_N")},
        {"scope": "UNCONST_GOOD_U", "value": un.get("GOOD_BRANCH_U_EXIT_N")},
        {"scope": "UNCONST_DIP_U", "value": un.get("DIP_BRANCH_U_EXIT_N")},
        {"scope": "UNCONST_EARLY_U", "value": un.get("EARLY_BRANCH_U_EXIT_N")},
    ]
    core_added = [
        {"scope": "OVERALL", "arm": "control", **dict(summary.get("CONTROL_ECONOMICS") or {})},
        {"scope": "OVERALL", "arm": "treatment", **dict(summary.get("TREATMENT_ECONOMICS") or {})},
        {"scope": "CORE", "arm": "control", **dict(summary.get("CORE_CONTROL_ECONOMICS") or {})},
        {"scope": "CORE", "arm": "treatment", **dict(summary.get("CORE_TREATMENT_ECONOMICS") or {})},
        {"scope": "ADDED", "arm": "control", **dict(summary.get("ADDED_CONTROL_ECONOMICS") or {})},
        {"scope": "ADDED", "arm": "treatment", **dict(summary.get("ADDED_TREATMENT_ECONOMICS") or {})},
        {"scope": "INCREMENTAL", "arm": "treatment", **dict(summary.get("INCREMENTAL_ECONOMICS") or {})},
    ]
    ctrl_tr = list(summary.get("trades_control") or [])
    treat_tr = list(summary.get("trades_treatment") or [])
    trade_rows = [_trade_sheet_row("control", t) for t in ctrl_tr] + [_trade_sheet_row("treatment", t) for t in treat_tr]
    day_rows = list((summary.get("DAY_ROBUSTNESS") or {}).get("daily") or [])
    return {
        "summary": kv_rows(req),
        "integrity": kv_rows({**dict(leak), **dict(reporting), "NON_INTERFERENCE_PASS": req.get("NON_INTERFERENCE_PASS")}),
        "provenance": list(summary.get("PROVENANCE") or []) or [{"empty": True}],
        "trades": trade_rows or [{"empty": True}],
        "path": path_rows,
        "core_added": core_added,
        "day_robustness": day_rows or [{"empty": True}],
        "series": list(report.get("series_history") or []) or [{"empty": True}],
        "decision": kv_rows(dict(report.get("decision") or {})),
    }


def _fmt(v: Any) -> str:
    if v is None:
        return "None"
    if isinstance(v, float):
        return f"{v:.4f}" if abs(v) < 1000 else f"{v:,.2f}"
    return str(v)


def _count_line(label: str, pack: dict[str, Any]) -> str:
    return (
        f"{label}: signal=`{pack.get('signal_n')}` candidate=`{pack.get('candidate_n')}` "
        f"accepted=`{pack.get('accepted_entry_n')}` fill=`{pack.get('fill_n')}` "
        f"Branch_U_exit=`{pack.get('Branch_U_exit_n')}` session_close=`{pack.get('session_close_exit_n')}` "
        f"slot_release=`{pack.get('slot_release_n')}` CAP_reject=`{pack.get('CAP_reject_n')}` "
        f"same_symbol_reject=`{pack.get('same_symbol_reject_n')}`"
    )


def build_markdown(report: dict[str, Any]) -> str:
    req = dict(report.get("required") or {})
    dec = dict(report.get("decision") or {})
    summary = dict(report.get("summary") or {})
    oc = dict(summary.get("CONTROL_ECONOMICS") or {})
    ot = dict(summary.get("TREATMENT_ECONOMICS") or {})
    attr = dict(summary.get("ATTRIBUTION") or {})
    occ = dict(summary.get("WINNER_SAFETY_OCCUPANCY") or {})
    day = dict(summary.get("DAY_ROBUSTNESS") or {})
    cc = dict(summary.get("CONTROL_COUNTS") or {})
    tc = dict(summary.get("TREATMENT_COUNTS") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{req.get('VERDICT')}**",
        f"CASE: `{dec.get('CASE')}`",
        f"HOLDOUT_CLASS: `{req.get('HOLDOUT_CLASS')}` TRUE_OOS=`{req.get('TRUE_OOS')}` CERTIFIED=`{dec.get('CERTIFIED')}`",
        f"NEXT: {req.get('NEXT')}",
        "",
        "Frozen Branch U only. Control = occupancy + session-close. Treatment = Control + UNPROVEN `U_BB_LOWER_BREAK`. "
        "All complete sealed AM captures after 20260827. Today excluded. No PnL date filter. No Branch P. No sizing.",
        "",
        f"HOLDOUT_DAYS = `{req.get('HOLDOUT_DAYS')}`",
        f"APPENDED_DAYS = `{req.get('APPENDED_DAYS')}` DATE_CHERRY_PICK = `{req.get('DATE_CHERRY_PICK')}`",
        f"NEXT_CAPTURE = `{req.get('NEXT_CAPTURE_CANDIDATE')}` status=`{req.get('NEXT_CAPTURE_STATUS')}`",
        f"OCCUPANCY_SOT_PARITY = `{req.get('OCCUPANCY_SOT_PARITY')}` PRE_DIVERGENCE_PARITY = `{req.get('PRE_DIVERGENCE_PARITY')}`",
        _count_line("CONTROL", cc),
        _count_line("TREATMENT", tc),
        f"incremental_trade_n = `{req.get('incremental_trade_n')}` control_only_trade_n = `{req.get('control_only_trade_n')}`",
        "",
        "## Attribution",
        "",
        f"A DIRECT_EXIT_DELTA = `{_fmt(attr.get('DIRECT_EXIT_DELTA'))}`",
        f"B SLOT_RELEASE_DOWNSTREAM_DELTA = `{_fmt(attr.get('SLOT_RELEASE_DOWNSTREAM_DELTA'))}`",
        f"C DISPLACED_TRADE_DELTA = `{_fmt(attr.get('DISPLACED_TRADE_DELTA'))}`",
        f"D TOTAL_CAUSAL_DELTA = `{_fmt(attr.get('TOTAL_CAUSAL_DELTA'))}`",
        "",
        "## Economics (100 shares, fees excluded, CAP=5)",
        "",
        f"Control total PnL = `{_fmt(oc.get('TOTAL_PNL_YEN'))}` PF = `{_fmt(oc.get('PF'))}` maxDD = `{_fmt(oc.get('REALIZED_MAX_DD'))}`",
        f"Treatment total PnL = `{_fmt(ot.get('TOTAL_PNL_YEN'))}` PF = `{_fmt(ot.get('PF'))}` maxDD = `{_fmt(ot.get('REALIZED_MAX_DD'))}`",
        f"DELTA_TOTAL_PNL = `{_fmt(req.get('DELTA_TOTAL_PNL'))}` DELTA_PF = `{_fmt(req.get('DELTA_PF'))}` DELTA_MAX_DD = `{_fmt(req.get('DELTA_MAX_DD'))}`",
        "",
        "## Winner safety (diagnostic labels, not rules)",
        "",
        f"Occupancy GOOD/DIP/EARLY/PTF/OTHER U EXIT = `{occ.get('GOOD_BRANCH_U_EXIT_N')}` / "
        f"`{occ.get('DIP_BRANCH_U_EXIT_N')}` / `{occ.get('EARLY_BRANCH_U_EXIT_N')}` / "
        f"`{occ.get('PTF_BRANCH_U_EXIT_N')}` / `{occ.get('OTHER_BRANCH_U_EXIT_N')}`",
        f"BRANCH_U_EXIT_N = `{occ.get('BRANCH_U_EXIT_N')}` (direction gate requires >= 10)",
        "",
        "## Provenance",
        "",
    ]
    for p in list(summary.get("PROVENANCE") or []):
        lines.append(
            f"- {p.get('date')}: previously_used=`{p.get('previously_used')}` used_by=`{p.get('used_by')}` reason=`{p.get('reason')}`"
        )
    lines.extend(["", "## Days", ""])
    for r in list(day.get("daily") or []):
        lines.append(
            f"- {r.get('date')}: control=`{_fmt(r.get('control_pnl'))}` treatment=`{_fmt(r.get('treatment_pnl'))}` "
            f"delta=`{_fmt(r.get('delta_pnl'))}` U_EXIT_n=`{r.get('branch_u_exit_n')}` "
            f"incremental_fills=`{r.get('additional_fill_n')}` cap_reject_delta=`{r.get('cap_reject_delta')}` "
            f"direct=`{_fmt(r.get('direct_exit_delta'))}` downstream=`{_fmt(r.get('downstream_delta'))}` `{r.get('side')}`"
        )
    lines.extend(
        [
            "",
            f"ENTRY_CHANGED=`{req.get('ENTRY_CHANGED')}` SIZING_CHANGED=`{req.get('SIZING_CHANGED')}` "
            f"NON_INTERFERENCE_PASS=`{req.get('NON_INTERFERENCE_PASS')}`",
            "",
            "STOP.",
            "",
        ]
    )
    return "\n".join(lines)


def terminate_holdout_series() -> dict[str, Any]:
    """Keep ACCUMULATING as the formal verdict. Add TERMINATED_INCONCLUSIVE status only."""
    from research.simple_tech_redesign.causal_board_rca_spec import (
        HOLDOUT_STATUS,
        HOLDOUT_TERMINATION_REASON,
        HOLDOUT_VERDICT_KEPT,
    )

    jp = BRANCH_U_HOLDOUT_OUT / "report.json"
    mp = BRANCH_U_HOLDOUT_OUT / "report.md"
    xp = BRANCH_U_HOLDOUT_OUT / "audit.xlsx"
    if not jp.is_file():
        return {"ok": False, "blocker": "holdout_report_missing"}
    body = json.loads(jp.read_text(encoding="utf-8"))
    req = dict(body.get("required") or {})
    if str(req.get("VERDICT") or "") != HOLDOUT_VERDICT_KEPT:
        return {"ok": False, "blocker": f"holdout_verdict_not_kept:{req.get('VERDICT')}"}
    req["HOLDOUT_STATUS"] = HOLDOUT_STATUS
    req["HOLDOUT_TERMINATION_REASON"] = HOLDOUT_TERMINATION_REASON
    req["FORWARD_BURNED"] = True
    req["FORMAL_DIRECTION_VERDICT"] = HOLDOUT_VERDICT_KEPT
    req["NOT_CONTRADICTED"] = True
    req["NOT_SUPPORTED"] = True
    body["required"] = req
    body["holdout_termination"] = {
        "HOLDOUT_STATUS": HOLDOUT_STATUS,
        "REASON": HOLDOUT_TERMINATION_REASON,
        "VERDICT_UNCHANGED": HOLDOUT_VERDICT_KEPT,
        "FORWARD_BURNED_DAYS": list(req.get("HOLDOUT_DAYS") or []),
    }
    jp.write_text(json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    md = mp.read_text(encoding="utf-8") if mp.is_file() else ""
    block = (
        "\n## Termination (not a direction verdict)\n\n"
        f"Status: **{HOLDOUT_STATUS}**\n\n"
        f"Formal VERDICT unchanged: `{HOLDOUT_VERDICT_KEPT}`.\n"
        "This is not CONTRADICTED and not SUPPORTED. The precommitted 10-event direction gate was not reached.\n"
        f"Reason: `{HOLDOUT_TERMINATION_REASON}`.\n"
        "The four holdout days are burned diagnostic data for the causal board RCA. 20260903 was not used.\n"
    )
    if "## Termination" not in md:
        if md.startswith("# "):
            first, rest = md.split("\n", 1)
            md = first + "\n" + block + rest
        else:
            md = block + md
        mp.write_text(md, encoding="utf-8")
    if xp.is_file():
        wb = load_workbook(xp)
        if "termination" in wb.sheetnames:
            del wb["termination"]
        ws = wb.create_sheet("termination")
        _sheet(
            ws,
            kv_rows(
                {
                    "HOLDOUT_STATUS": HOLDOUT_STATUS,
                    "REASON": HOLDOUT_TERMINATION_REASON,
                    "VERDICT_UNCHANGED": HOLDOUT_VERDICT_KEPT,
                    "NOT_CONTRADICTED": True,
                    "NOT_SUPPORTED": True,
                }
            ),
        )
        wb.save(xp)
    return {"ok": True, "status": HOLDOUT_STATUS, "verdict": HOLDOUT_VERDICT_KEPT}
