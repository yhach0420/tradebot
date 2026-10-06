"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.r11_online_episode_causality_audit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Frozen_Source",
    "Primitive_Onsets",
    "Offline_Clusters",
    "Streaming_Clusters",
    "Knowability",
    "Lookahead_Violations",
    "Candidate_Parity",
    "Decision_Time_Parity",
    "Causal_Delayed_Replay",
    "First_Event_Canary",
    "Lunch_Exit_Audit",
    "TimeStop_Audit",
    "Trade_Identity",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _econ(pack: dict[str, Any]) -> dict[str, Any]:
    return {
        "trade_n": pack.get("trade_n"),
        "day_n": pack.get("day_n"),
        "symbol_n": pack.get("symbol_n"),
        "mean_x0_bps": pack.get("mean_x0_bps"),
        "mean_x1_bps": pack.get("mean_x1_bps"),
        "profit_factor": pack.get("profit_factor"),
        "candidate_n": pack.get("candidate_n"),
        "skipped": pack.get("skipped"),
        "block_mean_x0": pack.get("block_mean_x0"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    k = dict(report.get("knowability") or {})
    p = dict(report.get("parity") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "strategy_id": report.get("strategy_id"),
                "old_confirmation_opened": False,
                "frozen_validation_opened": False,
                "parameter_changed": False,
            }
        ),
        "Frozen_Source": _kv({"strategy_id": report.get("strategy_id"), "episode_n": (report.get("episode_set") or {}).get("episode_n")}),
        "Primitive_Onsets": _kv({"same_as_atlas": True, "gap_min": 10, "last_event_is_offline_trigger": True}),
        "Offline_Clusters": _kv(dict(report.get("episode_set") or {})),
        "Streaming_Clusters": _kv(
            {
                "ONLINE_EVENT_PARITY": p.get("ONLINE_EVENT_PARITY"),
                "stream_last_n": p.get("ONLINE_EVENT_PARITY_n_stream"),
                "offline_last_n": p.get("ONLINE_EVENT_PARITY_n_offline"),
                "overlap": p.get("ONLINE_EVENT_PARITY_overlap"),
            }
        ),
        "Knowability": _kv(k) + list(report.get("knowability_samples") or []),
        "Lookahead_Violations": _kv(
            {
                "episode_n": k.get("lookahead_violation_episode_n"),
                "candidate_n": k.get("lookahead_violation_candidate_n"),
                "trade_n": k.get("lookahead_violation_trade_n"),
                "lag_min_counter": k.get("lag_min_counter"),
            }
        ),
        "Candidate_Parity": _kv({kk: p.get(kk) for kk in ("OFFLINE_SOURCE_PARITY", "ONLINE_CANDIDATE_PARITY", "ONLINE_EVENT_PARITY")}),
        "Decision_Time_Parity": _kv(
            {
                "ONLINE_DECISION_TIME_PARITY": p.get("ONLINE_DECISION_TIME_PARITY"),
                "ONLINE_DECISION_TIME_MATCH_N": p.get("ONLINE_DECISION_TIME_MATCH_N"),
            }
        ),
        "Causal_Delayed_Replay": [_econ(dict(report.get("causal_delayed_replay") or {}))],
        "First_Event_Canary": _kv(dict(k.get("first_event_canary") or {})) + [_econ(dict(report.get("first_event_refractory_canary") or {}))],
        "Lunch_Exit_Audit": _kv(dict(report.get("lunch_exit_audit") or {})),
        "TimeStop_Audit": _kv(dict(report.get("timestop_audit") or {})),
        "Trade_Identity": _kv(
            {
                "ENTRY_FILL_PARITY": p.get("ENTRY_FILL_PARITY"),
                "EXIT_PARITY": p.get("EXIT_PARITY"),
                "PORTFOLIO_ORDER_PARITY": p.get("PORTFOLIO_ORDER_PARITY"),
            }
        ),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0"}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# R11_ONLINE_EPISODE_CAUSALITY_AUDIT_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            d.get("INTERPRETATION") or "",
            "",
            f"Can last event be known at its own timestamp? **{a.get('Can_the_cluster_last_event_be_known_at_its_own_event_timestamp')}**",
            f"How: {a.get('How')}",
            f"Offline generator inspects later events? **{a.get('Does_current_offline_generator_inspect_later_events_to_determine_which_event_survives_clustering')}**",
            f"lookahead_violation episode/candidate/trade: **{a.get('lookahead_violation_episode_n')}** / **{a.get('lookahead_violation_candidate_n')}** / **{a.get('lookahead_violation_trade_n')}**",
            f"Streaming 4149/4149 at identical decision times? **{a.get('Can_streaming_reproduce_4149_4149_at_identical_decision_timestamps')}**",
            f"908/908 without retroactive assignment? **{a.get('Can_it_reproduce_908_908_without_retroactive_event_assignment')}**",
            f"Earliest causal decision: `{a.get('If_not_earliest_causal_decision_timestamp')}`",
            f"Causal delayed R11 n/X0/X1/PF: **{a.get('Causal_delayed_R11_trade_n')}** / **{a.get('Causal_delayed_R11_X0')}** / **{a.get('Causal_delayed_R11_X1')}** / **{a.get('Causal_delayed_R11_PF')}**",
            f"First-event refractory cand/trade/overlap: **{a.get('First_event_refractory_candidate_n')}** / **{a.get('First_event_refractory_trade_n')}** / `{a.get('First_event_overlap_with_frozen_R11')}`",
            f"Why session_flat ~11:30: {a.get('Why_session_flat_around_1130')}",
            f"Lunch flatten part of tested strategy? **{a.get('Is_lunch_flatten_part_of_tested_strategy')}**",
            f"time_stop 20: **{a.get('Is_time_stop_20_wall_clock_or_observed_bars')}**",
            f"OFFLINE_SOURCE_PARITY **{a.get('OFFLINE_SOURCE_PARITY')}** ONLINE_EVENT_PARITY **{a.get('ONLINE_EVENT_PARITY')}** ONLINE_DECISION_TIME_PARITY **{a.get('ONLINE_DECISION_TIME_PARITY')}** ONLINE_CANDIDATE_PARITY **{a.get('ONLINE_CANDIDATE_PARITY')}** ENTRY_FILL_PARITY **{a.get('ENTRY_FILL_PARITY')}** EXIT_PARITY **{a.get('EXIT_PARITY')}** PORTFOLIO_ORDER_PARITY **{a.get('PORTFOLIO_ORDER_PARITY')}**",
            f"Any parameter changed? **{a.get('Any_parameter_changed')}** threshold tuned? **{a.get('Any_threshold_tuned')}**",
            f"Old Confirmation opened? **{a.get('Old_Confirmation_opened')}** Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
            f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
