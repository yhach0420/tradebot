"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_direct_exec_u_target_precommit import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_direct_exec_u_target_precommit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Target",
    "Model",
    "Selection",
    "DevelopmentArms",
    "FutureGates",
    "TargetExamples",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    arms = g.get("ALLOWED_DEVELOPMENT_ARMS")
    arms_s = ", ".join(str(a) for a in arms) if isinstance(arms, (list, tuple)) else str(arms)
    return "\n".join(
        [
            "# AM DIRECT EXECUTION-ADJUSTED UPSIDE TARGET PRECOMMIT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. Candidate-level TARGET_EXEC_U = I(fill5) * U_FILL.",
            "RandomForestRegressor frozen from WAIT5 session-target learnability.",
            "pred_EXEC_U Top3. Primary reference FILL_ONLY. Stage2 closed.",
            "No training this run. No PnL. No Exact.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"SESSION: {g.get('SESSION')}",
            "",
            f"DEV_WAIT_SEC: {float(g.get('DEV_WAIT_SEC') or 0):.1f}",
            f"RUNTIME_WAIT_SEC: {float(g.get('RUNTIME_WAIT_SEC') or 0):.1f}",
            "",
            f"DIRECT_TARGET: {g.get('DIRECT_TARGET')}",
            f"DIRECT_TARGET_FORMULA: {g.get('DIRECT_TARGET_FORMULA')}",
            f"NONFILL_TARGET_VALUE: {float(g.get('NONFILL_TARGET_VALUE') if g.get('NONFILL_TARGET_VALUE') is not None else 0):.1f}",
            "",
            f"MODEL_FAMILY: {g.get('MODEL_FAMILY')}",
            f"MODEL_HYPERPARAMETERS_FROZEN: {_tf(g.get('MODEL_HYPERPARAMETERS_FROZEN'))}",
            "",
            f"REPRESENTATION_N: {g.get('REPRESENTATION_N')}",
            "",
            f"FINAL_SELECTION_N: {g.get('FINAL_SELECTION_N')}",
            f"SELECTION_RULE: {g.get('SELECTION_RULE')}",
            "",
            f"STAGE1_ALLOWED: {_tf(g.get('STAGE1_ALLOWED'))}",
            f"STAGE2_ALLOWED: {_tf(g.get('STAGE2_ALLOWED'))}",
            "",
            f"D_TRAINING_TARGET_ALLOWED: {_tf(g.get('D_TRAINING_TARGET_ALLOWED'))}",
            f"D_ROLE: {g.get('D_ROLE')}",
            "",
            f"FILL_LOSS_ALLOWED_IN_PASS_GATE: {_tf(g.get('FILL_LOSS_ALLOWED_IN_PASS_GATE'))}",
            "",
            f"ALLOWED_DEVELOPMENT_ARMS: [{arms_s}]",
            "",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("PRIMARY_FINDING") or ""),
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No EXEC_U training. No development performance run. No Stage2 restart.",
            "No W5 runtime adoption. Runtime WAIT_SEC remains 1.0.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
