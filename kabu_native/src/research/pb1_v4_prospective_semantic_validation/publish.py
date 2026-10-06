"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_prospective_semantic_validation.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Hash_Verification",
    "Session_Manifest",
    "Quality_Gate",
    "Prospective_State",
    "Stopping_Rule",
    "Safety",
)
STRIP = {"isolation_before", "isolation_after", "minutes"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    sess = dict(report.get("session_result") or {})
    return {
        "SPEC_CHANGED": False,
        "V4_CHANGED": False,
        "THRESHOLD_RETUNED": False,
        "PRECOMMIT_CHANGED": False,
        "PNL_USED": False,
        "MFE_MAE_USED": False,
        "FUTURE_OUTCOME_USED": False,
        "PROSPECTIVE_DATA_OPENED": sess.get("PROSPECTIVE_DATA_OPENED"),
        "SESSION_STATE": sess.get("SESSION_STATE"),
        "SESSION_ACCEPTED_FOR_PROSPECTIVE": sess.get("SESSION_ACCEPTED_FOR_PROSPECTIVE"),
        "INELIGIBLE_SESSION": sess.get("INELIGIBLE_SESSION"),
        "session_date": sess.get("session_date"),
        "machine_sha_verified": sess.get("machine_sha_verified"),
        "precommit_sha_verified": sess.get("precommit_sha_verified"),
        "data_complete": sess.get("data_complete"),
        "candidate_day_n": sess.get("candidate_day_n"),
        "cumulative_candidate_day_n": sess.get("cumulative_candidate_day_n"),
        "eligible_session_n": sess.get("eligible_session_n"),
        "hidden1m_mismatch_n": sess.get("hidden1m_mismatch_n"),
        "invariant_violation_n": sess.get("invariant_violation_n"),
        "VERDICT": (report.get("decision") or {}).get("VERDICT"),
        "NEXT": (report.get("decision") or {}).get("NEXT"),
        "submit/cancel/live": "0/0/0",
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        "Manifest": [{"key": k, "value": v} for k, v in (report.get("answers") or {}).items()],
        "Hash_Verification": [{"key": k, "value": v} for k, v in dict(report.get("hash_gate") or {}).items()],
        "Session_Manifest": [{"key": k, "value": v} for k, v in dict(report.get("session_result") or {}).items()],
        "Quality_Gate": [{"key": k, "value": v} for k, v in dict(report.get("quality") or {}).items()],
        "Prospective_State": list(report.get("state_log_rows") or [])[:200] or [{"empty": True, "note": "no_prospective_candidate_rows"}],
        "Stopping_Rule": [{"key": k, "value": v} for k, v in dict(report.get("stopping") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in dict(report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = [
        "# PB1 V4 prospective semantic validation",
        "",
        "Operational session summary only. No per-case semantic verdict. No PnL.",
        "",
    ]
    for k, v in ans.items():
        md.append(f"- **{k}** `{v}`")
    md += ["", "STOP."]
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
