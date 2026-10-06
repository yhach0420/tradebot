"""Write report.json / report.md / audit.xlsx only. No mass CSV."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_correction_v3_delta_rca.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Parent_Hashes",
    "Delta_Scope",
    "7011_20241205_Trace",
    "PostDominant_Persistence",
    "7011_20250523_5M_Path",
    "Renewal_Event_Audit",
    "Range_Acceptance_Audit",
    "4063_Control",
    "3382_Control",
    "Natural_Stale_Controls",
    "Natural_Pause_Controls",
    "Numeric_Predicate_Provenance",
    "Protected_Targets_Impact",
    "Safety",
)
STRIP = {"_markdown", "snap", "setups", "e0_events", "e1_events", "funnel_days", "hidden_1m_checks", "minutes", "isolation_before", "isolation_after"}


def finite_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): finite_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [finite_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(finite_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def kv(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    ans = dict(report.get("answers") or {})
    q = {k: (v.get("answer") if isinstance(v, dict) and "answer" in v else v) for k, v in ans.items()}
    return {
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "SPEC_CHANGE_REQUIRED": dec.get("SPEC_CHANGE_REQUIRED"),
        "Q1": q.get("Q1"),
        "Q2": q.get("Q2"),
        "Q3": q.get("Q3"),
        "Q4": q.get("Q4"),
        "Q5": q.get("Q5"),
        "Q6_10:09_GEOMETRIC": ((ans.get("Q6") or {}).get("10:09") or {}).get("GEOMETRIC_EXTENSION"),
        "Q6_10:09_REAL": ((ans.get("Q6") or {}).get("10:09") or {}).get("REAL_RENEWED_AUCTION"),
        "Q6_10:19_GEOMETRIC": ((ans.get("Q6") or {}).get("10:19") or {}).get("GEOMETRIC_EXTENSION"),
        "Q6_10:19_REAL": ((ans.get("Q6") or {}).get("10:19") or {}).get("REAL_RENEWED_AUCTION"),
        "Q7": q.get("Q7"),
        "Q8": q.get("Q8"),
        "Q9": q.get("Q9"),
        "Q10": q.get("Q10"),
        "CODE_CHANGED": False,
        "submit/cancel/live": "0/0/0",
        **{f"full_{k}": ans.get(k) for k in ("Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7", "Q8", "Q9", "Q10")},
    }


def build_sheets(report: dict[str, Any], *, db: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    da = dict(report.get("delta_a") or {})
    ans = dict(report.get("answers") or {})
    dec = dict(report.get("decision") or {})
    hashes = dict(report.get("hashes") or {})
    nat = dict((db.get("natural_controls") or {}))
    p7011 = list((db.get("7011_20250523") or {}).get("path") or [])
    p4063 = list((db.get("4063_20251118") or {}).get("path") or [])
    p3382 = list((db.get("3382_20241004") or {}).get("path") or [])
    renew = dict(db.get("renewal_event_audit") or {})
    return {
        "Manifest": [
            {
                "PROGRAM_ID": report.get("PROGRAM_ID"),
                "ANALYSIS_ID": report.get("ANALYSIS_ID"),
                "VERDICT": dec.get("VERDICT"),
                "NEXT": dec.get("NEXT"),
                "SPEC_CHANGE_REQUIRED": dec.get("SPEC_CHANGE_REQUIRED"),
                "SOURCE_SHA256": hashes.get("SOURCE_SHA256"),
                "rca_set_n": report.get("rca_set_n"),
                "diagnosis_only": True,
                "v4_implemented": False,
            }
        ],
        "Parent_Hashes": kv(
            {
                "SPEC_SHA256": hashes.get("SPEC_SHA256"),
                "CORRECTION_V2_SHA256": hashes.get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V2_SHA256"),
                "CORRECTION_V3_SHA256": hashes.get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V3_SHA256"),
                "PARENT_MACHINE_SHA256": hashes.get("PB1_V4_CLARIFIED_MACHINE_SHA256"),
                "CORRECTED_SHA256": hashes.get("V4_CORRECTED_MACHINE_SHA256"),
                "LEAKED_SHA256": hashes.get("V4_LEGACY_LEAKED_IMPLEMENTATION"),
                "SPEC_CHANGED": False,
                "CORRECTION_V2_CHANGED": False,
                "CORRECTION_V3_CHANGED": False,
            }
        ),
        "Delta_Scope": [
            {"delta": "A", "case": "7011/20241205", "issue": "NO_VALID path loses POST_DOMINANT class", "do_not": "restore TRUE"},
            {"delta": "B", "case": "7011/20250523", "issue": "leftover range not recognized as ACTIVE/THESIS death", "do_not": "N-bar expiry"},
            {"delta": "C", "case": "V3 promoted RCA predicates", "issue": "numeric provenance", "do_not": "retune after 88"},
        ],
        "7011_20241205_Trace": list(da.get("traces") or []),
        "PostDominant_Persistence": [dict(da.get("drop_locus") or {}), dict(da.get("desired_contract") or {})],
        "7011_20250523_5M_Path": p7011,
        "Renewal_Event_Audit": [{"time": k, **dict(v)} for k, v in renew.items()],
        "Range_Acceptance_Audit": [
            {
                "time": r.get("time"),
                "progress_class": r.get("progress_class"),
                "candidate_death_events": r.get("candidate_death_events"),
                "v3_stale": r.get("v3_stale"),
                "overlap_prev_bar": r.get("overlap_prev_bar"),
                "overlap_recent_active_range": r.get("overlap_recent_active_range"),
                "close_beyond_prior_active": r.get("close_beyond_prior_active"),
                "GEOMETRIC_EXTENSION": r.get("GEOMETRIC_EXTENSION"),
                "REAL_RENEWED_AUCTION": r.get("REAL_RENEWED_AUCTION"),
            }
            for r in p7011
        ],
        "4063_Control": p4063,
        "3382_Control": p3382,
        "Natural_Stale_Controls": list(nat.get("stale_after_thesis_reached") or []),
        "Natural_Pause_Controls": [
            {k: x.get(k) for k in ("symbol", "date", "THESIS_LIVE", "THESIS_LOST_AT", "THESIS_LOST_REASON", "e1_entry_t", "pause_then_real", "leftover_then_real")}
            for x in list(nat.get("long_no_progress_then_real_breakout") or [])
        ],
        "Numeric_Predicate_Provenance": list(report.get("numeric_predicate_provenance") or []),
        "Protected_Targets_Impact": list(report.get("protected_targets_impact") or []),
        "Safety": kv(
            {
                "CODE_CHANGED": False,
                "SPEC_CHANGED": False,
                "CORRECTION_V2_CHANGED": False,
                "CORRECTION_V3_CHANGED": False,
                "PROSPECTIVE_DATA_OPENED": False,
                "OLD_CONFIRMATION_OPENED": False,
                "FROZEN_VALIDATION_OPENED": False,
                "FUTURE_OUTCOME_USED": False,
                "PNL_USED": False,
                "MFE_MAE_USED": False,
                "GRID_SEARCH": False,
                "THRESHOLD_RETUNED": False,
                "submit/cancel/live": "0/0/0",
                "Q1": (ans.get("Q1") or {}).get("answer") if isinstance(ans.get("Q1"), dict) else ans.get("Q1"),
                "Q10": (ans.get("Q10") or {}).get("answer") if isinstance(ans.get("Q10"), dict) else ans.get("Q10"),
            }
        ),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = finite_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    answers = dict(report.get("answers") or {})
    dec = dict(report.get("decision") or {})
    lines = [
        "# PB1_V4_CLARIFIED_MACHINE_CORRECTION_V3_DELTA_RCA_V1",
        "",
        f"VERDICT: {dec.get('VERDICT')}",
        f"SPEC_CHANGE_REQUIRED: {dec.get('SPEC_CHANGE_REQUIRED')}",
        f"NEXT: {dec.get('NEXT')}",
        "",
        "Diagnosis only. No machine change. No spec change. No V3.1. No V4 implementation. No threshold retune. No prospective. No PnL.",
        "",
    ]
    for k in ("Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7", "Q8", "Q9", "Q10"):
        v = answers.get(k)
        lines.append(f"## {k}")
        lines.append("")
        lines.append(f"`{json.dumps(finite_sanitize(v), ensure_ascii=False)[:4000]}`")
        lines.append("")
    lines.extend(
        [
            "## SAFETY",
            "",
            "CODE_CHANGED = false",
            "SPEC_CHANGED = false",
            "CORRECTION_V2_CHANGED = false",
            "CORRECTION_V3_CHANGED = false",
            "PROSPECTIVE_DATA_OPENED = false",
            "OLD_CONFIRMATION_OPENED = false",
            "FROZEN_VALIDATION_OPENED = false",
            "FUTURE_OUTCOME_USED = false",
            "PNL_USED = false",
            "MFE_MAE_USED = false",
            "GRID_SEARCH = false",
            "THRESHOLD_RETUNED = false",
            "submit/cancel/live = 0/0/0",
            "",
            "STOP.",
            "",
        ]
    )
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
