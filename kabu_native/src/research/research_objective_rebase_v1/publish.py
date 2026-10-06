"""Write report.json / report.md / audit.xlsx only. No mass CSV."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.research_objective_rebase_v1 import ANALYSIS_ID
from research.research_objective_rebase_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "source_inventory",
    "architecture_to_lineage",
    "lineage_dependencies",
    "l1_technical",
    "l2_recovery",
    "l3_activity",
    "c4_overlay",
    "failure_taxonomy",
    "evidence_levels",
    "aggregate_edge",
    "generalization",
    "value_capture",
    "pseudo_replication_guard",
    "closed_lineage_guard",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def taxonomy_rows(pack: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for lid, blob in (
        ("L1_TECHNICAL_PRICE_STATE", pack.get("l1") or {}),
        ("L2_RECOVERY_RECLAIM_PATH", pack.get("l2") or {}),
        ("L3_ACTIVITY_ONSET", pack.get("l3") or {}),
    ):
        tax = dict(blob.get("taxonomy") or {})
        rows.append(
            {
                "LINEAGE_ID": lid,
                "EVIDENCE_LEVEL": blob.get("EVIDENCE_LEVEL"),
                "PRIMARY_FAILURE_PATTERN": blob.get("PRIMARY_FAILURE_PATTERN"),
                **tax,
            }
        )
    c4 = dict(pack.get("c4") or {})
    rows.append(
        {
            "LINEAGE_ID": "O1_PORTFOLIO_CROWDING_OVERLAY",
            "EVIDENCE_LEVEL": "NOT_A_LINEAGE",
            "PRIMARY_FAILURE_PATTERN": "ABSOLUTE_ECONOMIC_PASS_ZERO",
            "COVERAGE_FAILURE": "NOT_SUPPORTED",
            "ABSOLUTE_EDGE_FAILURE": "SUPPORTED",
            "DAY_SIGN_FAILURE": "NOT_COMPUTED",
            "BEST_DAY_DEPENDENCE": "NOT_COMPUTED",
            "DRAWDOWN_DOMINANCE": "NOT_COMPUTED",
            "SYMBOL_CONCENTRATION": "NOT_COMPUTED",
            "SELECTION_INSTABILITY": "NOT_APPLICABLE",
            "PORTFOLIO_INTERACTION": "SUPPORTED",
            "POST_FILL_VALUE_CAPTURE_EVIDENCE": "NOT_SUPPORTED",
            "INTEGRITY_LIMIT": "NOT_SUPPORTED",
            "OVERLAY_RESCUE": c4.get("OVERLAY_RESCUE"),
        }
    )
    return rows


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT') or a.get('50_VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT') or a.get('51_NEXT'))}",
        "",
        f"PRIMARY_DEFICIENCY: {_fmt(d.get('PRIMARY_DEFICIENCY') or a.get('31_PRIMARY_DEFICIENCY'))}",
        "",
        "NEW_REPLAY: false",
        "",
        "NEW_PNL_SIMULATION: false",
        "",
        "Holdout read: false",
        "",
        "Stress read: false",
        "",
        "## Answers 1-51",
        "",
    ]
    for k, v in a.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))
