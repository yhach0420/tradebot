"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.higher_magnitude_state_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Raw_Events",
    "Episodes",
    "Episode_Sequences",
    "Episode_Dedup",
    "Path_Magnitude",
    "Pullback_Reclaim",
    "Compression_Breakout",
    "Sector_Leader",
    "CrossSection_Ranking",
    "Rank_Response",
    "Sequence_Increment",
    "Complete_Strategy_1",
    "Complete_Strategy_2",
    "D1_D4",
    "Day_Distribution",
    "Symbol_Sector",
    "Execution_Sensitivity",
    "Secondary_Confirmation",
    "Frozen_Validation_Status",
    "Live_20260914",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _cand(c: dict[str, Any] | None, idx: int) -> list[dict[str, Any]]:
    if not c:
        return [{"candidate": idx, "status": "none"}]
    spec = dict(c.get("spec") or {})
    econ = dict(c.get("discovery_economics") or {})
    promo = dict(c.get("promotion") or {})
    return _kv({**{k: spec.get(k) for k in ("candidate_id", "spec_sha256", "archetype", "thesis", "event_sequence", "ranking_rule")}, **promo, **{k: econ.get(k) for k in ("trade_n", "day_n", "mean_x0_bps", "mean_x1_bps", "profit_factor", "max_dd_daily_mean_bps", "block_mean_x0", "sector_specific_strategy")}})


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    bind = dict(report.get("bind") or {})
    split = dict(report.get("split") or {})
    raw = dict(report.get("raw_events") or {})
    eps = dict(report.get("episodes") or {})
    seq = dict(report.get("sequence_increment") or {})
    arch = dict(report.get("archetype_magnitude") or {})
    rank = dict(report.get("rank_response") or {})
    absr = dict(report.get("absolute_vs_rank") or {})
    cand = dict(report.get("candidates") or {})
    promoted = list(cand.get("promoted") or [])
    proposals = list(report.get("proposal_summaries") or [])
    sec = dict(report.get("secondary_confirmation") or {})
    live = dict(report.get("live_20260914") or {})
    c1 = promoted[0] if len(promoted) > 0 else None
    c2 = promoted[1] if len(promoted) > 1 else None
    return {
        "Binding": _kv(
            {
                "parent_verdict": report.get("parent_verdict_accepted"),
                "split_sha": split.get("split_sha256"),
                "bind_ok": bind.get("ok"),
                "cs1_cs3_rescued": report.get("cs1_cs3_rescued"),
                "old_confirmation_used_to_design": report.get("old_confirmation_used_to_design"),
                "kabu_50": report.get("kabu_50_applied"),
            }
        )
        + [{"key": "rejected_cs", "value": x} for x in list(report.get("rejected_cs") or [])],
        "Raw_Events": _kv({"n": raw.get("n"), "event_sha": raw.get("event_sha")}) + [{"family": k, "n": v} for k, v in sorted(dict(raw.get("families") or {}).items())],
        "Episodes": _kv(eps),
        "Episode_Sequences": list(seq.get("by_class") or [{"empty": True}]),
        "Episode_Dedup": _kv({"rule": eps.get("dedup_rule"), "raw_assigned": eps.get("raw_events_assigned_to_episodes"), "unique_episode_n": eps.get("unique_episode_n"), "future_used": eps.get("future_used_for_episode_boundary")}),
        "Path_Magnitude": list(seq.get("by_class") or [{"empty": True}]),
        "Pullback_Reclaim": _kv(dict(arch.get("CONTROLLED_PULLBACK_RECLAIM") or {"empty": True})),
        "Compression_Breakout": _kv(dict(arch.get("COMPRESSION_BREAKOUT_EXPANSION") or {"empty": True})),
        "Sector_Leader": _kv(dict(arch.get("SECTOR_LEADER_TRANSITION") or {"empty": True})),
        "CrossSection_Ranking": _kv(absr),
        "Rank_Response": _kv({k: rank.get(k) for k in ("n", "ordered_rank_response", "quintile_means_best_to_worst", "quintile_n", "quintile_mfe", "k_predeclared", "k_searched_1_to_30", "by_k", "did_not_use_future_in_rank")}),
        "Sequence_Increment": list(seq.get("increments") or [{"empty": True}]) + _kv({"sequence_order_adds_value": seq.get("sequence_order_adds_value"), "largest_class": seq.get("largest_class")}),
        "Complete_Strategy_1": _cand(c1, 1),
        "Complete_Strategy_2": _cand(c2, 2),
        "D1_D4": [
            {"candidate_id": p.get("candidate_id"), "mode": "ranked_top3", **dict((p.get("ranked") or {}).get("block_mean_x0") or {})}
            for p in proposals
        ]
        or [{"empty": True}],
        "Day_Distribution": [
            {"candidate_id": p.get("candidate_id"), "day_n": (p.get("ranked") or {}).get("day_n"), "trades_per_day_mean": (p.get("ranked") or {}).get("trades_per_day_mean"), "daily_positive_share": (p.get("ranked") or {}).get("daily_positive_share"), "months": (p.get("ranked") or {}).get("month_mean_x0")}
            for p in proposals
        ]
        or [{"empty": True}],
        "Symbol_Sector": [
            {"candidate_id": p.get("candidate_id"), "symbol_n": (p.get("ranked") or {}).get("symbol_n"), "sector_n": (p.get("ranked") or {}).get("sector_n"), "top_symbol_share": (p.get("ranked") or {}).get("top_symbol_share_of_positive_bps"), "top_sector": (p.get("ranked") or {}).get("top_sector"), "top_sector_share": (p.get("ranked") or {}).get("top_sector_share_of_positive_bps"), "sector_specific": (p.get("ranked") or {}).get("sector_specific_strategy")}
            for p in proposals
        ]
        or [{"empty": True}],
        "Execution_Sensitivity": [
            {"candidate_id": p.get("candidate_id"), "X0": (p.get("ranked") or {}).get("mean_x0_bps"), "X1_8bps_tax": (p.get("ranked") or {}).get("mean_x1_bps"), "absolute_X0": (p.get("absolute") or {}).get("mean_x0_bps"), "top1_X0": (p.get("top1") or {}).get("mean_x0_bps"), "not_bid_ask": True}
            for p in proposals
        ]
        or [{"empty": True}],
        "Secondary_Confirmation": _kv({k: sec.get(k) for k in ("ran", "label", "old_confirmation_used_to_design", "any_passed", "certifies", "frozen_validation_opened")})
        + list(sec.get("rows") or [{"status": "not_run"}]),
        "Frozen_Validation_Status": _kv(dict(report.get("frozen_validation") or {})),
        "Live_20260914": _kv(
            {
                "classification": live.get("classification"),
                "fed_into_discovery_tuning": live.get("fed_into_discovery_tuning"),
                "futures_capture_running": live.get("futures_capture_running"),
                "breadth_capture_running": live.get("breadth_capture_running"),
                "futures": live.get("futures"),
                "breadth": live.get("breadth"),
            }
        ),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# HIGHER_MAGNITUDE_STATE_DISCOVERY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            d.get("INTERPRETATION") or "",
            "",
            f"Parent accepted: `{report.get('parent_verdict_accepted')}`",
            f"CS1–CS3 rescued? **{a.get('cs1_cs3_rescued')}**",
            f"Raw event N: **{a.get('raw_event_n')}**",
            f"Unique episode N: **{a.get('unique_episode_n')}**",
            f"Future in episode definition? **{a.get('future_outcomes_affected_episode_definition')}**",
            f"Sequence order adds value? **{a.get('sequence_order_adds_value')}**",
            f"Rank-response ordered? **{a.get('cross_sectional_ranking_ordered')}**",
            f"Best gross X0: `{a.get('best_gross_x0_opportunity_magnitude')}`",
            f"Best complete-strategy X0/X1: **{a.get('best_complete_strategy_x0')}** / **{a.get('best_complete_strategy_x1')}**",
            f"Exceeds prior 1–6bps? **{a.get('any_state_exceeds_prior_1_6bps')}**",
            f"Promoted n: **{a.get('complete_strategy_candidates_n')}**",
            f"Old Confirmation used to design? **{a.get('old_confirmation_used_to_design')}**",
            f"Secondary run? **{a.get('secondary_confirmation_run')}** result `{a.get('secondary_confirmation_result')}`",
            f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
            f"Kabu 50? **{a.get('kabu_50_applied')}**",
            f"20260914: `{a.get('live_20260914')}`",
            f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
