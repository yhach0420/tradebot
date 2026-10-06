"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_complete_strategy_economic_failure_decomposition.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Baseline",
    "Trade_Path",
    "Entry_Edge",
    "MFE_MAE",
    "Exit_Timing",
    "Exit_Reasons",
    "ASF_Walk",
    "Recross_Walk",
    "Price_Notional",
    "Normalized_Bps",
    "Long_Short",
    "E0_E1",
    "Seed_Location",
    "Portfolio_Blocks",
    "Session_Flat",
    "Execution_Tax",
    "Loss_Waterfall",
    "Mechanism_Verdict",
    "Safety",
)
STRIP = {
    "isolation_before",
    "isolation_after",
    "minutes",
    "recs",
    "funnel",
    "setups",
    "candidates",
    "e0_events",
    "e1_events",
}
PATH_COLS = (
    "date", "symbol", "entry_type", "side", "entry_t", "entry_px", "exit_t", "exit_reason",
    "THESIS_LOST_REASON", "entry_class", "MFE_bps", "MAE_bps", "realized_gross_bps", "gross_bps",
    "net_bps", "MFE_capture_ratio", "time_to_MFE", "time_to_MAE", "MFE_before_THESIS_LOST",
    "MAE_before_THESIS_LOST", "minutes_MFE_to_THESIS_LOST", "minutes_THESIS_LOST_to_fill",
    "ret_bps_1m", "ret_bps_3m", "ret_bps_5m", "ret_bps_10m", "ret_bps_20m", "ret_bps_30m", "ret_bps_60m",
    "net_pnl_yen", "seed_family", "location_family", "entry_notional_yen",
)


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _kv(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    mech = dict(report.get("mechanism") or {})
    return {
        "VERDICT": mech.get("VERDICT"),
        "CONCLUSION": mech.get("CONCLUSION"),
        "NEXT": mech.get("NEXT"),
        "V1_VERDICT": "PB1_V4_COMPLETE_STRATEGY_ECONOMIC_CONFIRMATION1_FAIL_V1",
        "V1_VERDICT_CHANGED": False,
        "CONFIRMATION1_RESCORED": False,
        "COMPLETE_STRATEGY_SHA256": "556542319d22dff40cc1758b24d44d2ecb1985b8595927ca8f80618126961bf8",
        "ECONOMIC_DEVELOPMENT_EXPOSED": True,
        "FROZEN_VALIDATION_ECONOMIC_OPENED": False,
        "PROSPECTIVE_DATA_OPENED": False,
        "causal_repair_clear": mech.get("causal_repair_clear"),
        "true_flags": mech.get("true_flags"),
        "submit/cancel/live": "0/0/0",
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    paths = list(report.get("paths") or [])
    compact = [{k: r.get(k) for k in PATH_COLS} for r in paths]
    slices = dict(report.get("slices") or {})
    return {
        "Manifest": _kv(dict(report.get("answers") or {})),
        "Baseline": _kv(dict(report.get("baseline") or {})),
        "Trade_Path": compact or [{"empty": True}],
        "Entry_Edge": _kv(dict(slices.get("entry_class") or {})),
        "MFE_MAE": compact or [{"empty": True}],
        "Exit_Timing": [
            {
                "date": r.get("date"),
                "symbol": r.get("symbol"),
                "entry_t": r.get("entry_t"),
                "MFE_t": r.get("MFE_t"),
                "MFE_bps": r.get("MFE_bps"),
                "THESIS_LOST_AT": r.get("THESIS_LOST_AT"),
                "exit_t": r.get("exit_t"),
                "realized_gross_bps": r.get("realized_gross_bps"),
                "MFE_to_exit_giveback_bps": r.get("MFE_to_exit_giveback_bps"),
                "MFE_capture_ratio": r.get("MFE_capture_ratio"),
                "minutes_MFE_to_THESIS_LOST": r.get("minutes_MFE_to_THESIS_LOST"),
                "minutes_THESIS_LOST_to_fill": r.get("minutes_THESIS_LOST_to_fill"),
            }
            for r in paths
        ]
        or [{"empty": True}],
        "Exit_Reasons": [{"reason": k, **v} for k, v in dict(slices.get("exit_reasons") or {}).items()] or [{"empty": True}],
        "ASF_Walk": list(report.get("asf_walk") or []) or [{"empty": True}],
        "Recross_Walk": list(report.get("recross_walk") or []) or [{"empty": True}],
        "Price_Notional": list(slices.get("notional") or []) or [{"empty": True}],
        "Normalized_Bps": _kv(dict(slices.get("normalized") or {})),
        "Long_Short": [{"side": "LONG", **dict((slices.get("long_short") or {}).get("bull") or (slices.get("long_short") or {}).get("long") or {})}, {"side": "SHORT", **dict((slices.get("long_short") or {}).get("bear") or (slices.get("long_short") or {}).get("short") or {})}],
        "E0_E1": [{"entry_type": "E0", **dict((slices.get("e0_e1") or {}).get("E0") or {})}, {"entry_type": "E1", **dict((slices.get("e0_e1") or {}).get("E1") or {})}],
        "Seed_Location": (
            [{"kind": "seed", "name": k, **v} for k, v in dict(slices.get("seed") or {}).items()]
            + [{"kind": "location", "name": k, **v} for k, v in dict(slices.get("location") or {}).items()]
        )
        or [{"empty": True}],
        "Portfolio_Blocks": list(report.get("portfolio_blocks") or []) or [{"empty": True}],
        "Session_Flat": list(report.get("session_flat") or []) or [{"empty": True}],
        "Execution_Tax": _kv(dict(slices.get("execution_tax") or {})),
        "Loss_Waterfall": list((report.get("waterfall") or {}).get("primary") or [])
        + [{"category": "PORTFOLIO_BLOCKING", **dict((report.get("waterfall") or {}).get("PORTFOLIO_BLOCKING") or {})}],
        "Mechanism_Verdict": _kv({k: v for k, v in dict(report.get("mechanism") or {}).items() if k != "waterfall"}),
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    mech = dict(report.get("mechanism") or {})
    ev = dict(mech.get("evidence") or {})
    md = [
        "# PB1 V4 Complete Strategy — Economic Failure Decomposition",
        "",
        "V1 Confirmation 1 FAIL is not reversed and is not rescored.",
        "Confirmation 1 is now `ECONOMIC_DEVELOPMENT_EXPOSED`. Frozen Validation economics remain sealed.",
        "",
        f"**V1 VERDICT (permanent)** `{ans.get('V1_VERDICT')}`",
        f"**RCA VERDICT** `{ans.get('VERDICT')}`",
        f"**CONCLUSION** `{ans.get('CONCLUSION')}`",
        f"**NEXT** `{ans.get('NEXT')}`",
        "",
        "## Evidence",
        "",
        f"- median MFE bps `{ev.get('median_MFE_bps')}`",
        f"- median realized gross bps `{ev.get('median_realized_gross_bps')}`",
        f"- mean net bps `{ev.get('mean_net_bps')}` vs mean net yen `{ev.get('mean_net_yen')}`",
        f"- class counts `{json.dumps(ev.get('class_n'), ensure_ascii=False)}`",
        f"- ASF+recross n `{ev.get('asf_recross_n')}` late `{ev.get('asf_recross_late_n')}`",
        f"- high-price yen `{ev.get('high_px_net_pnl_yen')}` n `{ev.get('high_px_n')}`",
        f"- blocked CF net `{ev.get('blocked_counterfactual_net_pnl_yen')}` (`COUNTERFACTUAL_NOT_STRATEGY_RESULT`)",
        f"- flags `{json.dumps(mech.get('flags'), ensure_ascii=False)}`",
        "",
        "No ENTRY/EXIT/symbol/E0-E1/LONG-SHORT/cost/CAP change in this task.",
        "",
        "STOP.",
        "",
    ]
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
