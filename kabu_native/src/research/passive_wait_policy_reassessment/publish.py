"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.passive_wait_policy_reassessment import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "passive_wait_policy_reassessment"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "W1Parity",
    "Waits",
    "Incremental",
    "QualityGroups",
    "TimeToFill",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# PASSIVE WAIT POLICY REASSESSMENT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine. Corrected Passive Fill frozen. t0 limit frozen.",
            "WAIT_SEC dimension only: 1 / 2 / 5 / 10. No interpolation. No repricing.",
            "Standalone fillability. No POSITION_CAP / occupancy / reentry.",
            "No model. No feature. No PnL. No Exact. Runtime WAIT unchanged.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"W1_ANY_RATE: {_fmt(g.get('W1_ANY_RATE'))}",
            f"W2_ANY_RATE: {_fmt(g.get('W2_ANY_RATE'))}",
            f"W5_ANY_RATE: {_fmt(g.get('W5_ANY_RATE'))}",
            f"W10_ANY_RATE: {_fmt(g.get('W10_ANY_RATE'))}",
            "",
            f"W1_MULTI_RATE: {_fmt(g.get('W1_MULTI_RATE'))}",
            f"W2_MULTI_RATE: {_fmt(g.get('W2_MULTI_RATE'))}",
            f"W5_MULTI_RATE: {_fmt(g.get('W5_MULTI_RATE'))}",
            f"W10_MULTI_RATE: {_fmt(g.get('W10_MULTI_RATE'))}",
            "",
            f"W1_VALID_BOARD_RATE: {_fmt(g.get('W1_VALID_BOARD_RATE'))}",
            f"W2_VALID_BOARD_RATE: {_fmt(g.get('W2_VALID_BOARD_RATE'))}",
            f"W5_VALID_BOARD_RATE: {_fmt(g.get('W5_VALID_BOARD_RATE'))}",
            f"W10_VALID_BOARD_RATE: {_fmt(g.get('W10_VALID_BOARD_RATE'))}",
            "",
            f"TIME_TO_FILL_MEDIAN_SEC: {_fmt(g.get('TIME_TO_FILL_MEDIAN_SEC'))}",
            f"TIME_TO_FILL_P90_SEC: {_fmt(g.get('TIME_TO_FILL_P90_SEC'))}",
            "",
            f"W2_NEW_FILL_N: {g.get('W2_NEW_FILL_N')}",
            f"W5_NEW_FILL_N: {g.get('W5_NEW_FILL_N')}",
            f"W10_NEW_FILL_N: {g.get('W10_NEW_FILL_N')}",
            "",
            f"W2_NEW_POSTFILL_MFE: {_fmt(g.get('W2_NEW_POSTFILL_MFE'))}",
            f"W2_NEW_POSTFILL_DOWNSIDE: {_fmt(g.get('W2_NEW_POSTFILL_DOWNSIDE'))}",
            "",
            f"W5_NEW_POSTFILL_MFE: {_fmt(g.get('W5_NEW_POSTFILL_MFE'))}",
            f"W5_NEW_POSTFILL_DOWNSIDE: {_fmt(g.get('W5_NEW_POSTFILL_DOWNSIDE'))}",
            "",
            f"W10_NEW_POSTFILL_MFE: {_fmt(g.get('W10_NEW_POSTFILL_MFE'))}",
            f"W10_NEW_POSTFILL_DOWNSIDE: {_fmt(g.get('W10_NEW_POSTFILL_DOWNSIDE'))}",
            "",
            f"W1_FILLABLE_JOINT_RATE: {_fmt(g.get('W1_FILLABLE_JOINT_RATE'))}",
            f"W2_FILLABLE_JOINT_RATE: {_fmt(g.get('W2_FILLABLE_JOINT_RATE'))}",
            f"W5_FILLABLE_JOINT_RATE: {_fmt(g.get('W5_FILLABLE_JOINT_RATE'))}",
            f"W10_FILLABLE_JOINT_RATE: {_fmt(g.get('W10_FILLABLE_JOINT_RATE'))}",
            "",
            f"W1_JOINT_SURVIVAL: {_fmt(g.get('W1_JOINT_SURVIVAL'))}",
            f"W2_JOINT_SURVIVAL: {_fmt(g.get('W2_JOINT_SURVIVAL'))}",
            f"W5_JOINT_SURVIVAL: {_fmt(g.get('W5_JOINT_SURVIVAL'))}",
            f"W10_JOINT_SURVIVAL: {_fmt(g.get('W10_JOINT_SURVIVAL'))}",
            "",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No WAIT adopted. No interpolation. No repricing. No price-policy test.",
            "Runtime WAIT_SEC remains 1.0. No runtime candidate.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
