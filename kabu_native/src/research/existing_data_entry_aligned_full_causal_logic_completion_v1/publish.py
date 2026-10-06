"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.existing_data_entry_aligned_full_causal_logic_completion_v1 import ANALYSIS_ID
from research.existing_data_entry_aligned_full_causal_logic_completion_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "states",
    "state_roles",
    "entry_library",
    "entry_theses",
    "structural_exclusions",
    "exit_mapping",
    "candidate_library",
    "duplicate_check",
    "integrity",
    "canary",
    "coverage",
    "economics",
    "daily",
    "symbols",
    "causal_ex_top",
    "blocks",
    "selection",
    "selected_logic",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    theses = list(report.get("entry_theses") or [])
    lib = list(report.get("library_rows") or [])
    evals = list(report.get("economics_rows") or [])
    daily_rows = list(report.get("daily_rows") or [])
    block_rows = list(report.get("block_rows") or [])
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "states": list(report.get("states") or []),
        "state_roles": list(report.get("state_roles") or []),
        "entry_library": [{"ENTRY_ID": r.get("ENTRY_ID"), "ENTRY_TYPE": r.get("ENTRY_TYPE"), "A": r.get("A"), "B": r.get("B")} for r in theses],
        "entry_theses": theses,
        "structural_exclusions": [r for r in theses if not r.get("COMPLETE_STRATEGY_ELIGIBLE")],
        "exit_mapping": [
            {
                "ENTRY_ID": r.get("ENTRY_ID"),
                "HOLDING_THESIS_STATES": r.get("HOLDING_THESIS_STATES"),
                "INVALIDATION_CONDITION": r.get("INVALIDATION_CONDITION"),
                "TECHNICAL_EXIT_ID": r.get("TECHNICAL_EXIT_ID"),
                "EXIT_VARIANTS_PER_ENTRY": r.get("EXIT_VARIANTS_PER_ENTRY"),
            }
            for r in theses
            if r.get("COMPLETE_STRATEGY_ELIGIBLE")
        ],
        "candidate_library": lib,
        "duplicate_check": _kv({k: v for k, v in dict(report.get("duplicate_check") or {}).items() if k != "KEPT"}),
        "integrity": _kv(dict(report.get("integrity") or {})),
        "canary": _kv(dict(report.get("canary") or {})),
        "coverage": evals,
        "economics": evals,
        "daily": daily_rows,
        "symbols": list(report.get("symbol_rows") or []),
        "causal_ex_top": [
            {
                "STRATEGY_ID": r.get("STRATEGY_ID"),
                "top_symbol": r.get("top_symbol"),
                "top_symbol_pnl": r.get("top_symbol_pnl"),
                "CAUSAL_EX_TOP1_PNL": r.get("CAUSAL_EX_TOP1_PNL"),
                "G6": (r.get("g_table") or {}).get("G6"),
            }
            for r in evals
        ],
        "blocks": block_rows,
        "selection": _kv(dict(report.get("selection") or {})),
        "selected_logic": _kv(dict(report.get("selected_logic") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"LOGIC_COMPLETE: {d.get('LOGIC_COMPLETE')}",
            "TRUE_OOS: false",
            "CERTIFIED: false",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- selectable N: {a.get('16_selectable_complete_strategy_N')}",
            f"- canary PASS: {a.get('24_canary_PASS')}",
            f"- coverage PASS N: {a.get('25_coverage_PASS_N')}",
            f"- G1-G5 N: {a.get('26_G1_G5_N')}",
            f"- robust N: {a.get('28_robust_DEV_qualified_N')}",
            f"- selected: {a.get('29_selected_STRATEGY_ID')}",
            f"- selected EXIT: {a.get('37_selected_TECHNICAL_EXIT')}",
            "",
            "No Holdout/Stress/future. No Sizing. No Z3 candidate EXIT. No EXIT grid.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
