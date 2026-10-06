"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.freeze_group_mechanism_definitions_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Frozen_BG_CONT_VWAP",
    "Closed_Group_Pairs",
    "Trade_Ledger_486",
    "Symbol_Contribution",
    "8136_Diagnostic",
    "Winner_Concentration",
    "Entry_Path_Class",
    "MFE_MAE",
    "Giveback",
    "Exit_Reasons",
    "D2_D3_D4",
    "Group_Membership",
    "Cost_Headroom",
    "HM1_Comparison",
    "Root_Cause",
    "Safety",
)
LEDGER_KEYS = (
    "date",
    "symbol",
    "block",
    "event_time",
    "exit_hh",
    "exit_reason",
    "x0_bps",
    "x1_bps",
    "hold_min",
    "trade_mfe_bps",
    "trade_mae_bps",
    "time_to_mfe_min",
    "time_to_mae_min",
    "mfe_before_mae",
    "giveback_bps",
    "giveback_class",
    "entry_path_class",
    "time_above_vwap_min",
    "time_above_entry_min",
    "time_reclaim_to_vwap_loss_min",
    "meaningful_mfe_before_exit",
    "immediate_vwap_loss",
    "recovered_after_exit",
    "p1m_close_bps",
    "p2m_close_bps",
    "p3m_close_bps",
    "p5m_close_bps",
    "p10m_close_bps",
    "preq_group",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    rca = dict(report.get("rca") or {})
    ledger = [{k: r.get(k) for k in LEDGER_KEYS} for r in list(report.get("ledger") or [])]
    hm1 = dict(report.get("hm1_reference") or {})
    econ = dict(report.get("economics") or {})
    hashes = dict(report.get("mechanism_hashes") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "frozen_mechanism_id": report.get("frozen_mechanism_id"),
                "split_sha256": report.get("split_sha256"),
                "block_sha256": report.get("block_sha256"),
                "parameter_changed": False,
                "promoted": False,
                "frozen_validation_opened": False,
            }
        ),
        "Frozen_BG_CONT_VWAP": _kv({**dict(report.get("frozen_spec") or {}), **hashes}),
        "Closed_Group_Pairs": list(report.get("closed_pairs") or [{"empty": True}]),
        "Trade_Ledger_486": ledger or [{"empty": True}],
        "Symbol_Contribution": list(rca.get("symbol_contribution") or [{"empty": True}]),
        "8136_Diagnostic": _kv(
            {
                "with_all": rca.get("with_all"),
                "exclude_8136": rca.get("exclude_8136"),
                "only_8136": rca.get("only_8136"),
                "no_8136_strategy_created": True,
                "diagnostic_only": True,
            }
        ),
        "Winner_Concentration": _kv(dict(rca.get("winner_concentration") or {})),
        "Entry_Path_Class": list(rca.get("entry_path_classes") or [{"empty": True}]),
        "MFE_MAE": [
            {k: r.get(k) for k in ("date", "symbol", "trade_mfe_bps", "trade_mae_bps", "time_to_mfe_min", "time_to_mae_min", "mfe_before_mae", "mae_before_mfe", "time_above_entry_min", "time_above_vwap_min", "p1m_close_bps", "p2m_close_bps", "p3m_close_bps", "p5m_close_bps", "p10m_close_bps", "x0_bps")}
            for r in list(report.get("ledger") or [])
        ]
        or [{"empty": True}],
        "Giveback": _kv(dict(rca.get("giveback") or {})),
        "Exit_Reasons": _kv(dict(rca.get("exits") or {})),
        "D2_D3_D4": list(rca.get("block_evolution") or [{"empty": True}]),
        "Group_Membership": list(rca.get("group_membership") or [{"empty": True}]),
        "Cost_Headroom": _kv(dict(rca.get("cost_headroom") or {})),
        "HM1_Comparison": [
            {
                "candidate": "BG_CONT_VWAP",
                "X0": econ.get("mean_x0_bps"),
                "X1": econ.get("mean_x1_bps"),
                "PF": econ.get("profit_factor"),
                "median": econ.get("median_x0_bps"),
                "hit_rate": econ.get("hit_rate"),
                "DD": econ.get("max_dd_daily_mean_bps"),
                "symbol_concentration": econ.get("top_symbol_share_of_positive_bps"),
                "top_winner_dependence": (rca.get("winner_concentration") or {}).get("share_positive_pnl_top_5pct"),
                "block_stability": econ.get("block_mean_x0"),
                "combined": False,
            },
            {
                "candidate": "HM1_CONTROLLED_PULLBACK_RECLAIM",
                "X0": hm1.get("mean_x0_bps"),
                "X1": hm1.get("mean_x1_bps"),
                "PF": hm1.get("profit_factor"),
                "median": hm1.get("median_x0_bps"),
                "hit_rate": hm1.get("hit_rate"),
                "DD": hm1.get("max_dd_daily_mean_bps"),
                "symbol_concentration": hm1.get("top_symbol_share_of_positive_bps"),
                "top_winner_dependence": "unknown_without_hm1_ledger",
                "block_stability": hm1.get("block_mean_x0"),
                "combined": False,
                "tuned": False,
            },
        ],
        "Root_Cause": _kv(dict(rca.get("root_cause") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0", "promoted": False, "v27_bolted": False}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        "# FREEZE_GROUP_MECHANISM_DEFINITIONS_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        f"BG_CONT_VWAP frozen unchanged? **{a.get('Was_BG_CONT_VWAP_frozen_unchanged')}**",
        f"Mechanism hash: `{a.get('Exact_mechanism_hash')}`",
        f"Strategy parameter changed? **{a.get('Did_any_strategy_parameter_change')}**",
        f"Replay matches parent 486 / X0? **{a.get('replay_matches_parent_486')}** n=**{a.get('trade_n')}**",
        "",
        f"Symbols with positive gross edge: **{a.get('How_many_symbols_actually_contribute_positive_gross_edge')}** `{a.get('positive_gross_symbols')}`",
        f"8136 dominant? **{a.get('Is_8136_dominant')}**",
        f"Exclude 8136 (diagnosis only): `{a.get('Result_excluding_8136_diagnosis_only')}`",
        "",
        f"Top 1% winner share of positive PnL: **{a.get('Top_1pct_winner_contribution_of_positive_PnL')}**",
        f"Top 5%: **{a.get('Top_5pct_winner_contribution_of_positive_PnL')}**",
        f"Top 10%: **{a.get('Top_10pct_winner_contribution_of_positive_PnL')}**",
        f"Never profitable frac: **{a.get('losers_never_profitable_frac')}**",
        f"Profitable then lose frac: **{a.get('losers_profitable_then_lose_frac')}**",
        f"Stall n: **{a.get('stall_n')}**",
        f"Giveback average/median: **{a.get('Average_giveback')}** / **{a.get('Median_giveback')}**",
        f"VWAP_LOSS role: **{a.get('Does_VWAP_LOSS_primarily')}**",
        "",
        f"Why median negative / mean +6.79: {a.get('Why_median_negative_while_mean_positive')}",
        f"D4 composition vs mechanism: **{a.get('Composition_change_or_stronger_mechanism')}**",
        f"Stable without top few winners? **{a.get('Is_edge_stable_without_top_few_winners')}** (mean without top 5% = `{a.get('mean_x0_without_top_5pct')}`)",
        "",
        f"Root cause: **{a.get('Root_cause')}** (primary `{a.get('Root_cause_primary')}`)",
        f"EXIT research justified next? **{a.get('Is_EXIT_research_justified_next')}**",
        "",
        f"Any Complete Strategy promoted? **{a.get('Any_Complete_Strategy_promoted')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
        f"New paid data? **{a.get('New_paid_data')}**",
        f"Kabu50? **{a.get('Kabu50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
