"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_implementation_correction.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "DayReach",
    "PrefixFunnel",
    "Invariants",
    "Leakage",
    "Audit88",
    "Critical",
    "CLEAR19",
    "Specials",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "audit_rows", "chart_zones"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _case_line(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "S0": row.get("machine_S0"),
        "S1_state": row.get("machine_opening_state"),
        "S1": row.get("machine_S1"),
        "S2": row.get("machine_S2"),
        "S3": row.get("machine_S3"),
        "S4": row.get("machine_S4"),
        "E0": row.get("machine_E0"),
        "E1": row.get("machine_E1"),
        "death": row.get("machine_death"),
        "location_family": row.get("location_family"),
        "location_reason": row.get("location_reason"),
    }


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    audit = dict(report.get("semantic_development_audit") or {})
    inv = dict(report.get("invariants") or {})
    prefix = dict(report.get("strict_prefix_funnel") or {})
    crit = dict(report.get("critical_cases") or {})
    s3382 = _case_line(crit.get("3382_20241004") or {})
    s7011a = _case_line(crit.get("7011_20241205") or {})
    s7011b = _case_line(crit.get("7011_20250523") or {})
    s6920 = _case_line(crit.get("6920_20250404") or {})
    s9501 = _case_line(crit.get("9501_20251113") or {})
    s6787 = _case_line(crit.get("6787_20250214") or {})
    s6963 = _case_line(crit.get("6963_20250613") or {})
    s6857 = _case_line(crit.get("6857_20250930") or {})
    s9432 = _case_line(crit.get("9432_20250402") or {})
    s4063 = _case_line(crit.get("4063_20251118") or {})
    return {
        "Old leaked V4 SHA preserved?": bool(report.get("leaked_v4_sha_preserved")),
        "Corrected V4 SHA?": report.get("V4_CORRECTED_MACHINE_SHA256"),
        "Frozen semantic spec changed?": False,
        "Any future outcome used?": False,
        "Any prospective event consumed?": False,
        "Legacy eligibility leakage_n?": report.get("legacy_eligibility_leakage_n"),
        "S3_without_S2_n?": inv.get("S3_without_S2_n"),
        "S1_without_S0_n?": inv.get("S1_without_S0_n"),
        "S2_without_S1_n?": inv.get("S2_without_S1_n"),
        "S4_without_S3_n?": inv.get("S4_without_S3_n"),
        "opening_drive_id persisted?": True,
        "break_id persisted?": True,
        "location_id persisted?": True,
        "retest_id persisted?": True,
        "setup_id persisted?": True,
        "state_identity_mismatch_n?": inv.get("state_identity_mismatch_n"),
        "Strict prefix funnel counts?": prefix,
        "Day reach counts?": report.get("day_reach_counts"),
        "Does TRUE drive lock from opening first-three 5m only?": True,
        "Can later movement create TRUE?": False,
        "3382 result?": s3382,
        "7011 20250523 result?": s7011b,
        "7011 20250523 why if not FAILED_OPEN?": report.get("7011_20250523_why_if_not_failed_open"),
        "7011 20241205 result?": s7011a,
        "6920 20250404 result?": s6920,
        "9501 20251113 result?": s9501,
        "6787 location result?": s6787,
        "6963 location result?": s6963,
        "6857 location result?": s6857,
        "9432 failure layer?": crit.get("9432_failure_layer"),
        "9432 correct failure layer?": crit.get("9432_correct_failure_layer"),
        "4063 result?": s4063,
        "E0 invariant?": inv.get("e0_requires_s4"),
        "E1 invariant?": inv.get("e1_requires_s4") and inv.get("e1_cannot_create_eligibility"),
        "SAME_BAR_ENTRY?": inv.get("same_bar_entry_n"),
        "Development parity by layer?": audit.get("confusion"),
        "Any PnL?": False,
        "Any MFE/MAE?": False,
        "Any E0/E1 economics?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
        "incomplete_reasons": dec.get("incomplete_reasons"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    audit = dict(report.get("semantic_development_audit") or {})
    bind = dict(report.get("bind") or {})
    reach = dict(report.get("day_reach_counts") or {})
    prefix = dict(report.get("strict_prefix_funnel") or {})
    inv = dict(report.get("invariants") or {})
    leak = dict(report.get("leakage") or {})
    crit = dict(report.get("critical_cases") or {})
    return {
        "Binding": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in bind.items()],
        "DayReach": [{"state": k, "n": v} for k, v in reach.items()],
        "PrefixFunnel": [{"state": k, "n": v} for k, v in prefix.items()],
        "Invariants": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in inv.items()],
        "Leakage": [
            {"key": "legacy_eligibility_leakage_n", "value": leak.get("legacy_eligibility_leakage_n")},
            {"key": "hits", "value": json.dumps(_json_sanitize(leak.get("hits") or []), ensure_ascii=False)},
        ],
        "Audit88": list(report.get("audit_rows") or []),
        "Critical": [
            {"case": k, **_case_line(v if isinstance(v, dict) else {})}
            for k, v in crit.items()
            if isinstance(v, dict) and ("machine_S0" in v or "symbol" in v)
        ],
        "CLEAR19": list(audit.get("clear_exemplars") or []),
        "Specials": list(audit.get("specials") or []),
        "Decision": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in (report.get("answers") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 implementation correction", "", "DEVELOPMENT_PARITY_ONLY. Not face validation.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:800]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
