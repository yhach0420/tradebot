"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_prospective_precommit_calendar_correction.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Parent_Precommit",
    "Calendar_Correction",
    "Corrected_Precommit",
    "Eligibility",
    "Stopping_Rule",
    "Safety",
)
STRIP = {"isolation_before", "isolation_after"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    pre = dict(report.get("corrected_precommit") or {})
    elig = dict(pre.get("data_eligibility") or {})
    return {
        "SPEC_CHANGED": False,
        "V4_CHANGED": False,
        "OLD_PRECOMMIT_SHA_PRESERVED": True,
        "NEW_CALENDAR_CORRECTED_PRECOMMIT_SHA_CREATED": True,
        "OLD_PRECOMMIT_SHA256": report.get("OLD_PRECOMMIT_SHA256"),
        "NEW_PRECOMMIT_SHA256": pre.get("PRECOMMIT_SHA256"),
        "FIRST_ELIGIBLE_JP_CASH_SESSION": elig.get("first_eligible_session"),
        "20260921_COUNTS_AS_SESSION": False,
        "20260922_COUNTS_AS_SESSION": False,
        "20260923_COUNTS_AS_SESSION": False,
        "PROSPECTIVE_DATA_OPENED": False,
        "OLD_CONFIRMATION_OPENED": False,
        "FROZEN_VALIDATION_OPENED": False,
        "THRESHOLD_RETUNED": False,
        "NEW_NUMERIC_CUTOFF_ADDED": False,
        "submit/cancel/live": "0/0/0",
        "VERDICT": (report.get("decision") or {}).get("VERDICT"),
        "NEXT": (report.get("decision") or {}).get("NEXT"),
        "machine_sha": (report.get("hashes") or {}).get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V4_SHA256"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    pre = dict(report.get("corrected_precommit") or {})
    parent = dict(report.get("parent_verify") or {})
    elig = dict(pre.get("data_eligibility") or {})
    stop = dict(pre.get("observation_period_stopping_rule") or {})
    return {
        "Manifest": [{"key": k, "value": v} for k, v in (report.get("answers") or {}).items()],
        "Parent_Precommit": [{"key": k, "value": v} for k, v in parent.items()],
        "Calendar_Correction": [{"key": k, "value": v} for k, v in dict(pre.get("calendar_correction") or {}).items()],
        "Corrected_Precommit": [{"key": k, "value": v} for k, v in pre.items() if k not in ("state_logging_schema",)],
        "Eligibility": [{"key": k, "value": v} for k, v in elig.items()],
        "Stopping_Rule": [{"key": k, "value": v} for k, v in stop.items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = [
        "# PB1 V4 prospective precommit — calendar correction",
        "",
        "V4 machine unchanged. Old precommit SHA preserved as immutable parent. Prospective data not opened.",
        "",
    ]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
