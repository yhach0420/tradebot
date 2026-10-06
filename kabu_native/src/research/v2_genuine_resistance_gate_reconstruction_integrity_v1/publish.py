"""Integrity repair report. Candidate economics are included only after every gate is zero."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from research.event_time_impulse_complete_strategy_v2.identity import bind, sha256_obj
from research.symbol_setup_baseline_complete_strategy_precommit.contract import EXTENSION17, ORIGINAL18
from research.v2_genuine_resistance_gate_complete_strategy_v1 import DETECTOR_SHA256, EXPECTED_PF, EXPECTED_PNL_YEN, EXPECTED_TRADE_N, STRATEGY_SHA256, SIGNAL_SHA256
from research.v2_genuine_resistance_gate_complete_strategy_v1.publish import _econ, _period, _ratchets, _reason
from research.v2_genuine_resistance_gate_complete_strategy_v1.scan import scan
from research.v2_structural_resistance_detector_validation_v1.rules import detector_document

OUT = Path("results/research/v2_genuine_resistance_gate_reconstruction_integrity_v1")


def publish() -> dict[str, Any]:
    OUT.mkdir(parents=True, exist_ok=True)
    bound = bind()
    detector_sha = sha256_obj(detector_document())
    raw = scan()
    base = [r for r in raw["baseline"] if r.get("pnl_yen") is not None]
    uids = list(raw["all_signal_uids"])
    uids2 = list(raw["all_signal_uids_again"])
    unique = len(set(uids))
    duplicate = len(uids) - unique
    trade_uids = [r.get("signal_uid") for r in base]
    mapped = sum(1 for u in trade_uids if u)
    multi = len(trade_uids) - len(set(u for u in trade_uids if u))
    unmapped = len(base) - mapped
    base_econ = _econ(base)
    parity = (
        base_econ["trade_n"] == EXPECTED_TRADE_N
        and abs(base_econ["pnl_yen"] - EXPECTED_PNL_YEN) < 1e-6
        and base_econ["pf"] is not None
        and abs(float(base_econ["pf"]) - EXPECTED_PF) < 1e-12
        and len({(r["date"], r["session"], r["symbol"], r["entry_t"], r["exit_t"], r["pnl_yen"]) for r in base}) == len(base)
    )
    reconstructions_match = uids == uids2
    integrity = (
        len(uids) == 11930
        and unique == 11930
        and duplicate == 0
        and reconstructions_match
        and mapped == 11902
        and unmapped == 0
        and multi == 0
        and int(raw["classified_trade_n"]) == 11902
        and int(raw["detector_miss"]) == 0
        and int(raw["ratchet_observer_mismatch"]) == 0
        and parity
        and detector_sha == DETECTOR_SHA256
        and bound["strategy_sha256"] == STRATEGY_SHA256
    )
    ambiguous = raw["ambiguous"]
    report: dict[str, Any] = {
        "study": "REPAIR_RECONSTRUCTION_INTEGRITY_ONLY_V1",
        "integrity_cleared": integrity,
        "reconstruction_abort": not integrity,
        "baseline": base_econ,
        "trade_by_trade_parity": parity,
        "full_signal_n": len(uids),
        "unique_signal_uid_n": unique,
        "duplicate_signal_uid_n": duplicate,
        "reconstructions_identical": reconstructions_match,
        "mapped_trade_n": mapped,
        "unmapped_trade_n": unmapped,
        "multi_mapped_trade_n": multi,
        "ambiguous_timestamp_n": len(ambiguous),
        "ambiguous_resolved_n": sum(1 for row in ambiguous if row.get("chosen_signal_uid")),
        "ambiguous_unresolved_n": sum(1 for row in ambiguous if not row.get("chosen_signal_uid")),
        "ambiguous_cases": ambiguous,
        "classified_trade_n": raw["classified_trade_n"],
        "detector_miss": raw["detector_miss"],
        "detector_multi_match": multi,
        "previous_ratchet_observer_mismatch_n": 2,
        "ratchet_observer_mismatch": raw["ratchet_observer_mismatch"],
        "old_observer_rca": raw["old_observer_mismatches"],
        "detector_sha256": detector_sha,
        "detector_unchanged": detector_sha == DETECTOR_SHA256,
        "strategy_sha256": bound["strategy_sha256"],
        "signal_sha256": bound["signal_sha256"],
        "v2_changed": False,
        "new_data_acquired": False,
        "submit_cancel_live": [0, 0, 0],
    }
    if integrity:
        cand = [r for r in raw["candidate"] if r.get("pnl_yen") is not None]
        b100 = raw["baseline_100"]
        c100 = raw["candidate_100"]
        dates = sorted({str(r["date"]) for r in base})
        folds = [set(dates[i * len(dates) // 3:(i + 1) * len(dates) // 3]) for i in range(3)]
        cand_econ = _econ(cand)
        base_block = base_econ
        periods = {
            "ORIGINAL18": {"baseline": _period(base, set(ORIGINAL18)), "candidate": _period(cand, set(ORIGINAL18))},
            "EXTENSION17": {"baseline": _period(base, set(EXTENSION17)), "candidate": _period(cand, set(EXTENSION17))},
            "FOLD1": {"baseline": _period(base, folds[0]), "candidate": _period(cand, folds[0])},
            "FOLD2": {"baseline": _period(base, folds[1]), "candidate": _period(cand, folds[1])},
            "FOLD3": {"baseline": _period(base, folds[2]), "candidate": _period(cand, folds[2])},
        }
        b100e, c100e = _econ(b100), _econ(c100)
        pf_up = (cand_econ["pf"] or 0) > (base_block["pf"] or 0)
        pnl_up = cand_econ["pnl_yen"] > base_block["pnl_yen"]
        period_ok = all((p["candidate"]["pnl_yen"] >= p["baseline"]["pnl_yen"]) and (p["candidate"]["pf"] or 0) >= (p["baseline"]["pf"] or 0) for p in periods.values())
        lat_ok = c100e["pnl_yen"] > b100e["pnl_yen"] and (c100e["pf"] or 0) >= (b100e["pf"] or 0)
        broad = pf_up and pnl_up and period_ok and lat_ok
        collapsed = pf_up and cand_econ["pnl_yen"] < base_block["pnl_yen"] * 0.5
        if broad:
            verdict, nxt = "V2_GENUINE_RESISTANCE_GATE_SUPPORTED_V1", "AUDIT_WHAT_EDGE_REMAINS_IN_INCIDENTAL_HIGH_ENTRIES_V1"
        elif collapsed:
            verdict, nxt = "V2_GENUINE_RESISTANCE_GATE_PRECISION_RECALL_TRADEOFF_V1", "KEEP_V2_UNCHANGED"
        elif not period_ok:
            verdict, nxt = "V2_GENUINE_RESISTANCE_GATE_NOT_ROBUST_V1", "KEEP_V2_UNCHANGED"
        else:
            verdict, nxt = "V2_GENUINE_RESISTANCE_ASSOCIATION_NOT_ACTIONABLE_V1", "KEEP_V2_UNCHANGED"
        report["verdict"] = verdict
        report["next"] = nxt
        report["candidate_economics"] = {
            "full_signal_n": raw["full_signal_n"],
            "accepted_n": raw["accepted_n"],
            "rejected_n": int(raw["full_signal_n"]) - int(raw["accepted_n"]),
            "baseline": base_block,
            "candidate": cand_econ,
            "periods": periods,
            "ratchet_groups": _ratchets(cand),
            "exit_reasons": {
                "BREAK_SUPPORT_FAILURE": _reason(cand, "BREAK_SUPPORT_FAILURE"),
                "IMPULSE_EXHAUSTED": _reason(cand, "IMPULSE_EXHAUSTED"),
            },
            "asof_100": {"baseline": b100e, "candidate": c100e},
            "complete_strategy_improvement": broad,
            "robust_across_periods": period_ok,
            "latency_100_maintained": lat_ok,
            "gate_actionable": broad,
            "note": "Recomputed after integrity clearance. Prior 3013/8917 counts are not used.",
        }
    else:
        report["verdict"] = "INTEGRITY_STILL_BLOCKED_V1"
        report["next"] = "REPAIR_RECONSTRUCTION_INTEGRITY_ONLY"
        report["candidate_economics"] = None
    public = report
    (OUT / "report.json").write_text(json.dumps(public, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(
        f"# {report['verdict']}\n\nNEXT: {report['next']}\n\nintegrity_cleared={integrity}\nreconstruction_abort={not integrity}\n",
        encoding="utf-8",
    )
    wb = Workbook()
    sheet = wb.active
    sheet.title = "summary"
    for key in ("verdict", "next", "integrity_cleared", "reconstruction_abort", "full_signal_n", "unique_signal_uid_n", "duplicate_signal_uid_n", "mapped_trade_n", "unmapped_trade_n", "multi_mapped_trade_n", "classified_trade_n", "detector_miss", "detector_multi_match", "ratchet_observer_mismatch"):
        sheet.append([key, public[key]])
    amb = wb.create_sheet("ambiguous_signals")
    amb.append(["json"])
    for row in ambiguous:
        amb.append([json.dumps(row, default=str)])
    rca = wb.create_sheet("ratchet_rca")
    rca.append(["json"])
    for row in raw["old_observer_mismatches"]:
        rca.append([json.dumps(row, default=str)])
    if public.get("candidate_economics"):
        econ = wb.create_sheet("candidate_after_integrity")
        econ.append([json.dumps(public["candidate_economics"], default=str)])
    wb.save(OUT / "audit.xlsx")
    return {"verdict": report["verdict"], "next": report["next"], "integrity_cleared": integrity}


if __name__ == "__main__":
    print(publish())
