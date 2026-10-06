"""Write report.json / report.md / audit.xlsx only. No mass CSV. No Runtime write."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.target_price_contract_v3 import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "target_price_contract_v3"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)

SHEET_ORDER = (
    "Summary",
    "Three_Contracts",
    "Qty_Sites",
    "Plus600",
    "Age_Dist",
    "Sensitivity",
    "Locked",
    "Waterfall",
    "First_Fail",
    "By_Anchor",
    "By_Day",
    "By_Symbol",
    "Missingness_SMD",
    "Score_Decile",
    "Cohort",
    "Label_Stability",
    "Decision",
    "B_Status",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def flatten_smd(smd_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in smd_rows:
        rec = {"feature": r.get("feature"), "smd_valid_minus_null": r.get("smd_valid_minus_null"), "flag": r.get("flag")}
        for side in ("valid", "null"):
            q = r.get(side) or {}
            for k, v in q.items():
                rec[f"{side}_{k}"] = v
        out.append(rec)
    return out


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("gates") or {}
    ch = report.get("chosen") or {}
    return "\n".join(
        [
            "# TARGET PRICE CONTRACT V3",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE TARGET AUDIT ONLY. C14 / Runtime / CLOCK / ENTRY / EXIT unchanged.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "C REBUILD was not started this run. Model search was not started.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"SESSION_SOT_VALID: {g.get('SESSION_SOT_VALID')}",
            f"PM_CONTINUOUS_END: {g.get('PM_CONTINUOUS_END')}",
            f"SELECTED_MARK_CONTRACT: {g.get('SELECTED_MARK_CONTRACT')}",
            f"SELECTED_MAX_MARK_AGE_SEC: {g.get('SELECTED_MAX_MARK_AGE_SEC')}",
            f"ROWS_TOTAL: {g.get('ROWS_TOTAL')}",
            f"PRIMARY_ROWS: {g.get('PRIMARY_ROWS')}",
            f"PRIMARY_COVERAGE_RATE: {g.get('PRIMARY_COVERAGE_RATE')}",
            f"MEDIAN_TARGET_VALID_N_PER_COHORT: {g.get('MEDIAN_TARGET_VALID_N_PER_COHORT')}",
            f"P10_TARGET_VALID_N_PER_COHORT: {g.get('P10_TARGET_VALID_N_PER_COHORT')}",
            f"TARGET_MISSINGNESS_SELECTION_BIAS: {g.get('TARGET_MISSINGNESS_SELECTION_BIAS')}",
            f"SCORE_DECILE_COVERAGE_RANGE: {g.get('SCORE_DECILE_COVERAGE_RANGE')}",
            f"SYMBOL_COVERAGE_RANGE: {g.get('SYMBOL_COVERAGE_RANGE')}",
            f"ITAYOSE_BASE_ROW_N: {g.get('ITAYOSE_BASE_ROW_N')}",
            f"CLOSING_AUCTION_ENDPOINT_ROW_N: {g.get('CLOSING_AUCTION_ENDPOINT_ROW_N')}",
            f"FUTURE_EVENT_USED: {g.get('FUTURE_EVENT_USED')}",
            f"TARGET_COHORT_COVERAGE_ADEQUATE: {g.get('TARGET_COHORT_COVERAGE_ADEQUATE')}",
            f"TARGET_V3_READY_FOR_C_REBUILD: {g.get('TARGET_V3_READY_FOR_C_REBUILD')}",
            f"VERDICT: {g.get('VERDICT')}",
            f"RECOMMENDED_NEXT_STEP: {g.get('RECOMMENDED_NEXT_STEP')}",
            "",
            "## Contracts",
            "",
            "A MARKET PHASE: current TSE. AM 09:00-11:30, PM Zaraba 12:30-15:25,",
            "auction 15:25-15:30, close 15:30. e1_x22 PM 15:00 is not used for TARGET.",
            "Runtime Dual Lane SESSION_CLOSE 15:00 is unchanged.",
            "",
            "B PRICE MARK: M1 quote mid ask>bid, not itayose/special, no qty>=100.",
            "M2 CurrentPrice status in {1,2}. M3 = M1 else M2, same rule both sides.",
            "M0 = V2 last_executable_mid (qty+executable+5s), reference only.",
            "",
            "C EXECUTION: is_executable_continuous_board + WAIT_SEC + qty. Frozen.",
            "",
            f"Selection: {ch.get('selection')} mark={ch.get('mark')} age={ch.get('max_mark_age_sec')}",
            "Selector did not use PnL, Spearman vs returns, or TopK.",
            "",
            "## B status",
            "",
            "B0: HISTORICAL_REFERENCE_ONLY",
            "B1: B1_SCORE_THRESHOLD_NO_ROBUST_IMPROVEMENT",
            "per-feature threshold: DO_NOT_START",
            "",
            "## STOP",
            "",
            "Audit complete. C REBUILD not started. Runtime not changed.",
            "",
        ]
    )
