"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.anchor_economic_sensitivity import ANALYSIS_ID, DOCUMENT_ID, PRIMARY_SHIFTS, TASK_LABEL

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "anchor_economic_sensitivity"
JST = timezone(timedelta(hours=9))
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            if v == float("inf"):
                v = "Infinity"
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _d(block: dict, shift: str, key: str) -> Any:
    return (block.get(shift) or {}).get(key)


def build_report(*, inventory_check: dict[str, Any], prior_meta: dict[str, Any], result: dict[str, Any]) -> dict[str, Any]:
    iso = result["iso_decomp"]
    port = result["port_stats"]
    b = result["boundary"]
    clf = result["classify"]
    iso_all, iso_h, iso_d = iso["ALL"], iso["POST_FREEZE_HOLDOUT"], iso["DEVELOPMENT"]
    port_all = port["ALL"]
    tail_iso = result["tail_iso_holdout"]
    tail_port = result["tail_port_holdout"]

    required = {
        "COMMON_SELECTED_RATE_M5": _d(iso_all, "M5", "COMMON_SELECTED_RATE"),
        "COMMON_SELECTED_RATE_P5": _d(iso_all, "P5", "COMMON_SELECTED_RATE"),
        "COMMON_TIMING_PNL_DELTA_M5": _d(iso_all, "M5", "COMMON_TIMING_PNL_DELTA"),
        "COMMON_TIMING_PNL_DELTA_P5": _d(iso_all, "P5", "COMMON_TIMING_PNL_DELTA"),
        "REF_ONLY_PNL_M5": _d(iso_all, "M5", "REF_ONLY_PNL"),
        "SHIFT_ONLY_PNL_M5": _d(iso_all, "M5", "SHIFT_ONLY_PNL"),
        "REF_ONLY_PNL_P5": _d(iso_all, "P5", "REF_ONLY_PNL"),
        "SHIFT_ONLY_PNL_P5": _d(iso_all, "P5", "SHIFT_ONLY_PNL"),
        "FILL_RATE_REF": _d(iso_all, "M5", "FILL_RATE_REF"),
        "FILL_RATE_M5": _d(iso_all, "M5", "FILL_RATE_SHIFT"),
        "FILL_RATE_P5": _d(iso_all, "P5", "FILL_RATE_SHIFT"),
        "FILL_PRICE_DELTA_BPS_M5": _d(iso_all, "M5", "FILL_PRICE_DELTA_BPS"),
        "FILL_PRICE_DELTA_BPS_P5": _d(iso_all, "P5", "FILL_PRICE_DELTA_BPS"),
        "EXIT_REASON_SWITCH_RATE_M5": _d(iso_all, "M5", "EXIT_REASON_SWITCH_RATE"),
        "EXIT_REASON_SWITCH_RATE_P5": _d(iso_all, "P5", "EXIT_REASON_SWITCH_RATE"),
        "TOP5_BOUNDARY_SENSITIVITY": (b.get("ALL") or {}).get("M5", {}).get("TOP5_BOUNDARY_SENSITIVITY"),
        "FIRST_ENTRY_PNL_REF": _d(port_all, "REFERENCE", "FIRST_ENTRY_PNL"),
        "FIRST_ENTRY_PNL_M5": _d(port_all, "M5", "FIRST_ENTRY_PNL"),
        "FIRST_ENTRY_PNL_P5": _d(port_all, "P5", "FIRST_ENTRY_PNL"),
        "REENTRY_PNL_REF": _d(port_all, "REFERENCE", "REENTRY_PNL"),
        "REENTRY_PNL_M5": _d(port_all, "M5", "REENTRY_PNL"),
        "REENTRY_PNL_P5": _d(port_all, "P5", "REENTRY_PNL"),
        "FIRST_ENTRY_ONLY_REF_PNL": _d(port_all, "REFERENCE", "FIRST_ENTRY_ONLY_PNL"),
        "FIRST_ENTRY_ONLY_M5_PNL": _d(port_all, "M5", "FIRST_ENTRY_ONLY_PNL"),
        "FIRST_ENTRY_ONLY_P5_PNL": _d(port_all, "P5", "FIRST_ENTRY_ONLY_PNL"),
        "REENTRY_INTERACTION_CONTRIBUTION": result.get("REENTRY_INTERACTION_CONTRIBUTION"),
        "HOLDOUT_REFERENCE_PNL": _d(iso_h, "M5", "pnl_ref_selected"),
        "HOLDOUT_M5_PNL": _d(iso_h, "M5", "pnl_shift_selected"),
        "HOLDOUT_P5_PNL": _d(iso_h, "P5", "pnl_shift_selected"),
        "HOLDOUT_LOO_SIGN_STABILITY": (tail_iso.get("M5_vs_REF") or tail_iso.get("M5") or {}).get("HOLDOUT_LOO_SIGN_STABILITY"),
        "TAIL_CONCENTRATION": (tail_iso.get("M5_vs_REF") or tail_iso.get("M5") or {}).get("TAIL_CONCENTRATION"),
        "PRIMARY_ROOT_CAUSE": clf.get("PRIMARY_ROOT_CAUSE"),
        "SECONDARY_ROOT_CAUSE": clf.get("SECONDARY_ROOT_CAUSE"),
        "CROSS_SECTIONAL_RANK_ROBUSTNESS": clf.get("CROSS_SECTIONAL_RANK_ROBUSTNESS"),
        "ECONOMIC_TIMING_ROBUSTNESS": clf.get("ECONOMIC_TIMING_ROBUSTNESS"),
        "EXACT_CLOCK_OVERFIT_EVIDENCE": clf.get("EXACT_CLOCK_OVERFIT_EVIDENCE"),
        "verdict": clf.get("verdict"),
    }

    micro_ok = bool(_d(iso_all, "M5", "microstructure_available") or _d(iso_all, "P5", "microstructure_available"))
    if micro_ok:
        micro_note = (
            "bid/ask/spread/imbalance at anchor and +1m/+5m mid returns are diagnostic outcomes only "
            "(not used as ENTRY/EXIT rules). Fill latency is fill_time - t0."
        )
    else:
        micro_note = (
            "bid/ask/spread/imbalance and +1m/+5m returns were not persisted in the prior artifact. "
            "Fill latency uses fill_time - t0; fill price/limit deltas are present."
        )

    report = {
        "ANALYSIS_ID": ANALYSIS_ID,
        "DOCUMENT_ID": DOCUMENT_ID,
        "label": TASK_LABEL,
        "generated_at_jst": datetime.now(JST).isoformat(timespec="seconds"),
        "runtime_changed": False,
        "clock_grid_changed": False,
        "strategy_changed": False,
        "best_shift_adopted": False,
        "safety": {"submit": 0, "cancel": 0, "live": 0, "offline_research_only": True},
        "source_study": "results/research/anchor_timing_robustness",
        "day_contract": inventory_check,
        "prior_rank_summary": {
            "verdict": prior_meta.get("verdict"),
            "RANK_SPEARMAN_M5": prior_meta.get("RANK_SPEARMAN_M5"),
            "RANK_SPEARMAN_P5": prior_meta.get("RANK_SPEARMAN_P5"),
            "TOP5_OVERLAP_M5": prior_meta.get("TOP5_OVERLAP_M5"),
            "TOP5_OVERLAP_P5": prior_meta.get("TOP5_OVERLAP_P5"),
        },
        "PRIMARY_INTERPRETATION_PERIOD": "POST_FREEZE_HOLDOUT",
        "note_holdout_n": "n=7 holdout days; not claimed as statistically conclusive.",
        "counterfactual_note": (
            "ONE_ENTRY_PER_SYMBOL_PER_DAY is a ledger filter of existing full-grid fills "
            "(first fill_time per date×symbol). It is not an occupancy-causal re-replay and is not a strategy."
        ),
        "microstructure_note": micro_note,
        **required,
        "isolated_decomp": iso,
        "portfolio_reentry": {
            p: {sh: {k: v for k, v in st.items() if k != "lineage"} for sh, st in block.items()}
            for p, block in port.items()
        },
        "boundary": b,
        "holdout_daily_isolated": result["daily_iso_holdout"],
        "holdout_daily_portfolio": result["daily_port_holdout"],
        "holdout_loo_isolated": tail_iso,
        "holdout_loo_portfolio": tail_port,
        "classify": clf,
        "iso_lineage_m5": result["lineage"]["ALL"]["M5"],
        "iso_lineage_p5": result["lineage"]["ALL"]["P5"],
        "port_lineage_ref": result["port_lineage"]["ALL"]["REFERENCE"],
        "port_lineage_m5": result["port_lineage"]["ALL"]["M5"],
        "port_lineage_p5": result["port_lineage"]["ALL"]["P5"],
        "exit_transitions_m5": _d(iso_all, "M5", "exit_transition"),
        "exit_transitions_p5": _d(iso_all, "P5", "exit_transition"),
    }
    return json_sanitize(report)


def write_md(report: dict[str, Any]) -> str:
    iso = report.get("isolated_decomp") or {}
    clf = report.get("classify") or {}
    lines = [
        "# TRADEBOT — Anchor Economic Sensitivity Decomposition",
        "",
        "Offline causal diagnostic. CLOCK_GRID / ENTRY / EXIT / Runtime unchanged. No best-shift adoption.",
        "Primary interpretation: POST_FREEZE_HOLDOUT (n=7; not statistically conclusive).",
        "",
        f"verdict: {report.get('verdict')}",
        f"PRIMARY_ROOT_CAUSE: {report.get('PRIMARY_ROOT_CAUSE')}",
        f"SECONDARY_ROOT_CAUSE: {report.get('SECONDARY_ROOT_CAUSE')}",
        f"CROSS_SECTIONAL_RANK_ROBUSTNESS: {report.get('CROSS_SECTIONAL_RANK_ROBUSTNESS')}",
        f"ECONOMIC_TIMING_ROBUSTNESS: {report.get('ECONOMIC_TIMING_ROBUSTNESS')}",
        f"EXACT_CLOCK_OVERFIT_EVIDENCE: {report.get('EXACT_CLOCK_OVERFIT_EVIDENCE')}",
        "",
        "## REQUIRED",
        f"COMMON_SELECTED_RATE_M5: {report.get('COMMON_SELECTED_RATE_M5')}",
        f"COMMON_SELECTED_RATE_P5: {report.get('COMMON_SELECTED_RATE_P5')}",
        f"COMMON_TIMING_PNL_DELTA_M5: {report.get('COMMON_TIMING_PNL_DELTA_M5')}",
        f"COMMON_TIMING_PNL_DELTA_P5: {report.get('COMMON_TIMING_PNL_DELTA_P5')}",
        f"REF_ONLY_PNL_M5: {report.get('REF_ONLY_PNL_M5')}",
        f"SHIFT_ONLY_PNL_M5: {report.get('SHIFT_ONLY_PNL_M5')}",
        f"REF_ONLY_PNL_P5: {report.get('REF_ONLY_PNL_P5')}",
        f"SHIFT_ONLY_PNL_P5: {report.get('SHIFT_ONLY_PNL_P5')}",
        f"FILL_RATE_REF: {report.get('FILL_RATE_REF')}",
        f"FILL_RATE_M5: {report.get('FILL_RATE_M5')}",
        f"FILL_RATE_P5: {report.get('FILL_RATE_P5')}",
        f"FILL_PRICE_DELTA_BPS_M5: {report.get('FILL_PRICE_DELTA_BPS_M5')}",
        f"FILL_PRICE_DELTA_BPS_P5: {report.get('FILL_PRICE_DELTA_BPS_P5')}",
        f"EXIT_REASON_SWITCH_RATE_M5: {report.get('EXIT_REASON_SWITCH_RATE_M5')}",
        f"EXIT_REASON_SWITCH_RATE_P5: {report.get('EXIT_REASON_SWITCH_RATE_P5')}",
        f"TOP5_BOUNDARY_SENSITIVITY: {report.get('TOP5_BOUNDARY_SENSITIVITY')}",
        f"FIRST_ENTRY_PNL_REF: {report.get('FIRST_ENTRY_PNL_REF')}",
        f"FIRST_ENTRY_PNL_M5: {report.get('FIRST_ENTRY_PNL_M5')}",
        f"FIRST_ENTRY_PNL_P5: {report.get('FIRST_ENTRY_PNL_P5')}",
        f"REENTRY_PNL_REF: {report.get('REENTRY_PNL_REF')}",
        f"REENTRY_PNL_M5: {report.get('REENTRY_PNL_M5')}",
        f"REENTRY_PNL_P5: {report.get('REENTRY_PNL_P5')}",
        f"FIRST_ENTRY_ONLY_REF_PNL: {report.get('FIRST_ENTRY_ONLY_REF_PNL')}",
        f"FIRST_ENTRY_ONLY_M5_PNL: {report.get('FIRST_ENTRY_ONLY_M5_PNL')}",
        f"FIRST_ENTRY_ONLY_P5_PNL: {report.get('FIRST_ENTRY_ONLY_P5_PNL')}",
        f"REENTRY_INTERACTION_CONTRIBUTION: {report.get('REENTRY_INTERACTION_CONTRIBUTION')}",
        f"HOLDOUT_REFERENCE_PNL: {report.get('HOLDOUT_REFERENCE_PNL')}",
        f"HOLDOUT_M5_PNL: {report.get('HOLDOUT_M5_PNL')}",
        f"HOLDOUT_P5_PNL: {report.get('HOLDOUT_P5_PNL')}",
        f"HOLDOUT_LOO_SIGN_STABILITY: {report.get('HOLDOUT_LOO_SIGN_STABILITY')}",
        f"TAIL_CONCENTRATION: {report.get('TAIL_CONCENTRATION')}",
        "",
        "## Isolated additive decomp (ALL / DEV / HOLDOUT) M5",
    ]
    for pname in ("ALL", "DEVELOPMENT", "POST_FREEZE_HOLDOUT"):
        d = (iso.get(pname) or {}).get("M5") or {}
        lines.append(
            f"- {pname}: observed={d.get('TOTAL_SHIFT_MINUS_REFERENCE_PNL')} "
            f"A_timing={d.get('COMMON_SYMBOL_TIMING_EFFECT')} "
            f"B_replace={d.get('SELECTION_REPLACEMENT_EFFECT')} "
            f"C_fill={d.get('FILL_CONVERSION_EFFECT')} residual={d.get('residual')}"
        )
    lines += ["", "## Isolated additive decomp P5"]
    for pname in ("ALL", "DEVELOPMENT", "POST_FREEZE_HOLDOUT"):
        d = (iso.get(pname) or {}).get("P5") or {}
        lines.append(
            f"- {pname}: observed={d.get('TOTAL_SHIFT_MINUS_REFERENCE_PNL')} "
            f"A_timing={d.get('COMMON_SYMBOL_TIMING_EFFECT')} "
            f"B_replace={d.get('SELECTION_REPLACEMENT_EFFECT')} "
            f"C_fill={d.get('FILL_CONVERSION_EFFECT')} residual={d.get('residual')}"
        )
    lines += [
        "",
        "## Development vs Holdout component signs (M5)",
    ]
    ddev = (iso.get("DEVELOPMENT") or {}).get("M5") or {}
    dho = (iso.get("POST_FREEZE_HOLDOUT") or {}).get("M5") or {}
    for key in (
        "COMMON_SYMBOL_TIMING_EFFECT",
        "SELECTION_REPLACEMENT_EFFECT",
        "FILL_CONVERSION_EFFECT",
        "TOTAL_SHIFT_MINUS_REFERENCE_PNL",
    ):
        a, b = ddev.get(key), dho.get(key)
        same = None
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            same = (a > 0) == (b > 0) if a != 0 and b != 0 else "zero_in_one_period"
        lines.append(f"- {key}: DEV={a} HOLDOUT={b} same_sign={same}")
    lines += [
        "",
        "## Reentry (full-grid ledger)",
        f"reentry_shrink_m5: {clf.get('reentry_shrink_m5')}",
        f"reentry_shrink_m5_holdout: {clf.get('reentry_shrink_m5_holdout')}",
        f"full_gap_m5: {clf.get('full_gap_m5')}",
        f"first_only_gap_m5: {clf.get('first_only_gap_m5')}",
        "",
        report.get("counterfactual_note"),
        "",
        report.get("microstructure_note"),
        "",
        "STOP. Do not change CLOCK_GRID / ENTRY / EXIT / reentry rules from this result.",
        "submit/cancel/live=0/0/0",
    ]
    return "\n".join(lines) + "\n"


def write_xlsx(report: dict[str, Any], path: Path) -> None:
    def flat_decomp() -> list[dict[str, Any]]:
        rows = []
        for period, block in (report.get("isolated_decomp") or {}).items():
            for sh, d in (block or {}).items():
                row = {"period": period, "shift_key": sh}
                for k, v in d.items():
                    if k in {"exit_transition", "fill_class_counts"}:
                        continue
                    row[k] = v
                rows.append(row)
        return rows or [{"status": "empty"}]

    def flat_port() -> list[dict[str, Any]]:
        rows = []
        for period, block in (report.get("portfolio_reentry") or {}).items():
            for sh, d in (block or {}).items():
                row = {"period": period, "shift_key": sh}
                for k, v in d.items():
                    if isinstance(v, dict) and k == "all":
                        for kk, vv in v.items():
                            row[f"all_{kk}"] = vv
                    else:
                        row[k] = v
                rows.append(row)
        return rows or [{"status": "empty"}]

    def daily(kind: str) -> list[dict[str, Any]]:
        src = report.get(kind) or {}
        rows = []
        for sh, lst in src.items():
            for r in lst or []:
                rows.append({"shift_key": sh, **r})
        return rows or [{"status": "empty"}]

    def loo(kind: str) -> list[dict[str, Any]]:
        src = report.get(kind) or {}
        rows = []
        for sh, body in src.items():
            for r in body.get("loo") or []:
                rows.append({"shift_key": sh, "TAIL_CONCENTRATION": body.get("TAIL_CONCENTRATION"), **r})
        return rows or [{"status": "empty"}]

    sheets = {
        "Summary": [
            {
                "verdict": report.get("verdict"),
                "PRIMARY_ROOT_CAUSE": report.get("PRIMARY_ROOT_CAUSE"),
                "SECONDARY_ROOT_CAUSE": report.get("SECONDARY_ROOT_CAUSE"),
                "CROSS_SECTIONAL_RANK_ROBUSTNESS": report.get("CROSS_SECTIONAL_RANK_ROBUSTNESS"),
                "ECONOMIC_TIMING_ROBUSTNESS": report.get("ECONOMIC_TIMING_ROBUSTNESS"),
                "EXACT_CLOCK_OVERFIT_EVIDENCE": report.get("EXACT_CLOCK_OVERFIT_EVIDENCE"),
                "COMMON_SELECTED_RATE_M5": report.get("COMMON_SELECTED_RATE_M5"),
                "COMMON_TIMING_PNL_DELTA_M5": report.get("COMMON_TIMING_PNL_DELTA_M5"),
                "REENTRY_INTERACTION_CONTRIBUTION": report.get("REENTRY_INTERACTION_CONTRIBUTION"),
                "HOLDOUT_LOO_SIGN_STABILITY": report.get("HOLDOUT_LOO_SIGN_STABILITY"),
                "TAIL_CONCENTRATION": report.get("TAIL_CONCENTRATION"),
                "TOP5_BOUNDARY_SENSITIVITY": report.get("TOP5_BOUNDARY_SENSITIVITY"),
            }
        ],
        "Common_Selected": [r for r in (report.get("iso_lineage_m5") or []) if r.get("class") == "COMMON_SELECTED"]
        + [r for r in (report.get("iso_lineage_p5") or []) if r.get("class") == "COMMON_SELECTED"]
        or [{"status": "empty"}],
        "Replacement": [r for r in (report.get("iso_lineage_m5") or []) if r.get("class") != "COMMON_SELECTED"]
        + [r for r in (report.get("iso_lineage_p5") or []) if r.get("class") != "COMMON_SELECTED"]
        or [{"status": "empty"}],
        "Fill": (report.get("iso_lineage_m5") or []) + (report.get("iso_lineage_p5") or []) or [{"status": "empty"}],
        "Exit_Transitions": [
            {"shift_key": "M5", **r} for r in (report.get("exit_transitions_m5") or [])
        ] + [
            {"shift_key": "P5", **r} for r in (report.get("exit_transitions_p5") or [])
        ] or [{"status": "empty"}],
        "Rank_Boundary": [
            {"period": p, "shift_key": sh, **d}
            for p, block in (report.get("boundary") or {}).items()
            for sh, d in (block or {}).items()
        ] or [{"status": "empty"}],
        "Reentry_Lineage": (report.get("port_lineage_ref") or [])
        + (report.get("port_lineage_m5") or [])
        + (report.get("port_lineage_p5") or [])
        or [{"status": "empty"}],
        "Holdout_Daily": daily("holdout_daily_isolated") + [
            {"source": "portfolio", **r} for r in daily("holdout_daily_portfolio")
        ],
        "Holdout_LOO": loo("holdout_loo_isolated") + [
            {"source": "portfolio", **r} for r in loo("holdout_loo_portfolio")
        ],
        "Isolated_Decomp": flat_decomp(),
        "Portfolio_Reentry": flat_port(),
    }
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(str(name)[:31])
        _sheet(ws, rows if isinstance(rows, list) else [rows])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_artifacts(report: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    jp, mp, xp = OUT / "report.json", OUT / "report.md", OUT / "audit.xlsx"
    slim = dict(report)
    for k in (
        "iso_lineage_m5",
        "iso_lineage_p5",
        "port_lineage_ref",
        "port_lineage_m5",
        "port_lineage_p5",
    ):
        slim.pop(k, None)
    jp.write_text(json.dumps(json_sanitize(slim), ensure_ascii=False, indent=2), encoding="utf-8")
    mp.write_text(write_md(report), encoding="utf-8")
    write_xlsx(report, xp)
    return {"report_json": str(jp), "report_md": str(mp), "audit_xlsx": str(xp)}
