"""Load prior robustness artifacts. No new capture scan. No Runtime."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from research.anchor_economic_sensitivity import DEVELOPMENT_DAYS, HOLDOUT_DAYS, X32_FORBIDDEN_FROM, X32_HISTORICAL_END

NATIVE = Path(__file__).resolve().parents[3]
PRIOR = NATIVE / "results" / "research" / "anchor_timing_robustness"
CACHE = PRIOR / "_work_cache"


def period_of(day: str) -> str:
    if day <= X32_HISTORICAL_END:
        return "DEVELOPMENT"
    if day >= X32_FORBIDDEN_FROM:
        return "POST_FREEZE_HOLDOUT"
    return "GAP_FREEZE_WINDOW"


def load_prior_json() -> dict[str, Any]:
    return json.loads((PRIOR / "report.json").read_text(encoding="utf-8"))


def _sheet(wb, name: str) -> list[dict[str, Any]]:
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c) if c is not None else "" for c in next(rows)]
    out = []
    for raw in rows:
        rec = {hdr[i]: raw[i] if i < len(raw) else None for i in range(len(hdr))}
        out.append(rec)
    return out


def _annotate_iso(iso: list[dict[str, Any]]) -> None:
    for r in iso:
        r["period"] = period_of(str(r.get("date") or ""))
        r["filled"] = bool(r.get("independent_filled") in (True, "True", "true", 1))


def _annotate_port(port: list[dict[str, Any]]) -> None:
    for r in port:
        r["period"] = period_of(str(r.get("date") or ""))
        r["anchor"] = str(r.get("anchor") or r.get("anchor_time") or "")


def trades_from_cache(cache_dir: Path | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    cache_dir = cache_dir or CACHE
    iso: list[dict[str, Any]] = []
    port: list[dict[str, Any]] = []
    rank_rows: list[dict[str, Any]] = []
    for fp in sorted(cache_dir.glob("*.json")):
        try:
            body = json.loads(fp.read_text(encoding="utf-8"))
        except Exception:
            continue
        if not body.get("ok"):
            continue
        primary = body.get("primary") or {}
        for t in primary.get("isolated_trades") or []:
            iso.append(dict(t))
        for k, p in (body.get("portfolios") or {}).items():
            for t in (p or {}).get("trades") or []:
                rec = dict(t)
                rec["shift_key"] = k
                port.append(rec)
        rank_rows.extend(list(primary.get("rank_rows") or []))
    _annotate_iso(iso)
    _annotate_port(port)
    return iso, port, rank_rows


def load_prior_trades() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    xlsx = PRIOR / "audit.xlsx"
    rank_rows: list[dict[str, Any]] = []
    cached = cache_complete()
    if cached.get("complete"):
        iso, port, rank_rows = trades_from_cache(CACHE)
        if iso and port:
            return iso, port, rank_rows
    if not xlsx.is_file():
        raise FileNotFoundError(
            f"missing {xlsx} and incomplete cache {CACHE}; re-run scripts/run_anchor_timing_robustness.py first"
        )
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    iso = _sheet(wb, "ISOLATED_TRADES")
    port = _sheet(wb, "PORTFOLIO_TRADES")
    if "RANK_COMPARE" in wb.sheetnames:
        # rank_rows not stored there; keep empty unless cache exists
        pass
    wb.close()
    _annotate_iso(iso)
    _annotate_port(port)
    return iso, port, rank_rows


def cache_complete(n_expected: int = 17) -> dict[str, Any]:
    days = []
    if CACHE.is_dir():
        for fp in sorted(CACHE.glob("*.json")):
            try:
                body = json.loads(fp.read_text(encoding="utf-8"))
            except Exception:
                continue
            if body.get("ok") and body.get("date"):
                days.append(str(body["date"]))
    days = sorted(set(days))
    return {
        "n": len(days),
        "days": days,
        "complete": len(days) >= n_expected,
        "cache_dir": str(CACHE),
    }


def assert_day_contract(iso: list[dict[str, Any]], port: list[dict[str, Any]]) -> dict[str, Any]:
    days = sorted({str(r.get("date")) for r in iso} | {str(r.get("date")) for r in port})
    missing_dev = [d for d in DEVELOPMENT_DAYS if d not in days]
    missing_h = [d for d in HOLDOUT_DAYS if d not in days]
    extra = [d for d in days if d not in DEVELOPMENT_DAYS and d not in HOLDOUT_DAYS]
    return {
        "days": days,
        "n": len(days),
        "development_ok": missing_dev == [],
        "holdout_ok": missing_h == [],
        "missing_development": missing_dev,
        "missing_holdout": missing_h,
        "extra_days": extra,
        "isolated_rows": len(iso),
        "portfolio_rows": len(port),
    }
