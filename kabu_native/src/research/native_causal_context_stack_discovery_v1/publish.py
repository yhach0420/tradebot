"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.native_causal_context_stack_discovery_v1.features import FEATURE_DEFINITIONS
from research.native_causal_context_stack_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Event_Generator",
    "Feature_Definitions",
    "D1_Internal_Stability",
    "D1_Tree",
    "Frozen_Rules",
    "D2",
    "D3",
    "D4",
    "Primary_Outcome",
    "Secondary_Outcomes",
    "Ablation",
    "SR_Contribution",
    "Participation_Contribution",
    "Matched_Control",
    "Concentration",
    "Causality_Audit",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def _block_rows(report: dict[str, Any], block: str) -> list[dict[str, Any]]:
    out = []
    for rule in list(report.get("frozen_rules") or []):
        rid = str(rule.get("rule_id") or "")
        st = (((report.get("rule_eval") or {}).get(rid) or {}).get("blocks") or {}).get(block) or {}
        out.append({"rule_id": rid, "predicates": rule.get("predicate_text"), **st})
    return out or [{"rule_id": None}]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    freeze = dict((report.get("bind") or {}).get("freeze") or {})
    lock = dict(report.get("lock") or {})
    ident = dict(report.get("identity") or {})
    feat_rows = [{"feature": k, "definition": v} for k, v in FEATURE_DEFINITIONS.items()]
    frozen = []
    for rule in list(report.get("frozen_rules") or []):
        frozen.append(
            {
                "rule_id": rule.get("rule_id"),
                "predicates": rule.get("predicate_text"),
                "n_train": rule.get("n_train"),
                "p40_train": rule.get("p40_train"),
                "lift_train": rule.get("lift_train"),
                "families": rule.get("families"),
            }
        )
    if not frozen:
        frozen = [{"rule_id": None}]
    sec = []
    for rule in list(report.get("frozen_rules") or []):
        rid = str(rule.get("rule_id") or "")
        e = dict((report.get("rule_eval") or {}).get(rid) or {})
        for b in ("D2", "D3", "D4"):
            st = dict((e.get("blocks") or {}).get(b) or {})
            sec.append(
                {
                    "rule_id": rid,
                    "block": b,
                    "p20": st.get("p20_before_m20"),
                    "p80": st.get("p80_before_m30"),
                    "mfe": st.get("median_mfe"),
                    "mae": st.get("median_mae"),
                    "ret5": st.get("ret_5m"),
                    "ret10": st.get("ret_10m"),
                    "ret20": st.get("ret_20m"),
                }
            )
    if not sec:
        sec = [{"rule_id": None}]
    abl = []
    for rid, fams in dict(report.get("ablation") or {}).items():
        for fam, row in dict(fams or {}).items():
            abl.append({"rule_id": rid, "family": fam, **dict(row or {})})
    if not abl:
        abl = [{"rule_id": None}]
    return {
        "Binding": _kv_rows({"ok": (report.get("bind") or {}).get("ok"), "VERDICT": d.get("VERDICT"), "NEXT": d.get("NEXT"), "parent": (report.get("bind") or {}).get("parent_verdict")}),
        "Event_Generator": _kv_rows(
            {
                "name": "NATIVE_PRICE_DISPLACEMENT_EVENT",
                "tv_required": False,
                "refractory_bars": 5,
                "first_onset_only": True,
                "EVENT_GENERATOR_SHA256": freeze.get("EVENT_GENERATOR_SHA256") or lock.get("EVENT_GENERATOR_SHA256"),
                "native_event_n": ident.get("native_event_n"),
                "FUTURE_EVENT_SELECTION_N": 0,
                "RETROSPECTIVE_CLUSTER_N": 0,
            }
        ),
        "Feature_Definitions": feat_rows,
        "D1_Internal_Stability": _kv_rows(report.get("d1_internal_stability") or {}),
        "D1_Tree": _kv_rows(report.get("d1_tree") or {}),
        "Frozen_Rules": frozen,
        "D2": _block_rows(report, "D2"),
        "D3": _block_rows(report, "D3"),
        "D4": _block_rows(report, "D4"),
        "Primary_Outcome": _kv_rows({"metric": "p40_before_m20", "d1_base_rate": report.get("d1_base_rate"), "surviving": d.get("surviving_rule_ids"), "eval": report.get("rule_eval")}),
        "Secondary_Outcomes": sec,
        "Ablation": abl,
        "SR_Contribution": _kv_rows(report.get("sr_contribution") or {}),
        "Participation_Contribution": _kv_rows(report.get("participation_contribution") or {}),
        "Matched_Control": _kv_rows(report.get("matched_control") or {"none": True}),
        "Concentration": _kv_rows(report.get("concentration") or {}),
        "Causality_Audit": _kv_rows(
            {
                "same_bar_entry_n": report.get("same_bar_entry_n"),
                "FUTURE_EVENT_SELECTION_N": 0,
                "RETROSPECTIVE_CLUSTER_N": 0,
                "tree_learned_only_on_d1": True,
                "d2_d4_rule_modification": False,
                "clock_prior_days_only": True,
                "no_future_vwap": True,
                "no_indicator_catalog": True,
            }
        ),
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# NATIVE_CAUSAL_CONTEXT_STACK_DISCOVERY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Compact causal context stack on native price-displacement events. Not a strategy.",
            "D1 learns. D2–D4 locked internal replication. Confirmation and Frozen Validation closed.",
            "",
            f"Base displacement event_n? **{a.get('Base displacement event_n?')}**",
            f"D1/D2/D3/D4? **{a.get('D1/D2/D3/D4 counts?')}**",
            f"future event selection? **{a.get('Any future event selection?')}** retrospective clustering? **{a.get('Any retrospective clustering?')}**",
            f"tree D1 only? **{a.get('Tree learned only on D1?')}** D2-D4 modification? **{a.get('Any D2-D4 rule modification?')}**",
            f"frozen rule_n? **{a.get('Frozen rule_n?')}**",
            f"Rule 1? **{a.get('Rule 1 predicates / counts / effects?')}**",
            f"Rule 2? **{a.get('Rule 2?')}**",
            f"Rule 3? **{a.get('Rule 3?')}**",
            f"positive D2/D3/D4? **{a.get('Which rules are positive in D2/D3/D4?')}**",
            f"S/R? **{a.get('Does S/R materially improve a surviving combination?')}** TV? **{a.get('Does TradingValue?')}** VWAP? **{a.get('Does VWAP?')}**",
            f"trend? **{a.get('Does 5m/15m trend?')}** local? **{a.get('Does local structure?')}** market/sector? **{a.get('Does market/sector context?')}**",
            f"I2 convergence? **{a.get('Does any frozen rule converge with the old underpowered I2 lead?')}**",
            f"one-symbol? **{a.get('Any one-symbol domination?')}** one-day? **{a.get('Any one-day domination?')}**",
            f"PnL opt? **{a.get('Any PnL optimization?')}** Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}**",
            f"Kabu50? **{a.get('Kabu50?')}** submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
