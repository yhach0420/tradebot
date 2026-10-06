"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.sr_a2_bracketed_structure_complete_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "PostHoc_Status",
    "Eligibility",
    "Eligibility_Timestamp_Audit",
    "Target_Identity",
    "Target_Availability",
    "No_Target_Skips",
    "Confound_Audit",
    "Adjusted_Diagnostic",
    "Path_Mechanism",
    "Entry_Lineage",
    "Exit_Lineage",
    "Trades",
    "CAP_Blocked",
    "Occupancy",
    "Tie_Order",
    "D1_D2_D3_D4",
    "Bps_Economics",
    "Yen_Economics",
    "Notional",
    "Target_Distance",
    "Exit_Attribution",
    "Session_Close",
    "MFE_MAE",
    "Giveback",
    "Concentration",
    "Tail_Robustness",
    "Freeze",
    "Safety",
)
STRIP = {"_markdown", "_rows"}
TRADE_COLS = (
    "signal_id",
    "symbol",
    "date",
    "block",
    "side",
    "decision_bar",
    "decision_px",
    "entry_t",
    "entry_price",
    "target_zone_id",
    "target_price",
    "exit_reason",
    "exit_t",
    "exit_price",
    "gross_yen",
    "stress_yen",
    "realized_bps",
    "mfe_bps",
    "hold_min",
    "attribution",
)


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _rows(items: list[dict[str, Any]], cols: tuple[str, ...]) -> list[dict[str, Any]]:
    return [{c: r.get(c) for c in cols} for r in items]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    rows = list(report.get("_rows") or [])
    trades = list((report.get("PRIMARY") or {}).get("_trades") or [])
    e = dict((report.get("PRIMARY") or {}).get("economics") or {})
    d = dict(report.get("decision") or {})
    skips = [r for r in rows if not r.get("opposing_zone_available")]
    elig = [r for r in rows if r.get("opposing_zone_available")]
    d14 = [{"block": b, **dict((e.get("blocks") or {}).get(b) or {})} for b in ("D1", "D2", "D3", "D4")]
    conc = dict(e.get("concentration") or {})
    tail = []
    for k in ("top1_trade_removed", "top5_trades_removed", "top5pct_winners_removed", "best_day_removed", "top_symbol_removed"):
        tail.append({"cut": k, **dict(conc.get(k) or {})})
    ties = [
        {"order": "PRIMARY", **{k: v for k, v in dict((report.get("PRIMARY") or {}).get("economics") or {}).items() if k in {"stress_yen", "pf_stress", "trades", "mean_stress_bps"}}},
        {"order": "REVERSE", **{k: v for k, v in dict((report.get("REVERSE") or {}).get("economics") or {}).items() if k in {"stress_yen", "pf_stress", "trades", "mean_stress_bps"}}},
        {"order": "HASH", **{k: v for k, v in dict((report.get("HASH") or {}).get("economics") or {}).items() if k in {"stress_yen", "pf_stress", "trades", "mean_stress_bps"}}},
    ]
    return {
        "Binding": _kv_rows({"parent": report.get("parent_verdict"), "bind": (report.get("bind") or {}).get("ok"), "VERDICT": d.get("VERDICT"), "NEXT": d.get("NEXT")}),
        "PostHoc_Status": _kv_rows({"POST_HOC_DEVELOPMENT_ARCHITECTURE": True, "D1_D4": "ALL_DEVELOPMENT", "not_validation": True, "c1_closed": True}),
        "Eligibility": _kv_rows({"eligible_n": report.get("eligible_n"), "skip_n": report.get("skip_n"), "a2_signal_n": (report.get("counts") or {}).get("a2_signal_n"), "price_kind": "decision_bar_close"}),
        "Eligibility_Timestamp_Audit": _kv_rows(
            {
                "TARGET_ELIGIBILITY_USES_ENTRY_OPEN_N": report.get("TARGET_ELIGIBILITY_USES_ENTRY_OPEN_N"),
                "TARGET_ELIGIBILITY_FUTURE_BAR_N": report.get("TARGET_ELIGIBILITY_FUTURE_BAR_N"),
                "TARGET_FUTURE_LEAKAGE_N": report.get("target_future_leakage_n"),
                "disagrees_with_entry_open_n": report.get("eligibility_disagrees_with_entry_open_n"),
                "SAME_BAR_ENTRY_N": report.get("same_bar_entry_n"),
            }
        ),
        "Target_Identity": _rows(elig, ("signal_id", "target_zone_id", "target_price", "target_activated_at", "target_distance_bps")),
        "Target_Availability": _kv_rows({"available_n": len(elig), "skip_n": len(skips), "target_exit_rate": e.get("target_exit_rate"), "target_reached_n": e.get("target_reached_n")}),
        "No_Target_Skips": _rows(skips, ("signal_id", "symbol", "date", "side", "decision_px", "skip_reason")),
        "Confound_Audit": _kv_rows(report.get("confound") or {}),
        "Adjusted_Diagnostic": _kv_rows(report.get("adjusted_diagnostic") or {}),
        "Path_Mechanism": _kv_rows(report.get("path_mechanism") or {}),
        "Entry_Lineage": _rows(elig, ("signal_id", "decision_bar", "decision_available_at", "decision_px", "entry_t", "entry_price", "used_entry_open_for_eligibility", "same_bar_entry")),
        "Exit_Lineage": _rows(trades, ("signal_id", "entry_thesis", "entry_zone_id", "target_zone_id", "exit_reason", "exit_decision_time", "exit_execution_time", "exit_retroactive")),
        "Trades": _rows(trades, TRADE_COLS),
        "CAP_Blocked": _rows(list((report.get("PRIMARY") or {}).get("cap_blocked") or []), ("signal_id", "symbol", "date", "entry_t", "block_reason")),
        "Occupancy": _kv_rows({k: e.get(k) for k in ("occupancy_mean_at_entry", "slot_utilization", "reentry_n", "cap_blocked_n", "same_symbol_blocked_n")}),
        "Tie_Order": ties,
        "D1_D2_D3_D4": d14,
        "Bps_Economics": _kv_rows({k: e.get(k) for k in ("mean_gross_bps", "median_gross_bps", "mean_stress_bps", "median_stress_bps", "pf_stress")}),
        "Yen_Economics": _kv_rows({k: e.get(k) for k in ("gross_yen", "stress_yen", "pf_gross", "pf_stress", "median_trade", "maxDD_yen", "positive_day_rate", "break_even_cost_bps")}),
        "Notional": _kv_rows({"shares": 100, "mean_notional_available": (report.get("confound") or {}).get("mean_notional_available"), "mean_notional_skip": (report.get("confound") or {}).get("mean_notional_skip"), "notional_confound": (report.get("confound") or {}).get("notional_confound")}),
        "Target_Distance": _kv_rows({"mean_bps": e.get("mean_target_distance_bps"), "median_bps": e.get("median_target_distance_bps")}),
        "Exit_Attribution": _kv_rows(e.get("attribution") or {}),
        "Session_Close": _kv_rows({"session_close_n": e.get("session_close_n"), "invalidation_n": e.get("invalidation_exit_n"), "target_n": e.get("target_exit_n"), "horizon_mismatch": d.get("STRUCTURAL_EXIT_HORIZON_MISMATCH")}),
        "MFE_MAE": _kv_rows({"mean_mfe": e.get("mean_mfe_bps"), "median_mfe": e.get("median_mfe_bps"), "mae_mean": e.get("mae_mean")}),
        "Giveback": _kv_rows({"mean": e.get("mean_giveback_bps"), "median": e.get("median_giveback_bps")}),
        "Concentration": _kv_rows({k: conc.get(k) for k in ("top1_yen", "top1_share", "top3_yen", "top5_yen", "top10_yen", "top1pct_yen", "top5pct_yen", "top_symbol", "best_day")}),
        "Tail_Robustness": tail,
        "Freeze": _kv_rows(report.get("freeze_candidate") or {}),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# SR_A2_BRACKETED_STRUCTURE_COMPLETE_STRATEGY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"classification: **{d.get('classification')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "POST_HOC_DEVELOPMENT_ARCHITECTURE = true. D1–D4 are ALL DEVELOPMENT. C1 closed.",
            "",
            f"eligibility before ENTRY? **{a.get('Was opposing-zone eligibility fully knowable before ENTRY?')}** USES_ENTRY_OPEN_N? **{a.get('TARGET_ELIGIBILITY_USES_ENTRY_OPEN_N?')}**",
            f"eligible? **{a.get('eligible signal_n?')}** skip? **{a.get('NO_OPPOSING_ZONE_SKIP_n?')}** trades? **{a.get('trade_n?')}** CAP blocked? **{a.get('CAP blocked?')}**",
            f"D2? **{a.get('D2 stress / PF?')}** D3? **{a.get('D3 stress / PF?')}** D4? **{a.get('D4 stress / PF?')}**",
            f"overall stress/PF? **{a.get('overall stress / PF?')}** bps? **{a.get('gross/stress bps?')}**",
            f"advantage in bps? **{a.get('Does advantage remain in bps?')}** notional confound? **{a.get('Any notional confound?')}** proxy? **{a.get('Any market/regime proxy evidence?')}**",
            f"target reached/rate? **{a.get('Target reached_n?')}** / **{a.get('Target exit rate?')}** acts as? **{a.get('Does target act as actual EXIT or only structural context?')}**",
            f"invalidation? **{a.get('Invalidation exits?')}** session-close? **{a.get('Session-close exits?')}** hold? **{a.get('Mean/median hold?')}**",
            f"tie fragile? **{a.get('CAP tie-order fragile?')}** C1 reopened? **{a.get('C1 reopened?')}** PnL rule? **{a.get('Any PnL-selected additional rule?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}** Kabu50? **{a.get('Kabu50 applied?')}** submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    if isinstance(payload.get("PRIMARY"), dict):
        payload["PRIMARY"] = {k: v for k, v in payload["PRIMARY"].items() if k != "_trades"}
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
