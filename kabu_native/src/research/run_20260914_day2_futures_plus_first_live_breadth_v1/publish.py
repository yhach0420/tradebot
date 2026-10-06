"""report.json / report.md / audit.xlsx. Minimize files. No many CSVs."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_x_stock_state_day2_confirmation_v1.publish import write_artifacts as write_day2
from research.run_20260914_day2_futures_plus_first_live_breadth_v1 import ANALYSIS_ID
from research.run_20260914_day2_futures_plus_first_live_breadth_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = ("Answers", "Transport", "Q1", "Q2", "Clocks", "Day2", "Decision", "Safety")


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    if isinstance(x, (dict, list, tuple)):
        return json.dumps(x, ensure_ascii=False, default=str)
    return str(x)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    op = dict(body.get("operator_monday") or {})
    t = dict(body.get("transport") or {})
    recon = dict(body.get("recon") or {})
    q1 = dict(recon.get("q1") or {})
    q2 = dict(recon.get("q2") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"clock: `{body.get('clock')}` today=`{body.get('today_jst')}` target=`{body.get('trading_date')}`",
        "",
        "20260911 existing-raw mining remains CLOSED.",
        "This file is the combined 20260914 EOD. ENTRY/EXIT stay false even if A/B.",
        "Day2 PASS means only the frozen cross-sectional interaction replicated.",
        "",
        "## Operator Monday sequence",
        "",
        f"- prepare: `{op.get('prepare')}`",
        f"- 07:55 futures: `{op.get('futures_live')}`",
        f"- 09:05 breadth: `{op.get('breadth_live')}`",
        f"- after 11:30 EOD: `{op.get('eod')}`",
        f"- types `{op.get('types')}` cadence `{op.get('cadence_sec')}` fail-soft `{op.get('fail_soft_sec')}`",
        "",
        "## A. Futures Day2",
        "",
        f"- 1 FULL?: `{_fmt(a.get('1_FULL'))}`",
        f"- 2 stock N?: `{a.get('2_stock_N')}`",
        f"- 3 NK present?: `{_fmt(a.get('3_NK_present'))}`",
        f"- 4 TOPIX present?: `{_fmt(a.get('4_TOPIX_present'))}`",
        f"- 5 Day2 BOTH_DOWN clock N?: `{a.get('5_Day2_BOTH_DOWN_clock_N')}`",
        f"- 6 TOP MID?: `{_fmt(a.get('6_TOP_MID'))}`",
        f"- 7 BOTTOM MID?: `{_fmt(a.get('7_BOTTOM_MID'))}`",
        f"- 8 TOP-BOTTOM?: `{_fmt(a.get('8_TOP_BOTTOM'))}`",
        f"- 9 BASE A?: `{_fmt(a.get('9_BASE_A'))}`",
        f"- 10 incremental lift?: `{_fmt(a.get('10_incremental_lift'))}`",
        f"- 11 TOP LONG mean?: `{_fmt(a.get('11_TOP_LONG_mean'))}`",
        f"- 12 TOP LONG median?: `{_fmt(a.get('12_TOP_LONG_median'))}`",
        f"- 13 primary PASS/FAIL?: `{a.get('13_primary_PASS_FAIL')}`",
        f"- 14 secondary rescue used?: `{_fmt(a.get('14_secondary_rescue_used'))}`",
        "",
        "## B. Breadth",
        "",
        f"- 15 transport FULL?: `{_fmt(a.get('15_transport_FULL'))}`",
        f"- 16 seven types present?: `{_fmt(a.get('16_seven_types_present'))}`",
        f"- 17 snapshot N by type?: `{_fmt(a.get('17_snapshot_N_by_type'))}`",
        f"- 18 empty rate?: `{_fmt(a.get('18_empty_rate'))}`",
        f"- 19 429 N?: `{a.get('19_429_N')}`",
        f"- 20 HTTP error N?: `{a.get('20_HTTP_error_N')}`",
        f"- 21 schema drift N?: `{a.get('21_schema_drift_N')}`",
        f"- 22 received_at continuity?: `{_fmt(a.get('22_received_at_continuity'))}`",
        f"- 23 Q1 incremental breadth evidence?: `{_fmt(a.get('23_Q1_incremental_breadth_evidence'))}`",
        f"- 24 Q2 activity-effect change?: `{_fmt(a.get('24_Q2_activity_effect_change'))}`",
        f"- 25 strongest explainable observation?: {a.get('25_strongest_explainable_observation')}",
        f"- 26 ENTRY built?: `{_fmt(a.get('26_ENTRY_built'))}`",
        f"- 27 EXIT built?: `{_fmt(a.get('27_EXIT_built'))}`",
        f"- 28 Runtime changed?: `{_fmt(a.get('28_Runtime_changed'))}`",
        f"- 29 Paper changed?: `{_fmt(a.get('29_Paper_changed'))}`",
        f"- 30 submit/cancel/live?: `{a.get('30_submit_cancel_live')}`",
        f"- 31 VERDICT?: `{a.get('31_VERDICT')}`",
        f"- 32 NEXT?: `{a.get('32_NEXT')}`",
        "",
        "## Transport detail",
        "",
        f"- collector clean exit: `{_fmt(t.get('collector_clean_exit'))}` pid alive `{_fmt(t.get('collector_pid_alive'))}`",
        f"- first `{t.get('first_received_at')}` last `{t.get('last_received_at')}`",
        "",
        "## Q1 / Q2",
        "",
        f"- recon ran: `{_fmt(recon.get('ran'))}` reason `{recon.get('reason')}`",
        f"- Q1: {q1.get('reason')}",
        f"- Q2: {q2.get('reason')}",
        "",
        "No large Futures × Breadth × Stock × horizon grid. Day1 outcomes were not used to tune thresholds.",
        "",
    ]
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any], *, day2_body: dict[str, Any] | None = None) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    day2_body = day2_body if day2_body is not None else body.get("_day2_full")
    slim = dict(body)
    slim.pop("_day2_full", None)
    payload = json_sanitize(slim)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(body), encoding="utf-8")
    recon = dict(payload.get("recon") or {})
    q1 = dict(recon.get("q1") or {})
    q2 = dict(recon.get("q2") or {})
    transport = dict(payload.get("transport") or {})
    clocks = []
    for r in recon.get("clock_rows") or []:
        clocks.append(
            {
                "clock": r.get("clock"),
                "agreement_180": r.get("agreement_180"),
                "LEADERSHIP_DIRECTION": r.get("LEADERSHIP_DIRECTION"),
                "EW_MID_10m": r.get("EW_MID_10m"),
                "EW_LONG_10m": r.get("EW_LONG_10m"),
                "TOP_BOTTOM_MID_10m": r.get("TOP_BOTTOM_MID_10m"),
                "TOP_BOTTOM_LONG_10m": r.get("TOP_BOTTOM_LONG_10m"),
                "TICK_PRESSURE": r.get("TICK_PRESSURE"),
                "SECTOR_DIRECTION": r.get("SECTOR_DIRECTION"),
                "LEADERSHIP_TURNOVER": r.get("LEADERSHIP_TURNOVER"),
            }
        )
    sheets = {
        "Answers": _kv(dict(payload.get("answers") or {})),
        "Transport": _kv({k: v for k, v in transport.items() if k != "by_type"}),
        "Q1": _kv({k: v for k, v in q1.items() if k != "horizons"}),
        "Q2": _kv(q2),
        "Clocks": clocks,
        "Day2": _kv(dict((payload.get("day2") or {}).get("answers") or {})),
        "Decision": _kv(dict(payload.get("decision") or {})),
        "Safety": [
            {"key": "ENTRY", "value": False},
            {"key": "EXIT", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "secondary_rescue_used", "value": False},
            {"key": "day1_mining_reopened", "value": False},
            {"key": "register_mutation_n", "value": (payload.get("mutations") or {}).get("register_mutation_n")},
            {"key": "unregister_n", "value": (payload.get("mutations") or {}).get("unregister_n")},
            {"key": "sendorder_n", "value": (payload.get("mutations") or {}).get("sendorder_n")},
            {"key": "submit_cancel_live", "value": payload.get("submit_cancel_live")},
        ],
    }
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    out = {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(xlsx),
    }
    if day2_body is not None:
        out.update({f"day2_{k}": v for k, v in write_day2(day2_body).items()})
    return out
