"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v3_2_opening_drive_and_reacceleration_semantics.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Rule_Diff",
    "Funnel",
    "Deaths",
    "Dropped_CLEAR",
    "Weak_Reclaim",
    "Confusion",
    "Unseen",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "events", "failed_push_archive", "funnel_days", "chart_zones"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    audit = dict(report.get("audit_old67") or {})
    conf = dict(audit.get("confusion") or {})
    weak = dict(audit.get("weak_reclaim_cause_counts") or {})
    unseen = dict(report.get("unseen") or {})
    dec = dict(report.get("decision") or {})
    drive_def = (
        "CLEAN_OPENING_DRIVE_V3: 09:14 close in directional half of frozen OR15; "
        "reject TWO_SIDED_OPEN (mid-band close with both-side excursions) and RANGE_OPEN; "
        "allow INITIAL_DIRECTION_CONTINUATION and EARLY_REVERSAL_THEN_DOMINANT_DRIVE; "
        "no session-open 0.50 net; no first-5m same-sign requirement; efficiency persisted not gated."
    )
    reaccel_def = (
        "REACCELERATION_TRIGGER_V2: after defended retest, close through retest micro structure "
        "AND directional close half of trigger bar AND body >= opposite wick "
        "AND trigger range >= retest bar range. reclaim_move/NORMAL_1M_RANGE persisted, not gated at 1.0."
    )
    return {
        "v31_unchanged": report.get("v31_unchanged"),
        "V3_2_MACHINE_SHA256": report.get("MACHINE_SHA256"),
        "any_future_outcome_used": False,
        "old67_semantic_development_only": True,
        "dropped_clear_n": audit.get("dropped_clear_n"),
        "dropped_clear_session_open_mismatch_n": audit.get("dropped_clear_session_open_mismatch_n"),
        "dropped_clear_other_n": audit.get("dropped_clear_other_n"),
        "opening_representation_best": audit.get("opening_representation_best"),
        "CLEAN_OPENING_DRIVE_V3": drive_def,
        "weak_tiny_overshoot": int(weak.get("tiny_overshoot") or 0),
        "weak_body": int(weak.get("weak_body") or 0),
        "weak_bad_close_location": int(weak.get("bad_close_location") or 0),
        "weak_no_expansion": int(weak.get("no_expansion") or 0),
        "weak_drift_crawl": int(weak.get("drift_crawl") or 0),
        "weak_late_stalled": int(weak.get("late_stalled") or 0),
        "weak_large_opposite_wick": int(weak.get("large_opposite_wick") or 0),
        "weak_other": int(weak.get("other") or 0),
        "REACCELERATION_TRIGGER_V2": reaccel_def,
        "reclaim_threshold_from_profit": False,
        "retest_5min_gate": False,
        "clock_0930_cutoff": False,
        "CLEAR_retained": conf.get("CLEAR_retained"),
        "CLEAR_lost": conf.get("CLEAR_lost"),
        "QUESTIONABLE_retained": conf.get("QUESTIONABLE_retained"),
        "QUESTIONABLE_removed": conf.get("QUESTIONABLE_removed"),
        "NOT_retained": conf.get("NOT_retained"),
        "NOT_removed": conf.get("NOT_removed"),
        "FACE_VALID_from_development_set": False,
        "post_20260911_unseen_candidate_n": unseen.get("post_20260911_historical_candidate_n"),
        "prospective_capture_unseen_candidate_n": unseen.get("prospective_capture_date_n"),
        "independent_face_sample_n": unseen.get("independent_face_sample_n"),
        "any_economic_path_test": False,
        "any_pnl": False,
        "old_confirmation_opened": False,
        "frozen_validation_opened": False,
        "submit_cancel_live": "0/0/0",
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "V3_2_event_n": report.get("setup_n"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    lines = [
        "# PB1_V3_2_OPENING_DRIVE_AND_REACCELERATION_SEMANTICS_V1",
        "",
        "Semantic development only. Frozen V3.1/V3/V2 unchanged. No PnL. FACE_VALID unknown.",
        "",
        f"V3.1 unchanged? **{a.get('v31_unchanged')}**",
        f"V3.2 SHA? `{a.get('V3_2_MACHINE_SHA256')}`",
        f"Any future outcome used? **{a.get('any_future_outcome_used')}**",
        f"Old67 used as semantic development only? **{a.get('old67_semantic_development_only')}**",
        "",
        f"Why were 9 prior CLEAR events dropped by V3.1? session-open/first-5m mismatch n=**{a.get('dropped_clear_session_open_mismatch_n')}**; OTHER n=**{a.get('dropped_clear_other_n')}** (cleared-zone MEANINGFUL_DEFENSE close-beyond, kept unchanged).",
        f"SESSION_OPEN anchor mismatch? **{a.get('dropped_clear_session_open_mismatch_n')}**",
        f"OTHER? **{a.get('dropped_clear_other_n')}**",
        f"What opening representation best matches human semantics? **{a.get('opening_representation_best')}**",
        "",
        f"CLEAN_OPENING_DRIVE_V3? {a.get('CLEAN_OPENING_DRIVE_V3')}",
        "",
        "WEAK_RECLAIM cause counts:",
        f"tiny overshoot? **{a.get('weak_tiny_overshoot')}**",
        f"weak body? **{a.get('weak_body')}**",
        f"bad close location? **{a.get('weak_bad_close_location')}**",
        f"no expansion? **{a.get('weak_no_expansion')}**",
        f"drift/crawl? **{a.get('weak_drift_crawl')}**",
        f"late/stalled? **{a.get('weak_late_stalled')}**",
        f"large opposite wick? **{a.get('weak_large_opposite_wick')}**",
        f"other? **{a.get('weak_other')}**",
        "",
        f"REACCELERATION_TRIGGER_V2? {a.get('REACCELERATION_TRIGGER_V2')}",
        f"Any reclaim threshold chosen from profit? **{a.get('reclaim_threshold_from_profit')}**",
        f"Any fixed 5m rule? **{a.get('retest_5min_gate')}**",
        f"Any hard 09:30 cutoff? **{a.get('clock_0930_cutoff')}**",
        "",
        f"CLEAR retained / lost? **{a.get('CLEAR_retained')}** / **{a.get('CLEAR_lost')}**",
        f"QUESTIONABLE retained / removed? **{a.get('QUESTIONABLE_retained')}** / **{a.get('QUESTIONABLE_removed')}**",
        f"NOT retained / removed? **{a.get('NOT_retained')}** / **{a.get('NOT_removed')}**",
        f"FACE_VALID from development set? **{a.get('FACE_VALID_from_development_set')}**",
        "",
        f"Post-20260911 unseen candidate_n? **{a.get('post_20260911_unseen_candidate_n')}**",
        f"Prospective-capture unseen candidate_n? **{a.get('prospective_capture_unseen_candidate_n')}**",
        f"Independent face sample_n? **{a.get('independent_face_sample_n')}**",
        "",
        f"Any economic path test? **{a.get('any_economic_path_test')}**",
        f"Any PnL? **{a.get('any_pnl')}**",
        f"Old Confirmation opened? **{a.get('old_confirmation_opened')}**",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"submit/cancel/live? **{a.get('submit_cancel_live')}**",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    audit = dict(report.get("audit_old67") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Rule_Diff": _kv_rows(report.get("rule_diff")),
        "Funnel": _kv_rows(report.get("funnel")),
        "Deaths": _kv_rows(report.get("deaths")),
        "Dropped_CLEAR": list(audit.get("dropped_clear_v31") or [_kv_rows({"empty": True})[0]]),
        "Weak_Reclaim": _kv_rows(audit.get("weak_reclaim_cause_counts") or {}),
        "Confusion": _kv_rows(audit.get("confusion") or {}),
        "Unseen": _kv_rows(report.get("unseen")),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
