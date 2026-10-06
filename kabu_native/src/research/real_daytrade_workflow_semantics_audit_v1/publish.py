"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.real_daytrade_workflow_semantics_audit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Interpretation",
    "Station_Config",
    "Daily_Convention",
    "Intraday_Convention",
    "One_Minute_Role",
    "Stock_Selection",
    "Indicator_Roles",
    "Face_Valid_Sample",
    "Human_Classification",
    "Prior_Machines",
    "Validation_Framework",
    "Matching_Audit",
    "Absolute_Gate",
    "Execution_Limits",
    "Canonical_Workflow",
    "Playbook_Candidates",
    "Causality_Audit",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    conv = dict(report.get("convention") or {})
    sample_rows = list(report.get("sample_rows") or [])
    slim = []
    for r in sample_rows:
        slim.append(
            {
                "sample_id": r.get("sample_id"),
                "family": r.get("family"),
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "t": r.get("t"),
                "direction": r.get("direction"),
                "daily_stack": r.get("daily_stack"),
                "human_class": r.get("human_class"),
                "human_name": r.get("human_name"),
                "human_reviewed": r.get("human_reviewed"),
                "chart": r.get("chart"),
            }
        )
    return {
        "Binding": _kv_rows({"ok": (report.get("bind") or {}).get("ok"), "parent_verdict": (report.get("bind") or {}).get("parent_verdict"), "VERDICT": d.get("VERDICT"), "NEXT": d.get("NEXT")}),
        "Interpretation": _kv_rows(report.get("interpretation") or {}),
        "Station_Config": _kv_rows(report.get("station") or {}),
        "Daily_Convention": _kv_rows(conv.get("daily") or {}),
        "Intraday_Convention": _kv_rows({"five_minute": conv.get("five_minute"), "fifteen_minute": conv.get("fifteen_minute")}),
        "One_Minute_Role": _kv_rows(conv.get("one_minute") or {}),
        "Stock_Selection": _kv_rows(report.get("availability") or {}),
        "Indicator_Roles": list(report.get("roles") or [{"item": None}]),
        "Face_Valid_Sample": slim or [{"sample_id": None}],
        "Human_Classification": _kv_rows(report.get("sample_summary") or {}),
        "Prior_Machines": list(report.get("prior_machines") or [{"machine": None}]),
        "Validation_Framework": _kv_rows(report.get("validation_framework") or {}),
        "Matching_Audit": _kv_rows(report.get("matching") or {}),
        "Absolute_Gate": _kv_rows(report.get("absolute_gate") or {}),
        "Execution_Limits": _kv_rows(report.get("execution_limits") or {}),
        "Canonical_Workflow": list((report.get("workflow") or {}).get("stack") or [{"step": None}]),
        "Playbook_Candidates": list(report.get("playbook_candidates") or [{"id": None}]),
        "Causality_Audit": _kv_rows(
            {
                "future_hidden_on_charts": True,
                "same_bar_entry_n": report.get("same_bar_entry_n"),
                "new_strategy_run": False,
                "counts_not_rerun": True,
                "sample": report.get("sample"),
            }
        ),
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# REAL_DAYTRADE_WORKFLOW_SEMANTICS_AUDIT_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Audit of whether research machines represent ordinary Japanese equity daytrade chart use.",
            "Not a strategy. Not an SMA rescue. Not a PnL test.",
            "",
            f"TESTED_MACHINE_FAILED? **{a.get('TESTED_MACHINE_FAILED')}**",
            f"STANDARD_MA_FAMILY_EXHAUSTED? **{a.get('STANDARD_MA_FAMILY_EXHAUSTED')}**",
            f"Correct timeframe for SMA5/25/75? **{a.get('Were we using SMA5/25/75 on the correct timeframe?')}**",
            f"Confused daily convention with intraday bars? **{a.get('Were we confusing daily convention with intraday bar periods?')}**",
            f"5m SMA25-pullback face-valid? **{a.get('Was the previous 5m SMA25-pullback machine face-valid?')}**",
            f"Sampled events a trader would call intended? **{a.get('What percentage of sampled events would a human trader actually call the intended setup?')}**",
            f"Matching over-conditioned? **{a.get('Were matching controls over-conditioning on variables that are part of the setup?')}**",
            f"Next playbooks? **{a.get('At most 3 next playbooks?')}**",
            f"No new strategy run? **{a.get('No new strategy run?')}** Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}**",
            f"submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
