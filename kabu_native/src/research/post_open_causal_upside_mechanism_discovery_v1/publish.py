"""Write report.json / report.md / audit.xlsx only. No CSV dump."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.post_open_causal_upside_mechanism_discovery_v1 import ANALYSIS_ID
from research.post_open_causal_upside_mechanism_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Anchor_Population",
    "Feature_Integrity",
    "Outcome_Distribution",
    "Univariate",
    "Quintiles",
    "Block_Stability",
    "TopSymbol_Exclusion",
    "LeaveOneDay",
    "Horizon_Robustness",
    "Tree",
    "Tree_Leaves",
    "LOBO",
    "Interactions",
    "Mechanism",
    "Next_Strategy",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    uni = list(report.get("univariate") or [])
    tree = dict(report.get("tree") or {})
    interp = dict(report.get("interpretation") or {})
    quint = []
    for u in uni:
        for q in list(u.get("quintiles") or []):
            quint.append({"feature": u.get("feature"), **q})
    return {
        "Summary": _kv(dict(report.get("answers") or {})),
        "Anchor_Population": _kv(dict(report.get("population") or {})),
        "Feature_Integrity": list(report.get("feature_integrity") or []) or [{"empty": True}],
        "Outcome_Distribution": _kv(dict(report.get("population") or {})),
        "Univariate": [
            {
                "feature": u.get("feature"),
                "spearman_markout": u.get("spearman_markout"),
                "spearman_path_edge": u.get("spearman_path_edge"),
                "candidate": u.get("mechanism_candidate"),
                "agree_m": u.get("block_agree_markout"),
                "agree_p": u.get("block_agree_path"),
            }
            for u in uni
        ]
        or [{"empty": True}],
        "Quintiles": quint or [{"empty": True}],
        "Block_Stability": [
            {"feature": u.get("feature"), **dict(u.get("block_markout_signs") or {})} for u in uni
        ]
        or [{"empty": True}],
        "TopSymbol_Exclusion": [
            {"feature": u.get("feature"), "holds": u.get("top_symbol_excluded_holds")} for u in uni
        ],
        "LeaveOneDay": [{"feature": u.get("feature"), "holds": u.get("leave_one_day_holds")} for u in uni],
        "Horizon_Robustness": [{"feature": u.get("feature"), "holds": u.get("horizon_5_15_holds")} for u in uni],
        "Tree": _kv({k: v for k, v in tree.items() if k != "leaves"}),
        "Tree_Leaves": list(tree.get("leaves") or []) or [{"empty": True}],
        "LOBO": list((report.get("lobo") or {}).get("folds") or []) or _kv(dict(report.get("lobo") or {})),
        "Interactions": list(report.get("interactions") or []) or [{"empty": True}],
        "Mechanism": _kv(interp),
        "Next_Strategy": _kv(
            {
                "PRECOMMIT_THIS_RUN": False,
                "ENTRY_THESIS": interp.get("ENTRY_THESIS"),
                "TECHNICAL_EXIT": interp.get("TECHNICAL_EXIT"),
                "NEXT": (report.get("decision") or {}).get("NEXT"),
            }
        ),
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        "TRUE_OOS: false",
        "CERTIFIED: false",
        "KIND: MECHANISM_DISCOVERY_ONLY",
        "CANDIDATE_STRATEGY_N: 0",
        "",
        str((report.get("interpretation") or {}).get("MARKET_MECHANISM") or ""),
        "",
    ]
    for i in range(1, 42):
        matches = [k for k in a if k.split("_", 1)[0] == str(i)]
        if matches:
            lines.append(f"{i}. {matches[0]}: {a.get(matches[0])}")
    lines.extend(["", "No Full Causal economics. No V5 rescue. No 10m EXIT freeze.", "", "STOP.", ""])
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
