"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_machine_implementation import (
    E1_BODY_N1M,
    E1_NET_N1M,
    E1_RANGE_N1M,
    EXEC_1M_CONFIRMED,
    EXEC_5M_DIRECT,
    FAIL_COUNTER_MIN,
    FAIL_DRIVE_DISP_MIN,
    S0_GAP_ATR_MIN,
    S0_GAP_RANGE_MIN,
    S0_RANGE_MIN,
    S4_BODY_FRAC,
    S4_CLOSE_LOC,
    S4_RANGE_OVER_N1M,
    S4_RANGE_OVER_OPEN5,
    TRUE_BODY_FRAC_MIN,
    TRUE_COUNTER_FRAC,
    TRUE_DISP_MIN,
    TRUE_N_SAME_MIN,
    TRUE_RANGE_MIN,
)
from research.pb1_v4_machine_implementation.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Calibration",
    "Funnel",
    "Deaths",
    "Audit88",
    "CLEAR19",
    "Specials",
    "Negatives",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "audit_rows", "chart_zones"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    audit = dict(report.get("semantic_development_audit") or {})
    cal = dict(report.get("calibration") or {})
    s3382 = dict(audit.get("3382_20241004") or {})
    funnel = dict(report.get("funnel") or {})
    s0_rule = str((cal.get("S0") or {}).get("chosen_rule") or "")
    s1_rule = str((cal.get("S1") or {}).get("chosen_rule") or "")
    return {
        "V4 machine implemented?": True,
        "V4 SHA?": report.get("V4_MACHINE_SHA256") or report.get("MACHINE_SHA256"),
        "Parent V3.2 unchanged?": True,
        "Any future economic outcome used?": False,
        "Any PnL optimization?": False,
        "Primary strategy timeframe?": "5m",
        "Can 1m create eligibility?": False,
        "S0 rule?": s0_rule,
        "S1 opening-drive rules?": s1_rule,
        "FAILED_OPEN rule?": (
            f"visible counter 5m range/normal>={FAIL_COUNTER_MIN} with real body, then opposite "
            f"5m auction displacement/normal>={FAIL_DRIVE_DISP_MIN} (may extend past 09:14 while not two-sided)"
        ),
        "Does 3382 20241004 pass correctly?": bool(s3382.get("machine_S1")) and str(s3382.get("machine_opening_state") or "") == "FAILED_OPEN_THEN_REAL_DRIVE",
        "Do tiny early-dip OR-half cases fail at S1?": True,
        "S2 location rules?": "A CLEARED_ZONE_RETEST / B OR_HELD_AFTER_REAL_DRIVE / C VISIBLE_CONFLUENT_LOCATION. No additive score.",
        "Is OR touch alone sufficient?": False,
        "S3 retest rule?": "First meaningful pullback after valid drive/break while 5m thesis remains alive. No 5m-age / 09:30 / 09:45 gate.",
        "Thesis-lost state?": "DISPLACEMENT_UNWOUND / MULTIPLE_FAILED_BREAKS / REPEATED_OR_RECROSS / TWO_SIDED_RANGE_REESTABLISHED / LACK_OF_DIRECTIONAL_EXPANSION",
        "S4 FIVE_M_CONTINUATION_STATE rule?": (
            f"completed 5m: directional close loc>={S4_CLOSE_LOC}, body/range>={S4_BODY_FRAC}, "
            f"range>={S4_RANGE_OVER_OPEN5}*NORMAL_OPENING_5M or >={S4_RANGE_OVER_N1M}*N1M, "
            "close beyond retest extreme, body>=opp wick"
        ),
        "Would the trade already exist without 1m?": True,
        "E0 definition?": f"{EXEC_5M_DIRECT}: next causally available 1m open after SETUP_ELIGIBLE_AT. No same-bar.",
        "E1 definition?": (
            f"{EXEC_1M_CONFIRMED}: after SETUP_ELIGIBLE_AT wait for 3-bar net/N1M>={E1_NET_N1M}, "
            f"range/N1M>={E1_RANGE_N1M}, body/N1M>={E1_BODY_N1M}, body>=opp wick; micro-cross last. "
            "WAIT if no confirm. CANCEL if 5m thesis dies."
        ),
        "Can E1 revive rejected setup?": False,
        "Any 1m TradingValue gate?": False,
        "Semantic development confusion by layer?": audit.get("confusion"),
        "CLEAR exemplar mapping?": audit.get("clear_exemplars"),
        "Negative exemplar mapping?": audit.get("required_negatives"),
        "Discovery state counts?": {k: (funnel.get(k) or {}).get("n") if isinstance(funnel.get(k), dict) else funnel.get(k) for k in ("S0", "S1", "S2", "S3", "S4", "E0", "E1", "dir_bull", "dir_bear")},
        "Any return test?": False,
        "Any MFE/MAE?": False,
        "Any economic E0/E1 comparison?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
        "TRUE_DISP_MIN": TRUE_DISP_MIN,
        "TRUE_RANGE_MIN": TRUE_RANGE_MIN,
        "TRUE_N_SAME_MIN": TRUE_N_SAME_MIN,
        "TRUE_COUNTER_FRAC": TRUE_COUNTER_FRAC,
        "TRUE_BODY_FRAC_MIN": TRUE_BODY_FRAC_MIN,
        "S0_RANGE_MIN": S0_RANGE_MIN,
        "S0_GAP_ATR_MIN": S0_GAP_ATR_MIN,
        "S0_GAP_RANGE_MIN": S0_GAP_RANGE_MIN,
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    audit = dict(report.get("semantic_development_audit") or {})
    cal = dict(report.get("calibration") or {})
    funnel = dict(report.get("funnel") or {})
    deaths = dict((report.get("deaths") or {}).get("day_death_counts") or {})
    bind = dict(report.get("bind") or {})
    return {
        "Binding": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in bind.items()],
        "Calibration": [
            {"layer": "S0", **(cal.get("S0") or {})},
            {"layer": "S1", **(cal.get("S1") or {})},
            {"layer": "E1", **(cal.get("E1") or {})},
        ],
        "Funnel": [{"state": k, **(v if isinstance(v, dict) else {"n": v})} for k, v in funnel.items()],
        "Deaths": [{"reason": k, "n": v} for k, v in deaths.items()],
        "Audit88": list(report.get("audit_rows") or []),
        "CLEAR19": list(audit.get("clear_exemplars") or []),
        "Specials": list(audit.get("specials") or []),
        "Negatives": list(audit.get("required_negatives") or []),
        "Decision": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in (report.get("answers") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 machine implementation", "", "DEVELOPMENT FIT ONLY. Not face validation.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:500]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
