"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_causal_path_failure_rca_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Funnel",
    "Stage_Path",
    "Consumption",
    "First_Passage",
    "Control_First_Passage",
    "Trigger_Split",
    "Failed_Push_Audit",
    "OR_Failure_Semantics",
    "Failure_Timing",
    "Block_Composition",
    "Composition_Standardization",
    "Retest_Age",
    "Clock",
    "Break_Quality",
    "Opening_Impulse",
    "In_Play",
    "Market_Regime",
    "Daily_Context",
    "Treated_vs_Control",
    "Human_Sample",
    "Cause_Hierarchy",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "events", "pairs", "recs", "eligible", "stages", "stage_paths", "market_days"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def json_sanitize(obj: Any) -> Any:
    return _json_sanitize(obj)


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    d = dict(report.get("decision") or {})
    st = dict(report.get("stage_edge") or {})
    cons = dict(report.get("consumption") or {})
    fp = dict((report.get("first_passage") or {}).get("by_block") or {})
    fpa = dict((report.get("first_passage") or {}).get("all") or {})
    cfp = dict((report.get("control_first_passage") or {}).get("all") or {})
    cfp_b = dict((report.get("control_first_passage") or {}).get("by_block") or {})
    trig = dict(report.get("trigger_split") or {})
    audit = dict(report.get("failed_push_audit") or {})
    orf = dict(report.get("or_failure_semantics") or {})
    ft = dict(report.get("failure_timing") or {})
    rw = dict(report.get("reweight") or {})
    age = dict(report.get("retest_age") or {})
    clk = dict(report.get("clocks") or {})
    mkt = dict(report.get("market_by_block") or {})
    tvc = dict(report.get("treated_vs_control") or {})
    human = dict(report.get("human") or {})
    hier = dict(report.get("hierarchy") or {})
    flags = dict(hier.get("flags") or {})

    def p1(block: str) -> Any:
        return (fp.get(block) or {}).get("P_plus_1_0R_before_fail")

    return {
        "favorable_path_first_appears": st.get("favorable_path_first_appears"),
        "stage_r10_p50": {k: ((st.get(k) or {}).get("r10_bps") or {}).get("p50") for k in ("S1", "S2", "S3", "S4", "S5", "S6")},
        "break_to_entry_bps_p50": (cons.get("break_to_entry") or {}).get("p50"),
        "break_to_entry_or_p50": (cons.get("break_to_entry_or") or {}).get("p50"),
        "max_ext_before_retest_p50": (cons.get("max_ext_before_retest") or {}).get("p50"),
        "plus_1R_before_OR_fail": {
            "all": fpa.get("P_plus_1_0R_before_fail"),
            "D2": p1("D2"),
            "D3": p1("D3"),
            "D4": p1("D4"),
        },
        "plus_1R_before_fail_beats_controls": flags.get("fp_beats_controls"),
        "control_plus_1R": cfp.get("P_plus_1_0R_before_fail"),
        "control_plus_1R_blocks": {b: (cfp_b.get(b) or {}).get("P_plus_1_0R_before_fail") for b in ("D2", "D3", "D4")},
        "reclaim_vs_failed_push_same_mechanism": trig.get("same_mechanism"),
        "failed_push_risk_invalid_why": audit.get("why"),
        "failed_push_mismatch": audit.get("MECHANISM_SEMANTICS_MISMATCH"),
        "or_fail_classes": orf.get("classes"),
        "or_fail_within": (ft.get("all") or {}),
        "reweight_instability_remains": rw.get("instability_remains"),
        "reweight": {b: rw.get(b) for b in ("D2", "D3", "D4")},
        "retest_age_r10": {k: ((age.get(k) or {}).get("r10_bps") or {}).get("p50") for k in ("0_5", "5_15", "15_30")},
        "late_break": (clk.get("break_clock") or {}),
        "market_breadth": {b: ((mkt.get(b) or {}).get("breadth_0915") or {}).get("p50") for b in ("D1", "D2", "D3", "D4")},
        "treated_vs_control_who": {b: (tvc.get(b) or {}).get("who") for b in ("D1", "D2", "D3", "D4")},
        "sample_n": human.get("sample_n"),
        "CLEAR_CONTINUATION_share": human.get("CLEAR_CONTINUATION_share"),
        "trigger_semantic_valid_share": human.get("trigger_semantic_valid_share"),
        "PRIMARY_CAUSE": hier.get("PRIMARY_CAUSE"),
        "SECONDARY_CAUSE": hier.get("SECONDARY_CAUSE"),
        "CONTRIBUTING_CAUSE": hier.get("CONTRIBUTING_CAUSE"),
        "any_eligibility_threshold_changed": False,
        "any_trigger_selected_by_pnl": False,
        "any_new_indicator_strategy": False,
        "any_pnl_optimization": False,
        "old_confirmation_opened": False,
        "frozen_validation_opened": False,
        "Kabu50": False,
        "submit_cancel_live": "0/0/0",
        "VERDICT": d.get("VERDICT"),
        "NEXT": d.get("NEXT"),
        "STOP_PB1": d.get("STOP_PB1"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    fun = dict(report.get("funnel") or {})
    lines = [
        "# PB1_CAUSAL_PATH_FAILURE_RCA_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        f"Funnel S0→S6: { {k: (fun.get(k) or {}).get('n') for k in ['S0','S1','S2','S3','S4','S5','S6']} }",
        f"Where does favorable path first appear? **{a.get('favorable_path_first_appears')}**",
        f"Stage 10m p50: {a.get('stage_r10_p50')}",
        f"How much move is consumed break→entry? median **{a.get('break_to_entry_bps_p50')}** bps / **{a.get('break_to_entry_or_p50')}** OR",
        f"Max post-break extension before retest p50: {a.get('max_ext_before_retest_p50')}",
        "",
        f"Does +1R happen before OR/retest failure? all={((a.get('plus_1R_before_OR_fail') or {}).get('all'))} "
        f"D2={((a.get('plus_1R_before_OR_fail') or {}).get('D2'))} "
        f"D3={((a.get('plus_1R_before_OR_fail') or {}).get('D3'))} "
        f"D4={((a.get('plus_1R_before_OR_fail') or {}).get('D4'))}",
        f"Does +1R-before-failure beat controls? **{a.get('plus_1R_before_fail_beats_controls')}** (control={a.get('control_plus_1R')})",
        f"Are RECLAIM and FAILED_PUSH the same mechanism? **{a.get('reclaim_vs_failed_push_same_mechanism')}**",
        f"Why are FAILED_PUSH RISK_INVALID? {a.get('failed_push_risk_invalid_why')}",
        f"MECHANISM_SEMANTICS_MISMATCH? **{a.get('failed_push_mismatch')}**",
        f"OR failure classes: {a.get('or_fail_classes')}",
        f"Failures within 1/3/5/10m: {a.get('or_fail_within')}",
        "",
        f"Does block instability remain after composition standardization? **{a.get('reweight_instability_remains')}**",
        f"Retest age 10m p50: {a.get('retest_age_r10')}",
        f"Market 09:15 breadth p50 by block: {a.get('market_breadth')}",
        f"Treated deteriorated or controls strengthened? {a.get('treated_vs_control_who')}",
        "",
        f"New blinded RCA sample_n? **{a.get('sample_n')}**",
        f"CLEAR_CONTINUATION share? **{a.get('CLEAR_CONTINUATION_share')}**",
        f"Trigger semantic-valid share? **{a.get('trigger_semantic_valid_share')}**",
        "",
        f"PRIMARY_CAUSE? **{a.get('PRIMARY_CAUSE')}**",
        f"SECONDARY_CAUSE? **{a.get('SECONDARY_CAUSE')}**",
        f"CONTRIBUTING_CAUSE? {a.get('CONTRIBUTING_CAUSE')}",
        "",
        f"Any eligibility threshold changed? {a.get('any_eligibility_threshold_changed')}",
        f"Any trigger selected by PnL? {a.get('any_trigger_selected_by_pnl')}",
        f"Any new indicator strategy? {a.get('any_new_indicator_strategy')}",
        f"Any PnL optimization? {a.get('any_pnl_optimization')}",
        f"Old Confirmation opened? {a.get('old_confirmation_opened')}",
        f"Frozen Validation opened? {a.get('frozen_validation_opened')}",
        f"Kabu50? {a.get('Kabu50')}",
        f"submit/cancel/live? {a.get('submit_cancel_live')}",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    human = dict(report.get("human") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Funnel": _kv_rows(report.get("funnel")),
        "Stage_Path": _kv_rows(report.get("stage_edge")),
        "Consumption": _kv_rows(report.get("consumption")),
        "First_Passage": _kv_rows(report.get("first_passage")),
        "Control_First_Passage": _kv_rows(report.get("control_first_passage")),
        "Trigger_Split": _kv_rows(report.get("trigger_split")),
        "Failed_Push_Audit": _kv_rows(report.get("failed_push_audit")),
        "OR_Failure_Semantics": _kv_rows(report.get("or_failure_semantics")),
        "Failure_Timing": _kv_rows(report.get("failure_timing")),
        "Block_Composition": _kv_rows(report.get("composition")),
        "Composition_Standardization": _kv_rows(report.get("reweight")),
        "Retest_Age": _kv_rows(report.get("retest_age")),
        "Clock": _kv_rows(report.get("clocks")),
        "Break_Quality": _kv_rows(report.get("break_quality")),
        "Opening_Impulse": _kv_rows(report.get("opening_impulse")),
        "In_Play": _kv_rows(report.get("in_play")),
        "Market_Regime": _kv_rows(report.get("market_by_block")),
        "Daily_Context": _kv_rows(report.get("daily_context")),
        "Treated_vs_Control": _kv_rows(report.get("treated_vs_control")),
        "Human_Sample": list(human.get("rows") or [_kv_rows(human)[0]]),
        "Cause_Hierarchy": _kv_rows(report.get("hierarchy")),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
