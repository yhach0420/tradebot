"""report.json / report.md / audit.xlsx only. No raw rewrite."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_microstructure_pressure_day1_v1 import ANALYSIS_ID
from research.futures_microstructure_pressure_day1_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = ("Answers", "Signed", "Clocks", "Decision", "Safety")


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"CASE: {d.get('CASE')}",
        f"NEXT: `{d.get('NEXT')}`",
        "",
        "Price-return lead is not retuned. Features are book pressure only.",
        "L1/DEPTH primary = 30s median. Quote pressure = 60s L1 update imbalance.",
        "Statistical unit is MINUTE CLOCK. Sign split >0 vs <0. Zero isolated.",
        "T+2m pressure vs stock 5m is NONCAUSAL_DIAGNOSTIC.",
        "Day2 frozen price-return test is unchanged.",
        "",
        f"- minute_clock_n: `{a.get('1_minute_clock_n')}`",
        f"- leakage_n: `{a.get('2_leakage_n')}`",
        "",
        "## Required answers",
        "",
        f"- 1_minute_clock_n: `{a.get('1_minute_clock_n')}`",
        f"- 2_leakage_n: `{a.get('2_leakage_n')}`",
        f"- 3_NK_L1 → NK 30s: `{_fmt(a.get('3_NK_L1_imbalance_NK_30s'))}` bps",
        f"- 4_NK_L1 → NK 60s: `{_fmt(a.get('4_NK_L1_imbalance_NK_60s'))}` bps",
        f"- 5_NK_depth10 → NK 30s: `{_fmt(a.get('5_NK_depth10_NK_30s'))}` bps",
        f"- 6_NK_depth10 → NK 60s: `{_fmt(a.get('6_NK_depth10_NK_60s'))}` bps",
        f"- 7_NK_quote_pressure → NK price: `{json.dumps(a.get('7_NK_quote_pressure_NK_price'), ensure_ascii=False, default=str)}`",
        f"- 8_TOPIX_equivalents: `{json.dumps(a.get('8_TOPIX_equivalents'), ensure_ascii=False, default=str)}`",
        f"- 9_strongest_Stage1_feature: `{a.get('9_strongest_Stage1_feature')}`",
        f"- 10_strongest_Stage1_direction: `{a.get('10_strongest_Stage1_direction')}`",
        f"- 11_stock_1m: `{_fmt(a.get('11_same_feature_stock_1m'))}` bps",
        f"- 12_stock_3m: `{_fmt(a.get('12_same_feature_stock_3m'))}` bps",
        f"- 13_stock_5m: `{_fmt(a.get('13_same_feature_stock_5m'))}` bps",
        f"- 14_strongest_stock_MID_shift: `{_fmt(a.get('14_strongest_stock_MID_shift'))}` bps",
        f"- 15_LONG_improvement: `{_fmt(a.get('15_LONG_improvement'))}` bps",
        f"- 16_SHORT_improvement: `{_fmt(a.get('16_SHORT_improvement'))}` bps",
        f"- 17_absolute_LONG_positive: `{_fmt(a.get('17_absolute_LONG_positive'))}`",
        f"- 18_absolute_SHORT_positive: `{_fmt(a.get('18_absolute_SHORT_positive'))}`",
        f"- 19_EARLY_direction: `{a.get('19_EARLY_direction')}`",
        f"- 20_LATE_direction: `{a.get('20_LATE_direction')}`",
        f"- 21_internal_hold: `{_fmt(a.get('21_internal_hold'))}`",
        f"- 22_T+2m_placebo: `{_fmt(a.get('22_Tplus2m_placebo'))}` bps",
        f"- 23_causal_beats_placebo: `{_fmt(a.get('23_causal_beats_placebo'))}`",
        f"- 24_prior_price_explains: `{json.dumps(a.get('24_prior_price_explains'), ensure_ascii=False, default=str)}`",
        f"- 25_exact_causal_sequence: {a.get('25_exact_causal_sequence')}",
        f"- 26_usable_mechanism_exists: `{_fmt(a.get('26_usable_mechanism_exists'))}`",
        f"- 27_ENTRY_built: `{_fmt(a.get('27_ENTRY_built'))}`",
        f"- 28_EXIT_built: `{_fmt(a.get('28_EXIT_built'))}`",
        f"- 29_Day2_frozen_price_return_test_changed: `{_fmt(a.get('29_Day2_frozen_price_return_test_changed'))}`",
        f"- 30_Runtime_changed: `{_fmt(a.get('30_Runtime_changed'))}`",
        f"- 31_Paper_changed: `{_fmt(a.get('31_Paper_changed'))}`",
        f"- 32_submit_cancel_live: `{a.get('32_submit_cancel_live')}`",
        f"- 33_VERDICT: `{a.get('33_VERDICT')}`",
        f"- 34_NEXT: `{a.get('34_NEXT')}`",
        "",
        "## Stage1 / Stage2 by frozen feature",
        "",
        "| feature | S1 30s bps | S1 60s bps | S1 pass | stk 1m | stk 3m | stk 5m | S2 pass | hold | placebo 5m | beats plc | candidate |",
        "|---|---:|---:|---|---:|---:|---:|---|---|---:|---|---|",
    ]
    for r in body.get("signed_rows") or []:
        s2 = r.get("stage2") or {}
        lines.append(
            "| {f} | {a} | {b} | {p1} | {m1} | {m3} | {m5} | {p2} | {h} | {plc} | {bp} | {c} |".format(
                f=r.get("feature"),
                a=_fmt((r.get("stage1_30") or {}).get("shift_bps")),
                b=_fmt((r.get("stage1_60") or {}).get("shift_bps")),
                p1=_fmt(r.get("stage1_pass")),
                m1=_fmt((s2.get("1m") or {}).get("shift")),
                m3=_fmt((s2.get("3m") or {}).get("shift")),
                m5=_fmt((s2.get("5m") or {}).get("shift")),
                p2=_fmt(r.get("stage2_pass")),
                h=_fmt(r.get("internal_hold")),
                plc=_fmt((r.get("placebo_5m") or {}).get("shift")),
                bp=_fmt(r.get("causal_beats_placebo")),
                c=_fmt(r.get("day1_candidate")),
            )
        )
    lines.extend(
        [
            "",
            "ENTRY: false",
            "EXIT: false",
            "Day2 frozen price-return test changed: false",
            "submit/cancel/live: 0/0/0",
            "",
        ]
    )
    return "\n".join(lines)


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    report = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(report), encoding="utf-8")
    signed_slim = []
    for r in report.get("signed_rows") or []:
        s2 = r.get("stage2") or {}
        signed_slim.append(
            {
                "feature": r.get("feature"),
                "s1_30_bps": (r.get("stage1_30") or {}).get("shift_bps"),
                "s1_60_bps": (r.get("stage1_60") or {}).get("shift_bps"),
                "s1_pass": r.get("stage1_pass"),
                "stk_1m": (s2.get("1m") or {}).get("shift"),
                "stk_3m": (s2.get("3m") or {}).get("shift"),
                "stk_5m": (s2.get("5m") or {}).get("shift"),
                "s2_pass": r.get("stage2_pass"),
                "hold": r.get("internal_hold"),
                "placebo_5m": (r.get("placebo_5m") or {}).get("shift"),
                "beats_placebo": r.get("causal_beats_placebo"),
                "candidate": r.get("day1_candidate"),
                "long_impr": r.get("long_impr_bps"),
                "short_impr": r.get("short_impr_bps"),
                "abs_long_pos": r.get("absolute_LONG_positive"),
                "abs_short_pos": r.get("absolute_SHORT_positive"),
            }
        )
    wb = Workbook()
    first = True
    sheets = {
        "Answers": _kv(dict(report.get("answers") or {})),
        "Signed": signed_slim,
        "Clocks": list(report.get("clocks_slim") or []),
        "Decision": _kv(dict(report.get("decision") or {})),
        "Safety": [
            {
                "submit_cancel_live": (report.get("answers") or {}).get("32_submit_cancel_live"),
                "ENTRY": False,
                "EXIT": False,
                "Day2_frozen_price_return_test_changed": False,
                "Runtime_changed": False,
                "Paper_changed": False,
            }
        ],
    }
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {"report_json": str(OUT / "report.json"), "report_md": str(OUT / "report.md"), "audit_xlsx": str(xlsx)}
