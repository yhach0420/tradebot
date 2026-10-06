"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_playbook_redesign_v3.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Rule_Diff",
    "Funnel",
    "Deaths",
    "Machine",
    "Human_Sample",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "events", "failed_push_archive", "funnel_days", "chart_zones"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def json_sanitize(obj: Any) -> Any:
    return _json_sanitize(obj)


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([r.get(c) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    deaths = dict(report.get("deaths") or {})
    human = dict(report.get("human") or {})
    dec = dict(report.get("decision") or {})
    return {
        "v2_unchanged": report.get("v2_unchanged"),
        "V3_MACHINE_SHA256": report.get("MACHINE_SHA256"),
        "PARENT_V2_SHA": report.get("PARENT_V2_SHA"),
        "rule_diff": report.get("rule_diff"),
        "FAILED_PUSH_removed_from_PB1": report.get("failed_push_removed_from_pb1"),
        "FAILED_PUSH_archived": report.get("failed_push_archived"),
        "any_return_used_to_remove_it": report.get("failed_push_removed_for_return"),
        "structural_route_defined": report.get("structural_route_defined"),
        "why_1R": report.get("why_1R"),
        "one_r_searched": report.get("one_r_searched"),
        "prior_zone_becomes_support": report.get("prior_zone_becomes_support"),
        "defended_location_defined": report.get("defended_location_defined"),
        "exhausted_retest_defined": report.get("exhausted_retest_defined"),
        "retest_5min_gate": report.get("retest_5min_gate"),
        "extension_bps_optimized": report.get("extension_bps_optimized"),
        "primary_trigger": report.get("primary_trigger"),
        "SAME_BAR_ENTRY_N": report.get("same_bar_entry_n"),
        "V3_event_n": report.get("setup_n"),
        "FAILED_PUSH_NOT_PB1": (deaths.get("FAILED_PUSH_NOT_PB1") or {}).get("count_field"),
        "NO_DEFENDED_LOCATION": (deaths.get("NO_DEFENDED_LOCATION") or {}).get("day_deaths"),
        "STRUCTURALLY_BLOCKED": (deaths.get("STRUCTURALLY_BLOCKED") or {}).get("day_deaths"),
        "MOVE_ALREADY_REACHED_STRUCTURE": (deaths.get("MOVE_ALREADY_REACHED_STRUCTURE") or {}).get("day_deaths"),
        "NO_RECLAIM": (deaths.get("NO_RECLAIM") or {}).get("day_deaths"),
        "RISK_INVALID_n": report.get("risk_invalid_n"),
        "sample_n": report.get("sample_n"),
        "actual_manual_blinded_review": human.get("actual_manual_blinded_review"),
        "CLEAR_CONTINUATION_n": human.get("CLEAR_CONTINUATION"),
        "CLEAR_CONTINUATION_share": human.get("clear_share"),
        "STRUCTURALLY_BLOCKED_leakage": human.get("structurally_blocked_leak_share"),
        "FRESH_RETEST_share": human.get("fresh_retest_share"),
        "VALID_RECLAIM_share": human.get("valid_reclaim_share"),
        "any_market_regime_gate": report.get("market_regime_gate"),
        "any_symbol_archetype_gate": report.get("symbol_archetype_gate"),
        "any_CLEAN_FLIP_winner_rule": report.get("clean_flip_winner_rule"),
        "any_pnl_test": report.get("pnl_test"),
        "any_threshold_optimized": False,
        "old_confirmation_opened": False,
        "frozen_validation_opened": False,
        "submit_cancel_live": "0/0/0",
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "FACE_VALID": dec.get("FACE_VALID"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    lines = [
        "# PB1_PLAYBOOK_REDESIGN_V3",
        "",
        "Minimum coherent redesign. Frozen V2 unchanged. No PnL.",
        "",
        f"V2 unchanged? **{a.get('v2_unchanged')}**",
        f"V3 machine SHA? `{a.get('V3_MACHINE_SHA256')}`",
        f"Exact V2→V3 rule changes? {a.get('rule_diff')}",
        "",
        f"FAILED_PUSH removed from PB1? **{a.get('FAILED_PUSH_removed_from_PB1')}**",
        f"FAILED_PUSH archived? **{a.get('FAILED_PUSH_archived')}**",
        f"Any return used to remove it? **{a.get('any_return_used_to_remove_it')}**",
        "",
        f"How is structural route defined? {a.get('structural_route_defined')}",
        f"Why exactly 1R? {a.get('why_1R')}",
        f"Was 1R searched? **{a.get('one_r_searched')}**",
        f"How is a prior S/R zone allowed to become support/resistance after the break? {a.get('prior_zone_becomes_support')}",
        f"How is defended location defined? {a.get('defended_location_defined')}",
        f"How is exhausted retest defined? {a.get('exhausted_retest_defined')}",
        f"Any 5-minute retest gate? **{a.get('retest_5min_gate')}**",
        f"Any extension-bps optimization? **{a.get('extension_bps_optimized')}**",
        "",
        f"Primary trigger? **{a.get('primary_trigger')}**",
        f"SAME_BAR_ENTRY_N? **{a.get('SAME_BAR_ENTRY_N')}**",
        f"V3 event_n? **{a.get('V3_event_n')}**",
        "",
        f"FAILED_PUSH_NOT_PB1? **{a.get('FAILED_PUSH_NOT_PB1')}**",
        f"NO_DEFENDED_LOCATION? **{a.get('NO_DEFENDED_LOCATION')}**",
        f"STRUCTURALLY_BLOCKED? **{a.get('STRUCTURALLY_BLOCKED')}**",
        f"MOVE_ALREADY_REACHED_STRUCTURE? **{a.get('MOVE_ALREADY_REACHED_STRUCTURE')}**",
        f"NO_RECLAIM? **{a.get('NO_RECLAIM')}**",
        f"RISK_INVALID_n? **{a.get('RISK_INVALID_n')}**",
        "",
        f"New face sample_n? **{a.get('sample_n')}**",
        f"Actual manual/blinded review? **{a.get('actual_manual_blinded_review')}**",
        f"CLEAR_CONTINUATION n/share? **{a.get('CLEAR_CONTINUATION_n')}** / **{a.get('CLEAR_CONTINUATION_share')}**",
        f"STRUCTURALLY_BLOCKED leakage? **{a.get('STRUCTURALLY_BLOCKED_leakage')}**",
        f"FRESH_RETEST share? **{a.get('FRESH_RETEST_share')}**",
        f"VALID_RECLAIM share? **{a.get('VALID_RECLAIM_share')}**",
        "",
        f"Any market-regime gate? **{a.get('any_market_regime_gate')}**",
        f"Any symbol/archetype gate? **{a.get('any_symbol_archetype_gate')}**",
        f"Any CLEAN_FLIP winner rule? **{a.get('any_CLEAN_FLIP_winner_rule')}**",
        f"Any PnL test? **{a.get('any_pnl_test')}**",
        f"Any threshold optimized? **{a.get('any_threshold_optimized')}**",
        f"Old Confirmation opened? **{a.get('old_confirmation_opened')}**",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"submit/cancel/live? **{a.get('submit_cancel_live')}**",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    human = dict(report.get("human") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Rule_Diff": _kv_rows(report.get("rule_diff")),
        "Funnel": _kv_rows(report.get("funnel")),
        "Deaths": _kv_rows(report.get("deaths")),
        "Machine": _kv_rows(
            {
                "MACHINE_SHA256": report.get("MACHINE_SHA256"),
                "PARENT_V2_SHA": report.get("PARENT_V2_SHA"),
                "STATE_MACHINE_TEXT": report.get("STATE_MACHINE_TEXT"),
            }
        ),
        "Human_Sample": list(human.get("rows") or [_kv_rows(human)[0]]),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
