"""Write report.json / report.md / audit.xlsx / universe_manifest.json only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.fixed_daytrade_universe_v1 import ANALYSIS_ID
from research.fixed_daytrade_universe_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "FINAL_UNIVERSE",
    "CANDIDATE_AUDIT",
    "LIQUIDITY_60D",
    "LIQUIDITY_30_30",
    "OPERABILITY",
    "LISTING_MASTER",
    "SECTOR_COVERAGE",
    "EXCLUSIONS",
    "SOURCE_AUDIT",
    "MANIFEST",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        "FINAL_UNIVERSE": list(report.get("final_universe") or []),
        "CANDIDATE_AUDIT": list(report.get("universe_candidates") or []),
        "LIQUIDITY_60D": list(report.get("liquidity_60d") or []),
        "LIQUIDITY_30_30": list(report.get("liquidity_30_30") or []),
        "OPERABILITY": list(report.get("operability") or []) or _kv(dict(report.get("operability_rules") or {})),
        "LISTING_MASTER": list(report.get("listing_master") or []),
        "SECTOR_COVERAGE": list(report.get("sector_coverage") or []),
        "EXCLUSIONS": list(report.get("exclusions") or []),
        "SOURCE_AUDIT": list(report.get("source_audit") or []),
        "MANIFEST": _kv(dict(report.get("manifest") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            str(report.get("threshold_rationale") or ""),
            "",
            "## Required answers",
            "",
            f"1. source period? {a.get('1_source_period')}",
            f"2. candidate N? **{a.get('2_candidate_N')}**",
            f"3. quantitative liquidity data available? **{a.get('3_quantitative_liquidity_data_available')}**",
            f"4. final N? {a.get('4_final_N')}",
            f"5. final symbols? {a.get('5_final_symbols')}",
            f"6. median trading value distribution? {a.get('6_median_trading_value_distribution')}",
            f"7. minimum p20 trading value? {a.get('7_minimum_p20_trading_value')}",
            f"8. session coverage minimum? {a.get('8_session_coverage_minimum')}",
            f"9. listing continuity pass N? {a.get('9_listing_continuity_pass_N')}",
            f"10. operational compatibility pass N? {a.get('10_operational_compatibility_pass_N')}",
            f"11. 100-share notional audit complete? **{a.get('11_100_share_notional_audit_complete')}**",
            f"12. trading-unit audit complete? **{a.get('12_trading_unit_audit_complete')}**",
            f"13. 9983 exact exclusion reason? {a.get('13_9983_exact_exclusion_reason')}",
            f"14. 6861 exact exclusion reason? {a.get('14_6861_exact_exclusion_reason')}",
            f"15. same rule applied to all candidates? **{a.get('15_same_rule_applied_to_all_candidates')}**",
            f"16. sector bucket N? **{a.get('16_sector_bucket_N')}**",
            f"17. any required sector lost? **{a.get('17_any_required_sector_lost')}**",
            f"18. any low-liquidity name retained only for sector coverage? **{a.get('18_any_low_liquidity_name_retained_only_for_sector_coverage')}**",
            f"19. outcome/PnL used? **{a.get('19_outcome_PnL_used')}**",
            f"20. future return used? **{a.get('20_future_return_used')}**",
            f"21. Paper result used? **{a.get('21_Paper_result_used')}**",
            f"22. PANEL_CONDITIONED? **{a.get('22_PANEL_CONDITIONED')}**",
            f"23. universe frozen? **{a.get('23_universe_frozen')}**",
            f"24. universe SHA? `{a.get('24_universe_SHA')}`",
            f"25. review cadence? {a.get('25_review_cadence')}",
            f"26. Phase0 autosplice inconsistency corrected? **{a.get('26_Phase0_autosplice_inconsistency_corrected')}**",
            f"27. strategy search started? **{a.get('27_strategy_search_started')}**",
            f"28. panel built? **{a.get('28_panel_built')}**",
            f"29. 20260914 live changed? **{a.get('29_20260914_live_changed')}**",
            f"30. Runtime changed? **{a.get('30_Runtime_changed')}**",
            f"31. Paper changed? **{a.get('31_Paper_changed')}**",
            f"32. submit/cancel/live? **{a.get('32_submit_cancel_live')}**",
            f"33. VERDICT? **{a.get('33_VERDICT')}**",
            f"34. NEXT? **{a.get('34_NEXT')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    (OUT / "universe_manifest.json").write_text(
        json.dumps(json_sanitize(dict(report.get("manifest") or {})), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
    assert (OUT / "universe_manifest.json").is_file()
