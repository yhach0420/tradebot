"""report.json / report.md / audit.xlsx only. No raw rewrite. No ranking."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_x_stock_state_interaction_day1_v1 import ANALYSIS_ID
from research.futures_x_stock_state_interaction_day1_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "Clocks",
    "Grid",
    "Passing",
    "Best",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def _grid_row(r: dict[str, Any]) -> dict[str, Any]:
    skip = {"clock_mid_spreads", "ctx_clocks"}
    return {k: v for k, v in r.items() if k not in skip}


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    best = dict(body.get("best") or {})
    c = dict(body.get("context_counts") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"CASE: {d.get('CASE')}",
        "",
        "Question: given the same frozen futures context, does stock-specific state",
        "split later TOP vs BOTTOM results, and does that split change with context?",
        "Primary metric is CROSS_SECTIONAL_SELECTOR_SPREAD (TOP16 − BOTTOM16), not the",
        "48-stock market mean. 20260911 FULL only. No ranking/breadth tape. No futures-only",
        "price or book-pressure retest. No ENTRY/EXIT. No threshold search.",
        "",
        "Causal rule: stock and futures `received_at <= T`. future leakage must be 0.",
        "Clocks match FUTURES_CONTEXT_DAY1_EFFECT_CHECK_V1 (13 anchors, 09:10–11:10).",
        "",
        "## Required answers",
        "",
        f"- 1 clock N: `{a.get('1_clock_n')}`",
        f"- 2 stock N: `{a.get('2_stock_n')}`",
        f"- 3 leakage N: `{a.get('3_leakage_n')}`",
        f"- 4 NK180 UP clock N: `{a.get('4_NK180_UP_clock_n')}`",
        f"- 5 NK180 DOWN clock N: `{a.get('5_NK180_DOWN_clock_n')}`",
        f"- 6 agreement bucket Ns: `{a.get('6_agreement_bucket_ns')}`",
        f"- 7 alignment bucket Ns: `{a.get('7_alignment_bucket_ns')}`",
        f"- 8 selectors tested: `{a.get('8_selectors_tested')}`",
        f"- 9 total frozen comparisons: `{a.get('9_total_frozen_comparisons')}`",
        f"- 10 best interaction selector: `{a.get('10_best_interaction_selector')}`",
        f"- 11 best futures context: `{a.get('11_best_futures_context')}`",
        f"- 12 best horizon: `{a.get('12_best_horizon')}`",
        f"- 13 TOP basket MID: `{_fmt(a.get('13_TOP_basket_MID'))}` bps",
        f"- 14 BOTTOM basket MID: `{_fmt(a.get('14_BOTTOM_basket_MID'))}` bps",
        f"- 15 TOP-BOTTOM MID spread: `{_fmt(a.get('15_TOP_BOTTOM_MID_spread'))}` bps",
        f"- 16 selector-only BASE spread: `{_fmt(a.get('16_selector_only_BASE_spread'))}` bps",
        f"- 17 incremental interaction lift: `{_fmt(a.get('17_incremental_interaction_lift'))}` bps",
        f"- 18 LONG TOP absolute mean/median: `{a.get('18_LONG_TOP_absolute_mean_median')}`",
        f"- 19 SHORT TOP absolute mean/median: `{a.get('19_SHORT_TOP_absolute_mean_median')}`",
        f"- 20 executable improvement: `{_fmt(a.get('20_executable_improvement'))}`",
        f"- 21 EARLY spread: `{_fmt(a.get('21_EARLY_spread'))}` bps",
        f"- 22 LATE spread: `{_fmt(a.get('22_LATE_spread'))}` bps",
        f"- 23 internal hold: `{_fmt(a.get('23_internal_hold'))}`",
        f"- 24 drop top clock result: `{a.get('24_drop_top_clock_result')}`",
        f"- 25 drop top2 clocks result: `{a.get('25_drop_top2_clocks_result')}`",
        f"- 26 drop top symbol result: `{a.get('26_drop_top_symbol_result')}`",
        f"- 27 drop top2 symbols result: `{a.get('27_drop_top2_symbols_result')}`",
        f"- 28 09:30/09:50 concentration: `{_fmt(a.get('28_0930_0950_concentration'))}`",
        f"- 29 context clock support sufficient: `{_fmt(a.get('29_context_clock_support_sufficient'))}`",
        f"- 30 exact mechanism sentence: {a.get('30_exact_mechanism_sentence')}",
        f"- 31 ENTRY thesis plausible: `{_fmt(a.get('31_ENTRY_thesis_plausible'))}`",
        f"- 32 ENTRY built: `{_fmt(a.get('32_ENTRY_built'))}`",
        f"- 33 EXIT built: `{_fmt(a.get('33_EXIT_built'))}`",
        f"- 34 Runtime changed: `{_fmt(a.get('34_Runtime_changed'))}`",
        f"- 35 Paper changed: `{_fmt(a.get('35_Paper_changed'))}`",
        f"- 36 submit/cancel/live: `{a.get('36_submit_cancel_live')}`",
        f"- 37 CASE: `{a.get('37_CASE')}`",
        f"- 38 VERDICT: `{a.get('38_VERDICT')}`",
        f"- 39 NEXT: `{a.get('39_NEXT')}`",
        "",
        "## Freeze",
        "",
        f"- selectors: {', '.join(body.get('selectors') or [])}",
        f"- context families: {', '.join(body.get('context_families') or [])}",
        f"- tercile: TOP {body.get('tercile_n')} / BOTTOM {body.get('tercile_n')} / middle excluded",
        f"- material lift: >= {body.get('material_bps')} bps vs BASE A",
        f"- candidate_n: {d.get('candidate_n')} executable_candidate_n: {d.get('executable_candidate_n')}",
        "",
        "## Context clock N",
        "",
        f"- NK_RET_180S UP={c.get('NK180_UP')} DOWN={c.get('NK180_DOWN')}",
        f"- AGREEMENT_180S BOTH_UP={c.get('AGREEMENT_BOTH_UP')} MIXED={c.get('AGREEMENT_MIXED')} BOTH_DOWN={c.get('AGREEMENT_BOTH_DOWN')}",
        f"- NK_VS_CASH_ALIGN_180S ALIGNED={c.get('ALIGN_ALIGNED')} DISAGREED={c.get('ALIGN_DISAGREED')}",
        "",
        "## Best row gates",
        "",
        f"- selector={best.get('selector')} context={best.get('family')}={best.get('bucket')} horizon={best.get('horizon')}",
        f"- BASE A (selector-only TOP-BOTTOM)={_fmt(best.get('base_a_mid_spread'))} bps",
        f"- BASE B (context 48-stock market mean)={_fmt(best.get('base_b_market_mid'))} bps",
        f"- INTERACTION (context TOP-BOTTOM)={_fmt(best.get('interaction_mid_spread'))} bps",
        f"- lift vs BASE A={_fmt(best.get('lift_mid'))} bps",
        f"- candidate={_fmt(best.get('candidate'))} absolute_executable_mean>0={_fmt(best.get('absolute_executable_mean_positive'))}",
        f"- strong_absolute (mean and median > 0): `{_fmt(best.get('strong_absolute'))}`",
        "",
        "## Passing rows (all gates)",
        "",
        "Day1 is exploratory. These rows all pass materiality, executable same-direction",
        "improvement, EARLY/LATE sign hold, clock exclusion, symbol exclusion, and",
        "context clock N >= 3. Best row alone is not a strategy.",
        "",
    ]
    passing = list(body.get("passing") or [])
    if passing:
        lines.append("| selector | context | horizon | N | TOP-BOTTOM | BASE A | lift | market mean | TOP LONG mean/median | abs exec |")
        lines.append("|---|---|---|---|---|---|---|---|---|---|")
        for r in passing:
            lng = f"{_fmt(r.get('top_long_mean'))}/{_fmt(r.get('top_long_median'))}"
            lines.append(
                f"| {r.get('selector')} | {r.get('family')}={r.get('bucket')} | {r.get('horizon')} | "
                f"{r.get('ranking_clock_n')} | {_fmt(r.get('interaction_mid_spread'))} | "
                f"{_fmt(r.get('base_a_mid_spread'))} | {_fmt(r.get('lift_mid'))} | "
                f"{_fmt(r.get('base_b_market_mid'))} | {lng} | {_fmt(r.get('absolute_executable_mean_positive'))} |"
            )
        lines.append("")
    spreads = dict(best.get("clock_mid_spreads") or {})
    if spreads:
        lines.extend(
            [
                "## Best-row per-clock TOP-BOTTOM MID (bps)",
                "",
            ]
        )
        for ck in sorted(spreads):
            lines.append(f"- {ck}: `{_fmt(spreads.get(ck))}`")
        lines.append("")
    lines.extend(
        [
            "ENTRY/EXIT remain false. Runtime/Paper unchanged. submit/cancel/live = 0/0/0.",
            "",
        ]
    )
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(body)
    payload = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(md, encoding="utf-8")
    wb = Workbook()
    sheets = {
        "Answers": _kv(dict(body.get("answers") or {})),
        "Clocks": list(body.get("clocks") or []),
        "Grid": [_grid_row(r) for r in (body.get("grid") or [])],
        "Passing": [_grid_row(r) for r in (body.get("passing") or [])],
        "Best": _kv({k: v for k, v in dict(body.get("best") or {}).items() if k != "clock_mid_spreads"}),
        "Decision": _kv(dict(body.get("decision") or {})),
        "Safety": [
            {"key": "ENTRY_built", "value": False},
            {"key": "EXIT_built", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "used_ranking", "value": False},
            {"key": "used_breadth", "value": False},
            {"key": "submit_cancel_live", "value": (body.get("answers") or {}).get("36_submit_cancel_live")},
            {"key": "future_leakage_n", "value": body.get("future_leakage_n")},
        ],
    }
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(xlsx),
    }
