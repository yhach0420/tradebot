"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.support_resistance_face_valid_first_interaction_rebuild_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Swing_Confirmation",
    "Reaction_Points",
    "Zone_Build",
    "Zone_Activation",
    "Zone_Invalidation",
    "Selected_Salient_Levels",
    "Broken_Level_State",
    "First_Interaction",
    "State_Machine",
    "Retest_Semantics",
    "Duplicate_Audit",
    "Timestamp_Audit",
    "New_Chart_Sample",
    "Face_Validity",
    "Matched_Test_Precommit",
    "Outcome_Precommit",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, (list, tuple)):
        return [json_sanitize(x) for x in got]
    return got


def _kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for k, v in d.items():
        if isinstance(v, (dict, list, tuple)):
            v = json.dumps(json_sanitize(v), ensure_ascii=False)[:32000]
        rows.append({"key": str(k), "value": v})
    return rows


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["key", "value"])
        ws.append(["empty", True])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    bind = dict(report.get("bind") or {})
    sw = dict(report.get("swing") or {})
    dens = dict(report.get("density") or {})
    gate = dict(report.get("gates") or {})
    ep = dict(report.get("episodes") or {})
    face = dict(report.get("face") or {})
    d = dict(report.get("decision") or {})
    safety = dict(report.get("safety") or {})
    return {
        "Binding": _kv_rows(
            {
                "analysis_id": report.get("analysis_id"),
                "parent_verdict": report.get("parent_verdict"),
                "bind_ok": bind.get("ok"),
                "old_no_info_revived": False,
                "no_pnl_parameter_tuning": True,
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
            }
        ),
        "Swing_Confirmation": _kv_rows(sw),
        "Reaction_Points": list(report.get("rx_compact") or []),
        "Zone_Build": _kv_rows({"half_atr": 0.15, "lookback_days": 60, "min_touches": 2, "independent_swing_cycles": True, **dens}),
        "Zone_Activation": _kv_rows({"activation": "available_from of 2nd independent confirmed swing", "retroactive_zone_n": gate.get("retroactive_zone_n")}),
        "Zone_Invalidation": _kv_rows({"rule": "prior completed close through far boundary", "already_broken_active_zone_n": gate.get("already_broken_active_zone_n")}),
        "Selected_Salient_Levels": list(report.get("rows_compact") or []),
        "Broken_Level_State": _kv_rows({"broken_does_not_auto_become_opposite_role": True, "flip_candidates_separately_labelled": True}),
        "First_Interaction": _kv_rows(ep),
        "State_Machine": _kv_rows(
            {
                "unit": "ONE symbol x zone x day first interaction",
                "primary": "FIRST_TEST -> REJECT or BREAK",
                "post_break": "ACCEPT2 / CLEAR / RETEST / HOLD or FAIL",
                "accept2_timestamp": "second confirming bar, never moved back to break",
            }
        ),
        "Retest_Semantics": _kv_rows(
            {
                "clear_required": True,
                "lingering_is_not_retest": True,
                "hold_and_fail_mutually_exclusive": True,
                "valid_retest_n": ep.get("valid_retest_n"),
                "retest_without_clear_n": gate.get("retest_without_clear_n"),
                "retest_hold_and_fail_overlap_n": gate.get("retest_hold_and_fail_overlap_n"),
            }
        ),
        "Duplicate_Audit": _kv_rows({"duplicate_first_interaction_n": gate.get("duplicate_first_interaction_n"), "canonical_record": True}),
        "Timestamp_Audit": _kv_rows({"retroactive_event_timestamp_n": gate.get("retroactive_event_timestamp_n"), "future_pivot_leakage_n": gate.get("future_pivot_leakage_n")}),
        "New_Chart_Sample": list(report.get("chart_index") or []),
        "Face_Validity": _kv_rows(face),
        "Matched_Test_Precommit": _kv_rows(dict(report.get("matched_test_precommit") or {})),
        "Outcome_Precommit": _kv_rows(dict(report.get("outcome_precommit") or {})),
        "Safety": _kv_rows(safety),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# SUPPORT_RESISTANCE_FACE_VALID_FIRST_INTERACTION_REBUILD_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            "This is not an economic test of support/resistance.",
            "",
            f"Confirmed swing highs? **{a.get('How many causal confirmed swing highs?')}** lows? **{a.get('swing lows?')}**",
            f"Selected resistance per symbol-day (median)? **{a.get('How many selected resistance levels per symbol-day?')}** support? **{a.get('support levels?')}**",
            f"already-broken active? **{a.get('Any already-broken active levels?')}** retroactive zones? **{a.get('Any retroactive zone creation?')}** future pivots? **{a.get('Any future pivot leakage?')}**",
            f"First-test episodes? **{a.get('How many first-test episodes?')}** raw events/episode? **{a.get('Raw events per episode?')}**",
            f"HOLD/FAIL overlap? **{a.get('Any RETEST_HOLD / FAILED_RETEST overlap?')}**",
            f"NEW 200 closer to discretionary S/R? **{a.get('Do the NEW 200 charts look materially closer to discretionary support/resistance?')}**",
            f"Clutter rate? **{a.get('Clutter rate?')}** relevant rate? **{a.get('Obvious/relevant rate?')}** stale stripes gone? **{a.get('Are stale stripes gone?')}**",
            f"No-level allowed? **{a.get('Does the detector sometimes correctly return NO LEVEL?')}**",
            f"PnL optimization? **{a.get('Any PnL optimization performed?')}** X0/X1 used to select rules? **{a.get('Any X0/X1 used to select rules?')}**",
            f"Old Confirmation opened? **{a.get('Old Confirmation opened?')}** Frozen Validation opened? **{a.get('Frozen Validation opened?')}**",
            f"submit/cancel/live: **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if k not in {"_markdown"}})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
