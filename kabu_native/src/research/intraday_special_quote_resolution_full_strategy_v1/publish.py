"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.full_causal_mechanism_discovery_v1.analyze import public_row
from research.intraday_special_quote_resolution_full_strategy_v1 import ANALYSIS_ID
from research.intraday_special_quote_resolution_full_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "data_boundary",
    "duplicate_check",
    "state_semantics",
    "trade_update_semantics",
    "trade_bar_semantics",
    "unit_tests",
    "canary",
    "precommit",
    "episodes",
    "releases",
    "accept_bars",
    "thesis_U",
    "thesis_D",
    "candidate_library",
    "execution",
    "portfolio",
    "coverage",
    "economics",
    "daily",
    "symbols",
    "causal_ex_top",
    "blocks",
    "selection",
    "selected_logic",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _episode_row(e: dict[str, Any]) -> dict[str, Any]:
    acc = dict(e.get("accept_bar") or {})
    return {
        "date": e.get("date"),
        "symbol": e.get("symbol"),
        "episode_seq": e.get("episode_seq"),
        "episode_id": e.get("episode_id"),
        "special_start_ingress": e.get("special_start_ingress"),
        "pre_special_price": e.get("pre_special_price"),
        "special_start_coobserved": e.get("special_start_coobserved"),
        "direction": e.get("direction"),
        "release_price": e.get("release_price"),
        "release_ingress": e.get("release_ingress"),
        "reinterrupted": e.get("reinterrupted"),
        "conflict": e.get("special_active_trade_conflict"),
        "no_entry_reason": e.get("no_entry_reason"),
        "signaled": e.get("signaled"),
        "accept_close": acc.get("close"),
        "accept_n": acc.get("TRADE_UPDATE_N"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    evals = [public_row(e) for e in list(report.get("candidate_evals") or [])]
    episodes = [_episode_row(e) for e in list(report.get("episodes") or [])]
    releases = [r for r in episodes if r.get("release_price") is not None]
    accepts = [r for r in episodes if r.get("accept_close") is not None]
    sigs = list(report.get("signals") or [])
    u_sig = [s for s in sigs if s.get("thesis") == "ISQ_UPWARD_RELEASE_ACCEPTANCE"]
    d_sig = [s for s in sigs if s.get("thesis") == "ISQ_DOWNWARD_RELEASE_FAILURE_RECLAIM"]
    daily_rows = []
    symbol_rows = []
    block_rows = []
    for ev in list(report.get("candidate_evals") or []):
        sid = ev.get("STRATEGY_ID")
        for day, pnl in dict(ev.get("daily") or {}).items():
            daily_rows.append({"STRATEGY_ID": sid, "date": day, "pnl": pnl})
        for t in list(ev.get("_trades") or []):
            symbol_rows.append(
                {
                    "STRATEGY_ID": sid,
                    "date": t.get("date"),
                    "symbol": t.get("symbol"),
                    "pnl": t.get("pnl_yen_100"),
                    "exit_reason": t.get("exit_reason"),
                }
            )
        for b in list((ev.get("blocks") or {}).get("blocks") or []):
            block_rows.append({"STRATEGY_ID": sid, **b})
    cov_rows = []
    econ_rows = []
    g6_rows = []
    for ev in evals:
        cov_rows.append(
            {
                "STRATEGY_ID": ev.get("STRATEGY_ID"),
                "C1": ev.get("C1"),
                "C2": ev.get("C2"),
                "C3": ev.get("C3"),
                "C4": ev.get("C4"),
                "coverage_ok": ev.get("coverage_ok"),
                "trade_n": ev.get("trade_n") or ev.get("TRADE_N"),
                "fill_day_n": ev.get("fill_day_n"),
                "trades_per_day": ev.get("trades_per_day"),
                "session_exit_unfilled_n": ev.get("session_exit_unfilled_n"),
            }
        )
        econ_rows.append(ev)
        g6_rows.append(
            {
                "STRATEGY_ID": ev.get("STRATEGY_ID"),
                "CAUSAL_EX_TOP1_PNL": ev.get("CAUSAL_EX_TOP1_PNL"),
                "CAUSAL_EX_TOP1_PF": ev.get("CAUSAL_EX_TOP1_PF"),
                "G6": (ev.get("g_table") or {}).get("G6") if isinstance(ev.get("g_table"), dict) else ev.get("G6"),
                "top_symbol": ev.get("top_symbol"),
                "METHOD": ev.get("CAUSAL_EX_TOP1_METHOD"),
            }
        )
    lib = list(((report.get("precommit") or {}).get("SPEC") or {}).get("CANDIDATES") or [])
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "data_boundary": _kv(dict(report.get("data_boundary") or {})),
        "duplicate_check": _kv(dict(report.get("duplicate_check") or {})),
        "state_semantics": _kv(dict(report.get("state_semantics") or {})),
        "trade_update_semantics": _kv(dict(report.get("trade_update_semantics") or {})),
        "trade_bar_semantics": _kv(dict(report.get("trade_bar_semantics") or {})),
        "unit_tests": list((report.get("unit_tests") or {}).get("rows") or []) or _kv(dict(report.get("unit_tests") or {})),
        "canary": _kv(dict(report.get("canary") or {})),
        "precommit": _kv(dict(report.get("precommit") or {})),
        "episodes": episodes,
        "releases": releases,
        "accept_bars": accepts,
        "thesis_U": u_sig,
        "thesis_D": d_sig,
        "candidate_library": lib or _kv({"CANDIDATES": (report.get("precommit") or {}).get("SPEC")}),
        "execution": _kv(dict(report.get("execution") or {})),
        "portfolio": _kv(dict(report.get("portfolio") or {})),
        "coverage": cov_rows,
        "economics": econ_rows,
        "daily": daily_rows,
        "symbols": symbol_rows,
        "causal_ex_top": g6_rows,
        "blocks": block_rows,
        "selection": _kv(dict(report.get("selection") or {})),
        "selected_logic": _kv(dict(report.get("selected_logic") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"LOGIC_COMPLETE: {d.get('LOGIC_COMPLETE')}",
            f"ROBUST_DEV_QUALIFIED: {d.get('ROBUST_DEV_QUALIFIED')}",
            f"OPENING_CURRENT_DATA_LINE_STATUS: {d.get('OPENING_CURRENT_DATA_LINE_STATUS')}",
            "TRUE_OOS: false",
            "CERTIFIED: false",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- parent pinned: {a.get('2_parent_verdict_pinned')}",
            f"- U/D duplicate: {a.get('8_U_exact_duplicate')}/{a.get('11_D_semantic_duplicate')}",
            f"- tests: {a.get('24_synthetic_tests_PASS_N_30')} all={a.get('25_tests_all_pass')}",
            f"- canary PASS: {a.get('31_canary_PASS')}",
            f"- episodes: {a.get('37_ISQ_episode_N')} U_sig={a.get('51_U_signal_N')} D_sig={a.get('52_D_signal_N')}",
            f"- selected: {a.get('80_selected_Strategy_ID')}",
            f"- LOGIC_COMPLETE: {a.get('95_LOGIC_COMPLETE')}",
            "",
            "No MBO. No futures. No Sizing. No Holdout/Stress/future. No V2 retune. Opening line CLOSED.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
