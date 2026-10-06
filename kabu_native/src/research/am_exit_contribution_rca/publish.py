"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize
from research.am_exit_contribution_rca import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_exit_contribution_rca"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "summary",
    "trade_paths",
    "c0_losses",
    "current_losses",
    "exit_reason",
    "giveback",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    q = report.get("questions") or {}
    lines = [
        "# AM EXIT CONTRIBUTION RCA V1",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "Frozen C0 / frozen C14. Diagnostic only. No EXIT or ENTRY change.",
        "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "",
        "## REQUIRED FINAL OUTPUT",
        "",
        f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
        "",
        f"CURRENT_TRADE_N: {_fmt(g.get('CURRENT_TRADE_N'))}",
        "",
        f"C0_AUGMENT_TRADE_N: {_fmt(g.get('C0_AUGMENT_TRADE_N'))}",
        "",
        f"C0_AUGMENT_WIN_N: {_fmt(g.get('C0_AUGMENT_WIN_N'))}",
        "",
        f"C0_AUGMENT_LOSS_N: {_fmt(g.get('C0_AUGMENT_LOSS_N'))}",
        "",
        f"C0_AUGMENT_FLAT_N: {_fmt(g.get('C0_AUGMENT_FLAT_N'))}",
        "",
        f"C0_AUGMENT_NET: {_fmt(g.get('C0_AUGMENT_NET'))}",
        "",
        f"C0_LOSS_ENTRY_NEVER_PROFITABLE_N: {_fmt(g.get('C0_LOSS_ENTRY_NEVER_PROFITABLE_N'))}",
        "",
        f"C0_LOSS_PROFIT_TO_LOSS_N: {_fmt(g.get('C0_LOSS_PROFIT_TO_LOSS_N'))}",
        "",
        f"C0_LOSS_RECOVERED_AFTER_EXIT_N: {_fmt(g.get('C0_LOSS_RECOVERED_AFTER_EXIT_N'))}",
        "",
        f"C0_LOSS_PARTIAL_RECOVERY_N: {_fmt(g.get('C0_LOSS_PARTIAL_RECOVERY_N'))}",
        "",
        f"C0_LOSS_GROSS_LOSS: {_fmt(g.get('C0_LOSS_GROSS_LOSS'))}",
        "",
        f"L1_GROSS_LOSS_SHARE: {_fmt(g.get('L1_GROSS_LOSS_SHARE'))}",
        "",
        f"L2_GROSS_LOSS_SHARE: {_fmt(g.get('L2_GROSS_LOSS_SHARE'))}",
        "",
        f"L3_GROSS_LOSS_SHARE: {_fmt(g.get('L3_GROSS_LOSS_SHARE'))}",
        "",
        f"L4_GROSS_LOSS_SHARE: {_fmt(g.get('L4_GROSS_LOSS_SHARE'))}",
        "",
        f"C0_MEDIAN_MFE: {_fmt(g.get('C0_MEDIAN_MFE'))}",
        "",
        f"C0_MEDIAN_MAE: {_fmt(g.get('C0_MEDIAN_MAE'))}",
        "",
        f"C0_MEDIAN_GIVEBACK: {_fmt(g.get('C0_MEDIAN_GIVEBACK'))}",
        "",
        f"C0_WIN_MEDIAN_CAPTURE_RATIO: {_fmt(g.get('C0_WIN_MEDIAN_CAPTURE_RATIO'))}",
        "",
        f"CURRENT_PROFIT_TO_LOSS_RATE: {_fmt(g.get('CURRENT_PROFIT_TO_LOSS_RATE'))}",
        "",
        f"C0_PROFIT_TO_LOSS_RATE: {_fmt(g.get('C0_PROFIT_TO_LOSS_RATE'))}",
        "",
        f"PRIMARY_LOSS_MECHANISM: {_fmt(g.get('PRIMARY_LOSS_MECHANISM'))}",
        "",
        f"EXIT_CONTRIBUTION_SUPPORTED: {_tf(g.get('EXIT_CONTRIBUTION_SUPPORTED'))}",
        "",
        f"C0_SPEC_SHA256_MATCH: {_tf(g.get('C0_SPEC_SHA256_MATCH'))}",
        "",
        f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
        "",
        f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
        "",
        f"VERDICT: {g.get('VERDICT')}",
        "",
        f"NEXT: {g.get('NEXT')}",
        "",
        "## PRIMARY QUESTIONS",
        "",
        f"Q1: {q.get('Q1')}",
        "",
        f"Q2: {q.get('Q2')}",
        "",
        f"Q3: {q.get('Q3')}",
        "",
        f"Q4: {q.get('Q4')}",
        "",
        f"Q5: {q.get('Q5')}",
        "",
        f"Q6: {q.get('Q6')}",
        "",
        f"Q7: {q.get('Q7')}",
        "",
        f"Q8: {q.get('Q8')}",
        "",
        "## STOP",
        "",
        "Diagnostic complete. C0 remains FROZEN_FOR_FUTURE_OOS_ONLY. No EXIT policy in this run.",
        "submit/cancel/live=0/0/0.",
        "",
    ]
    return "\n".join(lines)
