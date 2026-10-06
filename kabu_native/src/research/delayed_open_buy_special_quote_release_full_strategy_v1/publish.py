"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.delayed_open_buy_special_quote_release_full_strategy_v1 import ANALYSIS_ID
from research.delayed_open_buy_special_quote_release_full_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "data_boundary",
    "already_executed",
    "canary",
    "state_semantics",
    "strategy_precommit",
    "opening_episodes",
    "release_events",
    "accept_bars",
    "signals",
    "execution",
    "portfolio",
    "exits",
    "coverage",
    "economics",
    "daily",
    "symbols",
    "causal_ex_top",
    "blocks",
    "diagnostics",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ev = dict(report.get("selected_eval") or {})
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "data_boundary": _kv(dict(report.get("data_boundary") or {})),
        "already_executed": _kv(dict(report.get("already_executed") or {})),
        "canary": _kv(dict(report.get("canary") or {})),
        "state_semantics": _kv(dict(report.get("state_semantics") or {})),
        "strategy_precommit": _kv(dict(report.get("strategy_precommit") or {})),
        "opening_episodes": list(report.get("opening_episodes") or []),
        "release_events": list(report.get("release_events") or []),
        "accept_bars": list(report.get("accept_bars") or []),
        "signals": list(report.get("signal_rows") or []),
        "execution": _kv(dict(report.get("execution") or {})),
        "portfolio": _kv(dict(report.get("portfolio") or {})),
        "exits": _kv(dict(report.get("exits") or {})),
        "coverage": _kv({k: ev.get(k) for k in ("C1", "C2", "C3", "C4", "coverage_ok", "trade_n", "fill_day_n", "trades_per_day")}),
        "economics": _kv(ev) if ev else [{"empty": True}],
        "daily": list(report.get("daily_rows") or []),
        "symbols": list(report.get("symbol_rows") or []),
        "causal_ex_top": _kv(
            {
                "CAUSAL_EX_TOP1_PNL": ev.get("CAUSAL_EX_TOP1_PNL"),
                "CAUSAL_EX_TOP1_PF": ev.get("CAUSAL_EX_TOP1_PF"),
                "G6": (ev.get("g_table") or {}).get("G6"),
            }
        ),
        "blocks": list((ev.get("blocks") or {}).get("blocks") or []),
        "diagnostics": _kv(dict(report.get("diagnostics") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"LOGIC_COMPLETE: {d.get('LOGIC_COMPLETE')}",
            f"OPENING_CURRENT_DATA_LINE_STATUS: {d.get('OPENING_CURRENT_DATA_LINE_STATUS')}",
            "TRUE_OOS: false",
            "CERTIFIED: false",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- parent pinned: {a.get('2_parent_verdict_pinned')}",
            f"- duplicate: {a.get('7_exact_Complete_Strategy_duplicate')} semantic={a.get('8_semantic_duplicate')}",
            f"- canary PASS: {a.get('13_canary_PASS')}",
            f"- BUY direction proven: {a.get('18_BUY_direction_proven')}",
            f"- Coverage: {a.get('37_Coverage_PASS')}",
            f"- LOGIC_COMPLETE: {a.get('68_LOGIC_COMPLETE')}",
            "",
            "No MBO. No futures. No Sizing. No Holdout/Stress/future. No V2 retune.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
