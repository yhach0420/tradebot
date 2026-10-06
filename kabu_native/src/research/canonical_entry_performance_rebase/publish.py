"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.canonical_entry_performance_rebase import ANALYSIS_ID, C3_MODEL_STATUS

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "canonical_entry_performance_rebase"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Populations",
    "TargetIntegrity",
    "MatchedRanking",
    "PairedDays",
    "ActualExec",
    "MatchedExec",
    "Coverage",
    "ExitClasses",
    "FirstEntry",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    frozen = report.get("frozen") or {}
    return "\n".join(
        [
            "# CANONICAL ENTRY PERFORMANCE REBASE V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            f"C3_MODEL_STATUS: {C3_MODEL_STATUS}",
            "No new model. No C3 refit. No C4. No execution-aware objective.",
            "Exact Dual-Lane PnL is a PARITY ANCHOR, not a new SoT.",
            "ENTRY origin = PENDING.anchor. Rank then join TARGET V4. NO TARGET BACKFILL.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A0_PARITY: {_tf(g.get('A0_PARITY'))}",
            f"A0_CANONICAL: {g.get('A0_CANONICAL')}",
            "",
            f"A2_PARITY: {_tf(g.get('A2_PARITY'))}",
            f"A2_CANONICAL: {g.get('A2_CANONICAL')}",
            "",
            f"C3_FINAL_PARITY: {_tf(g.get('C3_FINAL_PARITY'))}",
            f"C3_FINAL_CANONICAL: {g.get('C3_FINAL_CANONICAL')}",
            "",
            f"C3_OOF_EXACT_PARITY: {_tf(g.get('C3_OOF_EXACT_PARITY'))}",
            f"C3_OOF_EXACT_CANONICAL: {g.get('C3_OOF_EXACT_CANONICAL')}",
            "",
            f"MATCHED_RANKING_ROWS: {g.get('MATCHED_RANKING_ROWS')}",
            f"CURRENT_ACTUAL_SCORABLE_ROWS: {g.get('CURRENT_ACTUAL_SCORABLE_ROWS')}",
            f"C3_ACTUAL_SCORABLE_ROWS: {g.get('C3_ACTUAL_SCORABLE_ROWS')}",
            f"C3_COVERAGE_LOSS_ROWS: {g.get('C3_COVERAGE_LOSS_ROWS')}",
            "",
            f"CURRENT_MATCHED_TOP1_UPLIFT: {g.get('CURRENT_MATCHED_TOP1_UPLIFT')}",
            f"CURRENT_MATCHED_TOP3_UPLIFT: {g.get('CURRENT_MATCHED_TOP3_UPLIFT')}",
            f"CURRENT_MATCHED_TOP5_UPLIFT: {g.get('CURRENT_MATCHED_TOP5_UPLIFT')}",
            "",
            f"C3_OOF_MATCHED_TOP1_UPLIFT: {g.get('C3_OOF_MATCHED_TOP1_UPLIFT')}",
            f"C3_OOF_MATCHED_TOP3_UPLIFT: {g.get('C3_OOF_MATCHED_TOP3_UPLIFT')}",
            f"C3_OOF_MATCHED_TOP5_UPLIFT: {g.get('C3_OOF_MATCHED_TOP5_UPLIFT')}",
            "",
            f"C3_OOF_TOP3_DELTA_VS_CURRENT: {g.get('C3_OOF_TOP3_DELTA_VS_CURRENT')}",
            "",
            f"TOP3_DELTA_POSITIVE_DAYS: {g.get('TOP3_DELTA_POSITIVE_DAYS')}",
            f"TOP3_DELTA_NEGATIVE_DAYS: {g.get('TOP3_DELTA_NEGATIVE_DAYS')}",
            "",
            f"CURRENT_ACTUAL_TOP3_WOULD_FILL: {g.get('CURRENT_ACTUAL_TOP3_WOULD_FILL')}",
            f"C3_OOF_ACTUAL_TOP3_WOULD_FILL: {g.get('C3_OOF_ACTUAL_TOP3_WOULD_FILL')}",
            f"C3_FINAL_ACTUAL_TOP3_WOULD_FILL: {g.get('C3_FINAL_ACTUAL_TOP3_WOULD_FILL')}",
            "",
            f"CURRENT_MATCHED_TOP3_FILL_TO_600: {g.get('CURRENT_MATCHED_TOP3_FILL_TO_600')}",
            f"C3_OOF_MATCHED_TOP3_FILL_TO_600: {g.get('C3_OOF_MATCHED_TOP3_FILL_TO_600')}",
            f"C3_FINAL_MATCHED_TOP3_FILL_TO_600: {g.get('C3_FINAL_MATCHED_TOP3_FILL_TO_600')}",
            "",
            f"C3_COVERAGE_SELECTION_DISPLACEMENT_N: {g.get('C3_COVERAGE_SELECTION_DISPLACEMENT_N')}",
            "",
            f"C3_EXIT_CLASS_A_N: {g.get('C3_EXIT_CLASS_A_N')}",
            f"C3_EXIT_CLASS_B_N: {g.get('C3_EXIT_CLASS_B_N')}",
            f"C3_EXIT_CLASS_C_N: {g.get('C3_EXIT_CLASS_C_N')}",
            f"C3_EXIT_CLASS_D_N: {g.get('C3_EXIT_CLASS_D_N')}",
            "",
            f"OLD_C3_INTERPRETATION: {g.get('OLD_C3_INTERPRETATION')}",
            f"CANONICAL_ENTRY_BASELINE_ESTABLISHED: {_tf(g.get('CANONICAL_ENTRY_BASELINE_ESTABLISHED'))}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## CONTRACT",
            "",
            "A0 = CURRENT ENTRY + CURRENT eligibility. No new executable-at-decision gate.",
            "A2 = CURRENT ENTRY + Exact executable-at-decision.",
            f"C3 results are {C3_MODEL_STATUS}.",
            "MATCHED_RANKING_POPULATION = Exact executable AND CURRENT score AND C3 OOF score.",
            "ACTUAL_STRATEGY_POPULATION uses each strategy's own score availability.",
            "Rank by score (symbol ASC tie-break), then join TARGET V4. No target backfill.",
            "",
            "## HISTORICAL VERDICTS MAINTAINED",
            "",
            f"C3: {frozen.get('C3')}",
            f"C3_AUDIT: {frozen.get('C3_AUDIT')}",
            f"C3_RECON: {frozen.get('C3_RECON')}",
            f"CONTRACT: {frozen.get('CONTRACT')}",
            "Future ENTRY research decision source = this Canonical Rebase, not old Panel ranking.",
            "",
            "## DECISION",
            "",
            f"VERDICT={g.get('VERDICT')}",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "C3 retrain is not started this run. C4 forbidden. Execution-aware objective forbidden.",
            "",
            "## STOP",
            "",
            "No C3 retrain. No C4. Runtime/C14 unchanged. CLOCK/EXIT/CAP/Fill unchanged.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
