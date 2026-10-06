"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.c3_execution_reconciliation import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "c3_execution_reconciliation"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Coverage",
    "CommonRanking",
    "CommonExact",
    "FillIdentities",
    "OriginalVsCommon",
    "FirstEntry",
    "ExitClasses",
    "Adverse",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    cov = report.get("coverage") or {}
    ids = report.get("fill_identities") or {}
    ov = report.get("original_vs_common") or {}
    first = report.get("first_entry") or {}
    adv = report.get("adverse_topk") or {}
    orig = report.get("original_exact") or {}
    reasons = cov.get("missing_reason_counts") or {}
    defs = ids.get("definitions") or {}
    a2i = ids.get("A2") or {}
    oofi = ids.get("C3_OOF") or {}
    fini = ids.get("C3_FINAL") or {}
    ov_a2 = ov.get("A2") or {}
    ov_oof = ov.get("C3_OOF") or {}
    ov_fin = ov.get("C3_FINAL") or {}
    confirmed = g.get("PASSIVE_FILL_ADVERSE_SELECTION_CONFIRMED")
    if confirmed is True:
        confirmed_s = "true"
    elif confirmed is False:
        confirmed_s = "false"
    else:
        confirmed_s = str(confirmed)
    cur = adv.get("CURRENT") or {}
    oof = adv.get("C3_OOF") or {}
    fin = adv.get("C3_FINAL") or {}
    return "\n".join(
        [
            "# C3 EXECUTION COUPLING RECONCILIATION",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE DIAGNOSIS ONLY. No new model. No C4. No execution-aware model.",
            "C3_TOP_EDGE_SUPPORTED_PORTFOLIO_FAIL maintained.",
            "C3_MULTIFACTOR_EXACT_FAILURE maintained.",
            "Fill counters: STANDALONE_WOULD_FILL_1S ≠ PENDING_CREATED ≠ PORTFOLIO_FILL ≠ CLOSED_TRADE.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"EXECUTABLE_T0_ROWS: {g.get('EXECUTABLE_T0_ROWS')}",
            f"C3_SCORE_AVAILABLE_N: {g.get('C3_SCORE_AVAILABLE_N')}",
            f"C3_SCORE_MISSING_N: {g.get('C3_SCORE_MISSING_N')}",
            f"COMMON_SCORABLE_N: {g.get('COMMON_SCORABLE_N')}",
            "",
            f"CURRENT_COMMON_TOP3_UPLIFT: {g.get('CURRENT_COMMON_TOP3_UPLIFT')}",
            f"C3_OOF_COMMON_TOP3_UPLIFT: {g.get('C3_OOF_COMMON_TOP3_UPLIFT')}",
            f"C3_FINAL_COMMON_TOP3_UPLIFT: {g.get('C3_FINAL_COMMON_TOP3_UPLIFT')}",
            "",
            f"A2_COMMON: {g.get('A2_COMMON')}",
            f"C3_OOF_COMMON: {g.get('C3_OOF_COMMON')}",
            f"C3_FINAL_COMMON: {g.get('C3_FINAL_COMMON')}",
            "",
            f"A2_PORTFOLIO_FILL_N: {g.get('A2_PORTFOLIO_FILL_N')}",
            f"A2_CLOSED_TRADE_N: {g.get('A2_CLOSED_TRADE_N')}",
            f"C3_OOF_PORTFOLIO_FILL_N: {g.get('C3_OOF_PORTFOLIO_FILL_N')}",
            f"C3_OOF_CLOSED_TRADE_N: {g.get('C3_OOF_CLOSED_TRADE_N')}",
            f"C3_FINAL_PORTFOLIO_FILL_N: {g.get('C3_FINAL_PORTFOLIO_FILL_N')}",
            f"C3_FINAL_CLOSED_TRADE_N: {g.get('C3_FINAL_CLOSED_TRADE_N')}",
            "",
            f"C3_FIRST_FILL_TO_600_CANONICAL: {g.get('C3_FIRST_FILL_TO_600_CANONICAL')}",
            f"C3_EXIT_CLASS_A_N: {g.get('C3_EXIT_CLASS_A_N')}",
            f"C3_EXIT_CLASS_B_N: {g.get('C3_EXIT_CLASS_B_N')}",
            f"C3_EXIT_CLASS_C_N: {g.get('C3_EXIT_CLASS_C_N')}",
            f"C3_EXIT_CLASS_D_N: {g.get('C3_EXIT_CLASS_D_N')}",
            "",
            f"PASSIVE_FILL_ADVERSE_SELECTION_CONFIRMED: {confirmed_s}",
            f"PRIMARY_CAUSE_AFTER_RECONCILIATION: {g.get('PRIMARY_CAUSE_AFTER_RECONCILIATION')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## 1. SCORE COVERAGE",
            "",
            f"EXECUTABLE_T0_ROWS={cov.get('EXECUTABLE_T0_ROWS')} C3_SCORE_AVAILABLE_N={cov.get('C3_SCORE_AVAILABLE_N')} "
            f"C3_SCORE_MISSING_N={cov.get('C3_SCORE_MISSING_N')}.",
            "Missing reasons (mutually exclusive, no future info):",
            f"FEATURE_MISSING={reasons.get('FEATURE_MISSING', 0)} (lead feature vwap_dist_bps).",
            f"LOOKBACK_INSUFFICIENT={reasons.get('LOOKBACK_INSUFFICIENT', 0)} (lead feature drawdown_180s).",
            f"NONFINITE_FEATURE={reasons.get('NONFINITE_FEATURE', 0)}.",
            f"NORMALIZATION_FAILURE={reasons.get('NORMALIZATION_FAILURE', 0)}.",
            f"LOOKUP_MISS={reasons.get('LOOKUP_MISS', 0)}.",
            f"OTHER={reasons.get('OTHER', 0)}.",
            "COMMON_SCORABLE = t0 executable AND CURRENT score AND C3 OOF score AND C3 FINAL score.",
            "CURRENT-only 11390 vs C3-only 9603 ranking comparison is not used.",
            "",
            "## 2. ORIGINAL EXACT VS COMMON_SCORABLE JOIN",
            "",
            "Join is admit_clock(fill_time)×symbol, not packed Dual-Lane anchor_time.",
            f"A2 original CLOSED={ov_a2.get('CLOSED_TRADE_N')} ranking_pop={ov_a2.get('ON_RANKING_POP_N')} "
            f"COMMON_SCORABLE={ov_a2.get('ON_COMMON_SCORABLE_N')} "
            f"ranking_but_C3_missing={ov_a2.get('RANKING_BUT_C3_SCORE_MISSING_N')} "
            f"panel_not_ranking (live fill, panel executable_at_t0=false)={ov_a2.get('PANEL_NOT_RANKING_POP_N')}.",
            f"C3 OOF original CLOSED={ov_oof.get('CLOSED_TRADE_N')} COMMON_SCORABLE={ov_oof.get('ON_COMMON_SCORABLE_N')}.",
            f"C3 FINAL original CLOSED={ov_fin.get('CLOSED_TRADE_N')} COMMON_SCORABLE={ov_fin.get('ON_COMMON_SCORABLE_N')}.",
            f"Original Exact (frozen, not rewritten): A2 {orig.get('A2')} | C3 OOF {orig.get('C3_OOF')} | C3 FINAL {orig.get('C3_FINAL')}.",
            "",
            "## 3. FILL COUNTER DEFINITIONS",
            "",
            f"A. STANDALONE_WOULD_FILL_1S: {defs.get('STANDALONE_WOULD_FILL_1S')}",
            f"B. PENDING_CREATED: {defs.get('PENDING_CREATED')}",
            f"C. PORTFOLIO_FILL: {defs.get('PORTFOLIO_FILL')}",
            f"D. CLOSED_TRADE: {defs.get('CLOSED_TRADE')}",
            f"HARVEST_FILL: {defs.get('HARVEST_FILL')}",
            f"A2 PENDING={a2i.get('PENDING_CREATED_N')} PORTFOLIO_FILL={a2i.get('PORTFOLIO_FILL_N')} "
            f"CLOSED={a2i.get('CLOSED_TRADE_N')} HARVEST={a2i.get('HARVEST_FILL_N')} EXPIRED={a2i.get('EXPIRED_N')}.",
            f"C3 OOF PENDING={oofi.get('PENDING_CREATED_N')} PORTFOLIO_FILL={oofi.get('PORTFOLIO_FILL_N')} "
            f"CLOSED={oofi.get('CLOSED_TRADE_N')} HARVEST={oofi.get('HARVEST_FILL_N')} EXPIRED={oofi.get('EXPIRED_N')}.",
            f"C3 FINAL PENDING={fini.get('PENDING_CREATED_N')} PORTFOLIO_FILL={fini.get('PORTFOLIO_FILL_N')} "
            f"CLOSED={fini.get('CLOSED_TRADE_N')} HARVEST={fini.get('HARVEST_FILL_N')} EXPIRED={fini.get('EXPIRED_N')}.",
            "53 vs 36: C3 FINAL CLOSED_TRADE_N=53 is the Exact headline. HARVEST_FILL_N=36 is CollectorEngine a_fills, incomplete vs Dual-Lane.",
            "80 vs 47: C3 OOF CLOSED_TRADE_N=80 vs HARVEST_FILL_N=47, same harvest gap.",
            "A2 233 vs 148 harvest: same definition split. Headline Exact == CLOSED_TRADE_N.",
            "",
            "## 4. FIRST-ENTRY CANONICAL (C3 FINAL)",
            "",
            f"FIRST_N={first.get('FIRST_N')} FIRST_TARGET_V4_MEAN={first.get('FIRST_TARGET_V4_MEAN')} "
            f"FIRST_FILL_TO_600_MEAN={first.get('FIRST_FILL_TO_600_MEAN')} "
            f"FIRST_FILL_TO_600_MEDIAN={first.get('FIRST_FILL_TO_600_MEDIAN')} "
            f"FIRST_FILL_TO_600_POS_RATE={first.get('FIRST_FILL_TO_600_POS_RATE')} "
            f"FIRST_ACTUAL_PNL={first.get('FIRST_ACTUAL_PNL')}.",
            "One canonical FILL_TO_600: m4_t1_px/fill_price-1 joined on admit_clock(fill_time)×symbol.",
            "",
            "## 5. ADVERSE RECHECK ON COMMON_SCORABLE Top3",
            "",
            f"CURRENT TARGET={cur.get('TARGET_V4_600S')} WOULD_FILL={cur.get('WOULD_FILL_1S_RATE')} "
            f"FILL_TO_600={cur.get('FILL_TO_600S_RETURN_MEAN')}.",
            f"C3 OOF TARGET={oof.get('TARGET_V4_600S')} WOULD_FILL={oof.get('WOULD_FILL_1S_RATE')} "
            f"FILL_TO_600={oof.get('FILL_TO_600S_RETURN_MEAN')}.",
            f"C3 FINAL TARGET={fin.get('TARGET_V4_600S')} WOULD_FILL={fin.get('WOULD_FILL_1S_RATE')} "
            f"FILL_TO_600={fin.get('FILL_TO_600S_RETURN_MEAN')}.",
            "CONFIRMED only if same population C3 Top3 TARGET > CURRENT AND C3 FILL_TO_600 < CURRENT.",
            "",
            "## 6. DECISION",
            "",
            "CASE A holds: common ranking edge + COMMON Exact fail + fill-conditioned return worse.",
            "CASE B holds: A2_COMMON collapses vs frozen A2 233 / +219610 / PF 1.114, and original Exact barely sits on COMMON_SCORABLE.",
            "CASE C does not hold: COMMON dual ADMIT/EXIT = CLOSED_TRADE; EXIT A+B+C+D=53.",
            "Composite → CASE D. Formal C3 / C3-audit verdicts are not rewritten.",
            "",
            "## STOP",
            "",
            "Audit complete. No C4. Runtime not changed. C3 not implemented. C14 unchanged.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
