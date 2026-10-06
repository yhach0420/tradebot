"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.true_cme_nq_es_cost_and_coverage_audit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Current_Driver_Map",
    "Tech_Unresolved_13",
    "Databento_Access",
    "Databento_Cost",
    "CME_Semantics",
    "Contract_Roll_Design",
    "KRX_Comparison",
    "TWSE_Comparison",
    "Paid_Data_Decision",
    "Deferred_Free_Drivers",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    p = dict(report.get("source_probe") or {})
    cmp = list(report.get("comparison") or [])
    krx = next((r for r in cmp if r.get("source") == "KRX_KOSPI200_SAMSUNG_SKHYNIX"), {})
    tw = next((r for r in cmp if r.get("source") == "TWSE_NATIVE_SEMICONDUCTOR"), {})
    bind = dict(report.get("bind") or {})
    return {
        "Binding": _kv(
            {
                "parent_accepted": report.get("parent_verdict_accepted"),
                "hkg_chi_closed_as": report.get("hkg_chi_closed_as"),
                "do_not_retune": True,
                "korea_not_tested": True,
                "nq_es_proxy_not_true_cme": True,
                "pool_n": bind.get("research_pool_n"),
                "playbooks_built": False,
                "purchase": False,
                "kabu_50": False,
            }
        ),
        "Current_Driver_Map": _kv(dict(report.get("current_driver_map") or {})),
        "Tech_Unresolved_13": list(report.get("tech_unresolved_13") or [{"empty": True}]),
        "Databento_Access": _kv(
            {
                "API_KEY_PRESENT": p.get("API_KEY_PRESENT"),
                "API_KEY_SOURCE_CLASS": p.get("API_KEY_SOURCE_CLASS"),
                "LOCAL_TRUE_CME_HISTORY_PRESENT": p.get("LOCAL_TRUE_CME_HISTORY_PRESENT"),
                "status": p.get("status"),
                "authenticated_metadata_called": p.get("authenticated_metadata_called"),
                "timeseries_called": p.get("timeseries_called"),
                "anonymous_get_cost_http": (p.get("anonymous_get_cost") or {}).get("status"),
                "catalog_nq_http": (p.get("catalog") or {}).get("nq_http"),
                "catalog_es_http": (p.get("catalog") or {}).get("es_http"),
                "local": p.get("local"),
            }
        ),
        "Databento_Cost": _kv(
            {
                "exact_nq_usd": p.get("exact_nq_usd"),
                "exact_es_usd": p.get("exact_es_usd"),
                "exact_combined_usd": p.get("exact_combined_usd"),
                "nq_parent_cost": p.get("nq_parent_cost"),
                "es_parent_cost": p.get("es_parent_cost"),
                "combined_parent_cost": p.get("combined_parent_cost"),
                "pricing_model": p.get("pricing"),
                "size": p.get("size"),
                "unit_prices": p.get("unit_prices"),
                "dataset_range": p.get("dataset_range"),
            }
        ),
        "CME_Semantics": _kv(dict(report.get("cme_semantics") or {})),
        "Contract_Roll_Design": _kv(dict(report.get("contract_roll_design") or {})),
        "KRX_Comparison": _kv(krx),
        "TWSE_Comparison": _kv(tw),
        "Paid_Data_Decision": _kv(dict(report.get("paid_data_decision") or {})),
        "Deferred_Free_Drivers": list(report.get("deferred_free_drivers") or [{"empty": True}]),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    p = dict(report.get("source_probe") or {})
    paid = dict(report.get("paid_data_decision") or {})
    lines = [
        "# TRUE_CME_NQ_ES_COST_AND_COVERAGE_AUDIT_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted parent: `HKG_CHINA_A50_SIMULTANEOUS_ONLY_V1`.",
        "HKG / CHI closed as: `NO_USABLE_CAUSAL_LEAD`. Do not retune lag, threshold, clock, or symbol/sector subset.",
        "USATECH.IDX-USD / USA500.IDX-USD remain `PROXY_NOT_CME_FUTURES`. That simultaneous tech result does not prove true NQ/ES lack lead.",
        "Korea remains not tested (source unavailable without purchase). Taiwan native remains unpaid/not entitled. TSM.US ADR is not Taiwan-session causal proof.",
        "This phase is source/cost audit only. No purchase. No paid download. No account creation. No entry playbooks.",
        "",
        f"Existing Databento API key? **{a.get('Existing_Databento_API_key')}**",
        f"Existing true CME files? **{a.get('Existing_true_CME_files')}**",
        f"True NQ 1m available? `{a.get('True_NQ_1m_available')}`",
        f"True ES 1m available? `{a.get('True_ES_1m_available')}`",
        f"Historical range sufficient? **{a.get('Historical_range_sufficient')}** (catalog since 2010-06-06 covers Discovery 2024-09-16 through 2025-11-26)",
        f"OHLCV-1m semantics proven from local bars? **{a.get('OHLCV_1m_semantics_proven')}** (documented: ts_event is interval start)",
        "",
        f"Exact estimated NQ cost? `{a.get('Exact_estimated_NQ_cost')}`",
        f"Exact estimated ES cost? `{a.get('Exact_estimated_ES_cost')}`",
        f"Combined cost? `{a.get('Combined_cost')}`",
        f"Authenticated metadata called? **{p.get('authenticated_metadata_called')}**  timeseries.get_range called? **False**",
        "",
        f"Any purchase? **{a.get('Any_purchase')}**",
        f"Any paid download? **{a.get('Any_paid_download')}**",
        f"Account creation performed? **{a.get('Account_creation_performed')}**",
        "",
        f"Can true CME directly resolve a material uncertainty left by the US proxy study? **{a.get('Can_true_CME_directly_resolve_material_US_proxy_uncertainty')}**",
        f"Expected information gain vs KRX: {a.get('Expected_information_gain_vs_KRX')}",
        f"Expected information gain vs TWSE: {a.get('Expected_information_gain_vs_TWSE')}",
        f"Cheapest high-information next dataset: **{a.get('Cheapest_high_information_next_dataset')}**",
        f"Paid-data acquisition justified enough to ask the user for approval? **{a.get('Paid_data_acquisition_justified_enough_to_ask_user')}**",
        f"If not approved, free driver next: **{a.get('If_not_approved_free_driver_next')}** (oil/commodity for resources / trading houses).",
        "",
        f"Tech 13 usable causal driver: **{a.get('tech_13_usable_causal_driver_n')}** / 13",
        f"Recommend true CME first this phase? **{paid.get('recommend_true_cme_first')}**",
        "",
        f"ENTRY playbook built? **{a.get('ENTRY_playbook_built')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Kabu 50? **{a.get('Kabu_50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "Do not buy anything. Do not auto-proceed to oil in this phase.",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
