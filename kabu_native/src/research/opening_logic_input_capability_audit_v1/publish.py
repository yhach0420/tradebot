"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.opening_logic_input_capability_audit_v1 import ANALYSIS_ID
from research.opening_logic_input_capability_audit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "data_boundary",
    "time_coverage",
    "preopen_fields",
    "timestamp_semantics",
    "auction_indicative",
    "symbol_trajectory",
    "actual_open",
    "prior_session",
    "futures_index",
    "cross_sectional_context",
    "gap_capability",
    "opening_objects",
    "confirmed_input_set",
    "missing_inputs",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    cov = dict(report.get("time_coverage") or {})
    ts = dict(report.get("timestamp_semantics") or {})
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "data_boundary": _kv(dict(report.get("data_boundary") or {})),
        "time_coverage": list(cov.get("EVENT_N_BY_INTERVAL_BY_DAY") or []) or _kv(cov),
        "preopen_fields": list(report.get("preopen_fields") or []),
        "timestamp_semantics": _kv(ts),
        "auction_indicative": _kv(dict(report.get("auction_indicative") or {})),
        "symbol_trajectory": _kv(dict(report.get("symbol_trajectory") or {})),
        "actual_open": _kv(dict(report.get("actual_open") or {})),
        "prior_session": _kv(dict(report.get("prior_session") or {})),
        "futures_index": list(report.get("futures_index") or []),
        "cross_sectional_context": _kv(dict(report.get("cross_sectional_context") or {})),
        "gap_capability": _kv(dict(report.get("gap_capability") or {})),
        "opening_objects": list(report.get("opening_objects") or []),
        "confirmed_input_set": list(report.get("confirmed_input_set") or []),
        "missing_inputs": list(report.get("missing_inputs") or []),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            "TRUE_OOS: false",
            "CERTIFIED: false",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- preopen day N: {a.get('8_preopen_Capture_day_N')}",
            f"- trajectory: {a.get('21_symbol_preopen_trajectory_reconstructable')}",
            f"- actual open: {a.get('22_actual_open_price_causal')}",
            f"- XS pressure: {a.get('31_cross_sectional_preopen_pressure_feasible')}",
            f"- SYMBOL_EXPECTED_GAP: {a.get('35_SYMBOL_EXPECTED_GAP_status')}",
            f"- MARKET_EXPECTED_GAP: {a.get('36_MARKET_EXPECTED_GAP_status')}",
            f"- RELATIVE_DISLOCATION: {a.get('37_RELATIVE_DISLOCATION_status')}",
            "",
            "No strategy. No thresholds. No PnL. No MBO. No Holdout/Stress/future.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
