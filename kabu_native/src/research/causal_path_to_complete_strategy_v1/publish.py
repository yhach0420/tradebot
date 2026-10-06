"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.causal_path_to_complete_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Data_Binding",
    "Discovery_Internal_Split",
    "Causal_Event_Universe",
    "Event_Primitives",
    "Forward_Path",
    "Path_Taxonomy",
    "Pullback_RCA",
    "Breakout_RCA",
    "Other_Event_RCA",
    "Market_State",
    "Sector_State",
    "Stock_State",
    "Continuous_Effects",
    "State_Transitions",
    "Diagnostic_Interactions",
    "Candidate_1",
    "Candidate_2",
    "Candidate_3",
    "Complete_Strategy_Economics",
    "Secondary_Confirmation",
    "Frozen_Validation_Status",
    "Live_20260914_Status",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _cand_sheet(c: dict[str, Any] | None, idx: int) -> list[dict[str, Any]]:
    if not c:
        return [{"candidate": idx, "status": "none"}]
    spec = dict(c.get("spec") or {})
    econ = dict(c.get("discovery_economics") or {})
    promo = dict(c.get("promotion") or {})
    return _kv(
        {
            "candidate_id": spec.get("candidate_id"),
            "spec_sha256": spec.get("spec_sha256"),
            "archetype": spec.get("archetype"),
            "thesis": spec.get("thesis"),
            "event_family": spec.get("event_family"),
            "ENTRY_setup": spec.get("setup"),
            "ENTRY_trigger": spec.get("trigger"),
            "entry_availability_time": spec.get("entry_availability_time"),
            "cuts": spec.get("cuts"),
            "EXIT": spec.get("technical_invalidation"),
            "exit_state_machine": spec.get("exit_state_machine"),
            "same_symbol": spec.get("same_symbol_behavior"),
            "reentry": spec.get("reentry_behavior"),
            "CAP": spec.get("CAP_semantics"),
            "occupancy": spec.get("occupancy"),
            "session_close": spec.get("session_close"),
            "copied_m4_m6": spec.get("copied_m4_m6"),
            "promoted": promo.get("promoted"),
            "fail_reasons": promo.get("fail_reasons"),
            **{k: econ.get(k) for k in ("trade_n", "day_n", "mean_x0_bps", "mean_x1_bps", "profit_factor", "hit_rate", "block_mean_x0", "top_symbol_share_of_positive_bps", "top_sector_share_of_positive_bps", "max_dd_daily_mean_bps")},
        }
    )


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    bind = dict(report.get("bind") or {})
    split = dict(report.get("split") or {})
    blocks = dict(report.get("blocks") or {})
    eu = dict(report.get("event_universe") or {})
    rca = dict(report.get("rca") or {})
    cand = dict(report.get("candidates") or {})
    promoted = list(cand.get("promoted") or [])
    proposals = list(cand.get("proposals") or [])
    sec = dict(report.get("secondary_confirmation") or {})
    live = dict(report.get("live_20260914") or {})
    fv = dict(report.get("frozen_validation") or {})
    other = dict(rca.get("other") or {})
    diags = list(rca.get("diagnostics") or [])
    c1 = promoted[0] if len(promoted) > 0 else (proposals[0] if len(proposals) > 0 else None)
    c2 = promoted[1] if len(promoted) > 1 else (proposals[1] if len(proposals) > 1 else None)
    c3 = promoted[2] if len(promoted) > 2 else (proposals[2] if len(proposals) > 2 else None)
    return {
        "Data_Binding": _kv(
            {
                "parent_verdict_accepted": report.get("parent_verdict_accepted"),
                "split_sha256": split.get("split_sha256"),
                "expected_split_sha256": report.get("expected_split_sha256"),
                "sha_match": split.get("sha_match"),
                "union_n": (bind.get("foundation") or {}).get("union_n") or bind.get("foundation", {}).get("union_n"),
                "m1_m11_rescued": report.get("m1_m11_rescued"),
                "old_confirmation_used_to_design": report.get("old_confirmation_used_to_design"),
                "kabu_50_applied": report.get("kabu_50_applied"),
                "bind_ok": bind.get("ok"),
                "bind_reason": bind.get("reason"),
            }
        )
        + [{"key": "rejected_m", "value": m} for m in list(report.get("m1_m11_rejected") or [])],
        "Discovery_Internal_Split": _kv({k: blocks.get(k) for k in ("ok", "n_blocks", "block_sha256", "random_split", "frozen_before_candidate_evaluation")})
        + [
            {"block_id": b.get("block_id"), "first": b.get("first"), "last": b.get("last"), "n": b.get("n")}
            for b in list(blocks.get("blocks") or [])
        ]
        + [{"block_id": b.get("block_id"), "i": i, "date": d} for b in list(report.get("block_date_lists") or []) for i, d in enumerate(list(b.get("dates") or []), start=1)],
        "Causal_Event_Universe": _kv(
            {
                "event_n": eu.get("event_n"),
                "event_id_sha256": eu.get("event_id_sha256"),
                "identities_frozen_before_outcomes": eu.get("identities_frozen_before_outcomes"),
                "future_used_as_decision_feature": eu.get("future_used_as_decision_feature"),
                "copied_m4_m6_thresholds": eu.get("copied_m4_m6_thresholds"),
            }
        )
        + [{"family": k, "n": v} for k, v in sorted(dict(eu.get("families") or {}).items())]
        + list(eu.get("sample") or []),
        "Event_Primitives": list(report.get("event_primitives") or [{"empty": True}]),
        "Forward_Path": _kv(dict(report.get("forward_path") or {})),
        "Path_Taxonomy": list(report.get("path_taxonomy") or [{"empty": True}]),
        "Pullback_RCA": _kv({k: v for k, v in dict(rca.get("pullback") or {}).items() if k not in {"top_features", "contrast"}})
        + [{"kind": "feature", **f} for f in list((rca.get("pullback") or {}).get("top_features") or [])]
        + [{"kind": "contrast", **f} for f in list((rca.get("pullback") or {}).get("contrast") or [])],
        "Breakout_RCA": _kv({k: v for k, v in dict(rca.get("breakout") or {}).items() if k not in {"top_features", "contrast"}})
        + [{"kind": "feature", **f} for f in list((rca.get("breakout") or {}).get("top_features") or [])]
        + [{"kind": "contrast", **f} for f in list((rca.get("breakout") or {}).get("contrast") or [])],
        "Other_Event_RCA": [{"family": fam, **(inf if isinstance(inf, dict) else {"value": inf})} for fam, inf in other.items()] or [{"empty": True}],
        "Market_State": list(report.get("market_state_sample") or [{"empty": True}]),
        "Sector_State": list(report.get("sector_state_sample") or [{"empty": True}]),
        "Stock_State": list(report.get("stock_state_sample") or [{"empty": True}]),
        "Continuous_Effects": list(report.get("continuous_effects") or [{"empty": True}]),
        "State_Transitions": list(report.get("state_transitions") or [{"empty": True}]),
        "Diagnostic_Interactions": [{"family": d.get("family"), "pairs": d.get("pairs"), "tree": d.get("tree"), "model_is_not_strategy": d.get("model_is_not_strategy")} for d in diags] or [{"empty": True}],
        "Candidate_1": _cand_sheet(c1, 1),
        "Candidate_2": _cand_sheet(c2, 2),
        "Candidate_3": _cand_sheet(c3, 3),
        "Complete_Strategy_Economics": [
            {
                "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                "spec_sha256": (p.get("spec") or {}).get("spec_sha256"),
                "promoted": (p.get("promotion") or {}).get("promoted"),
                "fail_reasons": (p.get("promotion") or {}).get("fail_reasons"),
                **{k: (p.get("discovery_economics") or {}).get(k) for k in ("trade_n", "day_n", "symbol_n", "sector_n", "mean_x0_bps", "mean_x1_bps", "profit_factor", "hit_rate", "daily_positive_share", "max_dd_daily_mean_bps", "top_symbol_share_of_positive_bps", "top_sector_share_of_positive_bps", "block_mean_x0", "block_positive_n")},
            }
            for p in proposals
        ]
        or [{"empty": True}],
        "Secondary_Confirmation": _kv({k: sec.get(k) for k in ("ran", "label", "old_confirmation_used_to_design", "any_passed", "certifies")})
        + list(sec.get("rows") or [{"status": "not_run"}]),
        "Frozen_Validation_Status": _kv(fv),
        "Live_20260914_Status": _kv(
            {
                "classification": live.get("classification"),
                "fed_into_discovery_tuning": live.get("fed_into_discovery_tuning"),
                "did_not_stop_captures": live.get("did_not_stop_captures"),
                "futures_capture_running": live.get("futures_capture_running"),
                "breadth_capture_running": live.get("breadth_capture_running"),
            }
        )
        + _kv({"futures": live.get("futures"), "breadth": live.get("breadth")}),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    eu = dict(report.get("event_universe") or {})
    lines = [
        "# CAUSAL_PATH_TO_COMPLETE_STRATEGY_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "## Bind / safety",
        "",
        f"Parent verdict accepted: `{report.get('parent_verdict_accepted')}`",
        f"Split SHA: `{a.get('VERDICT') and (report.get('split') or {}).get('split_sha256')}`",
        f"M1–M11 rescued? **{a.get('m1_m11_rescued')}**",
        f"Discovery-only design? **{a.get('discovery_only_logic_design')}**",
        f"Old Confirmation used to design? **{a.get('old_confirmation_used_to_design_rules')}**",
        f"Frozen Validation opened? **{a.get('frozen_validation_opened')}**",
        f"Future outcomes as decision features? **{a.get('future_outcomes_used_as_decision_time_features')}**",
        f"Kabu 50 applied? **{a.get('kabu_50_applied')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "## Event universe",
        "",
        f"Built before outcomes attached? **{a.get('causal_event_universe_built_before_outcomes_attached')}**",
        f"Event count: **{eu.get('event_n')}**",
        f"Families: `{json.dumps(eu.get('families') or {}, ensure_ascii=False)}`",
        f"Path classes: **{a.get('path_class_n')}**",
        "",
        "## RCA",
        "",
        f"Level useful: `{a.get('level_useful')}`",
        f"Transition useful: `{a.get('transition_useful')}`",
        f"Acceleration useful: `{a.get('acceleration_useful')}`",
        f"Market→sector→stock chain? **{a.get('stable_market_sector_stock_chain')}**",
        f"Stock-only chain? **{a.get('stable_stock_only_chain')}**",
        "",
        "### Pullback success vs failure (pre-event contrast)",
        "",
        f"`{json.dumps(a.get('what_distinguishes_successful_vs_failed_pullback'), ensure_ascii=False, default=str)[:4000]}`",
        "",
        "### Breakout success vs failure (pre-event contrast)",
        "",
        f"`{json.dumps(a.get('what_distinguishes_successful_vs_failed_breakout'), ensure_ascii=False, default=str)[:4000]}`",
        "",
        "## Complete strategy candidates",
        "",
        f"n={a.get('complete_strategy_candidates_n')}",
        f"Candidate 1: `{json.dumps(a.get('candidate_1'), ensure_ascii=False, default=str)[:2500]}`",
        f"Candidate 2: `{json.dumps(a.get('candidate_2'), ensure_ascii=False, default=str)[:2500]}`",
        f"Candidate 3: `{json.dumps(a.get('candidate_3'), ensure_ascii=False, default=str)[:2500]}`",
        f"Secondary confirmation any pass? **{a.get('any_candidate_passed_SECONDARY_CONFIRMATION_NOT_PRISTINE')}**",
        "",
        "## 20260914 prospective capture",
        "",
        f"Futures running? **{a.get('futures_capture_20260914_running')}**",
        f"Breadth running? **{a.get('breadth_capture_20260914_running')}**",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
