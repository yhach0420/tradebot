"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.hkg_china_a50_driver_response_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Preserved_Driver_Map",
    "HKG_Source",
    "CHI_Source",
    "Timestamp_Semantics",
    "Session_Semantics",
    "Foreign_Open_Event",
    "HKG_Response",
    "CHI_Response",
    "HKG_After_CHI",
    "CHI_After_HKG",
    "Reverse_Causality",
    "Sector_Response",
    "Stock_Raw",
    "Stock_MarketAdjusted",
    "Stock_SectorAdjusted",
    "UpDown_Asymmetry",
    "US_Risk_Comparison",
    "FX_Comparison",
    "D1_D4",
    "Driver_Coverage_105",
    "Tech_Missing_Driver_Data",
    "Unexplained_Groups",
    "Next_Driver",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _rho(rec: Any) -> Any:
    if isinstance(rec, dict):
        return rec.get("rho")
    return rec


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    stocks = list(report.get("stock_response") or [])
    hkg_sec = list(report.get("hkg_sector") or [])
    chi_sec = list(report.get("chi_sector") or [])
    inc = list(report.get("incremental_sector") or [])
    probe = dict(report.get("source_probe") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "korea_not_tested": True,
                "korea_not_economic_null": True,
                "TECH_DRIVER_BLOCKER_OPEN": True,
                "playbooks_built": False,
                "purchase": False,
                "kabu_50": False,
            }
        ),
        "Preserved_Driver_Map": _kv(dict(report.get("preserved_driver_map") or {})),
        "HKG_Source": _kv(dict(probe.get("hkg") or {})) + _kv(dict(report.get("hkg_meta") or {})),
        "CHI_Source": _kv(dict(probe.get("chi") or {})) + _kv(dict(report.get("chi_meta") or {})),
        "Timestamp_Semantics": _kv(dict(report.get("timestamp_semantics") or {})),
        "Session_Semantics": _kv(dict(report.get("session_semantics") or {})),
        "Foreign_Open_Event": list(report.get("foreign_open") or [{"empty": True}]),
        "HKG_Response": [
            {"sector": r.get("sector"), "lead_lag": r.get("lead_lag"), "fwd_rho_residual": r.get("fwd_rho_residual"), "usable": r.get("usable_before_stock_move")}
            for r in hkg_sec
        ]
        or [{"empty": True}],
        "CHI_Response": [
            {"sector": r.get("sector"), "lead_lag": r.get("lead_lag"), "fwd_rho_residual": r.get("fwd_rho_residual"), "usable": r.get("usable_before_stock_move")}
            for r in chi_sec
        ]
        or [{"empty": True}],
        "HKG_After_CHI": [
            {"sector": r.get("sector"), "hkg_after_chi_rho": _rho(r.get("hkg_after_chi_1m")), "hkg_incremental_after_chi": r.get("hkg_incremental_after_chi")}
            for r in inc
        ]
        or [{"empty": True}],
        "CHI_After_HKG": [
            {"sector": r.get("sector"), "chi_after_hkg_rho": _rho(r.get("chi_after_hkg_1m")), "chi_incremental_after_hkg": r.get("chi_incremental_after_hkg")}
            for r in inc
        ]
        or [{"empty": True}],
        "Reverse_Causality": (
            [{"driver": "HKG", "sector": r.get("sector"), "reverse_rho": r.get("reverse_rho"), "lead_lag": r.get("lead_lag")} for r in hkg_sec]
            + [{"driver": "CHI", "sector": r.get("sector"), "reverse_rho": r.get("reverse_rho"), "lead_lag": r.get("lead_lag")} for r in chi_sec]
        )
        or [{"empty": True}],
        "Sector_Response": inc or [{"empty": True}],
        "Stock_Raw": [
            {"symbol": r.get("symbol"), "sector": r.get("sector"), "hkg_raw_rho": _rho(r.get("hkg_raw_1m")), "chi_raw_rho": _rho(r.get("chi_raw_1m")), "class": r.get("class")}
            for r in stocks
        ]
        or [{"empty": True}],
        "Stock_MarketAdjusted": [
            {"symbol": r.get("symbol"), "hkg_mkt_rho": _rho(r.get("hkg_market_adj_1m")), "chi_mkt_rho": _rho(r.get("chi_market_adj_1m")), "class": r.get("class")}
            for r in stocks
        ]
        or [{"empty": True}],
        "Stock_SectorAdjusted": [
            {
                "symbol": r.get("symbol"),
                "hkg_sec_rho": _rho(r.get("hkg_sector_adj_1m")),
                "chi_sec_rho": _rho(r.get("chi_sector_adj_1m")),
                "hkg_after_chi_rho": _rho(r.get("hkg_after_chi_1m")),
                "chi_after_hkg_rho": _rho(r.get("chi_after_hkg_1m")),
                "class": r.get("class"),
                "usable_before_stock_move": r.get("usable_before_stock_move"),
            }
            for r in stocks
        ]
        or [{"empty": True}],
        "UpDown_Asymmetry": [
            {
                "sector": r.get("sector"),
                "hkg_up_mean": r.get("up_mean_hkg"),
                "hkg_down_mean": r.get("down_mean_hkg"),
                "chi_up_mean": r.get("up_mean_chi"),
                "chi_down_mean": r.get("down_mean_chi"),
                "d1_d4_hkg": r.get("d1_d4_hkg"),
                "d1_d4_chi": r.get("d1_d4_chi"),
            }
            for r in inc
        ]
        or [{"empty": True}],
        "US_Risk_Comparison": [
            {
                "symbol": r.get("symbol"),
                "class": r.get("class"),
                "hkg_after_es_stable": r.get("hkg_incremental_after_es"),
                "chi_after_es_stable": r.get("chi_incremental_after_es"),
                "usable": r.get("usable_before_stock_move"),
            }
            for r in stocks
            if r.get("symbol") in ("6779", "2737", "7717", "4461", "4092") or r.get("sector") in ("電気機器", "機械", "非鉄金属")
        ]
        or [{"empty": True}],
        "FX_Comparison": [
            {
                "symbol": r.get("symbol"),
                "class": r.get("class"),
                "hkg_after_fx_stable": r.get("hkg_incremental_after_fx"),
                "chi_after_fx_stable": r.get("chi_incremental_after_fx"),
                "usable": r.get("usable_before_stock_move"),
            }
            for r in stocks
            if r.get("symbol") in ("6787", "7717", "4092") or r.get("sector") in ("輸送用機器", "非鉄金属")
        ]
        or [{"empty": True}],
        "D1_D4": [
            {"symbol": r.get("symbol"), "d1_d4_hkg": r.get("d1_d4_hkg"), "d1_d4_chi": r.get("d1_d4_chi"), "d1_d4_hkg_after_chi": r.get("d1_d4_hkg_after_chi"), "class": r.get("class")}
            for r in stocks
        ]
        or [{"empty": True}],
        "Driver_Coverage_105": list(report.get("coverage_rows") or [{"empty": True}]),
        "Tech_Missing_Driver_Data": list(report.get("tech_missing_driver_data") or [{"empty": True}]),
        "Unexplained_Groups": _kv(dict(report.get("unexplained") or {})),
        "Next_Driver": _kv(dict(report.get("next_driver") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    q = dict(a.get("evidence_quality_breakdown") or {})
    lines = [
        "# HKG_CHINA_A50_DRIVER_RESPONSE_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted parent: `KOREA_INTRADAY_SOURCE_NOT_AVAILABLE_V1`.",
        "Korea was **not tested** (source unavailable without purchase). This is not an economic null.",
        "HKG/CHI are labeled Jetta proxies. They are **not** Korea, Taiwan, or CME substitutes.",
        "TECH_DRIVER_BLOCKER_OPEN remains true. FX and US-risk maps are preserved. No playbooks.",
        "",
        f"HKG source acquired? **{a.get('HKG_source_acquired')}**  CHI A50 acquired? **{a.get('CHI_A50_source_acquired')}**  Free? **{a.get('free')}**  Purchase? **{a.get('purchase')}**",
        f"Timestamp semantics proven? **{a.get('timestamp_semantics_proven')}**  Session from bars? **{a.get('actual_observable_session_from_bars')}**",
        f"Japan times that may use HKG: `{a.get('Japan_times_that_may_use_HKG')}`",
        f"Japan times that may use CHI: `{a.get('Japan_times_that_may_use_CHI')}`",
        "",
        f"HKG lead any Japan sector? **{a.get('HKG_leads_any_japan_sector')}**",
        f"CHI lead any Japan sector? **{a.get('CHI_leads_any_japan_sector')}**",
        f"Survives Japan-market removal: `{a.get('survives_japan_market_removal_sectors')}`",
        f"STRICT causal stock leads n=**{a.get('STRICT_causal_stock_leads_n')}** `{a.get('STRICT_causal_symbols')}`",
        "",
        f"Trading houses? **{a.get('explains_trading_houses')}**  Machinery? **{a.get('explains_machinery')}**  Nonferrous? **{a.get('explains_nonferrous')}**",
        f"Steel? **{a.get('explains_steel')}**  Shipping? **{a.get('explains_shipping')}**  Chemicals? **{a.get('explains_chemicals')}**",
        f"Explains any of 13 tech names? **{a.get('explains_any_of_13_tech')}**",
        "",
        f"Stocks with ≥1 usable external driver: **{a.get('stocks_with_at_least_one_usable_external_driver_n')}** / 105",
        f"Evidence quality: `{q}`",
        f"Remain unexplained: **{a.get('remain_unexplained_n')}**",
        f"Tech still highest-value unresolved? **{a.get('tech_still_highest_value_unresolved')}**",
        f"Paid Taiwan/Korea/CME would address a material missing group? **{a.get('paid_taiwan_korea_cme_would_address_material_missing_group')}** — do not purchase: **{a.get('do_not_purchase')}**",
        "",
        f"ENTRY playbooks built? **{a.get('ENTRY_playbooks_built')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used? **{a.get('Old_Confirmation_used')}**",
        f"Kabu 50? **{a.get('Kabu_50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "Preserve FX and US-risk maps. Tech blocker remains separately tracked. Do not build playbooks yet.",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
