"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_wait5_stage2_fill_edge_precommit import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_wait5_stage2_fill_edge_precommit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Interface",
    "DevelopmentArms",
    "FuturePassGate",
    "ConstructionExample",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _arms(v: Any) -> str:
    if isinstance(v, (list, tuple)):
        return "[" + ", ".join(str(x) for x in v) + "]"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# AM WAIT5 STAGE2 FILL-EDGE-PRESERVING OBJECTIVE PRECOMMIT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. Lock P_FILL ranks 1-2. Quality winner among ranks 3-5.",
            "Old free Top5→Top3 TWO_STAGE is closed. Min-rank quality definition unchanged.",
            "No training. No performance. No PnL. No Exact.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"SESSION: {g.get('SESSION')}",
            "",
            f"LOCKED_FILL_SLOTS: {g.get('LOCKED_FILL_SLOTS')}",
            f"QUALITY_SLOT_POOL: {_arms(g.get('QUALITY_SLOT_POOL'))}",
            f"MAX_MEMBERSHIP_SWAP_PER_COHORT: {g.get('MAX_MEMBERSHIP_SWAP_PER_COHORT')}",
            "",
            f"FILL_PRESERVING_BY_CONSTRUCTION: {_tf(g.get('FILL_PRESERVING_BY_CONSTRUCTION'))}",
            f"FILL_EDGE_PRESERVING_STRUCTURE: {_tf(g.get('FILL_EDGE_PRESERVING_STRUCTURE'))}",
            "",
            f"QUALITY_COMBINATION: {g.get('QUALITY_COMBINATION')}",
            "",
            f"ALLOWED_DEVELOPMENT_ARMS: {_arms(g.get('ALLOWED_DEVELOPMENT_ARMS'))}",
            "",
            f"OLD_TWO_STAGE_INTERFACE_ALLOWED: {_tf(g.get('OLD_TWO_STAGE_INTERFACE_ALLOWED'))}",
            f"FILL_LOSS_ALLOWED_IN_PASS_GATE: {_tf(g.get('FILL_LOSS_ALLOWED_IN_PASS_GATE'))}",
            f"SHORTLIST_SEARCH_ALLOWED: {_tf(g.get('SHORTLIST_SEARCH_ALLOWED'))}",
            f"PROBABILITY_THRESHOLD_ALLOWED: {_tf(g.get('PROBABILITY_THRESHOLD_ALLOWED'))}",
            f"U_D_WEIGHT_SEARCH_ALLOWED: {_tf(g.get('U_D_WEIGHT_SEARCH_ALLOWED'))}",
            "",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("PRIMARY_FINDING") or ""),
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No performance run. No new objective search. No new model.",
            "No W5 runtime adoption. Runtime WAIT_SEC remains 1.0.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
