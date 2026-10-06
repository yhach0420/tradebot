"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_failure_decomposition_v2 import ANALYSIS_ID
from research.am_entry_profit_improvement.publish import json_sanitize, kv_rows

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_entry_profit_failure_decomposition_v2"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "summary",
    "replacement_fill",
    "zero_as_safe",
    "price_scale",
    "current_veto",
    "loss_discrimination",
    "outside_augment",
    "oracle",
    "robustness",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# AM ENTRY PROFIT FAILURE DECOMPOSITION V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. Diagnostic classification of 27-spec replacement failure.",
            "No new model. No new feature. No new target. No veto/augment threshold. No policy.",
            "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED FINAL OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"POSITIVE_UTILITY_N: {_fmt(g.get('POSITIVE_UTILITY_N'))}",
            "",
            f"NEGATIVE_UTILITY_N: {_fmt(g.get('NEGATIVE_UTILITY_N'))}",
            "",
            f"SPEARMAN_FILL_VS_NET: {_fmt(g.get('SPEARMAN_FILL_VS_NET'))}",
            "",
            f"SPEARMAN_TRADE_N_VS_NET: {_fmt(g.get('SPEARMAN_TRADE_N_VS_NET'))}",
            "",
            f"ZERO_AS_SAFE_OUTCOME_SUPPORTED: {_tf(g.get('ZERO_AS_SAFE_OUTCOME_SUPPORTED'))}",
            "",
            f"PRICE_SCALE_TAIL_DISTORTION_SUPPORTED: {_tf(g.get('PRICE_SCALE_TAIL_DISTORTION_SUPPORTED'))}",
            "",
            f"CURRENT_VETO_SIGNAL_SUPPORTED: {_tf(g.get('CURRENT_VETO_SIGNAL_SUPPORTED'))}",
            "",
            f"OUTSIDE_CURRENT_AUGMENT_SIGNAL_SUPPORTED: {_tf(g.get('OUTSIDE_CURRENT_AUGMENT_SIGNAL_SUPPORTED'))}",
            "",
            f"ORACLE_NET_PNL: {_fmt(g.get('ORACLE_NET_PNL'))}",
            "",
            f"ORACLE_PF: {_fmt(g.get('ORACLE_PF'))}",
            "",
            f"ORACLE_MAX_DD: {_fmt(g.get('ORACLE_MAX_DD'))}",
            "",
            f"ORACLE_CURRENT_GATE_PASS: {_tf(g.get('ORACLE_CURRENT_GATE_PASS'))}",
            "",
            f"CURRENT_INTERNAL_ORACLE_NET_PNL: {_fmt(g.get('CURRENT_INTERNAL_ORACLE_NET_PNL'))}",
            "",
            f"CURRENT_INTERNAL_ORACLE_IMPROVEMENT: {_fmt(g.get('CURRENT_INTERNAL_ORACLE_IMPROVEMENT'))}",
            "",
            f"BEST_VETO_ARCHITECTURE_DIAGNOSTIC: {_fmt(g.get('BEST_VETO_ARCHITECTURE_DIAGNOSTIC'))}",
            "",
            f"BEST_AUGMENT_ARCHITECTURE_DIAGNOSTIC: {_fmt(g.get('BEST_AUGMENT_ARCHITECTURE_DIAGNOSTIC'))}",
            "",
            f"PRIMARY_FAILURE_MECHANISM: {_fmt(g.get('PRIMARY_FAILURE_MECHANISM'))}",
            "",
            f"NEXT: {_fmt(g.get('NEXT'))}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            "",
            f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## STOP",
            "",
            "Decomposition complete. 27-spec replacement architecture remains closed.",
            "No Runtime change. No W5 adoption. submit/cancel/live=0/0/0.",
            "",
        ]
    )
