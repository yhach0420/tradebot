"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_wait5_two_stage_development import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_wait5_two_stage_development"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Arms",
    "Distinctness",
    "Fillability",
    "ExecQuality",
    "Conditional",
    "ConsensusDays",
    "Gates",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# AM WAIT5 TWO-STAGE PRECOMMITTED DEVELOPMENT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. CONTROL / FILL_ONLY / TWO_STAGE. 9-rep median/consensus.",
            "No Exact. No PnL. No final model. Runtime WAIT_SEC remains 1.0.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"REPRESENTATION_N: {g.get('REPRESENTATION_N')}",
            "",
            f"CONTROL_FILL_RATE: {_fmt(g.get('CONTROL_FILL_RATE'))}",
            f"FILL_ONLY_FILL_RATE: {_fmt(g.get('FILL_ONLY_FILL_RATE'))}",
            f"TWO_STAGE_FILL_RATE: {_fmt(g.get('TWO_STAGE_FILL_RATE'))}",
            "",
            f"DELTA_FILL_RATE_VS_CONTROL: {_fmt(g.get('DELTA_FILL_RATE_VS_CONTROL'))}",
            f"DELTA_FILL_RATE_VS_FILL_ONLY: {_fmt(g.get('DELTA_FILL_RATE_VS_FILL_ONLY'))}",
            "",
            f"CONTROL_EXEC_U: {_fmt(g.get('CONTROL_EXEC_U'))}",
            f"CONTROL_EXEC_D: {_fmt(g.get('CONTROL_EXEC_D'))}",
            f"FILL_ONLY_EXEC_U: {_fmt(g.get('FILL_ONLY_EXEC_U'))}",
            f"FILL_ONLY_EXEC_D: {_fmt(g.get('FILL_ONLY_EXEC_D'))}",
            f"TWO_STAGE_EXEC_U: {_fmt(g.get('TWO_STAGE_EXEC_U'))}",
            f"TWO_STAGE_EXEC_D: {_fmt(g.get('TWO_STAGE_EXEC_D'))}",
            "",
            f"DELTA_EXEC_U_VS_CONTROL: {_fmt(g.get('DELTA_EXEC_U_VS_CONTROL'))}",
            f"DELTA_EXEC_D_VS_CONTROL: {_fmt(g.get('DELTA_EXEC_D_VS_CONTROL'))}",
            f"DELTA_EXEC_U_VS_FILL_ONLY: {_fmt(g.get('DELTA_EXEC_U_VS_FILL_ONLY'))}",
            f"DELTA_EXEC_D_VS_FILL_ONLY: {_fmt(g.get('DELTA_EXEC_D_VS_FILL_ONLY'))}",
            "",
            f"TWO_STAGE_NE_FILL_ONLY_COHORT_RATE: {_fmt(g.get('TWO_STAGE_NE_FILL_ONLY_COHORT_RATE'))}",
            "",
            f"TWO_STAGE_FILLABILITY_PASS: {_tf(g.get('TWO_STAGE_FILLABILITY_PASS'))}",
            f"STAGE2_QUALITY_INCREMENT_PASS: {_tf(g.get('STAGE2_QUALITY_INCREMENT_PASS'))}",
            f"AM_TWO_STAGE_DEVELOPMENT_PASS: {_tf(g.get('AM_TWO_STAGE_DEVELOPMENT_PASS'))}",
            "",
            f"U_POSITIVE_REP_N: {g.get('U_POSITIVE_REP_N')}",
            f"D_POSITIVE_REP_N: {g.get('D_POSITIVE_REP_N')}",
            "",
            f"U_POSITIVE_DAYS: {g.get('U_POSITIVE_DAYS')}",
            f"U_NEGATIVE_DAYS: {g.get('U_NEGATIVE_DAYS')}",
            f"D_POSITIVE_DAYS: {g.get('D_POSITIVE_DAYS')}",
            f"D_NEGATIVE_DAYS: {g.get('D_NEGATIVE_DAYS')}",
            "",
            f"U_EX_BEST_DAY: {_fmt(g.get('U_EX_BEST_DAY'))}",
            f"U_EX_TOP3_DAYS: {_fmt(g.get('U_EX_TOP3_DAYS'))}",
            f"D_EX_BEST_DAY: {_fmt(g.get('D_EX_BEST_DAY'))}",
            f"D_EX_TOP3_DAYS: {_fmt(g.get('D_EX_TOP3_DAYS'))}",
            "",
            f"PM_ROWS_USED_N: {g.get('PM_ROWS_USED_N')}",
            f"FUTURE_EVENT_USE_N: {g.get('FUTURE_EVENT_USE_N')}",
            f"TARGET_CONTAMINATION_N: {g.get('TARGET_CONTAMINATION_N')}",
            f"HELDOUT_FIT_LEAK_N: {g.get('HELDOUT_FIT_LEAK_N')}",
            f"STAGE2_NONFILL_TARGET_TRAIN_N: {g.get('STAGE2_NONFILL_TARGET_TRAIN_N')}",
            "",
            f"PRIMARY_FINDING: {g.get('PRIMARY_FINDING')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Exact. No PnL. No final model. No W5 runtime adoption. Runtime WAIT_SEC remains 1.0.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
