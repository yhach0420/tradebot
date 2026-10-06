"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.nq_es_sector_symbol_response_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "USDJPY_Preserved_Map",
    "USDJPY_Label_Reconciliation",
    "NQ_ES_Source_Audit",
    "Contract_Semantics",
    "Timestamp_Semantics",
    "Preopen_Response",
    "NQ_Impulse",
    "ES_Impulse",
    "NQ_After_ES",
    "ES_After_NQ",
    "Reverse_Causality",
    "Sector_Response",
    "Stock_Raw",
    "Stock_MarketAdjusted",
    "Stock_SectorAdjusted",
    "Tech_Name_Response",
    "Response_Groups",
    "D1_D4",
    "Unexplained_Groups",
    "Next_Driver",
    "True_Japan_Futures_Subtrack",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    nq_sec = list(report.get("nq_sector") or [])
    es_sec = list(report.get("es_sector") or [])
    inc = list(report.get("incremental_sector") or [])
    stocks = list(report.get("stock_response") or [])
    fx_map = dict(report.get("fx_response_map_v1") or {})
    groups = dict(report.get("groups") or {})
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "architecture": report.get("architecture"),
                "usdjpy_preserved": True,
                "kabu_50": False,
                "frozen_validation_opened": False,
                "old_confirmation_used_to_design": False,
                "playbooks_built": False,
            }
        ),
        "USDJPY_Preserved_Map": _kv(fx_map) or [{"empty": True}],
        "USDJPY_Label_Reconciliation": _kv(dict(report.get("usdjpy_label_reconciliation") or {})) or [{"empty": True}],
        "NQ_ES_Source_Audit": _kv(dict(report.get("source_probe") or {}))
        + _kv(dict(report.get("acquisition") or {}))
        + _kv(dict(report.get("nq_meta") or {}))
        + _kv(dict(report.get("es_meta") or {})),
        "Contract_Semantics": _kv(dict(report.get("contract_semantics") or {})),
        "Timestamp_Semantics": _kv(dict(report.get("timestamp_semantics") or {})),
        "Preopen_Response": list(report.get("preopen_incremental") or [{"empty": True}]),
        "NQ_Impulse": [
            {"sector": r.get("sector"), "lead_lag": r.get("lead_lag"), "fwd_rho_residual": r.get("fwd_rho_residual"), "usable": r.get("usable_before_stock_move")}
            for r in nq_sec
        ]
        or [{"empty": True}],
        "ES_Impulse": [
            {"sector": r.get("sector"), "lead_lag": r.get("lead_lag"), "fwd_rho_residual": r.get("fwd_rho_residual"), "usable": r.get("usable_before_stock_move")}
            for r in es_sec
        ]
        or [{"empty": True}],
        "NQ_After_ES": [
            {
                "sector": r.get("sector"),
                "nq_after_es_1m": r.get("nq_after_es_1m"),
                "nq_after_es_lead_lag": r.get("nq_after_es_lead_lag"),
                "nq_incremental_after_es": r.get("nq_incremental_after_es"),
            }
            for r in inc
        ]
        or [{"empty": True}],
        "ES_After_NQ": [
            {
                "sector": r.get("sector"),
                "es_after_nq_1m": r.get("es_after_nq_1m"),
                "es_incremental_after_nq": r.get("es_incremental_after_nq"),
            }
            for r in inc
        ]
        or [{"empty": True}],
        "Reverse_Causality": (
            [{"driver": "NQ", "sector": r.get("sector"), "reverse_rho": r.get("reverse_rho"), "lead_lag": r.get("lead_lag")} for r in nq_sec]
            + [{"driver": "ES", "sector": r.get("sector"), "reverse_rho": r.get("reverse_rho"), "lead_lag": r.get("lead_lag")} for r in es_sec]
        )
        or [{"empty": True}],
        "Sector_Response": inc or [{"empty": True}],
        "Stock_Raw": [
            {"symbol": r.get("symbol"), "sector": r.get("sector"), "nq_raw_1m": r.get("nq_raw_1m"), "es_raw_1m": r.get("es_raw_1m"), "class": r.get("class")}
            for r in stocks
        ]
        or [{"empty": True}],
        "Stock_MarketAdjusted": [
            {"symbol": r.get("symbol"), "nq_market_adj_1m": r.get("nq_market_adj_1m"), "es_market_adj_1m": r.get("es_market_adj_1m"), "class": r.get("class")}
            for r in stocks
        ]
        or [{"empty": True}],
        "Stock_SectorAdjusted": [
            {
                "symbol": r.get("symbol"),
                "nq_sector_adj_1m": r.get("nq_sector_adj_1m"),
                "es_sector_adj_1m": r.get("es_sector_adj_1m"),
                "nq_after_es_1m": r.get("nq_after_es_1m"),
                "es_after_nq_1m": r.get("es_after_nq_1m"),
                "class": r.get("class"),
                "usable_before_stock_move": r.get("usable_before_stock_move"),
                "simultaneous_common_news": r.get("simultaneous_common_news"),
            }
            for r in stocks
        ]
        or [{"empty": True}],
        "Tech_Name_Response": list(report.get("tech_name_response") or [{"empty": True}]),
        "Response_Groups": list(groups.get("groups") or [{"empty": True}]),
        "D1_D4": [
            {
                "symbol": r.get("symbol"),
                "d1_d4_nq": r.get("d1_d4_nq"),
                "d1_d4_es": r.get("d1_d4_es"),
                "d1_d4_nq_after_es": r.get("d1_d4_nq_after_es"),
                "class": r.get("class"),
            }
            for r in stocks
        ]
        or [{"empty": True}],
        "Unexplained_Groups": _kv(dict(report.get("unexplained") or {})),
        "Next_Driver": _kv(dict(report.get("next_driver") or {})),
        "True_Japan_Futures_Subtrack": _kv(dict(report.get("true_futures_subtrack") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        "# NQ_ES_SECTOR_SYMBOL_CAUSAL_RESPONSE_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted parent: `USDJPY_SECTOR_SPECIFIC_DRIVER_FOUND_V1` with narrow interpretation.",
        "USDJPY is a sector-specific overnight/preopen driver plus a small set of stock-level intraday direct candidates.",
        "USDJPY is **not** a global intraday driver for all 105 stocks.",
        "",
        f"USDJPY labels reconciled? **{a.get('USDJPY_labels_reconciled')}**",
        f"USABLE_DIRECT_CAUSAL_LEAD_N: **{a.get('USABLE_DIRECT_CAUSAL_LEAD_N')}** `{a.get('usable_direct_fx_symbols')}`",
        f"FX overnight group preserved? **{a.get('FX_overnight_group_preserved')}** sectors `{a.get('FX_overnight_sectors')}`",
        "",
        f"NQ source available? **{a.get('NQ_historical_source_available')}**  `{a.get('NQ_true_futures_or_proxy')}`  `{a.get('NQ_free_or_paid')}`",
        f"ES source available? **{a.get('ES_historical_source_available')}**  `{a.get('ES_true_futures_or_proxy')}`",
        f"Purchase? **{a.get('any_purchase')}**",
        f"Timestamp semantics proven? **{a.get('timestamp_semantics_proven')}**",
        f"Contract/roll semantics proven? **{a.get('contract_roll_semantics_proven')}**",
        "",
        f"NQ lead electrical machinery? **{a.get('NQ_leads_electrical_machinery')}**",
        f"ES lead electrical machinery? **{a.get('ES_leads_electrical_machinery')}**",
        f"NQ incremental after ES? **{a.get('NQ_incremental_after_ES')}**",
        f"6857 `{a.get('NQ_explains_6857')}`  8035 `{a.get('NQ_explains_8035')}`  6920 `{a.get('NQ_explains_6920')}`  285A0 `{a.get('NQ_explains_285A0')}`",
        f"Tech STRICT_CAUSAL_LEAD: `{a.get('tech_STRICT_CAUSAL_LEAD')}`",
        f"Tech SIMULTANEOUS_ONLY: `{a.get('tech_SIMULTANEOUS_ONLY')}`",
        f"Tech NO_STABLE_RESPONSE: `{a.get('tech_NO_STABLE_RESPONSE')}`",
        f"Preopen NQ relationship? **{a.get('preopen_NQ_relationship')}**",
        f"Overnight USDJPY interaction: `{a.get('overnight_USDJPY_interaction_changes_attribution')}`",
        f"New stable driver-response groups: **{a.get('new_stable_driver_response_groups_n')}**",
        f"Stocks with ≥1 causally usable external driver: **{a.get('stocks_with_at_least_one_causally_usable_external_driver_n')}** / 105",
        f"Remain unexplained: **{a.get('remain_unexplained_n')}**",
        f"Next driver family: **{a.get('next_driver_family')}** — {a.get('next_driver_why')}",
        "",
        f"ENTRY playbooks built? **{a.get('ENTRY_playbooks_built')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
        f"Kabu 50? **{a.get('Kabu_50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "Do not build playbooks yet. Preserve confirmed USDJPY groups.",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
