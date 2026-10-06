"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.support_resistance_first_interaction_matched_causal_test_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Detector_Freeze",
    "Direction",
    "Population_Primary",
    "Population_Secondary",
    "Matching",
    "Placebo",
    "Question_A",
    "Question_B",
    "Question_C",
    "Question_D",
    "Bootstrap",
    "Outcome_Metrics",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, (list, tuple)):
        return [json_sanitize(x) for x in got]
    return got


def _kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for k, v in d.items():
        if isinstance(v, (dict, list, tuple)):
            v = json.dumps(json_sanitize(v), ensure_ascii=False)[:32000]
        rows.append({"key": str(k), "value": v})
    return rows


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["key", "value"])
        ws.append(["empty", True])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_c)) + 2))


def _q_rows(q: dict[str, Any]) -> list[dict[str, Any]]:
    st = dict(q.get("d2_d4") or {})
    boot = dict(q.get("bootstrap_date") or {})
    d1 = dict(q.get("d1") or {})
    return _kv_rows(
        {
            "question": q.get("question"),
            "eval_blocks": q.get("eval_blocks"),
            "powered": q.get("powered"),
            "separation_on_primary_metric": q.get("separation_on_primary_metric"),
            "favorable_on_primary_metric": q.get("favorable_on_primary_metric"),
            "primary_metric": q.get("primary_metric"),
            "five_pp_continuation_not_used": q.get("five_pp_continuation_not_used"),
            "treatment_n": st.get("treatment_n"),
            "matched_n": st.get("matched_n"),
            "match_rate": st.get("match_rate"),
            "day_n": st.get("day_n"),
            "symbol_n": st.get("symbol_n"),
            "gap_p20_before_m20": st.get("gap_p20_before_m20"),
            "gap_p40_before_m20": st.get("gap_p40_before_m20"),
            "gap_median_mfe": st.get("gap_median_mfe"),
            "gap_median_end": st.get("gap_median_end"),
            "gap_mfe_before_mae": st.get("gap_mfe_before_mae"),
            "treatment": st.get("treatment"),
            "matched_control": st.get("matched_control"),
            "bootstrap_date": boot,
            "bootstrap_symbol": q.get("bootstrap_symbol"),
            "d1_diagnostic": d1,
        }
    )


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    bind = dict(report.get("bind") or {})
    freeze = dict(report.get("freeze") or {})
    d = dict(report.get("decision") or {})
    safety = dict(report.get("safety") or {})
    qs = dict(report.get("questions_primary") or {})
    return {
        "Binding": _kv_rows(
            {
                "analysis_id": report.get("analysis_id"),
                "parent_verdict": report.get("parent_verdict"),
                "bind_ok": bind.get("ok"),
                "identity": report.get("identity"),
                "is_strategy": False,
                "old_no_info_revived": False,
                "no_pnl_parameter_tuning": True,
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
                "INTERPRETATION": d.get("INTERPRETATION"),
            }
        ),
        "Detector_Freeze": _kv_rows(freeze),
        "Direction": _kv_rows(dict(report.get("direction") or {})),
        "Population_Primary": _kv_rows(dict(report.get("population_primary") or {})),
        "Population_Secondary": _kv_rows(dict(report.get("population_secondary") or {})),
        "Matching": _kv_rows(dict(report.get("matching") or {})),
        "Placebo": _kv_rows(dict(report.get("placebo") or {})),
        "Question_A": _q_rows(dict(qs.get("A") or {})),
        "Question_B": _q_rows(dict(qs.get("B") or {})),
        "Question_C": _q_rows(dict(qs.get("C") or {})),
        "Question_D": _q_rows(dict(qs.get("D") or {})),
        "Bootstrap": _kv_rows(
            {
                "primary_A": (qs.get("A") or {}).get("bootstrap_date"),
                "primary_B": (qs.get("B") or {}).get("bootstrap_date"),
                "primary_C": (qs.get("C") or {}).get("bootstrap_date"),
                "primary_D": (qs.get("D") or {}).get("bootstrap_date"),
                "cluster": "date then symbol",
                "eval_blocks": "D2 D3 D4",
                "d1": "diagnostic_only",
            }
        ),
        "Outcome_Metrics": list(report.get("pairs_compact") or []),
        "Safety": _kv_rows(safety),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    qa = a.get("Question A first-touch rejection?") or {}
    qb = a.get("Question B break continuation?") or {}
    qc = a.get("Question C retest hold?") or {}
    qd = a.get("Question D failed break?") or {}
    return "\n".join(
        [
            "# SUPPORT_RESISTANCE_FIRST_INTERACTION_MATCHED_CAUSAL_TEST_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            "This is research. It is not a strategy.",
            "",
            f"Detector frozen? **{a.get('Detector frozen exactly as rebuilt?')}**",
            f"DETECTOR_SHA256: `{a.get('DETECTOR_SHA256')}`",
            f"STATE_MACHINE_SHA256: `{a.get('STATE_MACHINE_SHA256')}`",
            f"Primary first-test n? **{a.get('Primary first-test n?')}** (expected {a.get('Expected primary first-test n?')}; identity ok? {a.get('Identity ok?')})",
            f"Secondary role-flip n? **{a.get('Secondary role-flip first-test n?')}** mixed into primary? **{a.get('Were broken role-flip candidates mixed into the primary test?')}**",
            f"Placebo A separation (diagnostic, not in verdict)? **{a.get('Placebo A separation diagnostic, not in verdict?')}**",
            f"Resistance first-touch favorable? **{a.get('Direction resistance first-touch favorable?')}** support? **{a.get('Direction support first-touch favorable?')}**",
            f"Q A p20-before-m20 gap: **{(qa or {}).get('gap_p20_before_m20')}** CI [{(qa or {}).get('ci95_lo')}, {(qa or {}).get('ci95_hi')}] powered={(qa or {}).get('powered')} sep={(qa or {}).get('separation')}",
            f"Q B p20-before-m20 gap: **{(qb or {}).get('gap_p20_before_m20')}** CI [{(qb or {}).get('ci95_lo')}, {(qb or {}).get('ci95_hi')}] powered={(qb or {}).get('powered')} sep={(qb or {}).get('separation')}",
            f"Q C p20-before-m20 gap: **{(qc or {}).get('gap_p20_before_m20')}** CI [{(qc or {}).get('ci95_lo')}, {(qc or {}).get('ci95_hi')}] powered={(qc or {}).get('powered')} sep={(qc or {}).get('separation')}",
            f"Q D p20-before-m20 gap: **{(qd or {}).get('gap_p20_before_m20')}** CI [{(qd or {}).get('ci95_lo')}, {(qd or {}).get('ci95_hi')}] powered={(qd or {}).get('powered')} sep={(qd or {}).get('separation')}",
            f"Primary metric? **{a.get('Primary metric?')}** 5pp sole gate? **{a.get('Was 5pp continuation used as the sole gate?')}**",
            f"PnL optimization? **{a.get('Any PnL optimization performed?')}** X0/X1 used to select rules? **{a.get('Any X0/X1 used to select rules?')}**",
            f"Old Confirmation opened? **{a.get('Old Confirmation opened?')}** Frozen Validation opened? **{a.get('Frozen Validation opened?')}**",
            f"Is this a strategy? **{a.get('Is this a strategy?')}** old no-info revived? **{a.get('Old no-info revived as economic null?')}**",
            f"submit/cancel/live: **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if k not in {"_markdown", "all_pairs"}})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
