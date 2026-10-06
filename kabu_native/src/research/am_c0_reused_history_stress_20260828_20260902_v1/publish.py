"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_c0_reused_history_stress_20260828_20260902_v1 import ANALYSIS_ID
from research.am_c0_reused_history_stress_20260828_20260902_v1.isolation import OUT
from research.am_entry_profit_improvement.publish import json_sanitize

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "precommit",
    "pin",
    "current",
    "augment",
    "overlay",
    "paired",
    "daily_pnl",
    "concentration",
    "preservation",
    "integrity",
    "economic_gates",
    "decision",
    "safety",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "LABEL: REUSED_HISTORY_STRESS",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"CASE: {_fmt(d.get('CASE'))}",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT'))}",
        "",
        "## Answers 1-41",
        "",
    ]
    order = [
        "1_prior_simple_tech_closure",
        "2_prior_p1_recovery_exhaustion",
        "3_why_c0_is_next",
        "4_c0_spec_sha_parity",
        "4_c0_spec_sha256",
        "5_c0_source_reproducible",
        "6_c14_reproducible",
        "7_development_model_frozen",
        "8_stress_refit_n",
        "9_stress_days",
        "10_TRUE_OOS",
        "11_label",
        "12_CURRENT_trades_pnl_pf_dd",
        "13_C0_candidate_admitted_fill_trade_n",
        "14_C0_pnl_pf_dd",
        "15_C0_win_loss",
        "16_OVERLAY_pnl_pf_dd",
        "17_delta_pnl_pf_dd",
        "18_paired_positive_negative_zero",
        "19_paired_median",
        "20_top_day_concentration",
        "21_top_symbol_concentration",
        "22_EX_TOP1",
        "23_EX_TOP2",
        "24_current_preservation_mismatch_counts",
        "25_integrity_gates",
        "26_economic_gates",
        "27_verdict",
        "28_c0_priority_maintained",
        "29_new_entry_family_design_allowed",
        "30_sizing_allowed",
        "31_ENTRY_Runtime_changed",
        "32_EXIT_Runtime_changed",
        "33_CAP_changed",
        "34_model_changed",
        "35_future_data_used",
        "36_MAX_RESEARCH_DATE",
        "37_prospective_suspended",
        "38_TRUE_OOS",
        "39_CERTIFIED",
        "40_submit_cancel_live",
        "41_next",
    ]
    for k in order:
        v = a.get(k)
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {_fmt(v) if not isinstance(v, str) else v}")
        lines.append("")
    lines.extend(
        [
            "STOP.",
            "",
        ]
    )
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))
