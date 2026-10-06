"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.one_minute_native_playbook_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Symbol_Behavior_Profile",
    "Sector_Behavior_Profile",
    "Leader_Laggard",
    "Opening_Behavior",
    "Intraday_Sequences",
    "Episode_Definitions",
    "Path_Outcomes",
    "Failure_Paths",
    "Behavior_Groups",
    "Group_Members",
    "Playbook_Candidates",
    "Complete_Strategy",
    "Execution",
    "D1_D4",
    "HM1_Reference",
    "External_Context_Optional",
    "Unsupported_Stocks",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    sym = list(report.get("symbol_profiles") or [])
    plays = list((report.get("playbooks") or {}).get("proposals") or [])
    promoted = list((report.get("playbooks") or {}).get("promoted") or [])
    groups = list((report.get("behavior_groups") or {}).get("groups") or [])
    members = []
    for g in groups:
        for s in list(g.get("symbols") or []):
            members.append({"group": g.get("group"), "symbol": s})
    seq = list(report.get("sequences") or [])
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "atlas_id": report.get("atlas_id"),
                "split_sha256": report.get("split_sha256"),
                "block_sha256": report.get("block_sha256"),
                "catalog_stopped": True,
                "no_new_paid_data": True,
                "hm1_tuned": False,
                "five_minute_grid": False,
                "kabu_50": False,
            }
        ),
        "Symbol_Behavior_Profile": [
            {k: r.get(k) for k in ("symbol", "sector", "episode_n", "continuation_share", "reversal_share", "favor_first_p", "lead_frac", "lag_frac", "opening_hold_p", "best_sequence", "failure_sequence", "behavior_group", "d1_d4_favor_agree")}
            for r in sym
        ]
        or [{"empty": True}],
        "Sector_Behavior_Profile": list(report.get("sector_profiles") or [{"empty": True}]),
        "Leader_Laggard": [
            {"symbol": r.get("symbol"), "sector": r.get("sector"), "lead_frac": r.get("lead_frac"), "lag_frac": r.get("lag_frac"), "lead_minus_lag": r.get("lead_minus_lag")}
            for r in sorted(sym, key=lambda z: -(z.get("lead_minus_lag") or 0))
        ]
        or [{"empty": True}],
        "Opening_Behavior": [
            {"symbol": r.get("symbol"), "opening_n": r.get("opening_n"), "opening_hold_p": r.get("opening_hold_p"), "opening_gap_up_hold_p": r.get("opening_gap_up_hold_p")}
            for r in sym
        ]
        or [{"empty": True}],
        "Intraday_Sequences": seq or [{"empty": True}],
        "Episode_Definitions": _kv(
            {
                "gap_min": (report.get("episode_meta") or {}).get("gap_min"),
                "future_in_boundary": False,
                "bar_start": True,
                "decision_available_at": "feature_bar_end_T+1m",
                "collapse": "same_symbol_events_within_10m",
            }
        ),
        "Path_Outcomes": [
            {"sequence": r.get("sequence"), "favor_first_p": r.get("favor_first_p"), "continuation_p": r.get("continuation_p"), "mean_mfe_bps": r.get("mean_mfe_bps"), "mean_mae_bps": r.get("mean_mae_bps")}
            for r in seq
        ]
        or [{"empty": True}],
        "Failure_Paths": list(report.get("worst_favor_sequences") or [{"empty": True}]),
        "Behavior_Groups": groups or [{"empty": True}],
        "Group_Members": members or [{"empty": True}],
        "Playbook_Candidates": [
            {
                "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                "sequence": (p.get("spec") or {}).get("sequence"),
                "promoted": (p.get("promotion") or {}).get("promoted"),
                "fail_reasons": (p.get("promotion") or {}).get("fail_reasons"),
                "X0": (p.get("economics") or {}).get("mean_x0_bps"),
                "X1": (p.get("economics") or {}).get("mean_x1_bps"),
                "trade_n": (p.get("economics") or {}).get("trade_n"),
            }
            for p in plays
        ]
        or [{"empty": True}],
        "Complete_Strategy": [
            {
                **dict(p.get("spec") or {}),
                "promoted": True,
                "X0": (p.get("economics") or {}).get("mean_x0_bps"),
                "X1": (p.get("economics") or {}).get("mean_x1_bps"),
            }
            for p in promoted
        ]
        or [{"none_promoted": True}],
        "Execution": [
            {
                "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                **{k: (p.get("economics") or {}).get(k) for k in ("mean_x0_bps", "mean_x1_bps", "profit_factor", "hit_rate", "trade_n", "day_n", "max_dd_daily_mean_bps", "exit_reasons")},
            }
            for p in plays
        ]
        or [{"empty": True}],
        "D1_D4": [
            {"candidate_id": (p.get("spec") or {}).get("candidate_id"), "block_mean_x0": (p.get("economics") or {}).get("block_mean_x0"), "block_positive_n": (p.get("economics") or {}).get("block_positive_n")}
            for p in plays
        ]
        or [{"empty": True}],
        "HM1_Reference": _kv(dict(report.get("hm1_reference") or {})),
        "External_Context_Optional": _kv(dict(report.get("external_optional") or {})),
        "Unsupported_Stocks": [{"symbol": s} for s in list(report.get("unsupported_symbols") or [])] or [{"empty": True}],
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0", "new_paid_data": False, "frozen_validation_opened": False}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        "# ONE_MINUTE_NATIVE_PLAYBOOK_DISCOVERY_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        f"Atlas: `{a.get('atlas_id')}`",
        "Course correction: native 105-stock 1-minute Discovery panel is the primary asset. External catalog expansion stopped. No new paid data.",
        "HM1_CONTROLLED_PULLBACK_RECLAIM is preserved as one playbook candidate and was not retuned.",
        "",
        f"Stable mechanisms n=**{a.get('stable_symbol_sector_behavior_mechanisms_n')}** `{a.get('stable_sequences')}`",
        f"Lead sector: `{a.get('stocks_repeatedly_lead_sector')}`",
        f"Lag / catch-up: `{a.get('stocks_repeatedly_lag_then_catch_up')}`",
        f"Continuation-type: `{a.get('continuation_type')}`",
        f"Mean-reversion-type: `{a.get('mean_reversion_type')}`",
        f"Stable opening: `{a.get('stable_opening_behavior')}`",
        f"Largest favorable-first sequences: `{a.get('largest_favorable_first_sequences')}`",
        f"Failing sequences: `{a.get('consistently_failing_sequences')}`",
        f"Behavior groups n=**{a.get('behavior_groups_n')}**",
        f"Complete playbooks clearing execution n=**{a.get('complete_playbook_candidates_clearing_execution_n')}** `{a.get('promoted_ids')}`",
        f"X0=**{a.get('X0')}** X1=**{a.get('X1')}** PF=**{a.get('PF')}** trade N=**{a.get('trade_N')}** D1-D4=`{a.get('D1_D4')}`",
        f"HM1 remains competitive? **{a.get('HM1_remains_competitive')}** (frozen X0=`{a.get('HM1_reference_x0')}` X1=`{a.get('HM1_reference_x1')}`)",
        f"Unsupported stocks n=**{a.get('unsupported_stock_n')}**",
        f"External driver required for any playbook? **{a.get('external_driver_required_for_any_playbook')}**",
        "",
        f"ENTRY playbooks built? **{a.get('ENTRY_playbooks_built')}**",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
        f"New paid data? **{a.get('New_paid_data')}**",
        f"5-minute grid? **{a.get('five_minute_grid')}**",
        f"Kabu 50? **False**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
