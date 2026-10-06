"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_wait5_fill_upside_geometry import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_wait5_fill_upside_geometry"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Spearman",
    "RankBuckets",
    "Swap",
    "Oracle",
    "Frontier",
    "DayRobustness",
    "Representations",
    "Gates",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# AM WAIT5 FILLABILITY-UPSIDE ARCHITECTURE GEOMETRY AUDIT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. Frozen 18-day OOF. CONTROL = CURRENT Top3. FILL_ONLY = P_FILL5 Top3.",
            "Oracle Top3 enumeration is diagnostic only. Not a strategy.",
            "No refit. No new model. No Stage2. No Exact. No PnL.",
            "Runtime WAIT_SEC remains 1.0. DEV_WAIT_SEC = 5.0 diagnostic only.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"P_FILL_U_SPEARMAN_MEDIAN: {_fmt(g.get('P_FILL_U_SPEARMAN_MEDIAN'))}",
            f"P_FILL_D_SPEARMAN_MEDIAN: {_fmt(g.get('P_FILL_D_SPEARMAN_MEDIAN'))}",
            "",
            f"R1_3_FILL_RATE: {_fmt(g.get('R1_3_FILL_RATE'))}",
            f"R1_3_COND_U: {_fmt(g.get('R1_3_COND_U'))}",
            f"R4_5_FILL_RATE: {_fmt(g.get('R4_5_FILL_RATE'))}",
            f"R4_5_COND_U: {_fmt(g.get('R4_5_COND_U'))}",
            "",
            f"FILL_NONWORSE_U_IMPROVE_COHORT_N: {g.get('FILL_NONWORSE_U_IMPROVE_COHORT_N')}",
            f"FILL_NONWORSE_U_IMPROVE_RATE: {_fmt(g.get('FILL_NONWORSE_U_IMPROVE_RATE'))}",
            f"SAME_FILL_U_IMPROVE_COHORT_N: {g.get('SAME_FILL_U_IMPROVE_COHORT_N')}",
            f"SAME_FILL_U_IMPROVE_RATE: {_fmt(g.get('SAME_FILL_U_IMPROVE_RATE'))}",
            "",
            f"FILL_NONWORSE_U_D_NONWORSE_RATE: {_fmt(g.get('FILL_NONWORSE_U_D_NONWORSE_RATE'))}",
            f"SAME_FILL_U_D_NONWORSE_RATE: {_fmt(g.get('SAME_FILL_U_D_NONWORSE_RATE'))}",
            "",
            f"ORACLE_EXEC_U_DELTA_MEDIAN: {_fmt(g.get('ORACLE_EXEC_U_DELTA_MEDIAN'))}",
            f"ORACLE_EXEC_U_EX_BEST_DAY: {_fmt(g.get('ORACLE_EXEC_U_EX_BEST_DAY'))}",
            f"ORACLE_EXEC_U_EX_TOP3_DAYS: {_fmt(g.get('ORACLE_EXEC_U_EX_TOP3_DAYS'))}",
            "",
            f"FILL_ONLY_DOMINATED_IN_FILL_U_RATE: {_fmt(g.get('FILL_ONLY_DOMINATED_IN_FILL_U_RATE'))}",
            "",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Exact. No PnL. No Stage2 restart. No EXEC_U training. No W5 runtime adoption.",
            "Runtime WAIT_SEC remains 1.0. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
