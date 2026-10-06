"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.reference_level_1m_price_action_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Reference_Level_Definitions",
    "Availability",
    "Previous_Day_Levels",
    "Opening_Range",
    "Gap_Structure",
    "VWAP",
    "MultiDay_Levels",
    "Level_Interaction_Events",
    "Break_Accept",
    "Break_Failure",
    "Retest",
    "Reclaim",
    "Path_Outcomes",
    "Symbol_Response",
    "Sector_Response",
    "Candidate_Playbooks",
    "Complete_Strategy",
    "Execution",
    "D1_D4",
    "Causality_Audit",
    "Session_Semantics",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _kinds(atlas: dict[str, Any], pred) -> list[dict[str, Any]]:
    return [r for r in list(atlas.get("level_kinds") or []) if pred(r)]


def _levels(atlas: dict[str, Any], family: str) -> list[dict[str, Any]]:
    return [r for r in list(atlas.get("levels") or []) if str(r.get("family")) == family]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    atlas = dict(report.get("atlas") or {})
    causal = dict(report.get("causality_audit") or {})
    sess = dict(report.get("session_semantics") or {})
    sel = dict(report.get("playbook_selection") or {})
    replay = dict(report.get("replay") or {})
    packs = list(replay.get("packs") or [])
    subgroups = list(report.get("subgroups") or [])
    d1d4 = list(report.get("d1_d4") or [])
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "r11_status": report.get("r11_status"),
                "r11_role": report.get("r11_role"),
                "r11_repaired": False,
                "old_confirmation_opened": False,
                "frozen_validation_opened": False,
            }
        ),
        "Reference_Level_Definitions": list(report.get("level_definitions") or []),
        "Availability": _kv(
            {
                "OR5": "feature_bar >= 09:05; range 09:00-09:04",
                "OR15": "feature_bar >= 09:15; range 09:00-09:14",
                "PDH": "completed prior Discovery day only; day-1 has no PDH",
                "D5": "needs 5 completed prior days",
                "D20": "needs 20 completed prior days",
                "GAP": "known at first session bar vs PDC",
                "MINUTE_BAR_VOLUME_PROFILE_APPROXIMATION": "not used as S/R truth; context omitted from first atlas pass",
            }
        ),
        "Previous_Day_Levels": _levels(atlas, "previous_day") + _kinds(atlas, lambda r: str(r.get("level_id")) in {"PDH", "PDL", "PDC"}),
        "Opening_Range": _levels(atlas, "opening_range") + _kinds(atlas, lambda r: str(r.get("level_id") or "").startswith("OR")),
        "Gap_Structure": _levels(atlas, "gap_structure") + _kinds(atlas, lambda r: "GAP" in str(r.get("level_id") or "") or "GAP" in str(r.get("event_kind") or "")),
        "VWAP": _kinds(atlas, lambda r: str(r.get("level_id")) == "VWAP"),
        "MultiDay_Levels": _levels(atlas, "multi_day") + _kinds(atlas, lambda r: str(r.get("level_id") or "").startswith("D")),
        "Level_Interaction_Events": list(report.get("comparisons") or []),
        "Break_Accept": _kinds(atlas, lambda r: "BREAK" in str(r.get("event_kind") or "") or str(r.get("event_kind") or "").startswith("ACCEPT")),
        "Break_Failure": _kinds(atlas, lambda r: "FAILED" in str(r.get("event_kind") or "") or "FAILURE" in str(r.get("event_kind") or "")),
        "Retest": _kinds(atlas, lambda r: "RETEST" in str(r.get("event_kind") or "")),
        "Reclaim": _kinds(atlas, lambda r: "RECLAIM" in str(r.get("event_kind") or "")),
        "Path_Outcomes": list(atlas.get("strongest_path_separation") or []),
        "Symbol_Response": [
            {"playbook_id": s.get("playbook_id"), "overall": s.get("overall"), "note": "symbol PnL not used as a pre-filter"}
            for s in subgroups
        ],
        "Sector_Response": [
            {"playbook_id": s.get("playbook_id"), "by_sector_top": s.get("by_sector_top"), "by_gap": s.get("by_gap")}
            for s in subgroups
        ],
        "Candidate_Playbooks": list(sel.get("all") or []),
        "Complete_Strategy": packs,
        "Execution": _kv(
            {
                "X0": "next causally executable bar open (available_at = feature_bar+1)",
                "X1": "X0 minus 8bps",
                "same_bar_execution": False,
                "occupancy": 3,
                "same_symbol": "one per symbol per day",
                "slot_release": "exit_hh > next candidate time",
                "reentry": "not same symbol same day",
            }
        ),
        "D1_D4": d1d4,
        "Causality_Audit": _kv(causal),
        "Session_Semantics": _kv(sess),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0"}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    fam = a.get("How_many_causal_events_for_each_reference_level_family") or {}
    lines = [
        "# REFERENCE_LEVEL_1M_PRICE_ACTION_DISCOVERY_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "R11_OFFLINE_CLUSTER_LOOKAHEAD_CONFIRMED_V1 accepted. NATIVE_1M_R11_VWAP_RECLAIM_V1 is INVALID as a causal strategy. Not repaired.",
        "R11 predicates retained only as HISTORICAL_RESEARCH_CLUE (VWAP / morning / reclaim).",
        "",
        f"Causal events by family: `{json.dumps(fam, ensure_ascii=False, default=str)}`",
        f"future-dependent event generation? **{a.get('Any_future_dependent_event_generation')}**",
        f"retroactive timestamp assignment? **{a.get('Any_retroactive_timestamp_assignment')}**",
        f"Strongest path-behavior levels: `{a.get('Which_levels_show_the_strongest_change_in_subsequent_path_behavior')}`",
        f"VWAP still special vs prior-day context? **{a.get('Does_VWAP_remain_special_after_prior_day_gap_OR_context')}**",
        f"Level × 1m price action improves path separation? **{a.get('Does_combining_a_level_with_1m_price_action_improve_path_separation')}**",
        f"Stable mechanisms D1-D4: `{a.get('Any_economically_explainable_mechanism_stable_across_D1_D4')}`",
        f"Candidate playbooks n=**{a.get('How_many_candidate_playbooks')}** ids=`{a.get('candidate_playbook_ids')}`",
        f"Complete-strategy X1>0? **{a.get('Any_complete_strategy_candidate_with_X1_gt_0')}** ids=`{a.get('complete_strategy_candidate_ids')}`",
        f"Dominated by one symbol? **{a.get('Any_candidate_dominated_by_one_symbol_or_top_winners')}** `{a.get('dominated_ids')}`",
        f"Lunch policy: **{a.get('lunch_policy')}** (architecture, not PnL)",
        f"Implicit 11:30 truncation? **{a.get('Any_implicit_1130_truncation')}**",
        f"Old Confirmation opened? **{a.get('Old_Confirmation_opened')}** Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"New paid data? **{a.get('New_paid_data')}** Kabu50? **{a.get('Kabu50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
