"""Publish verified-window runtime binding evidence. No strategy change."""
from __future__ import annotations

import json
import pickle
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook

from research.fixed_entry_support_runtime_binding_evidence_recovery_v1.contract import run_contract
from research.fixed_entry_support_runtime_binding_evidence_recovery_v1.evidence import classify_intervals, load_windows
from research.fixed_entry_support_runtime_binding_evidence_recovery_v1.scan import scan

OUT = Path("results/research/fixed_entry_support_runtime_binding_evidence_recovery_v1")


def _pf(rows: list[dict[str, Any]]) -> Optional[float]:
    loss = sum(-float(row["pnl_yen"]) for row in rows if float(row["pnl_yen"]) < 0)
    gain = sum(float(row["pnl_yen"]) for row in rows if float(row["pnl_yen"]) > 0)
    return gain / loss if loss else None


def _stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {"trade_n": len(rows), "pnl_yen": float(sum(float(row["pnl_yen"]) for row in rows)), "pf": _pf(rows)}


def _uid(row: dict[str, Any]) -> str:
    return f"{row['date']}|{row['session']}|{row['symbol']}|{int(row['signal_index'])}"


def _compare(reference: list[dict[str, Any]], implementation: list[dict[str, Any]], counts: dict[str, int]) -> dict[str, Any]:
    left = {_uid(row): row for row in reference}
    right = {_uid(row): row for row in implementation}
    for uid in set(left) & set(right):
        ref, got = left[uid], right[uid]
        if abs(float(ref["entry_t"]) - float(got["entry_t"])) > 1e-9 or abs(float(ref["entry_px"]) - float(got["entry_px"])) > 1e-6:
            counts["entry"] += 1
        if abs(float(ref["exit_t"]) - float(got["exit_t"])) > 1e-6 or abs(float(ref["exit_px"]) - float(got["exit_px"])) > 1e-4 or str(ref["reason"]) != str(got["reason"]):
            counts["exit"] += 1
        if str(ref["reason"]) != str(got["reason"]) and "SESSION_FLAT" in (str(ref["reason"]), str(got["reason"])):
            counts["session_close"] += 1
        supports = (float(ref["initial_support"]), float(ref["final_support"]), float(got["initial_support"]), float(got["final_support"]))
        if any(abs(value - supports[0]) > 1e-6 for value in supports):
            counts["support"] += 1
        if int(ref["reconfirm_raises"]) != int(got["reconfirm_raises"]):
            counts["reconfirm"] += 1
        if abs(float(ref["pnl_yen"]) - float(got["pnl_yen"])) > 1e-6:
            counts["pnl"] += 1
        if bool(ref["reentry"]) != bool(got["reentry"]):
            counts["reentry"] += 1
    return {
        "counts": counts,
        "exact_matched_trades": len(set(left) & set(right)),
        "reference_only": len(set(left) - set(right)),
        "implementation_only": len(set(right) - set(left)),
    }


def publish() -> dict[str, Any]:
    contract = run_contract()
    windows = load_windows()
    classified = classify_intervals(windows)
    cache = OUT / "_scan.pkl"
    payload = pickle.loads(cache.read_bytes()) if cache.exists() else None
    if payload is None:
        payload = scan()
        OUT.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(pickle.dumps(payload))
    compared = _compare(payload["reference"], payload["implementation"], dict(payload["signal_mismatches"]))
    support_ok = all(abs(float(row["initial_support"]) - float(row["final_support"])) <= 1e-6 for row in payload["reference"] + payload["implementation"])
    parity = all(value == 0 for value in compared["counts"].values()) and compared["reference_only"] == 0 and compared["implementation_only"] == 0
    if not contract["all_pass"]:
        verdict, nxt = "FIXED_ENTRY_SUPPORT_REGISTRATION_GATE_CONTRACT_FAIL_V1", "REPAIR_REGISTRATION_GATE_ONLY"
    elif not parity:
        verdict, nxt = "FIXED_ENTRY_SUPPORT_RUNTIME_BINDING_PARITY_FAIL_V1", "REPAIR_RUNTIME_BINDING_PARITY_ONLY"
    else:
        verdict, nxt = "FIXED_ENTRY_SUPPORT_RUNTIME_BINDING_EVIDENCE_SUFFICIENT_V1", "FIXED_ENTRY_SUPPORT_PAPER_PREFLIGHT_V1"
    report = {
        "study": "FIXED_ENTRY_SUPPORT_RUNTIME_BINDING_EVIDENCE_RECOVERY_V1",
        "verdict": verdict,
        "next": nxt,
        "candidate_status": [
            "CURRENT_STRATEGY_CANDIDATE" if parity and contract["all_pass"] else "CURRENT_RESEARCH_SELECTED_CANDIDATE",
            "EXISTING_DATA_VALIDATED",
            "IMPLEMENTATION_PARITY_PROVEN",
            "VERIFIED_WINDOW_RUNTIME_BINDING_PARITY_PROVEN" if parity else "VERIFIED_WINDOW_RUNTIME_BINDING_PARITY_NOT_PROVEN",
            "REGISTRATION_GATE_CONTRACT_PROVEN" if contract["all_pass"] else "REGISTRATION_GATE_CONTRACT_NOT_PROVEN",
            "FULL_HISTORICAL_RUNTIME_BINDING_UNRESOLVED",
        ],
        "full_historical_runtime_parity": "unresolved",
        "historical_evidence_gap_remains": True,
        "gap_is_missing_observability": True,
        "paper_preflight_allowed": bool(parity and contract["all_pass"]),
        "strategy_changed": False,
        "new_strategy_optimization": False,
        "new_market_data": False,
        "paper_executed": False,
        "submit_cancel_live": [0, 0, 0],
        "known_limitation": "FIXED_SUPPORT_IMPROVEMENT_SYMBOL_CONCENTRATED",
        "exact_recovered": [
            {
                "interval": f"{window['date']} {window['start']} -> {window['end']}",
                "status": "EXACT_RECOVERED",
                "source": window["source"],
                "timestamp": window["start"],
            }
            for window in windows
        ],
        "intervals": classified,
        "verified_registration_windows": len(windows),
        "verified_signals": payload["verified_signals"],
        "reference": _stats(payload["reference"]),
        "implementation": _stats(payload["implementation"]),
        "verified_window_trade_by_trade_parity": bool(parity),
        "active_support_invariant": bool(support_ok),
        "first_divergence": payload["first_divergence"],
        "contract": contract,
        **compared,
        "windows": [{key: window[key] for key in ("date", "start", "end", "symbol_n", "source", "confirmations")} for window in windows],
    }
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, indent=2, sort_keys=True), encoding="utf-8")
    (OUT / "report.md").write_text(_markdown(report), encoding="utf-8")
    wb = Workbook()
    wb.active.title = "summary"
    wb.active.append(["verdict", verdict, nxt])
    wb.active.append(["parity", parity])
    wb.active.append(["contract", contract["all_pass"]])
    ws = wb.create_sheet("windows")
    ws.append(["date", "start", "end", "symbol_n", "confirmations", "source"])
    for window in report["windows"]:
        ws.append([window["date"], window["start"], window["end"], window["symbol_n"], window["confirmations"], window["source"]])
    iv = wb.create_sheet("intervals")
    iv.append(["interval", "status", "reason", "source"])
    for row in classified:
        iv.append([row["interval"], row["status"], row["reason"], row["source"]])
    wb.save(OUT / "audit.xlsx")
    cache.unlink(missing_ok=True)
    print(f"VERDICT {verdict} NEXT {nxt} parity={parity} contract={contract['all_pass']}", flush=True)
    return report


def _markdown(report: dict[str, Any]) -> str:
    ref, impl = report["reference"], report["implementation"]
    return "\n".join([
        f"# {report['study']}",
        "",
        f"VERDICT: {report['verdict']}",
        f"NEXT: {report['next']}",
        "",
        "Strategy implementation parity from the prior freeze remains in force.",
        f"Verified-window runtime parity: {report['verified_window_trade_by_trade_parity']}.",
        f"Registration-gate contract: {report['contract']['all_pass']}.",
        "Full historical runtime binding remains unresolved.",
        "",
        f"Verified windows {report['verified_registration_windows']}. Verified signals {report['verified_signals']}.",
        f"Reference {ref['trade_n']} / {ref['pnl_yen']} / {ref['pf']}",
        f"Implementation {impl['trade_n']} / {impl['pnl_yen']} / {impl['pf']}",
        f"Mismatches {json.dumps(report['counts'])}",
        f"First divergence {report['first_divergence']}",
        "",
        "Unresolved registration intervals were not filled from Dynamic40, desired symbols, generation claims, or a later readback.",
        "submit/cancel/live = 0/0/0. Paper was not executed.",
        "",
    ])


if __name__ == "__main__":
    publish()
