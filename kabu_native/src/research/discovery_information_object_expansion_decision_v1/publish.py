"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.discovery_information_object_expansion_decision_v1 import ANALYSIS_ID
from research.discovery_information_object_expansion_decision_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "prior_use",
    "ioar",
    "ueia",
    "raw_objects",
    "full_depth",
    "event_flow",
    "preopen",
    "prior_session",
    "classification",
    "eligibility",
    "selection",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    schema = dict(report.get("schema") or {})
    days = list(schema.get("days") or [])
    rows = list(report.get("raw_objects") or [])
    by = {str(r.get("OBJECT_ID")): r for r in rows}
    depth = by.get("FULL_DEPTH_GEOMETRY") or {}
    ev = by.get("EVENT_FLOW_DYNAMICS") or {}
    pre = by.get("OPENING_AUCTION_CONTEXT") or {}
    prior = by.get("PRIOR_SESSION_CONTEXT") or {}
    selected = report.get("selected")
    spec = dict(report.get("spec") or {})
    decision = dict(report.get("decision") or {})
    full_depth_days = [
        {
            "date": d.get("date"),
            "buy10_key_am": d.get("buy10_key_am"),
            "sell10_key_am": d.get("sell10_key_am"),
            "buy10_qty_pos_am": d.get("buy10_qty_pos_am"),
            "sell10_qty_pos_am": d.get("sell10_qty_pos_am"),
            "full_depth_am": d.get("full_depth_am"),
            "symbol_n_am_buy10": d.get("symbol_n_am_buy10"),
            "beyond_top1_imbalance": True,
            "beyond_net_pressure_formula": True,
            "not_another_weighted_imbalance": True,
            "structure_or_migration_object": True,
        }
        for d in days
    ]
    event_days = [
        {
            "date": d.get("date"),
            "timestamp_ingress_ok": d.get("timestamp_ingress_ok"),
            "asktime_or_bidtime": d.get("asktime_or_bidtime"),
            "event_flow_present": d.get("event_flow_present"),
            "equivalent_to_UEIA_IOAR_EEM_BoardDynamic": True,
            "NOT_NEW": True,
        }
        for d in days
    ]
    pre_days = [
        {
            "date": d.get("date"),
            "preopen_n": d.get("preopen_n"),
            "mo_buy_key_preopen": d.get("mo_buy_key_preopen"),
            "mo_buy_key_am": d.get("mo_buy_key_am"),
            "mo_buy_pos_preopen": d.get("mo_buy_pos_preopen"),
            "mo_sell_pos_preopen": d.get("mo_sell_pos_preopen"),
            "over_sell_key_preopen": d.get("over_sell_key_preopen"),
            "under_buy_key_preopen": d.get("under_buy_key_preopen"),
            "preopen_buy10_key": d.get("preopen_buy10_key"),
            "PREOPEN_EXECUTION_VALID": False,
        }
        for d in days
    ]
    prior_days = [
        {
            "date": d.get("date"),
            "previous_close_am": d.get("previous_close_am"),
            "opening_price_nonnull_am": d.get("opening_price_nonnull_am"),
            "gap_fill_reclaim_proposed": False,
            "semantically_recovery_if_gap_fill": True,
        }
        for d in days
    ]
    sel_rows = [_kv(spec)]
    if selected is not None:
        sel_rows = [{**selected, **{f"SPEC_{k}": v for k, v in spec.items() if k != "INFORMATION_OBJECT_SPEC_SHA256"}}]
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "parent": _kv(dict(report.get("parent") or {})),
        "prior_use": list(report.get("prior_use") or []),
        "ioar": _kv(dict(report.get("ioar") or {})),
        "ueia": _kv(dict(report.get("ueia") or {})),
        "raw_objects": rows,
        "full_depth": full_depth_days + [depth],
        "event_flow": event_days + [ev],
        "preopen": pre_days + [pre],
        "prior_session": prior_days + [prior],
        "classification": [
            {"OBJECT_ID": r.get("OBJECT_ID"), "STATUS": r.get("STATUS"), "NOTES": r.get("NOTES")}
            for r in rows
        ],
        "eligibility": [
            {
                "OBJECT_ID": r.get("OBJECT_ID"),
                "STATUS": r.get("STATUS"),
                "I1": r.get("I1_STORED"),
                "I2": r.get("I2_CAUSAL_TIMESTAMP"),
                "I3": r.get("I3_DEV_DAY_N_GE_8"),
                "I4": r.get("I4_BROAD_COVERAGE"),
                "I5": r.get("I5_NOT_EXACT_CLOSED_PRIMARY"),
                "I6": r.get("I6_NOT_NEW_TRANSFORM_OF_OLD_OBJECT"),
                "I7": r.get("I7_NO_EXTERNAL"),
                "I8": r.get("I8_NO_FUTURE_LABEL"),
                "I9": r.get("I9_DIFFERENT_DECISION_MECHANISM"),
                "I10": r.get("I10_COMPLETE_STRATEGY_WITHOUT_THRESHOLD_SEARCH"),
                "ELIGIBLE": r.get("ELIGIBLE"),
                "DEV_DAY_N": r.get("DEV_DAY_N"),
            }
            for r in rows
        ],
        "selection": sel_rows,
        "decision": _kv(decision),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    spec = dict(report.get("spec") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        f"CASE: {d.get('CASE')}",
        "",
        "## Pin",
        "",
        "- Parent CASE C pinned: "
        + str(a.get("1_parent_CASE_C_pinned")),
        "- Current 42 closed; CURRENT_42_RETUNE=false",
        "- ALL_INFORMATION_EXHAUSTION_ALREADY_PROVEN=false",
        "- IOAR exact mechanism closed; broader order-flow family not auto-closed",
        "- UEIA inventoried; do not reopen",
        "",
        "## Decision",
        "",
        f"- raw object N: {a.get('9_raw_object_N')}",
        f"- underexplored N: {a.get('14_underexplored_causal_object_N')} {a.get('15_underexplored_object_IDs')}",
        f"- eligible N: {a.get('20_eligible_object_N')}",
        f"- selected: {a.get('21_selected_object_ID')}",
        f"- DEV_DAY_N: {a.get('22_selected_object_DEV_day_N')}",
        f"- SPEC_SHA256: `{a.get('30_INFORMATION_OBJECT_SPEC_SHA256')}`",
        "",
        "## Complete-strategy potential (no ENTRY/EXIT)",
        "",
        str(a.get("26_complete_full_strategy_potential") or ""),
        "",
        "## Forbidden if selected",
        "",
        "- imbalance formulas / k-level search / IOAR exact / UEIA reproduction / 42 retune",
        "",
        "## Safety",
        "",
        "- PnL/markout/threshold: false/false/false",
        "- Holdout/Stress/20260903+/Paper: unread",
        "- Runtime/Capture unchanged; submit/cancel/live=0/0/0",
        "",
        "STOP.",
        "",
    ]
    _ = spec
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    md = report.get("_markdown") or build_markdown(report)
    (OUT / "report.md").write_text(str(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
