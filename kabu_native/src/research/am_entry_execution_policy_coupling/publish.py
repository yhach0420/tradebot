"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_execution_policy_coupling import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_entry_execution_policy_coupling"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "SelectionSets",
    "PFill",
    "W5Cause",
    "Recovery",
    "WaitCounterfactual",
    "RecoveredQuality",
    "FillLoss",
    "Coupling",
    "DayRobustness",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# AM ENTRY / EXECUTION POLICY COUPLING REASSESSMENT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. Frozen DIRECT / FILL_ONLY OOF selection. Frozen waits W1/W2/W5/W10.",
            "No new model. No Stage2. No wait search. No W10 adoption.",
            "Runtime WAIT_SEC remains 1.0. DEV_WAIT_SEC = 5.0 diagnostic only.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"DIRECT_ONLY_P_FILL_MEDIAN: {_fmt(g.get('DIRECT_ONLY_P_FILL_MEDIAN'))}",
            f"FILL_ONLY_ONLY_P_FILL_MEDIAN: {_fmt(g.get('FILL_ONLY_ONLY_P_FILL_MEDIAN'))}",
            "",
            f"DIRECT_ONLY_W5_NONFILL_N: {_fmt(g.get('DIRECT_ONLY_W5_NONFILL_N'))}",
            f"DIRECT_ONLY_NO_VALID_BOARD_N: {_fmt(g.get('DIRECT_ONLY_NO_VALID_BOARD_N'))}",
            f"DIRECT_ONLY_NO_CROSS_N: {_fmt(g.get('DIRECT_ONLY_NO_CROSS_N'))}",
            "",
            f"DIRECT_ONLY_RECOVER_RATE_W10: {_fmt(g.get('DIRECT_ONLY_RECOVER_RATE_W10'))}",
            f"DIRECT_ONLY_NEVER_CROSS_10S_RATE: {_fmt(g.get('DIRECT_ONLY_NEVER_CROSS_10S_RATE'))}",
            "",
            f"FILL_ONLY_FILL_RATE_W1: {_fmt(g.get('FILL_ONLY_FILL_RATE_W1'))}",
            f"FILL_ONLY_FILL_RATE_W2: {_fmt(g.get('FILL_ONLY_FILL_RATE_W2'))}",
            f"FILL_ONLY_FILL_RATE_W5: {_fmt(g.get('FILL_ONLY_FILL_RATE_W5'))}",
            f"FILL_ONLY_FILL_RATE_W10: {_fmt(g.get('FILL_ONLY_FILL_RATE_W10'))}",
            "",
            f"DIRECT_FILL_RATE_W1: {_fmt(g.get('DIRECT_FILL_RATE_W1'))}",
            f"DIRECT_FILL_RATE_W2: {_fmt(g.get('DIRECT_FILL_RATE_W2'))}",
            f"DIRECT_FILL_RATE_W5: {_fmt(g.get('DIRECT_FILL_RATE_W5'))}",
            f"DIRECT_FILL_RATE_W10: {_fmt(g.get('DIRECT_FILL_RATE_W10'))}",
            "",
            f"DIRECT_MINUS_FILL_ONLY_W10: {_fmt(g.get('DIRECT_MINUS_FILL_ONLY_W10'))}",
            "",
            f"DIRECT_EXEC_U_DELTA_W10: {_fmt(g.get('DIRECT_EXEC_U_DELTA_W10'))}",
            f"DIRECT_EXEC_D_DELTA_W10: {_fmt(g.get('DIRECT_EXEC_D_DELTA_W10'))}",
            "",
            f"DIRECT_RECOVERED_N: {_fmt(g.get('DIRECT_RECOVERED_N'))}",
            f"RECOVERED_COND_U: {_fmt(g.get('RECOVERED_COND_U'))}",
            f"RECOVERED_COND_D: {_fmt(g.get('RECOVERED_COND_D'))}",
            f"RECOVERED_TIME_TO_FILL_MEDIAN: {_fmt(g.get('RECOVERED_TIME_TO_FILL_MEDIAN'))}",
            "",
            f"FILL_LOSS_LATE_RECOVERABLE_N: {_fmt(g.get('FILL_LOSS_LATE_RECOVERABLE_N'))}",
            f"FILL_LOSS_NEVER_CROSS_N: {_fmt(g.get('FILL_LOSS_NEVER_CROSS_N'))}",
            f"FILL_LOSS_NO_VALID_BOARD_N: {_fmt(g.get('FILL_LOSS_NO_VALID_BOARD_N'))}",
            "",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No wait change. No model change. No Exact. No PnL. No Stage2 restart.",
            "No W5/W10 runtime adoption. Runtime WAIT_SEC remains 1.0.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
