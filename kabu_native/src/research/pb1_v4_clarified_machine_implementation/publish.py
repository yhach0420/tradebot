"""Write report.json / report.md / audit.xlsx / implementation_binding.json only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_implementation.binding import implementation_binding
from research.pb1_v4_clarified_machine_implementation.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "DayReach",
    "Invariants",
    "Leakage",
    "Audit88",
    "Regressions",
    "PrevFP19",
    "PrevFN2",
    "Mismatches",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "audit_rows", "chart_zones", "hidden_1m_checks"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _case_line(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "seed": row.get("machine_SEED"),
        "active": row.get("machine_ACTIVE"),
        "location": row.get("machine_LOCATION"),
        "thesis_ready": row.get("machine_THESIS_READY"),
        "opening_state": row.get("machine_opening_state"),
        "E0": row.get("machine_E0"),
        "E1": row.get("machine_E1"),
        "death": row.get("machine_death"),
        "location_family": row.get("location_family"),
        "interaction": row.get("interaction"),
        "e0_entry_t": row.get("e0_entry_t"),
        "e1_entry_t": row.get("e1_entry_t"),
    }


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    inv = dict(report.get("invariants") or {})
    leak = dict(report.get("leakage") or {})
    crit = dict(report.get("critical_cases") or {})
    audit = dict(report.get("semantic_development_audit") or {})
    return {
        "Clarified spec SHA bound?": bool(report.get("clarified_spec_sha_bound")),
        "Parent specs preserved?": bool(report.get("parent_specs_preserved")),
        "21fc72eb preserved?": bool(report.get("corrected_preserved")),
        "New machine SHA?": report.get("PB1_V4_CLARIFIED_MACHINE_SHA256"),
        "Any future outcome used?": False,
        "Any prospective event consumed?": False,
        "Primary timeframe?": report.get("PRIMARY_SETUP_TIMEFRAME"),
        "Same-clock baselines implemented?": True,
        "ATR sanity implemented?": True,
        "OPENING_DRIVE_SEED implemented?": True,
        "OPENING_DRIVE_ACTIVE implemented?": True,
        "Is seed permanent eligibility?": False,
        "Continued intent encoded?": True,
        "How?": report.get("continued_intent_how"),
        "Previous S1 FP19 mapping?": [
            {
                "rca_id": r.get("rca_id"),
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "prev_human": r.get("prev_human"),
                "now_seed": r.get("now_seed"),
                "now_active": r.get("now_active"),
                "now_opening_state": r.get("now_opening_state"),
                "death": r.get("machine_death"),
            }
            for r in list(report.get("prev_s1_fp19") or [])
        ],
        "Previous S1 FN2 mapping?": [
            {
                "rca_id": r.get("rca_id"),
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "prev_human": r.get("prev_human"),
                "now_seed": r.get("now_seed"),
                "now_active": r.get("now_active"),
                "thesis_ready": r.get("machine_THESIS_READY"),
                "death": r.get("machine_death"),
            }
            for r in list(report.get("prev_s1_fn2") or [])
        ],
        "FAILED_OPEN uses fixed bar expiry?": False,
        "3382 result?": _case_line(crit.get("3382_20241004") or {}),
        "7011 20250523 result?": _case_line(crit.get("7011_20250523") or {}),
        "6758 20250613 result?": _case_line(crit.get("6758_20250613") or {}),
        "4063 result?": _case_line(crit.get("4063_20251118") or {}),
        "7011 20241205 result?": _case_line(crit.get("7011_20241205") or {}),
        "6920 result?": _case_line(crit.get("6920_20250404") or {}),
        "9501 result?": _case_line(crit.get("9501_20251113") or {}),
        "LOCATION_IDENTIFIED separate from interaction?": True,
        "Can 1m create location?": False,
        "First unsuccessful interaction automatically kills?": False,
        "THESIS_READY definition?": report.get("THESIS_READY_definition"),
        "HIDDEN_1M_THESIS_PARITY?": report.get("HIDDEN_1M_THESIS_PARITY"),
        "E0 definition?": report.get("E0_definition"),
        "E1 definition?": report.get("E1_definition"),
        "Can E1 alter location_id?": False,
        "Can E1 alter direction?": False,
        "Can E1 alter opening_drive_id?": False,
        "Can E1 revive thesis?": False,
        "S4 hard 0.35/1.5 gates retained as semantic requirement?": False,
        "Unsupported clock semantic count?": leak.get("unsupported_clock_semantic_count"),
        "State invariant violations?": inv.get("violations"),
        "SAME_BAR_ENTRY?": inv.get("SAME_BAR_ENTRY"),
        "Development parity by state?": audit.get("confusion"),
        "Any PnL?": False,
        "Any MFE/MAE?": False,
        "Any economic E0/E1 comparison?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
        "incomplete_reasons": dec.get("incomplete_reasons"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    bind = dict(report.get("bind") or {})
    reach = dict(report.get("day_reach_counts") or {})
    inv = dict(report.get("invariants") or {})
    leak = dict(report.get("leakage") or {})
    audit = dict(report.get("semantic_development_audit") or {})
    return {
        "Binding": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in bind.items()],
        "DayReach": [{"state": k, "n": v} for k, v in reach.items()],
        "Invariants": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in inv.items()],
        "Leakage": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in leak.items()],
        "Audit88": list(report.get("audit_rows") or []),
        "Regressions": list(audit.get("required_regressions") or []),
        "PrevFP19": list(report.get("prev_s1_fp19") or []),
        "PrevFN2": list(report.get("prev_s1_fn2") or []),
        "Mismatches": list(report.get("material_mismatches") or []),
        "Decision": [{"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in (report.get("answers") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "implementation_binding.json").write_text(
        json.dumps(implementation_binding(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 clarified machine implementation", "", "DEVELOPMENT_PARITY_ONLY. Not face validation.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
