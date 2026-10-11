"""Causal AUTO50 rebuild. Runtime code and older research reports stay untouched.

A target date D uses only the previous JPX session. Each symbol must have its
own bar on that date. Later stock splits are inverted so the traded price, not
today's back-adjusted close, is what the score and the G check see.
"""
from __future__ import annotations

import csv
import json
import math
import subprocess
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

from openpyxl import Workbook

from research.alphanumeric_universe_restoration_v1.normalize import (
    ResearchCodeRejected,
    research_normalize_code,
    research_yahoo_ticker,
)
from research.auto50_research_capture.mode import eligible_feature
from research.daytrade50_capital_efficient_selector_v2.features import (
    execution_suitability,
    selection_calendar,
    snapshot_days,
)
from research.daytrade50_capital_efficient_selector_v2.lineage import previous_jpx_session
from research.daytrade50_capital_efficient_selector_v2.policy import one_lot_notional
from research.daytrade50_capital_efficient_selector_v2.selector import select_top50
from research.daytrade50_capital_efficient_selector_v2.selector_v3 import select_top50_v3
from research.daytrade50_capital_efficient_selector_v2.target_blind_g1m import (
    _aggregate,
    _book_metrics,
    _feature,
    _series,
    _target_metrics,
)

NATIVE = Path(__file__).resolve().parents[3]
ROOT = NATIVE.parent
MASTER = ROOT / "data" / "jpx" / "tradable_symbols.csv"
REPORTS = NATIVE / "results" / "reports"
BUDGET = NATIVE / "runtime" / "system_gross_exposure_budget.json"
V2_REPORT = (
    NATIVE
    / "results"
    / "research"
    / "daytrade50_capital_efficient_selector_v2"
    / "target_blind_validation_cap500k_lev2x_g1m"
    / "report.json"
)
OUT = NATIVE / "results" / "research" / "auto50_causal_freshness_reconciliation_v1"
G = 1_000_000.0
FOCUS = ("285A", "278A", "6857", "5803")
PAPER_MARKERS = (
    "paper_trade_checked_runner",
    "market_ingress_service",
    "v1r_paper_primary_launcher",
    "run_core10_dynamic40",
    "run_small_paper_pilot",
)


def paper_alive() -> bool:
    script = (
        "Get-CimInstance Win32_Process -Filter \"Name='python.exe'\" | "
        "ForEach-Object { $_.CommandLine }"
    )
    completed = subprocess.run(
        ["powershell", "-NoProfile", "-Command", script],
        capture_output=True,
        text=True,
        check=False,
    )
    text = completed.stdout or ""
    return any(marker in text for marker in PAPER_MARKERS)


def load_master() -> tuple[list[dict[str, str]], list[str]]:
    active: list[dict[str, str]] = []
    rejected: list[str] = []
    with MASTER.open(encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle):
            flag = str(raw.get("is_active") or "").strip().lower()
            if flag not in {"1", "true", "yes"}:
                continue
            try:
                code = research_normalize_code(raw.get("symbol") or "")
            except ResearchCodeRejected as exc:
                rejected.append(str(exc))
                continue
            active.append(
                {
                    "code": code,
                    "exchange": str(raw.get("exchange") or "1"),
                    "market": str(raw.get("market") or "").strip().lower(),
                    "name": str(raw.get("name") or ""),
                }
            )
    return active, rejected


def capital_ok() -> dict[str, Any]:
    body = json.loads(BUDGET.read_text(encoding="utf-8"))
    base = body.get("BASE_CAPITAL_YEN")
    leverage = body.get("MAX_LEVERAGE_MULTIPLIER")
    requested = body.get("REQUESTED_GROSS_EXPOSURE_BUDGET_YEN")
    return {
        "match": base == 500000 and str(leverage) in {"2", "2.0"} and requested == 1000000,
        "BASE_CAPITAL_YEN": base,
        "MAX_LEVERAGE_MULTIPLIER": leverage,
        "G": requested,
        "POLICY_INSTANCE_SHA": body.get("POLICY_INSTANCE_SHA"),
    }


def close_match(left: float, right: float) -> bool:
    return abs(left - right) <= max(0.05, abs(right) * 1e-4)


def volume_match(left: float, right: float) -> bool:
    return abs(left - right) <= max(1.0, abs(right) * 0.002)


def download_panel(symbols: Sequence[str], start: date, end: date) -> tuple[dict[str, list[dict[str, Any]]], list[str]]:
    import yfinance as yf

    panel: dict[str, list[dict[str, Any]]] = {}
    failed: list[str] = []

    def parse_view(view: Any) -> list[dict[str, Any]]:
        if view is None or getattr(view, "empty", True):
            return []
        columns = set(getattr(view, "columns", []))
        if "Close" not in columns or "Stock Splits" not in columns:
            return []
        raw_rows = []
        for idx, row in view.iterrows():
            try:
                close = float(row["Close"])
                high = float(row["High"])
                low = float(row["Low"])
                volume = float(row["Volume"])
                split = float(row["Stock Splits"])
                adj = float(row["Adj Close"]) if "Adj Close" in columns else close
            except (TypeError, ValueError, KeyError):
                continue
            if close != close or close <= 0 or high != high or low != low or volume != volume:
                continue
            if split != split or split < 0:
                split = 0.0
            raw_rows.append(
                {
                    "date": idx.date().strftime("%Y%m%d"),
                    "raw_close": close,
                    "raw_high": high,
                    "raw_low": low,
                    "raw_volume": volume,
                    "split": split if split > 0 and abs(split - 1.0) > 1e-9 else 0.0,
                    "adj_close": adj if adj == adj else close,
                }
            )
        raw_rows.sort(key=lambda item: item["date"])
        factor_after = 1.0
        restored: list[dict[str, Any]] = []
        for item in reversed(raw_rows):
            factor = factor_after
            if item["split"]:
                factor_after *= item["split"]
            restored.append(
                {
                    "source_bar_date": item["date"],
                    "close": item["raw_close"] * factor,
                    "high": item["raw_high"] * factor,
                    "low": item["raw_low"] * factor,
                    "volume": item["raw_volume"] / factor if factor else item["raw_volume"],
                    "raw_close": item["raw_close"],
                    "factor": factor,
                    "adj_close": item["adj_close"],
                    "dividend_adjusted_diff": abs(item["adj_close"] - item["raw_close"]) > max(0.05, abs(item["raw_close"]) * 0.001),
                }
            )
        restored.sort(key=lambda item: item["source_bar_date"])
        return restored

    batch_size = 80
    batches = (len(symbols) + batch_size - 1) // batch_size
    for offset in range(0, len(symbols), batch_size):
        batch = list(symbols[offset : offset + batch_size])
        print(f"daily bars {offset // batch_size + 1}/{batches}", flush=True)
        try:
            frame = yf.download(
                batch,
                start=start.isoformat(),
                end=end.isoformat(),
                interval="1d",
                group_by="ticker",
                auto_adjust=False,
                actions=True,
                progress=False,
                threads=False,
            )
        except Exception:
            frame = None
        if frame is None or getattr(frame, "empty", True):
            failed.extend(symbol.split(".")[0] for symbol in batch)
            continue
        for symbol in batch:
            key = symbol.split(".")[0]
            try:
                view = frame[symbol]
            except Exception:
                failed.append(key)
                continue
            rows = parse_view(view)
            if rows:
                panel[key] = rows
            else:
                failed.append(key)
    return panel, failed


def bar_on(rows: Sequence[Mapping[str, Any]], day: str) -> Optional[Mapping[str, Any]]:
    for row in rows:
        if row["source_bar_date"] == day:
            return row
    return None


def buckets_for(rows: Sequence[Mapping[str, Any]]) -> dict[int, list[Mapping[str, Any]]]:
    found: dict[int, list[Mapping[str, Any]]] = defaultdict(list)
    for row in rows:
        found[round(float(row["close"]))].append(row)
    return found


def match_stored(bucket: Mapping[int, Sequence[Mapping[str, Any]]], close: float, volume: float) -> list[Mapping[str, Any]]:
    hits: list[Mapping[str, Any]] = []
    key = round(close)
    for probe in (key - 1, key, key + 1):
        for row in bucket.get(probe, ()):
            if close_match(float(row["close"]), close) and volume_match(float(row["volume"]), volume):
                hits.append(row)
    return hits


def split_label(factor: float) -> str:
    if factor > 1.0000001:
        return "SPLIT"
    if factor < 0.9999999:
        return "REVERSE_SPLIT"
    return "CLEAN"


def read_feature_rows(day: str) -> list[dict[str, str]]:
    path = REPORTS / f"features_{day}.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def audit_stored_prices(panel: Mapping[str, Sequence[Mapping[str, Any]]], sessions: Sequence[str]) -> dict[str, Any]:
    counts: Counter[str] = Counter()
    examples: list[dict[str, Any]] = []
    asof: dict[tuple[str, str], float] = {}
    conflicts: set[tuple[str, str]] = set()
    indexed = {symbol: buckets_for(rows) for symbol, rows in panel.items()}
    for day in snapshot_days(REPORTS):
        required = previous_jpx_session(sessions, day)
        for raw in read_feature_rows(day):
            code = str(raw.get("symbol") or "").split(".")[0].split("@")[0]
            if not code.isdigit():
                continue
            try:
                stored_close = float(raw["close"])
                stored_volume = float(raw["volume"])
            except (TypeError, ValueError, KeyError):
                counts["UNREADABLE"] += 1
                continue
            rows = panel.get(code)
            if not rows:
                counts["MISSING_SERIES"] += 1
                continue
            hits = match_stored(indexed[code], stored_close, stored_volume)
            if not hits:
                counts["PROVIDER_REVISION"] += 1
                if len(examples) < 40:
                    examples.append({"day": day, "symbol": code, "class": "PROVIDER_REVISION", "stored_close": stored_close})
                continue
            if len(hits) > 1:
                counts["AMBIGUOUS"] += 1
                continue
            bar = hits[0]
            label = split_label(float(bar["factor"]))
            if label == "CLEAN" and not close_match(float(bar["raw_close"]), stored_close):
                label = "PROVIDER_REVISION"
            counts[label] += 1
            if bar.get("dividend_adjusted_diff"):
                counts["DIVIDEND_ADJUSTMENT_PRESENT"] += 1
            matched = str(bar["source_bar_date"])
            if required and matched != required:
                counts["FETCH_TIME_DIFF"] += 1
            elif required and matched == required:
                counts["MATCHED_REQUIRED_SESSION"] += 1
            key = (code, matched)
            previous = asof.get(key)
            if previous is None:
                asof[key] = stored_close
            elif not close_match(previous, stored_close):
                conflicts.add(key)
                counts["CONFLICTING_STORED_CLOSE"] += 1
            if label in {"SPLIT", "REVERSE_SPLIT"} and len(examples) < 40:
                examples.append(
                    {
                        "day": day,
                        "symbol": code,
                        "class": label,
                        "stored_close": stored_close,
                        "raw_close": bar["raw_close"],
                        "restored_close": bar["close"],
                        "factor": bar["factor"],
                        "matched_session": matched,
                        "required_session": required,
                    }
                )
    explained = counts["CLEAN"] + counts["SPLIT"] + counts["REVERSE_SPLIT"] + counts["AMBIGUOUS"]
    unexplained = counts["PROVIDER_REVISION"]
    return {
        "counts": dict(counts),
        "explained_rate": (explained / (explained + unexplained)) if (explained + unexplained) else None,
        "examples": examples,
        "conflicts": conflicts,
        "witness_6379_20260717": _witness(panel),
    }


def _witness(panel: Mapping[str, Sequence[Mapping[str, Any]]]) -> dict[str, Any]:
    bar = bar_on(panel.get("6379") or [], "20260717")
    if bar is None:
        return {"pass": False}
    return {
        "pass": close_match(float(bar["close"]), 2488.0) and abs(float(bar["factor"]) - 3.0) < 1e-6,
        "restored_close": bar["close"],
        "raw_close": bar["raw_close"],
        "factor": bar["factor"],
        "stored_close_seen_on_20260722_file": 2488.0,
    }


def score_bar(code: str, meta: Mapping[str, str], bar: Mapping[str, Any]) -> Optional[dict[str, Any]]:
    close = float(bar["close"])
    high = float(bar["high"])
    low = float(bar["low"])
    volume = float(bar["volume"])
    if close <= 0 or high < low or volume < 0:
        return None
    atr = round((high - low) / close * 100.0, 6)
    trading = volume * close
    item = eligible_feature(
        {
            "symbol": code,
            "exchange": meta.get("exchange") or "1",
            "market": meta.get("market") or "",
            "close": close,
            "volume": volume,
            "trading_value": trading,
            "atr_pct": atr,
        }
    )
    if item is None:
        return None
    item["factor"] = float(bar["factor"])
    item["split_status"] = split_label(float(bar["factor"]))
    item["source_bar_date"] = str(bar["source_bar_date"])
    return item


def rank_items(items: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    ordered = [dict(item) for item in items]
    ordered.sort(key=lambda item: (-float(item["score"]), str(item["symbol"])))
    return ordered


def old_top50(day: str) -> list[dict[str, Any]]:
    ranked = []
    for raw in read_feature_rows(day):
        item = eligible_feature(raw)
        if item is not None:
            ranked.append(item)
    ranked.sort(key=lambda item: (-float(item["score"]), str(item["symbol"])))
    return ranked[:50]


def book_from_symbols(
    symbols: Sequence[str],
    closes: Mapping[str, float],
    targets: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    notionals = {symbol: one_lot_notional(float(closes[symbol])) for symbol in symbols if symbol in closes}
    suits = {}
    for symbol, price in closes.items():
        suit = execution_suitability(float(price))
        if suit is not None:
            suits[symbol] = suit
    return _book_metrics(symbols, notionals, suits, targets, G)


def position(ranked: Sequence[Mapping[str, Any]], code: str) -> Optional[int]:
    for index, item in enumerate(ranked, start=1):
        if item["symbol"] == code:
            return index
    return None


def run_days(
    days: Sequence[str],
    panel: Mapping[str, Sequence[Mapping[str, Any]]],
    meta: Mapping[str, Mapping[str, str]],
    sessions: Sequence[str],
    conflicts: set[tuple[str, str]],
    alpha: set[str],
) -> dict[str, Any]:
    series = {symbol: _series(rows) for symbol, rows in panel.items() if rows}
    daily = []
    focus = []
    books: dict[str, list[dict[str, Any]]] = {"old": [], "raw": [], "g1m": [], "v2": [], "v3": []}
    for day in days:
        required = previous_jpx_session(sessions, day)
        targets = {}
        for symbol, item in series.items():
            target = _target_metrics(item, day)
            if target is not None:
                targets[symbol] = target
        scored = []
        missing_numeric = 0
        missing_alpha = 0
        conflict_excluded = 0
        valid_numeric = 0
        valid_alpha = 0
        if required is None:
            daily.append({"target_trading_date": day, "official": False, "reason": "NO_PREVIOUS_SESSION"})
            continue
        for code, info in meta.items():
            if (code, required) in conflicts:
                conflict_excluded += 1
                continue
            bar = bar_on(panel.get(code) or [], required)
            if bar is None:
                if code in alpha:
                    missing_alpha += 1
                else:
                    missing_numeric += 1
                continue
            item = score_bar(code, info, bar)
            if item is None:
                continue
            if item["source_bar_date"] != required or item["source_bar_date"] >= day:
                continue
            scored.append(item)
            if code in alpha:
                valid_alpha += 1
            else:
                valid_numeric += 1
        raw = rank_items(scored)
        g_pool = [item for item in raw if one_lot_notional(float(item["prior_close"])) <= G]
        g_ranked = rank_items(g_pool)
        raw50 = raw[:50]
        g50 = g_ranked[:50]
        old = old_top50(day)
        v2_rows = []
        for symbol, item in series.items():
            info = meta.get(symbol)
            if info is None or (symbol, required) in conflicts:
                continue
            built = _feature(symbol, item, required, info.get("exchange") or "1", info.get("market") or "")
            if built is None:
                continue
            if built["SOURCE_BAR_DATE"] != required or built["WINDOW_END_SESSION"] >= day:
                continue
            v2_rows.append(built)
        v2 = select_top50(v2_rows, gross_budget_yen=G)
        v3 = select_top50_v3(v2_rows, gross_budget_yen=G)
        raw_ok = len(raw50) == 50 and all(item["source_bar_date"] == required for item in raw50)
        g_ok = len(g50) == 50 and all(item["source_bar_date"] == required for item in g50)
        g_violation = sum(1 for item in g50 if one_lot_notional(float(item["prior_close"])) > G)
        future = sum(1 for item in g50 if item["source_bar_date"] >= day)
        lag = 0 if future == 0 and all(item["source_bar_date"] == required for item in g50) else 1
        official = bool(g_ok and lag == 0 and g_violation == 0 and required < day)
        closes = {item["symbol"]: float(item["prior_close"]) for item in raw}
        old_closes = {item["symbol"]: float(item["prior_close"]) for item in old}
        raw_book = book_from_symbols([item["symbol"] for item in raw50], closes, targets)
        g_book = book_from_symbols([item["symbol"] for item in g50], closes, targets)
        old_book = book_from_symbols([item["symbol"] for item in old], old_closes, targets)
        v2_closes = {str(row["symbol"]): float(row["prior_close"]) for row in v2.get("rows") or []}
        v3_closes = {str(row["symbol"]): float(row["prior_close"]) for row in v3.get("rows") or []}
        v2_book = book_from_symbols(list(v2.get("symbols") or []), v2_closes, targets)
        v3_book = book_from_symbols(list(v3.get("symbols") or []), v3_closes, targets)
        if raw_book.get("range_median") is not None:
            books["raw"].append(raw_book)
        if official and g_book.get("range_median") is not None:
            books["g1m"].append(g_book)
        if old_book.get("range_median") is not None:
            books["old"].append(old_book)
        if v2.get("status") == "TOP50" and v2_book.get("range_median") is not None:
            books["v2"].append(v2_book)
        if v3.get("status") == "TOP50" and v3_book.get("range_median") is not None:
            books["v3"].append(v3_book)
        old_set = {item["symbol"] for item in old}
        raw_set = {item["symbol"] for item in raw50}
        g_set = {item["symbol"] for item in g50}
        split_counts = Counter(item["split_status"] for item in g50)
        daily.append(
            {
                "target_trading_date": day,
                "latest_required_session": required,
                "actual_source_session": required if g_ok else "",
                "lag_sessions": 0 if official else lag,
                "future_used": future,
                "valid_feature_count_numeric": valid_numeric,
                "valid_feature_count_alpha": valid_alpha,
                "split_adjustment_status": dict(split_counts),
                "missing_numeric": missing_numeric,
                "missing_alpha": missing_alpha,
                "conflict_excluded": conflict_excluded,
                "official_g1m": official,
                "raw_top50": raw_ok,
                "g1m_top50": len(g50) == 50,
                "g_violations": g_violation,
                "old_vs_raw_replacements": len(old_set.symmetric_difference(raw_set)) // 2,
                "old_vs_g1m_replacements": len(old_set.symmetric_difference(g_set)) // 2,
                "raw_alpha_selected": sorted(symbol for symbol in raw_set if symbol in alpha),
                "g1m_alpha_selected": sorted(symbol for symbol in g_set if symbol in alpha),
                "old_range": old_book.get("range_median"),
                "raw_range": raw_book.get("range_median"),
                "g1m_range": g_book.get("range_median"),
                "v2_range": v2_book.get("range_median"),
                "v3_range": v3_book.get("range_median"),
                "old_tv": old_book.get("trading_value_median"),
                "raw_tv": raw_book.get("trading_value_median"),
                "g1m_tv": g_book.get("trading_value_median"),
                "v2_tv": v2_book.get("trading_value_median"),
                "g1m_active": g_book.get("active_rate"),
                "v2_active": v2_book.get("active_rate"),
                "g1m_dead": g_book.get("dead_rate"),
                "v2_dead": v2_book.get("dead_rate"),
                "g1m_tick": g_book.get("tick_suitability_median"),
                "v2_tick": v2_book.get("tick_suitability_median"),
                "g1m_eligibility": g_book.get("one_lot_eligibility_rate"),
                "v2_eligibility": v2_book.get("one_lot_eligibility_rate"),
                "paired_g1m_minus_v2_range": (
                    None
                    if g_book.get("range_median") is None or v2_book.get("range_median") is None
                    else float(g_book["range_median"]) - float(v2_book["range_median"])
                ),
                "v2_status": v2.get("status"),
                "v3_status": v3.get("status"),
                "v3_reference_only": True,
            }
        )
        for code in FOCUS:
            info = meta.get(code, {})
            bar = bar_on(panel.get(code) or [], required)
            status = "CONFLICT" if (code, required) in conflicts else ("FEATURE_OK" if bar is not None else "MISSING_BAR")
            if code not in meta:
                status = "ABSENT_FROM_MASTER"
            lot = one_lot_notional(float(bar["close"])) if bar is not None else None
            raw_rank = position(raw, code)
            g_rank = position(g_ranked, code)
            if status != "FEATURE_OK":
                reason = status
            elif lot is not None and lot > G:
                reason = "ONE_LOT_ABOVE_G"
            elif g_rank is not None and g_rank <= 50:
                reason = "SELECTED"
            elif raw_rank is None:
                reason = "NOT_ELIGIBLE"
            else:
                reason = "RANK_OUTSIDE_TOP50"
            focus.append(
                {
                    "day": day,
                    "symbol": code,
                    "status": status,
                    "source_bar_date": None if bar is None else bar["source_bar_date"],
                    "lag_sessions": 0 if bar is not None and bar["source_bar_date"] == required else None,
                    "score": None if raw_rank is None else next(item["score"] for item in raw if item["symbol"] == code),
                    "raw_rank": raw_rank,
                    "raw_selected": bool(raw_rank and raw_rank <= 50),
                    "prior_close_x100": lot,
                    "g_eligible": bool(lot is not None and lot <= G and status == "FEATURE_OK"),
                    "g1m_selected": bool(g_rank and g_rank <= 50),
                    "g_exclusion_reason": reason,
                    "market": info.get("market"),
                }
            )
        print(f"day {day} official {official} valid {valid_numeric}+{valid_alpha}", flush=True)
    return {"daily": daily, "focus": focus, "books": {name: _aggregate(rows) for name, rows in books.items()}}


def judge(report: Mapping[str, Any]) -> tuple[str, str]:
    """Frozen before the ranks are interpreted. Already-seen days are not a new holdout."""
    price = report["price_audit"]
    daily = report["daily"]
    official = [row for row in daily if row.get("official_g1m")]
    integrity = (
        report["capital_match"]
        and report["witness_pass"]
        and price.get("explained_rate") is not None
        and float(price["explained_rate"]) >= 0.99
        and report["MASTER_ALPHANUMERIC_N"] == 162
        and report["parser_rejected"] == 0
        and len(official) == len(daily)
        and all(row.get("lag_sessions") == 0 and row.get("future_used") == 0 for row in official)
        and all(row.get("g_violations") == 0 for row in official)
        and all(row.get("g1m_eligibility") == 1 for row in official)
    )
    focus_ok = True
    for code in ("285A", "278A"):
        rows = [row for row in report["focus"] if row["symbol"] == code]
        if len(rows) != len(daily) or any(row["status"] == "ABSENT_FROM_MASTER" for row in rows):
            focus_ok = False
    books = report["books"]
    g1m = books.get("g1m") or {}
    v2 = books.get("v2") or {}
    range_ok = (
        g1m.get("range_median") is not None
        and v2.get("range_median") is not None
        and float(g1m["range_median"]) + 1e-9 >= float(v2["range_median"])
    )
    active_ok = g1m.get("active_rate") is not None and float(g1m["active_rate"]) >= 0.95
    dead_ok = g1m.get("dead_rate") is not None and float(g1m["dead_rate"]) <= 0.05
    if not (integrity and focus_ok and report["witness_pass"] and float(price.get("explained_rate") or 0) >= 0.99):
        return "AUTO50_G1M_CAUSAL_BASELINE_NOT_READY_V1", "REPAIR_CAUSAL_OR_PRICE_DATA"
    if not (range_ok and active_ok and dead_ok):
        return "AUTO50_G1M_CAUSAL_BASELINE_NOT_READY_V1", "DO_NOT_ADOPT_G1M_AS_RANGE_BASELINE"
    return "AUTO50_G1M_CAUSAL_BASELINE_READY_V1", "PREPARE_ECONOMIC_EVALUATION_PATH_WITHOUT_PNL_CLAIM"


def write_outputs(report: Mapping[str, Any]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    lines = [
        f"VERDICT: {report['verdict']}",
        f"NEXT: {report['next']}",
        "",
        "PAPER V16 は変更していません。Runtime の normalize_code も変更していません。",
        "CAPITAL: 500000 × 2.0 = 1000000",
        "submit/cancel/live: 0/0/0",
        f"used_as_independent_holdout: {report['used_as_independent_holdout']}",
        f"economic_superiority_claimed: {report['economic_superiority_claimed']}",
        f"target_days: {report['target_days']}",
        f"official_g1m_days: {report['official_g1m_days']}",
        f"lag_sessions_zero_days: {report['lag_zero_days']}",
        f"future_used: {report['future_used_total']}",
        f"price_explained_rate: {report['price_audit'].get('explained_rate')}",
        f"witness_6379: {report['witness_pass']}",
        f"FETCH_FAILED_N: {report['FETCH_FAILED_N']}",
        f"books: {json.dumps(report['books'], ensure_ascii=False)}",
        "V3 は不採用の参照です。Complete Strategy PnL は未検証です。",
    ]
    (OUT / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    book = Workbook()
    summary = book.active
    summary.title = "summary"
    for key, value in report.items():
        if key in {"daily", "focus", "price_audit"}:
            continue
        summary.append([key, json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else value])
    daily_sheet = book.create_sheet("daily")
    rows = report["daily"]
    if rows:
        daily_sheet.append(list(rows[0].keys()))
        for row in rows:
            daily_sheet.append([json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value for value in row.values()])
    focus_sheet = book.create_sheet("focus")
    if report["focus"]:
        focus_sheet.append(list(report["focus"][0].keys()))
        for row in report["focus"]:
            focus_sheet.append(list(row.values()))
    price = book.create_sheet("price")
    price.append(["item", "value"])
    for key, value in report["price_audit"]["counts"].items():
        price.append([key, value])
    price.append(["explained_rate", report["price_audit"].get("explained_rate")])
    examples = book.create_sheet("price_examples")
    if report["price_audit"]["examples"]:
        examples.append(list(report["price_audit"]["examples"][0].keys()))
        for row in report["price_audit"]["examples"]:
            examples.append(list(row.values()))
    failed = book.create_sheet("fetch_failed")
    failed.append(["symbol"])
    for symbol in report["fetch_failed_symbols"]:
        failed.append([symbol])
    book.save(OUT / "audit.xlsx")


def main() -> None:
    if paper_alive():
        report = {
            "verdict": "AUTO50_G1M_CAUSAL_BASELINE_NOT_READY_V1",
            "next": "RESUME_AFTER_PAPER_NORMAL_CLOSE",
            "reason": "PAPER_PROCESS_STILL_ALIVE",
            "PAPER_STOPPED": False,
            "RUNTIME_CHANGED": False,
            "submit_cancel_live": "0/0/0",
        }
        OUT.mkdir(parents=True, exist_ok=True)
        (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        (OUT / "report.md").write_text("VERDICT: AUTO50_G1M_CAUSAL_BASELINE_NOT_READY_V1\nNEXT: RESUME_AFTER_PAPER_NORMAL_CLOSE\n", encoding="utf-8")
        print("paper still alive", flush=True)
        return
    identity = capital_ok()
    active, rejected = load_master()
    alpha = {row["code"] for row in active if row["code"][3].isalpha()}
    meta = {row["code"]: row for row in active}
    print(f"universe {len(active)} alpha {len(alpha)} rejected {len(rejected)}", flush=True)
    panel, failed = download_panel([research_yahoo_ticker(row["code"]) for row in active], date(2026, 3, 1), date(2026, 10, 9))
    sessions = sorted(
        {
            row["source_bar_date"]
            for symbol in ("1301", "7203")
            for row in panel.get(symbol, [])
        }
    )
    saved = json.loads(V2_REPORT.read_text(encoding="utf-8"))
    unavailable = set(saved.get("target_close_unavailable_dates") or [])
    session_set = set(sessions)
    days = [day for day in selection_calendar(REPORTS) if day in session_set and day <= "20261007" and day not in unavailable]
    print(f"sessions {len(sessions)} days {len(days)} failed {len(failed)}", flush=True)
    price = audit_stored_prices(panel, sessions)
    conflicts = price.pop("conflicts")
    witness = price["witness_6379_20260717"]
    walked = run_days(days, panel, meta, sessions, conflicts, alpha)
    official = [row for row in walked["daily"] if row.get("official_g1m")]
    paired = [
        float(row["paired_g1m_minus_v2_range"])
        for row in walked["daily"]
        if row.get("paired_g1m_minus_v2_range") is not None and row.get("official_g1m")
    ]
    paired.sort()
    report = {
        "capital_match": identity["match"],
        "CAPITAL": {"BASE_CAPITAL_YEN": 500000, "MAX_LEVERAGE_MULTIPLIER": 2.0, "G": 1000000, "QTY": 100},
        "POLICY_INSTANCE_SHA": identity.get("POLICY_INSTANCE_SHA"),
        "MASTER_ALPHANUMERIC_N": len(alpha),
        "MASTER_ACTIVE_N": len(active),
        "parser_rejected": len(rejected),
        "FETCH_FAILED_N": len(failed),
        "fetch_failed_symbols": failed,
        "price_audit": price,
        "witness_pass": bool(witness.get("pass")),
        "target_days": len(days),
        "official_g1m_days": len(official),
        "lag_zero_days": sum(1 for row in walked["daily"] if row.get("lag_sessions") == 0),
        "future_used_total": sum(int(row.get("future_used") or 0) for row in walked["daily"]),
        "books": walked["books"],
        "paired_g1m_minus_v2_range_median": paired[len(paired) // 2] if paired else None,
        "old_vs_g1m_replacement_name_days": sum(int(row.get("old_vs_g1m_replacements") or 0) for row in walked["daily"]),
        "old_vs_raw_replacement_name_days": sum(int(row.get("old_vs_raw_replacements") or 0) for row in walked["daily"]),
        "g1m_alpha_selected_name_days": sum(len(row.get("g1m_alpha_selected") or []) for row in walked["daily"]),
        "raw_alpha_selected_name_days": sum(len(row.get("raw_alpha_selected") or []) for row in walked["daily"]),
        "used_as_independent_holdout": False,
        "economic_superiority_claimed": False,
        "v3_adopted": False,
        "score_formula_changed": False,
        "runtime_normalize_code_changed": False,
        "PAPER_V16": "UNCHANGED",
        "PAPER_STOPPED": False,
        "RUNTIME_CHANGED": False,
        "LIVE_ORDER_ENABLED": False,
        "submit_cancel_live": "0/0/0",
        "daily": walked["daily"],
        "focus": walked["focus"],
        "decision_rule": {
            "official_days_must_equal_target_days": True,
            "lag_sessions": 0,
            "future_used": 0,
            "g_violations": 0,
            "stored_price_explained_rate_min": 0.99,
            "g1m_range_median_must_be_at_least_v2": True,
            "active_rate_min": 0.95,
            "dead_rate_max": 0.05,
            "already_seen_dates_are_not_an_independent_holdout": True,
        },
    }
    verdict, nxt = judge(report)
    report["verdict"] = verdict
    report["next"] = nxt
    write_outputs(report)
    print(verdict, nxt, flush=True)


if __name__ == "__main__":
    main()
