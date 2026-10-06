"""Write report.json / report.md / audit.xlsx only. Do not rewrite V1 freeze artifacts."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.fixed_daytrade_universe_composition_audit_v1 import ANALYSIS_ID
from research.fixed_daytrade_universe_composition_audit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "SUMMARY",
    "TSE33_TOPIX17",
    "PROXIES",
    "FACTOR_COVERAGE",
    "MISSING_ROLES",
    "REDUNDANCY",
    "TOPIX_BETA",
    "V1_1_CANDIDATE",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    v11 = dict(report.get("v1_1_candidate") or {})
    return {
        "SUMMARY": _kv(dict(report.get("answers") or {})) + _kv(dict(report.get("decision") or {})),
        "TSE33_TOPIX17": list(report.get("mapping") or []),
        "PROXIES": list(report.get("proxies") or []),
        "FACTOR_COVERAGE": list(report.get("factor_coverage") or []),
        "MISSING_ROLES": list(report.get("missing_roles") or []),
        "REDUNDANCY": list(report.get("redundancy_clusters") or [])
        + [{"pair": True, **p} for p in list(report.get("redundancy_pairs_high") or [])],
        "TOPIX_BETA": list(dict(report.get("topix_beta") or {}).get("rows") or []) or _kv(dict(report.get("topix") or {})),
        "V1_1_CANDIDATE": list(v11.get("add") or []) or _kv({k: v for k, v in v11.items() if k != "proposed_symbols_if_accepted_without_drops"}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    cov = list(report.get("factor_coverage") or [])
    v11 = dict(report.get("v1_1_candidate") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        "",
        str(d.get("INTERPRETATION") or ""),
        "",
        f"V1 rewritten? **{a.get('v1_rewritten')}**. Universe SHA `{a.get('universe_sha')}`.",
        "",
        "## Factor side coverage",
        "",
    ]
    for r in cov:
        lines.append(
            f"- {r.get('factor')}: {r.get('status')} pos={r.get('pos_n')} {r.get('pos_symbols')} "
            f"neg={r.get('neg_n')} {r.get('neg_symbols')}"
        )
    lines += [
        "",
        "## Missing roles",
        "",
        f"Necessary absent: {a.get('necessary_missing_roles')}",
        "",
        "## V1.1 candidate (proposal only)",
        "",
        f"add_n={v11.get('add_n')} symbols={a.get('v1_1_adds')}",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
    assert (OUT / "report.json").is_file()
    freeze = OUT.parent / "fixed_daytrade_universe_v1" / "universe_manifest.json"
    assert freeze.is_file()
