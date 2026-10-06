"""Pin parent discovery identity. No V5 / Full Strategy freeze."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from research.discovery_search_space_reassessment_v1 import (
    REQUIRED_CANDIDATE_MECHANISM_N,
    REQUIRED_COVERAGE_QUALIFIED_N,
    REQUIRED_PARENT_VERDICT,
    REQUIRED_PASS_D1_D11_N,
    REQUIRED_RAW_PREDICATE_N,
)
from research.discovery_search_space_reassessment_v1.isolation import RESEARCH_ROOT

SOURCE_FILES = (
    "__init__.py",
    "spec.py",
    "isolation.py",
    "engine.py",
    "controls.py",
    "labels.py",
    "analyze.py",
    "publish.py",
    "__main__.py",
)

PARENT_REPORT = RESEARCH_ROOT / "profitable_move_mechanism_discovery_v1" / "report.json"
PARENT_AUDIT = RESEARCH_ROOT / "profitable_move_mechanism_discovery_v1" / "audit.xlsx"


def dumps_sha256(obj: Any) -> str:
    body = json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(body.encode("utf-8")).hexdigest()


def source_sha256() -> str:
    root = Path(__file__).resolve().parent
    h = hashlib.sha256()
    for name in SOURCE_FILES:
        h.update(name.encode("utf-8"))
        h.update((root / name).read_bytes())
    return h.hexdigest()


def load_parent_report() -> dict[str, Any]:
    if not PARENT_REPORT.is_file():
        return {}
    return json.loads(PARENT_REPORT.read_text(encoding="utf-8"))


def coverage_qualified_h5_means() -> list[float]:
    if not PARENT_AUDIT.is_file():
        return []
    from openpyxl import load_workbook

    wb = load_workbook(PARENT_AUDIT, read_only=True, data_only=True)
    rows = list(wb["markout_h5"].iter_rows(values_only=True))
    hdr = list(rows[0])
    out: list[float] = []
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        v = d.get("mean_markout_yen100")
        n = d.get("EVENT_N")
        try:
            nn = int(n) if n is not None else 0
            fv = float(v) if v is not None else float("nan")
        except (TypeError, ValueError):
            continue
        if nn >= 40 and fv == fv:
            out.append(fv)
    return out


def pin_parent() -> dict[str, Any]:
    prev = load_parent_report()
    d = dict(prev.get("decision") or {})
    verdict = str(d.get("VERDICT") or prev.get("VERDICT") or "")
    raw_n = int(prev.get("RAW_PREDICATE_N") or -1)
    cand_n = int(prev.get("CANDIDATE_MECHANISM_N") or -1)
    cov_n = int(prev.get("COVERAGE_QUALIFIED_N") or -1)
    pass_n = int(d.get("PASS_D1_D11_N") if d.get("PASS_D1_D11_N") is not None else -1)
    feat_raw = (prev.get("data") or {}).get("FEATURE_LOOKAHEAD_N")
    feat = int(feat_raw) if feat_raw is not None else -1
    frozen = d.get("FULL_STRATEGY_FROZEN")
    means = coverage_qualified_h5_means()
    all_neg = bool(means) and all(float(v) < 0.0 for v in means)
    ok = (
        verdict == REQUIRED_PARENT_VERDICT
        and raw_n == int(REQUIRED_RAW_PREDICATE_N)
        and cand_n == int(REQUIRED_CANDIDATE_MECHANISM_N)
        and cov_n == int(REQUIRED_COVERAGE_QUALIFIED_N)
        and pass_n == int(REQUIRED_PASS_D1_D11_N)
        and feat == 0
        and frozen is False
        and all_neg
        and len(means) == int(REQUIRED_COVERAGE_QUALIFIED_N)
    )
    return {
        "ok": bool(ok),
        "REQUIRED_PARENT_VERDICT": REQUIRED_PARENT_VERDICT,
        "observed_verdict": verdict,
        "RAW_PREDICATE_N": raw_n,
        "CANDIDATE_MECHANISM_N": cand_n,
        "COVERAGE_QUALIFIED_N": cov_n,
        "PASS_D1_D11_N": pass_n,
        "FEATURE_LOOKAHEAD_N": feat,
        "FULL_STRATEGY_FROZEN": frozen,
        "COVERAGE_QUALIFIED_H5_MEAN_N": int(len(means)),
        "ALL_COVERAGE_QUALIFIED_H5_ABS_MEAN_NEGATIVE": all_neg,
        "H5_ABS_MEAN_MAX": (max(means) if means else None),
        "H5_ABS_MEAN_MIN": (min(means) if means else None),
    }
