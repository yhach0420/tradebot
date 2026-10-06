"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.discovery_search_space_reassessment_v1 import ANALYSIS_ID
from research.discovery_search_space_reassessment_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent_pin",
    "p1",
    "p2",
    "negative_control",
    "sign_compatibility",
    "strategy_path",
    "execution_cost",
    "hard_label",
    "p1_trades",
    "p2_trades",
    "csb_trades",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _trade_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in rows:
        out.append(
            {
                "date": r.get("date"),
                "symbol": r.get("symbol"),
                "signal_t0": r.get("signal_t0"),
                "entry_fill_t": r.get("entry_fill_t"),
                "entry_fill_price": r.get("entry_fill_price"),
                "exit_fill_t": r.get("exit_fill_t"),
                "actual_pnl_yen100": r.get("actual_pnl_yen100"),
                "hold_sec": r.get("hold_sec"),
                "X1_FILL_COMPATIBLE": r.get("X1_FILL_COMPATIBLE"),
                "exec_h3": r.get("exec_yen100_h3"),
                "exec_h5": r.get("exec_yen100_h5"),
                "exec_h10": r.get("exec_yen100_h10"),
                "mid_h3": r.get("mid_yen100_h3"),
                "mid_h5": r.get("mid_yen100_h5"),
                "mid_h10": r.get("mid_yen100_h10"),
                "excess_h5": r.get("excess_yen100_h5"),
                "cross_tax_h5": r.get("cross_tax_yen100_h5"),
            }
        )
    return out or [{"empty": True}]


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT') or a.get('32_VERDICT'))}",
        "",
        f"FIXED_H5_HARD_GATE_VALID: {_fmt(d.get('FIXED_H5_HARD_GATE_VALID') or a.get('19_FIXED_H5_HARD_GATE_VALID'))}",
        "",
        f"SECONDARY_CAUSE: {_fmt(d.get('SECONDARY_CAUSE') or a.get('33_SECONDARY_CAUSE'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT') or a.get('34_NEXT'))}",
        "",
        "NEW_STRATEGY_CREATED: false",
        "",
        "## Answers",
        "",
    ]
    for k, v in a.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    dumped = json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n"
    (OUT / "report.json").write_text(dumped, encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_sheets(report: dict[str, Any], pack: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    p1 = dict(pack.get("P1") or {})
    p2 = dict(pack.get("P2") or {})
    neg = dict(pack.get("NEGATIVE") or {})
    labeled = dict(pack.get("labeled") or {})
    return {
        "answers": kv_rows(report.get("answers")),
        "objective_alignment": kv_rows(pack.get("objective_alignment")),
        "parent_pin": kv_rows(pack.get("pin")),
        "p1": kv_rows(p1),
        "p2": kv_rows(p2),
        "negative_control": kv_rows(neg),
        "sign_compatibility": [
            {"control": "P1", **dict(p1.get("signs") or {})},
            {"control": "P2", **dict(p2.get("signs") or {})},
            {"control": "CSB_V3", **dict(neg.get("signs") or {})},
        ],
        "strategy_path": [
            {"control": "P1", **dict(p1.get("path") or {})},
            {"control": "P2", **dict(p2.get("path") or {})},
            {"control": "CSB_V3", **dict(neg.get("path") or {})},
        ],
        "execution_cost": [
            {
                "control": "P1",
                **{f"exec_{k}": v for k, v in dict(p1.get("executable") or {}).items()},
                **{f"mid_{k}": v for k, v in dict(p1.get("mid") or {}).items()},
                **{f"tax_{k}": v for k, v in dict(p1.get("crossing_tax") or {}).items()},
            },
            {
                "control": "P2",
                **{f"exec_{k}": v for k, v in dict(p2.get("executable") or {}).items()},
                **{f"mid_{k}": v for k, v in dict(p2.get("mid") or {}).items()},
                **{f"tax_{k}": v for k, v in dict(p2.get("crossing_tax") or {}).items()},
            },
        ],
        "hard_label": kv_rows(pack.get("decision")),
        "p1_trades": _trade_rows(list(labeled.get("P1") or [])),
        "p2_trades": _trade_rows(list(labeled.get("P2") or [])),
        "csb_trades": _trade_rows(list(labeled.get("NEGATIVE") or [])),
        "decision": kv_rows(pack.get("decision")),
        "safety": kv_rows(report.get("safety")),
    }
