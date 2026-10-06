"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.entry_objective_redesign_c3 import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "entry_objective_redesign_c3"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Eligibility",
    "Current_Baseline",
    "Outer_Folds",
    "Spec_Selection",
    "TopK_OOF",
    "Daily_Paired",
    "Robustness",
    "A0_A2_Parity",
    "Exact_Portfolio",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                obj = float(obj)
            elif isinstance(obj, np.integer):
                return int(obj)
    except Exception:
        pass
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        elif isinstance(obj, tuple):
            flat[prefix] = list(obj)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# C3 ENTRY OBJECTIVE REDESIGN",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE RESEARCH ONLY. TARGET V4 M4_PERSISTENT frozen. CURRENT IRREGULAR CLOCK frozen.",
            "Ranking population = is_executable_continuous_board(t0). Primary metric = mean daily Top3 uplift.",
            "Ridge only. 27 precommitted specs. Nested leave-one-day-out. No RF. No new features.",
            "C2_OOF_RANKING_EDGE_NOT_STABLE maintained. B_NO_ROBUST_IMPROVEMENT maintained.",
            "REENTRY_ORDINAL_IS_CLOCK_PROXY maintained. B / UNIFORM10 are not C3 gates.",
            "C14 / Runtime / CLOCK / EXIT / CAP / re-entry / Fill / freshness 5s unchanged.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "TRUE_OOS=false. NEW_FORWARD_N=0. HISTORICAL DEVELOPMENT ONLY.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A0_PARITY: {g.get('A0_PARITY')}",
            f"A2_PARITY: {g.get('A2_PARITY')}",
            f"EXECUTABLE_T0_ROWS: {g.get('EXECUTABLE_T0_ROWS')}",
            f"CURRENT_TOP1_UPLIFT: {g.get('CURRENT_TOP1_UPLIFT')}",
            f"CURRENT_TOP3_UPLIFT: {g.get('CURRENT_TOP3_UPLIFT')}",
            f"CURRENT_TOP5_UPLIFT: {g.get('CURRENT_TOP5_UPLIFT')}",
            f"C3_TOP1_UPLIFT: {g.get('C3_TOP1_UPLIFT')}",
            f"C3_TOP3_UPLIFT: {g.get('C3_TOP3_UPLIFT')}",
            f"C3_TOP5_UPLIFT: {g.get('C3_TOP5_UPLIFT')}",
            f"CURRENT_MEAN_DAILY_SPEARMAN: {g.get('CURRENT_MEAN_DAILY_SPEARMAN')}",
            f"C3_MEAN_DAILY_SPEARMAN: {g.get('C3_MEAN_DAILY_SPEARMAN')}",
            f"TOP3_DELTA_MEAN: {g.get('TOP3_DELTA_MEAN')}",
            f"TOP3_DELTA_MEDIAN: {g.get('TOP3_DELTA_MEDIAN')}",
            f"TOP3_DELTA_POSITIVE_DAYS: {g.get('TOP3_DELTA_POSITIVE_DAYS')}",
            f"TOP3_DELTA_NEGATIVE_DAYS: {g.get('TOP3_DELTA_NEGATIVE_DAYS')}",
            f"TOP3_DELTA_EX_BEST_DAY: {g.get('TOP3_DELTA_EX_BEST_DAY')}",
            f"TOP3_DELTA_EX_TOP3_DAYS: {g.get('TOP3_DELTA_EX_TOP3_DAYS')}",
            f"MOST_COMMON_SPEC: {g.get('MOST_COMMON_SPEC')}",
            f"MOST_COMMON_SPEC_SHARE: {g.get('MOST_COMMON_SPEC_SHARE')}",
            f"OOF_RANKING_GATE_PASS: {g.get('OOF_RANKING_GATE_PASS')}",
            f"FINAL_SPEC: {g.get('FINAL_SPEC')}",
            f"C3_EXACT_RAN: {g.get('C3_EXACT_RAN')}",
            f"A0: {g.get('A0')}",
            f"A2: {g.get('A2')}",
            f"C3: {g.get('C3')}",
            f"C3_VS_A2: {g.get('C3_VS_A2')}",
            f"C3_VS_A0: {g.get('C3_VS_A0')}",
            f"TRUE_OOS: {g.get('TRUE_OOS')}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## STOP",
            "",
            "Audit complete. Runtime not changed. C3 not implemented live. C14 unchanged.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
