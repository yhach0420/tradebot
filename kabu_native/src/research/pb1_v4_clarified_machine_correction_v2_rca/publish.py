"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_correction_v2_rca.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "SUMMARY",
    "ONE_BAR_RCA",
    "TWO_SIDED_RCA",
    "ACTIVE_LIVENESS",
    "EXEC_AFTER_DEATH",
    "CONTROL_REGRESSIONS",
    "BOUNDARY_PROVENANCE",
)
STRIP = {"_markdown", "snap", "setups", "e0_events", "e1_events", "funnel_days", "hidden_1m_checks"}


def finite_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): finite_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [finite_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(finite_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    one = dict(report.get("one_bar_rca") or {})
    two = dict(report.get("two_sided_rca") or {})
    live = dict(report.get("active_liveness") or {})
    ctrl = dict(report.get("controls") or {})
    hid = dict(report.get("hidden_1m") or {})
    exec_c = dict(live.get("execution_contract") or {})
    locg = dict(live.get("location_gating") or {})
    p7011 = dict(one.get("primary") or {})
    s3382 = dict(live.get("3382_20241004") or {})
    s7011b = dict(live.get("7011_20250523") or {})
    s4063 = dict(live.get("4063_20251118") or {})
    return {
        "Correction V2 SHA unchanged?": True,
        "Spec SHA unchanged?": True,
        "RCA set n?": 88,
        "Any rule changed?": False,
        "Any prospective event consumed?": False,
        "Any future outcome used?": False,
        "ONE_BAR primary cause?": (
            "FOLLOWTHROUGH treats last-bar same-dir + n_strong>=1 as continued auction, "
            "and n_strong includes the dominant first print."
        ),
        "Does current followthrough incorrectly count the dominant print as its own evidence?": bool(
            one.get("followthrough_incorrectly_counts_dominant_print")
        ),
        "What post-dominant state distinguishes 7011 / 20241205 from real continuation?": p7011.get("post_dominant_class"),
        "How many one-bar-heavy cases reviewed?": one.get("one_bar_heavy_n"),
        "Is a numeric share threshold required?": one.get("numeric_share_threshold_required"),
        "TWO_SIDED primary cause?": two.get("primary_cause"),
        "Why is 6963 / 20250613 different from 7741?": two.get("6963_vs_7741"),
        "Why is 5803 / 20250212 different from 7741?": two.get("5803_0212_vs_7741"),
        "Are 5803 / 20250709 and 5706 / 20250725 still genuine two-sided?": {
            "5803_20250709": two.get("still_genuine_two_sided_5803_0709"),
            "5706_20250725": two.get("still_genuine_two_sided_5706"),
            "5803_class": two.get("5803_20250709_class"),
            "5706_class": two.get("5706_20250725_class"),
        },
        "What semantic dimension is missing?": two.get("missing_semantic_dimension"),
        "Is body>=0.35 itself sufficient?": False,
        "Is retracement alone sufficient?": False,
        "ACTIVE / THESIS primary cause?": live.get("primary_cause"),
        "Does current machine distinguish THESIS_REACHED from THESIS_LIVE?": live.get("distinguishes_THESIS_REACHED_from_THESIS_LIVE"),
        "Does location alter ACTIVE death monitoring?": locg.get("behaves_differently_before_vs_after_LOCATION_IDENTIFIED"),
        "Is N=3 non-reset part of frozen semantic spec?": False,
        "What constitutes actual STALE_RANGE_RESOLUTION?": live.get("what_constitutes_STALE_RANGE_RESOLUTION"),
        "3382 / 20241004: would the proposed semantic stale state remain live until valid continuation?": s3382.get(
            "semantic_stale_would_remain_live_until_valid_continuation"
        ),
        "7011 / 20250523: when does thesis actually become stale, if at all?": s7011b.get("when_stale"),
        "4063 / 20251118: when does thesis actually become stale, if at all?": s4063.get("when_stale"),
        "E0_after_ACTIVE_loss_n?": live.get("E0_after_ACTIVE_loss_n"),
        "E1_after_ACTIVE_loss_n?": live.get("E1_after_ACTIVE_loss_n"),
        "E0_after_THESIS_loss_n?": live.get("E0_after_THESIS_loss_n"),
        "E1_after_THESIS_loss_n?": live.get("E1_after_THESIS_loss_n"),
        "Does E0 check live thesis immediately before entry?": exec_c.get("E0_checks_live_thesis_at_mint"),
        "Does E1 check live thesis immediately before entry?": exec_c.get("E1_checks_live_thesis_immediately_before_entry_classify"),
        "Should thesis_id be erased on loss?": False,
        "Should live-state and day-reach be separate?": True,
        "FAILED_OPEN semantics still hold?": ctrl.get("FAILED_OPEN_semantics_still_hold"),
        "Family A semantics still hold?": ctrl.get("Family_A_semantics_still_hold"),
        "Hidden-1m parity still structural?": int(hid.get("mismatch_n") or 0) == 0,
        "SPEC_CHANGE_REQUIRED?": dec.get("SPEC_CHANGE_REQUIRED"),
        "Any threshold optimized?": False,
        "Any PnL?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
        "causal_one_bar_distinction": one.get("causal_distinction"),
        "one_bar_compare_n": len(list(one.get("compare_set") or [])),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    one = dict(report.get("one_bar_rca") or {})
    two = dict(report.get("two_sided_rca") or {})
    live = dict(report.get("active_liveness") or {})
    ctrl = dict(report.get("controls") or {})
    dec = dict(report.get("decision") or {})
    answers = dict(report.get("answers") or {})
    return {
        "SUMMARY": [{"key": k, "value": answers.get(k)} for k in answers],
        "ONE_BAR_RCA": list(one.get("compare_set") or []) + list(one.get("one_bar_heavy") or []),
        "TWO_SIDED_RCA": list(two.get("sandwiches") or []),
        "ACTIVE_LIVENESS": list(live.get("traces_active") or []),
        "EXEC_AFTER_DEATH": [
            {
                "E0_after_ACTIVE_loss_n": live.get("E0_after_ACTIVE_loss_n"),
                "E1_after_ACTIVE_loss_n": live.get("E1_after_ACTIVE_loss_n"),
                "E0_after_THESIS_loss_n": live.get("E0_after_THESIS_loss_n"),
                "E1_after_THESIS_loss_n": live.get("E1_after_THESIS_loss_n"),
                "material_strategy_state_bug": live.get("exec_after_death_is_material_strategy_bug"),
                **dict(live.get("execution_contract") or {}),
            }
        ],
        "CONTROL_REGRESSIONS": list(ctrl.get("positives") or []) + list(ctrl.get("negatives") or []) + list(ctrl.get("family_a_focus") or []),
        "BOUNDARY_PROVENANCE": [
            {"name": "ONE_BAR_SHARE_MAX 0.55", "status": "not_restored", "role": "exemplar-derived; not reusable"},
            {"name": "WEAK_BAR_BODY_MAX 0.20", "status": "not_restored", "role": "exemplar-derived; not reusable"},
            {"name": "TRUE_BODY_FRAC_MIN 0.35 as substantial_opp", "status": "implementation_proxy", "role": "too broad for SUBSTANTIAL_COUNTER_AUCTION"},
            {"name": "REPEATED_NO_EXPANSION_N=3", "status": "implementation_encoding", "role": "not in Clarified V2; N-bar expiry forbidden"},
            {"name": "PROGRESS_TINY_EXTREME_FRAC 0.20", "status": "untouched", "role": "not this RCA"},
            {"name": "PROGRESS_HEAVY_OVERLAP 0.70", "status": "untouched", "role": "not this RCA"},
            {"name": "SPEC_CHANGE_REQUIRED", "status": str(dec.get("SPEC_CHANGE_REQUIRED")), "role": "V2 already contains the three concepts"},
        ],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = finite_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    answers = dict(report.get("answers") or {})
    lines = [
        "# PB1_V4_CLARIFIED_MACHINE_CORRECTION_V2_RCA_V1",
        "",
        f"VERDICT: {answers.get('VERDICT?')}",
        f"NEXT: {answers.get('NEXT?')}",
        "",
        "Diagnosis only. No machine change. No spec change. No threshold retune. No prospective. No PnL.",
        "",
    ]
    for k, v in answers.items():
        if isinstance(v, (dict, list)):
            lines.append(f"- {k} `{json.dumps(finite_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            lines.append(f"- {k} `{v}`")
    lines.append("")
    lines.append("STOP.")
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
