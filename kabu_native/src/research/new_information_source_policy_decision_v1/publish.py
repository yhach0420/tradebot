"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.new_information_source_policy_decision_v1 import ANALYSIS_ID
from research.new_information_source_policy_decision_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "candidate_sources",
    "official_documentation",
    "source_semantics",
    "historical_live_equivalence",
    "timestamp_audit",
    "material_gain",
    "state_potential",
    "policy_gates",
    "selection",
    "restart_gate",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "parent": _kv(dict(report.get("parent") or {})),
        "candidate_sources": list(report.get("candidate_sources") or []),
        "official_documentation": list(report.get("official_documentation") or []),
        "source_semantics": list(report.get("source_semantics") or []),
        "historical_live_equivalence": list(report.get("historical_live_equivalence") or []),
        "timestamp_audit": list(report.get("timestamp_audit") or []),
        "material_gain": list(report.get("material_gain") or []),
        "state_potential": list(report.get("state_potential") or []),
        "policy_gates": list(report.get("policy_gates") or []),
        "selection": _kv(dict(report.get("selection") or {})),
        "restart_gate": _kv(dict(report.get("restart_gate") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"STRATEGY_RESEARCH_STATUS: {d.get('STRATEGY_RESEARCH_STATUS')}",
            f"RESTART_ALLOWED: {d.get('RESTART_ALLOWED')}",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- selected: {a.get('36_selected_source_family_ID')}",
            f"- policy SHA: `{a.get('41_EXTERNAL_INFORMATION_SOURCE_POLICY_SHA256')}`",
            f"- policy-eligible unique feeds: {a.get('35_policy_eligible_IDs')}",
            f"- S1/S2 consolidated: {d.get('S1_S2_CONSOLIDATED')}",
            f"- parent pinned: {a.get('1_parent_verdict_pinned')}",
            "",
            "No market data acquired. No subscription. No strategy. Do not consume 20260907+.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
