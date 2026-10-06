"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.passive_wait5_precommit import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "passive_wait5_precommit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Days",
    "Sessions",
    "Clocks",
    "Delay",
    "Quality",
    "Gates",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# PASSIVE WAIT5 PRECOMMIT DEVELOPMENT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CONTROL W1=1s. CANDIDATE W5=5s. W2/W10 reference only.",
            "Shortest evaluated WAIT that clears ANY>=0.50 and MULTI>=0.50.",
            "Corrected Passive Fill frozen. t0 limit frozen. No repricing.",
            "Runtime WAIT_SEC remains 1.0. W5 not adopted.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"WAIT_POLICY_CANDIDATE: {g.get('WAIT_POLICY_CANDIDATE')}",
            "",
            f"W5_ANY_RATE: {_fmt(g.get('W5_ANY_RATE'))}",
            f"W5_MULTI_RATE: {_fmt(g.get('W5_MULTI_RATE'))}",
            "",
            f"AM_W1_ANY: {_fmt(g.get('AM_W1_ANY'))}",
            f"AM_W5_ANY: {_fmt(g.get('AM_W5_ANY'))}",
            "",
            f"PM_W1_ANY: {_fmt(g.get('PM_W1_ANY'))}",
            f"PM_W5_ANY: {_fmt(g.get('PM_W5_ANY'))}",
            "",
            f"AM_W1_MULTI: {_fmt(g.get('AM_W1_MULTI'))}",
            f"AM_W5_MULTI: {_fmt(g.get('AM_W5_MULTI'))}",
            "",
            f"PM_W1_MULTI: {_fmt(g.get('PM_W1_MULTI'))}",
            f"PM_W5_MULTI: {_fmt(g.get('PM_W5_MULTI'))}",
            "",
            f"CLOCKS_ANY_IMPROVED_N: {g.get('CLOCKS_ANY_IMPROVED_N')}",
            f"CLOCKS_MULTI_IMPROVED_N: {g.get('CLOCKS_MULTI_IMPROVED_N')}",
            "",
            f"W5_NEW_FILL_N: {g.get('W5_NEW_FILL_N')}",
            f"W5_NEW_POSTFILL_MFE: {_fmt(g.get('W5_NEW_POSTFILL_MFE'))}",
            f"W5_NEW_POSTFILL_DOWNSIDE: {_fmt(g.get('W5_NEW_POSTFILL_DOWNSIDE'))}",
            "",
            f"W5_JOINT_SURVIVAL: {_fmt(g.get('W5_JOINT_SURVIVAL'))}",
            "",
            f"W5_PRECOMMIT_PASS: {_tf(g.get('W5_PRECOMMIT_PASS'))}",
            "",
            f"PRIMARY_FINDING: {g.get('PRIMARY_FINDING')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "Runtime WAIT_SEC remains 1.0. W5 not adopted.",
            "No Exact. No ENTRY retrain. No price-policy test.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
