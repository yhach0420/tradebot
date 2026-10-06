"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_fixed_spec_oof import ANALYSIS_ID
from research.am_entry_profit_improvement.publish import json_sanitize, kv_rows

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_entry_fixed_spec_oof"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "summary",
    "fixed_spec_oof",
    "daily_pnl_by_spec",
    "architecture_summary",
    "nested_failure",
    "tail_loss",
    "utility_scale",
    "fill_support",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# AM ENTRY FIXED-SPEC OOF MATRIX AND PROFIT FAILURE DECOMPOSITION V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. 27 frozen specs as independent policies. No inner selection.",
            "Each spec: 18 outer LODO days, stitch, same W5 / EXIT / occupancy / economic gate as V2.",
            "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED FINAL OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"SPEC_N: {_fmt(g.get('SPEC_N'))}",
            "",
            f"FIXED_SPEC_OOF_COMPLETE_N: {_fmt(g.get('FIXED_SPEC_OOF_COMPLETE_N'))}",
            "",
            f"PASS_SPEC_N: {_fmt(g.get('PASS_SPEC_N'))}",
            "",
            f"BEST_PASS_SPEC: {_fmt(g.get('BEST_PASS_SPEC'))}",
            "",
            f"BEST_PASS_NET_PNL: {_fmt(g.get('BEST_PASS_NET_PNL'))}",
            "",
            f"BEST_PASS_PF: {_fmt(g.get('BEST_PASS_PF'))}",
            "",
            f"BEST_PASS_MAX_DD: {_fmt(g.get('BEST_PASS_MAX_DD'))}",
            "",
            f"BEST_PASS_PAIRED_MEDIAN: {_fmt(g.get('BEST_PASS_PAIRED_MEDIAN'))}",
            "",
            f"BEST_PASS_EX_TOP3: {_fmt(g.get('BEST_PASS_EX_TOP3'))}",
            "",
            f"INNER_OUTER_SPEARMAN: {_fmt(g.get('INNER_OUTER_SPEARMAN'))}",
            "",
            f"INNER_OUTER_PEARSON: {_fmt(g.get('INNER_OUTER_PEARSON'))}",
            "",
            f"SELECTED_OUTER_PNL_RF: {_fmt(g.get('SELECTED_OUTER_PNL_RF'))}",
            "",
            f"SELECTED_OUTER_PNL_RIDGE: {_fmt(g.get('SELECTED_OUTER_PNL_RIDGE'))}",
            "",
            f"SELECTED_OUTER_PNL_PAIRWISE: {_fmt(g.get('SELECTED_OUTER_PNL_PAIRWISE'))}",
            "",
            f"TOP2_LOSS_SHARE: {_fmt(g.get('TOP2_LOSS_SHARE'))}",
            "",
            f"WORST_SYMBOL: {_fmt(g.get('WORST_SYMBOL'))}",
            "",
            f"WORST_SYMBOL_PNL: {_fmt(g.get('WORST_SYMBOL_PNL'))}",
            "",
            f"SPEARMAN_ABS_UTILITY_VS_FILL_PRICE: {_fmt(g.get('SPEARMAN_ABS_UTILITY_VS_FILL_PRICE'))}",
            "",
            f"PRIMARY_FAILURE_MECHANISM: {_fmt(g.get('PRIMARY_FAILURE_MECHANISM'))}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            "",
            f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            f"NEXT: {g.get('NEXT')}",
            "",
            "## STOP",
            "",
            "Fixed-spec 27×18 OOF complete. No Runtime change. No W5 adoption.",
            "No new feature. No new target. No post-hoc spec addition. submit/cancel/live=0/0/0.",
            "",
        ]
    )
