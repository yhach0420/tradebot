"""Nested-selector, tail-loss, utility-scale, and fill-support diagnostics. No policy change."""
from __future__ import annotations

from collections import defaultdict
from typing import Any

import numpy as np
from openpyxl import load_workbook

from research.am_entry_fixed_spec_oof import V2_ANALYSIS_ID
from research.am_entry_profit_improvement import ARCH_PAIR, ARCH_RF, ARCH_RIDGE, UTILITY_KEY
from research.canonical_entry_performance_rebase.analyze import _f
from research.entry_objective_redesign_c3.oof import spearman
from research.entry_rank_shape_audit.oof import pearson


def load_v2_audit(path: Any) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[dict[str, Any]]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            out[name] = []
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        recs = []
        for raw in rows[1:]:
            recs.append({header[i]: raw[i] if i < len(raw) else None for i in range(len(header))})
        out[name] = recs
    wb.close()
    return out


def nested_inner_outer(outer_folds: list[dict[str, Any]]) -> dict[str, Any]:
    xs: list[float] = []
    ys: list[float] = []
    rows = []
    by_arch_pnl: dict[str, float] = {ARCH_RF: 0.0, ARCH_RIDGE: 0.0, ARCH_PAIR: 0.0}
    by_arch_n: dict[str, int] = {ARCH_RF: 0, ARCH_RIDGE: 0, ARCH_PAIR: 0}
    by_arch_trades: dict[str, int] = {ARCH_RF: 0, ARCH_RIDGE: 0, ARCH_PAIR: 0}
    for r in outer_folds:
        inner = _f(r.get("mean_inner_pnl"))
        outer = _f(r.get("outer_pnl"))
        arch = str(r.get("architecture_id") or "")
        trades = int(r.get("outer_trades") or 0)
        if inner is not None and outer is not None:
            xs.append(float(inner))
            ys.append(float(outer))
        if arch in by_arch_pnl and outer is not None:
            by_arch_pnl[arch] += float(outer)
            by_arch_n[arch] += 1
            by_arch_trades[arch] += trades
        rows.append(
            {
                "outer_day": r.get("outer_day"),
                "spec_id": r.get("spec_id"),
                "architecture_id": arch,
                "representation_id": r.get("representation_id"),
                "mean_inner_pnl": inner,
                "outer_pnl": outer,
                "outer_trades": trades,
            }
        )
    return {
        "INNER_OUTER_SPEARMAN": spearman(xs, ys) if len(xs) >= 10 else None,
        "INNER_OUTER_PEARSON": pearson(xs, ys),
        "n_folds": len(xs),
        "SELECTED_OUTER_PNL_RF": by_arch_pnl[ARCH_RF],
        "SELECTED_OUTER_PNL_RIDGE": by_arch_pnl[ARCH_RIDGE],
        "SELECTED_OUTER_PNL_PAIRWISE": by_arch_pnl[ARCH_PAIR],
        "SELECTED_OUTER_TRADE_N_RF": by_arch_trades[ARCH_RF],
        "SELECTED_OUTER_TRADE_N_RIDGE": by_arch_trades[ARCH_RIDGE],
        "SELECTED_OUTER_TRADE_N_PAIRWISE": by_arch_trades[ARCH_PAIR],
        "SELECTED_FOLD_N_RF": by_arch_n[ARCH_RF],
        "SELECTED_FOLD_N_RIDGE": by_arch_n[ARCH_RIDGE],
        "SELECTED_FOLD_N_PAIRWISE": by_arch_n[ARCH_PAIR],
        "folds": rows,
        "V2_ANALYSIS_ID": V2_ANALYSIS_ID,
    }


def tail_loss(trades: list[dict[str, Any]]) -> dict[str, Any]:
    losses = []
    for t in trades:
        pnl = float(_f(t.get("pnl_yen_100")) or 0.0)
        if pnl < -1e-12:
            losses.append(dict(t, pnl_yen_100=pnl))
    losses.sort(key=lambda r: float(r.get("pnl_yen_100") or 0.0))
    gross_loss = abs(sum(float(r.get("pnl_yen_100") or 0.0) for r in losses))
    net = sum(float(_f(t.get("pnl_yen_100")) or 0.0) for t in trades)

    def _share(k: int) -> float | None:
        if gross_loss <= 1e-12:
            return None
        top = losses[:k]
        combined = abs(sum(float(r.get("pnl_yen_100") or 0.0) for r in top))
        return float(combined / gross_loss)

    by_sym: dict[str, float] = defaultdict(float)
    by_sym_n: dict[str, int] = defaultdict(int)
    for t in trades:
        sym = str(t.get("symbol") or "")
        by_sym[sym] += float(_f(t.get("pnl_yen_100")) or 0.0)
        by_sym_n[sym] += 1
    sym_rows = [
        {"symbol": s, "pnl_yen_100": by_sym[s], "trade_n": by_sym_n[s]}
        for s in sorted(by_sym, key=lambda x: (by_sym[x], x))
    ]
    worst = sym_rows[0] if sym_rows else {"symbol": None, "pnl_yen_100": None, "trade_n": 0}
    conc = None
    if gross_loss > 1e-12 and worst.get("pnl_yen_100") is not None:
        worst_pnl = float(worst["pnl_yen_100"])
        if worst_pnl < 0:
            conc = abs(worst_pnl) / gross_loss
    return {
        "TRADE_N": len(trades),
        "NET_PNL": net,
        "GROSS_LOSS": gross_loss,
        "LOSS_N": len(losses),
        "TOP1_LOSS_SHARE": _share(1),
        "TOP2_LOSS_SHARE": _share(2),
        "TOP3_LOSS_SHARE": _share(3),
        "TOP5_LOSS_SHARE": _share(5),
        "WORST_SYMBOL": worst.get("symbol"),
        "WORST_SYMBOL_PNL": worst.get("pnl_yen_100"),
        "WORST_SYMBOL_TRADE_N": worst.get("trade_n"),
        "WORST_SYMBOL_CONCENTRATION": conc,
        "worst_trades": [
            {
                "rank": i + 1,
                "date": t.get("date"),
                "symbol": t.get("symbol"),
                "anchor": t.get("anchor"),
                "pnl_yen_100": t.get("pnl_yen_100"),
                "fill_price": t.get("fill_price"),
            }
            for i, t in enumerate(losses[:5])
        ],
        "symbol_pnl": sym_rows,
    }


def _terciles(prices: np.ndarray, values: np.ndarray) -> list[dict[str, Any]]:
    if prices.size == 0:
        return []
    qs = np.quantile(prices, [1.0 / 3.0, 2.0 / 3.0])
    out = []
    masks = [
        ("T1_LOW", prices <= qs[0]),
        ("T2_MID", (prices > qs[0]) & (prices <= qs[1])),
        ("T3_HIGH", prices > qs[1]),
    ]
    for name, mask in masks:
        vv = values[mask]
        out.append(
            {
                "tercile": name,
                "n": int(mask.sum()),
                "price_lo": float(prices[mask].min()) if int(mask.sum()) else None,
                "price_hi": float(prices[mask].max()) if int(mask.sum()) else None,
                "UTILITY_ABS_MEDIAN": float(np.median(vv)) if vv.size else None,
                "UTILITY_ABS_P90": float(np.percentile(vv, 90)) if vv.size else None,
            }
        )
    return out


def utility_scale(rows: list[dict[str, Any]]) -> dict[str, Any]:
    fills = []
    for r in rows:
        if int(r.get("Y_FILL5") or 0) != 1:
            continue
        u = _f(r.get(UTILITY_KEY))
        px = _f(r.get("fill_price"))
        if u is None or px is None:
            continue
        fills.append(
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "fill_price": float(px),
                "utility": float(u),
                "abs_utility": abs(float(u)),
                "pnl_yen_100": float(u),
            }
        )
    if not fills:
        return {
            "SPEARMAN_ABS_UTILITY_VS_FILL_PRICE": None,
            "FILL_N": 0,
            "price_terciles": [],
            "loss_terciles": [],
        }
    prices = np.asarray([f["fill_price"] for f in fills], dtype=float)
    abs_u = np.asarray([f["abs_utility"] for f in fills], dtype=float)
    sp = spearman(prices.tolist(), abs_u.tolist())
    losses = [f for f in fills if float(f["utility"]) < -1e-12]
    loss_sp = None
    loss_terciles: list[dict[str, Any]] = []
    if losses:
        lp = np.asarray([f["fill_price"] for f in losses], dtype=float)
        lm = np.asarray([abs(float(f["utility"])) for f in losses], dtype=float)
        loss_sp = spearman(lp.tolist(), lm.tolist()) if len(losses) >= 10 else None
        loss_terciles = _terciles(lp, lm)
        for row in loss_terciles:
            row["metric"] = "LOSS_ABS"
    price_terciles = _terciles(prices, abs_u)
    for row in price_terciles:
        row["metric"] = "ABS_UTILITY"
    return {
        "SPEARMAN_ABS_UTILITY_VS_FILL_PRICE": sp,
        "SPEARMAN_ABS_LOSS_VS_FILL_PRICE": loss_sp,
        "FILL_N": len(fills),
        "LOSS_N": len(losses),
        "price_terciles": price_terciles,
        "loss_terciles": loss_terciles,
    }


def fill_row(name: str, pack: dict[str, Any]) -> dict[str, Any]:
    admitted = int(pack.get("admitted_n") or 0)
    fill_n = int(pack.get("fill_n") or 0)
    return {
        "arm": name,
        "ADMITTED_N": admitted,
        "FILL_N": fill_n,
        "EXPIRED_N": int(pack.get("expired_n") or 0),
        "TRADE_N": int(pack.get("trade_count") or 0),
        "OOF_FILL_RATE": (float(fill_n) / float(admitted)) if admitted else pack.get("fill_rate"),
        "NET_PNL": pack.get("net_pnl_yen_100"),
    }


def reconstruct_nested_fill(
    v2_outer: list[dict[str, Any]],
    spec_folds: dict[tuple[str, str], dict[str, Any]],
) -> dict[str, Any]:
    admitted = fill_n = expired = trades = 0
    pnl = 0.0
    mismatch = 0
    missing = 0
    rows = []
    for r in v2_outer:
        day = str(r.get("outer_day") or "")
        sid = str(r.get("spec_id") or "")
        rec = spec_folds.get((sid, day))
        v2_pnl = _f(r.get("outer_pnl"))
        if rec is None:
            missing += 1
            rows.append({"outer_day": day, "spec_id": sid, "status": "MISSING_FIXED_FOLD"})
            continue
        admitted += int(rec.get("admitted_n") or 0)
        fill_n += int(rec.get("fill_n") or 0)
        expired += int(rec.get("expired_n") or 0)
        trades += int(rec.get("trade_count") or 0)
        fpnl = float(rec.get("pnl_yen_100") or 0.0)
        pnl += fpnl
        if v2_pnl is not None and abs(fpnl - float(v2_pnl)) > 1e-6:
            mismatch += 1
        rows.append(
            {
                "outer_day": day,
                "spec_id": sid,
                "status": "OK",
                "admitted_n": rec.get("admitted_n"),
                "fill_n": rec.get("fill_n"),
                "expired_n": rec.get("expired_n"),
                "trade_count": rec.get("trade_count"),
                "fixed_pnl": fpnl,
                "v2_outer_pnl": v2_pnl,
            }
        )
    return {
        "arm": "NESTED_CANDIDATE",
        "ADMITTED_N": admitted,
        "FILL_N": fill_n,
        "EXPIRED_N": expired,
        "TRADE_N": trades,
        "OOF_FILL_RATE": (float(fill_n) / float(admitted)) if admitted else None,
        "NET_PNL": pnl,
        "NESTED_FIXED_DAY_PNL_MISMATCH_N": mismatch,
        "NESTED_FIXED_DAY_MISSING_N": missing,
        "days": rows,
    }
