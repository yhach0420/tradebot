"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.cause_first_mechanism_discovery_v1 import ANALYSIS_ID
from research.cause_first_mechanism_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Data_Split",
    "Availability",
    "Feature_Catalog",
    "Deployability",
    "Market_State",
    "Sector_State",
    "Stock_State",
    "Mechanism_Hypotheses",
    "Outcome_Diagnostics",
    "Temporal_Stability",
    "Symbol_Sector_Stability",
    "Rejected_Hypotheses",
    "Surviving_Hypotheses",
    "Next_Experiments",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _dates(name: str, dates: list[str]) -> list[dict[str, Any]]:
    return [{"partition": name, "i": i, "date": d} for i, d in enumerate(dates, start=1)]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    split = dict(report.get("split") or {})
    mech = dict(report.get("mechanisms") or {})
    hyps = list(mech.get("hypotheses") or [])
    avail = dict(report.get("availability") or {})
    avail_rows = [{"symbol": s, **dict(v)} for s, v in dict(avail.get("symbol_first_last") or {}).items()]
    outcome = []
    temporal = []
    conc = []
    for h in hyps:
        d = dict(h.get("discovery") or {})
        c = dict(h.get("confirmation") or {})
        outcome.append({"id": h.get("id"), "status": h.get("status"), "partition": "discovery", **{k: d.get(k) for k in ("treat_n", "ctrl_n", "mean_treat_x0_h15_bps", "mean_ctrl_x0_h15_bps", "mean_diff_x0_h15_bps", "hit_rate_treat", "mean_treat_x1_h15_bps")}})
        outcome.append({"id": h.get("id"), "status": h.get("status"), "partition": "confirmation", **{k: c.get(k) for k in ("treat_n", "ctrl_n", "mean_treat_x0_h15_bps", "mean_ctrl_x0_h15_bps", "mean_diff_x0_h15_bps", "hit_rate_treat", "mean_treat_x1_h15_bps")}})
        temporal.append({"id": h.get("id"), "status": h.get("status"), "discovery_day_stability": d.get("day_stability"), "discovery_day_n": d.get("day_n_for_gate"), "confirmation_day_stability": c.get("day_stability"), "robustness": h.get("robustness_clocks")})
        conc.append({"id": h.get("id"), "status": h.get("status"), "top_symbol": d.get("top_symbol"), "top_symbol_share": d.get("top_symbol_share_of_positive_bps"), "top_sector": d.get("top_sector"), "top_sector_share": d.get("top_sector_share_of_positive_bps")})
    hyp_rows = []
    for h in hyps:
        hyp_rows.append({k: h.get(k) for k in ("id", "layer", "hypothesis", "chain", "status", "reasons") if k in h or True})
        hyp_rows[-1]["discovery_diff_bps"] = (h.get("discovery") or {}).get("mean_diff_x0_h15_bps")
        hyp_rows[-1]["confirmation_diff_bps"] = (h.get("confirmation") or {}).get("mean_diff_x0_h15_bps")
    return {
        "Data_Split": _kv({k: split.get(k) for k in ("ok", "n_sessions", "split_sha256", "random_split", "frozen_validation_hidden_from_discovery")})
        + _dates("discovery", list(split.get("discovery_dates") or []))
        + _dates("confirmation", list(split.get("confirmation_dates") or []))
        + _dates("frozen_validation", list(split.get("frozen_validation_dates") or [])),
        "Availability": avail_rows or _kv(avail),
        "Feature_Catalog": list(report.get("feature_catalog") or []),
        "Deployability": _kv(dict(report.get("deployability") or {})),
        "Market_State": list(report.get("market_state_discovery_0930") or []),
        "Sector_State": list(report.get("sector_state_discovery_0930") or []),
        "Stock_State": list(report.get("stock_state_summary_discovery_0930") or []),
        "Mechanism_Hypotheses": hyp_rows,
        "Outcome_Diagnostics": outcome,
        "Temporal_Stability": temporal,
        "Symbol_Sector_Stability": conc,
        "Rejected_Hypotheses": list(mech.get("rejected") or []) or [{"empty": True}],
        "Surviving_Hypotheses": list(mech.get("surviving") or []) or [{"empty": True}],
        "Next_Experiments": list(report.get("next_experiments") or []),
        "Safety": _kv(dict(report.get("answers") or {})) + _kv(dict(report.get("decision") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    disc = dict(a.get("exact_discovery_dates") or {})
    conf = dict(a.get("exact_confirmation_dates") or {})
    val = dict(a.get("exact_frozen_validation_dates") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            "## Required answers",
            "",
            f"Exact Discovery dates? **{disc.get('first')}–{disc.get('last')} n={disc.get('n')}**",
            f"Exact Confirmation dates? **{conf.get('first')}–{conf.get('last')} n={conf.get('n')}**",
            f"Exact Frozen Validation dates? **{val.get('first')}–{val.get('last')} n={val.get('n')}**",
            f"Split SHA? `{a.get('split_sha')}`",
            f"Frozen Validation accessed during Discovery? **{a.get('frozen_validation_accessed_during_discovery')}**",
            f"105-stock pool used as mandatory final universe? **{a.get('pool_105_used_as_mandatory_final_universe')}**",
            f"Kabu 50-slot constraint applied to research? **{a.get('kabu_50_slot_constraint_applied_to_research')}**",
            f"Additional research data allowed? **{a.get('additional_research_data_allowed')}**",
            f"External historical context allowed? **{a.get('external_historical_context_allowed')}**",
            f"PnL allowed in strategy research? **{a.get('pnl_allowed_in_strategy_research')}**",
            f"ENTRY-only success allowed? **{a.get('entry_only_success_allowed')}**",
            f"Historical Bid/Ask inferred? **{a.get('historical_bid_ask_inferred')}**",
            f"Same-bar leakage? **{a.get('same_bar_leakage')}**",
            f"Panel-conditioned label retained? **{a.get('panel_conditioned_label_retained')}**",
            f"Strategy search started in this run? **{a.get('strategy_search_started_in_this_run')}**",
            f"Runtime modified? **{a.get('runtime_modified')}**",
            f"Paper modified? **{a.get('paper_modified')}**",
            f"Kabu registration modified? **{a.get('kabu_registration_modified')}**",
            f"20260914 live modified? **{a.get('live_20260914_modified')}**",
            f"submit/cancel/live? **{a.get('submit_cancel_live')}**",
            f"VERDICT? **{a.get('VERDICT')}**",
            f"NEXT? **{a.get('NEXT')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
    assert (OUT / "audit.xlsx").is_file()
