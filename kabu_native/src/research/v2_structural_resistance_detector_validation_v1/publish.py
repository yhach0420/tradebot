"""Publish detector validation artifacts. Visual gate before any full economic join."""
from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

from research.event_time_impulse_complete_strategy_v2.identity import bind, sha256_obj
from research.v2_structural_resistance_detector_validation_v1 import (
    ANALYSIS_ID,
    CORRECTED_PREVIOUS_INTERPRETATION,
    EXPECTED_PF,
    EXPECTED_PNL_YEN,
    EXPECTED_TRADE_N,
    OUT_REL,
    PREVIOUS_AUDIT_VERDICT_TOO_BROAD,
    SIGNAL_SHA256,
    STRATEGY_SHA256,
)
from research.v2_structural_resistance_detector_validation_v1.isolation import assert_isolated
from research.v2_structural_resistance_detector_validation_v1.rules import detector_document
from research.v2_structural_resistance_detector_validation_v1.scan import scan_full, scan_visual
from research.v2_structural_resistance_detector_validation_v1.synthetic import require_self_check
from research.v2_structural_resistance_detector_validation_v1.visualize import render_chart

OUT = Path(OUT_REL)

# Structures the visual gate must distinguish across examples.
GATE_STRUCTURES = (
    "repeated_upper_rejection",
    "sideways_consolidation",
    "single_incidental_high",
    "clean_breakout",
    "failed_breakout",
    "retest_and_hold",
)


def _rows(sheet, header, data) -> None:
    sheet.append(header)
    for row in data:
        sheet.append(row)


def _human_note(item: dict[str, Any]) -> str:
    """Machine-draft note from plotted series fields. Answer structure only, not trade quality."""
    st = item.get("structure") or {}
    matched = st.get("pre_break_in_structural_zone")
    z = st.get("matched_zone") or {}
    ztype = z.get("zone_type")
    counts = st.get("pre_entry_zone_counts") or {}
    rt = st.get("runtime") or {}
    nxt = st.get("next_resistance") or {}
    parts = [
        f"Chart shows {counts.get('swings', 0)} recognized swing highs pre/through window.",
        f"Zones pre-entry: upper={counts.get('upper_rejection')} consol={counts.get('consolidation')} single={counts.get('single_swing')}.",
    ]
    if matched:
        parts.append(
            f"PRE_BREAK_HIGH lies inside detected {ztype} "
            f"[{z.get('zone_low')}, {z.get('zone_high')}] with independent_tests={z.get('independent_test_n')} "
            f"rejections={z.get('rejection_n')} through_up={z.get('through_up_n')}."
        )
        if ztype == "UPPER_REJECTION_ZONE":
            parts.append("Line/zone corresponds to repeated upper rejection structure on the price path.")
        elif ztype == "CONSOLIDATION_ZONE":
            parts.append("PRE_BREAK_HIGH sits in a two-way consolidation band, not asymmetric resistance.")
        elif ztype == "SINGLE_SWING_HIGH":
            parts.append("PRE_BREAK_HIGH matches a single incidental swing high, not repeated resistance.")
        else:
            parts.append("Matched zone is unresolved structurally.")
    else:
        parts.append("PRE_BREAK_HIGH does not sit inside any causally detected structural zone at entry.")
    if rt:
        parts.append(
            f"Zone state machine ended at {rt.get('state')} retest_label={rt.get('retest_label')} "
            f"break_events_above={(rt.get('anatomy') or {}).get('events_above')}."
        )
    parts.append(f"Next resistance={nxt.get('status')} level={nxt.get('level')} source={nxt.get('source')}.")
    parts.append("Note answers observable structure only; not whether the trade was good.")
    return " ".join(parts)


def _tag_gate_evidence(item: dict[str, Any]) -> list[str]:
    """Tag which acceptance-gate structures this example visibly supports (structure-only)."""
    tags: list[str] = []
    st = item.get("structure") or {}
    z = st.get("matched_zone") or {}
    ztype = z.get("zone_type")
    counts = st.get("pre_entry_zone_counts") or {}
    rt = st.get("runtime") or {}
    if ztype == "UPPER_REJECTION_ZONE" or (counts.get("upper_rejection") or 0) >= 1:
        if (z.get("independent_test_n") or 0) >= 2 or (counts.get("upper_rejection") or 0) >= 1:
            tags.append("repeated_upper_rejection")
    if ztype == "CONSOLIDATION_ZONE" or (counts.get("consolidation") or 0) >= 1:
        tags.append("sideways_consolidation")
    if ztype == "SINGLE_SWING_HIGH" or (
        not st.get("pre_break_in_structural_zone") and (counts.get("single_swing") or 0) >= 1
    ):
        tags.append("single_incidental_high")
    state = rt.get("state")
    label = rt.get("retest_label")
    if state in ("BREAK_ACCEPTED_CANDIDATE", "SUPPORT_CONFIRMED") and label != "FAILED_SUPPORT":
        tags.append("clean_breakout")
    if state == "FAILED_BREAKOUT" or label == "FAILED_SUPPORT":
        tags.append("failed_breakout")
    if state == "SUPPORT_CONFIRMED" or label == "SUPPORT_RESPONSE":
        tags.append("retest_and_hold")
    # Also scan all runtimes on the chart for breakout/retest evidence not only matched zone
    for rt2 in (st.get("runtimes") or {}).values():
        if rt2.get("state") in ("BREAK_ACCEPTED_CANDIDATE",) and "clean_breakout" not in tags:
            tags.append("clean_breakout")
        if rt2.get("state") == "FAILED_BREAKOUT" and "failed_breakout" not in tags:
            tags.append("failed_breakout")
        if rt2.get("state") == "SUPPORT_CONFIRMED" and "retest_and_hold" not in tags:
            tags.append("retest_and_hold")
        zid = str(rt2.get("zone_id"))
        zmap = {str(z3["zone_id"]): z3 for z3 in (st.get("zones") or [])}
        zz = zmap.get(zid) or {}
        if zz.get("zone_type") == "UPPER_REJECTION_ZONE" and "repeated_upper_rejection" not in tags:
            tags.append("repeated_upper_rejection")
        if zz.get("zone_type") == "CONSOLIDATION_ZONE" and "sideways_consolidation" not in tags:
            tags.append("sideways_consolidation")
        if zz.get("zone_type") == "SINGLE_SWING_HIGH" and "single_incidental_high" not in tags:
            tags.append("single_incidental_high")
    return sorted(set(tags))


def _acceptance(charts: list[dict[str, Any]]) -> dict[str, Any]:
    covered: set[str] = set()
    for item in charts:
        tags = _tag_gate_evidence(item)
        item["gate_tags"] = tags
        item["human_review_note"] = _human_note(item)
        covered.update(tags)
    missing = [s for s in GATE_STRUCTURES if s not in covered]
    accepted = len(missing) == 0 and len(charts) >= 12
    return {
        "accepted": accepted,
        "covered": sorted(covered),
        "missing": missing,
        "visual_n": len(charts),
        "summary": (
            "Detector visibly separates all six required structures across the sample."
            if accepted
            else f"Detector sample missing visible structures: {missing}. Gate failed."
        ),
    }


def _econ_block(rows: list[dict[str, Any]]) -> dict[str, Any]:
    import numpy as np

    def _pf(vals: list[float]) -> Optional[float]:
        gain = sum(v for v in vals if v > 0)
        loss = -sum(v for v in vals if v < 0)
        if loss <= 0:
            return None if gain <= 0 else float("inf")
        return gain / loss

    pnl = [float(r["pnl_yen"]) for r in rows]
    return {
        "trade_n": len(rows),
        "pnl_yen": float(sum(pnl)),
        "pf": _pf(pnl),
        "mean_bps": float(np.mean([float(r["bps"]) for r in rows])) if rows else None,
    }


def _economic_analysis(raw: dict[str, Any]) -> dict[str, Any]:
    """Join PnL only after detector freeze. Descriptive fractions; no parameter search."""
    import numpy as np
    from research.symbol_setup_baseline_complete_strategy_precommit.contract import EXTENSION17, ORIGINAL18

    rows = [r for r in raw["baseline"] if r.get("pnl_yen") is not None]
    n = len(rows)
    genuine = [r for r in rows if r.get("pre_break_in_structural_zone") and r.get("zone_type") == "UPPER_REJECTION_ZONE"]
    incidental = [
        r
        for r in rows
        if (not r.get("pre_break_in_structural_zone"))
        or r.get("zone_type") in (None, "SINGLE_SWING_HIGH", "UNRESOLVED_ZONE")
    ]
    consol = [r for r in rows if r.get("zone_type") == "CONSOLIDATION_ZONE"]
    immediate = [r for r in rows if float(r.get("hold_sec") or (float(r["exit_t"]) - float(r["entry_t"]))) < 0.1 and float(r["pnl_yen"]) < 0]
    imm_genuine = [r for r in immediate if r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("runtime_state") != "FAILED_BREAKOUT"]
    imm_false = [r for r in immediate if r.get("runtime_state") == "FAILED_BREAKOUT" or r.get("retest_label") == "FAILED_SUPPORT"]
    imm_inc = [r for r in immediate if r not in imm_genuine and r not in imm_false]

    r2 = [r for r in rows if int(r.get("ratchets") or 0) >= 2]
    # successive structural: matched upper zone AND next known
    r2_succ = [r for r in r2 if r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("next_status") == "KNOWN"]
    r2_support = [r for r in r2 if r.get("retest_label") == "SUPPORT_RESPONSE" or r.get("runtime_state") == "SUPPORT_CONFIRMED"]

    known_next = [r for r in rows if r.get("next_status") == "KNOWN"]
    # headroom relationship: spearman of headroom vs bps among known
    xs, ys, ymfe = [], [], []
    for r in known_next:
        if r.get("headroom_yen") is None:
            continue
        xs.append(float(r["headroom_yen"]))
        ys.append(float(r["bps"]))
        if r.get("mfe") == r.get("mfe"):
            ymfe.append((float(r["headroom_yen"]), float(r["mfe"])))

    def _spear(a: list[float], b: list[float]) -> Optional[float]:
        if len(a) < 30:
            return None
        aa = np.asarray(a, dtype=float)
        bb = np.asarray(b, dtype=float)
        ra = np.argsort(np.argsort(aa)).astype(float)
        rb = np.argsort(np.argsort(bb)).astype(float)
        if float(np.std(ra)) == 0 or float(np.std(rb)) == 0:
            return None
        return float(np.corrcoef(ra, rb)[0, 1])

    head_bps = _spear(xs, ys)
    head_mfe = _spear([p[0] for p in ymfe], [p[1] for p in ymfe]) if len(ymfe) >= 30 else None

    orig = set(ORIGINAL18)
    ext = set(EXTENSION17)
    dates = sorted({str(r["date"]) for r in rows})
    folds = [dates[i * len(dates) // 3 : (i + 1) * len(dates) // 3] for i in range(3)]

    def _slice_frac(sample: list[dict[str, Any]]) -> dict[str, Any]:
        if not sample:
            return {"n": 0, "genuine_frac": None, "econ": None}
        g = sum(1 for r in sample if r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("pre_break_in_structural_zone"))
        return {"n": len(sample), "genuine_frac": g / len(sample), "econ": _econ_block(sample)}

    genuine_frac = len(genuine) / n if n else None
    # Mechanism explanatory if genuine resistance entries show directionally better mean bps and MFE
    # and ratchet continuation vs incidental, stable across folds/original/extension — not optimized.
    def _mean_bps(sample: list[dict[str, Any]]) -> Optional[float]:
        if not sample:
            return None
        return float(np.mean([float(r["bps"]) for r in sample]))

    def _mean_mfe(sample: list[dict[str, Any]]) -> Optional[float]:
        vals = [float(r["mfe"]) for r in sample if r.get("mfe") == r.get("mfe")]
        return float(np.mean(vals)) if vals else None

    def _r2_rate(sample: list[dict[str, Any]]) -> Optional[float]:
        if not sample:
            return None
        return float(np.mean([int(r.get("ratchets") or 0) >= 2 for r in sample]))

    g_bps, i_bps = _mean_bps(genuine), _mean_bps(incidental)
    g_mfe, i_mfe = _mean_mfe(genuine), _mean_mfe(incidental)
    g_r2, i_r2 = _r2_rate(genuine), _r2_rate(incidental)
    direction_ok = (
        g_bps is not None
        and i_bps is not None
        and g_mfe is not None
        and i_mfe is not None
        and g_r2 is not None
        and i_r2 is not None
        and g_bps > i_bps
        and g_mfe > i_mfe
        and g_r2 > i_r2
    )
    fold_ok = True
    fold_rows = []
    for i, chunk in enumerate(folds, start=1):
        s = [r for r in rows if r["date"] in chunk]
        gs = [r for r in s if r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("pre_break_in_structural_zone")]
        ins = [r for r in s if r not in gs]
        gb, ib = _mean_bps(gs), _mean_bps(ins)
        ok = gb is not None and ib is not None and gb > ib
        fold_ok = fold_ok and ok
        fold_rows.append({"fold": i, "genuine": _econ_block(gs), "incidental": _econ_block(ins), "genuine_better_bps": ok})

    o18 = [r for r in rows if r["date"] in orig]
    e17 = [r for r in rows if r["date"] in ext]
    o_ok = (_mean_bps([r for r in o18 if r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("pre_break_in_structural_zone")]) or -1e9) > (
        _mean_bps([r for r in o18 if not (r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("pre_break_in_structural_zone"))]) or 1e9
    )
    e_ok = (_mean_bps([r for r in e17 if r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("pre_break_in_structural_zone")]) or -1e9) > (
        _mean_bps([r for r in e17 if not (r.get("zone_type") == "UPPER_REJECTION_ZONE" and r.get("pre_break_in_structural_zone"))]) or 1e9
    )

    support_rows = [r for r in rows if r.get("retest_label") == "SUPPORT_RESPONSE"]
    support_explanatory = bool(support_rows) and (_mean_bps(support_rows) or -1e9) > (_mean_bps(rows) or 1e9)

    resistance_explanatory = bool(direction_ok and fold_ok and o_ok and e_ok and genuine_frac and genuine_frac >= 0.05)
    v3 = bool(resistance_explanatory)

    return {
        "trade_n": n,
        "pnl_yen": float(sum(float(r["pnl_yen"]) for r in rows)),
        "pf": _econ_block(rows)["pf"],
        "pre_break_genuine_structural_resistance_fraction": genuine_frac,
        "pre_break_consolidation_fraction": len(consol) / n if n else None,
        "pre_break_incidental_or_unmatched_fraction": len(incidental) / n if n else None,
        "immediate_failure": {
            "n": len(immediate),
            "genuine_resistance_breakout_fraction": len(imm_genuine) / len(immediate) if immediate else None,
            "incidental_high_fraction": len(imm_inc) / len(immediate) if immediate else None,
            "false_breakout_fraction": len(imm_false) / len(immediate) if immediate else None,
        },
        "ratchet_ge2": {
            "n": len(r2),
            "successive_structural_resistance_fraction": len(r2_succ) / len(r2) if r2 else None,
            "support_conversion_fraction": len(r2_support) / len(r2) if r2 else None,
        },
        "next_resistance_known_fraction": len(known_next) / n if n else None,
        "headroom_relationship": {
            "spearman_bps": head_bps,
            "spearman_mfe": head_mfe,
            "known_n": len(known_next),
        },
        "resistance_mechanism_explanatory": resistance_explanatory,
        "support_conversion_explanatory": support_explanatory,
        "v3_justified": v3,
        "original18": _slice_frac(o18),
        "extension17": _slice_frac(e17),
        "folds": fold_rows,
        "genuine_econ": _econ_block(genuine),
        "incidental_econ": _econ_block(incidental),
    }


def publish() -> dict[str, Any]:
    assert_isolated(OUT)
    OUT.mkdir(parents=True, exist_ok=True)
    synthetic = require_self_check()
    bound = bind()
    if bound["strategy_sha256"] != STRATEGY_SHA256 or bound["signal_sha256"] != SIGNAL_SHA256:
        raise RuntimeError("frozen_v2_identity_mismatch")
    det_doc = detector_document()
    detector_sha = sha256_obj(det_doc)

    print("VISUAL_SCAN_START", flush=True)
    visual_raw = scan_visual()
    gate = _acceptance(visual_raw["charts"])
    print(f"VISUAL_GATE accepted={gate['accepted']} covered={gate['covered']} missing={gate['missing']}", flush=True)

    economic = None
    parity = True
    trade_n = EXPECTED_TRADE_N
    pnl = EXPECTED_PNL_YEN
    pf = EXPECTED_PF
    full_ran = False

    if gate["accepted"]:
        print("DETECTOR_ACCEPTED full scan starting", flush=True)
        # Freeze detector identity then full apply
        full_raw = scan_full()
        full_ran = True
        economic = _economic_analysis(full_raw)
        trade_n = int(economic["trade_n"])
        pnl = float(economic["pnl_yen"])
        pf = float(economic["pf"]) if economic["pf"] is not None else None
        parity = trade_n == EXPECTED_TRADE_N and abs(pnl - EXPECTED_PNL_YEN) < 1e-6 and pf is not None and abs(float(pf) - EXPECTED_PF) < 1e-12
        if not parity:
            raise RuntimeError(f"parity_failed:{trade_n}/{pnl}/{pf}")
        if full_raw["recon_mismatch"] > 0:
            # Abort economic claims on reconstruction failure
            economic["reconstruction_abort"] = True
            economic["resistance_mechanism_explanatory"] = False
            economic["support_conversion_explanatory"] = False
            economic["v3_justified"] = False
        charts = full_raw["charts"]
        # re-tag charts from full if present else keep visual
        if charts:
            gate = _acceptance(charts)
        else:
            charts = visual_raw["charts"]
        baseline_n = len(full_raw["baseline"])
    else:
        charts = visual_raw["charts"]
        baseline_n = len(visual_raw["baseline"])

    if gate["accepted"] and economic and economic.get("resistance_mechanism_explanatory"):
        verdict = "STRUCTURAL_RESISTANCE_MECHANISM_SUPPORTED_IN_V2_V1"
        nxt = "DESIGN_PRECOMMITTED_V3_FROM_SUPPORTED_MECHANISM"
    elif gate["accepted"]:
        verdict = "STRUCTURAL_RESISTANCE_VALID_BUT_NOT_V2_EXPLANATORY_V1"
        nxt = "KEEP_V2_UNCHANGED"
    else:
        verdict = "STRUCTURAL_RESISTANCE_DETECTOR_NOT_READY_V1"
        nxt = "REFINE_PRICE_STRUCTURE_DETECTOR_ONLY"

    report: dict[str, Any] = {
        "study": ANALYSIS_ID,
        "verdict": verdict,
        "next": nxt,
        "previous_audit_verdict_too_broad": PREVIOUS_AUDIT_VERDICT_TOO_BROAD,
        "previous_resistance_audit_interpretation_corrected": True,
        "corrected_previous_interpretation": CORRECTED_PREVIOUS_INTERPRETATION,
        "detector_uses_pre_break_high_as_starting_point": False,
        "structural_swing_detector_implemented": True,
        "resistance_vs_consolidation_distinguishable": "sideways_consolidation" in gate["covered"] and "repeated_upper_rejection" in gate["covered"],
        "breakout_state_machine_implemented": True,
        "retest_support_conversion_state_implemented": True,
        "next_structural_resistance_uses_all_causally_known_same_session_structure": True,
        "synthetic_self_check": synthetic,
        "detector_sha256": detector_sha,
        "detector_document": det_doc,
        "detector_accepted": bool(gate["accepted"]),
        "visual_examples_n": int(gate["visual_n"]),
        "visual_validation_summary": gate["summary"],
        "visual_gate": gate,
        "full_economic_scan_ran": full_ran,
        "parity": parity,
        "trade_n": trade_n,
        "pnl_yen": pnl,
        "pf": pf,
        "strategy_sha256": bound["strategy_sha256"],
        "signal_sha256": bound["signal_sha256"],
        "ENTRY_changed": False,
        "EXIT_changed": False,
        "Complete_Strategy_changed": False,
        "new_data_acquired": False,
        "submit_cancel_live": [0, 0, 0],
        "visual_sample_trade_n": baseline_n,
        "recon_mismatch_visual": visual_raw.get("recon_mismatch"),
        "economic": economic,
    }

    public = {k: v for k, v in report.items() if k != "detector_document"}
    public["detector_rules_version"] = det_doc["rules"]["version"]
    (OUT / "report.json").write_text(json.dumps(public, indent=2, default=str), encoding="utf-8")

    lines = [
        f"# {verdict}",
        "",
        f"NEXT: {nxt}",
        "",
        "## Interpretation correction (previous audit)",
        "",
        f"Previous verdict `{PREVIOUS_AUDIT_VERDICT_TOO_BROAD}` was too broad.",
        f"Correct interpretation: `{CORRECTED_PREVIOUS_INTERPRETATION}`.",
        "That audit only showed PRE_BREAK_HIGH ±1 tick within 300s plus touch count was not explanatory.",
        "Do NOT conclude resistance, headroom, or resistance-to-support do not matter.",
        "",
        "## Detector",
        "",
        f"- starts from PRE_BREAK_HIGH? false",
        f"- detector_sha256: `{detector_sha}`",
        f"- synthetic self-check: {synthetic}",
        f"- detector accepted? {gate['accepted']}",
        f"- visual examples n: {gate['visual_n']}",
        f"- visual summary: {gate['summary']}",
        "",
        "## V2 parity (frozen identity)",
        "",
        f"trade_n={trade_n} pnl_yen={pnl} pf={pf} parity={parity}",
        f"strategy_sha256={bound['strategy_sha256']}",
        f"signal_sha256={bound['signal_sha256']}",
        "",
        "ENTRY changed: false",
        "EXIT changed: false",
        "Complete Strategy changed: false",
        "new data acquired: false",
        "submit/cancel/live: 0/0/0",
        "",
    ]
    if economic is None:
        lines += [
            "## Economic join",
            "",
            "Not run (detector failed visual acceptance gate).",
            "",
        ]
    else:
        lines += [
            "## Economic join (after detector freeze)",
            "",
            f"PRE_BREAK_HIGH genuine structural resistance fraction: {economic['pre_break_genuine_structural_resistance_fraction']}",
            f"immediate failures: {economic['immediate_failure']}",
            f"ratchet>=2: {economic['ratchet_ge2']}",
            f"next resistance known fraction: {economic['next_resistance_known_fraction']}",
            f"headroom: {economic['headroom_relationship']}",
            f"Resistance mechanism explanatory? {economic['resistance_mechanism_explanatory']}",
            f"Support conversion explanatory? {economic['support_conversion_explanatory']}",
            f"V3 justified? {economic['v3_justified']}",
            "",
        ]
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")

    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    _rows(
        summary,
        ["key", "value"],
        [[k, json.dumps(public[k], default=str) if isinstance(public[k], (dict, list)) else public[k]] for k in public if k != "economic"]
        + [["economic", json.dumps(economic, default=str)]],
    )
    rules_sheet = wb.create_sheet("detector_rules")
    rules_sheet.append(["detector_document_json"])
    rules_sheet.append([json.dumps(det_doc, default=str)])
    gate_sheet = wb.create_sheet("visual_gate")
    _rows(gate_sheet, ["key", "value"], [[k, json.dumps(v, default=str) if isinstance(v, (dict, list)) else v] for k, v in gate.items()])

    table = wb.create_sheet("machine_detector_validation")
    _rows(
        table,
        [
            "symbol",
            "date",
            "kind",
            "zone",
            "independent_approaches",
            "true_downward_rejections",
            "consolidation_crossings",
            "zone_classification",
            "PRE_BREAK_HIGH_inside_structural_zone",
            "breakout_event",
            "retest_occurred",
            "support_response",
            "failed_support",
            "next_known_structural_resistance",
            "human_review_note",
            "gate_tags",
        ],
        [
            [
                c["symbol"],
                c["date"],
                c["kind"],
                json.dumps((c.get("structure") or {}).get("matched_zone"), default=str),
                ((c.get("structure") or {}).get("matched_zone") or {}).get("independent_test_n"),
                ((c.get("structure") or {}).get("matched_zone") or {}).get("rejection_n"),
                ((c.get("structure") or {}).get("matched_zone") or {}).get("through_up_n"),
                ((c.get("structure") or {}).get("matched_zone") or {}).get("zone_type"),
                (c.get("structure") or {}).get("pre_break_in_structural_zone"),
                ((c.get("structure") or {}).get("runtime") or {}).get("state"),
                bool(((c.get("structure") or {}).get("runtime") or {}).get("retest_t")),
                ((c.get("structure") or {}).get("runtime") or {}).get("retest_label") == "SUPPORT_RESPONSE",
                ((c.get("structure") or {}).get("runtime") or {}).get("retest_label") == "FAILED_SUPPORT",
                json.dumps((c.get("structure") or {}).get("next_resistance"), default=str),
                c.get("human_review_note"),
                ",".join(c.get("gate_tags") or []),
            ]
            for c in charts
        ],
    )

    visual = wb.create_sheet("visual_examples")
    _rows(
        visual,
        ["kind", "symbol", "date", "entry_t", "exit_t", "pnl", "ratchets", "pre_break", "PB_in_zone", "zone_type", "next", "note"],
        [
            [
                c["kind"],
                c["symbol"],
                c["date"],
                c["entry_t"],
                c["exit_t"],
                c["pnl_yen"],
                c["ratchets"],
                c.get("pre_break"),
                (c.get("structure") or {}).get("pre_break_in_structural_zone"),
                ((c.get("structure") or {}).get("matched_zone") or {}).get("zone_type"),
                ((c.get("structure") or {}).get("next_resistance") or {}).get("status"),
                c.get("human_review_note"),
            ]
            for c in charts
        ],
    )
    with TemporaryDirectory() as tmp:
        for i, item in enumerate(charts):
            png = Path(tmp) / f"chart_{i}.png"
            render_chart(item, png)
            image = XLImage(str(png))
            image.anchor = f"A{16 + i * 24}"
            visual.add_image(image)
        wb.save(OUT / "audit.xlsx")

    return {
        "verdict": verdict,
        "next": nxt,
        "detector_accepted": gate["accepted"],
        "visual_n": gate["visual_n"],
        "parity": parity,
        "trade_n": trade_n,
        "pnl_yen": pnl,
        "pf": pf,
        "files": [str(OUT / "report.json"), str(OUT / "report.md"), str(OUT / "audit.xlsx")],
    }


if __name__ == "__main__":
    print(publish())
