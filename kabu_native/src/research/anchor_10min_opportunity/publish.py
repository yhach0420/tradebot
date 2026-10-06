"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.anchor_10min_opportunity import (
    ANALYSIS_ID,
    DATA_STATUS,
    DOCUMENT_ID,
    FUTURE_FORWARD_REQUIRED,
    NEW_GRID_TRUE_HOLDOUT,
    TASK_LABEL,
)
from research.anchor_10min_opportunity.grids import grid_contract

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "anchor_10min_opportunity"
JST = timezone(timedelta(hours=9))
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            if v == float("inf"):
                v = "Infinity"
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _sess_line(st: dict[str, Any] | None) -> str:
    if not st:
        return "empty"
    return f"n={st.get('trades')} pnl={st.get('pnl')} PF={st.get('PF')}"


def build_report(*, provenance: dict[str, Any], inventory_check: dict[str, Any], analysis: dict[str, Any]) -> dict[str, Any]:
    g = grid_contract()
    ref, aug, uni = analysis["ref_stats"], analysis["aug_stats"], analysis["uni_stats"]
    iso, tail, u10n = analysis["iso_added_stats"], analysis["iso_tail_stats"], analysis["iso_u10new_stats"]
    de = analysis["decomp"]
    missed = analysis["missed"]
    fr = analysis["first_reentry"]
    required = {
        "REFERENCE_ANCHORS": g["REFERENCE_ANCHORS"],
        "PHASE_A_ADDED_ANCHORS": g["PHASE_A_ADDED_ANCHORS"],
        "REFERENCE_TRADES": ref.get("trades"),
        "REFERENCE_PNL": ref.get("pnl"),
        "REFERENCE_PF": ref.get("PF"),
        "REFERENCE_MAXDD": ref.get("maxDD"),
        "AUGMENTED_TRADES": aug.get("trades"),
        "AUGMENTED_PNL": aug.get("pnl"),
        "AUGMENTED_PF": aug.get("PF"),
        "AUGMENTED_MAXDD": aug.get("maxDD"),
        "ADDED_ANCHOR_SELECTIONS": iso.get("selected"),
        "ADDED_ANCHOR_FILLS": iso.get("fills"),
        "ADDED_ANCHOR_FILL_RATE": iso.get("fill_rate"),
        "ADDED_ANCHOR_DIRECT_PNL": iso.get("direct_pnl"),
        "DISPLACED_LATER_ENTRY_N": de.get("DISPLACED_LATER_ENTRY_N"),
        "BLOCKED_LATER_ENTRY_N": de.get("BLOCKED_LATER_ENTRY_N"),
        "SAME_SYMBOL_STATE_EFFECT": de.get("SAME_SYMBOL_STATE_EFFECT"),
        "CAPACITY_EFFECT": de.get("CAPACITY_EFFECT"),
        "REENTRY_EFFECT": de.get("REENTRY_EFFECT"),
        "PHASE_A_PNL_DELTA": de.get("observed"),
        "PHASE_A_RESIDUAL": de.get("residual"),
        "PHASE_A_VERDICT": analysis.get("PHASE_A_VERDICT"),
        "UNIFORM10_ANCHORS": g["UNIFORM10_ANCHORS"],
        "UNIFORM10_TRADES": uni.get("trades"),
        "UNIFORM10_PNL": uni.get("pnl"),
        "UNIFORM10_PF": uni.get("PF"),
        "UNIFORM10_MAXDD": uni.get("maxDD"),
        "UNIFORM10_PNL_PER_ANCHOR": uni.get("pnl_per_anchor"),
        "UNIFORM10_PNL_PER_FILL": uni.get("pnl_per_fill"),
        "REFERENCE_FIRST_ENTRY_PNL": (fr.get("REFERENCE") or {}).get("FIRST_ENTRY_PNL"),
        "UNIFORM10_FIRST_ENTRY_PNL": (fr.get("UNIFORM10") or {}).get("FIRST_ENTRY_PNL"),
        "REFERENCE_REENTRY_PNL": (fr.get("REFERENCE") or {}).get("REENTRY_PNL"),
        "UNIFORM10_REENTRY_PNL": (fr.get("UNIFORM10") or {}).get("REENTRY_PNL"),
        "REFERENCE_DAY_POSITIVE_RATE": ref.get("day_positive_rate"),
        "UNIFORM10_DAY_POSITIVE_RATE": uni.get("day_positive_rate"),
        "AM_RESULT": f"REF {_sess_line(ref.get('AM'))} | U10 {_sess_line(uni.get('AM'))}",
        "PM_RESULT": f"REF {_sess_line(ref.get('PM'))} | U10 {_sess_line(uni.get('PM'))}",
        "TAIL_RESULT": (
            f"isolated_tail fills={tail.get('fills')} pnl={tail.get('pnl')} PF={tail.get('PF')} "
            f"| portfolio SESSION_TAIL REF={_sess_line((ref.get('tod') or {}).get('SESSION_TAIL'))} "
            f"U10={_sess_line((uni.get('tod') or {}).get('SESSION_TAIL'))}"
        ),
        "HISTORICAL_10MIN_SUPPORT": analysis.get("HISTORICAL_10MIN_SUPPORT"),
        "MISSED_OPPORTUNITY_EVIDENCE": analysis.get("MISSED_OPPORTUNITY_EVIDENCE"),
        "FUTURE_FORWARD_REQUIRED": FUTURE_FORWARD_REQUIRED,
        "verdict": analysis.get("verdict"),
    }
    report = {
        "ANALYSIS_ID": ANALYSIS_ID,
        "DOCUMENT_ID": DOCUMENT_ID,
        "label": TASK_LABEL,
        "DATA_STATUS": DATA_STATUS,
        "NEW_GRID_TRUE_HOLDOUT": NEW_GRID_TRUE_HOLDOUT,
        "generated_at_jst": datetime.now(JST).isoformat(timespec="seconds"),
        "runtime_changed": False,
        "clock_grid_changed": False,
        "strategy_changed": False,
        "best_interval_adopted": False,
        "best_start_minute_searched": False,
        "safety": {"submit": 0, "cancel": 0, "live": 0, "offline_research_only": True},
        "SOURCE_FILE": provenance.get("SOURCE_FILE"),
        "SOURCE_SHA": provenance.get("SOURCE_SHA"),
        "ANCHOR_SHA": provenance.get("ANCHOR_SHA"),
        "CANONICAL_CLOCK_GRID": provenance.get("CANONICAL_CLOCK_GRID"),
        "grid_contract": g,
        "day_contract": inventory_check,
        "note": (
            "All 17 days are HISTORICAL_DIAGNOSTIC. 20260810+ is not NEW_GRID_TRUE_HOLDOUT. "
            "Do not change CLOCK_GRID from this result. FUTURE_FORWARD_REQUIRED=true."
        ),
        **required,
        "decomp": de,
        "iso_added_stats": iso,
        "iso_tail_stats": tail,
        "iso_u10new_stats": u10n,
        "ref_stats": {k: v for k, v in ref.items() if k not in {"AM", "PM", "tod"}},
        "aug_stats": {k: v for k, v in aug.items() if k not in {"AM", "PM", "tod"}},
        "uni_stats": {k: v for k, v in uni.items() if k not in {"AM", "PM", "tod"}},
        "am_pm": {"REFERENCE": {"AM": ref.get("AM"), "PM": ref.get("PM")}, "UNIFORM10": {"AM": uni.get("AM"), "PM": uni.get("PM")}},
        "tod": {"REFERENCE": ref.get("tod"), "UNIFORM10": uni.get("tod"), "AUGMENTED": aug.get("tod")},
        "missed": missed,
        "first_reentry": analysis.get("first_reentry"),
        "daily": analysis.get("daily"),
        "lineage": analysis.get("lineage"),
        "iso_added": analysis.get("iso_added"),
        "iso_tail": analysis.get("iso_tail"),
        "ref_trades": analysis.get("ref_trades"),
        "aug_trades": analysis.get("aug_trades"),
        "uni_trades": analysis.get("uni_trades"),
    }
    return json_sanitize(report)


def write_md(report: dict[str, Any]) -> str:
    keys = [
        "REFERENCE_ANCHORS", "PHASE_A_ADDED_ANCHORS",
        "REFERENCE_TRADES", "REFERENCE_PNL", "REFERENCE_PF", "REFERENCE_MAXDD",
        "AUGMENTED_TRADES", "AUGMENTED_PNL", "AUGMENTED_PF", "AUGMENTED_MAXDD",
        "ADDED_ANCHOR_SELECTIONS", "ADDED_ANCHOR_FILLS", "ADDED_ANCHOR_FILL_RATE", "ADDED_ANCHOR_DIRECT_PNL",
        "DISPLACED_LATER_ENTRY_N", "BLOCKED_LATER_ENTRY_N",
        "SAME_SYMBOL_STATE_EFFECT", "CAPACITY_EFFECT", "REENTRY_EFFECT",
        "PHASE_A_PNL_DELTA", "PHASE_A_RESIDUAL", "PHASE_A_VERDICT",
        "UNIFORM10_ANCHORS", "UNIFORM10_TRADES", "UNIFORM10_PNL", "UNIFORM10_PF", "UNIFORM10_MAXDD",
        "UNIFORM10_PNL_PER_ANCHOR", "UNIFORM10_PNL_PER_FILL",
        "REFERENCE_FIRST_ENTRY_PNL", "UNIFORM10_FIRST_ENTRY_PNL",
        "REFERENCE_REENTRY_PNL", "UNIFORM10_REENTRY_PNL",
        "REFERENCE_DAY_POSITIVE_RATE", "UNIFORM10_DAY_POSITIVE_RATE",
        "AM_RESULT", "PM_RESULT", "TAIL_RESULT",
        "HISTORICAL_10MIN_SUPPORT", "MISSED_OPPORTUNITY_EVIDENCE", "FUTURE_FORWARD_REQUIRED", "verdict",
    ]
    lines = [
        "# TRADEBOT — Fixed 10-minute anchor opportunity",
        "",
        "Offline historical diagnostic. CLOCK_GRID / ENTRY / EXIT / Runtime unchanged.",
        "Not a start-minute or interval search. 20260810+ is not NEW_GRID_TRUE_HOLDOUT.",
        "",
        f"DATA_STATUS: {report.get('DATA_STATUS')}",
        f"SOURCE_FILE: {report.get('SOURCE_FILE')}",
        f"SOURCE_SHA: {report.get('SOURCE_SHA')}",
        "",
        "## REQUIRED",
    ]
    for k in keys:
        lines.append(f"{k}: {report.get(k)}")
    de = report.get("decomp") or {}
    lines += [
        "",
        "## Phase A decomp",
        f"DIRECT_NEW_ENTRY_VALUE: {de.get('DIRECT_NEW_ENTRY_VALUE')}",
        f"DISPLACED_ENTRY_EFFECT: {de.get('DISPLACED_ENTRY_EFFECT')}",
        f"CAPACITY_EFFECT: {de.get('CAPACITY_EFFECT')}",
        f"SAME_SYMBOL_STATE_EFFECT: {de.get('SAME_SYMBOL_STATE_EFFECT')}",
        f"REENTRY_EFFECT: {de.get('REENTRY_EFFECT')}",
        f"residual: {de.get('residual')}",
        f"observed: {de.get('observed')}",
        "",
        "STOP. Do not change CLOCK_GRID. FUTURE_FORWARD_REQUIRED=true.",
        "submit/cancel/live=0/0/0",
    ]
    return "\n".join(lines) + "\n"


def write_xlsx(report: dict[str, Any], path: Path) -> None:
    def daily_rows() -> list[dict[str, Any]]:
        rows = []
        for sh, lst in (report.get("daily") or {}).items():
            for r in lst or []:
                rows.append({"variant": sh, **r})
        return rows or [{"status": "empty"}]

    def fr_rows() -> list[dict[str, Any]]:
        rows = []
        for sh, st in (report.get("first_reentry") or {}).items():
            rec = {"variant": sh}
            for k, v in (st or {}).items():
                if k == "all" and isinstance(v, dict):
                    for kk, vv in v.items():
                        rec[f"all_{kk}"] = vv
                else:
                    rec[k] = v
            rows.append(rec)
        return rows or [{"status": "empty"}]

    def ampm_rows() -> list[dict[str, Any]]:
        rows = []
        for var, body in (report.get("am_pm") or {}).items():
            for sess, st in (body or {}).items():
                rec = {"variant": var, "session": sess}
                rec.update(st or {})
                rows.append(rec)
        return rows or [{"status": "empty"}]

    sheets = {
        "Summary": [{
            "verdict": report.get("verdict"),
            "PHASE_A_VERDICT": report.get("PHASE_A_VERDICT"),
            "HISTORICAL_10MIN_SUPPORT": report.get("HISTORICAL_10MIN_SUPPORT"),
            "MISSED_OPPORTUNITY_EVIDENCE": report.get("MISSED_OPPORTUNITY_EVIDENCE"),
            "FUTURE_FORWARD_REQUIRED": report.get("FUTURE_FORWARD_REQUIRED"),
            "REFERENCE_PNL": report.get("REFERENCE_PNL"),
            "AUGMENTED_PNL": report.get("AUGMENTED_PNL"),
            "UNIFORM10_PNL": report.get("UNIFORM10_PNL"),
            "PHASE_A_PNL_DELTA": report.get("PHASE_A_PNL_DELTA"),
            "PHASE_A_RESIDUAL": report.get("PHASE_A_RESIDUAL"),
        }],
        "PhaseA_AddedAnchors": report.get("iso_added") or [{"status": "empty"}],
        "Incremental_Lineage": report.get("lineage") or [{"status": "empty"}],
        "Opportunity_Decomp": [report.get("decomp") or {"status": "empty"}],
        "Uniform10": [report.get("uni_stats") or {"status": "empty"}],
        "Daily": daily_rows(),
        "FirstEntry_Reentry": fr_rows(),
        "AM_PM": ampm_rows(),
        "Tail": (report.get("iso_tail") or []) or [{"status": "empty"}],
    }
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(str(name)[:31])
        _sheet(ws, rows if isinstance(rows, list) else [rows])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_artifacts(report: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    jp, mp, xp = OUT / "report.json", OUT / "report.md", OUT / "audit.xlsx"
    slim = dict(report)
    for k in ("lineage", "iso_added", "iso_tail", "ref_trades", "aug_trades", "uni_trades"):
        slim.pop(k, None)
    jp.write_text(json.dumps(json_sanitize(slim), ensure_ascii=False, indent=2), encoding="utf-8")
    mp.write_text(write_md(report), encoding="utf-8")
    write_xlsx(report, xp)
    return {"report_json": str(jp), "report_md": str(mp), "audit_xlsx": str(xp)}
