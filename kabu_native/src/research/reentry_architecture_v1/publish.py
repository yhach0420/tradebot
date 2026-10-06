"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.reentry_architecture_v1 import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "reentry_architecture_v1"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Episode_Ledger",
    "Sequence",
    "Prior_Outcome",
    "Prior_Exit",
    "Anchor_Distance",
    "Score_Rank_Change",
    "A_Variants",
    "B_Variants",
    "Executability_Interaction",
    "Daily",
    "Tail_Robustness",
    "Mechanism",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                obj = float(obj)
            elif isinstance(obj, np.integer):
                return int(obj)
    except Exception:
        pass
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    q = report.get("answers") or {}
    return "\n".join(
        [
            "# RE-ENTRY ARCHITECTURE CAUSAL AUDIT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE Exact Dual-Lane. Occupancy approximation was not used.",
            "Trade deletion counterfactuals were not used. R1/R2/R3 replay from event start.",
            "C2_OOF_RANKING_EDGE_NOT_STABLE maintained. B2 formal verdict unchanged.",
            "A2 standalone verdict unchanged. C14 / Runtime / CLOCK / ENTRY / EXIT / Fill SoT unchanged.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A0_PARITY: {g.get('A0_PARITY')}",
            f"B0_PARITY: {g.get('B0_PARITY')}",
            f"A_FIRST_PNL: {g.get('A_FIRST_PNL')}",
            f"A_REENTRY_PNL: {g.get('A_REENTRY_PNL')}",
            f"B_FIRST_PNL: {g.get('B_FIRST_PNL')}",
            f"B_REENTRY_PNL: {g.get('B_REENTRY_PNL')}",
            f"A_REENTRY1_PNL: {g.get('A_REENTRY1_PNL')}",
            f"A_REENTRY2_PNL: {g.get('A_REENTRY2_PNL')}",
            f"A_REENTRY3PLUS_PNL: {g.get('A_REENTRY3PLUS_PNL')}",
            f"AFTER_WIN_REENTRY_PNL: {g.get('AFTER_WIN_REENTRY_PNL')}",
            f"AFTER_LOSS_REENTRY_PNL: {g.get('AFTER_LOSS_REENTRY_PNL')}",
            f"PRIMARY_REENTRY_FAILURE_MODE: {g.get('PRIMARY_REENTRY_FAILURE_MODE')}",
            f"R1_A_PNL: {g.get('R1_A_PNL')}",
            f"R1_A_PF: {g.get('R1_A_PF')}",
            f"R1_A_DD: {g.get('R1_A_DD')}",
            f"R2_A_PNL: {g.get('R2_A_PNL')}",
            f"R2_A_PF: {g.get('R2_A_PF')}",
            f"R2_A_DD: {g.get('R2_A_DD')}",
            f"R3_A_PNL: {g.get('R3_A_PNL')}",
            f"R3_A_PF: {g.get('R3_A_PF')}",
            f"R3_A_DD: {g.get('R3_A_DD')}",
            f"R1_B_PNL: {g.get('R1_B_PNL')}",
            f"R1_B_PF: {g.get('R1_B_PF')}",
            f"R1_B_DD: {g.get('R1_B_DD')}",
            f"R2_B_PNL: {g.get('R2_B_PNL')}",
            f"R2_B_PF: {g.get('R2_B_PF')}",
            f"R2_B_DD: {g.get('R2_B_DD')}",
            f"R3_B_PNL: {g.get('R3_B_PNL')}",
            f"R3_B_PF: {g.get('R3_B_PF')}",
            f"R3_B_DD: {g.get('R3_B_DD')}",
            f"BEST_MECHANISM_SUPPORTED_VARIANT: {g.get('BEST_MECHANISM_SUPPORTED_VARIANT')}",
            f"CROSS_CLOCK_SUPPORTED: {g.get('CROSS_CLOCK_SUPPORTED')}",
            f"EXECUTABILITY_PLUS_REENTRY_PROMISING: {g.get('EXECUTABILITY_PLUS_REENTRY_PROMISING')}",
            f"TRUE_OOS: {g.get('TRUE_OOS')}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            f"RECOMMENDED_NEXT_STEP: {g.get('RECOMMENDED_NEXT_STEP')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## Q1–Q6",
            "",
            f"Q1: {q.get('Q1')}",
            f"Q2: {q.get('Q2')}",
            f"Q3: {q.get('Q3')}",
            f"Q4: {q.get('Q4')}",
            f"Q5: {q.get('Q5')}",
            f"Q6: {q.get('Q6')}",
            "",
            "## STOP",
            "",
            "Runtime not changed. C14 not changed. C2 not changed. C3 not started.",
            "B2 formal verdict not rewritten. A2 standalone verdict not rewritten.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
