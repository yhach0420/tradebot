"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.existing_data_strategy_research_stop_reassessment_v1 import ANALYSIS_ID
from research.existing_data_strategy_research_stop_reassessment_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "current_closure",
    "architecture_inventory",
    "method_taxonomy",
    "e1_model_audit",
    "joint_full_strategy_audit",
    "method_equivalence",
    "eligibility",
    "selection",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    elig = dict(report.get("eligibility") or {})
    freeze = dict(report.get("freeze") or {})
    methods = list(report.get("method_taxonomy") or [])
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "current_closure": _kv(
            {
                **dict(report.get("parent_fdg") or {}),
                **{f"object_{k}": v for k, v in dict(report.get("parent_object") or {}).items()},
            }
        ),
        "architecture_inventory": list(report.get("architecture_inventory") or []),
        "method_taxonomy": methods,
        "e1_model_audit": _kv(dict(report.get("e1_model_audit") or {})),
        "joint_full_strategy_audit": _kv(dict(report.get("joint_full_strategy_audit") or {})),
        "method_equivalence": _kv(
            {
                "H5_vs_Full_Causal": "not equivalent",
                "prewritten_library_vs_data_learned": "not automatically equivalent",
                "bounded_data_driven_equivalent_prior": None,
                "collapses_to_M2_if_bounded_on_existing_vocab": True,
            }
        ),
        "eligibility": [dict(elig.get("bounded") or {})] if elig.get("bounded") else _kv(elig),
        "selection": _kv(freeze),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            "",
            f"- FDG pinned: {a.get('1_current_FDG_verdict_pinned')} integrity={a.get('2_FDG_integrity_PASS')} canary={a.get('3_FDG_canary_PASS')}",
            f"- FDG PnL/PF/days: {a.get('4_FDG_total_PnL')} / {a.get('5_FDG_PF')} / {a.get('6_FDG_days')}",
            f"- FDG rescued: false",
            f"- FULL_DEPTH_GEOMETRY justified domain exhausted: {a.get('8_FULL_DEPTH_GEOMETRY_exhausted_within_justified_domain')}",
            f"- closed architecture N: {a.get('10_closed_architecture_family_N')}",
            f"- remaining eligible method N: {a.get('29_remaining_eligible_method_N')}",
            f"- SPEC: `{a.get('41_REMAINING_RESEARCH_METHOD_SPEC_SHA256')}`",
            "",
            "With current sealed DEV, current available information, and already-tested research methods, there is no justified next strategy search.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
