"""report.json / report.md / audit.xlsx. Day2 primary freeze untouched."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_reversal_confirmed_entry_day1_v1 import ANALYSIS_ID
from research.futures_reversal_confirmed_entry_day1_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "Events",
    "Episodes",
    "OriginalFive",
    "HorizonStats",
    "ParentAudit",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"CASE: {d.get('CASE')}",
        "",
        "Question: after BOTH_DOWN arms high-activity names, does waiting until",
        "BOTH_UP is actually observed (received_at ≤ ENTRY) leave TOP LONG mean",
        "and median above zero? Precursors P1–P4/C1/C2 are not inputs.",
        "Day2 primary freeze is unchanged. No ENTRY/EXIT. 10m is diagnostic.",
        "",
        "## Required answers",
        "",
        f"- 1 BOTH_DOWN episode N: `{a.get('1_BOTH_DOWN_episode_N')}`",
        f"- 2 armed N: `{a.get('2_armed_N')}`",
        f"- 3 expired N: `{a.get('3_expired_N')}`",
        f"- 4 confirmed BOTH_UP trigger N: `{a.get('4_confirmed_BOTH_UP_trigger_N')}`",
        f"- 5 trigger timestamps: `{a.get('5_trigger_timestamps')}`",
        f"- 6 arm→trigger latencies: `{a.get('6_arm_to_trigger_latencies_sec')}`",
        f"- 7 future leakage N: `{a.get('7_future_leakage_N')}`",
        f"- 8 TOP16 selection time causal: `{_fmt(a.get('8_TOP16_selection_time_causal'))}`",
        f"- 9 TOP LONG 1m mean/median: `{a.get('9_TOP_LONG_1m_mean_median')}`",
        f"- 10 TOP LONG 3m: `{a.get('10_TOP_LONG_3m')}`",
        f"- 11 TOP LONG 5m: `{a.get('11_TOP_LONG_5m')}`",
        f"- 12 TOP LONG 10m: `{a.get('12_TOP_LONG_10m')}`",
        f"- 13 BOTTOM LONG 10m: `{a.get('13_BOTTOM_LONG_10m')}`",
        f"- 14 TOP-BOTTOM LONG 10m: `{_fmt(a.get('14_TOP_BOTTOM_LONG_10m'))}`",
        f"- 15 ALL48 LONG 10m: `{a.get('15_ALL48_LONG_10m')}`",
        f"- 16 TOP incremental vs ALL48: `{_fmt(a.get('16_TOP_incremental_vs_ALL48'))}`",
        f"- 17 TOP positive event N: `{a.get('17_TOP_positive_event_N')}`",
        f"- 18 TOP negative event N: `{a.get('18_TOP_negative_event_N')}`",
        f"- 19 original 09:20 episode: `{a.get('19_original_09:20_episode')}`",
        f"- 20 original 10:20 episode: `{a.get('20_original_10:20_episode')}`",
        f"- 21 09:40: `{a.get('21_09:40')}`",
        f"- 22 10:30: `{a.get('22_10:30')}`",
        f"- 23 10:40: `{a.get('23_10:40')}`",
        f"- 24 edge survives confirmation delay: `{_fmt(a.get('24_edge_survives_confirmation_delay'))}`",
        f"- 25 rebound already priced: `{_fmt(a.get('25_rebound_already_priced'))}`",
        f"- 26 reversal confirmation sufficient: `{_fmt(a.get('26_reversal_confirmation_sufficient'))}`",
        f"- 27 exact causal mechanism: {a.get('27_exact_causal_mechanism_sentence')}",
        f"- 28 ENTRY thesis plausible: `{_fmt(a.get('28_ENTRY_thesis_plausible'))}`",
        f"- 29 ENTRY built: `{_fmt(a.get('29_ENTRY_built'))}`",
        f"- 30 EXIT built: `{_fmt(a.get('30_EXIT_built'))}`",
        f"- 31 existing Day2 primary changed: `{_fmt(a.get('31_existing_Day2_primary_changed'))}`",
        f"- 32 Runtime changed: `{_fmt(a.get('32_Runtime_changed'))}`",
        f"- 33 Paper changed: `{_fmt(a.get('33_Paper_changed'))}`",
        f"- 34 submit/cancel/live: `{a.get('34_submit_cancel_live')}`",
        f"- 35 CASE: `{a.get('35_CASE')}`",
        f"- 36 VERDICT: `{a.get('36_VERDICT')}`",
        f"- 37 NEXT: `{a.get('37_NEXT')}`",
        "",
        "No confirmation-threshold search. No TopK/selector/expiry/horizon retune.",
        "Precursor report.json was not overwritten; parent_consistency_audit records",
        "the POST_HOC_REVERSAL_DEPENDENT flag correction.",
        "",
    ]
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize(body)
    (OUT / "report.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(build_markdown(body), encoding="utf-8")
    orig = body.get("original_five") or {}
    orig_rows = list(orig.values()) if isinstance(orig, dict) else list(orig)
    hs = dict(body.get("horizon_stats") or {})
    sheets = {
        "Answers": _kv(dict(body.get("answers") or {})),
        "Events": list(body.get("events") or []),
        "Episodes": list(body.get("episodes") or []),
        "OriginalFive": orig_rows,
        "HorizonStats": [{"horizon": k, **(v if isinstance(v, dict) else {"value": v})} for k, v in hs.items()],
        "ParentAudit": _kv(dict(body.get("parent_consistency_audit") or {})),
        "Decision": _kv(dict(body.get("decision") or {})),
        "Safety": [
            {"key": "ENTRY_built", "value": False},
            {"key": "EXIT_built", "value": False},
            {"key": "Day2_primary_changed", "value": False},
            {"key": "Day2_substitutions_allowed", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "submit_cancel_live", "value": (body.get("answers") or {}).get("34_submit_cancel_live")},
            {"key": "future_leakage_n", "value": body.get("future_leakage_n")},
            {"key": "one_trigger_per_episode", "value": body.get("one_trigger_per_episode")},
            {"key": "confirmation_received_at_le_entry", "value": body.get("confirmation_received_at_le_entry")},
            {"key": "USE_PRECURSORS", "value": False},
        ],
    }
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(xlsx),
    }
