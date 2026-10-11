"""Add alphanumeric names to the research universe. Runtime code is not edited.

Numeric feature rows stay the stored CSV values. Alphanumeric rows join the
same source bar as that file. V2 and V3 scores are not retuned.
"""
from __future__ import annotations

import csv
import json
import math
import subprocess
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

from openpyxl import Workbook

from research.alphanumeric_universe_restoration_v1.normalize import (
    ResearchCodeRejected,
    research_normalize_code,
    research_yahoo_ticker,
)
from research.auto50_research_capture.mode import eligible_feature
from research.daytrade50_capital_efficient_selector_v2.features import selection_calendar, snapshot_days
from research.daytrade50_capital_efficient_selector_v2.lineage import (
    download_daily_bars,
    jpx_sessions,
    previous_jpx_session,
)
from research.daytrade50_capital_efficient_selector_v2.policy import one_lot_notional
from research.daytrade50_capital_efficient_selector_v2.selector import select_top50
from research.daytrade50_capital_efficient_selector_v2.selector_v3 import select_top50_v3
from research.daytrade50_capital_efficient_selector_v2.target_blind_g1m import (
    _aggregate,
    _book_metrics,
    _feature,
    _market_map,
    _series,
    _target_metrics,
)

NATIVE = Path(__file__).resolve().parents[3]
ROOT = NATIVE.parent
MASTER = ROOT / "data" / "jpx" / "tradable_symbols.csv"
REPORTS = NATIVE / "results" / "reports"
OUT = NATIVE / "results" / "research" / "auto50_alphanumeric_universe_restoration_v1"
AUTO50_DATES = NATIVE / "results" / "research" / "auto50_optional_fixed_slots_v1" / "report.json"
V2_REPORT = (
    NATIVE
    / "results"
    / "research"
    / "daytrade50_capital_efficient_selector_v2"
    / "target_blind_validation_cap500k_lev2x_g1m"
    / "report.json"
)
V3_REPORT = NATIVE / "results" / "research" / "daytrade50_selector_v3_long_horizon_correction" / "report.json"
G = 1_000_000.0
FOCUS = ("285A", "278A", "6857", "5803")
PAPER_MARKERS = (
    "paper_trade_checked_runner",
    "market_ingress_service",
    "v1r_paper_primary_launcher",
    "run_core10_dynamic40",
    "run_small_paper_pilot",
)


def paper_commands() -> list[str]:
    script = (
        "Get-CimInstance Win32_Process -Filter \"Name='python.exe'\" | "
        "ForEach-Object { $_.CommandLine }"
    )
    completed = subprocess.run(
        ["powershell", "-NoProfile", "-Command", script],
        capture_output=True,
        text=True,
        check=False,
    )
    lines = [line.strip() for line in (completed.stdout or "").splitlines() if line.strip()]
    return [line for line in lines if any(marker in line for marker in PAPER_MARKERS)]


def load_master() -> tuple[list[dict[str, str]], list[str]]:
    active: list[dict[str, str]] = []
    rejected: list[str] = []
    with MASTER.open(encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle):
            flag = str(raw.get("is_active") or "").strip().lower()
            if flag not in {"1", "true", "yes"}:
                continue
            try:
                code = research_normalize_code(raw.get("symbol") or "")
            except ResearchCodeRejected as exc:
                rejected.append(str(exc))
                continue
            active.append(
                {
                    "code": code,
                    "exchange": str(raw.get("exchange") or "1"),
                    "market": str(raw.get("market") or "").strip().lower(),
                    "name": str(raw.get("name") or ""),
                }
            )
    return active, rejected


def alpha_codes(active: Sequence[Mapping[str, str]]) -> list[dict[str, str]]:
    return [row for row in active if len(row["code"]) == 4 and row["code"][3].isalpha()]


def download_recorded(symbols: Sequence[str], start: date, end: date) -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, str]]]:
    """Download free daily bars and keep a row for every requested name."""
    import yfinance as yf

    panel: dict[str, list[dict[str, Any]]] = {}
    notes: list[dict[str, str]] = []

    def one(symbol: str) -> tuple[list[dict[str, Any]], str]:
        try:
            frame = yf.download(
                symbol,
                start=start.isoformat(),
                end=end.isoformat(),
                interval="1d",
                auto_adjust=False,
                progress=False,
                threads=False,
            )
        except Exception as exc:
            return [], f"FETCH_FAILED:{type(exc).__name__}"
        if frame is None or getattr(frame, "empty", True):
            return [], "NO_YAHOO_BARS"
        rows: list[dict[str, Any]] = []
        dropped_close = 0
        for idx, row in frame.iterrows():
            try:
                close = float(row["Close"])
                high = float(row["High"])
                low = float(row["Low"])
                volume = float(row["Volume"])
            except (TypeError, ValueError, KeyError):
                dropped_close += 1
                continue
            if close != close or close <= 0 or high != high or low != low or volume != volume:
                dropped_close += 1
                continue
            rows.append(
                {
                    "source_bar_date": idx.date().strftime("%Y%m%d"),
                    "close": close,
                    "high": high,
                    "low": low,
                    "volume": volume,
                }
            )
        rows.sort(key=lambda item: item["source_bar_date"])
        if not rows:
            return [], "NO_COMPLETE_BARS" if dropped_close else "NO_YAHOO_BARS"
        return rows, "FETCH_SUCCESS"

    batch_size = 40
    for offset in range(0, len(symbols), batch_size):
        batch = list(symbols[offset : offset + batch_size])
        print(f"alpha batch {offset // batch_size + 1}/{(len(symbols) + batch_size - 1) // batch_size}", flush=True)
        try:
            frame = yf.download(
                batch,
                start=start.isoformat(),
                end=end.isoformat(),
                interval="1d",
                group_by="ticker",
                auto_adjust=False,
                progress=False,
                threads=False,
            )
        except Exception as exc:
            frame = None
            print(f"batch failed {type(exc).__name__}; retrying singles", flush=True)
        if frame is None or getattr(frame, "empty", True):
            for symbol in batch:
                rows, status = one(symbol)
                panel[symbol.split(".")[0]] = rows
                notes.append({"symbol": symbol, "status": status})
            continue
        for symbol in batch:
            key = symbol.split(".")[0]
            try:
                view = frame[symbol]
                view = view.dropna(how="all")
            except Exception:
                rows, status = one(symbol)
                panel[key] = rows
                notes.append({"symbol": symbol, "status": status})
                continue
            rows = []
            dropped_close = 0
            for idx, row in view.iterrows():
                try:
                    close = float(row["Close"])
                    high = float(row["High"])
                    low = float(row["Low"])
                    volume = float(row["Volume"])
                except (TypeError, ValueError, KeyError):
                    dropped_close += 1
                    continue
                if close != close or close <= 0 or high != high or low != low or volume != volume:
                    dropped_close += 1
                    continue
                rows.append(
                    {
                        "source_bar_date": idx.date().strftime("%Y%m%d"),
                        "close": close,
                        "high": high,
                        "low": low,
                        "volume": volume,
                    }
                )
            rows.sort(key=lambda item: item["source_bar_date"])
            if rows:
                panel[key] = rows
                notes.append({"symbol": symbol, "status": "FETCH_SUCCESS", "dropped_incomplete_rows": str(dropped_close)})
            else:
                single, status = one(symbol)
                panel[key] = single
                notes.append({"symbol": symbol, "status": status, "dropped_incomplete_rows": str(dropped_close)})
    return panel, notes


def read_csv_rows(day: str) -> list[dict[str, str]]:
    path = REPORTS / f"features_{day}.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def witness_bar(rows: Sequence[Mapping[str, str]], symbol: str) -> Optional[tuple[float, float]]:
    for row in rows:
        code = str(row.get("symbol") or "").split(".")[0].split("@")[0]
        if code == symbol:
            try:
                return float(row["close"]), float(row["volume"])
            except (TypeError, ValueError):
                return None
    return None


def match_source_date(series_rows: Sequence[Mapping[str, Any]], close: float, volume: float) -> list[str]:
    found = []
    for row in series_rows:
        if abs(float(row["close"]) - close) <= 0.02 and abs(float(row["volume"]) - volume) <= 1:
            found.append(str(row["source_bar_date"]))
    return found


def bar_on(rows: Sequence[Mapping[str, Any]], day: str) -> Optional[Mapping[str, Any]]:
    for row in rows:
        if str(row["source_bar_date"]) == day:
            return row
    return None


def score_bar(row: Mapping[str, Any]) -> Optional[dict[str, float]]:
    close = float(row["close"])
    high = float(row["high"])
    low = float(row["low"])
    volume = float(row["volume"])
    if close <= 0 or high < low:
        return None
    atr = round((high - low) / close * 100.0, 6)
    trading = volume * close
    if trading <= 0:
        return None
    score = round(atr * math.log10(trading), 6)
    return {"close": close, "volume": volume, "trading_value": trading, "atr_pct": atr, "score": score}


def rank_eligible(rows: Sequence[Mapping[str, str]], *, gross: Optional[float]) -> list[dict[str, Any]]:
    ranked = []
    for row in rows:
        item = eligible_feature(row)
        if item is None:
            continue
        if gross is not None and one_lot_notional(float(item["prior_close"])) > gross:
            continue
        ranked.append(item)
    ranked.sort(key=lambda item: (-float(item["score"]), str(item["symbol"])))
    return ranked


def day_status(rows: Sequence[Mapping[str, Any]], source_day: str) -> str:
    if not rows:
        return "FETCH_FAILED"
    first = str(rows[0]["source_bar_date"])
    if first > source_day:
        return "NOT_YET_LISTED"
    if bar_on(rows, source_day) is None:
        return "MISSING_BAR"
    return "FEATURE_OK"


def run_auto50(
    days: Sequence[str],
    alpha_panel: Mapping[str, Sequence[Mapping[str, Any]]],
    numeric_panel: Mapping[str, Sequence[Mapping[str, Any]]],
    alpha_meta: Mapping[str, Mapping[str, str]],
    sessions: Sequence[str],
) -> dict[str, Any]:
    coverage = []
    focus_rows = []
    parity_days = []
    mismatches = []
    changed_name_days = 0
    alpha_selected_name_days = 0
    raw_top50_days = 0
    g_top50_days = 0
    for day in days:
        stored = read_csv_rows(day)
        witness = witness_bar(stored, "1301")
        hits = match_source_date(numeric_panel.get("1301") or [], *(witness or (0.0, 0.0))) if witness else []
        source = hits[0] if len(hits) == 1 else ""
        required = previous_jpx_session(sessions, day)
        lag = None
        if source and required:
            lag = sum(1 for session in sessions if source < session <= required)
        compared = 0
        close_match = 0
        score_match = 0
        value_diff = 0
        if source:
            for raw in stored:
                code = str(raw.get("symbol") or "").split(".")[0]
                if not code.isdigit():
                    continue
                panel_row = bar_on(numeric_panel.get(code) or [], source)
                if panel_row is None:
                    continue
                built = score_bar(panel_row)
                if built is None:
                    continue
                compared += 1
                try:
                    same_close = abs(float(raw["close"]) - built["close"]) <= 0.02
                    same_volume = abs(float(raw["volume"]) - built["volume"]) <= 1
                except (TypeError, ValueError):
                    same_close = False
                    same_volume = False
                if same_close and same_volume:
                    close_match += 1
                    try:
                        same_score = abs(float(raw["volatility_liquidity_score"]) - built["score"]) <= 0.000002
                    except (TypeError, ValueError):
                        same_score = False
                    if same_score:
                        score_match += 1
                    elif len(mismatches) < 30:
                        mismatches.append(
                            {
                                "day": day,
                                "symbol": code,
                                "kind": "SCORE",
                                "stored": raw.get("volatility_liquidity_score"),
                                "rebuilt": built["score"],
                            }
                        )
                else:
                    value_diff += 1
                    if len(mismatches) < 30:
                        mismatches.append(
                            {
                                "day": day,
                                "symbol": code,
                                "kind": "BAR",
                                "stored_close": raw.get("close"),
                                "rebuilt_close": built["close"],
                            }
                        )
        extra = []
        ok = 0
        for code, meta in alpha_meta.items():
            status = day_status(alpha_panel.get(code) or [], source) if source else "SOURCE_BAR_UNMATCHED"
            built = score_bar(bar_on(alpha_panel.get(code) or [], source)) if status == "FEATURE_OK" else None
            if built is not None:
                ok += 1
                extra.append(
                    {
                        "symbol": f"{code}.T",
                        "exchange": meta["exchange"],
                        "market": meta["market"],
                        "close": built["close"],
                        "volume": built["volume"],
                        "trading_value": built["trading_value"],
                        "atr_pct": built["atr_pct"],
                        "volatility_liquidity_score": built["score"],
                    }
                )
        old_raw = rank_eligible(stored, gross=None)
        new_raw = rank_eligible([*stored, *extra], gross=None)
        old_g = rank_eligible(stored, gross=G)
        new_g = rank_eligible([*stored, *extra], gross=G)
        old_set = {item["symbol"] for item in old_raw[:50]}
        new_set = {item["symbol"] for item in new_raw[:50]}
        entered = sorted(new_set - old_set)
        alpha_in = [code for code in entered if code in alpha_meta]
        changed_name_days += len(old_set.symmetric_difference(new_set))
        alpha_selected_name_days += sum(1 for item in new_raw[:50] if item["symbol"] in alpha_meta)
        raw_top50_days += int(len(new_raw) >= 50)
        g_top50_days += int(len(new_g) >= 50)
        coverage.append(
            {
                "day": day,
                "source_bar_date": source,
                "source_match_count": len(hits),
                "required_previous_session": required,
                "lag_sessions": lag,
                "future_used": bool(source and required and source > required),
                "alpha_feature_ok": ok,
                "alpha_universe": len(alpha_meta),
                "numeric_compared": compared,
                "numeric_close_volume_match": close_match,
                "numeric_score_match": score_match,
                "numeric_value_diff": value_diff,
                "raw_changed_names": len(old_set.symmetric_difference(new_set)) // 2,
                "raw_alpha_selected": alpha_in,
                "g1m_size": len(new_g[:50]),
            }
        )
        parity_days.append(close_match / compared if compared else 0.0)

        def position(ranked: Sequence[Mapping[str, Any]], code: str) -> Optional[int]:
            for index, item in enumerate(ranked, start=1):
                if item["symbol"] == code:
                    return index
            return None

        for code in FOCUS:
            status = day_status(alpha_panel.get(code) or [], source) if code in alpha_meta and source else ""
            if code not in alpha_meta:
                status = "NUMERIC_CONTROL"
                built = None
                stored_row = next((row for row in stored if str(row.get("symbol") or "").startswith(code)), None)
            else:
                built = score_bar(bar_on(alpha_panel.get(code) or [], source)) if status == "FEATURE_OK" else None
                stored_row = None
            close = None
            if built is not None:
                close = built["close"]
            elif stored_row is not None:
                try:
                    close = float(stored_row["close"])
                except (TypeError, ValueError):
                    close = None
            lot = one_lot_notional(close) if close is not None else None
            raw_rank = position(new_raw, code)
            g_rank = position(new_g, code)
            if code in alpha_meta and status != "FEATURE_OK":
                g_reason = status
            elif lot is not None and lot > G:
                g_reason = "ONE_LOT_ABOVE_G"
            elif g_rank is not None and g_rank <= 50:
                g_reason = "SELECTED"
            elif raw_rank is None:
                g_reason = "NOT_ELIGIBLE"
            else:
                g_reason = "RANK_OUTSIDE_TOP50"
            focus_rows.append(
                {
                    "book": "AUTO50",
                    "day": day,
                    "symbol": code,
                    "features_present": status == "FEATURE_OK" or stored_row is not None,
                    "fetch_status": status or ("IN_STORED_FEATURES" if stored_row else "ABSENT"),
                    "score": (built or {}).get("score") if built else (stored_row or {}).get("volatility_liquidity_score"),
                    "raw_rank": raw_rank,
                    "raw_selected": bool(raw_rank and raw_rank <= 50),
                    "prior_close_x100": lot,
                    "g_eligible": bool(lot is not None and lot <= G and (status in {"", "FEATURE_OK", "NUMERIC_CONTROL"})),
                    "g1m_selected": bool(g_rank and g_rank <= 50),
                    "g_exclusion_reason": g_reason,
                    "old_raw_rank": position(old_raw, code),
                }
            )
    rates = sorted(parity_days)
    median_rate = rates[len(rates) // 2] if rates else 0.0
    score_rates = []
    for row in coverage:
        matched = int(row["numeric_close_volume_match"])
        score_rates.append((int(row["numeric_score_match"]) / matched) if matched else 0.0)
    score_rates.sort()
    score_median = score_rates[len(score_rates) // 2] if score_rates else 0.0
    return {
        "days": coverage,
        "focus": focus_rows,
        "mismatches": mismatches,
        "parity_median_close_volume_match": median_rate,
        "parity_median_score_match": score_median,
        "parity_pass": median_rate >= 0.98 and score_median >= 0.999,
        "changed_symbol_slots": changed_name_days // 2,
        "alpha_selected_name_days": alpha_selected_name_days,
        "raw_top50_days": raw_top50_days,
        "g1m_top50_days": g_top50_days,
        "day_count": len(days),
    }


def aggregate_metric(rows: Sequence[Mapping[str, Any]], key: str) -> Optional[float]:
    values = sorted(float(row[key]) for row in rows if row.get(key) is not None)
    if not values:
        return None
    return values[len(values) // 2]


def run_retrospective(
    days: Sequence[str],
    panel: Mapping[str, Sequence[Mapping[str, Any]]],
    meta: Mapping[str, tuple[str, str]],
    sessions: Sequence[str],
    alpha_set: set[str],
    label: str,
) -> dict[str, Any]:
    series = {symbol: _series(rows) for symbol, rows in panel.items() if len(rows) >= 20}
    daily = []
    v2_base_ranges = []
    v2_new_ranges = []
    v3_base_ranges = []
    v3_new_ranges = []
    change_v2 = []
    change_v3 = []
    alpha_v2 = 0
    alpha_v3 = 0
    for day in days:
        required = previous_jpx_session(sessions, day)
        if required is None:
            continue
        base_rows = []
        all_rows = []
        targets = {}
        for symbol, item in series.items():
            exchange, market = meta.get(symbol, ("1", ""))
            built = _feature(symbol, item, required, exchange, market)
            target = _target_metrics(item, day)
            if target is not None:
                targets[symbol] = target
            if built is None:
                continue
            if built["SOURCE_BAR_DATE"] != required or built["WINDOW_END_SESSION"] >= day:
                continue
            all_rows.append(built)
            if symbol not in alpha_set:
                base_rows.append(built)
        base_v2 = select_top50(base_rows, gross_budget_yen=G)
        new_v2 = select_top50(all_rows, gross_budget_yen=G)
        base_v3 = select_top50_v3(base_rows, gross_budget_yen=G)
        new_v3 = select_top50_v3(all_rows, gross_budget_yen=G)

        def metrics(chosen: Mapping[str, Any]) -> dict[str, Any]:
            symbols = list(chosen.get("symbols") or [])
            notionals = {}
            suits = {}
            for row in chosen.get("rows") or []:
                notionals[str(row["symbol"])] = one_lot_notional(float(row["prior_close"]))
                if row.get("suitability_feature") is not None:
                    suits[str(row["symbol"])] = float(row["suitability_feature"])
            return _book_metrics(symbols, notionals, suits, targets, G)

        left = metrics(base_v2)
        right = metrics(new_v2)
        left3 = metrics(base_v3)
        right3 = metrics(new_v3)
        if left.get("range_median") is not None:
            v2_base_ranges.append(left)
            v2_new_ranges.append(right)
        if left3.get("range_median") is not None:
            v3_base_ranges.append(left3)
            v3_new_ranges.append(right3)
        base_symbols = set(base_v2.get("symbols") or [])
        new_symbols = set(new_v2.get("symbols") or [])
        base3 = set(base_v3.get("symbols") or [])
        new3 = set(new_v3.get("symbols") or [])
        v2_delta = len(base_symbols.symmetric_difference(new_symbols)) // 2
        v3_delta = len(base3.symmetric_difference(new3)) // 2
        change_v2.append(v2_delta)
        change_v3.append(v3_delta)
        alpha_v2 += sum(1 for symbol in new_symbols if symbol in alpha_set)
        alpha_v3 += sum(1 for symbol in new3 if symbol in alpha_set)
        daily.append(
            {
                "book": label,
                "day": day,
                "required_session": required,
                "v2_changed_names": v2_delta,
                "v3_changed_names": v3_delta,
                "v2_alpha_selected": sorted(symbol for symbol in new_symbols if symbol in alpha_set),
                "v3_alpha_selected": sorted(symbol for symbol in new3 if symbol in alpha_set),
                "v2_base_range": left.get("range_median"),
                "v2_corrected_range": right.get("range_median"),
                "v3_base_range": left3.get("range_median"),
                "v3_corrected_range": right3.get("range_median"),
                "v2_base_tv": left.get("trading_value_median"),
                "v2_corrected_tv": right.get("trading_value_median"),
                "v2_status": new_v2.get("status"),
                "v3_status": new_v3.get("status"),
            }
        )
    return {
        "label": label,
        "days": len(daily),
        "used_as_independent_holdout": False,
        "reassessment": "CORRECTED_UNIVERSE_RETROSPECTIVE_REASSESSMENT",
        "v2_changed_names_sum": sum(change_v2),
        "v3_changed_names_sum": sum(change_v3),
        "v2_alpha_selected_name_days": alpha_v2,
        "v3_alpha_selected_name_days": alpha_v3,
        "v2_base": _aggregate(v2_base_ranges),
        "v2_corrected": _aggregate(v2_new_ranges),
        "v3_base": _aggregate(v3_base_ranges),
        "v3_corrected": _aggregate(v3_new_ranges),
        "daily": daily,
    }


def write_outputs(report: Mapping[str, Any], fetch_rows: Sequence[Mapping[str, Any]], focus: Sequence[Mapping[str, Any]], auto_days: Sequence[Mapping[str, Any]], retro_days: Sequence[Mapping[str, Any]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [
        f"VERDICT: {report['verdict']}",
        f"NEXT: {report['next']}",
        "",
        "PAPER V16 は変更していません。Runtime の normalize_code も変更していません。",
        "CAPITAL: 500000 × 2.0 = 1000000",
        f"MASTER_ALPHANUMERIC_N: {report['MASTER_ALPHANUMERIC_N']}",
        f"FETCH_SUCCESS_N: {report['FETCH_SUCCESS_N']}",
        f"FETCH_FAILED_N: {report['FETCH_FAILED_N']}",
        f"NUMERIC_SYMBOL_PARITY: {report['NUMERIC_SYMBOL_PARITY']}",
        f"AUTO50_RAW changed name-days: {report['AUTO50_RAW']['changed_name_days']}",
        f"AUTO50_RAW alphanumeric selections: {report['AUTO50_RAW']['alphanumeric_selection_name_days']}",
        f"285A raw ranks: {report['AUTO50_RAW']['rank_285A']}",
        f"278A raw ranks: {report['AUTO50_RAW']['rank_278A']}",
        f"AUTO50_G1M top50 days: {report['AUTO50_G1M']['top50_success_days']}",
        f"285A G exclusion: {report['AUTO50_G1M']['exclusion_285A']}",
        f"278A G exclusion: {report['AUTO50_G1M']['exclusion_278A']}",
        "V2/V3 は CORRECTED_UNIVERSE_RETROSPECTIVE_REASSESSMENT です。独立 holdout には使っていません。",
        "既存の MIXED と V3 NOT_SUPPORTED は上書きしていません。",
        f"PAPER_STOPPED: {report['PAPER_STOPPED']}",
        f"RUNTIME_CHANGED: {report['RUNTIME_CHANGED']}",
        f"LIVE_ORDER_ENABLED: {report['LIVE_ORDER_ENABLED']}",
        "submit/cancel/live: 0/0/0",
    ]
    (OUT / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    book = Workbook()
    summary = book.active
    summary.title = "summary"
    for key, value in report.items():
        if isinstance(value, (dict, list)):
            summary.append([key, json.dumps(value, ensure_ascii=False)])
        else:
            summary.append([key, value])
    fetch = book.create_sheet("fetch")
    fetch.append(["symbol", "name", "market", "status", "first_bar", "last_bar", "bars"])
    for row in fetch_rows:
        fetch.append([row.get("symbol"), row.get("name"), row.get("market"), row.get("status"), row.get("first_bar"), row.get("last_bar"), row.get("bars")])
    auto = book.create_sheet("auto50")
    if auto_days:
        auto.append(list(auto_days[0].keys()))
        for row in auto_days:
            auto.append([json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value for value in row.values()])
    focus_sheet = book.create_sheet("focus")
    if focus:
        focus_sheet.append(list(focus[0].keys()))
        for row in focus:
            focus_sheet.append(list(row.values()))
    retro = book.create_sheet("v2_v3")
    if retro_days:
        retro.append(list(retro_days[0].keys()))
        for row in retro_days:
            retro.append([json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value for value in row.values()])
    book.save(OUT / "audit.xlsx")


def main() -> None:
    alive = paper_commands()
    if alive:
        report = {
            "verdict": "ALPHANUMERIC_UNIVERSE_RESTORATION_PENDING",
            "next": "RESUME_ISOLATED_REBUILD_AFTER_PAPER_NORMAL_CLOSE",
            "reason": "PAPER_PROCESS_STILL_ALIVE",
            "PAPER_STOPPED": False,
            "RUNTIME_CHANGED": False,
            "LIVE_ORDER_ENABLED": False,
            "submit_cancel_live": "0/0/0",
        }
        write_outputs(report, [], [], [], [])
        print("paper still alive", flush=True)
        return
    active, rejected = load_master()
    alpha = alpha_codes(active)
    alpha_meta = {row["code"]: row for row in alpha}
    print(f"alpha {len(alpha)} rejected {len(rejected)}", flush=True)
    panel, notes = download_recorded([research_yahoo_ticker(row["code"]) for row in alpha], date(2020, 1, 1), date(2026, 10, 9))
    note_by_symbol = {row["symbol"].split(".")[0]: row["status"] for row in notes}
    fetch_rows = []
    success = 0
    failed = 0
    for row in alpha:
        bars = panel.get(row["code"]) or []
        status = note_by_symbol.get(row["code"], "NO_YAHOO_BARS" if not bars else "FETCH_SUCCESS")
        if bars and status == "FETCH_SUCCESS":
            success += 1
        else:
            failed += 1
            status = status if status != "FETCH_SUCCESS" else "NO_YAHOO_BARS"
        fetch_rows.append(
            {
                "symbol": row["code"],
                "name": row["name"],
                "market": row["market"],
                "status": status if not bars else "FETCH_SUCCESS",
                "first_bar": bars[0]["source_bar_date"] if bars else "",
                "last_bar": bars[-1]["source_bar_date"] if bars else "",
                "bars": len(bars),
                "listing_censored": bool(bars and bars[0]["source_bar_date"] <= "20200106"),
            }
        )
    auto50_days = [row["date"] for row in json.loads(AUTO50_DATES.read_text(encoding="utf-8"))["daily"]]
    v3_days = json.loads(V3_REPORT.read_text(encoding="utf-8"))["holdout_provenance"]["holdout_dates"]
    numeric_symbols = sorted({code for day in snapshot_days(REPORTS) for code in []})
    numeric_symbols = []
    seen = set()
    for day in snapshot_days(REPORTS):
        for raw in read_csv_rows(day):
            code = str(raw.get("symbol") or "").split(".")[0]
            if code.isdigit() and code not in seen:
                seen.add(code)
                numeric_symbols.append(code)
    print(f"numeric download {len(numeric_symbols)}", flush=True)
    numeric_panel = download_daily_bars([f"{code}.T" for code in numeric_symbols], date(2025, 12, 1), date(2026, 10, 9))
    witness = {
        "1301": [(row["source_bar_date"], float(row["close"]), float(row["volume"])) for row in numeric_panel.get("1301") or []],
        "7203": [(row["source_bar_date"], float(row["close"]), float(row["volume"])) for row in numeric_panel.get("7203") or []],
    }
    sessions = jpx_sessions(witness)
    auto = run_auto50(auto50_days, panel, numeric_panel, alpha_meta, sessions)
    print(
        f"auto50 parity {auto['parity_pass']} match {auto['parity_median_close_volume_match']:.4f} score {auto['parity_median_score_match']:.4f}",
        flush=True,
    )
    meta = _market_map(REPORTS)
    for row in alpha:
        meta[row["code"]] = (row["exchange"], row["market"])
    merged = dict(numeric_panel)
    for code, rows in panel.items():
        merged[code] = list(rows)
    v2_saved = json.loads(V2_REPORT.read_text(encoding="utf-8"))
    session_set = set(sessions)
    unavailable = set(v2_saved.get("target_close_unavailable_dates") or [])
    calendar = [day for day in selection_calendar(REPORTS) if day in session_set and day <= "20261007"]
    metric_days = [day for day in calendar if day not in unavailable]
    print(f"v2 days {len(metric_days)} v3 days {len(v3_days)}", flush=True)
    v2 = run_retrospective(metric_days, merged, meta, sessions, set(alpha_meta), "V2_EVALUATION")
    v3 = run_retrospective(v3_days, merged, meta, sessions, set(alpha_meta), "V3_HOLDOUT")
    saved_v2_range = v2_saved["comparison"]["DAYTRADE50_V2"]["range_median"]
    saved_v3 = json.loads(V3_REPORT.read_text(encoding="utf-8"))
    saved_v3_range = saved_v3["books"]["V3"]["range_median"]
    parity = "PASS" if auto["parity_pass"] else "FAIL"
    restored = parity == "PASS" and failed >= 0 and success > 0 and not rejected
    report = {
        "verdict": "ALPHANUMERIC_UNIVERSE_RESTORED_V1" if restored else "ALPHANUMERIC_UNIVERSE_RESTORATION_INCOMPLETE_V1",
        "next": "JUDGE_WHETHER_CORRECTED_AUTO50_IS_THE_SELECTION_BASELINE" if restored else "REPAIR_DATA_BEFORE_SELECTOR_WORK",
        "MASTER_ALPHANUMERIC_N": len(alpha),
        "MASTER_FOUR_DIGIT_N": sum(1 for row in active if row["code"].isdigit()),
        "MASTER_REJECTED_N": len(rejected),
        "FETCH_SUCCESS_N": success,
        "FETCH_FAILED_N": failed,
        "FEATURE_COVERAGE_BY_DAY": auto["days"],
        "NUMERIC_SYMBOL_PARITY": parity,
        "parity_median_close_volume_match": auto["parity_median_close_volume_match"],
        "parity_median_score_match": auto["parity_median_score_match"],
        "parity_mismatches": auto["mismatches"],
        "AUTO50_RAW": {
            "changed_name_days": auto["changed_symbol_slots"],
            "alphanumeric_selection_name_days": auto["alpha_selected_name_days"],
            "top50_days": auto["raw_top50_days"],
            "day_count": auto["day_count"],
            "rank_285A": [row["raw_rank"] for row in auto["focus"] if row["symbol"] == "285A"],
            "rank_278A": [row["raw_rank"] for row in auto["focus"] if row["symbol"] == "278A"],
            "selected_285A": [row["day"] for row in auto["focus"] if row["symbol"] == "285A" and row["raw_selected"]],
            "selected_278A": [row["day"] for row in auto["focus"] if row["symbol"] == "278A" and row["raw_selected"]],
        },
        "AUTO50_G1M": {
            "top50_success_days": auto["g1m_top50_days"],
            "day_count": auto["day_count"],
            "g_yen": G,
            "exclusion_285A": [row["g_exclusion_reason"] for row in auto["focus"] if row["symbol"] == "285A"],
            "exclusion_278A": [row["g_exclusion_reason"] for row in auto["focus"] if row["symbol"] == "278A"],
            "selected_285A": [row["day"] for row in auto["focus"] if row["symbol"] == "285A" and row["g1m_selected"]],
            "selected_278A": [row["day"] for row in auto["focus"] if row["symbol"] == "278A" and row["g1m_selected"]],
        },
        "controls": {
            "6857_raw_rank": [row["raw_rank"] for row in auto["focus"] if row["symbol"] == "6857"],
            "5803_raw_rank": [row["raw_rank"] for row in auto["focus"] if row["symbol"] == "5803"],
            "6857_old_raw_rank": [row["old_raw_rank"] for row in auto["focus"] if row["symbol"] == "6857"],
            "5803_old_raw_rank": [row["old_raw_rank"] for row in auto["focus"] if row["symbol"] == "5803"],
        },
        "effects": {
            "alphanumeric_admission": "AUTO50_RAW membership versus the stored numeric Top50",
            "fetch_failures": failed,
            "g_eligibility": "AUTO50_G1M applies prior close x 100 <= 1000000 after the same score",
            "rank_changes": "control ranks are old versus new on the same stored numeric scores",
        },
        "V2_V3": {
            "recomputed": True,
            "used_as_independent_holdout": False,
            "label": "CORRECTED_UNIVERSE_RETROSPECTIVE_REASSESSMENT",
            "saved_v2_verdict": "DAYTRADE50_V2_TARGET_BLIND_SELECTOR_MIXED",
            "saved_v3_verdict": "DAYTRADE50_SELECTOR_V3_BLIND_NOT_SUPPORTED_V1",
            "saved_v2_range_median": saved_v2_range,
            "recomputed_v2_numeric_only_range_median": v2["v2_base"].get("range_median"),
            "recomputed_v2_with_alpha_range_median": v2["v2_corrected"].get("range_median"),
            "saved_v3_range_median": saved_v3_range,
            "recomputed_v3_numeric_only_range_median": v3["v3_base"].get("range_median"),
            "recomputed_v3_with_alpha_range_median": v3["v3_corrected"].get("range_median"),
            "saved_v2_on_v3_holdout_range_median": saved_v3["books"]["V2"]["range_median"],
            "recomputed_v2_on_holdout_numeric_only_range_median": v3["v2_base"].get("range_median"),
            "v2": {key: value for key, value in v2.items() if key != "daily"},
            "v3": {key: value for key, value in v3.items() if key != "daily"},
            "score_formulas_changed": False,
            "adopted": False,
        },
        "CAPITAL": {"BASE_CAPITAL_YEN": 500000, "MAX_LEVERAGE_MULTIPLIER": 2.0, "G": 1000000},
        "PAPER_V16": "UNCHANGED",
        "PAPER_STOPPED": False,
        "RUNTIME_CHANGED": False,
        "runtime_normalize_code_changed": False,
        "LIVE_ORDER_ENABLED": False,
        "submit_cancel_live": "0/0/0",
        "auto_adjust": False,
        "future_used": sum(1 for row in auto["days"] if row["future_used"]),
        "research_normalizer": "src/research/alphanumeric_universe_restoration_v1/normalize.py",
    }
    write_outputs(report, fetch_rows, auto["focus"], auto["days"], [*v2["daily"], *v3["daily"]])
    print(report["verdict"], report["NUMERIC_SYMBOL_PARITY"], success, failed, flush=True)


if __name__ == "__main__":
    main()
