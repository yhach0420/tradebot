"""Write report.json / report.md / audit.xlsx only. No mass CSV. No Runtime write."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.executable_target_v2_coverage_audit import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "executable_target_v2_coverage_audit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)

SHEET_ORDER = (
    "Summary",
    "Session_Inventory",
    "Fifteen00_Class",
    "Anchor_Plus600",
    "Waterfall",
    "First_Fail",
    "V2_Null_Reasons",
    "By_Anchor",
    "By_Day",
    "By_Symbol",
    "Endpoint_Lookup",
    "Missingness_SMD",
    "Score_Decile",
    "Cohort",
    "Decision",
    "B_Status",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def flatten_smd(smd_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in smd_rows:
        rec = {
            "feature": r.get("feature"),
            "smd_valid_minus_null": r.get("smd_valid_minus_null"),
            "flag": r.get("flag"),
        }
        for side in ("valid", "null"):
            q = r.get(side) or {}
            for k, v in q.items():
                rec[f"{side}_{k}"] = v
        out.append(rec)
    return out


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("gates") or {}
    lines = [
        "# EXECUTABLE TARGET V2 SESSION + COVERAGE INTEGRITY AUDIT",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "OFFLINE AUDIT ONLY. C14 / Runtime / CLOCK / ENTRY / EXIT unchanged.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "C REBUILD V2 was not started this run.",
        "",
        "## REQUIRED OUTPUT",
        "",
        f"SESSION_SOT_VALID: {g.get('SESSION_SOT_VALID')}",
        f"PM_MARKET_SESSION_END: {g.get('PM_MARKET_SESSION_END')}",
        f"PM_CONTINUOUS_END: {g.get('PM_CONTINUOUS_END')}",
        f"PM_CLOSING_AUCTION: {g.get('PM_CLOSING_AUCTION')}",
        f"STALE_SESSION_CONSTANT: {g.get('STALE_SESSION_CONSTANT')}",
        f"15_00_PLUS600_STATUS: {g.get('15_00_PLUS600_STATUS')}",
        f"15_10_PLUS600_STATUS: {g.get('15_10_PLUS600_STATUS')}",
        f"15_20_PLUS600_STATUS: {g.get('15_20_PLUS600_STATUS')}",
        f"ROWS_TOTAL: {g.get('ROWS_TOTAL')}",
        f"ROWS_EXECUTABLE_T0: {g.get('ROWS_EXECUTABLE_T0')}",
        f"PRIMARY_ROWS: {g.get('PRIMARY_ROWS')}",
        f"PRIMARY_COVERAGE_RATE: {g.get('PRIMARY_COVERAGE_RATE')}",
        f"TOP_NULL_REASON: {g.get('TOP_NULL_REASON')}",
        f"TOP_NULL_REASON_N: {g.get('TOP_NULL_REASON_N')}",
        f"ENDPOINT_LOOKUP_SEMANTIC: {g.get('ENDPOINT_LOOKUP_SEMANTIC')}",
        f"ENDPOINT_LOOKUP_DEFECT: {g.get('ENDPOINT_LOOKUP_DEFECT')}",
        f"TARGET_MISSINGNESS_SELECTION_BIAS: {g.get('TARGET_MISSINGNESS_SELECTION_BIAS')}",
        f"MEDIAN_TARGET_VALID_N_PER_COHORT: {g.get('MEDIAN_TARGET_VALID_N_PER_COHORT')}",
        f"TARGET_COHORT_COVERAGE_ADEQUATE: {g.get('TARGET_COHORT_COVERAGE_ADEQUATE')}",
        f"TARGET_V2_COVERAGE_ACCEPTABLE: {g.get('TARGET_V2_COVERAGE_ACCEPTABLE')}",
        f"VERDICT: {g.get('VERDICT')}",
        f"RECOMMENDED_NEXT_STEP: {g.get('RECOMMENDED_NEXT_STEP')}",
        "",
        "## TSE vs V2 session SoT",
        "",
        "Current TSE cash equity: AM 09:00-11:30, PM 12:30-15:30, Zaraba until 15:25,",
        "closing auction 15:25-15:30. V2 PRIMARY reuses e1_x22 PM_SESSION_CLOSE_HM=15:00",
        "as continuous_session_end. That 15:00 is Dual Lane SESSION_CLOSE / last CLOCK",
        "slot (strategy cutoff) and a pre-2024 TSE close leftover. It is not current",
        "market close and not Zaraba end. Using it as PRIMARY t+600 market phase is a",
        "stale-session defect. Dual Lane EXIT 15:00 is left frozen; it must not be",
        "copied into market-session semantics.",
        "",
        "HORIZON_SEC=600 is 10 minutes. 15:00+600=15:10 (TSE Zaraba, V2 excluded).",
        "15:10+600=15:20 (TSE Zaraba, V2 excluded). 15:20+600=15:30 (closing-auction",
        "/ market close endpoint; continuous-mid Primary correctly excluded under TSE).",
        "",
        "## Endpoint lookup (code SoT, unchanged this audit)",
        "",
        "last_executable_mid: np.searchsorted(t, t_at, side='right')-1 then walk back.",
        "Accept last valid executable continuous board with event_t <= t_at and",
        "(t_at - event_t) <= BOARD_FRESHNESS_SEC (5.0). Not exact timestamp, not first",
        ">= t_at, not nearest. No walk-back past 5s (that would be last-stale-board).",
        "No itayose / locked / special / OPENS_WITHIN_1S.",
        "",
        "## B status",
        "",
        "B0: historical reference only.",
        "B1: B1_SCORE_THRESHOLD_NO_ROBUST_IMPROVEMENT.",
        "Per-feature hard threshold research: not started.",
        "",
        "## STOP",
        "",
        "Audit complete. C REBUILD V2 not started. Runtime not changed.",
        "",
    ]
    return "\n".join(lines)
