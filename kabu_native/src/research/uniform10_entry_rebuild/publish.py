"""Write uniform10_entry_rebuild report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.uniform10_entry_rebuild import ANALYSIS_ID, NEW_FORWARD_N, PRIMARY_TARGET, UNIFORM10

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "uniform10_entry_rebuild"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _kv(ws, data: dict[str, Any]) -> None:
    ws.append(["key", "value"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for k, v in data.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        ws.append([k, v])
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 80


def write_artifacts(report: dict[str, Any], extra_sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(report), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    _kv(
        ws,
        {
            k: v
            for k, v in report.items()
            if k not in {"_markdown"} and not isinstance(v, (dict, list))
        },
    )
    for name, rows in extra_sheets.items():
        ws2 = wb.create_sheet(name[:31])
        _sheet(ws2, rows if isinstance(rows, list) else [rows])
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    abc = report.get("abc") or {}
    a, b, c = abc.get("A") or {}, abc.get("B") or {}, abc.get("C") or {}
    cand = report.get("NEW_RUNTIME_CANDIDATE") or {}
    base = report.get("CORRECTED_BASELINE") or {}
    lines = [
        "# UNIFORM10 ENTRY rebuild (corrected Passive Fill)",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        f"PRIMARY_TARGET: {PRIMARY_TARGET}",
        f"UNIFORM10_N: {len(UNIFORM10)}",
        "CLOCK interval/start not searched. Runtime CLOCK_GRID unchanged.",
        "SAFETY: submit/cancel/live=0/0/0. Runtime activation forbidden.",
        "",
        "## NEW_RUNTIME_CANDIDATE",
        "",
        f"- {cand.get('candidate_id')} sha=`{cand.get('candidate_sha')}`",
        "",
        "## CORRECTED_BASELINE (A / current irregular CLOCK)",
        "",
        f"- trades/PnL/PF/maxDD: {base.get('trades')} / {base.get('pnl')} / {base.get('PF')} / {base.get('maxDD')}",
        "",
        f"OLD_RESULTS_STATUS: {report.get('OLD_RESULTS_STATUS')}",
        f"20260827_CAPTURE: {report.get('20260827_CAPTURE')}",
        "",
        "## A / B / C",
        "",
        f"- A: PnL={a.get('pnl')} PF={a.get('PF')} maxDD={a.get('maxDD')} trades={a.get('trades')}",
        f"- B: PnL={b.get('pnl')} PF={b.get('PF')} maxDD={b.get('maxDD')} trades={b.get('trades')}",
        f"- C: PnL={c.get('pnl')} PF={c.get('PF')} maxDD={c.get('maxDD')} trades={c.get('trades')}",
        "",
        f"UP_MOVER_RANKING_SUPPORTED: {report.get('UP_MOVER_RANKING_SUPPORTED')}",
        f"HISTORICAL_ROBUSTNESS: {report.get('HISTORICAL_ROBUSTNESS')}",
        f"NEW_FORWARD_N: {NEW_FORWARD_N}",
        "",
        f"verdict: **{report.get('verdict')}**",
        "",
        "STOP. UNIFORM10 / rebuilt ENTRY not written to Runtime.",
        "",
    ]
    return "\n".join(lines)
