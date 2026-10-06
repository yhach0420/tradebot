"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.systematic_state_transition_library_precommit_v1 import ANALYSIS_ID
from research.systematic_state_transition_library_precommit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "inventory_correction",
    "pfq_closure",
    "or_closure",
    "state_registry",
    "raw_library",
    "duplicate_map",
    "closed_lineage",
    "final_library",
    "execution",
    "exit",
    "portfolio",
    "folds",
    "gates",
    "hashes",
    "decision",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    reg = list(d.get("state_registry") or [])
    final = list(d.get("final_library") or [])
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"CASE: {_fmt(d.get('CASE_NAME') or d.get('CASE'))}",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT'))}",
        "",
        "ECONOMICS_RUN: false",
        "",
        "PRIOR_84_GRID_EXECUTED: false",
        "",
        "PRIOR_84_GRID_REJECTED_AS_REDUNDANT: true",
        "",
        "EXISTING_ARCHITECTURE_EXHAUSTED: true",
        "",
        "THRESHOLD_TUNED_THIS_RUN: false",
        "",
        "## Answers 1-48",
        "",
    ]
    for k, v in a.items():
        if str(k).startswith("_"):
            continue
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("## State registry")
    lines.append("")
    lines.append("| STATE_ID | FAMILY | SOURCE_FUNCTION | LOOKBACK | TUNED |")
    lines.append("|---|---|---|---|---|")
    for r in reg:
        lines.append(
            f"| {r['STATE_ID']} | {r['FAMILY']} | {r['SOURCE_FUNCTION']} | {r['LOOKBACK']} | false |"
        )
    lines.append("")
    lines.append("## Final candidate library")
    lines.append("")
    lines.append("| CANDIDATE_ID | TEMPLATE | STATE_A | STATE_B | EXEC | EXIT | CAP |")
    lines.append("|---|---|---|---|---|---|---|")
    for c in final:
        lines.append(
            f"| {c['CANDIDATE_ID']} | {c['TEMPLATE']} | {c['STATE_A']} | {c.get('STATE_B') or ''} | "
            f"{c['EXECUTION']} | {c['EXIT']} | {c['CAP']} |"
        )
    lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))


__all__ = [
    "SHEET_ORDER",
    "build_markdown",
    "json_sanitize",
    "kv_rows",
    "write_artifacts",
]
