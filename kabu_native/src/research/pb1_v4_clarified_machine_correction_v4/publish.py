"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_correction_v4.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Parent_Hashes",
    "V4_Change_Summary",
    "PostDominant_Persistence",
    "7011_20241205",
    "AuctionEnd_State",
    "FailedProbe_Reaccepted",
    "7011_20250523",
    "3382_FullWalk_Control",
    "4063_Control",
    "V3_to_V4_Changed_Rows",
    "Protected_14_Checks",
    "Reached_vs_Live",
    "Thesis_Loss_Events",
    "Hidden1M",
    "Invariants",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "audit_rows", "chart_zones", "hidden_1m_checks"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    inv = dict(report.get("invariants") or {})
    hid = dict(report.get("hidden_1m_recompute") or {})
    hashes = dict(report.get("hashes") or {})
    return {
        "SPEC_CHANGED": False,
        "CORRECTION_V2_CHANGED": False,
        "CORRECTION_V3_CHANGED": False,
        "NEW_V4_IDENTITY_CREATED": True,
        "PROSPECTIVE_DATA_OPENED": False,
        "OLD_CONFIRMATION_OPENED": False,
        "FROZEN_VALIDATION_OPENED": False,
        "FUTURE_OUTCOME_USED": False,
        "PNL_USED": False,
        "MFE_MAE_USED": False,
        "GRID_SEARCH": False,
        "THRESHOLD_RETUNED": False,
        "NEW_NUMERIC_CUTOFF_ADDED": False,
        "THRESHOLD_OPTIMIZED": False,
        "submit/cancel/live": "0/0/0",
        "SPEC_SHA256": hashes.get("SPEC_SHA256") or report.get("EXPECTED_SPEC_SHA256"),
        "CORRECTION_V2_SHA256": "916ec521c6830d0b3aed6079ff40e1d043cc52c00bbd4c880108d39ced95536c",
        "CORRECTION_V3_SHA256": "299fb0caeccdf725914da65aaaf59e980315881b523331ae64936a1b37f4dcb0",
        "NEW_V4_SHA256": hashes.get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V4_SHA256")
        or report.get("PB1_V4_CLARIFIED_MACHINE_CORRECTION_V4_SHA256"),
        "SPEC_CHANGE_REQUIRED": False,
        "Hidden-1m mismatch_n": hid.get("mismatch_n"),
        "SAME_BAR_ENTRY": inv.get("SAME_BAR_ENTRY"),
        "invariant_violations": inv.get("violations"),
        "target_pass_n": (dec.get("target_checks") or {}).get("pass_n"),
        "target_fail_n": (dec.get("target_checks") or {}).get("fail_n"),
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "incomplete_reasons": dec.get("incomplete_reasons"),
        "FULL_WALK_SOURCE_OF_TRUTH": True,
        "SIMPLIFIED_REPLAY_USED_AS_ORACLE": False,
    }


def _pick(rows: list[dict[str, Any]], crit: dict[str, Any], sym: str, date: str) -> dict[str, Any]:
    hit = next((r for r in rows if str(r.get("symbol")) == sym and str(r.get("date")) == date), None)
    return dict(hit or crit.get(f"{sym}_{date}") or {"symbol": sym, "date": date})


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    rows = list(report.get("audit_rows") or [])
    hashes = dict(report.get("hashes") or {})
    targets = list((report.get("target_checks") or {}).get("checks") or [])
    crit = dict(report.get("critical_cases") or {})
    cmp = dict(report.get("changed_row_manifest") or {})
    loss_rows = [r for r in rows if r.get("machine_THESIS_LOST") or r.get("machine_THESIS_LOST_AT")]
    r7011m = _pick(rows, crit, "7011", "20241205")
    r7011 = _pick(rows, crit, "7011", "20250523")
    r3382 = _pick(rows, crit, "3382", "20241004")
    r4063 = _pick(rows, crit, "4063", "20251118")
    return {
        "Manifest": [{"key": k, "value": v} for k, v in (report.get("answers") or {}).items()],
        "Parent_Hashes": [{"key": k, "value": v} for k, v in hashes.items()],
        "V4_Change_Summary": [
            {"change": "A", "what": "POST_DOMINANT_PATH_STATE persists on TRUE/FAILED/NO_VALID. Representation only."},
            {"change": "B", "what": "FAILED_BREAK_REACCEPTED / RANGE_ACCEPTED added under AUCTION_ENDED_AS_RANGE beside V3 STALE."},
            {"delta_a_report_only_n": cmp.get("delta_a_report_only_n"), "delta_b_lifecycle_n": cmp.get("delta_b_lifecycle_n"), "failed_break_reaccepted_n": cmp.get("failed_break_reaccepted_n")},
        ],
        "PostDominant_Persistence": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "SEED": r.get("machine_SEED"),
                "opening": r.get("machine_opening_state"),
                "post_dominant_class": r.get("post_dominant_class") or r.get("POST_DOMINANT_PATH_STATE"),
                "ACTIVE_LIVE": r.get("machine_ACTIVE_LIVE"),
                "THESIS_LIVE": r.get("machine_THESIS_LIVE"),
            }
            for r in rows
        ],
        "7011_20241205": [r7011m],
        "AuctionEnd_State": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "THESIS_LOST_REASON": r.get("machine_THESIS_LOST_REASON"),
                "auction_end_family": r.get("auction_end_family"),
                "THESIS_LOST_AT": r.get("machine_THESIS_LOST_AT"),
                "THESIS_LIVE": r.get("machine_THESIS_LIVE"),
                "THESIS_REACHED": r.get("machine_THESIS_REACHED"),
            }
            for r in loss_rows
        ]
        or [{"empty": True}],
        "FailedProbe_Reaccepted": [
            r
            for r in rows
            if str(r.get("machine_THESIS_LOST_REASON") or "") == "FAILED_BREAK_REACCEPTED"
        ]
        or [{"empty": True, "note": "no FAILED_BREAK_REACCEPTED rows"}],
        "7011_20250523": [r7011],
        "3382_FullWalk_Control": [r3382],
        "4063_Control": [r4063],
        "V3_to_V4_Changed_Rows": list(cmp.get("rows") or []) or [{"empty": True}],
        "Protected_14_Checks": targets or [{"empty": True}],
        "Reached_vs_Live": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "ACTIVE_REACHED": r.get("machine_ACTIVE_REACHED"),
                "ACTIVE_LIVE": r.get("machine_ACTIVE_LIVE"),
                "THESIS_REACHED": r.get("machine_THESIS_REACHED"),
                "THESIS_LIVE": r.get("machine_THESIS_LIVE"),
                "thesis_id": r.get("thesis_id"),
                "opening_drive_id": r.get("opening_drive_id"),
            }
            for r in rows
        ],
        "Thesis_Loss_Events": loss_rows or [{"empty": True}],
        "Hidden1M": [
            {"key": k, "value": v} for k, v in (report.get("hidden_1m_recompute") or {}).items() if k != "mismatches"
        ]
        or [{"empty": True}],
        "Invariants": [{"key": k, "value": v} for k, v in (report.get("invariants") or {}).items()],
        "Safety": [{"key": k, "value": v} for k, v in (report.get("safety") or {}).items()],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = ["# PB1 V4 clarified machine correction V4", "", "DEVELOPMENT_PARITY_ONLY. Not face validation. First candidate frozen.", ""]
    for k, v in ans.items():
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
