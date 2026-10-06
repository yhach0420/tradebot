"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_information_expansion import ANALYSIS_ID
from research.am_entry_profit_improvement.publish import json_sanitize

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_entry_information_expansion"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "precommit",
    "arms",
    "expansion",
    "outer_folds",
    "daily_pnl",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# AM ENTRY INFORMATION EXPANSION V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CURRENT-preserving profitable-fill classification. 4 arms × 18-day OOF.",
            "JOINT_SCORE = P_WIN - P_LOSS. X14 bundle frozen. No subset search.",
            "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED FINAL OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"ARM_N: {_fmt(g.get('ARM_N'))}",
            "",
            f"OUTER_FOLD_N: {_fmt(g.get('OUTER_FOLD_N'))}",
            "",
            f"BASE_LOGIT_NET: {_fmt(g.get('BASE_LOGIT_NET'))}",
            "",
            f"BASE_RF_NET: {_fmt(g.get('BASE_RF_NET'))}",
            "",
            f"X14_LOGIT_NET: {_fmt(g.get('X14_LOGIT_NET'))}",
            "",
            f"X14_RF_NET: {_fmt(g.get('X14_RF_NET'))}",
            "",
            f"BASE_LOGIT_PF: {_fmt(g.get('BASE_LOGIT_PF'))}",
            "",
            f"BASE_RF_PF: {_fmt(g.get('BASE_RF_PF'))}",
            "",
            f"X14_LOGIT_PF: {_fmt(g.get('X14_LOGIT_PF'))}",
            "",
            f"X14_RF_PF: {_fmt(g.get('X14_RF_PF'))}",
            "",
            f"BEST_ARM: {_fmt(g.get('BEST_ARM'))}",
            "",
            f"BEST_ARM_NET: {_fmt(g.get('BEST_ARM_NET'))}",
            "",
            f"BEST_ARM_PF: {_fmt(g.get('BEST_ARM_PF'))}",
            "",
            f"BEST_ARM_MAX_DD: {_fmt(g.get('BEST_ARM_MAX_DD'))}",
            "",
            f"BEST_ARM_PAIRED_MEDIAN: {_fmt(g.get('BEST_ARM_PAIRED_MEDIAN'))}",
            "",
            f"BEST_ARM_EX_TOP3: {_fmt(g.get('BEST_ARM_EX_TOP3'))}",
            "",
            f"PASS_ARM_N: {_fmt(g.get('PASS_ARM_N'))}",
            "",
            f"CURRENT_PRESERVATION_PASS_ALL: {_tf(g.get('CURRENT_PRESERVATION_PASS_ALL'))}",
            "",
            f"X14_INFORMATION_INCREMENTAL: {_tf(g.get('X14_INFORMATION_INCREMENTAL'))}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            "",
            f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            f"NEXT: {g.get('NEXT')}",
            "",
            "## STOP",
            "",
            "4-arm profitable-fill OOF complete. No Runtime change.",
            "No utility regression. No P_FILL ranking. submit/cancel/live=0/0/0.",
            "",
        ]
    )
