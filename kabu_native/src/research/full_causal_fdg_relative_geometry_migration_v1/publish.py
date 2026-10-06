"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.full_causal_fdg_relative_geometry_migration_v1 import ANALYSIS_ID, STRATEGY_ID
from research.full_causal_fdg_relative_geometry_migration_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "object_boundary",
    "strategy_spec",
    "relative_geometry",
    "migration_events",
    "baseline_thesis",
    "timestamp_semantics",
    "structural_coverage",
    "integrity",
    "canary",
    "trades",
    "economics",
    "daily",
    "symbols",
    "causal_ex_top",
    "blocks",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    ev = dict(report.get("evaluated") or {})
    trades = list(report.get("trades") or [])
    daily = ev.get("daily") or {}
    blocks = dict(ev.get("blocks") or {})
    g = dict(ev.get("g_table") or {})
    spec = dict(report.get("spec") or {})
    struct = dict(report.get("structural") or {})
    parent = {
        **dict(report.get("parent_object") or {}),
        **{f"static_{k}": v for k, v in dict(report.get("parent_static") or {}).items()},
    }
    integ = report.get("integrity") or {}
    gates = list(integ.get("gates") or [])
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "parent": _kv(parent),
        "object_boundary": _kv(
            dict(report.get("object_boundary") or {"SELECTED_OBJECT_ID": "FULL_DEPTH_GEOMETRY", "STATIC_SPAN_FAMILY_CLOSED": True})
        ),
        "strategy_spec": _kv(spec),
        "relative_geometry": _kv(
            {
                "BID_REL_K": spec.get("BID_REL_K"),
                "ASK_REL_K": spec.get("ASK_REL_K"),
                "translation_invariant": True,
                "DEEP_RANKS_REQUIRED": True,
                "qty_in_alpha": False,
            }
        ),
        "migration_events": _kv(
            {
                "BID_GEOMETRY_CONTRACTS": spec.get("BID_GEOMETRY_CONTRACTS"),
                "ASK_GEOMETRY_EXPANDS": spec.get("ASK_GEOMETRY_EXPANDS"),
                "FAVORABLE_RELATIVE_GEOMETRY_MIGRATION": spec.get("FAVORABLE_RELATIVE_GEOMETRY_MIGRATION"),
                "ONSET": spec.get("ONSET"),
                "VALID_CONSECUTIVE_PAIR_N": struct.get("VALID_CONSECUTIVE_PAIR_N"),
                "UNKNOWN_PAIR_N": struct.get("UNKNOWN_PAIR_N"),
                "JOINT_FAVORABLE_MIGRATION_EVENT_N": struct.get("JOINT_FAVORABLE_MIGRATION_EVENT_N"),
                "FALSE_TO_TRUE_SIGNAL_N": struct.get("FALSE_TO_TRUE_SIGNAL_N"),
            }
        ),
        "baseline_thesis": _kv(
            {
                "BASELINE": spec.get("BASELINE"),
                "THESIS_VALID": spec.get("THESIS_VALID"),
                "TECHNICAL_EXIT": spec.get("TECHNICAL_EXIT"),
                "TIMEOUT": False,
                "ALTERNATE_EXIT": False,
            }
        ),
        "timestamp_semantics": _kv(
            {
                "event_clock": "causal ingress received_at",
                "levels_2_10": "inherit carrying quote-update",
                "freshness": "AskTime / BidTime then ingress",
                "CurrentPriceTime": "forbidden",
                "PREOPEN_EXECUTION_VALID": False,
            }
        ),
        "structural_coverage": _kv(struct),
        "integrity": gates if gates else _kv(dict(integ)),
        "canary": _kv(dict(report.get("canary") or {})),
        "trades": trades[:5000] if trades else [{"empty": True}],
        "economics": _kv({k: ev.get(k) for k in ("TOTAL_PNL", "PF", "MaxDD", "trade_n", "signal_n") if ev}),
        "daily": [{"date": k, "pnl": v} for k, v in daily.items()] if daily else [{"empty": True}],
        "symbols": [{"top_symbol": ev.get("top_symbol"), "top_symbol_pnl": ev.get("top_symbol_pnl")}],
        "causal_ex_top": _kv({"CAUSAL_EX_TOP1_PNL": ev.get("CAUSAL_EX_TOP1_PNL"), "G6": g.get("G6")}),
        "blocks": list(blocks.get("blocks") or [{"S1": blocks.get("S1"), "S2": blocks.get("S2")}]),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"STRATEGY_ID: {STRATEGY_ID}",
            "",
            f"- SPEC: `{a.get('27_FULL_STRATEGY_SPEC_SHA256_FDG_MIGRATION_V1')}`",
            f"- integrity: {a.get('41_integrity_PASS_N_total')}",
            f"- canary PASS: {a.get('47_canary_parity_PASS')}",
            f"- FALSE→TRUE / days / symbols: {a.get('35_FALSE_TO_TRUE_SIGNAL_N')} / {a.get('36_SIGNAL_DAY_N')} / {a.get('37_SIGNAL_SYMBOL_N')}",
            f"- trade_n / TOTAL_PNL / PF: {a.get('50_trade_n')} / {a.get('64_TOTAL_PNL')} / {a.get('65_PF')}",
            f"- EXHAUSTED: {a.get('81_FULL_DEPTH_GEOMETRY_JUSTIFIED_MECHANISM_SPACE_EXHAUSTED')}",
            "",
            "Within sealed DEV Capture and this precommitted non-threshold, non-quantity, full-10-level FULL_DEPTH_GEOMETRY domain, CASE B/C/D/E exhausts the justified mechanism space. Do not invent another span, ratio, k-level, weighting, migration direction, persistence, window, or EXIT.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
