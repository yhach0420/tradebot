"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize
from research.am_exit_research_final_decision import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_exit_research_final_decision"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "required",
    "closed_lines",
    "evidence",
    "future_oos",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(52, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    lines = [
        "# AM EXIT RESEARCH FINAL DECISION V1",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "Development EXIT close. Original C0 + C14 retained.",
        "Mechanism weakness is recorded. No EXIT replacement is supported.",
        "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
        "",
        "## REQUIRED FINAL OUTPUT",
        "",
        f"AM_EXIT_MECHANISM_WEAKNESS_FOUND: {_tf(g.get('AM_EXIT_MECHANISM_WEAKNESS_FOUND'))}",
        "",
        f"AM_EXIT_FULL_GATE_POLICY_FOUND: {_tf(g.get('AM_EXIT_FULL_GATE_POLICY_FOUND'))}",
        "",
        f"AM_EXIT_SUPPORTED_REPLACEMENT_FOUND: {_tf(g.get('AM_EXIT_SUPPORTED_REPLACEMENT_FOUND'))}",
        "",
        f"PRIMARY_EXIT_WEAKNESS: {_fmt(g.get('PRIMARY_EXIT_WEAKNESS'))}",
        "",
        f"C0_L2_GROSS_LOSS_SHARE: {_fmt(g.get('C0_L2_GROSS_LOSS_SHARE'))}",
        "",
        f"PTL_ARCHITECTURE_SUPPORTED: {_tf(g.get('PTL_ARCHITECTURE_SUPPORTED'))}",
        "",
        f"CONTINUATION_ARCHITECTURE_SUPPORTED: {_tf(g.get('CONTINUATION_ARCHITECTURE_SUPPORTED'))}",
        "",
        f"BEST_TESTED_EXIT_FOR_C0: {_fmt(g.get('BEST_TESTED_EXIT_FOR_C0'))}",
        "",
        f"ORIGINAL_C0_C14_RETAINED: {_tf(g.get('ORIGINAL_C0_C14_RETAINED'))}",
        "",
        f"C0_PROSPECTIVE_SPEC_SHA256: {_fmt(g.get('C0_PROSPECTIVE_SPEC_SHA256'))}",
        "",
        f"AM_EXISTING_18D_EXIT_SEARCH_CLOSED: {_tf(g.get('AM_EXISTING_18D_EXIT_SEARCH_CLOSED'))}",
        "",
        f"RUNTIME_EXIT_CHANGE_ALLOWED: {_tf(g.get('RUNTIME_EXIT_CHANGE_ALLOWED'))}",
        "",
        f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
        "",
        f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
        "",
        f"VERDICT: {g.get('VERDICT')}",
        "",
        f"NEXT: {g.get('NEXT')}",
        "",
        "## CLOSED RESEARCH LINES",
        "",
    ]
    for line in list(report.get("closed_lines") or []):
        lines.append(str(line))
        lines.append("")
    lines.extend(
        [
            "## INTERPRETATION",
            "",
            "C14 is not certified perfect. Tested replacements did not beat C14 on this 18-day set.",
            "C0 future OOS keeps original C14. Do not mix PTL or forced-750 into that OOS.",
            "",
            "## STOP",
            "",
            "Existing 18-day AM EXIT search is closed.",
            "Original C0 remains FROZEN_FOR_FUTURE_OOS_ONLY with C14.",
            "submit/cancel/live=0/0/0.",
            "",
        ]
    )
    return "\n".join(lines)
