"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.reentry_clock_interaction_v2 import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "reentry_clock_interaction_v2"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Parity",
    "Clock_Class",
    "RE2_Cross",
    "Overlap",
    "Lineage",
    "Ordinal_Shift",
    "Elapsed",
    "Common_Anchor",
    "Setup_Reuse",
    "RE2_Ledger",
    "Mechanism",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                obj = float(obj)
            elif isinstance(obj, np.integer):
                return int(obj)
    except Exception:
        pass
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    q = report.get("answers") or {}
    return "\n".join(
        [
            "# RE-ENTRY CLOCK INTERACTION CAUSAL AUDIT V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE Exact Dual-Lane caches. No new policy. Occupancy approximation was not used.",
            "Nearest-anchor matching was not used. R2 is not a Runtime candidate.",
            "V1 REENTRY_POLICY_CLOCK_DEPENDENT maintained. C2/B2/A2+R2 frozen.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A_REENTRY2_N: {g.get('A_REENTRY2_N')}",
            f"A_REENTRY2_PNL: {g.get('A_REENTRY2_PNL')}",
            f"RE2_AFTER_WIN_N: {g.get('RE2_AFTER_WIN_N')}",
            f"RE2_AFTER_WIN_PNL: {g.get('RE2_AFTER_WIN_PNL')}",
            f"RE2_SCORE_NOT_IMPROVED_N: {g.get('RE2_SCORE_NOT_IMPROVED_N')}",
            f"RE2_SCORE_NOT_IMPROVED_PNL: {g.get('RE2_SCORE_NOT_IMPROVED_PNL')}",
            f"RE2_PRICE_ABOVE_N: {g.get('RE2_PRICE_ABOVE_N')}",
            f"RE2_PRICE_ABOVE_PNL: {g.get('RE2_PRICE_ABOVE_PNL')}",
            f"RE2_TRIPLE_OVERLAP_N: {g.get('RE2_TRIPLE_OVERLAP_N')}",
            f"RE2_TRIPLE_OVERLAP_PNL: {g.get('RE2_TRIPLE_OVERLAP_PNL')}",
            f"RE2_UNEXPLAINED_REMAINDER_N: {g.get('RE2_UNEXPLAINED_REMAINDER_N')}",
            f"RE2_UNEXPLAINED_REMAINDER_PNL: {g.get('RE2_UNEXPLAINED_REMAINDER_PNL')}",
            f"COMMON_ANCHOR_N: {g.get('COMMON_ANCHOR_N')}",
            f"A_ONLY_ANCHOR_N: {g.get('A_ONLY_ANCHOR_N')}",
            f"B_ONLY_ANCHOR_N: {g.get('B_ONLY_ANCHOR_N')}",
            f"ORDINAL_SHIFT_EPISODES_N: {g.get('ORDINAL_SHIFT_EPISODES_N')}",
            f"CLOCK_INSERTED_FILL_N: {g.get('CLOCK_INSERTED_FILL_N')}",
            f"CLOCK_REMOVED_FILL_N: {g.get('CLOCK_REMOVED_FILL_N')}",
            f"OCCUPANCY_CASCADE_EPISODES_N: {g.get('OCCUPANCY_CASCADE_EPISODES_N')}",
            f"A_RE2_BECOMES_B_RE1_N: {g.get('A_RE2_BECOMES_B_RE1_N')}",
            f"A_RE2_BECOMES_B_RE2_N: {g.get('A_RE2_BECOMES_B_RE2_N')}",
            f"A_RE2_BECOMES_B_RE3PLUS_N: {g.get('A_RE2_BECOMES_B_RE3PLUS_N')}",
            f"A_RE2_NO_DIRECT_MATCH_N: {g.get('A_RE2_NO_DIRECT_MATCH_N')}",
            f"A_ELAPSED_TIME: {g.get('A_ELAPSED_TIME')}",
            f"B_ELAPSED_TIME: {g.get('B_ELAPSED_TIME')}",
            f"CLOCK_EXPLAINS_ORDINAL_REVERSAL: {g.get('CLOCK_EXPLAINS_ORDINAL_REVERSAL')}",
            f"ORDINAL_EFFECT_INDEPENDENT: {g.get('ORDINAL_EFFECT_INDEPENDENT')}",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"RECOMMENDED_NEXT_RESEARCH: {g.get('RECOMMENDED_NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## Q1–Q7",
            "",
            f"Q1: {q.get('Q1')}",
            f"Q2: {q.get('Q2')}",
            f"Q3: {q.get('Q3')}",
            f"Q4: {q.get('Q4')}",
            f"Q5: {q.get('Q5')}",
            f"Q6: {q.get('Q6')}",
            f"Q7: {q.get('Q7')}",
            "",
            "## STOP",
            "",
            "Runtime not changed. R2 not adopted. C3 not started. No new policy.",
            "C14 / CLOCK / ENTRY / EXIT / Fill unchanged. Paper/OPVAL not operated.",
            "submit/cancel/live=0/0/0.",
            "",
        ]
    )
