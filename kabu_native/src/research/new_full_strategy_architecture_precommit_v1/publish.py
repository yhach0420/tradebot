"""Write report.json / report.md / audit.xlsx only. No mass CSV."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.new_full_strategy_architecture_precommit_v1 import ANALYSIS_ID
from research.new_full_strategy_architecture_precommit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "old_rca_stop",
    "closed_lineages",
    "design_constraints",
    "architecture_proposals",
    "novelty_audit",
    "eligibility",
    "selection",
    "frozen_strategy",
    "data_access",
    "decision",
    "safety",
)

PNL_FORBIDDEN_IN_DUMP = (
    "TOTAL_PNL",
    "EX_BEST",
    "CAUSAL_EX_TOP1",
    "FOLD_SELECTED_TEST_TOTAL_PNL",
    "top_symbol_pnl",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def proposal_sheet_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for p in rows:
        out.append(
            {
                "ARCHITECTURE_ID": p.get("ARCHITECTURE_ID"),
                "CORE_MECHANISM": p.get("CORE_MECHANISM"),
                "CLOSED_LINEAGE_MATCH": p.get("CLOSED_LINEAGE_MATCH"),
                "DISCRETIONARY_DOF": p.get("DISCRETIONARY_DOF"),
                "SIGNAL_PRIMITIVE_N": p.get("SIGNAL_PRIMITIVE_N"),
                "STATE_MACHINE_SIZE": p.get("STATE_MACHINE_SIZE"),
                "COVERAGE_PLAUSIBILITY_RANK": p.get("COVERAGE_PLAUSIBILITY_RANK"),
                "IMPLEMENTATION_AMBIGUITY_RANK": p.get("IMPLEMENTATION_AMBIGUITY_RANK"),
                "VWAP_ENTRY_USED": p.get("VWAP_ENTRY_USED"),
                "VWAP_EXIT_USED": p.get("VWAP_EXIT_USED"),
                "BOARD_PRIMARY_ALPHA": p.get("BOARD_PRIMARY_ALPHA"),
                "ENTRY_SIGNAL": (p.get("ENTRY_STATE_MACHINE") or {}).get("SIGNAL"),
                "EXIT_ID": (p.get("EXIT_STATE_MACHINE") or {}).get("EXIT_ID"),
                "JOINT_WITH_ENTRY": (p.get("EXIT_STATE_MACHINE") or {}).get("JOINT_WITH_ENTRY"),
                "CAP": (p.get("PORTFOLIO") or {}).get("CAP"),
                "same_symbol": (p.get("PORTFOLIO") or {}).get("same_symbol"),
                "REENTRY": (p.get("PORTFOLIO") or {}).get("REENTRY"),
                "STRUCTURAL_NOVELTY_REASON": p.get("STRUCTURAL_NOVELTY_REASON"),
                "EXPECTED_COVERAGE_REASON": p.get("EXPECTED_COVERAGE_REASON"),
            }
        )
    return out


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT') or a.get('49_VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT') or a.get('50_NEXT'))}",
        "",
        f"SELECTED_ARCHITECTURE_ID: {_fmt(d.get('SELECTED_ARCHITECTURE_ID') or a.get('18_selected_architecture_ID'))}",
        "",
        f"FULL_STRATEGY_SPEC_SHA256: {_fmt(d.get('FULL_STRATEGY_SPEC_SHA256') or a.get('31_FULL_STRATEGY_SPEC_SHA256'))}",
        "",
        "NEW_ARCHITECTURE_ECONOMICS_RUN: false",
        "",
        "NEW_REPLAY: false",
        "",
        "## Answers 1-50",
        "",
    ]
    for k, v in a.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    dumped = json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n"
    compact = dumped.replace(" ", "")
    for key in PNL_FORBIDDEN_IN_DUMP:
        if f'"{key}"' in dumped or f'"{key}":' in compact:
            raise RuntimeError(f"PNL_KEY_IN_REPORT {key}")
    (OUT / "report.json").write_text(dumped, encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))
