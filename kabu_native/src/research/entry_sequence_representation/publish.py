"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.entry_sequence_representation import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "entry_sequence_representation"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "S0",
    "S1",
    "Gates",
    "Daily",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# CAUSAL ENTRY SEQUENCE REPRESENTATION PROBE V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine. Common 600s. Exact executable-at-decision.",
            "Direct joint label frozen. RandomForestClassifier frozen. Top3 frozen.",
            "S0 = A0 / F2_UNION|none. S1 = A0 + 37x7 causal PRICE+BOARD sequence.",
            "Window t0-180s, 5-second grid, 37 marks. Flattened into the same RF.",
            "No volume / XS / engineered rebound / technical channels.",
            "No Exact. No PnL. No temporal deep model.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"S0_MFE_DELTA: {_fmt(g.get('S0_MFE_DELTA'))}",
            f"S0_DOWNSIDE_DELTA: {_fmt(g.get('S0_DOWNSIDE_DELTA'))}",
            f"S0_JOINT_RATE: {_fmt(g.get('S0_JOINT_RATE'))}",
            "",
            f"SEQUENCE_ROWS: {g.get('SEQUENCE_ROWS')}",
            f"SEQUENCE_FEATURE_N: {g.get('SEQUENCE_FEATURE_N')}",
            "",
            f"S1_MFE_DELTA: {_fmt(g.get('S1_MFE_DELTA'))}",
            f"S1_DOWNSIDE_DELTA: {_fmt(g.get('S1_DOWNSIDE_DELTA'))}",
            f"S1_JOINT_RATE: {_fmt(g.get('S1_JOINT_RATE'))}",
            "",
            f"DELTA_MFE_VS_S0: {_fmt(g.get('DELTA_MFE_VS_S0'))}",
            f"DELTA_DOWNSIDE_VS_S0: {_fmt(g.get('DELTA_DOWNSIDE_VS_S0'))}",
            f"DELTA_JOINT_RATE_VS_S0: {_fmt(g.get('DELTA_JOINT_RATE_VS_S0'))}",
            "",
            f"S1_MFE_POS_DAYS: {g.get('S1_MFE_POS_DAYS')}",
            f"S1_MFE_NEG_DAYS: {g.get('S1_MFE_NEG_DAYS')}",
            f"S1_DOWNSIDE_POS_DAYS: {g.get('S1_DOWNSIDE_POS_DAYS')}",
            f"S1_DOWNSIDE_NEG_DAYS: {g.get('S1_DOWNSIDE_NEG_DAYS')}",
            "",
            f"S1_MFE_EX_BEST_DAY: {_fmt(g.get('S1_MFE_EX_BEST_DAY'))}",
            f"S1_MFE_EX_TOP3_DAYS: {_fmt(g.get('S1_MFE_EX_TOP3_DAYS'))}",
            f"S1_DOWNSIDE_EX_BEST_DAY: {_fmt(g.get('S1_DOWNSIDE_EX_BEST_DAY'))}",
            f"S1_DOWNSIDE_EX_TOP3_DAYS: {_fmt(g.get('S1_DOWNSIDE_EX_TOP3_DAYS'))}",
            "",
            f"S1_ROC_AUC: {_fmt(g.get('S1_ROC_AUC'))}",
            f"S1_AVERAGE_PRECISION: {_fmt(g.get('S1_AVERAGE_PRECISION'))}",
            "",
            f"FUTURE_EVENT_USE_N: {g.get('FUTURE_EVENT_USE_N')}",
            f"SESSION_CARRY_N: {g.get('SESSION_CARRY_N')}",
            f"ITAYOSE_STATE_USE_N: {g.get('ITAYOSE_STATE_USE_N')}",
            f"SPECIAL_STATE_USE_N: {g.get('SPECIAL_STATE_USE_N')}",
            f"TARGET_CONTAMINATION_N: {g.get('TARGET_CONTAMINATION_N')}",
            "",
            f"SEQUENCE_REPRESENTATION_PASS: {_tf(g.get('SEQUENCE_REPRESENTATION_PASS'))}",
            "",
            f"PRIMARY_FINDING: {g.get('PRIMARY_FINDING')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Exact. No runtime candidate. No temporal deep model.",
            "Grid/window/channels frozen. No lag selection. No importance rerun.",
            "Runtime/C14 unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
