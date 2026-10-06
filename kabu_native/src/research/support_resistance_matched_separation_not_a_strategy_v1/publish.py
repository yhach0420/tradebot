"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.support_resistance_matched_separation_not_a_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Freeze",
    "A_Matchability",
    "C_Matchability",
    "Matched_Unmatched",
    "Propensity_Overlap",
    "OverlapWeighted_A",
    "OverlapWeighted_C",
    "Placebo_Calibration",
    "TwoWay_Inference",
    "Multiple_Testing",
    "A1_Passive",
    "A2_Confirmation",
    "A2_Executable_Path",
    "C1_Executable_Path",
    "Bar_Time_Lineage",
    "Resistance_Support",
    "D2_D3_D4",
    "Exit_Architecture",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, (list, tuple)):
        return [json_sanitize(x) for x in got]
    return got


def _kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for k, v in d.items():
        if isinstance(v, (dict, list, tuple)):
            v = json.dumps(json_sanitize(v), ensure_ascii=False)[:32000]
        rows.append({"key": str(k), "value": v})
    return rows


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["key", "value"])
        ws.append(["empty", True])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    bind = dict(report.get("bind") or {})
    freeze = dict(report.get("freeze") or {})
    d = dict(report.get("decision") or {})
    inf = dict(report.get("inference") or {})
    return {
        "Binding": _kv_rows({"analysis_id": report.get("analysis_id"), "parent_verdict": report.get("parent_verdict"), "bind_ok": bind.get("ok"), "VERDICT": d.get("VERDICT"), "NEXT": d.get("NEXT")}),
        "Freeze": _kv_rows(freeze),
        "A_Matchability": _kv_rows(dict(report.get("matchability_A") or {})),
        "C_Matchability": _kv_rows(dict(report.get("matchability_C") or {})),
        "Matched_Unmatched": _kv_rows(
            {
                "A": {k: (report.get("matchability_A") or {}).get(k) for k in ("matched_p20", "unmatched_p20", "p20_gap_matched_minus_unmatched", "MATCHABILITY_SELECTION_IS_MATERIAL")},
                "C": {k: (report.get("matchability_C") or {}).get(k) for k in ("matched_p20", "unmatched_p20", "p20_gap_matched_minus_unmatched", "MATCHABILITY_SELECTION_IS_MATERIAL")},
            }
        ),
        "Propensity_Overlap": _kv_rows({"A_ess": (report.get("overlap_A") or {}).get("effective_sample_size"), "C_ess": (report.get("overlap_C") or {}).get("effective_sample_size"), "model": "l2_logistic", "outcome_in_fit": False}),
        "OverlapWeighted_A": _kv_rows(dict(report.get("overlap_A") or {})),
        "OverlapWeighted_C": _kv_rows(dict(report.get("overlap_C") or {})),
        "Placebo_Calibration": _kv_rows(dict(report.get("placebo_calibration") or {})),
        "TwoWay_Inference": _kv_rows(dict(inf.get("by_question") or {})),
        "Multiple_Testing": _kv_rows(dict(inf.get("multiple_testing") or {})),
        "A1_Passive": _kv_rows(dict(report.get("A1") or {})),
        "A2_Confirmation": _kv_rows(dict(report.get("A2") or {})),
        "A2_Executable_Path": _kv_rows(dict((report.get("A2") or {}).get("event") or {})),
        "C1_Executable_Path": _kv_rows(dict((report.get("C1") or {}).get("event") or {})),
        "Bar_Time_Lineage": _kv_rows({"SAME_BAR_ENTRY_N": report.get("same_bar_entry_n"), "bar_start": True, "lunch": report.get("lunch_policy"), "session_flat": "15:20"}),
        "Resistance_Support": _kv_rows(dict(report.get("direction_split") or {})),
        "D2_D3_D4": _kv_rows(dict(report.get("blocks") or {})),
        "Exit_Architecture": _kv_rows(dict(report.get("exit_architecture") or {})),
        "Decision": _kv_rows({**d, "gate_A": report.get("gate_A"), "gate_C": report.get("gate_C")}),
        "Safety": _kv_rows(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# SUPPORT_RESISTANCE_MATCHED_SEPARATION_NOT_A_STRATEGY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Matched path separation is not a complete strategy.",
            "",
            f"A matchability material? **{a.get('Is A matchability selection material?')}** C? **{a.get('Is C matchability selection material?')}**",
            f"Overlap-weighted agrees? **{a.get('Does overlap-weighted full common-support analysis agree with matched pairs?')}**",
            f"ESS? **{a.get('Effective sample size?')}** extreme weights? **{a.get('Any extreme weights?')}**",
            f"Two-way A? **{a.get('Does true two-way clustered inference still support A?')}** C? **{a.get('C?')}** holm? **{a.get('After multiple-testing adjustment?')}**",
            f"A real/placebo ratio? **{a.get('A real/placebo magnitude ratio?')}**",
            f"A2 retains? **{a.get('Does A2 retain separation after waiting for confirmed rejection?')}** consumed? **{a.get('How much A edge is consumed before executable entry?')}**",
            f"C1 retains? **{a.get('Does C1 retain separation from next executable bar?')}**",
            f"A1 fillability? **{a.get('A1 passive approximation fillability?')}**",
            f"same-bar entry? **{a.get('Any same-bar entry?')}**",
            f"D2/D3/D4? **{a.get('D2 / D3 / D4 same direction?')}**",
            f"Support/resistance same? **{a.get('Support and resistance same?')}**",
            f"PnL optimization? **{a.get('Any PnL optimization?')}** strategy selected? **{a.get('Any strategy selected?')}**",
            f"Old Confirmation opened? **{a.get('Old Confirmation opened?')}** Frozen Validation opened? **{a.get('Frozen Validation opened?')}**",
            f"submit/cancel/live: **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if k not in {"_markdown", "all_pairs"}})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
