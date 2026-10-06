"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_opening_range_causal_path_test_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Causal_Audit",
    "Entry",
    "Absolute_Path",
    "D1",
    "D2",
    "D3",
    "D4",
    "Bull_Bear",
    "Structural_Risk",
    "OR_Failure",
    "Retest_Extreme",
    "Target_Semantics",
    "Target_Path",
    "Risk_Set_Control",
    "Matching_Balance",
    "Incremental_Path",
    "Trigger_Types",
    "Opening_Quality",
    "Retest_Age",
    "Reward_Geometry",
    "Execution_Consumption",
    "Cost_Feasibility",
    "Failure_Attribution",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "events", "pairs", "recs", "eligible"}
EVENT_COLS = (
    "symbol",
    "date",
    "block",
    "direction",
    "trigger_t",
    "entry_t",
    "trigger_primary",
    "in_play_reason",
    "break_to_retest_minutes",
    "away_n",
    "break_beyond",
    "impulse_mag",
    "daily_bias",
    "r5_bps",
    "r10_bps",
    "r20_bps",
    "MFE_bps",
    "MAE_bps",
    "MFE_over_R",
    "MAE_over_R",
    "R_bps",
    "risk_state",
    "or_accept_fail",
    "retest_extreme_breach",
    "target_kind",
    "target_distance_R",
    "TARGET_HIT_BEFORE_OR_FAILURE",
    "trigger_to_entry_signed_bps",
)


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def json_sanitize(obj: Any) -> Any:
    return _json_sanitize(obj)


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {
                "key": k,
                "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v,
            }
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_c)) + 2))


def _event_rows(events: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for e in events:
        rows.append({k: e.get(k) for k in EVENT_COLS})
    return rows or [{"empty": True}]


def build_markdown(report: dict[str, Any]) -> str:
    d = dict(report.get("decision") or {})
    a = dict(report.get("answers") or {})
    abs_p = dict(a.get("absolute_PB1") or {})
    d2 = dict(abs_p.get("D2") or {})
    d3 = dict(abs_p.get("D3") or {})
    d4 = dict(abs_p.get("D4") or {})
    t2e = a.get("trigger_to_entry_bps") or {}
    know = dict(report.get("causal_knowability") or {})
    lines = [
        "# PB1_OPENING_RANGE_CAUSAL_PATH_TEST_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        "Four questions are not merged.",
        "A. FACE VALIDITY: already established by V2.",
        f"B. CAUSAL KNOWABILITY: {know.get('answer')}",
        f"C. INCREMENTAL INFORMATION D2/D3/D4 claimable: {a.get('incremental_information_D2_D3_D4')}",
        f"D. ABSOLUTE TRADE UTILITY D2/D3/D4 positive: {a.get('positive_absolute_path_D2_D3_D4')}",
        "",
        f"Parent machine unchanged? {a.get('parent_machine_unchanged')}",
        f"setup_n? {a.get('setup_n')}",
        f"next-open executable n? {a.get('next_open_executable_n')}",
        f"trigger→entry bps mean/median? {t2e.get('mean')} / {t2e.get('p50')}",
        "",
        f"Absolute PB1 D2 5m/10m/20m p50: {d2.get('r5_p50')} / {d2.get('r10_p50')} / {d2.get('r20_p50')}",
        f"Absolute PB1 D3 5m/10m/20m p50: {d3.get('r5_p50')} / {d3.get('r10_p50')} / {d3.get('r20_p50')}",
        f"Absolute PB1 D4 5m/10m/20m p50: {d4.get('r5_p50')} / {d4.get('r10_p50')} / {d4.get('r20_p50')}",
        f"MFE / MAE p50: {(a.get('MFE') or {}).get('p50')} / {(a.get('MAE') or {}).get('p50')}",
        f"MFE/R p50: {(a.get('MFE_over_R') or {}).get('p50')}",
        "",
        f"OR acceptance failure rate? {a.get('OR_acceptance_failure_rate')}",
        f"Retest-extreme breach rate? {a.get('retest_extreme_breach_rate')}",
        f"TARGET_AHEAD n? {a.get('TARGET_AHEAD_n')}",
        f"NO_PREKNOWN_TARGET_AHEAD n? {a.get('NO_PREKNOWN_TARGET_AHEAD_n')}",
        f"Target-before-failure rate? {a.get('target_before_failure_rate')}",
        "",
        f"Risk-set matched n/rate? {a.get('risk_set_matched_n')} / {a.get('risk_set_match_rate')}",
        f"Matched 10m gap D2/D3/D4 p50: {(a.get('matched_10m_gap') or {}).get('D2')} / "
        f"{(a.get('matched_10m_gap') or {}).get('D3')} / {(a.get('matched_10m_gap') or {}).get('D4')}",
        "",
        f"Does PB1 have positive ABSOLUTE path in D2/D3/D4? {a.get('positive_absolute_path_D2_D3_D4')}",
        f"Does PB1 add incremental information in D2/D3/D4? {a.get('incremental_information_D2_D3_D4')}",
        f"Does path magnitude plausibly exceed cost scale? {a.get('path_magnitude_plausibly_exceeds_cost_scale')}",
        "",
        f"Reclaim trigger: {a.get('reclaim_trigger')}",
        f"Failed-push trigger: {a.get('failed_push_trigger')}",
        f"Bull: {a.get('bull')}  Bear: {a.get('bear')}",
        "",
        f"Any future control selection? {a.get('any_future_control_selection')}",
        f"Any treatment variable matched away? {a.get('any_treatment_variable_matched_away')}",
        f"Any threshold retune? {a.get('any_threshold_retune')}",
        f"Any PnL optimization? {a.get('any_pnl_optimization')}",
        f"Old Confirmation opened? {a.get('old_confirmation_opened')}",
        f"Frozen Validation opened? {a.get('frozen_validation_opened')}",
        f"Kabu50? {a.get('Kabu50')}",
        f"submit/cancel/live? {a.get('submit_cancel_live')}",
        "",
        f"VERDICT? {a.get('VERDICT')}",
        f"NEXT? {a.get('NEXT')}",
        "STOP.",
        "",
        f"Interpretation: {d.get('interpretation')}",
        "",
        "Human 48-chart review was not used as an outcome-selected filter.",
        "V2 QUESTIONABLE / RANGE_NOISE share is a contamination limitation, not a purge.",
        "No CAP / occupancy / PF / position sizing / complete strategy PnL.",
        "",
    ]
    return "\n".join(lines)


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    byb = dict(report.get("by_block") or {})
    side = dict(report.get("bull_bear") or {})
    trig = dict(report.get("trigger_types") or {})
    inc = dict(report.get("incremental_path") or {})
    events = list(report.get("events") or [])
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Causal_Audit": _kv_rows(
            {
                "face_validity": report.get("face_validity"),
                "causal_knowability": report.get("causal_knowability"),
                "matching_variable_roles": report.get("matching_variable_roles"),
            }
        ),
        "Entry": _event_rows(events),
        "Absolute_Path": _kv_rows(report.get("absolute_path")),
        "D1": _kv_rows(byb.get("D1") or {}),
        "D2": _kv_rows(byb.get("D2") or {}),
        "D3": _kv_rows(byb.get("D3") or {}),
        "D4": _kv_rows(byb.get("D4") or {}),
        "Bull_Bear": _kv_rows(side),
        "Structural_Risk": _kv_rows(report.get("structural_risk")),
        "OR_Failure": _kv_rows(report.get("or_failure")),
        "Retest_Extreme": _kv_rows(report.get("retest_extreme")),
        "Target_Semantics": _kv_rows(report.get("target_semantics")),
        "Target_Path": _kv_rows(report.get("target_path")),
        "Risk_Set_Control": _kv_rows(report.get("risk_set_control")),
        "Matching_Balance": _kv_rows(report.get("matching_balance")),
        "Incremental_Path": _kv_rows(inc),
        "Trigger_Types": _kv_rows(
            {
                "RECLAIM_RETEST_MICRO_HIGH": trig.get("RECLAIM_RETEST_MICRO_HIGH"),
                "FAILED_PUSH_THEN_CLOSE_BACK": trig.get("FAILED_PUSH_THEN_CLOSE_BACK"),
                "same_mechanism": trig.get("same_mechanism"),
                "do_not_select_the_better_one": True,
            }
        ),
        "Opening_Quality": _kv_rows(report.get("opening_quality")),
        "Retest_Age": _kv_rows(report.get("retest_age")),
        "Reward_Geometry": _kv_rows(report.get("reward_geometry")),
        "Execution_Consumption": _kv_rows(report.get("execution_consumption")),
        "Cost_Feasibility": _kv_rows(report.get("cost_feasibility")),
        "Failure_Attribution": _kv_rows(report.get("failure_attribution")),
        "Decision": _kv_rows(report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, sheets.get(name) or [])
    wb.save(OUT / "audit.xlsx")
