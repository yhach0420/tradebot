"""V19 structure hashes, trade-by-trade parity, causal audit. No new economic gate."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from research.simple_tech_entry_family.v13_analyze import _canon_num, _close, set_hash
from research.simple_tech_exit_family.v14_analyze import _finite
from research.simple_tech_exit_family.v18_analyze import latency_pack
from research.simple_tech_exit_family.v19_spec import (
    E4_FILLED_N_EXPECTED,
    PARITY_ABS_TOL,
    PARITY_PF_TOL,
    PARITY_TS_TOL,
    PARITY_YEN_TOL,
    V18_AVG_PNL_YEN_100_EXPECTED,
    V18_DROP_TOP3_SYMBOL_EXPECTED,
    V18_DROP_TOP_SYMBOL_EXPECTED,
    V18_EX_BEST_DAY_EXPECTED,
    V18_EX_TOP3_DAY_EXPECTED,
    V18_FLAT_N_EXPECTED,
    V18_LAT_MAX_EXPECTED,
    V18_LAT_MEAN_EXPECTED,
    V18_LAT_MEDIAN_EXPECTED,
    V18_LAT_P75_EXPECTED,
    V18_LOSS_N_EXPECTED,
    V18_MAX_DRAWDOWN_YEN_100_EXPECTED,
    V18_MEAN_BPS_EXPECTED,
    V18_MEDIAN_BPS_EXPECTED,
    V18_NEG_DAY_EXPECTED,
    V18_POS_DAY_EXPECTED,
    V18_PROFIT_FACTOR_YEN_100_EXPECTED,
    V18_SESSION_CLAMP_N_EXPECTED,
    V18_TOTAL_PNL_YEN_100_EXPECTED,
    V18_TRADE_N_EXPECTED,
    V18_WIN_N_EXPECTED,
    V18_WIN_RATE_EXPECTED,
)


def _int(v: Any, default: int = -1) -> int:
    if v is None:
        return int(default)
    try:
        return int(v)
    except (TypeError, ValueError):
        return int(default)


def _ts_close(a: Any, b: Any, tol: float = PARITY_TS_TOL) -> bool:
    if not _finite(a) or not _finite(b):
        return a is None and b is None
    return abs(float(a) - float(b)) <= float(tol)


def _yen_close(a: Any, b: Any) -> bool:
    return _close(a, b, PARITY_YEN_TOL)


def trade_key(row: dict[str, Any]) -> tuple[str, str, float]:
    ft = float(row.get("fill_t") or 0.0)
    return (str(row.get("date") or ""), str(row.get("symbol") or "").replace(".T", ""), round(ft, 6))


def exit_input_fill_tuples(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    out = []
    for r in rows:
        out.append(
            (
                str(r.get("date") or ""),
                str(r.get("symbol") or "").replace(".T", ""),
                _canon_num(r.get("fill_t")),
                _canon_num(r.get("fill_price")),
            )
        )
    return sorted(out)


def scheduled_exit_tuples(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    out = []
    for r in rows:
        out.append(
            (
                str(r.get("date") or ""),
                str(r.get("symbol") or "").replace(".T", ""),
                _canon_num(r.get("fill_t")),
                _canon_num(r.get("scheduled_exit_time")),
            )
        )
    return sorted(out)


def actual_exit_tuples(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    out = []
    for r in rows:
        out.append(
            (
                str(r.get("date") or ""),
                str(r.get("symbol") or "").replace(".T", ""),
                _canon_num(r.get("fill_t")),
                _canon_num(r.get("fill_price")),
                _canon_num(r.get("scheduled_exit_time")),
                _canon_num(r.get("actual_exit_quote_time")),
                _canon_num(r.get("exit_bid")),
            )
        )
    return sorted(out)


def identity_hashes(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "EXIT_INPUT_FILL_SET_HASH": set_hash(exit_input_fill_tuples(rows)),
        "SCHEDULED_EXIT_SET_HASH": set_hash(scheduled_exit_tuples(rows)),
        "ACTUAL_EXIT_SET_HASH": set_hash(actual_exit_tuples(rows)),
        "N": len(rows),
    }


def load_v18_official_trades(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Trades" not in wb.sheetnames:
            return []
        ws = wb["Trades"]
        recs = list(ws.iter_rows(values_only=True))
        if not recs:
            return []
        headers = [str(h) if h is not None else "" for h in recs[0]]
        out = []
        for rec in recs[1:]:
            d = {headers[i]: rec[i] if i < len(rec) else None for i in range(len(headers))}
            if not d.get("date") and not d.get("symbol"):
                continue
            d["date"] = str(d.get("date") or "")
            d["symbol"] = str(d.get("symbol") or "").replace(".T", "")
            out.append(d)
        return out
    finally:
        wb.close()


def trade_by_trade_parity(v19: list[dict[str, Any]], v18: list[dict[str, Any]]) -> dict[str, Any]:
    old = {trade_key(r): r for r in v18}
    details: list[dict[str, Any]] = []
    n = 0
    for r in v19:
        key = trade_key(r)
        prev = old.get(key)
        if not prev:
            n += 1
            details.append({"key": key, "reason": "MISSING_IN_V18"})
            continue
        checks = {
            "fill_price": _close(r.get("fill_price"), prev.get("fill_price"), PARITY_ABS_TOL),
            "scheduled_exit_time": _ts_close(r.get("scheduled_exit_time"), prev.get("scheduled_exit_time")),
            "actual_exit_quote_time": _ts_close(r.get("actual_exit_quote_time"), prev.get("actual_exit_quote_time")),
            "exit_bid": _close(r.get("exit_bid"), prev.get("exit_bid"), PARITY_ABS_TOL),
            "pnl_bps": _close(r.get("exit_bps"), prev.get("exit_bps"), PARITY_ABS_TOL),
            "pnl_yen_100": _yen_close(r.get("pnl_yen_100"), prev.get("pnl_yen_100")),
        }
        if not all(checks.values()):
            n += 1
            details.append({"key": key, "failed": [k for k, ok in checks.items() if not ok]})
    extra = 0
    new_keys = {trade_key(r) for r in v19}
    for r in v18:
        if trade_key(r) not in new_keys:
            extra += 1
            details.append({"key": trade_key(r), "reason": "MISSING_IN_V19"})
    n += extra
    return {
        "EXIT_TRADE_BY_TRADE_PARITY": bool(n == 0 and len(v19) == int(E4_FILLED_N_EXPECTED) and len(v18) == int(E4_FILLED_N_EXPECTED)),
        "EXIT_MISMATCH_N": int(n),
        "V19_N": len(v19),
        "V18_N": len(v18),
        "details": details[:20],
    }


def causal_exit_audit(rows: list[dict[str, Any]]) -> dict[str, Any]:
    pre = miss = last_before = carry = 0
    for r in rows:
        if r.get("EXIT_MISS"):
            miss += 1
        if r.get("LAST_BEFORE_USED_AS_EXIT"):
            last_before += 1
        sched = r.get("scheduled_exit_time")
        actual = r.get("actual_exit_quote_time")
        sess = r.get("sess_end")
        if (not _finite(actual)) or (not _finite(sched)):
            if r.get("PRE_DECISION_EXIT") or (not _finite(actual)):
                pre += 1
            continue
        if float(actual) + 1e-12 < float(sched) or r.get("PRE_DECISION_EXIT"):
            pre += 1
        if _finite(sess) and float(actual) > float(sess) + 1e-12:
            carry += 1
    pass_ok = (
        len(rows) == int(E4_FILLED_N_EXPECTED)
        and pre == 0
        and carry == 0
        and last_before == 0
        and miss == 0
    )
    return {
        "CAUSAL_EXIT_AUDIT_PASS": bool(pass_ok),
        "PRE_DECISION_EXIT_N": int(pre),
        "FUTURE_CARRY_BACK_N": int(carry),
        "LAST_BEFORE_USED_AS_EXIT_N": int(last_before),
        "EXIT_BID_MISS_N": int(miss),
        "AUDITED_N": len(rows),
    }


def numeric_parity(pack: dict[str, Any]) -> dict[str, Any]:
    pf = pack.get("PROFIT_FACTOR_YEN_100")
    checks = {
        "TRADE_N": _int(pack.get("TRADE_N")) == int(V18_TRADE_N_EXPECTED),
        "WIN_N": _int(pack.get("WIN_N")) == int(V18_WIN_N_EXPECTED),
        "LOSS_N": _int(pack.get("LOSS_N")) == int(V18_LOSS_N_EXPECTED),
        "FLAT_N": _int(pack.get("FLAT_N")) == int(V18_FLAT_N_EXPECTED),
        "WIN_RATE": _close(pack.get("WIN_RATE"), V18_WIN_RATE_EXPECTED, PARITY_ABS_TOL),
        "MEAN_BPS": _close(pack.get("MEAN_BPS"), V18_MEAN_BPS_EXPECTED, PARITY_ABS_TOL),
        "MEDIAN_BPS": _close(pack.get("MEDIAN_BPS"), V18_MEDIAN_BPS_EXPECTED, PARITY_ABS_TOL),
        "TOTAL_PNL_YEN_100": _yen_close(pack.get("TOTAL_PNL_YEN_100"), V18_TOTAL_PNL_YEN_100_EXPECTED),
        "AVG_PNL_YEN_100": _yen_close(pack.get("AVG_PNL_YEN_100"), V18_AVG_PNL_YEN_100_EXPECTED),
        "PROFIT_FACTOR_YEN_100": _close(pf, V18_PROFIT_FACTOR_YEN_100_EXPECTED, PARITY_PF_TOL),
        "MAX_DRAWDOWN_YEN_100": _yen_close(pack.get("MAX_DRAWDOWN_YEN_100"), V18_MAX_DRAWDOWN_YEN_100_EXPECTED),
        "POSITIVE_DAY_N": _int(pack.get("POSITIVE_DAY_N")) == int(V18_POS_DAY_EXPECTED),
        "NEGATIVE_DAY_N": _int(pack.get("NEGATIVE_DAY_N")) == int(V18_NEG_DAY_EXPECTED),
        "EX_BEST_DAY": _close(pack.get("EX_BEST_DAY"), V18_EX_BEST_DAY_EXPECTED, PARITY_ABS_TOL),
        "EX_TOP3_DAY": _close(pack.get("EX_TOP3_DAY"), V18_EX_TOP3_DAY_EXPECTED, PARITY_ABS_TOL),
        "DROP_TOP_SYMBOL": _close(pack.get("DROP_TOP_SYMBOL"), V18_DROP_TOP_SYMBOL_EXPECTED, PARITY_ABS_TOL),
        "DROP_TOP3_SYMBOL": _close(pack.get("DROP_TOP3_SYMBOL"), V18_DROP_TOP3_SYMBOL_EXPECTED, PARITY_ABS_TOL),
    }
    return {"V18_NUMERIC_PARITY": all(checks.values()), "checks": checks}


def latency_parity(rows: list[dict[str, Any]]) -> dict[str, Any]:
    lat = latency_pack(rows)
    checks = {
        "mean": _close(lat.get("mean"), V18_LAT_MEAN_EXPECTED, PARITY_ABS_TOL),
        "median": _close(lat.get("median"), V18_LAT_MEDIAN_EXPECTED, PARITY_ABS_TOL),
        "p75": _close(lat.get("p75"), V18_LAT_P75_EXPECTED, PARITY_ABS_TOL),
        "max": _close(lat.get("max"), V18_LAT_MAX_EXPECTED, PARITY_ABS_TOL),
        "SESSION_CLAMP_N": _int(lat.get("SESSION_CLAMP_N")) == int(V18_SESSION_CLAMP_N_EXPECTED),
        "N": _int(lat.get("N")) == int(V18_TRADE_N_EXPECTED),
    }
    return {"V18_LATENCY_PARITY": all(checks.values()), "LATENCY": lat, "checks": checks}


def decision_case(
    *,
    identity_ok: bool,
    integrity_ok: bool,
    policy_ok: bool,
    hashes: dict[str, Any],
    tbt: dict[str, Any],
    causal: dict[str, Any],
    numeric: dict[str, Any],
    latency: dict[str, Any],
) -> dict[str, Any]:
    structure_ok = bool(
        identity_ok
        and integrity_ok
        and policy_ok
        and _int(hashes.get("N"), 0) == int(E4_FILLED_N_EXPECTED)
        and bool(tbt.get("EXIT_TRADE_BY_TRADE_PARITY"))
        and _int(tbt.get("EXIT_MISMATCH_N"), -1) == 0
        and bool(causal.get("CAUSAL_EXIT_AUDIT_PASS"))
        and bool(numeric.get("V18_NUMERIC_PARITY"))
        and bool(latency.get("V18_LATENCY_PARITY"))
    )
    if not structure_ok:
        fails = []
        if not identity_ok:
            fails.append("identity")
        if not integrity_ok:
            fails.append("integrity")
        if not policy_ok:
            fails.append("exit_policy")
        if not tbt.get("EXIT_TRADE_BY_TRADE_PARITY"):
            fails.append(f"trade_mismatch_n={tbt.get('EXIT_MISMATCH_N')}")
        if not causal.get("CAUSAL_EXIT_AUDIT_PASS"):
            fails.append("causal_audit")
        if not numeric.get("V18_NUMERIC_PARITY"):
            fails.append("numeric")
        if not latency.get("V18_LATENCY_PARITY"):
            fails.append("latency")
        return {
            "PASS": False,
            "VERDICT": "SIMPLE_TECH_V19_EXIT_STRUCTURE_VERIFICATION_FAILED",
            "NEXT": f"STOP. EXIT structure verification failed ({', '.join(fails)}). Do not freeze EXIT execution.",
            "ENTRY_SIGNAL_SPEC_FROZEN_DEVELOPMENT": True,
            "ENTRY_EXECUTION_SPEC_FROZEN_DEVELOPMENT": True,
            "EXIT_SIGNAL_SPEC_FROZEN_DEVELOPMENT": False,
            "EXIT_EXECUTION_SPEC_FROZEN_DEVELOPMENT": False,
            "DEVELOPMENT_STRATEGY_STACK": None,
            "FAILS": fails,
        }
    return {
        "PASS": True,
        "VERDICT": "SIMPLE_TECH_V19_ENTRY_EXIT_DEVELOPMENT_FREEZE_READY",
        "NEXT": "STOP. Frozen ENTRY + frozen EXIT. Next run only: portfolio economics of the frozen stack. TRUE_OOS=false. STRATEGY_CERTIFIED=false.",
        "ENTRY_SIGNAL_SPEC_FROZEN_DEVELOPMENT": True,
        "ENTRY_EXECUTION_SPEC_FROZEN_DEVELOPMENT": True,
        "EXIT_SIGNAL_SPEC_FROZEN_DEVELOPMENT": True,
        "EXIT_EXECUTION_SPEC_FROZEN_DEVELOPMENT": True,
        "DEVELOPMENT_STRATEGY_STACK": "T3_PULLBACK_RCI__E4_INSIDE1_W5__FIXED180_FIRST_CAUSAL_BID",
        "FAILS": [],
    }
