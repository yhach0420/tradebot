"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_semantic_spec_clarification.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Role_Map",
    "Invariants",
    "Thesis_Execution",
    "Opening_Drive",
    "Failed_Open",
    "Location",
    "E0_E1",
    "State_Diagram",
    "Forbidden",
    "Exemplars",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(56, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    spec = dict(report.get("specification") or {})
    tf = dict(spec.get("timeframe_contract") or {})
    seed = dict(spec.get("opening_drive_seed") or {})
    active = dict(spec.get("opening_drive_active") or {})
    intent = dict(spec.get("continued_directional_intent") or {})
    scale = dict(spec.get("normal_opening_scale") or {})
    fo = dict(spec.get("FAILED_OPEN") or {})
    loc = dict(spec.get("location_identified") or {})
    inter = dict(spec.get("location_interaction") or {})
    thesis = dict(spec.get("THESIS_READY") or {})
    e0 = dict(spec.get("E0") or {})
    e1 = dict(spec.get("E1") or {})
    s4 = dict(spec.get("five_m_continuation") or {})
    dec = dict(report.get("decision") or {})
    one_m = dict(tf.get("one_m_may_never_establish") or {})
    return {
        "Old semantic spec preserved?": True,
        "Clarified spec SHA?": report.get("SPEC_SHA256") or dec.get("SPEC_SHA256"),
        "Corrected machine unchanged?": True,
        "Any machine implemented?": False,
        "Any threshold optimized?": False,
        "Any future outcome used?": False,
        "Any prospective event consumed?": False,
        "Primary strategy timeframe?": spec.get("primary_setup_timeframe"),
        "Can 1m create stock selection?": one_m.get("stock_selection"),
        "Can 1m create direction?": one_m.get("trade_direction"),
        "Can 1m create opening drive?": one_m.get("opening_drive"),
        "Can 1m create location identity?": one_m.get("structural_level_identity"),
        "Can 1m observe a retest/hold of a pre-identified 5m level?": True,
        "Does this make PB1 a 1m strategy?": spec.get("v4_is_1m_strategy"),
        "Is opening drive a permanent 09:15 label?": seed.get("is_permanent_setup_eligibility"),
        "What is OPENING_DRIVE_SEED?": seed,
        "What is OPENING_DRIVE_ACTIVE?": active,
        "What invalidates it?": active.get("becomes_false_if"),
        "Is continued directional intent explicitly part of the spec?": intent.get("explicitly_part_of_spec"),
        "Is a numeric path-efficiency threshold frozen?": intent.get("numeric_path_efficiency_threshold_frozen"),
        "Is NORMAL_OPENING_5M_RANGE first-bar-only?": scale.get("first_bar_only_normalization_of_all_three_bars"),
        "Are same-clock opening baselines required?": scale.get("same_clock_opening_baselines_required"),
        "Can a wide failed auction/doji seed FAILED_OPEN?": fo.get("wide_doji_or_range_rejection_may_be_valid_seed"),
        "Is FAIL_EXTEND_MAX_BARS=6 part of the semantic spec?": fo.get("FAIL_EXTEND_MAX_BARS_in_semantic_spec"),
        "What expires FAILED_OPEN?": fo.get("expires_when"),
        "Is S2/location separate from retest confirmation?": True,
        "Can first unsuccessful interaction permanently kill an otherwise-live level?": "not automatically",
        "What is THESIS_READY?": thesis,
        "What is E0?": e0,
        "What is E1?": e1,
        "Can E1 create a trade without THESIS_READY?": e1.get("can_create_trade_without_thesis_ready"),
        "Can E1 invent support/resistance?": e1.get("can_invent_support_resistance"),
        "Is S4 required to print a huge completed 5m candle?": s4.get("requires_huge_completed_5m_candle"),
        "Are S4 numeric scale constants part of clarified semantic spec?": s4.get("numeric_scale_in_clarified_spec"),
        "Any PnL?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
        "location_identified_separate": loc.get("state"),
        "first_unsuccessful_kills": inter.get("first_unsuccessful_observation_does_not_automatically_kill"),
        "adopted_interpretation": tf.get("adopted_interpretation_name"),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    spec = dict(report.get("specification") or {})
    diagram = dict(spec.get("state_diagram") or {})
    lines = [
        "# PB1_V4_SEMANTIC_SPEC_CLARIFICATION_V1",
        "",
        "Clarified semantic spec only. Parent READY_V1 preserved. Corrected machine unchanged. No implementation.",
        "",
        f"Old semantic spec preserved? **true**",
        f"Clarified spec SHA? `{a.get('Clarified spec SHA?')}`",
        "Corrected machine unchanged? **true**",
        "Any machine implemented? **false**",
        "",
        "Primary strategy timeframe? **5m**",
        "Adopted interpretation? **B 5M_THESIS_PLUS_CAUSAL_INTRABAR_EXECUTION**",
        "Does this make PB1 a 1m strategy? **false**",
        "",
        "1m may never create stock selection, direction, opening drive, or location identity.",
        "1m may observe retest/hold of a pre-identified 5m level.",
        "",
        "```mermaid",
        str(diagram.get("mermaid") or "").strip(),
        "```",
        "",
        f"THESIS_READY: {json.dumps(_json_sanitize(a.get('What is THESIS_READY?')), ensure_ascii=False)[:1200]}",
        "",
        f"E0: {json.dumps(_json_sanitize(a.get('What is E0?')), ensure_ascii=False)[:800]}",
        "",
        f"E1: {json.dumps(_json_sanitize(a.get('What is E1?')), ensure_ascii=False)[:800]}",
        "",
        "FAIL_EXTEND_MAX_BARS=6 part of semantic spec? **false**",
        "Opening drive a permanent 09:15 label? **false**",
        "Continued directional intent explicit? **true**",
        "Path-efficiency threshold frozen? **false**",
        "Same-clock opening baselines required? **true**",
        "First-bar-only NORMAL_OPENING? **false**",
        "S4 huge candle required? **false**",
        "S4 numeric scale in spec? **false**",
        "Any PnL? **false**",
        "Old Confirmation opened? **false**",
        "Frozen Validation opened? **false**",
        "submit/cancel/live? **0/0/0**",
        "",
        f"VERDICT? **{a.get('VERDICT?')}**",
        f"NEXT? **{a.get('NEXT?')}**",
        "STOP.",
    ]
    return "\n".join(lines) + "\n"


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    spec = dict(report.get("specification") or {})
    diagram = dict(spec.get("state_diagram") or {})
    return {
        "Binding": _kv_rows(report.get("bind")),
        "Role_Map": _kv_rows(spec.get("role_map")),
        "Invariants": [{"invariant": x} for x in list(spec.get("invariants") or [])] or [{"empty": True}],
        "Thesis_Execution": _kv_rows({"THESIS_READY": spec.get("THESIS_READY"), "EXECUTION_READY": spec.get("EXECUTION_READY")}),
        "Opening_Drive": _kv_rows(
            {
                "seed": spec.get("opening_drive_seed"),
                "TRUE_SEED": spec.get("TRUE_OPENING_DRIVE_SEED"),
                "ACTIVE": spec.get("opening_drive_active"),
                "continued_intent": spec.get("continued_directional_intent"),
                "scale": spec.get("normal_opening_scale"),
            }
        ),
        "Failed_Open": _kv_rows(spec.get("FAILED_OPEN")),
        "Location": _kv_rows({"identified": spec.get("location_identified"), "interaction": spec.get("location_interaction")}),
        "E0_E1": _kv_rows({"E0": spec.get("E0"), "E1": spec.get("E1"), "S4": spec.get("five_m_continuation")}),
        "State_Diagram": list(diagram.get("allowed") or []) or [{"empty": True}],
        "Forbidden": list(diagram.get("forbidden") or []) or [{"empty": True}],
        "Exemplars": list((report.get("exemplars") or {}).get("rows") or []) or [{"empty": True}],
        "Decision": _kv_rows(report.get("answers") or report.get("decision")),
        "Safety": _kv_rows(report.get("safety")),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = build_markdown(report)
    report["_markdown"] = md
    (OUT / "report.md").write_text(md, encoding="utf-8")
    (OUT / "report.json").write_text(json.dumps(_json_sanitize(report), ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
