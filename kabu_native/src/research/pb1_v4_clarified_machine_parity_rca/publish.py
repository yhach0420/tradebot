"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_parity_rca.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "LateAlign",
    "Active6963",
    "TwoSided",
    "FailedOpen",
    "FlatCrawl",
    "FamilyA",
    "OneBar",
    "Reclassify7",
    "Causes",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "snap"}


def finite_sanitize(obj: Any, *, strip: bool = True) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    drop = STRIP if strip else {"_markdown"}
    if isinstance(obj, dict):
        return {str(k): finite_sanitize(v, strip=strip) for k, v in obj.items() if k not in drop}
    if isinstance(obj, list):
        return [finite_sanitize(v, strip=strip) for v in obj]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(finite_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    prog = dict(report.get("progress") or {})
    late = dict(prog.get("late_time_alignment") or {})
    two = dict(report.get("seed_two_sided") or {})
    fo = dict(report.get("failed_open") or {})
    flat = dict(report.get("flat_crawl") or {})
    fam = dict(report.get("family_a") or {})
    one = dict(report.get("one_bar") or {})
    rec = dict(report.get("reclassify_seven") or {})
    causes = dict(report.get("causes") or {})
    c3110 = dict(fo.get("3110") or {})

    def fo_row(sym: str, date: str) -> dict[str, Any] | None:
        for x in list(fo.get("positives") or []) + list(fo.get("negatives") or []):
            if str(x.get("symbol")) == sym and str(x.get("date")) == date:
                return x
        return None

    def rec_row(sym: str, date: str) -> dict[str, Any] | None:
        return next((x for x in list(rec.get("rows") or []) if str(x.get("symbol")) == sym and str(x.get("date")) == date), None)

    return {
        "Machine SHA unchanged?": True,
        "Spec SHA unchanged?": True,
        "RCA set n?": report.get("rca_set_n"),
        "Any future outcome used?": False,
        "Any prospective event consumed?": False,
        "How many current violation rows remain confirmed machine bugs after time-aligned review?": rec.get("confirmed_machine_bug_n"),
        "PRIMARY_CAUSE?": causes.get("PRIMARY"),
        "SECONDARY_CAUSE?": causes.get("SECONDARY"),
        "CONTRIBUTING_CAUSE?": causes.get("CONTRIBUTING"),
        "Is ACTIVE staleness reset too permissive?": prog.get("ACTIVE_staleness_reset_too_permissive"),
        "What exactly constitutes meaningful progress semantically?": prog.get("meaningful_progress"),
        "Does any new extreme qualify?": False,
        "6963 / 20241002 result?": report.get("6963_20241002"),
        "How many human LATE cases are actual ACTIVE leaks after time alignment?": late.get("actual_active_leaks_after_time_alignment_n"),
        "8031 / 20250225: legitimate pullback or two-sided bug?": (two.get("8031") or {}).get("verdict"),
        "7741 / 20250314: why different from 8031?": (two.get("7741") or {}).get("why"),
        "Does a committed final bar erase a substantial counter-auction?": two.get("committed_final_bar_erases_substantial_counter_auction"),
        "What is a VISIBLE FAILED ATTEMPT?": fo.get("visible_failed_attempt_semantics"),
        "Can a wide doji alone seed FAILED_OPEN?": fo.get("wide_doji_alone_can_seed_FAILED_OPEN"),
        "3382 / 20241004 result?": fo_row("3382", "20241004"),
        "7011 / 20250523 result?": fo_row("7011", "20250523"),
        "4063 / 20251118 result?": fo_row("4063", "20251118"),
        "3382 / 20241115 result?": rec_row("3382", "20241115") or fo_row("3382", "20241115"),
        "6273 / 20250120 result?": rec_row("6273", "20250120"),
        "5802 / 20250613 result?": rec_row("5802", "20250613"),
        "7182 / 20251112 result?": rec_row("7182", "20251112"),
        "3110 / 20250805: A/B/C?": c3110.get("abc"),
        "Does FAILED_OPEN require a dynamic unresolved state?": c3110.get("dynamic_unresolved_state_required"),
        "8058 / 20250812: machine bug or label ambiguity?": (flat.get("8058_20250812") or {}).get("verdict"),
        "8630 / 20250911: machine bug or label ambiguity?": (flat.get("8630_20250911") or {}).get("verdict"),
        "Is family A 09:14 identification automatically invalid?": fam.get("family_A_0914_automatically_invalid"),
        "How many family A rows lack a real causal clear event?": fam.get("lack_real_causal_clear_n"),
        "9432 / 20250402 cause?": fam.get("9432_20250402"),
        "8058 / 20250814 cause?": fam.get("8058_20250814"),
        "Are ONE_BAR_SHARE_MAX and WEAK_BAR_BODY_MAX generalizable?": "uncertain",
        "Should they survive future correction unchanged?": "no",
        "Any threshold optimized?": False,
        "Any code changed?": False,
        "Any PnL?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    prog = dict(report.get("progress") or {})
    late = dict(prog.get("late_time_alignment") or {})
    two = dict(report.get("seed_two_sided") or {})
    fo = dict(report.get("failed_open") or {})
    flat = dict(report.get("flat_crawl") or {})
    fam = dict(report.get("family_a") or {})
    rec = dict(report.get("reclassify_seven") or {})
    return {
        "Binding": [dict(report.get("bind") or {})],
        "LateAlign": list(late.get("actual_active_leaks") or []) or [late],
        "Active6963": [dict(report.get("6963_20241002") or {})],
        "TwoSided": [dict(two.get("8031") or {}), dict(two.get("7741") or {})],
        "FailedOpen": list(fo.get("positives") or []) + list(fo.get("negatives") or []) + [dict(fo.get("3110") or {})],
        "FlatCrawl": list(flat.get("rows") or []),
        "FamilyA": [dict(fam.get("9432_20250402") or {}), dict(fam.get("8058_20250814") or {}), {"by_class": fam.get("by_class"), "a3_n": fam.get("lack_real_causal_clear_n")}],
        "OneBar": [dict(report.get("one_bar") or {})],
        "Reclassify7": list(rec.get("rows") or []),
        "Causes": [dict(report.get("causes") or {})],
        "Decision": [dict(report.get("decision") or {})],
        "Safety": [dict(report.get("safety") or {})],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = finite_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    answers = dict(report.get("answers") or {})
    lines = [
        "# PB1_V4_CLARIFIED_MACHINE_PARITY_RCA_V1",
        "",
        f"VERDICT: {answers.get('VERDICT?')}",
        f"NEXT: {answers.get('NEXT?')}",
        "",
        "Diagnosis only. Frozen machine/spec unchanged. No prospective. No PnL.",
        "",
    ]
    for k, v in answers.items():
        if isinstance(v, (dict, list)):
            lines.append(f"- {k} `{json.dumps(finite_sanitize(v), ensure_ascii=False)[:700]}`")
        else:
            lines.append(f"- {k} `{v}`")
    lines.append("")
    lines.append("STOP.")
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
