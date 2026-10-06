"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.temporal_model_probe import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "temporal_model_probe"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Seeds",
    "Gates",
    "Daily",
    "Integrity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# TEMPORAL JOINT MODEL ARCHITECTURE PROBE V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine. Common 600s. Exact executable-at-decision.",
            "Direct joint label frozen. Sequence 37x7 frozen. F2_UNION static 11 frozen.",
            "Model: small causal 1D TCN only. 3 seeds median/consensus. No seed selection.",
            "No Exact. No PnL. No other temporal family.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            f"SEED_N: {g.get('SEED_N')}",
            "",
            f"MEDIAN_TCN_MFE_DELTA: {_fmt(g.get('MEDIAN_TCN_MFE_DELTA'))}",
            f"MEDIAN_TCN_DOWNSIDE_DELTA: {_fmt(g.get('MEDIAN_TCN_DOWNSIDE_DELTA'))}",
            f"MEDIAN_TCN_JOINT_RATE: {_fmt(g.get('MEDIAN_TCN_JOINT_RATE'))}",
            "",
            f"SEEDS_MFE_POSITIVE_N: {g.get('SEEDS_MFE_POSITIVE_N')}",
            f"SEEDS_DOWNSIDE_POSITIVE_N: {g.get('SEEDS_DOWNSIDE_POSITIVE_N')}",
            "",
            f"CONSENSUS_MFE_POS_DAYS: {g.get('CONSENSUS_MFE_POS_DAYS')}",
            f"CONSENSUS_MFE_NEG_DAYS: {g.get('CONSENSUS_MFE_NEG_DAYS')}",
            f"CONSENSUS_DOWNSIDE_POS_DAYS: {g.get('CONSENSUS_DOWNSIDE_POS_DAYS')}",
            f"CONSENSUS_DOWNSIDE_NEG_DAYS: {g.get('CONSENSUS_DOWNSIDE_NEG_DAYS')}",
            "",
            f"CONSENSUS_MFE_EX_BEST_DAY: {_fmt(g.get('CONSENSUS_MFE_EX_BEST_DAY'))}",
            f"CONSENSUS_MFE_EX_TOP3_DAYS: {_fmt(g.get('CONSENSUS_MFE_EX_TOP3_DAYS'))}",
            f"CONSENSUS_DOWNSIDE_EX_BEST_DAY: {_fmt(g.get('CONSENSUS_DOWNSIDE_EX_BEST_DAY'))}",
            f"CONSENSUS_DOWNSIDE_EX_TOP3_DAYS: {_fmt(g.get('CONSENSUS_DOWNSIDE_EX_TOP3_DAYS'))}",
            "",
            f"DELTA_MFE_VS_FLAT_RF: {_fmt(g.get('DELTA_MFE_VS_FLAT_RF'))}",
            f"DELTA_DOWNSIDE_VS_FLAT_RF: {_fmt(g.get('DELTA_DOWNSIDE_VS_FLAT_RF'))}",
            f"DELTA_JOINT_RATE_VS_FLAT_RF: {_fmt(g.get('DELTA_JOINT_RATE_VS_FLAT_RF'))}",
            "",
            f"TEMPORAL_ORDER_INCREMENTAL: {_tf(g.get('TEMPORAL_ORDER_INCREMENTAL'))}",
            "",
            f"MEDIAN_ROC_AUC: {_fmt(g.get('MEDIAN_ROC_AUC'))}",
            f"MEDIAN_AVERAGE_PRECISION: {_fmt(g.get('MEDIAN_AVERAGE_PRECISION'))}",
            "",
            f"TEMPORAL_MODEL_PASS: {_tf(g.get('TEMPORAL_MODEL_PASS'))}",
            "",
            f"PRIMARY_FINDING: {g.get('PRIMARY_FINDING')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No Exact. No runtime candidate. No other temporal model.",
            "Architecture/kernel/dilation/epochs frozen. No seed selection.",
            "Runtime/C14 unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
