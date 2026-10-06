"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.c3_canonical_retrain import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "c3_canonical_retrain"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Integrity",
    "OOFGate",
    "PairedDays",
    "Folds",
    "Coverage",
    "Parity",
    "Exact",
    "ExecCoupling",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# C3 CANONICAL RETRAIN SAME PROTOCOL V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine training data. Same 27 Ridge specs. Nested OOF. Rank then join TARGET.",
            "No RF. No pairwise. No new feature. No execution-aware objective. No PnL model selection.",
            "Legacy C3 coefficients were not reused.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"CANONICAL_TRAINING_ROWS: {g.get('CANONICAL_TRAINING_ROWS')}",
            f"MATCHED_RANKING_ROWS: {g.get('MATCHED_RANKING_ROWS')}",
            "",
            f"CURRENT_TOP1_UPLIFT: {g.get('CURRENT_TOP1_UPLIFT')}",
            f"CURRENT_TOP3_UPLIFT: {g.get('CURRENT_TOP3_UPLIFT')}",
            f"CURRENT_TOP5_UPLIFT: {g.get('CURRENT_TOP5_UPLIFT')}",
            "",
            f"CANON_C3_TOP1_UPLIFT: {g.get('CANON_C3_TOP1_UPLIFT')}",
            f"CANON_C3_TOP3_UPLIFT: {g.get('CANON_C3_TOP3_UPLIFT')}",
            f"CANON_C3_TOP5_UPLIFT: {g.get('CANON_C3_TOP5_UPLIFT')}",
            "",
            f"TOP3_DELTA_MEAN: {g.get('TOP3_DELTA_MEAN')}",
            f"TOP3_DELTA_MEDIAN: {g.get('TOP3_DELTA_MEDIAN')}",
            f"TOP3_DELTA_POSITIVE_DAYS: {g.get('TOP3_DELTA_POSITIVE_DAYS')}",
            f"TOP3_DELTA_NEGATIVE_DAYS: {g.get('TOP3_DELTA_NEGATIVE_DAYS')}",
            f"TOP3_DELTA_EX_BEST_DAY: {g.get('TOP3_DELTA_EX_BEST_DAY')}",
            f"TOP3_DELTA_EX_TOP3_DAYS: {g.get('TOP3_DELTA_EX_TOP3_DAYS')}",
            "",
            f"CANON_C3_MEAN_DAILY_SPEARMAN: {g.get('CANON_C3_MEAN_DAILY_SPEARMAN')}",
            "",
            f"MOST_COMMON_SPEC: {g.get('MOST_COMMON_SPEC')}",
            f"MOST_COMMON_SPEC_SHARE: {g.get('MOST_COMMON_SPEC_SHARE')}",
            "",
            f"SPEC_NATIVE_TOP3_DELTA: {g.get('SPEC_NATIVE_TOP3_DELTA')}",
            f"COMMON_POP_TOP3_DELTA: {g.get('COMMON_POP_TOP3_DELTA')}",
            f"EDGE_DEPENDS_ON_COVERAGE_SELECTION: {_tf(g.get('EDGE_DEPENDS_ON_COVERAGE_SELECTION'))}",
            "",
            f"OOF_RANKING_GATE_PASS: {_tf(g.get('OOF_RANKING_GATE_PASS'))}",
            "",
            f"A0_PARITY: {_tf(g.get('A0_PARITY'))}",
            f"A2_PARITY: {_tf(g.get('A2_PARITY'))}",
            "",
            f"FINAL_SPEC: {g.get('FINAL_SPEC')}",
            f"FINAL_EXACT_RAN: {_tf(g.get('FINAL_EXACT_RAN'))}",
            f"OOF_EXACT_RAN: {_tf(g.get('OOF_EXACT_RAN'))}",
            "",
            f"CANON_C3_OOF_EXACT: {g.get('CANON_C3_OOF_EXACT')}",
            f"CANON_C3_FINAL_EXACT: {g.get('CANON_C3_FINAL_EXACT')}",
            "",
            f"EXACT_SUCCESS_BAR_PASS: {_tf(g.get('EXACT_SUCCESS_BAR_PASS'))}",
            f"PASSIVE_FILL_ADVERSE_SELECTION: {g.get('PASSIVE_FILL_ADVERSE_SELECTION')}",
            "",
            f"LEGACY_C3_COMPARISON: {g.get('LEGACY_C3_COMPARISON')}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No C4. No execution-aware model auto-start. Runtime/C14 unchanged.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
