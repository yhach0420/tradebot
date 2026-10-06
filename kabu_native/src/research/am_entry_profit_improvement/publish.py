"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_entry_profit_improvement"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "summary",
    "outer_folds",
    "inner_selection",
    "spec_performance",
    "daily_pnl",
    "trades",
    "integrity",
    "Manifest",
    "Parity",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# AM ENTRY PROFIT IMPROVEMENT PROGRAM V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. Nested 27-spec utility program. Label = I(W5 fill) × current EXIT pnl_yen_100.",
            "Inner selection = portfolio net_pnl_yen_100. Outer held-out never retunes the spec.",
            "Execution = Corrected Passive Fill W5 + frozen EXIT + occupancy. Runtime WAIT_SEC remains 1.0.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0. Formal candidate=0.",
            "",
            "## REQUIRED FINAL OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"OUTER_FOLD_N: {_fmt(g.get('OUTER_FOLD_N'))}",
            f"ARCHITECTURE_N: {_fmt(g.get('ARCHITECTURE_N'))}",
            f"REPRESENTATION_N: {_fmt(g.get('REPRESENTATION_N'))}",
            f"SPEC_N: {_fmt(g.get('SPEC_N'))}",
            "",
            f"CURRENT_TRADE_N: {_fmt(g.get('CURRENT_TRADE_N'))}",
            f"CURRENT_NET_PNL: {_fmt(g.get('CURRENT_NET_PNL'))}",
            f"CURRENT_PF: {_fmt(g.get('CURRENT_PF'))}",
            f"CURRENT_MAX_DD: {_fmt(g.get('CURRENT_MAX_DD'))}",
            "",
            f"FILL_ONLY_TRADE_N: {_fmt(g.get('FILL_ONLY_TRADE_N'))}",
            f"FILL_ONLY_NET_PNL: {_fmt(g.get('FILL_ONLY_NET_PNL'))}",
            f"FILL_ONLY_PF: {_fmt(g.get('FILL_ONLY_PF'))}",
            f"FILL_ONLY_MAX_DD: {_fmt(g.get('FILL_ONLY_MAX_DD'))}",
            "",
            f"CANDIDATE_TRADE_N: {_fmt(g.get('CANDIDATE_TRADE_N'))}",
            f"CANDIDATE_NET_PNL: {_fmt(g.get('CANDIDATE_NET_PNL'))}",
            f"CANDIDATE_PF: {_fmt(g.get('CANDIDATE_PF'))}",
            f"CANDIDATE_MAX_DD: {_fmt(g.get('CANDIDATE_MAX_DD'))}",
            "",
            f"DELTA_PNL_VS_CURRENT: {_fmt(g.get('DELTA_PNL_VS_CURRENT'))}",
            f"DELTA_PF_VS_CURRENT: {_fmt(g.get('DELTA_PF_VS_CURRENT'))}",
            f"DELTA_DD_VS_CURRENT: {_fmt(g.get('DELTA_DD_VS_CURRENT'))}",
            "",
            f"DELTA_PNL_VS_FILL_ONLY: {_fmt(g.get('DELTA_PNL_VS_FILL_ONLY'))}",
            "",
            f"PAIRED_POS_DAYS: {_fmt(g.get('PAIRED_POS_DAYS'))}",
            f"PAIRED_NEG_DAYS: {_fmt(g.get('PAIRED_NEG_DAYS'))}",
            f"PAIRED_ZERO_DAYS: {_fmt(g.get('PAIRED_ZERO_DAYS'))}",
            "",
            f"PAIRED_MEDIAN_DAILY_DELTA: {_fmt(g.get('PAIRED_MEDIAN_DAILY_DELTA'))}",
            "",
            f"EX_BEST_DAY_PNL_DELTA: {_fmt(g.get('EX_BEST_DAY_PNL_DELTA'))}",
            f"EX_TOP3_DAYS_PNL_DELTA: {_fmt(g.get('EX_TOP3_DAYS_PNL_DELTA'))}",
            "",
            f"SELECTED_ARCH_COUNTS: {_fmt(g.get('SELECTED_ARCH_COUNTS'))}",
            "",
            f"SELECTED_REP_COUNTS: {_fmt(g.get('SELECTED_REP_COUNTS'))}",
            "",
            f"AM_ENTRY_PROFIT_IMPROVEMENT_PASS: {_tf(g.get('AM_ENTRY_PROFIT_IMPROVEMENT_PASS'))}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            f"NEXT: {g.get('NEXT')}",
            "",
            "## STOP",
            "",
            "Nested 27-spec program complete. No Runtime change. No W5 adoption.",
            "No new feature. No new target. submit/cancel/live=0/0/0.",
            "",
        ]
    )
