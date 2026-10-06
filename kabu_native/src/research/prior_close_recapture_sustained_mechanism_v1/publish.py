"""Write report.json / report.md / audit.xlsx only. No CSV dump."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.prior_close_recapture_sustained_mechanism_v1 import ANALYSIS_ID
from research.prior_close_recapture_sustained_mechanism_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Labels",
    "Feature_Integrity",
    "Univariate",
    "Block_Stability",
    "Monotonicity",
    "TopSymbol_Exclusion",
    "LeaveOneBlockOut",
    "ShallowTree",
    "Mechanism_Interpretation",
    "Next_Strategy",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    uni = list(report.get("univariate") or [])
    mono = list(report.get("monotonicity") or [])
    lobo_pack = dict(report.get("lobo") or {})
    interp = dict(report.get("interpretation") or {})
    nxt = dict(interp.get("NEXT_STRATEGY") or {}) or {"empty": True}
    return {
        "Summary": _kv(dict(report.get("answers") or {})),
        "Labels": _kv(dict(report.get("labels") or {})),
        "Feature_Integrity": list(report.get("feature_integrity") or []) or [{"empty": True}],
        "Univariate": [
            {
                "feature": u.get("feature"),
                "S_n": (u.get("sustained") or {}).get("n"),
                "S_median": (u.get("sustained") or {}).get("median"),
                "S_q25": (u.get("sustained") or {}).get("q25"),
                "S_q75": (u.get("sustained") or {}).get("q75"),
                "F_n": (u.get("failed") or {}).get("n"),
                "F_median": (u.get("failed") or {}).get("median"),
                "F_q25": (u.get("failed") or {}).get("q25"),
                "F_q75": (u.get("failed") or {}).get("q75"),
                "median_diff": u.get("median_diff_S_minus_F"),
                "direction": u.get("direction"),
                "block_agree_n": u.get("block_agree_n"),
                "candidate": u.get("mechanism_candidate"),
            }
            for u in uni
        ]
        or [{"empty": True}],
        "Block_Stability": [
            {"feature": u.get("feature"), **dict(u.get("block_signs") or {}), "agree": u.get("block_agree_n")}
            for u in uni
        ]
        or [{"empty": True}],
        "Monotonicity": [
            {
                "feature": m.get("feature"),
                "ok": m.get("ok"),
                "monotonic_up": m.get("monotonic_up"),
                "monotonic_down": m.get("monotonic_down"),
                "q4_minus_q1": m.get("q4_minus_q1_sustained_rate"),
                "bins": m.get("bins"),
                "edges": m.get("quartile_edges"),
            }
            for m in mono
        ]
        or [{"empty": True}],
        "TopSymbol_Exclusion": [
            {
                "feature": u.get("feature"),
                "top_symbol_excluded_holds": u.get("top_symbol_excluded_holds"),
                "leave_one_day_holds": u.get("leave_one_day_holds"),
            }
            for u in uni
        ]
        or [{"empty": True}],
        "LeaveOneBlockOut": list(lobo_pack.get("folds") or []) or _kv(lobo_pack),
        "ShallowTree": _kv(dict(report.get("tree") or {})),
        "Mechanism_Interpretation": _kv(interp),
        "Next_Strategy": _kv(nxt) if isinstance(nxt, dict) else [{"value": nxt}],
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: {d.get('NEXT')}",
        f"CASE: {d.get('CASE')}",
        "TRUE_OOS: false",
        "CERTIFIED: false",
        "KIND: MECHANISM_DISCOVERY_ONLY",
        "",
        str((report.get("interpretation") or {}).get("MARKET_MECHANISM") or d.get("INTERPRETATION") or ""),
        "",
    ]
    for i in range(1, 18):
        matches = [k for k in a if k.split("_", 1)[0] == str(i)]
        if matches:
            lines.append(f"{i}. {matches[0]}: {a.get(matches[0])}")
    lines.extend(["", "No V5 rescue. No Full Causal economics in this run.", "", "STOP.", ""])
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
