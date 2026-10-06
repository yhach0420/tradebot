"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.systematic_state_transition_full_strategy_v1 import ANALYSIS_ID
from research.systematic_state_transition_full_strategy_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "canary",
    "ranking",
    "winner",
    "folds",
    "winner_blocks",
    "frozen_strategy",
    "leakage",
    "hashes",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    can = dict(report.get("canary") or {})
    ranking = list(report.get("ranking") or [])
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        "TRUE_OOS: false",
        "",
        "CERTIFIED: false",
        "",
        f"CASE: {_fmt(d.get('CASE'))}",
        "",
        f"VERDICT: {_fmt(d.get('VERDICT'))}",
        "",
        f"NEXT: {_fmt(d.get('NEXT'))}",
        "",
        f"CANARY_ID: {can.get('CANARY_ID')}",
        "",
        f"CANARY_HARD_PASS: {_fmt(can.get('HARD_PASS'))}",
        "",
        f"ECONOMICS_OPENED: {_fmt(d.get('ECONOMICS_OPENED'))}",
        "",
        "LEGACY_E4_X3_CANARY: false",
        "",
        "CAUSAL_EX_TOP1_PART_OF_STRATEGY: false",
        "",
        "## Answers",
        "",
    ]
    for k, v in a.items():
        if str(k).startswith("_") or k in {"51_complete_ranking_table", "9_exact_state_registry", "26_complete_architecture_inventory_table"}:
            continue
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        lines.append(f"{k}: {v if isinstance(v, str) else _fmt(v)}")
        lines.append("")
    lines.append("## Ranking")
    lines.append("")
    lines.append("| candidate | gate | trades | pnl | PF | G6 | score |")
    lines.append("|---|---|---|---|---|---|---|")
    for r in ranking:
        g6 = (r.get("g_table") or {}).get("G6_CAUSAL_EX_TOP1")
        lines.append(
            f"| {r.get('candidate_id')} | {r.get('gate')} | {r.get('trade_n')} | {r.get('pnl')} | {r.get('PF')} | {g6} | {r.get('score')} |"
        )
    lines.append("")
    lines.append("STOP.")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")
    extra = [p for p in OUT.iterdir() if p.is_file() and p.name not in {"report.json", "report.md", "audit.xlsx"}]
    if extra:
        raise RuntimeError("OUT_FILE_COUNT " + ",".join(p.name for p in extra))
