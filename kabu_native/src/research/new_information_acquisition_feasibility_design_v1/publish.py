"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.new_information_acquisition_feasibility_design_v1 import ANALYSIS_ID
from research.new_information_acquisition_feasibility_design_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "source_boundary",
    "protocol_messages",
    "event_order",
    "historical_access",
    "live_access",
    "historical_live_pair",
    "timestamp_semantics",
    "date_exposure",
    "clean_history",
    "reference_data",
    "book_reconstruction",
    "storage_compute",
    "licensing",
    "cost_access",
    "provider_normalization",
    "feasibility_gates",
    "restart_gate",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    hist = {**dict(report.get("historical_access") or {}), **{"pcap_" + k: v for k, v in dict(report.get("historical_pcap") or {}).items()}}
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "parent": _kv(dict(report.get("parent") or {})),
        "source_boundary": _kv(dict(report.get("source_boundary") or {})),
        "protocol_messages": _kv(dict(report.get("protocol_messages") or {})),
        "event_order": _kv(dict(report.get("event_order") or {})),
        "historical_access": _kv(hist),
        "live_access": list(report.get("live_access") or []),
        "historical_live_pair": _kv(dict(report.get("historical_live_pair") or {})),
        "timestamp_semantics": _kv(dict(report.get("timestamp_semantics") or {})),
        "date_exposure": list(report.get("date_exposure_rows") or []) or _kv(dict(report.get("date_exposure") or {})),
        "clean_history": _kv(dict(report.get("clean_history") or {})),
        "reference_data": list(report.get("reference_data") or []),
        "book_reconstruction": _kv(dict(report.get("book_reconstruction") or {})),
        "storage_compute": _kv(dict(report.get("storage_compute") or {})),
        "licensing": _kv(dict(report.get("licensing") or {})),
        "cost_access": _kv(dict(report.get("cost_access") or {})),
        "provider_normalization": _kv(dict(report.get("provider_normalization") or {})),
        "feasibility_gates": _kv(dict(report.get("feasibility_gates") or {})),
        "restart_gate": _kv(dict(report.get("restart_gate") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"STRATEGY_RESEARCH_STATUS: PAUSED",
            f"RESTART_ALLOWED: false",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- selected: {a.get('2_selected_source_family')}",
            f"- policy SHA exact: {a.get('3_source_policy_SHA_exact')}",
            f"- event order replayable: {a.get('16_causal_event_order_replayable')}",
            f"- exposed date N: {a.get('27_project_exposed_date_N')} range {a.get('28_project_exposed_date_range')}",
            f"- clean historical 30d structurally possible: {a.get('31_at_least_30_clean_days_structurally_possible')}",
            f"- failed gates: {d.get('FAILED_GATES')}",
            "",
            "No market data purchased. No collector. No strategy. Do not consume 20260907+.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
