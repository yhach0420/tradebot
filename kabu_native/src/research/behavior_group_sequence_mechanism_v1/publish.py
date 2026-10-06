"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.behavior_group_sequence_mechanism_v1.isolation import OUT
from research.behavior_group_sequence_mechanism_v1.replay import SPECS

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Prequential_Group_Map",
    "Group_Definitions",
    "Sector_Leader_State",
    "Laggard_State",
    "Catchup_Transition",
    "Sequence_Group_Pairs",
    "Global_vs_Group",
    "Complete_Strategy",
    "Execution",
    "BreakEven_Cost",
    "D1_D4",
    "Symbol_Concentration",
    "HM1_Comparison",
    "Failure_Mechanisms",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _econ_keys(p: dict[str, Any]) -> dict[str, Any]:
    e = dict(p.get("economics") or {})
    s = dict(p.get("spec") or {})
    return {
        "candidate_id": s.get("candidate_id"),
        "sequence": s.get("sequence"),
        "pair": s.get("pair"),
        "group_filter": s.get("group_filter"),
        "window": s.get("window"),
        "diagnostic": s.get("diagnostic"),
        "promoted": (p.get("promotion") or {}).get("promoted"),
        "fail_reasons": (p.get("promotion") or {}).get("fail_reasons"),
        "trade_n": e.get("trade_n"),
        "day_n": e.get("day_n"),
        "symbol_n": e.get("symbol_n"),
        "X0": e.get("mean_x0_bps"),
        "X1": e.get("mean_x1_bps"),
        "PF": e.get("profit_factor"),
        "hit_rate": e.get("hit_rate"),
        "daily_mean": e.get("daily_mean_bps"),
        "daily_median": e.get("daily_median_bps"),
        "max_DD": e.get("max_dd_daily_mean_bps"),
        "MFE": e.get("mean_mfe_bps"),
        "MAE": e.get("mean_mae_bps"),
        "favorable_first": e.get("favor_first_p"),
        "time_to_favorable": e.get("mean_time_to_favorable"),
        "time_to_failure": e.get("mean_time_to_failure"),
        "exit_reason": e.get("exit_reasons"),
        "break_even_execution_bps": e.get("break_even_execution_bps"),
        "top_symbol_share": e.get("top_symbol_share_of_positive_bps"),
        "block_mean_x0": e.get("block_mean_x0"),
        "CAP": s.get("occupancy"),
        "same_symbol": e.get("same_symbol"),
        "session_close": e.get("session_close"),
        "x1_tax_bps": e.get("x1_tax_bps"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    preq = dict(report.get("prequential") or {})
    plays = list((report.get("candidates") or {}).get("proposals") or [])
    atlas = dict(report.get("atlas_reference") or {})
    hm1 = dict(report.get("hm1_reference") or {})
    by_seq = dict((report.get("walk_meta") or {}).get("by_sequence") or {})
    conc = []
    for p in plays:
        e = dict(p.get("economics") or {})
        mp = dict(e.get("symbol_n_map") or {})
        top = sorted(mp.items(), key=lambda kv: -int(kv[1]))[:15]
        for i, (sym, n) in enumerate(top, start=1):
            conc.append(
                {
                    "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                    "rank": i,
                    "symbol": sym,
                    "trade_n": n,
                    "top_symbol_share_of_positive_bps": e.get("top_symbol_share_of_positive_bps"),
                }
            )
    hm_rows = []
    for p in plays:
        e = dict(p.get("economics") or {})
        s = dict(p.get("spec") or {})
        x0 = e.get("mean_x0_bps")
        be = e.get("break_even_execution_bps")
        hm_rows.append(
            {
                "candidate_id": s.get("candidate_id"),
                "candidate_x0": x0,
                "candidate_x1": e.get("mean_x1_bps"),
                "HM1_x0": hm1.get("mean_x0_bps"),
                "HM1_x1": hm1.get("mean_x1_bps"),
                "beats_hm1_gross": (x0 is not None and hm1.get("mean_x0_bps") is not None and float(x0) > float(hm1["mean_x0_bps"])),
                "break_even_execution_bps": be,
                "canonical_tax_bps": 8.0,
                "near_threshold_like_hm1": (be is not None and 0 < float(be) <= 8.0),
                "hm1_retuned": False,
            }
        )
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "atlas_id": report.get("atlas_id"),
                "split_sha256": report.get("split_sha256"),
                "block_sha256": report.get("block_sha256"),
                "prequential": True,
                "full_discovery_group_leakage": False,
                "catalog_stopped": True,
                "no_new_paid_data": True,
                "hm1_tuned": False,
                "five_minute_grid": False,
                "kabu_50": False,
                "zero_trade_cause": report.get("zero_trade_cause"),
                "implementation_mismatch_fixed": report.get("implementation_mismatch_fixed"),
            }
        ),
        "Prequential_Group_Map": list(preq.get("rows") or [{"empty": True}]),
        "Group_Definitions": list(preq.get("definitions") or [{"empty": True}]),
        "Sector_Leader_State": list(report.get("leader_state") or [{"empty": True}]),
        "Laggard_State": list(report.get("laggard_state") or [{"empty": True}]),
        "Catchup_Transition": (
            [{"sequence": k, "event_n": v, "future_catchup_used_for_entry": False} for k, v in by_seq.items()]
            or [{"empty": True}]
        ),
        "Sequence_Group_Pairs": [
            {
                "candidate_id": s.get("candidate_id"),
                "sequence": s.get("sequence"),
                "pair": s.get("pair"),
                "group_filter": s.get("group_filter"),
                "window": s.get("window"),
                "exit_kind": s.get("exit_kind"),
                "diagnostic": s.get("diagnostic"),
                "thesis": s.get("thesis"),
                "giant_matrix": False,
            }
            for s in SPECS
        ],
        "Global_vs_Group": list(report.get("global_vs_group") or [{"empty": True}]),
        "Complete_Strategy": [_econ_keys(p) for p in plays] or [{"empty": True}],
        "Execution": [
            {
                "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                "X0_gross": (p.get("economics") or {}).get("mean_x0_bps"),
                "X1_execution_adjusted": (p.get("economics") or {}).get("mean_x1_bps"),
                "x1_tax_bps": 8.0,
                "tax_changed_to_rescue": False,
                "not_bid_ask": True,
                "fill": "next_bar_open",
            }
            for p in plays
        ]
        or [{"empty": True}],
        "BreakEven_Cost": [
            {
                "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                "mean_x0_bps": (p.get("economics") or {}).get("mean_x0_bps"),
                "break_even_execution_bps": (p.get("economics") or {}).get("break_even_execution_bps"),
                "canonical_tax_bps": 8.0,
                "far_below_vs_near_hm1": (
                    "near_threshold"
                    if (p.get("economics") or {}).get("break_even_execution_bps") is not None
                    and 3.0 <= float((p.get("economics") or {}).get("break_even_execution_bps") or 0) <= 8.0
                    else "far_below_or_negative"
                ),
            }
            for p in plays
        ]
        or [{"empty": True}],
        "D1_D4": [
            {
                "candidate_id": (p.get("spec") or {}).get("candidate_id"),
                "window": (p.get("spec") or {}).get("window"),
                "block_mean_x0": (p.get("economics") or {}).get("block_mean_x0"),
                "block_positive_n": (p.get("economics") or {}).get("block_positive_n"),
                "d1_used_for_group_eval": False,
            }
            for p in plays
        ]
        or [{"empty": True}],
        "Symbol_Concentration": conc or [{"empty": True}],
        "HM1_Comparison": hm_rows or [{"empty": True}],
        "Failure_Mechanisms": [
            {
                "sequence": "BREAKOUT20",
                "role_candidate": "VETO / failure-state / EXIT evidence",
                "causal_role_proven": False,
                "applied_as_post_hoc_filter": False,
                "note": "Atlas often ranks BREAKOUT20 / COMPRESSION_THEN_BREAKOUT / OTHER_VOL_EXPAND as weak. Not auto-filters.",
            },
            {
                "sequence": "COMPRESSION_THEN_BREAKOUT",
                "role_candidate": "VETO / failure-state",
                "causal_role_proven": False,
                "applied_as_post_hoc_filter": False,
            },
            {
                "sequence": "OTHER_VOL_EXPAND",
                "role_candidate": "VETO / failure-state",
                "causal_role_proven": False,
                "applied_as_post_hoc_filter": False,
            },
            {
                "atlas_reference_members_not_used_for_entry": True,
                "full_discovery_catchup_n": len(list((atlas.get("members_full_discovery_reference_only") or {}).get("SECTOR_LAGGARD_CATCHUP") or [])),
            },
        ],
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0", "new_paid_data": False, "frozen_validation_opened": False, "kabu_50": False}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    lines = [
        "# BEHAVIOR_GROUP_SEQUENCE_MECHANISM_V1",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: **{d.get('NEXT')}**",
        "",
        d.get("INTERPRETATION") or "",
        "",
        "Accepted atlas: `ONE_MINUTE_BEHAVIOR_ATLAS_READY_NO_COMPLETE_PLAYBOOK_V1`. Do not conclude native 1-minute information has no edge.",
        "",
        f"Prequential group assignment? **{a.get('Was_behavior_group_assignment_made_prequentially')}**",
        f"Full-Discovery group leakage? **{a.get('Any_full_Discovery_group_leakage')}**",
        f"Why PB_NATIVE_LAG_CATCHUP had zero trades: {a.get('Why_did_prior_PB_NATIVE_LAG_CATCHUP_have_zero_trades')}",
        f"Implementation mismatch fixed? **{a.get('Was_that_implementation_mismatch_fixed')}**",
        "",
        f"SECTOR_LAGGARD_CATCHUP + catch-up trigger produces trades? **{a.get('SECTOR_LAGGARD_CATCHUP_plus_catchup_trigger_produces_trades')}**",
        f"BG_CATCHUP_RS_TURN trade_n=**{a.get('BG_CATCHUP_RS_TURN_trade_n')}** day_n=**{a.get('BG_CATCHUP_RS_TURN_day_n')}** symbol_n=**{a.get('BG_CATCHUP_RS_TURN_symbol_n')}**",
        f"X0=**{a.get('BG_CATCHUP_RS_TURN_X0')}** X1=**{a.get('BG_CATCHUP_RS_TURN_X1')}** PF=**{a.get('BG_CATCHUP_RS_TURN_PF')}** D2-D4=`{a.get('BG_CATCHUP_RS_TURN_D1_D4')}`",
        f"GM_CATCHUP_RS_TURN trade_n=**{a.get('GM_CATCHUP_RS_TURN_trade_n')}**",
        "",
        f"Laggard + VWAP reclaim beat global VWAP reclaim? **{a.get('Does_laggard_VWAP_reclaim_beat_global_VWAP_reclaim')}**",
        f"CONTINUATION improve impulse/reclaim? **{a.get('Does_CONTINUATION_conditioning_improve_impulse_or_reclaim')}**",
        f"OPENING_MOMENTUM improve OPENING_GAP_HOLD? **{a.get('Does_OPENING_MOMENTUM_conditioning_improve_OPENING_GAP_HOLD')}**",
        f"SECTOR_LEADER improve continuation? **{a.get('Does_SECTOR_LEADER_conditioning_improve_continuation')}**",
        "",
        f"Any Complete Strategy X1 > 0? **{a.get('Any_Complete_Strategy_X1_gt_0')}**",
        f"Closest X1: `{a.get('closest_X1_id')}` = **{a.get('closest_X1')}** break-even execution bps=**{a.get('closest_break_even_execution_bps')}**",
        f"Beat frozen HM1 gross? **{a.get('Does_any_new_candidate_beat_frozen_HM1')}** `{a.get('candidates_beating_HM1_gross')}`",
        f"HM1 frozen X0=**{a.get('HM1_X0')}** X1=**{a.get('HM1_X1')}** (not retuned)",
        f"Distinct behavior groups with executable edge: **{a.get('How_many_distinct_behavior_groups_have_executable_edge')}** `{a.get('executable_groups')}`",
        f"Material improvement n=**{a.get('material_improvement_n')}** `{a.get('material_pairs')}`",
        f"Frozen mechanism definitions: `{a.get('frozen_mechanism_definitions')}`",
        "",
        f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
        f"Old Confirmation used to design? **{a.get('Old_Confirmation_used_to_design')}**",
        f"New paid data? **{a.get('New_paid_data')}**",
        f"Kabu50? **{a.get('Kabu50')}**",
        f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
        "",
        "STOP.",
        "",
    ]
    return "\n".join(lines)


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
