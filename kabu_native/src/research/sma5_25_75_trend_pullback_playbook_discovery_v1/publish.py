"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.sma5_25_75_trend_pullback_playbook_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "State_Machine",
    "D1_Freeze",
    "Setups",
    "A_vs_B",
    "SMA25_Pullback",
    "VWAP",
    "SR",
    "Participation",
    "Location",
    "Execution",
    "MFE_MAE",
    "Structural_Risk",
    "Concentration",
    "D1",
    "D2",
    "D3",
    "D4",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    freeze = dict(report.get("freeze") or {})
    s = dict(report.get("setups") or {})
    ab = dict(report.get("a_vs_b") or {})
    a = dict(s.get("A") or {})
    ab_rows = []
    for b in ("D2", "D3", "D4"):
        ab_rows.append({"block": b, **dict((ab.get("blocks") or {}).get(b) or {})})
    sma = dict((report.get("sma25_pullback") or {}).get("matched") or {})
    sma_rows = [{"block": b, **dict((sma.get("blocks") or {}).get(b) or {})} for b in ("D2", "D3", "D4")]
    block_rows = {}
    for b in ("D1", "D2", "D3", "D4"):
        block_rows[b] = _kv_rows(dict((s.get("by_block") or {}).get(b) or {}))
    return {
        "Binding": _kv_rows(
            {
                "ok": (report.get("bind") or {}).get("ok"),
                "parent_verdict": (report.get("bind") or {}).get("parent_verdict"),
                "peer_rescue": False,
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
            }
        ),
        "State_Machine": _kv_rows({k: freeze.get(k) for k in ("playbook", "sma_periods", "trend", "pullback", "trigger", "entry", "invalidation", "same_bar_entry")}),
        "D1_Freeze": _kv_rows({"FREEZE_SHA256": freeze.get("FREEZE_SHA256"), "MACHINE_SHA256": freeze.get("MACHINE_SHA256"), "d1_only": True, "d2_not_read": True, "locked": freeze.get("locked")}),
        "Setups": [
            {"kind": "A", **dict(s.get("A") or {})},
            {"kind": "B", **dict(s.get("B") or {})},
            {"kind": "C", **dict(s.get("C") or {})},
        ],
        "A_vs_B": ab_rows or [{"block": None}],
        "SMA25_Pullback": sma_rows or [{"block": None}],
        "VWAP": _kv_rows(report.get("vwap") or {}),
        "SR": _kv_rows(report.get("sr") or {}),
        "Participation": _kv_rows(report.get("participation") or {}),
        "Location": list(report.get("location") or [{"cat": None}]),
        "Execution": _kv_rows(
            {
                "same_bar_entry_n": report.get("same_bar_entry_n"),
                "signal_to_entry_bps": a.get("signal_to_entry_bps"),
                "dead_before_entry_n": report.get("dead_before_entry_n"),
            }
        ),
        "MFE_MAE": _kv_rows({"mfe": a.get("mfe"), "mae": a.get("mae"), "time_to_mfe": a.get("time_to_mfe"), "time_to_invalidation": a.get("time_to_invalidation")}),
        "Structural_Risk": _kv_rows({"risk_bps": a.get("risk_bps"), "mfe_over_risk": a.get("mfe_over_risk")}),
        "Concentration": _kv_rows(report.get("concentration") or {}),
        "D1": block_rows.get("D1") or [{"key": "empty"}],
        "D2": block_rows.get("D2") or [{"key": "empty"}],
        "D3": block_rows.get("D3") or [{"key": "empty"}],
        "D4": block_rows.get("D4") or [{"key": "empty"}],
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# SMA5_25_75_TREND_PULLBACK_PLAYBOOK_DISCOVERY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "TREND → PULLBACK → HOLD/RECLAIM → REACCELERATION on native 1-minute SMA5/25/75.",
            "Not a period search. Not a peer-propagation rescue. Not a Complete Strategy.",
            "",
            f"setup_n? **{a.get('setup_n?')}** bull/bear? **{a.get('bull / bear n?')}**",
            f"D1/D2/D3/D4? **{a.get('D1/D2/D3/D4?')}**",
            f"5/25/75 adds to same 1m trigger? **{a.get('Does 5/25/75 aligned structure add to the same 1m trigger?')}**",
            f"SMA25 pullback matters? **{a.get('Does SMA25 pullback matter?')}**",
            f"VWAP adds? **{a.get('Does VWAP add inside the setup?')}** S/R adds? **{a.get('Does S/R add inside the setup?')}**",
            f"participation contraction→expansion adds? **{a.get('Does participation contraction→expansion add?')}**",
            f"next-open MFE/MAE? **{a.get('What is executable next-open MFE/MAE?')}**",
            f"MFE/risk? **{a.get('What is structural-risk-normalized MFE?')}**",
            f"move before entry? **{a.get('How much move occurs before executable entry?')}**",
            f"A vs B 10m D2/D3/D4? **{a.get('A vs B D2/D3/D4 10m gap?')}**",
            f"future setup selection? **{a.get('Any future setup selection?')}** MA period opt? **{a.get('Any MA period optimization?')}** PnL tune? **{a.get('Any threshold PnL tuning?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}**",
            f"submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
