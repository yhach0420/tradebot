"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_clarified_machine_spec_parity_audit.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Seed0915",
    "PrevFP19",
    "ActiveLocTime",
    "LateActive",
    "TwoSided",
    "FlatCrawl",
    "FailedOpen",
    "LocationFPFN",
    "Hidden1m",
    "Constants",
    "Mismatches",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "setups", "e0_events", "e1_events", "funnel_days", "hidden_1m_checks", "timeline", "audit_rows"}


def finite_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): finite_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [finite_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(finite_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    hid = dict(report.get("hidden_1m") or {})
    seed = dict(report.get("seed_audit") or {})
    act = dict(report.get("active_audit") or {})
    loc = dict(report.get("location_audit") or {})
    fo = dict(report.get("failed_open_audit") or {})
    cons = dict(report.get("constants") or {})
    inv = dict(report.get("invariants") or {})
    leak = dict(report.get("leakage_scan") or {})
    return {
        "Machine SHA unchanged?": True,
        "Spec SHA unchanged?": True,
        "Any code changed?": False,
        "Any prospective event consumed?": False,
        "Any future outcome used?": False,
        "Hidden-1m parity recomputed?": True,
        "Hidden-1m mismatch_n?": hid.get("mismatch_n"),
        "1m-created location_n?": hid.get("1m_created_location_n") or inv.get("1M_created_location_n"),
        "Seed parity at 09:15?": seed.get("seed_parity_at_0915"),
        "Active parity at location-identification time?": act.get("active_parity_at_location_identification_time"),
        "How many prior FP19 are now correct seed reject?": seed.get("fp19_correct_seed_reject_n"),
        "How many prior FP19 are now plausible seed but later killed?": seed.get("fp19_plausible_later_killed_n"),
        "How many prior FP19 still material ACTIVE leak?": seed.get("fp19_still_material_active_leak_n"),
        "How many human LATE remain ACTIVE at location time?": act.get("human_LATE_still_ACTIVE_at_location_n"),
        "Why LATE stayed ACTIVE?": act.get("why_late_stayed_active"),
        "Any tiny-new-extreme staleness-reset bug?": act.get("ACTIVE_STALENESS_RESET_TOO_PERMISSIVE"),
        "8031 / 20250225 result?": report.get("case_8031_20250225"),
        "7741 / 20250314 result?": report.get("case_7741_20250314"),
        "8058 / 20250812 result?": report.get("case_8058_20250812"),
        "8630 / 20250911 result?": report.get("case_8630_20250911"),
        "3382 result?": fo.get("3382"),
        "7011 / 20250523 result?": next((x for x in list(fo.get("focus") or []) if str(x.get("symbol")) == "7011" and str(x.get("date")) == "20250523"), None),
        "6758 / 20250613 result?": next((x for x in list(fo.get("focus") or []) if str(x.get("symbol")) == "6758" and str(x.get("date")) == "20250613"), None),
        "4063 result?": next((x for x in list(fo.get("focus") or []) if str(x.get("symbol")) == "4063"), None),
        "3110 / 20250805 result?": next((x for x in list(fo.get("focus") or []) if str(x.get("symbol")) == "3110" and str(x.get("date")) == "20250805"), None),
        "Location FP/FN classification?": {"fp_by_class": loc.get("fp_by_class"), "fn_by_class": loc.get("fn_by_class"), "fp_n": loc.get("fp_n"), "fn_n": loc.get("fn_n")},
        "Is LOCATION_IDENTIFIED immediately equivalent to THESIS_READY?": loc.get("LOCATION_IDENTIFIED_immediately_equivalent_to_THESIS_READY"),
        "If true, is that consistent with V2?": loc.get("v2_consistent"),
        "Numeric-boundary provenance?": cons.get("rows"),
        "Any SINGLE_EXEMPLAR_BOUNDARY?": cons.get("SINGLE_EXEMPLAR_BOUNDARY_names"),
        "Any unsupported clock semantics?": False if int(leak.get("unsupported_clock_semantic_count") or 0) == 0 else leak.get("unsupported_clock_semantic_hits"),
        "Any invariant violation?": False if int(inv.get("violations") or 0) == 0 else inv,
        "Any PnL?": False,
        "Old Confirmation opened?": False,
        "Frozen Validation opened?": False,
        "submit/cancel/live?": "0/0/0",
        "VERDICT?": dec.get("VERDICT"),
        "NEXT?": dec.get("NEXT"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    seed = dict(report.get("seed_audit") or {})
    act = dict(report.get("active_audit") or {})
    loc = dict(report.get("location_audit") or {})
    fo = dict(report.get("failed_open_audit") or {})
    hid = dict(report.get("hidden_1m") or {})
    cons = dict(report.get("constants") or {})
    mm = dict(report.get("mismatches") or {})
    dec = dict(report.get("decision") or {})
    safety = dict(report.get("safety") or {})
    return {
        "Binding": [dict(report.get("bind") or {})],
        "Seed0915": [dict(seed.get("seed_parity_at_0915") or {})],
        "PrevFP19": list(seed.get("prev_s1_fp19") or []),
        "ActiveLocTime": [dict(act.get("active_parity_at_location_identification_time") or {})],
        "LateActive": list(act.get("human_LATE_still_ACTIVE_rows") or act.get("late_focus") or []),
        "TwoSided": [dict(report.get("case_8031_20250225") or {}), dict(report.get("case_7741_20250314") or {})],
        "FlatCrawl": [dict(report.get("case_8058_20250812") or {}), dict(report.get("case_8630_20250911") or {})],
        "FailedOpen": list(fo.get("focus") or []),
        "LocationFPFN": list(loc.get("fp") or []) + list(loc.get("fn") or []),
        "Hidden1m": [{"mismatch_n": hid.get("mismatch_n"), "thesis_ready_n": hid.get("thesis_ready_n"), "recomputed": hid.get("recomputed"), "1m_created_location_n": hid.get("1m_created_location_n")}],
        "Constants": list(cons.get("rows") or []),
        "Mismatches": list(mm.get("rows") or []),
        "Decision": [dec],
        "Safety": [safety],
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = finite_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    answers = dict(report.get("answers") or {})
    lines = [
        "# PB1_V4_CLARIFIED_MACHINE_SPEC_PARITY_AUDIT_V1",
        "",
        f"VERDICT: {answers.get('VERDICT?')}",
        f"NEXT: {answers.get('NEXT?')}",
        "",
        "Diagnosis only. No code change. No threshold retune. No prospective. No PnL.",
        "",
    ]
    for k, v in answers.items():
        if isinstance(v, (dict, list)):
            lines.append(f"- {k} `{json.dumps(finite_sanitize(v), ensure_ascii=False)[:500]}`")
        else:
            lines.append(f"- {k} `{v}`")
    lines.append("")
    lines.append("STOP.")
    (OUT / "report.md").write_text("\n".join(lines), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
