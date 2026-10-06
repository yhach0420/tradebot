"""Write report.json / report.md / audit.xlsx only. No mass CSV. No Runtime write."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.target_price_contract_v4 import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "target_price_contract_v4"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)

SHEET_ORDER = (
    "Summary",
    "Inventory",
    "Code_Semantics",
    "Intervals",
    "Next_Event",
    "Comparison",
    "Waterfall",
    "First_Fail",
    "By_Anchor",
    "By_Day",
    "By_Symbol",
    "Missingness_SMD",
    "Score_Decile",
    "Cohort",
    "Label_Stability",
    "Mark_Age",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def flatten_smd(smd_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for r in smd_rows:
        rec = {"feature": r.get("feature"), "smd_valid_minus_null": r.get("smd_valid_minus_null"), "flag": r.get("flag")}
        for side in ("valid", "null"):
            q = r.get(side) or {}
            for k, v in q.items():
                rec[f"{side}_{k}"] = v
        out.append(rec)
    return out


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("gates") or {}
    return "\n".join(
        [
            "# TARGET PRICE CONTRACT V4",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE PERSISTENT MARKET STATE VALIDITY AUDIT ONLY.",
            "C14 / Runtime / CLOCK / ENTRY / EXIT / CAP / re-entry / Dual Lane SESSION_CLOSE unchanged.",
            "Corrected Fill SoT unchanged: is_executable_continuous_board, WAIT_SEC=1.0, fill_price=limit_price.",
            "Execution freshness unchanged (BOARD_FRESHNESS_SEC_V1R=5.0).",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "C REBUILD was not started this run. B research was not restarted. Model search was not started.",
            "Age search was not expanded beyond 60s. New mark-source hunt was not started.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"MARK_STATE_PERSISTS_UNTIL_NEXT_EVENT: {g.get('MARK_STATE_PERSISTS_UNTIL_NEXT_EVENT')}",
            f"ARBITRARY_MARK_AGE_CUTOFF_REQUIRED: {g.get('ARBITRARY_MARK_AGE_CUTOFF_REQUIRED')}",
            f"M4_PRIMARY_ROWS: {g.get('M4_PRIMARY_ROWS')}",
            f"M4_COVERAGE_RATE: {g.get('M4_COVERAGE_RATE')}",
            f"M4_MEDIAN_COHORT: {g.get('M4_MEDIAN_COHORT')}",
            f"M4_P10_COHORT: {g.get('M4_P10_COHORT')}",
            f"M4_EVENT_RATE_SMD: {g.get('M4_EVENT_RATE_SMD')}",
            f"M4_SCORE_DECILE_COVERAGE_RANGE: {g.get('M4_SCORE_DECILE_COVERAGE_RANGE')}",
            f"M4_SYMBOL_COVERAGE_RANGE: {g.get('M4_SYMBOL_COVERAGE_RANGE')}",
            f"M4_TARGET_MISSINGNESS_SELECTION_BIAS: {g.get('M4_TARGET_MISSINGNESS_SELECTION_BIAS')}",
            f"M4_T0_MARK_AGE_P90: {g.get('M4_T0_MARK_AGE_P90')}",
            f"M4_T600_MARK_AGE_P90: {g.get('M4_T600_MARK_AGE_P90')}",
            f"COMMON_5S_M4_LABEL_SPEARMAN: {g.get('COMMON_5S_M4_LABEL_SPEARMAN')}",
            f"TARGET_V4_READY_FOR_C_REBUILD: {g.get('TARGET_V4_READY_FOR_C_REBUILD')}",
            f"VERDICT: {g.get('VERDICT')}",
            f"RECOMMENDED_NEXT_STEP: {g.get('RECOMMENDED_NEXT_STEP')}",
            "",
            "## Core question",
            "",
            "kabu PUSH / Capture board state is last-observed until the next event (A),",
            "not unknown after a fixed number of seconds (B).",
            "BOARD_FRESHNESS_SEC=5 was not reused as a historical mark TTL.",
            "Execution freshness remains 5s and was not changed.",
            "",
            "## Persistence evidence",
            "",
            "Code: ingest_push keeps the last row until the next append. No timer zeros the book.",
            "fresh_sec is quote-clock vs recv on the same PUSH, not inter-event TTL.",
            "Global sequence +1 holes are accepted-universe global seq, not per-symbol Capture gaps.",
            "Quiet symbol time is not classified as CAPTURE_GAP.",
            "",
            "## STOP",
            "",
            "Audit complete. C REBUILD not started. Runtime not changed. Execution freshness not changed.",
            "",
        ]
    )
