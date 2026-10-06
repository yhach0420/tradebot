"""Write report.json / report.md / audit.xlsx / research_pool_manifest.json / minute_panel_manifest.json."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.daytrade_historical_research_foundation_v2 import ANALYSIS_ID
from research.daytrade_historical_research_foundation_v2.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "SUMMARY",
    "RESEARCH_POOL",
    "ELIGIBILITY",
    "MINUTE_COVERAGE",
    "CONTEXT",
    "MARKET_STATE",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    panel = dict(report.get("panel") or {})
    return {
        "SUMMARY": _kv(dict(report.get("answers") or {})) + _kv(dict(report.get("decision") or {})),
        "RESEARCH_POOL": list(report.get("pool_union_rows") or []),
        "ELIGIBILITY": list(report.get("eligibility") or []),
        "MINUTE_COVERAGE": list(panel.get("coverage") or []) or _kv(panel),
        "CONTEXT": list(report.get("context_only_proposal") or []) or _kv(dict(report.get("runtime_draft_not_frozen") or {})),
        "MARKET_STATE": _kv(dict(report.get("reference_market_state") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            "## Required answers",
            "",
            f"Minute entitlement active? **{a.get('minute_entitlement_active')}**",
            f"Historical candidate N before ranking? **{a.get('historical_candidate_n_before_ranking')}**",
            f"Liquidity top80 N? **{a.get('liquidity_top80_n')}**",
            f"Activity top40 N? **{a.get('activity_top40_n')}**",
            f"Union Research Pool N? **{a.get('union_research_pool_n')}**",
            f"Actual history first/last? {a.get('actual_history_first_last')}",
            f"Minute Time semantics? **{a.get('minute_time_semantics')}**",
            f"Tick cross-check pass? **{a.get('tick_crosscheck_pass')}**",
            f"Per-symbol minute coverage? {a.get('per_symbol_minute_coverage')}",
            f"Trade eligibility metrics complete? **{a.get('trade_eligibility_metrics_complete')}**",
            f"Any actual Bid/Ask historical data available? **{a.get('bid_ask_historical_available')}**",
            f"Reference market state buildable? **{a.get('reference_market_state_buildable')}**",
            f"Strategy search started? **{a.get('strategy_search_started')}**",
            f"PnL used? **{a.get('pnl_used')}**",
            f"Future return used? **{a.get('future_return_used')}**",
            f"Old V1 modified? **{a.get('old_v1_modified')}**",
            f"Runtime modified? **{a.get('runtime_modified')}**",
            f"Paper modified? **{a.get('paper_modified')}**",
            f"20260914 live modified? **{a.get('live_20260914_modified')}**",
            f"submit/cancel/live? **{a.get('submit_cancel_live')}**",
            f"V1.1 adopted? **{a.get('v1_1_adopted')}**",
            f"PANEL_CONDITIONED_AS_OF_202609? **{a.get('panel_conditioned_as_of_202609')}**",
            f"VERDICT? **{a.get('VERDICT')}**",
            f"NEXT? **{a.get('NEXT')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    pool_manifest = {
        "panel_conditioned_as_of_202609": True,
        "claim_selectable_in_2025": False,
        "v1_modified": False,
        "v1_1_adopted": False,
        "pool": dict(report.get("pool") or {}),
        "symbols": [r.get("symbol") for r in list(report.get("pool_union_rows") or [])],
    }
    (OUT / "research_pool_manifest.json").write_text(json.dumps(json_sanitize(pool_manifest), ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "minute_panel_manifest.json").write_text(
        json.dumps(json_sanitize(dict(report.get("panel") or {"ingested": False})), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
    assert (OUT / "audit.xlsx").is_file()
    assert (OUT / "research_pool_manifest.json").is_file()
    assert (OUT / "minute_panel_manifest.json").is_file()
