"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.r14_trend_incrementality_rca_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Feature_Semantics",
    "Direction_Audit",
    "Base",
    "T15",
    "T15_P5",
    "R14_Full",
    "Nested_Contrasts",
    "Matched_R14",
    "Matched_T15",
    "Overlap_Weighted",
    "R15_Dose_Response",
    "Threshold_Locality",
    "D1_Stability",
    "D2_D3_D4",
    "MFE_MAE",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def _rule_rows(pack: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for b in ("D1", "D2", "D3", "D4"):
        st = dict((pack or {}).get(b) or {})
        rows.append({"block": b, **st})
    return rows or [{"block": None}]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    nested = dict(report.get("nested") or {})
    d24 = []
    for name in ("BASE", "T15", "T15_P5", "R14_FULL"):
        for b in ("D2", "D3", "D4"):
            st = dict((nested.get(name) or {}).get(b) or {})
            d24.append({"rule": name, "block": b, **st})
    if not d24:
        d24 = [{"rule": None}]
    mfe = []
    for name in ("BASE", "T15", "T15_P5", "R14_FULL"):
        for b in ("D2", "D3", "D4"):
            st = dict((nested.get(name) or {}).get(b) or {})
            mfe.append({"rule": name, "block": b, "median_mfe": st.get("median_mfe"), "median_mae": st.get("median_mae")})
    if not mfe:
        mfe = [{"rule": None}]
    nest = []
    for k, v in dict(report.get("nested_contrasts") or {}).items():
        nest.append({"contrast": k, **dict(v or {})})
    if not nest:
        nest = [{"contrast": None}]
    dose_rows = []
    pooled = dict((report.get("r15_dose_response") or {}).get("d2_d4_pooled") or {}).get("cells") or {}
    for q, st in pooled.items():
        dose_rows.append({"quintile": q, **dict(st or {})})
    if not dose_rows:
        dose_rows = [{"quintile": None}]
    return {
        "Binding": _kv_rows({"ok": (report.get("bind") or {}).get("ok"), "current_judgment": (report.get("bind") or {}).get("current_judgment"), "TREE_SHA256": (report.get("bind") or {}).get("TREE_SHA256"), "EVENT_GENERATOR_SHA256": (report.get("bind") or {}).get("EVENT_GENERATOR_SHA256")}),
        "Feature_Semantics": _kv_rows(report.get("feature_semantics") or {}),
        "Direction_Audit": _kv_rows(report.get("direction_audit") or {}),
        "Base": _rule_rows(nested.get("BASE") or {}),
        "T15": _rule_rows(nested.get("T15") or {}),
        "T15_P5": _rule_rows(nested.get("T15_P5") or {}),
        "R14_Full": _rule_rows(nested.get("R14_FULL") or {}),
        "Nested_Contrasts": nest,
        "Matched_R14": _kv_rows(report.get("matched_r14") or {}),
        "Matched_T15": _kv_rows(report.get("matched_t15") or {}),
        "Overlap_Weighted": _kv_rows({"T15": report.get("overlap_t15"), "R14": report.get("overlap_r14")}),
        "R15_Dose_Response": dose_rows,
        "Threshold_Locality": _kv_rows(report.get("threshold_locality") or {}),
        "D1_Stability": _kv_rows(report.get("d1_stability") or {}),
        "D2_D3_D4": d24,
        "MFE_MAE": mfe,
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# R14_TREND_INCREMENTALITY_RCA_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Parent interpretation downgraded to NATIVE_CONTEXT_STACK_PARTIAL_MECHANISM_V1.",
            "Frozen R14. No threshold retune. No complete strategy.",
            "",
            f"r15 direction-normalized? **{a.get('Is r15 direction-normalized?')}** prior5_ret? **{a.get('Is prior5_ret direction-normalized?')}**",
            f"Bullish R14? **{a.get('Bullish R14 effect?')}**",
            f"Bearish R14? **{a.get('Bearish R14 effect?')}**",
            f"T15 D2/D3/D4 gaps? **{a.get('T15 D2 gap?')}** / **{a.get('T15 D3 gap?')}** / **{a.get('T15 D4 gap?')}**",
            f"R14 D2/D3/D4 gaps? **{a.get('R14_FULL D2 gap?')}** / **{a.get('R14_FULL D3 gap?')}** / **{a.get('R14_FULL D4 gap?')}**",
            f"prior5_ret adds? **{a.get('Does prior5_ret add beyond T15?')}** range adds? **{a.get('Does prior5_range_rel add beyond T15_P5?')}** FULL vs T15? **{a.get('Does FULL add beyond T15?')}**",
            f"Matched T15 p40 gap? **{a.get('Matched T15 p40 gap?')}** Matched R14? **{a.get('Matched R14 p40 gap?')}**",
            f"Overlap T15? **{a.get('Overlap-weighted T15 gap?')}** Overlap R14? **{a.get('Overlap-weighted R14 gap?')}**",
            f"dose-response monotonic? **{a.get('Is r15 dose-response monotonic?')}** threshold? **{a.get('Does the exact D1 threshold look like a state transition or smooth continuation?')}**",
            f"D1 stability? **{a.get('D1 stability classification?')}**",
            f"new feature? **{a.get('Any new feature added?')}** retune? **{a.get('Any threshold retuned?')}** PnL? **{a.get('Any PnL optimization?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}** submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
