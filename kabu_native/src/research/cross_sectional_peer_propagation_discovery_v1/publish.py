"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.cross_sectional_peer_propagation_discovery_v1.features import FEATURE_DEFINITIONS
from research.cross_sectional_peer_propagation_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Peer_Universe",
    "Peer_Exclusion",
    "Feature_Definitions",
    "D1_State_Boundaries",
    "P1",
    "P2",
    "P3",
    "True_Lead",
    "Simultaneous",
    "Matched_Control",
    "D2",
    "D3",
    "D4",
    "Offset_Map",
    "Sector_vs_Market",
    "MFE_MAE",
    "Barrier_Races",
    "SR_Diagnostic",
    "TV_Diagnostic",
    "Concentration",
    "Causality_Audit",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def _mech_block(report: dict[str, Any], block: str) -> list[dict[str, Any]]:
    out = []
    for m, pack in dict(report.get("mechanisms") or {}).items():
        st = (((pack.get("true_lead") or {}).get("counts") or {}).get(block) or {})
        g = ((((pack.get("true_lead") or {}).get("matched") or {}).get("blocks") or {}).get(block) or {})
        out.append({"mech": m, **st, "matched_gap_10m": g.get("gap_10m"), "matched_n": g.get("matched_n")})
    return out or [{"mech": None}]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    freeze = dict(report.get("freeze") or {})
    mech = dict(report.get("mechanisms") or {})
    uni = dict(report.get("peer_universe") or {})
    feat = [{"feature": k, "definition": v} for k, v in FEATURE_DEFINITIONS.items()]
    p_rows = {}
    true_rows = []
    sim_rows = []
    match_rows = []
    off_rows = []
    svm_rows = []
    mfe_rows = []
    bar_rows = []
    sr_rows = []
    tv_rows = []
    conc_rows = []
    for m, pack in mech.items():
        tl = dict(pack.get("true_lead") or {})
        mc = dict(tl.get("matched") or {})
        p_rows[m] = _kv_rows(
            {
                "definition": pack.get("definition"),
                "survives": pack.get("survives"),
                "d1": (tl.get("counts") or {}).get("D1"),
                "d2_d4": tl.get("d2_d4"),
                "matched": {k: mc.get(k) for k in ("treated_n", "matched_n", "match_rate", "gap_10m", "all_eval_positive", "all_eval_coherent", "positive_blocks")},
            }
        )
        true_rows.append({"mech": m, **dict(tl.get("d2_d4") or {})})
        sim_rows.append({"mech": m, **dict((pack.get("simultaneous") or {}).get("d2_d4") or {})})
        for b in ("D2", "D3", "D4"):
            st = dict((mc.get("blocks") or {}).get(b) or {})
            match_rows.append({"mech": m, "block": b, **st})
            c = dict((tl.get("counts") or {}).get(b) or {})
            mfe_rows.append({"mech": m, "block": b, "mfe": c.get("median_mfe"), "mae": c.get("median_mae")})
            bar_rows.append({"mech": m, "block": b, "p20": c.get("p20_before_m20"), "p40": c.get("p40_before_m20"), "p80": c.get("p80_before_m30")})
        for row in list(pack.get("offset_map") or []) or [{"offset": None, "mech": m}]:
            off_rows.append({"mech": m, **dict(row)})
        svm_rows.append({"mech": m, **dict(pack.get("sector_vs_market") or {})})
        for row in list(pack.get("sr_diagnostic") or []) or [{"sr_bin": None}]:
            sr_rows.append({"mech": m, **dict(row)})
        tv_rows.append({"mech": m, **dict(pack.get("tv_diagnostic") or {})})
        conc_rows.append({"mech": m, **{k: json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v for k, v in dict(pack.get("concentration") or {}).items()}})
    return {
        "Binding": _kv_rows({"ok": (report.get("bind") or {}).get("ok"), "VERDICT": d.get("VERDICT"), "NEXT": d.get("NEXT"), "parent": (report.get("bind") or {}).get("parent_verdict")}),
        "Peer_Universe": list(uni.get("rows") or [{"target": None}]),
        "Peer_Exclusion": _kv_rows({"TARGET_INCLUDED_IN_PEER_METRIC_N": report.get("TARGET_INCLUDED_IN_PEER_METRIC_N"), "future_peer_information_n": 0, "performance_selected": False}),
        "Feature_Definitions": feat,
        "D1_State_Boundaries": _kv_rows({k: freeze.get(k) for k in ("ok", "high", "low", "n", "FREEZE_SHA256", "p_definitions", "d1_only", "d2_not_read")}),
        "P1": p_rows.get("P1") or [{"key": "empty"}],
        "P2": p_rows.get("P2") or [{"key": "empty"}],
        "P3": p_rows.get("P3") or [{"key": "empty"}],
        "True_Lead": true_rows or [{"mech": None}],
        "Simultaneous": sim_rows or [{"mech": None}],
        "Matched_Control": match_rows or [{"mech": None}],
        "D2": _mech_block(report, "D2"),
        "D3": _mech_block(report, "D3"),
        "D4": _mech_block(report, "D4"),
        "Offset_Map": off_rows or [{"mech": None}],
        "Sector_vs_Market": svm_rows or [{"mech": None}],
        "MFE_MAE": mfe_rows or [{"mech": None}],
        "Barrier_Races": bar_rows or [{"mech": None}],
        "SR_Diagnostic": sr_rows or [{"mech": None}],
        "TV_Diagnostic": tv_rows or [{"mech": None}],
        "Concentration": conc_rows or [{"mech": None}],
        "Causality_Audit": _kv_rows(
            {
                "TARGET_INCLUDED_IN_PEER_METRIC_N": report.get("TARGET_INCLUDED_IN_PEER_METRIC_N"),
                "future_peer_information_n": 0,
                "same_bar_outcome_n": report.get("same_bar_outcome_n"),
                "d1_frozen_before_d2": True,
                "displacement_used": False,
                "lunch_counted_as_trading_minutes": False,
            }
        ),
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# CROSS_SECTIONAL_PEER_PROPAGATION_DISCOVERY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Sector-peer lead/follow on native 1-minute panel. Not a strategy.",
            "Direction-aligned compact stack closed. R13/R11/R14 not repaired.",
            "",
            f"Target excluded from every peer metric? **{a.get('Target excluded from every peer metric?')}**",
            f"future peer information? **{a.get('Any future peer information?')}** same-bar outcome? **{a.get('Any same-bar outcome included?')}**",
            f"P1/P2/P3? **{a.get('P1/P2/P3 definitions?')}**",
            f"D1 freeze before D2-D4? **{a.get('D1 boundaries frozen before D2-D4?')}**",
            f"P1? **{a.get('P1')}**",
            f"P2? **{a.get('P2')}**",
            f"P3? **{a.get('P3')}**",
            f"matched D2/D3/D4? **{a.get('Any mechanism positive against matched same-symbol controls in D2/D3/D4?')}**",
            f"true lead vs simultaneous? **{a.get('True peer lead or simultaneous movement?')}**",
            f"offset map? **{a.get('What does offset map show?')}**",
            f"sector vs market? **{a.get('Sector-specific or broad-market effect?')}**",
            f"S/R diagnostic? **{a.get('Does S/R change the result diagnostically?')}**",
            f"target TV diagnostic? **{a.get('Does target participation change it diagnostically?')}**",
            f"one-symbol? **{a.get('Any one-symbol dominance?')}** one-sector? **{a.get('Any one-sector dominance?')}** one-day? **{a.get('Any one-day dominance?')}**",
            f"threshold retune? **{a.get('Any threshold retune?')}** PnL? **{a.get('Any PnL optimization?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}** Kabu50? **{a.get('Kabu50?')}**",
            f"submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
