"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.native_1m_path_state_strategy_freeze_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Frozen_R11",
    "Episode_Generator",
    "Episode_Causality",
    "Feature_Semantics",
    "VWAP_Semantics",
    "Reclaim_Semantics",
    "Timestamp_Lineage",
    "Entry_Execution",
    "Exit_Execution",
    "Portfolio_State",
    "Source_Replay",
    "Frozen_Replay",
    "Trade_Identity_908",
    "EventGated_Canary",
    "Manual_Lineage_30",
    "Robustness",
    "Symbol_Contribution",
    "Sector_Contribution",
    "Manifest",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _econ(pack: dict[str, Any]) -> dict[str, Any]:
    return {
        "trade_n": pack.get("trade_n"),
        "day_n": pack.get("day_n"),
        "symbol_n": pack.get("symbol_n"),
        "mean_x0_bps": pack.get("mean_x0_bps"),
        "mean_x1_bps": pack.get("mean_x1_bps"),
        "profit_factor": pack.get("profit_factor"),
        "hit_rate": pack.get("hit_rate"),
        "candidate_n": pack.get("candidate_n"),
        "skipped": pack.get("skipped"),
        "block_mean_x0": pack.get("block_mean_x0"),
        "d2_d3": pack.get("d2_d3"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    a = dict(report.get("answers") or {})
    feat = dict(report.get("feature_semantics") or {})
    lin30 = dict(report.get("manual_lineage_30") or {})
    fr = dict(report.get("frozen_replay") or {})
    src = dict(report.get("source_replay") or {})
    ident_rows = list(fr.get("eval_trades") or [])
    lineage_rows = []
    for r in list(lin30.get("rows") or []):
        ev = dict(r.get("R11_evaluation") or {})
        ts = dict(r.get("timestamp_lineage") or {})
        lineage_rows.append(
            {
                "block": r.get("block"),
                "date": r.get("date"),
                "symbol": r.get("symbol"),
                "reclaim_bar": r.get("reclaim_bar"),
                "available_at": r.get("feature_available_at"),
                "dist_vwap": ev.get("dist_vwap"),
                "mins_from_open": ev.get("mins_from_open"),
                "vwap_reclaim": ev.get("vwap_reclaim"),
                "match": ev.get("match"),
                "entry_fill_time": r.get("ENTRY_fill_time"),
                "entry_fill_price": r.get("ENTRY_fill_price"),
                "exit_trigger": r.get("EXIT_trigger"),
                "exit_fill_time": r.get("EXIT_fill_time"),
                "exit_fill_price": r.get("EXIT_fill_price"),
                "available_at_le_decision": ts.get("available_at_le_decision"),
                "prior_vwap_state": r.get("prior_vwap_state"),
                "raw_bars": r.get("raw_minute_bars_before_and_reclaim"),
                "subsequent": r.get("subsequent_bars_to_exit"),
            }
        )
    return {
        "Binding": _kv(
            {
                "parent": report.get("parent_verdict_accepted"),
                "strategy_id": report.get("strategy_id"),
                "promoted": False,
                "frozen_validation_opened": False,
                "old_confirmation_opened": False,
                "design_evidence_not_certification": True,
            }
        ),
        "Frozen_R11": _kv(dict((report.get("human_spec") or {}).get("r11_human") or {}))
        + [{"feature": p.get("feature"), "op": p.get("op"), "threshold": p.get("threshold")} for p in list((report.get("human_spec") or {}).get("r11_machine") or [])],
        "Episode_Generator": _kv(dict(report.get("episode_generator") or {})),
        "Episode_Causality": _kv(dict(report.get("episode_causality") or {})),
        "Feature_Semantics": _kv(feat),
        "VWAP_Semantics": _kv(dict(feat.get("dist_vwap") or {})),
        "Reclaim_Semantics": _kv(dict(feat.get("vwap_reclaim") or {})),
        "Timestamp_Lineage": _kv(dict(report.get("timestamp_lineage_audit") or {})),
        "Entry_Execution": _kv({"next_bar_open": True, "same_bar_fill": False, "BAR_START": True, "entry_cutoff": "14:50"}),
        "Exit_Execution": _kv({"kind": "reclaim", "vwap_loss": True, "session_flat": "15:20", "time_stop_min": 20, "v27": False}),
        "Portfolio_State": _kv({"CAP": 3, "same_symbol_day": True, "x1_tax_bps": 8.0, **dict(fr.get("skipped") or {})}),
        "Source_Replay": [_econ(src)],
        "Frozen_Replay": [_econ(fr)],
        "Trade_Identity_908": ident_rows or [{"empty": True}],
        "EventGated_Canary": _kv(dict(report.get("event_gated_canary") or {})),
        "Manual_Lineage_30": lineage_rows or [{"empty": True}],
        "Robustness": _kv({k: v for k, v in dict(report.get("robustness") or {}).items() if k not in {"time_of_day", "sector_contribution", "tail"}}),
        "Symbol_Contribution": list(report.get("symbol_contribution") or [{"empty": True}]),
        "Sector_Contribution": list(report.get("sector_contribution") or [{"empty": True}]),
        "Manifest": _kv(dict(report.get("manifest") or {})),
        "Safety": _kv(a) + _kv(dict(report.get("decision") or {})) + _kv({"submit_cancel_live": "0/0/0", "v27_bolted": False, "promoted": False}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    man = dict(report.get("manifest") or {})
    return "\n".join(
        [
            "# NATIVE_1M_PATH_STATE_STRATEGY_FREEZE_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            d.get("INTERPRETATION") or "",
            "",
            f"Frozen strategy ID: **{a.get('Exact_frozen_strategy_ID')}**",
            f"R11 predicates frozen exactly? **{a.get('R11_predicates_frozen_exactly')}**",
            f"Machine dist_vwap threshold: `{a.get('Machine_threshold_for_dist_vwap')}`",
            f"Episode generator frozen? **{a.get('Episode_generator_frozen')}**",
            f"Episode generator hash: `{a.get('Episode_generator_hash')}`",
            f"Future information in episode-start construction? **{a.get('Any_future_information_in_episode_start_construction')}**",
            f"R11 depends on future path labels at runtime? **{a.get('Does_R11_depend_on_future_path_labels_at_runtime')}**",
            f"Runtime candidate generation equals research? **{a.get('Does_runtime_candidate_generation_equal_research_candidate_generation')}**",
            f"Every-minute scanning a different strategy? **{a.get('Would_every_minute_scanning_be_a_different_strategy')}**",
            f"Feature available_at causal? **{a.get('Feature_available_at_causal')}**",
            f"Any same-bar execution? **{a.get('Any_same_bar_execution')}**",
            f"VWAP formula frozen? **{a.get('VWAP_formula_frozen')}**",
            f"VWAP reclaim semantics frozen? **{a.get('VWAP_reclaim_semantics_frozen')}**",
            f"EXIT semantics frozen? **{a.get('EXIT_semantics_frozen')}**",
            f"CAP / occupancy frozen? **{a.get('CAP_occupancy_frozen')}**",
            f"908/908 trade identity? **{a.get('trade_identity_908_908')}** `{a.get('TRADE_IDENTITY_MATCH')}`",
            f"X0 exact parity? **{a.get('X0_exact_parity')}** `{a.get('X0')}`",
            f"X1 exact parity? **{a.get('X1_exact_parity')}** `{a.get('X1')}`",
            f"D2 exact? **{a.get('D2_exact')}** D3? **{a.get('D3_exact')}** D4? **{a.get('D4_exact')}**",
            f"D2+D3 exact? **{a.get('D2_D3_exact')}**",
            f"Top-day / top-symbol robustness: `{a.get('Top_day_top_symbol_robustness')}`",
            f"Any parameter changed? **{a.get('Any_parameter_changed')}**",
            f"Any new feature added? **{a.get('Any_new_feature_added')}**",
            f"BG_CONT_VWAP reopened? **{a.get('BG_CONT_VWAP_reopened')}**",
            f"Frozen Validation opened? **{a.get('Frozen_Validation_opened')}**",
            f"Old Confirmation opened? **{a.get('Old_Confirmation_opened')}**",
            f"New paid data? **{a.get('New_paid_data')}**",
            f"Kabu50? **{a.get('Kabu50')}**",
            f"submit/cancel/live: **{a.get('submit_cancel_live')}**",
            "",
            f"ENTRY_EVENT_SOURCE_ID: `{a.get('ENTRY_EVENT_SOURCE_ID')}`",
            f"complete_strategy_hash: `{man.get('complete_strategy_hash')}`",
            f"Canary Jaccard: `{a.get('canary_jaccard')}` different=`{a.get('canary_different')}`",
            f"FUTURE_IN_EPISODE_START_N: **{a.get('FUTURE_IN_EPISODE_START_N')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
