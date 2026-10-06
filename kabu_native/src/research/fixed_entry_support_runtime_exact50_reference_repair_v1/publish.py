"""Publish the exact-50 reference repair and the paired Frozen V2 comparison."""
from __future__ import annotations

import json
import pickle
from pathlib import Path
from typing import Any, Optional

import numpy as np
from openpyxl import Workbook

from research.event_time_impulse_complete_strategy_v2.scan import _dates
from research.fixed_entry_support_runtime_exact50_reference_repair_v1.scan import scan
from research.symbol_setup_baseline_complete_strategy_precommit.contract import EXTENSION17, ORIGINAL18

OUT = Path("results/research/fixed_entry_support_runtime_exact50_reference_repair_v1")


def _pf(rows: list[dict[str, Any]]) -> Optional[float]:
    loss = sum(-float(row["pnl_yen"]) for row in rows if float(row["pnl_yen"]) < 0)
    gain = sum(float(row["pnl_yen"]) for row in rows if float(row["pnl_yen"]) > 0)
    return gain / loss if loss else None


def _realized_dd(rows: list[dict[str, Any]]) -> float:
    equity = peak = worst = 0.0
    for row in sorted(rows, key=lambda item: (float(item["exit_t"]), str(item["symbol"]))):
        equity += float(row["pnl_yen"])
        peak = max(peak, equity)
        worst = min(worst, equity - peak)
    return worst


def _stats(rows: list[dict[str, Any]], mtm: Optional[float] = None) -> dict[str, Any]:
    pnl = [float(row["pnl_yen"]) for row in rows]
    bps = [float(row["bps"]) for row in rows if row.get("bps") is not None]
    hold = [float(row["hold_sec"]) for row in rows if row.get("hold_sec") is not None]
    gain = float(sum(value for value in pnl if value > 0))
    loss = float(sum(value for value in pnl if value < 0))
    return {
        "trade_n": len(rows),
        "pnl_yen": float(sum(pnl)),
        "pf": _pf(rows),
        "gross_profit_yen": gain,
        "gross_loss_yen": loss,
        "mean_bps": None if not bps else float(np.mean(bps)),
        "median_bps": None if not bps else float(np.median(bps)),
        "max_realized_drawdown_yen": _realized_dd(rows),
        "mtm_max_drawdown_yen": mtm,
        "mean_hold_sec": None if not hold else float(np.mean(hold)),
        "median_hold_sec": None if not hold else float(np.median(hold)),
    }


def _uid(row: dict[str, Any]) -> str:
    return f"{row['date']}|{row['session']}|{row['symbol']}|{int(row['signal_index'])}"


def _compare(reference: list[dict[str, Any]], implementation: list[dict[str, Any]], counts: dict[str, int]) -> dict[str, Any]:
    left = {_uid(row): row for row in reference}
    right = {_uid(row): row for row in implementation}
    for uid in set(left) & set(right):
        ref, got = left[uid], right[uid]
        if abs(float(ref["entry_t"]) - float(got["entry_t"])) > 1e-9 or abs(float(ref["entry_px"]) - float(got["entry_px"])) > 1e-6:
            counts["entry"] += 1
        if abs(float(ref["exit_t"]) - float(got["exit_t"])) > 1e-6 or abs(float(ref["exit_px"]) - float(got["exit_px"])) > 1e-4 or str(ref["reason"]) != str(got["reason"]):
            counts["exit"] += 1
        if str(ref["reason"]) != str(got["reason"]) and "SESSION_FLAT" in (str(ref["reason"]), str(got["reason"])):
            counts["session_close"] += 1
        supports = (float(ref["initial_support"]), float(ref["final_support"]), float(got["initial_support"]), float(got["final_support"]))
        if any(abs(value - supports[0]) > 1e-6 for value in supports):
            counts["support"] += 1
        if int(ref["reconfirm_raises"]) != int(got["reconfirm_raises"]):
            counts["reconfirm"] += 1
        if abs(float(ref["pnl_yen"]) - float(got["pnl_yen"])) > 1e-6:
            counts["pnl"] += 1
        if bool(ref["reentry"]) != bool(got["reentry"]):
            counts["reentry"] += 1
    only_ref = len(set(left) - set(right))
    only_impl = len(set(right) - set(left))
    if only_ref or only_impl:
        counts["eligibility"] += only_ref + only_impl
    return {"counts": counts, "exact_matched_trades": len(set(left) & set(right)), "reference_only": only_ref, "implementation_only": only_impl}


def _period(rows: list[dict[str, Any]], dates: set[str]) -> dict[str, Any]:
    chosen = [row for row in rows if str(row["date"]) in dates]
    body = _stats(chosen)
    return {"trade_n": body["trade_n"], "pnl_yen": body["pnl_yen"], "pf": body["pf"]}


def _folds(dates: list[str]) -> dict[str, set[str]]:
    width = len(dates) // 3
    return {
        "FOLD1": set(dates[:width]),
        "FOLD2": set(dates[width: 2 * width]),
        "FOLD3": set(dates[2 * width:]),
    }


def publish() -> dict[str, Any]:
    cache = OUT / "_scan.pkl"
    payload = pickle.loads(cache.read_bytes()) if cache.exists() else None
    if payload is None:
        payload = scan()
        OUT.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(pickle.dumps(payload))
    compared = _compare(payload["reference"], payload["implementation"], dict(payload["signal_mismatches"]))
    counts = compared["counts"]
    parity = all(value == 0 for value in counts.values()) and compared["reference_only"] == 0 and compared["implementation_only"] == 0
    support_ok = all(abs(float(row["initial_support"]) - float(row["final_support"])) <= 1e-6 for row in payload["reference"] + payload["implementation"])
    fixed = _stats(payload["reference"], payload["fixed_mtm"]["mtm_max_drawdown_yen"])
    impl = _stats(payload["implementation"])
    frozen = _stats(payload["frozen_v2"], payload["v2_mtm"]["mtm_max_drawdown_yen"])
    pnl_delta = fixed["pnl_yen"] - frozen["pnl_yen"]
    pf_delta = None if fixed["pf"] is None or frozen["pf"] is None else fixed["pf"] - frozen["pf"]
    bps_delta = None if fixed["mean_bps"] is None or frozen["mean_bps"] is None else fixed["mean_bps"] - frozen["mean_bps"]
    mtm_delta = None if fixed["mtm_max_drawdown_yen"] is None or frozen["mtm_max_drawdown_yen"] is None else fixed["mtm_max_drawdown_yen"] - frozen["mtm_max_drawdown_yen"]
    advantage = pnl_delta > 0
    prior = payload.get("prior_divergence") or {}
    prior_resolved = bool(prior) and prior.get("reference") == prior.get("runtime") == "EXACT50_FAIL_CLOSED"
    reentry_downstream = counts["reentry"] == 0
    if not parity:
        verdict, nxt = "FIXED_ENTRY_SUPPORT_RUNTIME_BINDING_PARITY_FAIL_V2", "RCA_FIRST_REMAINING_RUNTIME_DIVERGENCE_ONLY"
    elif advantage:
        verdict, nxt = "FIXED_ENTRY_SUPPORT_RUNTIME_BINDING_PARITY_PASS_V1", "FIXED_ENTRY_SUPPORT_PAPER_PREFLIGHT_V1"
    else:
        verdict, nxt = "FIXED_ENTRY_SUPPORT_RUNTIME_PARITY_PASS_BUT_EDGE_NOT_PRESERVED_V1", "KEEP_FROZEN_V2_AS_RUNTIME_BASELINE"
    periods = {"ORIGINAL18": set(ORIGINAL18), "EXTENSION17": set(EXTENSION17), **_folds(list(payload["dates"]))}
    paired = {}
    for name, dates in periods.items():
        left, right = _period(payload["frozen_v2"], dates), _period(payload["reference"], dates)
        paired[name] = {"frozen_v2": left, "fixed_support": right, "pnl_delta": right["pnl_yen"] - left["pnl_yen"]}
    report = {
        "study": "FIXED_ENTRY_SUPPORT_RUNTIME_EXACT50_REFERENCE_REPAIR_V1",
        "verdict": verdict,
        "next": nxt,
        "reference_bug_repaired": True,
        "runtime_exact50_guard_changed": False,
        "registration_guard_changed": False,
        "strategy_changed": False,
        "new_strategy_optimization": False,
        "new_market_data": False,
        "paper_executed": False,
        "submit_cancel_live": [0, 0, 0],
        "verified_windows": payload["verified_windows"],
        "verified_signals": payload["verified_signals"],
        "fixed_support_corrected_reference": fixed,
        "fixed_support_runtime_implementation": impl,
        "frozen_v2_runtime_boundary": frozen,
        "trade_by_trade_parity": bool(parity),
        "active_support_invariant": bool(support_ok),
        "prior_divergence": prior,
        "prior_divergence_resolved": bool(prior_resolved),
        "reentry_mismatch_downstream_of_eligibility": bool(reentry_downstream),
        "candidate_minus_v2": {
            "pnl_delta": pnl_delta,
            "pf_delta": pf_delta,
            "mean_bps_delta": bps_delta,
            "mtm_dd_delta": mtm_delta,
        },
        "advantage_preserved": bool(advantage and parity),
        "verified_window_runtime_binding_parity_proven": bool(parity),
        "full_historical_runtime_binding": "unresolved",
        "paper_preflight_allowed": bool(parity and advantage),
        "periods": paired,
        "period_note": "Period rows contain only trades inside verified registration windows. Unresolved dates are not filled.",
        "first_divergence": payload["first_divergence"],
        **compared,
    }
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "report.json").write_text(json.dumps(report, indent=2, sort_keys=True), encoding="utf-8")
    (OUT / "report.md").write_text(_markdown(report), encoding="utf-8")
    wb = Workbook()
    wb.active.title = "summary"
    wb.active.append(["verdict", verdict, nxt])
    wb.active.append(["parity", parity, "advantage", advantage])
    wb.active.append(["fixed_n", fixed["trade_n"], "fixed_pnl", fixed["pnl_yen"], "fixed_pf", fixed["pf"]])
    wb.active.append(["v2_n", frozen["trade_n"], "v2_pnl", frozen["pnl_yen"], "v2_pf", frozen["pf"]])
    wb.active.append(["pnl_delta", pnl_delta, "mtm_delta", mtm_delta])
    ws = wb.create_sheet("periods")
    ws.append(["period", "v2_n", "v2_pnl", "v2_pf", "fixed_n", "fixed_pnl", "fixed_pf", "pnl_delta"])
    for name, row in paired.items():
        ws.append([name, row["frozen_v2"]["trade_n"], row["frozen_v2"]["pnl_yen"], row["frozen_v2"]["pf"], row["fixed_support"]["trade_n"], row["fixed_support"]["pnl_yen"], row["fixed_support"]["pf"], row["pnl_delta"]])
    mm = wb.create_sheet("mismatches")
    mm.append(["field", "count"])
    for key, value in counts.items():
        mm.append([key, value])
    wb.save(OUT / "audit.xlsx")
    cache.unlink(missing_ok=True)
    print(f"VERDICT {verdict} NEXT {nxt} parity={parity} advantage={advantage}", flush=True)
    return report


def _markdown(report: dict[str, Any]) -> str:
    fixed = report["fixed_support_corrected_reference"]
    impl = report["fixed_support_runtime_implementation"]
    frozen = report["frozen_v2_runtime_boundary"]
    delta = report["candidate_minus_v2"]
    lines = [
        f"# {report['study']}",
        "",
        f"VERDICT: {report['verdict']}",
        f"NEXT: {report['next']}",
        "",
        "The reference now requires the day-fixed Dynamic40 set to equal the verified Kabu registered set before any symbol is admitted.",
        "resolve_registered_probe_symbol was not modified.",
        "",
        f"Corrected reference {fixed['trade_n']} / {fixed['pnl_yen']} / {fixed['pf']}",
        f"Runtime implementation {impl['trade_n']} / {impl['pnl_yen']} / {impl['pf']}",
        f"Trade-by-trade parity {report['trade_by_trade_parity']}",
        f"Frozen V2 on the same boundary {frozen['trade_n']} / {frozen['pnl_yen']} / {frozen['pf']} MTM {frozen['mtm_max_drawdown_yen']}",
        f"Fixed-support MTM {fixed['mtm_max_drawdown_yen']}",
        f"Candidate minus V2 PnL {delta['pnl_delta']} PF {delta['pf_delta']} mean bps {delta['mean_bps_delta']} MTM DD {delta['mtm_dd_delta']}",
        "",
        report["period_note"],
        "Full historical registration remains unresolved.",
        "submit/cancel/live = 0/0/0. Paper was not executed.",
        "",
    ]
    return "\n".join(lines)


if __name__ == "__main__":
    publish()
