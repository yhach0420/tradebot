"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_c0_indicator_exit import ANALYSIS_ID, ARCHITECTURE_ID
from research.am_entry_profit_improvement.publish import json_sanitize

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_c0_indicator_exit"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "precommit",
    "non_interference",
    "training_population",
    "folds",
    "model",
    "trade_exit",
    "fixed26",
    "daily_pnl",
    "economics",
    "overfit_audit",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def kv_rows(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": d.get(k)} for k in d.keys()]


def build_markdown(report: dict[str, Any]) -> str:
    req = dict(report.get("required") or {})
    dec = dict(report.get("decision") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"ARCHITECTURE_ID: `{ARCHITECTURE_ID}`",
        f"EXIT_INDICATOR_SPEC_SHA256: `{req.get('EXIT_INDICATOR_SPEC_SHA256')}`",
        f"VERDICT: **{dec.get('VERDICT') or req.get('VERDICT')}**",
        f"NEXT: `{dec.get('NEXT') or req.get('NEXT')}`",
        "",
        "## Required output",
        "",
    ]
    for k in (
        "BASE_PARITY",
        "TIME_FEATURE_N",
        "HOLDING_TIME_RULE_N",
        "OUTER_FOLD_N",
        "TRAIN_TRADE_N",
        "OOF_AUC",
        "FIXED_TRADE_N",
        "FIXED_DELTA_NET",
        "L2_BASE_N",
        "L2_IMPROVED_N",
        "WINNER_PREMATURE_EXIT_N",
        "EIND_AUGMENT_TRADE_N",
        "EIND_AUGMENT_NET",
        "EIND_AUGMENT_PF",
        "EIND_OVERLAY_NET",
        "EIND_OVERLAY_PF",
        "EIND_OVERLAY_DD",
        "PAIRED_POS_DAYS",
        "PAIRED_NEG_DAYS",
        "PAIRED_ZERO_DAYS",
        "PAIRED_MEDIAN",
        "EX_BEST",
        "EX_TOP3",
        "BEST_DAY_SHARE",
        "TOP3_DAY_SHARE",
        "DELTA_NET_VS_C0_C14",
        "DELTA_PF_VS_C0_C14",
        "DELTA_DD_VS_C0_C14",
        "DELTA_PAIRED_MEDIAN_VS_C0_C14",
        "CURRENT_PRESERVATION_PASS",
        "NON_INTERFERENCE_PASS",
        "RUNTIME_PID_BEFORE",
        "RUNTIME_PID_AFTER",
        "CAPTURE_PID_BEFORE",
        "CAPTURE_PID_AFTER",
        "RUNTIME_HEARTBEAT_ADVANCED",
        "CAPTURE_ADVANCED",
        "EIND_FULL_PASS",
        "TRUE_OOS",
        "NEW_FORWARD_N",
    ):
        lines.append(f"- {k}: `{req.get(k)}`")
    if req.get("STOP_REASON"):
        lines.extend(["", f"STOP_REASON: {req.get('STOP_REASON')}"])
    lines.extend(
        [
            "",
            "## Notes",
            "",
            "- Diagnostic metrics (AUC, L2 counts, premature-exit) are not pass/fail.",
            "- TRUE_OOS is false. This 18-day set is a burned development set.",
            "- Runtime ENTRY/EXIT and Capture were not modified by this research process.",
            "",
        ]
    )
    return "\n".join(lines) + "\n"
