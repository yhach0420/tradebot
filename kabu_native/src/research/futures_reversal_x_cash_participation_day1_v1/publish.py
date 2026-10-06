"""report.json / report.md / audit.xlsx. Day2 primary freeze untouched."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_reversal_x_cash_participation_day1_v1 import ANALYSIS_ID
from research.futures_reversal_x_cash_participation_day1_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "EventTable",
    "TrueVsFalse",
    "Binding",
    "Robustness",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"CASE: {d.get('CASE')}",
        "",
        "Question: at the frozen BOTH_UP confirmation, is the cash market and",
        "the activity-selected TOP16 actually participating in recovery?",
        "CASH_CONFIRM = CASH1 AND CASH2 AND CASH3. 180s only. No new futures",
        "features. Day2 primary freeze unchanged. No ENTRY/EXIT.",
        "",
        "## Required answers",
        "",
        f"- 1 trigger N: `{a.get('1_trigger_N')}`",
        f"- 2 future leakage N: `{a.get('2_future_leakage_N')}`",
        f"- 3 CASH1 TRUE N: `{a.get('3_CASH1_TRUE_N')}`",
        f"- 4 CASH2 TRUE N: `{a.get('4_CASH2_TRUE_N')}`",
        f"- 5 CASH3 TRUE N: `{a.get('5_CASH3_TRUE_N')}`",
        f"- 6 CASH_CONFIRM TRUE N: `{a.get('6_CASH_CONFIRM_TRUE_N')}`",
        f"- 7 09:21:46 CASH states: `{a.get('7_09:21:46_CASH_states')}`",
        f"- 8 10:32:39 CASH states: `{a.get('8_10:32:39_CASH_states')}`",
        f"- 9 CASH_CONFIRM rejects 10:30: `{_fmt(a.get('9_CASH_CONFIRM_rejects_10:30'))}`",
        f"- 10 TRUE timestamps: `{a.get('10_TRUE_event_timestamps')}`",
        f"- 11 FALSE timestamps: `{a.get('11_FALSE_event_timestamps')}`",
        f"- 12 TRUE TOP LONG 1m: `{a.get('12_TRUE_TOP_LONG_1m')}`",
        f"- 13 3m: `{a.get('13_TRUE_TOP_LONG_3m')}`",
        f"- 14 5m: `{a.get('14_TRUE_TOP_LONG_5m')}`",
        f"- 15 10m mean/median: `{a.get('15_TRUE_TOP_LONG_10m_mean_median')}`",
        f"- 16 FALSE TOP LONG 10m: `{a.get('16_FALSE_TOP_LONG_10m')}`",
        f"- 17 TRUE BOTTOM LONG: `{a.get('17_TRUE_BOTTOM_LONG_10m')}`",
        f"- 18 TRUE TOP-BOTTOM LONG: `{_fmt(a.get('18_TRUE_TOP_BOTTOM_LONG_10m'))}`",
        f"- 19 TRUE ALL48 LONG: `{a.get('19_TRUE_ALL48_LONG_10m')}`",
        f"- 20 TRUE event win rate: `{a.get('20_TRUE_event_win_rate')}`",
        f"- 21 TRUE stock-level win rate: `{a.get('21_TRUE_stock_level_win_rate')}`",
        f"- 22 drop best event: `{a.get('22_drop_best_event')}`",
        f"- 23 drop worst event: `{a.get('23_drop_worst_event')}`",
        f"- 24 drop top symbol: `{a.get('24_drop_top_symbol')}`",
        f"- 25 drop top2 symbols: `{a.get('25_drop_top2_symbols')}`",
        f"- 26 N sufficient: `{_fmt(a.get('26_N_sufficient'))}`",
        f"- 27 exact causal mechanism: {a.get('27_exact_causal_mechanism')}",
        f"- 28 ENTRY thesis plausible: `{_fmt(a.get('28_ENTRY_thesis_plausible'))}`",
        f"- 29 ENTRY built: `{_fmt(a.get('29_ENTRY_built'))}`",
        f"- 30 EXIT built: `{_fmt(a.get('30_EXIT_built'))}`",
        f"- 31 Day2 primary changed: `{_fmt(a.get('31_Day2_primary_changed'))}`",
        f"- 32 Runtime changed: `{_fmt(a.get('32_Runtime_changed'))}`",
        f"- 33 Paper changed: `{_fmt(a.get('33_Paper_changed'))}`",
        f"- 34 submit/cancel/live: `{a.get('34_submit_cancel_live')}`",
        f"- 35 CASE: `{a.get('35_CASE')}`",
        f"- 36 VERDICT: `{a.get('36_VERDICT')}`",
        f"- 37 NEXT: `{a.get('37_NEXT')}`",
        "",
        "No CASH combination search. No 09:20-only rule. No Day2 substitution.",
        "",
    ]
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize(body)
    (OUT / "report.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(build_markdown(body), encoding="utf-8")
    ts = dict(body.get("true_stats") or {})
    fs = dict(body.get("false_stats") or {})
    sheets = {
        "Answers": _kv(dict(body.get("answers") or {})),
        "EventTable": list(body.get("event_table") or []),
        "TrueVsFalse": [
            {"split": "TRUE", **{k: v for k, v in ts.items()}},
            {"split": "FALSE_10m", **(fs.get("10m") or {})},
        ],
        "Binding": _kv(dict(body.get("binding") or {})),
        "Robustness": _kv(dict(body.get("robustness") or {})),
        "Decision": _kv(dict(body.get("decision") or {})),
        "Safety": [
            {"key": "ENTRY_built", "value": False},
            {"key": "EXIT_built", "value": False},
            {"key": "Day2_primary_changed", "value": False},
            {"key": "Day2_substitutions_allowed", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "submit_cancel_live", "value": (body.get("answers") or {}).get("34_submit_cancel_live")},
            {"key": "future_leakage_n", "value": body.get("future_leakage_n")},
            {"key": "NEW_FUTURES_FEATURES", "value": False},
            {"key": "close_day1_feature_mining", "value": (body.get("decision") or {}).get("close_day1_feature_mining")},
        ],
    }
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(xlsx),
    }
