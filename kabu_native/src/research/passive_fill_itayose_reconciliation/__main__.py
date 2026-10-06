"""Offline Passive Fill itayose reconciliation. Paper/Ingress not started. No UNIFORM10."""
from __future__ import annotations

import json
import os
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

NATIVE = Path(__file__).resolve().parents[3]
SRC = NATIVE / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
if str(NATIVE / "scripts") not in sys.path:
    sys.path.insert(0, str(NATIVE / "scripts"))

from research.anchor_vs_event_driven.run_comparison import (  # noqa: E402
    classify_capture,
    find_capture_dir,
    historical_universe,
)
from research.passive_fill_itayose_reconciliation import (  # noqa: E402
    C13_SHA,
    FULL14,
    MAX_WORKERS,
)
from research.passive_fill_itayose_reconciliation.analyze import (  # noqa: E402
    big_winner_audit,
    compare_fills,
    headline,
    pick_verdict,
    reassess,
    subset_days,
)
from research.passive_fill_itayose_reconciliation.publish import (  # noqa: E402
    OUT,
    build_report,
    write_artifacts,
)
from research.passive_fill_itayose_reconciliation.replay import process_day  # noqa: E402
from small_paper.v1r_activation_binding import (  # noqa: E402
    candidate_source_digest,
    collect_runtime_inventory,
    file_sha256,
)
from openpyxl import load_workbook


def _sheet_rows(xlsx: Path, name: str) -> list[dict]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[name]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c) if c is not None else "" for c in next(rows_iter, [])]
    out = []
    for raw in rows_iter:
        if not header or header[0] == "empty":
            break
        rec = {header[i]: raw[i] if i < len(raw) else None for i in range(len(header))}
        out.append(rec)
    wb.close()
    return out


def _regen_from_xlsx() -> dict | None:
    xlsx = OUT / "audit.xlsx"
    if not xlsx.is_file():
        return None
    old_trades = _sheet_rows(xlsx, "trades_old")
    new_trades = _sheet_rows(xlsx, "trades_new")
    old_fills = _sheet_rows(xlsx, "fills_old")
    new_fills = _sheet_rows(xlsx, "fills_new")
    inv = _sheet_rows(xlsx, "inventory")
    if not old_trades and not old_fills:
        return None
    return {
        "old_trades": old_trades,
        "new_trades": new_trades,
        "old_fills": old_fills,
        "new_fills": new_fills,
        "inventory": inv,
        "day_replay": _sheet_rows(xlsx, "day_replay"),
    }


def _paper_stopped() -> dict:
    import subprocess

    try:
        raw = subprocess.check_output(
            ["powershell", "-NoProfile", "-Command",
             "Get-CimInstance Win32_Process | Where-Object { $_.Name -match 'python' } | "
             "Select-Object ProcessId,CommandLine | ConvertTo-Json -Compress"],
            text=True,
            timeout=20,
        )
    except Exception as exc:
        return {"ok": True, "note": f"process_scan_failed:{exc}", "hits": []}
    hits = []
    try:
        rows = json.loads(raw or "[]")
    except Exception:
        rows = []
    if isinstance(rows, dict):
        rows = [rows]
    for row in rows or []:
        cmd = str(row.get("CommandLine") or "")
        if "market_ingress" in cmd or "pilot_runner" in cmd or "paper_full_day" in cmd or "v1r_paper" in cmd:
            hits.append({"pid": row.get("ProcessId"), "cmd": cmd[:240]})
    return {"ok": len(hits) == 0, "hits": hits}


def _identity() -> dict:
    inv = collect_runtime_inventory()
    digest = candidate_source_digest(inv)
    changed = []
    for rel in (
        "src/research/e1_x34a_execution_policy/arms.py",
        "src/research/e1_x34a_execution_policy/executable_board.py",
        "src/small_paper/v1r_native_entry_live.py",
        "src/small_paper/v1r_activation_binding.py",
        "src/research/e1_x28_executable_joint/board.py",
        "scripts/run_p0_4_exact_vs_fast_parity.py",
    ):
        fp = NATIVE / rel
        changed.append({"rel": rel, "exists": fp.is_file(), "sha256": file_sha256(fp) if fp.is_file() else None})
    return {
        "C13_SHA": C13_SHA,
        "working_tree_matches_C13": digest == C13_SHA,
        "candidate_source_digest": digest,
        "changed_files": changed,
        "NEW_RUNTIME_CANDIDATE_IDENTITY_REQUIRED": True,
        "NEW_RUNTIME_CANDIDATE_FROZEN": False,
        "C13_OVERWRITTEN": False,
    }


def _inventory() -> tuple[list[dict], list[dict]]:
    mc = NATIVE / "data" / "market_capture"
    days = sorted(p.name for p in mc.iterdir() if p.is_dir() and p.name.isdigit() and len(p.name) == 8) if mc.is_dir() else []
    rows = []
    jobs = []
    for day in days:
        cap = find_capture_dir(day)
        uni, uni_src = historical_universe(day, cap or (mc / day))
        row = classify_capture(day, cap, len(uni))
        row["universe_source"] = uni_src
        rows.append(row)
        usable_job = cap is not None and uni and (
            bool(row.get("full")) or (day in FULL14 and bool(row.get("usable")))
        )
        if usable_job:
            jobs.append(
                {
                    "date": day,
                    "capture_path": str(cap),
                    "universe": uni,
                    "universe_source": uni_src,
                    "capture_class": row.get("capture_class"),
                }
            )
    return rows, jobs


def main() -> int:
    os.environ["PYTHONPATH"] = (
        f"{SRC};{NATIVE.parent}" if os.name == "nt" else f"{SRC}:{NATIVE.parent}"
    )
    print("SAFETY submit/cancel/live=0/0/0 offline_replay_only", flush=True)
    print("UNIFORM10 HOLD. CLOCK_GRID / ENTRY / EXIT unchanged.", flush=True)
    paper = _paper_stopped()
    if not paper.get("ok"):
        print("BLOCKED: paper/ingress process still running", paper.get("hits"), flush=True)
        return 2
    ident = _identity()
    print("working_tree_matches_C13", ident.get("working_tree_matches_C13"), flush=True)
    print("digest", ident.get("candidate_source_digest"), flush=True)

    inv, jobs = _inventory()
    print(f"inventory days={len(inv)} FULL jobs={len(jobs)}", flush=True)
    for row in inv:
        print(
            f"  {row.get('date')} {row.get('capture_class')} uni={row.get('universe_n')} {row.get('exclusion_reason') or ''}",
            flush=True,
        )
    only = str(os.environ.get("ITAYOSE_DAYS") or "").strip()
    if only:
        want = {x.strip() for x in only.split(",") if x.strip()}
        jobs = [j for j in jobs if j["date"] in want]
        print(f"ITAYOSE_DAYS filter -> {len(jobs)} jobs", flush=True)

    regen = str(os.environ.get("ITAYOSE_REGEN") or "").strip().lower() in {"1", "true", "yes"}
    day_rows: list[dict] = []
    old_fills: list[dict] = []
    new_fills: list[dict] = []
    old_trades: list[dict] = []
    new_trades: list[dict] = []
    old_expired: list[dict] = []
    new_expired: list[dict] = []
    failed: list[dict] = []
    if regen:
        cached = _regen_from_xlsx()
        if not cached:
            print("ITAYOSE_REGEN set but audit.xlsx missing/empty", flush=True)
            return 1
        print("REGEN from audit.xlsx (no causal re-stream)", flush=True)
        old_fills = cached["old_fills"]
        new_fills = cached["new_fills"]
        old_trades = cached["old_trades"]
        new_trades = cached["new_trades"]
        day_rows = cached.get("day_replay") or []
        jobs = jobs or [{"date": "regen"}]
    elif jobs:
        workers = min(MAX_WORKERS, len(jobs))
        with ProcessPoolExecutor(max_workers=workers) as ex:
            futs = {ex.submit(process_day, job): job["date"] for job in jobs}
            for fut in as_completed(futs):
                day = futs[fut]
                try:
                    body = fut.result()
                except Exception as exc:
                    body = {"ok": False, "date": day, "blocker": f"{type(exc).__name__}:{exc}"}
                day_rows.append(body)
                print(f"done {day} ok={body.get('ok')} blocker={body.get('blocker')}", flush=True)
        day_rows.sort(key=lambda r: str(r.get("date") or ""))
        for row in day_rows:
            if not row.get("ok"):
                failed.append({"date": row.get("date"), "blocker": row.get("blocker")})
                continue
            old = row.get("old") or {}
            new = row.get("new") or {}
            old_fills.extend(old.get("fills") or [])
            new_fills.extend(new.get("fills") or [])
            old_trades.extend(old.get("trades") or [])
            new_trades.extend(new.get("trades") or [])
            old_expired.extend(old.get("expired") or [])
            new_expired.extend(new.get("expired") or [])

    fill_cmp = compare_fills(old_fills, new_fills)
    old_full14 = subset_days(old_trades, FULL14)
    new_full14 = subset_days(new_trades, FULL14)
    winners = big_winner_audit(old_fills, new_fills, old_trades, new_trades, old_expired, new_expired)
    reass = reassess(
        old_trades=old_trades,
        new_trades=new_trades,
        fill_cmp=fill_cmp,
        old_full14=old_full14,
        new_full14=new_full14,
    )
    replay_ok = not failed and bool(jobs)
    verdict, reason = pick_verdict(
        replay_ok=replay_ok,
        fill_cmp=fill_cmp,
        old_full14=old_full14,
        new_full14=new_full14,
        winners=winners,
        old_replay_vs_prior=reass.get("old_replay_vs_prior_full14_pnl"),
        old_dev_pnl=(reass.get("items") or {}).get("DEV10 vs POST7 decay", {}).get("old_dev"),
        old_post_pnl=(reass.get("items") or {}).get("DEV10 vs POST7 decay", {}).get("old_post"),
    )

    inv_summary = {
        "days_seen": len(inv),
        "full_n": sum(1 for r in inv if r.get("full")),
        "jobs": len(jobs),
        "failed": failed,
        "full_dates": [r.get("date") for r in inv if r.get("full")],
        "paper_stopped": paper,
    }
    report_body = {
        "verdict": verdict,
        "verdict_reason": reason,
        "identity": ident,
        "inventory_summary": inv_summary,
        "fill_cmp": fill_cmp,
        "reassess": reass,
        "big_winners": winners,
        "all_full_old": headline(old_trades),
        "all_full_new": headline(new_trades),
    }
    report = build_report(report_body)
    extra = {
        "inventory": [
            {k: r.get(k) for k in (
                "date", "capture_class", "universe_n", "full", "exclusion_reason",
                "first_event_time", "last_event_time",
            )}
            for r in inv
        ],
        "fills_old": old_fills,
        "fills_new": new_fills,
        "removed_invalid": fill_cmp.get("removed_rows") or [],
        "by_day": fill_cmp.get("by_day") or [],
        "by_symbol": fill_cmp.get("by_symbol") or [],
        "by_anchor": fill_cmp.get("by_anchor") or [],
        "anchor_0905": [
            x
            for x in (fill_cmp.get("by_anchor") or [])
            if str(x.get("anchor") or "") == "09:05"
        ]
        + [
            {
                "scope": "09:05_totals",
                **(fill_cmp.get("09:05") or {}),
            }
        ],
        "big_winners": winners,
        "trades_old": [
            {k: t.get(k) for k in (
                "date", "symbol", "session", "anchor_time", "candidate_rank", "score",
                "limit", "limit_price", "fill_time", "fill_price", "fill_event_time",
                "exit_time", "exit_price", "exit_reason", "pnl_yen_100",
                "board_execution_state", "fill_class", "cross_ask", "cross_ask_qty",
                "AskSign", "opening_status", "OpeningPrice",
            )}
            for t in old_trades
        ],
        "trades_new": [
            {k: t.get(k) for k in (
                "date", "symbol", "session", "anchor_time", "candidate_rank", "score",
                "limit", "limit_price", "fill_time", "fill_price", "fill_event_time",
                "exit_time", "exit_price", "exit_reason", "pnl_yen_100",
                "board_execution_state", "fill_class", "cross_ask", "cross_ask_qty",
                "AskSign", "opening_status", "OpeningPrice",
            )}
            for t in new_trades
        ],
        "reassess": [
            {"item": k, **(v if isinstance(v, dict) else {"value": v})}
            for k, v in (reass.get("items") or {}).items()
        ],
        "day_replay": [
            {
                "date": r.get("date"),
                "ok": r.get("ok"),
                "blocker": r.get("blocker"),
                "old_fills": (r.get("old") or {}).get("fills_n") if isinstance(r.get("old"), dict) else r.get("old_fills"),
                "new_fills": (r.get("new") or {}).get("fills_n") if isinstance(r.get("new"), dict) else r.get("new_fills"),
                "old_pnl": (r.get("old") or {}).get("pnl") if isinstance(r.get("old"), dict) else r.get("old_pnl"),
                "new_pnl": (r.get("new") or {}).get("pnl") if isinstance(r.get("new"), dict) else r.get("new_pnl"),
                "old_trades": (r.get("old") or {}).get("trade_n") if isinstance(r.get("old"), dict) else r.get("old_trades"),
                "new_trades": (r.get("new") or {}).get("trade_n") if isinstance(r.get("new"), dict) else r.get("new_trades"),
                "elapsed_sec": r.get("elapsed_sec"),
            }
            for r in day_rows
        ],
    }
    write_artifacts(report, extra_sheets=extra)
    print("OUT", OUT, flush=True)
    print("verdict", verdict, flush=True)
    print(reason, flush=True)
    print("STOP. UNIFORM10 not started.", flush=True)
    return 0 if replay_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
