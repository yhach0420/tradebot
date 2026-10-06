"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.passive_fill_itayose_reconciliation import (
    ANALYSIS_ID,
    C13_ID,
    C13_SHA,
    EXECUTION_SEMANTICS_DEFECT_FIXED,
    STRATEGY_RETUNED,
    TASK_LABEL,
    UNIFORM10_HOLD,
    V25_ID,
    V25_SHA,
    WAIT_SEC,
)

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "passive_fill_itayose_reconciliation"
JST = timezone(timedelta(hours=9))
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    return obj


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            if v == float("inf"):
                v = "Infinity"
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(_k)) + 2))


def _kv_sheet(ws, data: dict[str, Any]) -> None:
    ws.append(["key", "value"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for k, v in data.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        ws.append([k, v])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 80


def build_report(body: dict[str, Any]) -> dict[str, Any]:
    now = datetime.now(JST).isoformat(timespec="seconds")
    verdict = body["verdict"]
    reason = body["verdict_reason"]
    fill = body["fill_cmp"]
    reass = body["reassess"]
    ident = body["identity"]
    h_old = reass.get("old_headline_full14") or {}
    h_new = reass.get("new_headline_full14") or {}
    fam_o = reass.get("old_family") or {}
    fam_n = reass.get("new_family") or {}
    return {
        "ANALYSIS_ID": ANALYSIS_ID,
        "TASK_LABEL": TASK_LABEL,
        "generated_at": now,
        "submit_cancel_live": "0/0/0",
        "offline_replay_only": True,
        "STRATEGY_RETUNED": STRATEGY_RETUNED,
        "EXECUTION_SEMANTICS_DEFECT_FIXED": EXECUTION_SEMANTICS_DEFECT_FIXED,
        "UNIFORM10_HOLD": UNIFORM10_HOLD,
        "UNIFORM10_ENTRY_FULL_DATA_REBUILD": "NOT_STARTED",
        "WAIT_SEC": WAIT_SEC,
        "CLOCK_GRID_CHANGED": False,
        "ENTRY_SCORE_CHANGED": False,
        "EXIT_CHANGED": False,
        "fill_price_policy": "NO_PRICE_IMPROVEMENT fill_price=limit_price",
        "C13_ID": C13_ID,
        "C13_SHA": C13_SHA,
        "C13_OVERWRITTEN": False,
        "V25_ID": V25_ID,
        "V25_SHA": V25_SHA,
        "working_tree_matches_C13": ident.get("working_tree_matches_C13"),
        "NEW_RUNTIME_CANDIDATE_IDENTITY_REQUIRED": True,
        "NEW_RUNTIME_CANDIDATE_FROZEN": False,
        "identity": ident,
        "confirmed_defect": {
            "symbol": "5801",
            "date": "20260827",
            "anchor": "09:05",
            "limit": 4045,
            "reported_fill": "09:05:00.252 @ 4045",
            "AskSign": "0102",
            "real_opening": "09:06:00 OpeningPrice=4097",
            "verdict": "INVALID_FILL_RUNTIME_DEFECT",
        },
        "inventory": body.get("inventory_summary"),
        "fill_comparison": {
            "old_fills": fill.get("old_fills"),
            "new_fills": fill.get("new_fills"),
            "removed_invalid_fills": fill.get("removed_invalid_fills"),
            "added_fills": fill.get("added_fills"),
            "occupancy_displaced_valid": fill.get("occupancy_displaced_valid"),
            "old_class_counts": fill.get("old_class_counts"),
            "09:05": fill.get("09:05"),
        },
        "by_day": fill.get("by_day"),
        "by_symbol": fill.get("by_symbol"),
        "by_anchor": fill.get("by_anchor"),
        "full14_old": h_old,
        "full14_new": h_new,
        "all_full_old": body.get("all_full_old"),
        "all_full_new": body.get("all_full_new"),
        "families_old": fam_o,
        "families_new": fam_n,
        "entry_kind_old": reass.get("old_entry_kind"),
        "entry_kind_new": reass.get("new_entry_kind"),
        "concentration_old": reass.get("old_concentration"),
        "concentration_new": reass.get("new_concentration"),
        "big_winners": body.get("big_winners"),
        "reassess": reass.get("items"),
        "old_replay_vs_prior_full14_pnl": reass.get("old_replay_vs_prior_full14_pnl"),
        "verdict": verdict,
        "verdict_reason": reason,
        "STOP": True,
        "do_not_start_UNIFORM10": True,
    }


def write_md(report: dict[str, Any]) -> str:
    fill = report.get("fill_comparison") or {}
    f0905 = fill.get("09:05") or {}
    ident = report.get("identity") or {}
    reass = report.get("reassess") or {}
    lines = [
        "# TRADEBOT — Passive Fill itayose reconciliation",
        "",
        "Offline replay only. CLOCK_GRID / ENTRY score / EXIT unchanged. submit/cancel/live=0/0/0.",
        "",
        f"STRATEGY_RETUNED: {report.get('STRATEGY_RETUNED')}",
        f"EXECUTION_SEMANTICS_DEFECT_FIXED: {report.get('EXECUTION_SEMANTICS_DEFECT_FIXED')}",
        f"UNIFORM10_HOLD: {report.get('UNIFORM10_HOLD')}",
        f"WAIT_SEC: {report.get('WAIT_SEC')}",
        "",
        "## Verdict",
        f"verdict: {report.get('verdict')}",
        f"reason: {report.get('verdict_reason')}",
        "STOP. Do not start UNIFORM10 ENTRY FULL-DATA REBUILD.",
        "",
        "## Identity",
        f"C13_ID: {report.get('C13_ID')}",
        f"C13_SHA: {report.get('C13_SHA')}",
        f"C13_OVERWRITTEN: {report.get('C13_OVERWRITTEN')}",
        f"working_tree_matches_C13: {report.get('working_tree_matches_C13')}",
        f"NEW_RUNTIME_CANDIDATE_IDENTITY_REQUIRED: {report.get('NEW_RUNTIME_CANDIDATE_IDENTITY_REQUIRED')}",
        f"NEW_RUNTIME_CANDIDATE_FROZEN: {report.get('NEW_RUNTIME_CANDIDATE_FROZEN')}",
        f"digest: {ident.get('candidate_source_digest')}",
        "",
        "## Confirmed defect",
        "5801 @ 20260827 09:05 limit 4045 reported fill 09:05:00.252 on itayose/special board "
        "(CurrentPrice/OpeningPrice/TradingVolume null, AskSign=0102). Real open 09:06 OpeningPrice=4097.",
        "Canonical gate: `is_executable_continuous_board` (kabu STATION OpenAPI Sign/Status mapping).",
        "Pre-open / itayose / special quote is not Passive Fill evidence. WAIT_SEC remains 1.0.",
        "fill_price remains limit_price (NO_PRICE_IMPROVEMENT).",
        "",
        "## Fill comparison (all FULL captures)",
        f"old_fills: {fill.get('old_fills')}",
        f"new_fills: {fill.get('new_fills')}",
        f"removed_invalid_fills: {fill.get('removed_invalid_fills')}",
        f"added_fills: {fill.get('added_fills')}",
        f"occupancy_displaced_valid: {fill.get('occupancy_displaced_valid')}",
        f"old_class_counts: {json.dumps(fill.get('old_class_counts') or {}, ensure_ascii=False)}",
        "",
        "### 09:05",
        f"old_fills: {f0905.get('old_fills')}",
        f"new_fills: {f0905.get('new_fills')}",
        f"removed_invalid_fills: {f0905.get('removed_invalid_fills')}",
        f"old_class_counts: {json.dumps(f0905.get('old_class_counts') or {}, ensure_ascii=False)}",
        "",
        "## FULL14 causal replay",
        f"old pnl: {(report.get('full14_old') or {}).get('pnl')}  trades: {(report.get('full14_old') or {}).get('trades')}  PF: {(report.get('full14_old') or {}).get('PF')}  maxDD: {(report.get('full14_old') or {}).get('maxDD')}",
        f"new pnl: {(report.get('full14_new') or {}).get('pnl')}  trades: {(report.get('full14_new') or {}).get('trades')}  PF: {(report.get('full14_new') or {}).get('PF')}  maxDD: {(report.get('full14_new') or {}).get('maxDD')}",
        f"old vs prior Fixed 2289100: {report.get('old_replay_vs_prior_full14_pnl')}",
        "",
        "## All FULL captures causal replay",
        f"old pnl: {(report.get('all_full_old') or {}).get('pnl')}  trades: {(report.get('all_full_old') or {}).get('trades')}  PF: {(report.get('all_full_old') or {}).get('PF')}  maxDD: {(report.get('all_full_old') or {}).get('maxDD')}",
        f"new pnl: {(report.get('all_full_new') or {}).get('pnl')}  trades: {(report.get('all_full_new') or {}).get('trades')}  PF: {(report.get('all_full_new') or {}).get('PF')}  maxDD: {(report.get('all_full_new') or {}).get('maxDD')}",
        "",
        "## Families (FULL14)",
        f"OPEN_EARLY old: {(report.get('families_old') or {}).get('OPEN_EARLY')} new: {(report.get('families_new') or {}).get('OPEN_EARLY')}",
        f"NORMAL old: {(report.get('families_old') or {}).get('NORMAL_SESSION')} new: {(report.get('families_new') or {}).get('NORMAL_SESSION')}",
        f"TAIL old: {(report.get('families_old') or {}).get('SESSION_TAIL')} new: {(report.get('families_new') or {}).get('SESSION_TAIL')}",
        f"09:05 old: {(report.get('families_old') or {}).get('09:05')} new: {(report.get('families_new') or {}).get('09:05')}",
        f"AM old: {(report.get('families_old') or {}).get('AM')} new: {(report.get('families_new') or {}).get('AM')}",
        f"PM old: {(report.get('families_old') or {}).get('PM')} new: {(report.get('families_new') or {}).get('PM')}",
        "",
        "## First-entry / re-entry (FULL14)",
        f"old: {report.get('entry_kind_old')}",
        f"new: {report.get('entry_kind_new')}",
        "",
        "## Concentration (FULL14)",
        f"old: {json.dumps(report.get('concentration_old') or {}, ensure_ascii=False, default=str)}",
        f"new: {json.dumps(report.get('concentration_new') or {}, ensure_ascii=False, default=str)}",
        "",
        "## Big winners",
    ]
    for w in report.get("big_winners") or []:
        lines.append(
            f"- {w.get('symbol')} {w.get('date')} {w.get('anchor')}: "
            f"open={w.get('opening_time')} px={w.get('opening_price')} "
            f"anchor_state={w.get('anchor_board_state')} limit={w.get('limit')} "
            f"old_fill={w.get('old_fill')}@{w.get('old_fill_price')} class={w.get('old_fill_class')} "
            f"corrected={w.get('corrected')} new_fill={w.get('new_fill')} "
            f"old_pnl={w.get('old_pnl')} corrected_pnl={w.get('corrected_pnl')}"
        )
    lines += ["", "## Reassess prior conclusions"]
    for name, rec in reass.items():
        if isinstance(rec, dict):
            note = rec.get("note") or ""
            payload = {k: v for k, v in rec.items() if k not in {"status", "note"}}
            lines.append(
                f"- {name}: {rec.get('status')} {json.dumps(payload, ensure_ascii=False, default=str)}"
                + (f" — {note}" if note else "")
            )
        else:
            lines.append(f"- {name}: {rec}")
    lines += [
        "",
        "## Safety",
        "Paper not started. Ingress not started. submit/cancel/live=0/0/0.",
        "C13 snapshot not overwritten. New candidate not frozen.",
        "",
    ]
    return "\n".join(lines) + "\n"


def write_artifacts(report: dict[str, Any], *, extra_sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    # Keep the folder to the three required files only.
    for p in list(OUT.iterdir()) if OUT.is_dir() else []:
        if p.name not in {"report.json", "report.md", "audit.xlsx"}:
            if p.is_file():
                p.unlink()
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(report), indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    (OUT / "report.md").write_text(write_md(report), encoding="utf-8")
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "summary"
    _kv_sheet(
        ws0,
        {
            "verdict": report.get("verdict"),
            "verdict_reason": report.get("verdict_reason"),
            "STRATEGY_RETUNED": report.get("STRATEGY_RETUNED"),
            "EXECUTION_SEMANTICS_DEFECT_FIXED": report.get("EXECUTION_SEMANTICS_DEFECT_FIXED"),
            "UNIFORM10_HOLD": report.get("UNIFORM10_HOLD"),
            "C13_OVERWRITTEN": report.get("C13_OVERWRITTEN"),
            "NEW_RUNTIME_CANDIDATE_IDENTITY_REQUIRED": report.get("NEW_RUNTIME_CANDIDATE_IDENTITY_REQUIRED"),
            "old_fills": (report.get("fill_comparison") or {}).get("old_fills"),
            "new_fills": (report.get("fill_comparison") or {}).get("new_fills"),
            "removed_invalid_fills": (report.get("fill_comparison") or {}).get("removed_invalid_fills"),
            "09:05_old": ((report.get("fill_comparison") or {}).get("09:05") or {}).get("old_fills"),
            "09:05_new": ((report.get("fill_comparison") or {}).get("09:05") or {}).get("new_fills"),
            "09:05_removed": ((report.get("fill_comparison") or {}).get("09:05") or {}).get("removed_invalid_fills"),
            "full14_old_pnl": (report.get("full14_old") or {}).get("pnl"),
            "full14_new_pnl": (report.get("full14_new") or {}).get("pnl"),
            "full14_old_trades": (report.get("full14_old") or {}).get("trades"),
            "full14_new_trades": (report.get("full14_new") or {}).get("trades"),
        },
    )
    for name, rows in extra_sheets.items():
        ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    xlsx = OUT / "audit.xlsx"
    if xlsx.exists():
        xlsx.unlink()
    wb.save(xlsx)
