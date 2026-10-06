"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize
from research.am_event_sequence_architecture import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_event_sequence_architecture"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "precommit",
    "channels",
    "arms",
    "increment",
    "seq_diagnostic",
    "selection_changed",
    "outer_folds",
    "daily_pnl",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def _arm_block(g: dict[str, Any], prefix: str) -> list[str]:
    return [
        f"{prefix}_NET: {_fmt(g.get(f'{prefix}_NET'))}",
        "",
        f"{prefix}_PF: {_fmt(g.get(f'{prefix}_PF'))}",
        "",
        f"{prefix}_DD: {_fmt(g.get(f'{prefix}_DD'))}",
        "",
        f"{prefix}_PAIRED_MEDIAN: {_fmt(g.get(f'{prefix}_PAIRED_MEDIAN'))}",
        "",
    ]


def _delta_block(g: dict[str, Any], prefix: str) -> list[str]:
    return [
        f"{prefix}_DELTA_NET: {_fmt(g.get(f'{prefix}_DELTA_NET'))}",
        "",
        f"{prefix}_DELTA_PF: {_fmt(g.get(f'{prefix}_DELTA_PF'))}",
        "",
        f"{prefix}_DELTA_DD: {_fmt(g.get(f'{prefix}_DELTA_DD'))}",
        "",
        f"{prefix}_DELTA_PAIRED_MEDIAN: {_fmt(g.get(f'{prefix}_DELTA_PAIRED_MEDIAN'))}",
        "",
        f"{prefix}_DELTA_POS_MINUS_NEG: {_fmt(g.get(f'{prefix}_DELTA_POS_MINUS_NEG'))}",
        "",
        f"{prefix}_DELTA_EX_TOP3: {_fmt(g.get(f'{prefix}_DELTA_EX_TOP3'))}",
        "",
    ]


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    lines = [
        "# AM EVENT SEQUENCE ARCHITECTURE V1",
        "",
        f"ANALYSIS_ID: {ANALYSIS_ID}",
        "Cross-fitted CAUSAL_GRU16 sequence posterior appended after X14 / regime extras.",
        "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
        "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0. RAW23 closed.",
        "",
        "## REQUIRED FINAL OUTPUT",
        "",
        f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
        "",
        f"SEQUENCE_LEN: {_fmt(g.get('SEQUENCE_LEN'))}",
        "",
        f"SEQUENCE_CHANNEL_N: {_fmt(g.get('SEQUENCE_CHANNEL_N'))}",
        "",
        f"SEQUENCE_MODEL: {g.get('SEQUENCE_MODEL')}",
        "",
        f"ARM_N: {_fmt(g.get('ARM_N'))}",
        "",
        f"OUTER_FOLD_N: {_fmt(g.get('OUTER_FOLD_N'))}",
        "",
    ]
    for prefix in ("B0", "S0", "B1", "S1"):
        lines.extend(_arm_block(g, prefix))
    lines.extend(_delta_block(g, "S0"))
    lines.extend(_delta_block(g, "S1"))
    lines.extend(
        [
            f"PASS_ARM_N: {_fmt(g.get('PASS_ARM_N'))}",
            "",
            f"BEST_PASS_ARM: {_fmt(g.get('BEST_PASS_ARM'))}",
            "",
            f"SEQUENCE_INCREMENTAL_SUPPORTED: {_tf(g.get('SEQUENCE_INCREMENTAL_SUPPORTED'))}",
            "",
            f"CURRENT_PRESERVATION_PASS_ALL: {_tf(g.get('CURRENT_PRESERVATION_PASS_ALL'))}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            "",
            f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            f"NEXT: {g.get('NEXT')}",
            "",
            "## STOP",
            "",
            "Four-arm 18-day OOF complete. No Runtime change. No GRU/channel/window search.",
            "RAW23 closed. submit/cancel/live=0/0/0.",
            "",
        ]
    )
    return "\n".join(lines)
