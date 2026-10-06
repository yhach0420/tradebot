"""report.json / report.md / audit.xlsx. No raw rewrite. No weekend leftover research."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.day2_futures_plus_market_breadth_acquisition_v1 import ANALYSIS_ID
from research.day2_futures_plus_market_breadth_acquisition_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = ("Answers", "Coexistence", "Decision", "Safety")


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _fmt(x: Any) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, (dict, list, tuple)):
        return json.dumps(x, ensure_ascii=False, default=str)
    return str(x)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    seq = dict(body.get("monday_sequence") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"CASE: {d.get('CASE')}",
        f"NEXT: `{d.get('NEXT')}`",
        "",
        "Raw /ranking jsonl is unchanged. ETF/ETN stay in the tape.",
        "Primary features are equity-only robust metrics. Type6/7 raw means are nonprimary.",
        "Type14/15 are one 33-sector universe. 20260912 probe is schema proof only.",
        "Day2 frozen futures price-return test is unchanged. No strategy. No ENTRY/EXIT.",
        "",
        "## Required answers",
        "",
    ]
    for i in range(1, 20):
        key = next((k for k in a if k.startswith(f"{i}_")), None)
        if key:
            lines.append(f"- {key}: `{_fmt(a.get(key))}`")
    lines.extend(
        [
            "",
            "## Monday sequence",
            "",
            f"- 07:55 futures: `{seq.get('07:55')}`",
            f"- 09:05 breadth: `{seq.get('09:05')}`",
            f"- futures PID: `{seq.get('futures_pid')}`",
            f"- breadth PID: `{seq.get('breadth_pid')}`",
            f"- cadence: `{seq.get('cadence_sec')}` fail-soft `{seq.get('fail_soft_sec')}`",
            "",
            f"submit/cancel/live: `{body.get('submit_cancel_live')}`",
            "ENTRY: false",
            "EXIT: false",
            "STRATEGY_BUILT: false",
            "",
        ]
    )
    return "\n".join(lines)


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    report = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(report), encoding="utf-8")
    wb = Workbook()
    first = True
    sheets = {
        "Answers": _kv(dict(report.get("answers") or {})),
        "Coexistence": _kv(dict(report.get("coexistence") or {})),
        "Decision": _kv(dict(report.get("decision") or {})),
        "Safety": [
            {
                "submit_cancel_live": report.get("submit_cancel_live"),
                "register_n": (report.get("answers") or {}).get("15_register_mutation"),
                "sendorder_n": (report.get("answers") or {}).get("16_sendorder"),
                "ENTRY": False,
                "EXIT": False,
                "STRATEGY_BUILT": False,
                "WEEKEND_PROBE_IS_RESEARCH_OUTCOME": False,
            }
        ],
    }
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {"report_json": str(OUT / "report.json"), "report_md": str(OUT / "report.md"), "audit_xlsx": str(xlsx)}
