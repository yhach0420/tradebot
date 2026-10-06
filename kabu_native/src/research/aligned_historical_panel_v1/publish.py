"""Write report.json / report.md / audit.xlsx / panel_manifest.json only. No mass CSV."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.aligned_historical_panel_v1 import ANALYSIS_ID
from research.aligned_historical_panel_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "UNIVERSE_BIND",
    "EQUITY_MINUTE",
    "TIME_SEMANTICS",
    "TICK_CROSSCHECK",
    "EXTERNAL",
    "SPLIT",
    "TECHNICAL",
    "QUALITY",
    "MANIFEST",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    t = dict(report.get("time_semantics") or {})
    return {
        "UNIVERSE_BIND": _kv({k: v for k, v in dict(report.get("universe_bind") or {}).items() if k != "symbols"})
        + [{"key": "symbols", "value": ",".join(dict(report.get("universe_bind") or {}).get("symbols") or [])}],
        "EQUITY_MINUTE": _kv(dict(report.get("equity_minute") or {})),
        "TIME_SEMANTICS": _kv({k: v for k, v in t.items() if k != "tick_crosscheck"}),
        "TICK_CROSSCHECK": _kv(dict(t.get("tick_crosscheck") or {})),
        "EXTERNAL": list(dict(report.get("external") or {}).get("rows") or []),
        "SPLIT": _kv(dict(report.get("split") or {})),
        "TECHNICAL": _kv(dict(report.get("technical_foundation") or {})),
        "QUALITY": _kv(dict(report.get("quality") or {})),
        "MANIFEST": _kv(dict(report.get("manifest") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            "## Required answers",
            "",
            f"1. Frozen Universe manifest verified? **{a.get('1_frozen_universe_manifest_verified')}**",
            f"2. final symbol N? **{a.get('2_final_symbol_N')}**",
            f"3. universe SHA? `{a.get('3_universe_SHA')}`",
            f"4. equity minute entitlement available? **{a.get('4_equity_minute_entitlement_available')}**",
            f"5. minute history first/last? {a.get('5_minute_history_first_last')}",
            f"6. each symbol coverage? {a.get('6_each_symbol_coverage')}",
            f"7. equity minute Time semantics? **{a.get('7_equity_minute_Time_semantics')}**",
            f"8. Tick cross-check performed? **{a.get('8_tick_crosscheck_performed')}**",
            f"9. same-bar leakage protected? **{a.get('9_same_bar_leakage_protected')}**",
            f"10. corporate action handling? {a.get('10_corporate_action_handling')}",
            f"11. NK/TOPIX source? {a.get('11_NK_TOPIX_source')}",
            f"12. futures contracts kept separate? **{a.get('12_futures_contracts_kept_separate')}**",
            f"13. USDJPY source / range? {a.get('13_USDJPY_source_range')}",
            f"14. ES/NQ source / range? {a.get('14_ES_NQ_source_range')}",
            f"15. WTI source / range? {a.get('15_WTI_source_range')}",
            f"16. rates source status? {a.get('16_rates_source_status')}",
            f"17. Asia source status? {a.get('17_Asia_source_status')}",
            f"18. canonical UTC/JST conversion pass? **{a.get('18_canonical_UTC_JST_conversion_pass')}**",
            f"19. common overlap period? {a.get('19_common_overlap_period')}",
            f"20. chronological split exact dates? {a.get('20_chronological_split_exact_dates')}",
            f"21. Discovery N days? {a.get('21_Discovery_N_days')}",
            f"22. Confirmation N days? {a.get('22_Confirmation_N_days')}",
            f"23. Frozen Validation N days? {a.get('23_Frozen_Validation_N_days')}",
            f"24. missing minute policy? {a.get('24_missing_minute_policy')}",
            f"25. source quality failures? {a.get('25_source_quality_failures')}",
            f"26. panel manifest SHA? `{a.get('26_panel_manifest_SHA')}`",
            f"27. strategy search started? **{a.get('27_strategy_search_started')}**",
            f"28. PnL used? **{a.get('28_PnL_used')}**",
            f"29. Runtime changed? **{a.get('29_Runtime_changed')}**",
            f"30. Paper changed? **{a.get('30_Paper_changed')}**",
            f"31. 20260914 live changed? **{a.get('31_20260914_live_changed')}**",
            f"32. submit/cancel/live? **{a.get('32_submit_cancel_live')}**",
            f"33. VERDICT? **{a.get('33_VERDICT')}**",
            f"34. NEXT? **{a.get('34_NEXT')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    (OUT / "panel_manifest.json").write_text(
        json.dumps(json_sanitize(dict(report.get("manifest") or {})), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
    assert (OUT / "panel_manifest.json").is_file()
