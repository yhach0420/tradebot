"""Bind FIXED_DAYTRADE_UNIVERSE_V1 from freeze artifacts. Do not modify the universe."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import load_workbook

from research.fixed_daytrade_universe_v1 import CASE_FROZEN, UNIVERSE_ID
from research.fixed_daytrade_universe_v1.analyze import universe_identity_sha
from research.aligned_historical_panel_v1.isolation import FREEZE_OUT


def bind_frozen_universe() -> dict[str, Any]:
    report_path = FREEZE_OUT / "report.json"
    manifest_path = FREEZE_OUT / "universe_manifest.json"
    xlsx_path = FREEZE_OUT / "audit.xlsx"
    missing = [str(p) for p in (report_path, manifest_path, xlsx_path) if not p.is_file()]
    if missing:
        return {
            "ok": False,
            "verified": False,
            "reason": "freeze_artifacts_missing",
            "missing": missing,
            "did_not_change_universe": True,
        }
    report = json.loads(report_path.read_text(encoding="utf-8"))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    answers = dict(report.get("answers") or {})
    decision = dict(report.get("decision") or {})
    symbols = [str(s) for s in (manifest.get("symbols_sorted") or decision.get("final_symbols") or [])]
    sha = str(manifest.get("identity_sha256") or "")
    recomputed = universe_identity_sha(symbols)
    frozen = bool(manifest.get("frozen")) and bool(decision.get("universe_frozen"))
    xlsx_sheets: list[str] = []
    xlsx_final_n = None
    liq_fail: list[str] = []
    exclusions: list[dict[str, Any]] = []
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        xlsx_sheets = list(wb.sheetnames)
        if "FINAL_UNIVERSE" in wb.sheetnames:
            ws = wb["FINAL_UNIVERSE"]
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                hdr = [str(c) if c is not None else "" for c in rows[0]]
                body = [dict(zip(hdr, r)) for r in rows[1:] if r and r[0]]
                names = [str(r.get("symbol") or "") for r in body if r.get("symbol")]
                xlsx_final_n = len(names)
                if sorted(n for n in names if n) != symbols:
                    xlsx_final_n = -1
        if "EXCLUSIONS" in wb.sheetnames:
            ws = wb["EXCLUSIONS"]
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                hdr = [str(c) if c is not None else "" for c in rows[0]]
                for raw in rows[1:]:
                    rec = dict(zip(hdr, raw))
                    sym = str(rec.get("symbol") or "")
                    if sym in {"9983", "6861"}:
                        exclusions.append({"symbol": sym, "reason": rec.get("freeze_reason") or rec.get("reason")})
        if "LIQUIDITY_60D" in wb.sheetnames:
            ws = wb["LIQUIDITY_60D"]
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                hdr = [str(c) if c is not None else "" for c in rows[0]]
                for raw in rows[1:]:
                    rec = dict(zip(hdr, raw))
                    if rec.get("liquidity_pass") is False:
                        liq_fail.append(str(rec.get("symbol") or ""))
    finally:
        wb.close()
    xlsx_symbols_ok = xlsx_final_n == 45
    ok = (
        frozen
        and str(manifest.get("universe_id")) == UNIVERSE_ID
        and str(decision.get("VERDICT")) == CASE_FROZEN
        and len(symbols) == 45
        and sha == recomputed
        and sha != ""
        and bool(answers.get("22_PANEL_CONDITIONED")) is True
        and xlsx_symbols_ok
        and "FINAL_UNIVERSE" in xlsx_sheets
        and "LIQUIDITY_60D" in xlsx_sheets
        and "EXCLUSIONS" in xlsx_sheets
        and not liq_fail
    )
    return {
        "ok": ok,
        "verified": ok,
        "did_not_change_universe": True,
        "console_verdict_not_used": True,
        "artifacts": {
            "report": str(report_path),
            "manifest": str(manifest_path),
            "audit_xlsx": str(xlsx_path),
            "xlsx_sheets": xlsx_sheets,
        },
        "universe_id": manifest.get("universe_id"),
        "frozen": frozen,
        "symbol_n": len(symbols),
        "symbols": symbols,
        "identity_sha256": sha,
        "recomputed_sha256": recomputed,
        "sha_match": sha == recomputed,
        "source_period": manifest.get("source_period"),
        "panel_conditioned_as_of": manifest.get("panel_conditioned_as_of"),
        "claim_selectable_in_2025": manifest.get("claim_selectable_in_2025"),
        "verdict": decision.get("VERDICT"),
        "xlsx_final_n": xlsx_final_n,
        "liquidity_fail_symbols": liq_fail,
        "liquidity_gates": {
            "research_listing_pass_n": answers.get("9_listing_continuity_pass_N"),
            "session_coverage_minimum": answers.get("8_session_coverage_minimum"),
            "min_p20_trading_value": answers.get("7_minimum_p20_trading_value"),
            "median_va_distribution": answers.get("6_median_trading_value_distribution"),
            "liquidity_60d_fail_n": len(liq_fail),
            "all_45_passed_liquidity_and_listing": (
                answers.get("9_listing_continuity_pass_N") == 45
                and answers.get("8_session_coverage_minimum") == 1.0
                and not liq_fail
            ),
        },
        "exclusions_9983_6861": exclusions,
        "listing_continuity_pass_n": answers.get("9_listing_continuity_pass_N"),
        "session_coverage_minimum": answers.get("8_session_coverage_minimum"),
        "min_p20_trading_value": answers.get("7_minimum_p20_trading_value"),
        "median_va_distribution": answers.get("6_median_trading_value_distribution"),
        "PANEL_CONDITIONED": bool(answers.get("22_PANEL_CONDITIONED")),
        "reason": None if ok else "frozen_universe_identity_or_manifest_invalid",
    }


assert UNIVERSE_ID == "FIXED_DAYTRADE_UNIVERSE_V1"
assert CASE_FROZEN == "FIXED_DAYTRADE_UNIVERSE_FROZEN_V1"
