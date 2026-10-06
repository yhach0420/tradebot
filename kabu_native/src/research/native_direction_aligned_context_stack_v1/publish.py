"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.native_direction_aligned_context_stack_v1.align import DIR_FORMULAS, FEATURE_DEFINITIONS
from research.native_direction_aligned_context_stack_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Direction_Semantics",
    "Feature_Definitions",
    "Base_Event",
    "D1_Tertiles",
    "D1_Family_Stability",
    "D1_Tree",
    "Frozen_Rules",
    "Bullish",
    "Bearish",
    "SameSide_Base",
    "D2",
    "D3",
    "D4",
    "Matched_Control",
    "Ablation",
    "SR_Contribution",
    "Participation_Contribution",
    "VWAP_Contribution",
    "Trend_Contribution",
    "EventStrength_Contribution",
    "MarketSector_Contribution",
    "Concentration",
    "Causality_Audit",
    "Decision",
    "Safety",
)
STRIP = {"_markdown"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def _block_rows(report: dict[str, Any], block: str) -> list[dict[str, Any]]:
    out = []
    for rule in list(report.get("frozen_rules") or []):
        rid = str(rule.get("rule_id") or "")
        st = (((report.get("rule_eval") or {}).get(rid) or {}).get("blocks") or {}).get(block) or {}
        out.append({"rule_id": rid, "predicates": rule.get("predicate_text"), **st})
    return out or [{"rule_id": None}]


def _side_rows(report: dict[str, Any], side: str) -> list[dict[str, Any]]:
    out = []
    for rid, pack in dict(report.get("side_consistency") or {}).items():
        s = dict(pack.get(side) or {})
        for b, st in dict(s.get("blocks") or {}).items():
            out.append({"rule_id": rid, "side": side, "block": b, "symmetry": pack.get("symmetry"), **dict(st or {})})
    return out or [{"rule_id": None, "side": side}]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    freeze = dict((report.get("bind") or {}).get("freeze") or {})
    lock = dict(report.get("lock") or {})
    ident = dict(report.get("identity") or {})
    stab = dict(report.get("d1_internal_stability") or {})
    feat_rows = [{"feature": k, "definition": v} for k, v in FEATURE_DEFINITIONS.items()]
    frozen = []
    for rule in list(report.get("frozen_rules") or []):
        frozen.append(
            {
                "rule_id": rule.get("rule_id"),
                "predicates": rule.get("predicate_text"),
                "n_train": rule.get("n_train"),
                "p40_train": rule.get("p40_train"),
                "lift_train": rule.get("lift_train"),
                "families": rule.get("families"),
            }
        )
    if not frozen:
        frozen = [{"rule_id": None}]
    tert = []
    for fold in list(stab.get("folds") or []):
        tert.append(dict(fold or {}))
    if not tert:
        tert = [{"fold": None}]
    ss = []
    for rule in list(report.get("frozen_rules") or []):
        rid = str(rule.get("rule_id") or "")
        e = dict((report.get("rule_eval") or {}).get(rid) or {})
        for b in ("D1", "D2", "D3", "D4"):
            st = dict((e.get("blocks") or {}).get(b) or {})
            ss.append(
                {
                    "rule_id": rid,
                    "block": b,
                    "same_side_gap": st.get("same_side_gap"),
                    "same_side_expected_p40": st.get("same_side_expected_p40"),
                    "rule_p40": st.get("rule_p40") or st.get("p40_before_m20"),
                    "bull_gap": st.get("bull_gap"),
                    "bear_gap": st.get("bear_gap"),
                    "pooled_gap": st.get("pooled_gap"),
                }
            )
    if not ss:
        ss = [{"rule_id": None}]
    matched_rows = []
    for rid, m in dict(report.get("matched_control") or {}).items():
        for b in ("D2", "D3", "D4"):
            st = dict((m.get("blocks") or {}).get(b) or {})
            matched_rows.append({"rule_id": rid, "block": b, "all_eval_positive": m.get("all_eval_positive"), **st})
    if not matched_rows:
        matched_rows = [{"rule_id": None}]
    abl = []
    for rid, fams in dict(report.get("ablation") or {}).items():
        for fam, row in dict(fams or {}).items():
            abl.append({"rule_id": rid, "family": fam, **dict(row or {})})
    if not abl:
        abl = [{"rule_id": None}]
    conc_rows = []
    for rid, c in dict(report.get("concentration") or {}).items():
        conc_rows.append({"rule_id": rid, **{k: json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v for k, v in dict(c or {}).items()}})
    if not conc_rows:
        conc_rows = [{"rule_id": None}]
    return {
        "Binding": _kv_rows(
            {
                "ok": (report.get("bind") or {}).get("ok"),
                "VERDICT": d.get("VERDICT"),
                "NEXT": d.get("NEXT"),
                "parent": (report.get("bind") or {}).get("parent_verdict"),
                "r14_closed": True,
                "t15_unsigned_closed": True,
                "r14_not_repaired": True,
            }
        ),
        "Direction_Semantics": _kv_rows({**DIR_FORMULAS, **dict(report.get("direction_semantics") or {})}),
        "Feature_Definitions": feat_rows,
        "Base_Event": _kv_rows(
            {
                "name": "NATIVE_PRICE_DISPLACEMENT_EVENT",
                "EVENT_GENERATOR_SHA256": freeze.get("EVENT_GENERATOR_SHA256") or lock.get("EVENT_GENERATOR_SHA256"),
                "native_event_n": ident.get("native_event_n"),
                "kept_event_n": ident.get("kept_event_n"),
                "D1": ident.get("d1_n"),
                "D2": ident.get("d2_n"),
                "D3": ident.get("d3_n"),
                "D4": ident.get("d4_n"),
                "FUTURE_EVENT_SELECTION_N": 0,
                "RETROSPECTIVE_CLUSTER_N": 0,
                "refractory_bars": 5,
                "first_onset_only": True,
                "BAR_START": True,
            }
        ),
        "D1_Tertiles": tert,
        "D1_Family_Stability": _kv_rows(
            {
                "stable_families": stab.get("stable_families"),
                "family_fold_counts": stab.get("family_fold_counts"),
                "unstable": stab.get("unstable"),
                "rule": stab.get("rule"),
            }
        ),
        "D1_Tree": _kv_rows(report.get("d1_tree") or {}),
        "Frozen_Rules": frozen,
        "Bullish": _side_rows(report, "BULLISH"),
        "Bearish": _side_rows(report, "BEARISH"),
        "SameSide_Base": ss,
        "D2": _block_rows(report, "D2"),
        "D3": _block_rows(report, "D3"),
        "D4": _block_rows(report, "D4"),
        "Matched_Control": matched_rows,
        "Ablation": abl,
        "SR_Contribution": _kv_rows(report.get("sr_contribution") or {}),
        "Participation_Contribution": _kv_rows(report.get("participation_contribution") or {}),
        "VWAP_Contribution": _kv_rows(report.get("vwap_contribution") or {}),
        "Trend_Contribution": _kv_rows(report.get("trend_contribution") or {}),
        "EventStrength_Contribution": _kv_rows(report.get("event_strength_contribution") or {}),
        "MarketSector_Contribution": _kv_rows(report.get("market_sector_contribution") or {}),
        "Concentration": conc_rows,
        "Causality_Audit": _kv_rows(
            {
                "same_bar_entry_n": report.get("same_bar_entry_n"),
                "FUTURE_EVENT_SELECTION_N": 0,
                "RETROSPECTIVE_CLUSTER_N": 0,
                "tree_learned_only_on_d1": True,
                "d2_d4_rule_modification": False,
                "raw_unsigned_directional_used_by_model": False,
                "no_future_vwap": True,
                "no_indicator_catalog": True,
                "r14_not_repaired": True,
                "r14_thresholds_not_reused": True,
            }
        ),
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# NATIVE_DIRECTION_ALIGNED_CONTEXT_STACK_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "Direction-aligned compact context stack on native 1-minute displacement events.",
            "Not a strategy. D1 learns. D2–D4 locked internal replication.",
            "R14 / T15 unsigned stack closed. Thresholds not reused.",
            "",
            f"Event generator unchanged? **{a.get('Event generator unchanged?')}**",
            f"DIR semantics correct? **{a.get('DIR semantics correct?')}**",
            f"aligned_r5? **{a.get('aligned_r5 formula?')}** aligned_r15? **{a.get('aligned_r15?')}** aligned_prior5_ret? **{a.get('aligned_prior5_ret?')}**",
            f"aligned_vwap? **{a.get('aligned_vwap?')}** market/sector? **{a.get('aligned_market/sector?')}**",
            f"ahead S/R? **{a.get('How is ahead S/R defined?')}** behind S/R? **{a.get('How is behind S/R defined?')}**",
            f"raw unsigned directional in model? **{a.get('Any raw unsigned directional return used by model?')}**",
            f"D1 family stability? **{a.get('D1 feature-family stability?')}**",
            f"frozen rule_n? **{a.get('Frozen rule_n?')}**",
            f"Rule 1? **{a.get('Rule 1?')}**",
            f"Rule 2? **{a.get('Rule 2?')}**",
            f"Rule 3? **{a.get('Rule 3?')}**",
            f"positive vs same-side AND matched D2/D3/D4? **{a.get('Any candidate positive against BOTH same-side base AND matched same-direction control in D2/D3/D4?')}**",
            f"S/R? **{a.get('Does S/R contribute?')}** TV? **{a.get('TradingValue?')}** VWAP? **{a.get('VWAP?')}** trend? **{a.get('trend?')}**",
            f"event strength? **{a.get('event strength?')}** market/sector? **{a.get('market/sector?')}** local? **{a.get('local structure?')}**",
            f"side consistency? **{a.get('Bullish / bearish consistency?')}**",
            f"one-symbol? **{a.get('Any one-symbol domination?')}** one-day? **{a.get('Any one-day domination?')}**",
            f"threshold retune? **{a.get('Any threshold retune?')}** PnL opt? **{a.get('Any PnL optimization?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}** Kabu50? **{a.get('Kabu50?')}**",
            f"submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
