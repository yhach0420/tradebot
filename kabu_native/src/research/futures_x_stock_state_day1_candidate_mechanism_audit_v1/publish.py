"""report.json / report.md / audit.xlsx / day2_freeze_manifest.json. No raw rewrite."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_x_stock_state_day1_candidate_mechanism_audit_v1 import ANALYSIS_ID, DAY2_ANALYSIS_ID
from research.futures_x_stock_state_day1_candidate_mechanism_audit_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Answers",
    "PerClock",
    "Path",
    "Distribution",
    "Winners",
    "Correlations",
    "Noncausal",
    "Decision",
    "Day2Freeze",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def _clock_rows(per_clock: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for c in per_clock:
        nc = dict(c.get("noncausal") or {})
        rows.append(
            {
                "clock": c.get("clock"),
                "TOP_MID": (c.get("TOP") or {}).get("MID", {}).get("mean"),
                "MIDDLE_MID": (c.get("MIDDLE") or {}).get("MID", {}).get("mean"),
                "BOTTOM_MID": (c.get("BOTTOM") or {}).get("MID", {}).get("mean"),
                "TOP_BOTTOM_MID": c.get("TOP_BOTTOM_MID"),
                "TOP_LONG_mean": (c.get("TOP") or {}).get("LONG", {}).get("mean"),
                "TOP_LONG_median": (c.get("TOP") or {}).get("LONG", {}).get("median"),
                "TOP_LONG_pos": (c.get("TOP") or {}).get("LONG", {}).get("positive_n"),
                "TOP_LONG_neg": (c.get("TOP") or {}).get("LONG", {}).get("negative_n"),
                "BOTTOM_LONG_mean": (c.get("BOTTOM") or {}).get("LONG", {}).get("mean"),
                "TOP_BOTTOM_LONG": c.get("TOP_BOTTOM_LONG"),
                "monotonic_mid": c.get("monotonic_mid"),
                "spearman_mid": c.get("spearman_mid"),
                "spearman_long": c.get("spearman_long"),
                "pre_TOP": (c.get("pre_t_mid_bps") or {}).get("TOP"),
                "pre_BOTTOM": (c.get("pre_t_mid_bps") or {}).get("BOTTOM"),
                "agree_1m": nc.get("agreement_t_plus_1m"),
                "agree_3m": nc.get("agreement_t_plus_3m"),
                "label_1m": nc.get("label_1m"),
                "label_3m": nc.get("label_3m"),
                "nk_fwd_1m": nc.get("nk_fwd_1m_bps"),
                "nk_fwd_3m": nc.get("nk_fwd_3m_bps"),
            }
        )
    return rows


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    frozen = dict(body.get("frozen") or {})
    path = ((a.get("2_TOP_MIDDLE_BOTTOM_per_clock_path") or {}).get("pooled") or {})
    top_inc = ((path.get("TOP") or {}).get("incremental_mid") or {})
    top_cum = ((path.get("TOP") or {}).get("cumulative_mid") or {})
    bot_cum = ((path.get("BOTTOM") or {}).get("cumulative_mid") or {})
    mono = dict(a.get("3_tercile_monotonicity") or {})
    rev = dict(a.get("24_future_futures_reversal_dependence") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"mechanism: `{d.get('mechanism_class')}`",
        "",
        "Frozen candidate only. No new selector, context, threshold, TopK, or horizon search.",
        f"context=`{frozen.get('context')}` selector=`{frozen.get('selector')}` "
        f"horizon=`{frozen.get('horizon')}` clocks=`{frozen.get('clocks')}`",
        "10m is a diagnostic outcome, not an EXIT. ENTRY/EXIT remain false.",
        "Post-T futures agreement is NONCAUSAL_MECHANISM_DIAGNOSTIC only.",
        "",
        "## Required answers",
        "",
        f"- 1 frozen context clocks: `{a.get('1_frozen_context_clocks')}`",
        f"- 2 TOP/MIDDLE/BOTTOM per-clock path: TOP cum 1m/3m/5m/10m="
        f"{_fmt(top_cum.get('1m'))}/{_fmt(top_cum.get('3m'))}/{_fmt(top_cum.get('5m'))}/{_fmt(top_cum.get('10m'))}; "
        f"incremental 0→1/1→3/3→5/5→10="
        f"{_fmt(top_inc.get('0_to_1m'))}/{_fmt(top_inc.get('1_to_3m'))}/"
        f"{_fmt(top_inc.get('3_to_5m'))}/{_fmt(top_inc.get('5_to_10m'))}",
        f"- 3 tercile monotonicity: pooled={mono.get('pooled_monotonic_mid')} "
        f"clocks_monotone={mono.get('clocks_monotone_mid_n')}/5 LONG={mono.get('pooled_monotonic_long')}",
        f"- 4 per-clock Spearman MID: `{a.get('4_per_clock_spearman_MID')}`",
        f"- 5 per-clock Spearman LONG: `{a.get('5_per_clock_spearman_LONG')}`",
        f"- 6 pooled TOP LONG N: `{a.get('6_pooled_TOP_LONG_N')}`",
        f"- 7 pooled LONG mean: `{_fmt(a.get('7_pooled_LONG_mean'))}` bps",
        f"- 8 median: `{_fmt(a.get('8_median'))}` bps",
        f"- 9 trimmed mean: `{_fmt(a.get('9_trimmed_mean'))}` bps",
        f"- 10 win rate: `{_fmt(a.get('10_win_rate'))}`",
        f"- 11 top1 winner contribution: `{a.get('11_top1_winner_contribution')}`",
        f"- 12 top2 contribution: `{a.get('12_top2_contribution')}`",
        f"- 13 top3 contribution: `{a.get('13_top3_contribution')}`",
        f"- 14 drop top1 LONG mean/median: `{a.get('14_drop_top1_LONG_mean_median')}`",
        f"- 15 drop top2 LONG mean/median: `{a.get('15_drop_top2_LONG_mean_median')}`",
        f"- 16 drop top3 LONG mean/median: `{a.get('16_drop_top3_LONG_mean_median')}`",
        f"- 17 drop top clock absolute LONG: `{a.get('17_drop_top_clock_absolute_LONG')}`",
        f"- 18 drop top2 clocks absolute LONG: `{a.get('18_drop_top2_clocks_absolute_LONG')}`",
        f"- 19 activity selector correlations: mean matrix in report.json / audit.xlsx",
        f"- 20 one latent family?: `{a.get('20_one_latent_family')}`",
        f"- 21 TOP pre-T MID state: `{_fmt(a.get('21_TOP_pre_T_MID_state'))}` bps",
        f"- 22 BOTTOM pre-T MID state: `{_fmt(a.get('22_BOTTOM_pre_T_MID_state'))}` bps",
        f"- 23 continuation or reversal?: `{a.get('23_continuation_or_reversal')}`",
        f"- 24 future futures reversal dependence?: `{rev}`",
        f"- 25 mechanism classification: `{a.get('25_mechanism_classification')}`",
        f"- 26 exact mechanism sentence: {a.get('26_exact_mechanism_sentence')}",
        f"- 27 Day2 primary manifest frozen?: `{_fmt(a.get('27_Day2_primary_manifest_frozen'))}`",
        f"- 28 Day2 substitutions allowed?: `{_fmt(a.get('28_Day2_substitutions_allowed'))}`",
        f"- 29 ENTRY built?: `{_fmt(a.get('29_ENTRY_built'))}`",
        f"- 30 EXIT built?: `{_fmt(a.get('30_EXIT_built'))}`",
        f"- 31 Runtime changed?: `{_fmt(a.get('31_Runtime_changed'))}`",
        f"- 32 Paper changed?: `{_fmt(a.get('32_Paper_changed'))}`",
        f"- 33 submit/cancel/live?: `{a.get('33_submit_cancel_live')}`",
        f"- 34 VERDICT?: `{a.get('34_VERDICT')}`",
        f"- 35 NEXT?: `{a.get('35_NEXT')}`",
        "",
        "## BOTTOM cumulative MID (pooled)",
        "",
        f"1m/3m/5m/10m = {_fmt(bot_cum.get('1m'))}/{_fmt(bot_cum.get('3m'))}/"
        f"{_fmt(bot_cum.get('5m'))}/{_fmt(bot_cum.get('10m'))}",
        "",
        f"Day2 freeze file: `{DAY2_ANALYSIS_ID}` primary is OBSERVED_TRADE_N_180S × BOTH_DOWN × 10m LONG TOP16. No substitutions.",
        "",
    ]
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(body), encoding="utf-8")
    manifest = json_sanitize(body.get("day2_manifest") or {})
    (OUT / "day2_freeze_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8"
    )
    per_clock = list(body.get("per_clock") or [])
    corr = dict((body.get("activity_corr_mean") or {}))
    corr_rows = []
    for a, row in corr.items():
        if isinstance(row, dict):
            rec = {"selector": a}
            rec.update(row)
            corr_rows.append(rec)
    sheets = {
        "Answers": _kv(dict(body.get("answers") or {})),
        "PerClock": _clock_rows(per_clock),
        "Path": [
            {
                "clock": c.get("clock"),
                "top_0_1": ((c.get("path") or {}).get("TOP") or {}).get("incremental_mid", {}).get("0_to_1m"),
                "top_1_3": ((c.get("path") or {}).get("TOP") or {}).get("incremental_mid", {}).get("1_to_3m"),
                "top_3_5": ((c.get("path") or {}).get("TOP") or {}).get("incremental_mid", {}).get("3_to_5m"),
                "top_5_10": ((c.get("path") or {}).get("TOP") or {}).get("incremental_mid", {}).get("5_to_10m"),
                "top_10m": ((c.get("path") or {}).get("TOP") or {}).get("cumulative_mid", {}).get("10m"),
                "bottom_10m": ((c.get("path") or {}).get("BOTTOM") or {}).get("cumulative_mid", {}).get("10m"),
            }
            for c in per_clock
        ],
        "Distribution": _kv(dict(body.get("long_distribution") or {})),
        "Winners": list(body.get("winner_contrib") or [])[:20],
        "Correlations": corr_rows,
        "Noncausal": [
            {"clock": c.get("clock"), **(c.get("noncausal") or {}), "TOP_BOTTOM_MID": c.get("TOP_BOTTOM_MID")}
            for c in per_clock
        ],
        "Decision": _kv(dict(body.get("decision") or {})),
        "Day2Freeze": _kv(dict(body.get("day2_manifest") or {})),
        "Safety": [
            {"key": "ENTRY_built", "value": False},
            {"key": "EXIT_built", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "Day2_substitutions_allowed", "value": False},
            {"key": "submit_cancel_live", "value": (body.get("answers") or {}).get("33_submit_cancel_live")},
            {"key": "future_leakage_n", "value": body.get("future_leakage_n")},
            {"key": "NONCAUSAL_MECHANISM_DIAGNOSTIC", "value": True},
        ],
    }
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(xlsx),
        "day2_freeze_manifest": str(OUT / "day2_freeze_manifest.json"),
    }
