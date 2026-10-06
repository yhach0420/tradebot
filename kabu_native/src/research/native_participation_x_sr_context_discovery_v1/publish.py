"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.native_participation_x_sr_context_discovery_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Binding",
    "Development_Status",
    "SR_Freeze",
    "TV_Causal_Baseline",
    "Volume_Diagnostic",
    "Price_Displacement",
    "First_Onset",
    "I1_Definition",
    "I2_Definition",
    "I3_Definition",
    "Away_From_SR_Control",
    "Placebo",
    "Match_Quality",
    "I1_Result",
    "I2_Result",
    "I3_Result",
    "D1_D4",
    "MFE_MAE",
    "Barrier_Races",
    "Market_Sector",
    "Concentration",
    "Multiple_Testing",
    "Causality_Audit",
    "Decision",
    "Safety",
)
STRIP = {"_markdown", "_pairs", "_away", "_placebo"}


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items()}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _kv_rows(d: Any) -> list[dict[str, Any]]:
    if isinstance(d, dict):
        return [
            {"key": k, "value": json.dumps(_json_sanitize(v), ensure_ascii=False) if isinstance(v, (dict, list)) else v}
            for k, v in d.items()
        ]
    return [{"key": "value", "value": d}]


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append(
            [
                r.get(c)
                if not isinstance(r.get(c), (dict, list))
                else json.dumps(_json_sanitize(r.get(c)), ensure_ascii=False)[:32000]
                for c in cols
            ]
        )
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(c)) + 2))


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    d = dict(report.get("decision") or {})
    freeze = dict((report.get("bind") or {}).get("freeze") or {})
    counts = dict(report.get("counts") or {})
    d14 = []
    for qid in ("I1", "I2", "I3"):
        for b, g in dict(((report.get("advance") or {}).get(qid) or {}).get("block_gaps") or {}).items():
            d14.append({"family": qid, "block": b, "p20_gap": g})
    mq = []
    for qid in ("I1", "I2", "I3"):
        cov = dict((report.get("coverage") or {}).get(qid) or {})
        mq.append({"family": qid, **{k: v for k, v in cov.items() if k != "balance"}, "balance": cov.get("balance")})
    if not mq:
        mq = [{"family": None}]
    return {
        "Binding": _kv_rows({"ok": (report.get("bind") or {}).get("ok"), "VERDICT": d.get("VERDICT"), "NEXT": d.get("NEXT")}),
        "Development_Status": _kv_rows({"POST_HOC": False, "D1_D4": "ALL_DEVELOPMENT", "not_validation": True, "not_holdout": True}),
        "SR_Freeze": _kv_rows(freeze),
        "TV_Causal_Baseline": _kv_rows({"lookback_days": 20, "min_obs": 10, "expand_pctl": 0.80, "future_normalization": False}),
        "Volume_Diagnostic": _kv_rows(report.get("volume_diagnostic") or {}),
        "Price_Displacement": _kv_rows({"prior_highs": 5, "close_loc_bull": 0.75, "close_loc_bear": 0.25, "range_median_bars": 20, "no_grid": True}),
        "First_Onset": _kv_rows({"refractory_bars": 5, "suppressed_n": counts.get("refractory_suppressed_n"), "native_event_n": counts.get("native_event_n")}),
        "I1_Definition": _kv_rows({"rule": "first-test unresolved bounce direction + first native participation at/after interaction"}),
        "I2_Definition": _kv_rows({"rule": "native event bar is first close through far boundary; S/R alone cannot create I2"}),
        "I3_Definition": _kv_rows({"rule": "after HOLD known, first LATER native event in breakout direction; event_time > hold_time"}),
        "Away_From_SR_Control": _kv_rows({"away_n": counts.get("away_n"), **dict(report.get("participation_alone") or {})}),
        "Placebo": _kv_rows(report.get("placebo") or {}),
        "Match_Quality": mq,
        "I1_Result": _kv_rows({"pooled": report.get("I1"), "by_side": report.get("I1_by_side"), "coverage": (report.get("coverage") or {}).get("I1")}),
        "I2_Result": _kv_rows({"pooled": report.get("I2"), "coverage": (report.get("coverage") or {}).get("I2")}),
        "I3_Result": _kv_rows({"pooled": report.get("I3"), "hold_without": report.get("i3_hold_without_participation"), "coverage": (report.get("coverage") or {}).get("I3")}),
        "D1_D4": d14,
        "MFE_MAE": _kv_rows({qid: ((report.get("advance") or {}).get(qid) or {}).get("path_extras") for qid in ("I1", "I2", "I3")}),
        "Barrier_Races": _kv_rows({"break_without_participation_n": report.get("break_without_participation_n"), "hold_without_i3": report.get("i3_hold_without_participation")}),
        "Market_Sector": _kv_rows(report.get("market_sector") or {"matched_on_mkt_sign_and_sec_sign": True, "not_used_as_filter": True}),
        "Concentration": _kv_rows({qid: {"symbol": ((report.get("advance") or {}).get(qid) or {}).get("symbol_conc"), "day": ((report.get("advance") or {}).get(qid) or {}).get("day_conc")} for qid in ("I1", "I2", "I3")}),
        "Multiple_Testing": _kv_rows(report.get("multiple_testing") or {}),
        "Causality_Audit": _kv_rows({"same_bar_entry_n": report.get("same_bar_entry_n"), "future_normalization": False, "retrospective": False, "clock_prior_days_only": True}),
        "Decision": _kv_rows(d),
        "Safety": _kv_rows(report.get("safety") or {}),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            "# NATIVE_PARTICIPATION_X_SR_CONTEXT_DISCOVERY_V1",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: **{d.get('NEXT')}**",
            "",
            "S/R is context. Native participation + displacement is the event. Not a strategy.",
            "D1–D4 are ALL DEVELOPMENT.",
            "",
            f"TV event_n? **{a.get('Primary TradingValue event_n?')}** symbols? **{a.get('symbol_n?')}** days? **{a.get('day_n?')}**",
            f"future norm? **{a.get('Any future normalization?')}** retrospective? **{a.get('Any retrospective event selection?')}** same-bar? **{a.get('Any same-bar assumption?')}**",
            f"I1? **{a.get('I1')}**",
            f"I2? **{a.get('I2')}**",
            f"I3? **{a.get('I3')}**",
            f"participation alone? **{a.get('Does native participation alone separate?')}** SR adds? **{a.get('Does S/R add information conditional on native participation?')}**",
            f"interaction vs placebo? **{a.get('Does S/R × participation show an interaction larger than placebo?')}**",
            f"Volume same direction? **{a.get('Does Volume show the same direction as TradingValue?')}**",
            f"A2 strategy reopened? **{a.get('A2 reopened as strategy?')}** C1? **{a.get('C1 reopened?')}** PnL opt? **{a.get('Any PnL optimization?')}**",
            f"Confirmation? **{a.get('Old Confirmation opened?')}** FV? **{a.get('Frozen Validation opened?')}** submit/cancel/live? **{a.get('submit/cancel/live?')}**",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize({k: v for k, v in report.items() if k not in STRIP})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")
