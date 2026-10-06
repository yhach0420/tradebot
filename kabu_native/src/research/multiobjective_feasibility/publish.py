"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.multiobjective_feasibility import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "multiobjective_feasibility"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "OutcomeCorr",
    "DailyCorr",
    "ParetoCurrent",
    "SpecConflict",
    "EnrichmentDays",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.8g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    dec = report.get("decision") or {}
    return "\n".join(
        [
            "# MULTI-OBJECTIVE ENTRY PARETO FEASIBILITY AUDIT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "CanonicalEngine. Common 600s cohort. Exact executable-at-decision.",
            "U=MFE_600. D=DOWNSIDE_AVOID_600. Existing T1/T2 Ridge OOF scores only.",
            "Actual Pareto is an oracle diagnostic. Not a runtime rule. No knee/weight/threshold.",
            "JOINT_FEASIBLE_MIN_RATE=0.50 frozen before seeing results.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"OUTCOME_UD_SPEARMAN_MEAN: {_fmt(g.get('OUTCOME_UD_SPEARMAN_MEAN'))}",
            f"OUTCOME_UD_SPEARMAN_MEDIAN: {_fmt(g.get('OUTCOME_UD_SPEARMAN_MEDIAN'))}",
            f"NEGATIVE_CORR_COHORT_N: {g.get('NEGATIVE_CORR_COHORT_N')}",
            "",
            f"PARETO_FRONT_SIZE_MEAN: {_fmt(g.get('PARETO_FRONT_SIZE_MEAN'))}",
            f"PARETO_FRONT_SHARE_MEAN: {_fmt(g.get('PARETO_FRONT_SHARE_MEAN'))}",
            "",
            f"CURRENT_TOP3_DOMINATED_N: {g.get('CURRENT_TOP3_DOMINATED_N')}",
            f"CURRENT_TOP3_DOMINATED_RATE: {_fmt(g.get('CURRENT_TOP3_DOMINATED_RATE'))}",
            "",
            f"JOINT_IMPROVEMENT_AVAILABLE_COHORT_N: {g.get('JOINT_IMPROVEMENT_AVAILABLE_COHORT_N')}",
            f"JOINT_IMPROVEMENT_AVAILABLE_RATE: {_fmt(g.get('JOINT_IMPROVEMENT_AVAILABLE_RATE'))}",
            "",
            f"MEDIAN_T1_T2_PREDICTED_SPEARMAN: {_fmt(g.get('MEDIAN_T1_T2_PREDICTED_SPEARMAN'))}",
            f"MEDIAN_T1_T2_TOP3_OVERLAP_RATE: {_fmt(g.get('MEDIAN_T1_T2_TOP3_OVERLAP_RATE'))}",
            "",
            f"PREDICTED_PARETO_SIZE_MEAN: {_fmt(g.get('PREDICTED_PARETO_SIZE_MEAN'))}",
            "",
            f"PARETO_MFE_DELTA_VS_CURRENT: {_fmt(g.get('PARETO_MFE_DELTA_VS_CURRENT'))}",
            f"PARETO_DOWNSIDE_DELTA_VS_CURRENT: {_fmt(g.get('PARETO_DOWNSIDE_DELTA_VS_CURRENT'))}",
            "",
            f"JOINT_OBJECTIVE_FEASIBLE: {_tf(g.get('JOINT_OBJECTIVE_FEASIBLE'))}",
            "",
            f"PRIMARY_MECHANISM: {g.get('PRIMARY_MECHANISM')}",
            f"NEXT_RESEARCH: {g.get('NEXT_RESEARCH')}",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## DECISION",
            "",
            f"CASE={dec.get('CASE')}",
            str(dec.get("note") or ""),
            "",
            "## STOP",
            "",
            "No multi-objective strategy. No weight search. No Exact. No Pareto runtime rule.",
            "Runtime/C14 unchanged. Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
