"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_current_utility_augment import ANALYSIS_ID
from research.am_entry_profit_improvement.publish import json_sanitize, kv_rows

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "am_current_utility_augment"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "precommit",
    "outer_folds",
    "augment_candidates",
    "augment_trades",
    "daily_pnl",
    "price_tail",
    "current_preservation",
    "integrity",
)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    return str(v)


def _fmt(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:.16g}"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    return "\n".join(
        [
            "# AM CURRENT-PRESERVING SELECTIVE UTILITY AUGMENT V1",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "AM only. CURRENT Top3 first. Ridge-utility median ensemble, 1 augment per cohort.",
            "Eligibility: outside CURRENT, AVAILABLE_REP_N>=6, AUG_SCORE>0, POSITIVE_REP_N>=6.",
            "Runtime WAIT_SEC remains 1.0. TRUE_OOS=false. NEW_FORWARD_N=0.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED FINAL OUTPUT",
            "",
            f"BASE_PARITY: {_tf(g.get('BASE_PARITY'))}",
            "",
            f"PRECOMMIT_SPEC_SHA256: {_fmt(g.get('PRECOMMIT_SPEC_SHA256'))}",
            "",
            f"OUTER_FOLD_N: {_fmt(g.get('OUTER_FOLD_N'))}",
            "",
            f"AUGMENT_CANDIDATE_N: {_fmt(g.get('AUGMENT_CANDIDATE_N'))}",
            "",
            f"AUGMENT_ADMITTED_N: {_fmt(g.get('AUGMENT_ADMITTED_N'))}",
            "",
            f"AUGMENT_FILL_N: {_fmt(g.get('AUGMENT_FILL_N'))}",
            "",
            f"AUGMENT_TRADE_N: {_fmt(g.get('AUGMENT_TRADE_N'))}",
            "",
            f"AUGMENT_FILL_RATE: {_fmt(g.get('AUGMENT_FILL_RATE'))}",
            "",
            f"AUGMENT_NET_PNL: {_fmt(g.get('AUGMENT_NET_PNL'))}",
            "",
            f"AUGMENT_PF: {_fmt(g.get('AUGMENT_PF'))}",
            "",
            f"AUGMENT_MAX_DD: {_fmt(g.get('AUGMENT_MAX_DD'))}",
            "",
            f"CURRENT_NET_PNL: {_fmt(g.get('CURRENT_NET_PNL'))}",
            "",
            f"CURRENT_PF: {_fmt(g.get('CURRENT_PF'))}",
            "",
            f"CURRENT_MAX_DD: {_fmt(g.get('CURRENT_MAX_DD'))}",
            "",
            f"OVERLAY_NET_PNL: {_fmt(g.get('OVERLAY_NET_PNL'))}",
            "",
            f"OVERLAY_PF: {_fmt(g.get('OVERLAY_PF'))}",
            "",
            f"OVERLAY_MAX_DD: {_fmt(g.get('OVERLAY_MAX_DD'))}",
            "",
            f"DELTA_PNL_VS_CURRENT: {_fmt(g.get('DELTA_PNL_VS_CURRENT'))}",
            "",
            f"DELTA_PF_VS_CURRENT: {_fmt(g.get('DELTA_PF_VS_CURRENT'))}",
            "",
            f"DELTA_DD_VS_CURRENT: {_fmt(g.get('DELTA_DD_VS_CURRENT'))}",
            "",
            f"PAIRED_POS_DAYS: {_fmt(g.get('PAIRED_POS_DAYS'))}",
            "",
            f"PAIRED_NEG_DAYS: {_fmt(g.get('PAIRED_NEG_DAYS'))}",
            "",
            f"PAIRED_ZERO_DAYS: {_fmt(g.get('PAIRED_ZERO_DAYS'))}",
            "",
            f"PAIRED_MEDIAN_DAILY_DELTA: {_fmt(g.get('PAIRED_MEDIAN_DAILY_DELTA'))}",
            "",
            f"EX_BEST_DAY_PNL_DELTA: {_fmt(g.get('EX_BEST_DAY_PNL_DELTA'))}",
            "",
            f"EX_TOP3_DAYS_PNL_DELTA: {_fmt(g.get('EX_TOP3_DAYS_PNL_DELTA'))}",
            "",
            f"CURRENT_ADMISSION_LOST_N: {_fmt(g.get('CURRENT_ADMISSION_LOST_N'))}",
            "",
            f"CURRENT_FILL_LOST_N: {_fmt(g.get('CURRENT_FILL_LOST_N'))}",
            "",
            f"CURRENT_EXIT_MISMATCH_N: {_fmt(g.get('CURRENT_EXIT_MISMATCH_N'))}",
            "",
            f"CURRENT_PNL_MISMATCH_N: {_fmt(g.get('CURRENT_PNL_MISMATCH_N'))}",
            "",
            f"CURRENT_CAP_INTERFERENCE_N: {_fmt(g.get('CURRENT_CAP_INTERFERENCE_N'))}",
            "",
            f"CURRENT_SAME_SYMBOL_INTERFERENCE_N: {_fmt(g.get('CURRENT_SAME_SYMBOL_INTERFERENCE_N'))}",
            "",
            f"CURRENT_PRESERVATION_PASS: {_tf(g.get('CURRENT_PRESERVATION_PASS'))}",
            "",
            f"AUGMENT_OVERLAY_PASS: {_tf(g.get('AUGMENT_OVERLAY_PASS'))}",
            "",
            f"TRUE_OOS: {_tf(g.get('TRUE_OOS'))}",
            "",
            f"NEW_FORWARD_N: {_fmt(g.get('NEW_FORWARD_N'))}",
            "",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            f"NEXT: {g.get('NEXT')}",
            "",
            "## STOP",
            "",
            "CURRENT-preserving augment OOF complete. No Runtime change. No W5 adoption.",
            "No representation selection. No threshold search. submit/cancel/live=0/0/0.",
            "",
        ]
    )
