"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.uniform10_entry_rebuild_v2 import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "uniform10_entry_rebuild_v2"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "Eligibility",
    "Feature_Inventory",
    "OOF_Folds",
    "OOF_Ranking",
    "OOF_Days",
    "TOD",
    "Robustness",
    "Comparison_ABC",
    "C2_Portfolio",
    "C2_Tail",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                obj = float(obj)
            elif isinstance(obj, np.integer):
                return int(obj)
    except Exception:
        pass
    if isinstance(obj, float):
        if obj == float("inf"):
            return "Infinity"
        if obj == float("-inf"):
            return "-Infinity"
        if obj != obj:
            return None
        return obj
    return obj


def kv_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    flat: dict[str, Any] = {}

    def _flat(prefix: str, obj: Any) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _flat(f"{prefix}.{k}" if prefix else str(k), v)
        elif isinstance(obj, tuple):
            flat[prefix] = list(obj)
        else:
            flat[prefix] = obj

    _flat("", data)
    return [{"key": k, "value": v} for k, v in flat.items()]


def _sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("gates") or {}
    return "\n".join(
        [
            "# UNIFORM10 ENTRY REBUILD V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE RESEARCH ONLY. TARGET V4 M4_PERSISTENT is the frozen label SoT.",
            "C_REBUILD_V1_INVALID_TARGET_CONTAMINATED remains SUPERSEDED. C V2 does not inherit V1 weights.",
            "C14 / Runtime / CLOCK / ENTRY / EXIT / CAP / re-entry / Dual Lane SESSION_CLOSE unchanged.",
            "Execution freshness 5s unchanged and not used as a label TTL.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "TRUE_OOS=false. NEW_FORWARD_N=0. HISTORICAL DEVELOPMENT ONLY.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"TARGET_V4_ROWS: {g.get('TARGET_V4_ROWS')}",
            f"STRUCTURAL_INELIGIBLE_ROWS: {g.get('STRUCTURAL_INELIGIBLE_ROWS')}",
            f"UNEXPECTED_TARGET_MISSING: {g.get('UNEXPECTED_TARGET_MISSING')}",
            f"OOF_MEAN_DAILY_SPEARMAN: {g.get('OOF_MEAN_DAILY_SPEARMAN')}",
            f"OOF_MEDIAN_DAILY_SPEARMAN: {g.get('OOF_MEDIAN_DAILY_SPEARMAN')}",
            f"OOF_POSITIVE_DAY_COUNT: {g.get('OOF_POSITIVE_DAY_COUNT')}",
            f"CURRENT_SCORE_OOF_SPEARMAN: {g.get('CURRENT_SCORE_OOF_SPEARMAN')}",
            f"NEW_SCORE_OOF_SPEARMAN: {g.get('NEW_SCORE_OOF_SPEARMAN')}",
            f"ALL_TARGET: {g.get('ALL_TARGET')}",
            f"TOP10_TARGET: {g.get('TOP10_TARGET')}",
            f"TOP5_TARGET: {g.get('TOP5_TARGET')}",
            f"TOP3_TARGET: {g.get('TOP3_TARGET')}",
            f"TOP1_TARGET: {g.get('TOP1_TARGET')}",
            f"SELECTED_FEATURES: {g.get('SELECTED_FEATURES')}",
            f"SELECTED_MODEL: {g.get('SELECTED_MODEL')}",
            f"C2_EXACT: {g.get('C2_EXACT')}",
            f"C2_POSITIVE_DAY_RATE: {g.get('C2_POSITIVE_DAY_RATE')}",
            f"C2_MEDIAN_DAILY_PNL: {g.get('C2_MEDIAN_DAILY_PNL')}",
            f"C2_PNL_EX_TOP3_TRADES: {g.get('C2_PNL_EX_TOP3_TRADES')}",
            f"C2_PNL_EX_TOP3_DAYS: {g.get('C2_PNL_EX_TOP3_DAYS')}",
            f"C2_PNL_EX_TOP_SYMBOL: {g.get('C2_PNL_EX_TOP_SYMBOL')}",
            f"A: {g.get('A')}",
            f"B: {g.get('B')}",
            f"C2: {g.get('C2')}",
            f"TRUE_OOS: {g.get('TRUE_OOS')}",
            f"NEW_FORWARD_N: {g.get('NEW_FORWARD_N')}",
            f"VERDICT: {g.get('VERDICT')}",
            f"RECOMMENDED_NEXT_STEP: {g.get('RECOMMENDED_NEXT_STEP')}",
            "",
            "## STOP",
            "",
            "Audit complete. Runtime not changed. No new Runtime candidate. C not implemented live.",
            "",
        ]
    )
