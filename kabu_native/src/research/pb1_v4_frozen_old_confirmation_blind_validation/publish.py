"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _base_sanitize
from research.pb1_v4_frozen_old_confirmation_blind_validation.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Manifest",
    "Precommit",
    "Data_Eligibility",
    "Candidate_Days",
    "State_Transitions",
    "Causal_Timestamps",
    "Hidden1M",
    "Thesis_Loss_Events",
    "Invariants",
    "Semantic_Adjudication",
    "Failure_Cases",
    "Safety",
)
STRIP = {
    "isolation_before",
    "isolation_after",
    "funnel_days",
    "setups",
    "e0_events",
    "e1_events",
    "minutes",
    "walked_raw",
    "candidate_days",
}
CANDIDATE_COLS = (
    "candidate_day_id",
    "symbol",
    "date",
    "WHY_THIS_STOCK",
    "OPENING_DRIVE_SEED",
    "OPENING_DRIVE_REACHED",
    "OPENING_DRIVE_LIVE",
    "LOCATION_IDENTIFIED",
    "location_family",
    "location_id",
    "THESIS_REACHED",
    "THESIS_LIVE",
    "THESIS_LOST",
    "THESIS_LOST_AT",
    "THESIS_LOST_REASON",
    "auction_end_family",
    "E0",
    "E1",
    "execution_id",
    "opening_state",
    "information_available_at",
    "event_completed_at",
    "state_changed_at",
    "entry_allowed_at",
    "same_bar_entry",
    "backdating",
)


def _json_sanitize(obj: Any) -> Any:
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    out = _base_sanitize(obj)
    if isinstance(out, float) and not math.isfinite(out):
        return None
    if isinstance(out, dict):
        return {str(k): _json_sanitize(v) for k, v in out.items() if k not in STRIP}
    if isinstance(out, list):
        return [_json_sanitize(v) for v in out]
    return out


def _excel_cell(v: Any) -> Any:
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(_json_sanitize(v), ensure_ascii=False)[:32000]
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    cols: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([_excel_cell(r.get(c)) for c in cols])
    for i, _c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(str(_c)) + 2))


def _kv(d: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": k, "value": v} for k, v in d.items()]


def _compact_candidate(row: dict[str, Any]) -> dict[str, Any]:
    return {k: row.get(k) for k in CANDIDATE_COLS}


def build_answers(report: dict[str, Any]) -> dict[str, Any]:
    dec = dict(report.get("decision") or {})
    adj = dict(report.get("adjudication") or {})
    ident = dict(report.get("identity") or {})
    pre = dict(report.get("precommit") or {})
    return {
        "VERDICT": dec.get("VERDICT"),
        "NEXT": dec.get("NEXT"),
        "FROZEN_IDENTITY": ident.get("frozen_identity") or pre.get("machine_identity"),
        "machine_sha": ident.get("machine_sha"),
        "source_inventory_sha": ident.get("source_inventory_sha"),
        "PRECOMMIT_SHA256": pre.get("PRECOMMIT_SHA256"),
        "OLD_CONFIRMATION_OPENED": bool(report.get("OLD_CONFIRMATION_OPENED")),
        "DEVELOPMENT_EXPOSED_AFTER_REVIEW": bool(adj.get("DEVELOPMENT_EXPOSED_AFTER_REVIEW") or report.get("OLD_CONFIRMATION_OPENED")),
        "FROZEN_VALIDATION_OPENED": False,
        "PROSPECTIVE_DATA_OPENED": False,
        "V4_CHANGED": False,
        "SPEC_CHANGED": False,
        "THRESHOLD_RETUNED": False,
        "PNL_USED": False,
        "MFE_MAE_USED": False,
        "FUTURE_OUTCOME_USED": False,
        "submit/cancel/live": "0/0/0",
        "candidate_day_n": (adj.get("funnel_counts") or {}).get("candidate_day_n"),
        "WHY_THIS_STOCK": (adj.get("funnel_counts") or {}).get("WHY_THIS_STOCK"),
        "SEED": (adj.get("funnel_counts") or {}).get("SEED"),
        "ACTIVE_REACHED": (adj.get("funnel_counts") or {}).get("ACTIVE_REACHED"),
        "ACTIVE_LIVE": (adj.get("funnel_counts") or {}).get("ACTIVE_LIVE"),
        "LOCATION_IDENTIFIED": (adj.get("funnel_counts") or {}).get("LOCATION_IDENTIFIED"),
        "THESIS_REACHED": (adj.get("funnel_counts") or {}).get("THESIS_REACHED"),
        "THESIS_LIVE": (adj.get("funnel_counts") or {}).get("THESIS_LIVE"),
        "THESIS_LOST": (adj.get("funnel_counts") or {}).get("THESIS_LOST"),
        "E0": (adj.get("funnel_counts") or {}).get("E0"),
        "E1": (adj.get("funnel_counts") or {}).get("E1"),
    }


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    adj = dict(report.get("adjudication") or {})
    pre = dict(report.get("precommit") or {})
    funnel = list(report.get("candidate_days") or [])
    why_rows = [_compact_candidate(r) for r in funnel if r.get("WHY_THIS_STOCK")]
    lost_rows = [
        _compact_candidate(r)
        for r in funnel
        if r.get("THESIS_LOST") and r.get("THESIS_REACHED")
    ]
    hid = dict(report.get("hidden1m") or adj.get("hidden1m") or {})
    elig = list(report.get("eligibility") or [])
    zeros = dict(adj.get("REQUIRED_ZERO") or {})
    return {
        "Manifest": _kv(dict(report.get("answers") or {})),
        "Precommit": _kv({k: v for k, v in pre.items() if k not in ("state_logging_schema",)}),
        "Data_Eligibility": elig or [{"empty": True}],
        "Candidate_Days": why_rows or [{"empty": True}],
        "State_Transitions": why_rows or [{"empty": True}],
        "Causal_Timestamps": [
            {
                "symbol": r.get("symbol"),
                "date": r.get("date"),
                "candidate_day_id": r.get("candidate_day_id"),
                "information_available_at": r.get("information_available_at"),
                "event_completed_at": r.get("event_completed_at"),
                "state_changed_at": r.get("state_changed_at"),
                "entry_allowed_at": r.get("entry_allowed_at"),
                "same_bar_entry": r.get("same_bar_entry"),
                "backdating": r.get("backdating"),
            }
            for r in why_rows
        ]
        or [{"empty": True}],
        "Hidden1M": (
            _kv({k: v for k, v in hid.items() if k != "mismatches"}) + list(hid.get("mismatches") or [])
        )
        or [{"HIDDEN_1M_THESIS_PARITY": hid.get("HIDDEN_1M_THESIS_PARITY_recomputed")}],
        "Thesis_Loss_Events": lost_rows or [{"empty": True}],
        "Invariants": _kv(zeros) + _kv({"HIDDEN_1M_THESIS_PARITY": hid.get("HIDDEN_1M_THESIS_PARITY_recomputed")}),
        "Semantic_Adjudication": list(adj.get("questions") or []) or [{"empty": True}],
        "Failure_Cases": list(adj.get("failure_cases") or []) or [{"empty": True, "n": 0}],
        "Safety": _kv(dict(report.get("safety") or {})),
    }


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _json_sanitize(report)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    ans = dict(report.get("answers") or {})
    md = [
        "# PB1 V4 Frozen Old Confirmation — blind historical validation",
        "",
        "Frozen V4 was not mutated. Frozen Validation and prospective were not opened.",
        "Purpose: semantic / causal generalization, not profitability.",
        "",
        f"**VERDICT** `{ans.get('VERDICT')}`",
        f"**NEXT** `{ans.get('NEXT')}`",
        "",
    ]
    for k, v in ans.items():
        if k in ("VERDICT", "NEXT"):
            continue
        if isinstance(v, (dict, list)):
            md.append(f"- **{k}** `{json.dumps(_json_sanitize(v), ensure_ascii=False)[:900]}`")
        else:
            md.append(f"- **{k}** `{v}`")
    md.append("")
    md.append("STOP.")
    (OUT / "report.md").write_text("\n".join(md), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _write_sheet(ws, list(sheets.get(name) or []))
    wb.save(OUT / "audit.xlsx")


def read_sheet_dicts(path, name: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    cols = [str(c) if c is not None else "" for c in rows[0]]
    out = []
    for raw in rows[1:]:
        if raw is None or all(v is None for v in raw):
            continue
        out.append({cols[i]: raw[i] if i < len(raw) else None for i in range(len(cols)) if cols[i]})
    return out

