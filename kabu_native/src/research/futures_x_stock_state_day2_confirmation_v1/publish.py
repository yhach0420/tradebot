"""report.json / report.md / audit.xlsx. No raw rewrite. No Day1 mining."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.futures_x_stock_state_day2_confirmation_v1 import ANALYSIS_ID
from research.futures_x_stock_state_day2_confirmation_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = ("Answers", "Primary", "Gates", "Secondary", "Decision", "Safety")


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _fmt(x: Any, nd: int = 3) -> str:
    if x is None:
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    if isinstance(x, (dict, list, tuple)):
        return json.dumps(x, ensure_ascii=False, default=str)
    return str(x)


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_markdown(body: dict[str, Any]) -> str:
    d = dict(body.get("decision") or {})
    a = dict(body.get("answers") or {})
    g = dict((body.get("primary_gates") or {}).get("gates") or {})
    lines = [
        f"# {ANALYSIS_ID}",
        "",
        f"VERDICT: **{d.get('VERDICT')}**",
        f"NEXT: `{d.get('NEXT')}`",
        f"day: `{body.get('trading_date')}` status: `{body.get('status')}`",
        "",
        "Frozen primary only: AGREEMENT_180S=BOTH_DOWN × OBSERVED_TRADE_N_180S TOP16/BOTTOM16 × 10m LONG TOP16.",
        "substitutions_allowed=false. Secondary rows cannot rescue a fail.",
        "PASS means Day1 cross-sectional interaction replicated. PASS does not certify ENTRY, EXIT, or Paper.",
        "20260911 mining remains closed.",
        "",
        "## Required answers",
        "",
        f"- 1 FULL?: `{_fmt(a.get('1_FULL'))}`",
        f"- 2 stock N?: `{a.get('2_stock_N')}`",
        f"- 3 NK present?: `{_fmt(a.get('3_NK_present'))}`",
        f"- 4 TOPIX present?: `{_fmt(a.get('4_TOPIX_present'))}`",
        f"- 5 Day2 BOTH_DOWN clock N?: `{a.get('5_Day2_BOTH_DOWN_clock_N')}`",
        f"- 6 TOP MID?: `{_fmt(a.get('6_TOP_MID'))}`",
        f"- 7 BOTTOM MID?: `{_fmt(a.get('7_BOTTOM_MID'))}`",
        f"- 8 TOP-BOTTOM?: `{_fmt(a.get('8_TOP_BOTTOM'))}`",
        f"- 9 BASE A?: `{_fmt(a.get('9_BASE_A'))}`",
        f"- 10 incremental lift?: `{_fmt(a.get('10_incremental_lift'))}`",
        f"- 11 TOP LONG mean?: `{_fmt(a.get('11_TOP_LONG_mean'))}`",
        f"- 12 TOP LONG median?: `{_fmt(a.get('12_TOP_LONG_median'))}`",
        f"- 13 primary PASS/FAIL?: `{a.get('13_primary_PASS_FAIL')}`",
        f"- 14 secondary rescue used?: `{_fmt(a.get('14_secondary_rescue_used'))}`",
        "",
        "## Frozen gates C1-C7",
        "",
        f"gate_status: `{(body.get('primary_gates') or {}).get('gate_status')}` "
        f"evaluated=`{_fmt((body.get('primary_gates') or {}).get('evaluated'))}` "
        f"PASS=`{_fmt((body.get('primary_gates') or {}).get('PASS'))}` "
        f"FAIL=`{_fmt((body.get('primary_gates') or {}).get('FAIL'))}`",
        "",
    ]
    for k, v in g.items():
        lines.append(f"- {k}: `{_fmt(v)}`")
    strong = (body.get("primary_gates") or {}).get("strong_confirmation")
    lines.extend(
        [
            "",
            f"strong TOP LONG median > 0: `{_fmt(None if strong is None else (strong or {}).get('TOP_LONG_median_gt_0'))}`",
            "",
            f"ENTRY: false  EXIT: false  submit/cancel/live: `{body.get('submit_cancel_live')}`",
            "",
        ]
    )
    return "\n".join(lines) + "\n"


def write_artifacts(body: dict[str, Any]) -> dict[str, str]:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize(body)
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    (OUT / "report.md").write_text(build_markdown(body), encoding="utf-8")
    primary = dict(payload.get("primary_row") or {})
    sheets = {
        "Answers": _kv(dict(payload.get("answers") or {})),
        "Primary": _kv({k: v for k, v in primary.items() if k not in ("members",)}),
        "Gates": _kv(dict((payload.get("primary_gates") or {}).get("gates") or {})),
        "Secondary": list(payload.get("secondary_rows") or []),
        "Decision": _kv(dict(payload.get("decision") or {})),
        "Safety": [
            {"key": "ENTRY", "value": False},
            {"key": "EXIT", "value": False},
            {"key": "Runtime_changed", "value": False},
            {"key": "Paper_changed", "value": False},
            {"key": "substitutions_allowed", "value": False},
            {"key": "secondary_rescue_used", "value": False},
            {"key": "submit_cancel_live", "value": payload.get("submit_cancel_live")},
            {"key": "day1_mining_closed", "value": True},
        ],
    }
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        _sheet(ws, sheets.get(name) or [])
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    return {
        "report_json": str(OUT / "report.json"),
        "report_md": str(OUT / "report.md"),
        "audit_xlsx": str(xlsx),
    }
