"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.entry_decision_population_contract import ANALYSIS_ID

NATIVE = Path(__file__).resolve().parents[3]
OUT = NATIVE / "results" / "research" / "entry_decision_population_contract"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "Summary",
    "Manifest",
    "A2Origins",
    "C3Origins",
    "Causes155",
    "ThreeWay",
    "Missing1787",
    "CoverageByAnchor",
    "ScoreParity",
    "Jaccard",
    "QA",
    "Decision",
    "Safety",
)


def json_sanitize(obj: Any) -> Any:
    try:
        import numpy as np

        if isinstance(obj, np.generic):
            if isinstance(obj, np.bool_):
                return bool(obj)
            if isinstance(obj, np.floating):
                x = float(obj)
                return None if not np.isfinite(x) else x
            if isinstance(obj, np.integer):
                return int(obj)
            return obj.item()
        if isinstance(obj, np.ndarray):
            return [json_sanitize(v) for v in obj.tolist()]
    except Exception:
        pass
    if isinstance(obj, dict):
        return {str(k): json_sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_sanitize(v) for v in obj]
    if isinstance(obj, float) and obj != obj:
        return None
    return obj


def kv_rows(d: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not d:
        return [{"key": "empty", "value": True}]
    return [{"key": k, "value": v} for k, v in d.items()]


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    body = {k: v for k, v in report.items() if k != "_markdown"}
    (OUT / "report.json").write_text(
        json.dumps(json_sanitize(body), ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    (OUT / "report.md").write_text(str(report.get("_markdown") or ""), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        rows = sheets.get(name) or [{"empty": True}]
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(name[:31])
        _sheet(ws, rows)
    wb.save(OUT / "audit.xlsx")


def _tf(v: Any) -> str:
    if v is True:
        return "true"
    if v is False:
        return "false"
    return str(v)


def build_markdown(report: dict[str, Any]) -> str:
    g = report.get("required") or {}
    q = report.get("questions") or {}
    c155 = report.get("a2_155") or {}
    miss = report.get("c3_missing") or {}
    return "\n".join(
        [
            "# ENTRY DECISION POPULATION CONTRACT AUDIT V2",
            "",
            f"ANALYSIS_ID: {ANALYSIS_ID}",
            "OFFLINE CONTRACT VERIFICATION. No new model. No C4. No execution-aware objective.",
            "C3_TOP_EDGE_SUPPORTED_PORTFOLIO_FAIL / C3_MULTIFACTOR_EXACT_FAILURE /",
            "C3_EXECUTION_COUPLING_NOT_YET_RESOLVED maintained.",
            "ENTRY origin = PENDING.anchor WAIT window. fill_time / nearest-anchor are not decision keys.",
            "SAFETY: submit/cancel/live=0/0/0. Paper=0. OPVAL=0.",
            "",
            "## REQUIRED OUTPUT",
            "",
            f"A2_CLOSED_N: {g.get('A2_CLOSED_N')}",
            f"A2_ORIGIN_FOUND_N: {g.get('A2_ORIGIN_FOUND_N')}",
            f"A2_ORIGIN_MISSING_N: {g.get('A2_ORIGIN_MISSING_N')}",
            f"A2_EXACT_EXEC_TRUE_N: {g.get('A2_EXACT_EXEC_TRUE_N')}",
            f"A2_EXACT_EXEC_FALSE_N: {g.get('A2_EXACT_EXEC_FALSE_N')}",
            f"A2_RAW_PANEL_EXEC_MISMATCH_N: {g.get('A2_RAW_PANEL_EXEC_MISMATCH_N')}",
            f"A2_RAW_EXACT_EXEC_MISMATCH_N: {g.get('A2_RAW_EXACT_EXEC_MISMATCH_N')}",
            f"A2_155_PRIMARY_CAUSE: {g.get('A2_155_PRIMARY_CAUSE')}",
            f"A2_EXEC_TRUE_THEN_NONEXEC_AT_FILL_N: {g.get('A2_EXEC_TRUE_THEN_NONEXEC_AT_FILL_N')}",
            "",
            f"C3_OOF_CLOSED_N: {g.get('C3_OOF_CLOSED_N')}",
            f"C3_OOF_ORIGIN_FOUND_N: {g.get('C3_OOF_ORIGIN_FOUND_N')}",
            f"C3_OOF_SCORE_PRESENT_AT_ORIGIN_N: {g.get('C3_OOF_SCORE_PRESENT_AT_ORIGIN_N')}",
            f"C3_FINAL_CLOSED_N: {g.get('C3_FINAL_CLOSED_N')}",
            f"C3_FINAL_ORIGIN_FOUND_N: {g.get('C3_FINAL_ORIGIN_FOUND_N')}",
            f"C3_FINAL_SCORE_PRESENT_AT_ORIGIN_N: {g.get('C3_FINAL_SCORE_PRESENT_AT_ORIGIN_N')}",
            f"C3_SCORE_MISSING_RUNTIME_HANDLING: {g.get('C3_SCORE_MISSING_RUNTIME_HANDLING')}",
            "",
            f"CURRENT_SCORE_MAX_ABS_DIFF: {g.get('CURRENT_SCORE_MAX_ABS_DIFF')}",
            f"C3_OOF_SCORE_MAX_ABS_DIFF: {g.get('C3_OOF_SCORE_MAX_ABS_DIFF')}",
            f"C3_FINAL_SCORE_MAX_ABS_DIFF: {g.get('C3_FINAL_SCORE_MAX_ABS_DIFF')}",
            "",
            f"DECISION_POPULATION_JACCARD: {g.get('DECISION_POPULATION_JACCARD')}",
            f"RESEARCH_ONLY_N: {g.get('RESEARCH_ONLY_N')}",
            f"EXACT_ONLY_N: {g.get('EXACT_ONLY_N')}",
            f"EXACT_CONTRACT_MISMATCH: {_tf(g.get('EXACT_CONTRACT_MISMATCH'))}",
            f"DECISION_POPULATION_MATCH: {_tf(g.get('DECISION_POPULATION_MATCH'))}",
            f"EXECUTION_AWARE_OBJECTIVE_ALLOWED: {_tf(g.get('EXECUTION_AWARE_OBJECTIVE_ALLOWED'))}",
            f"PRIMARY_CAUSE: {g.get('PRIMARY_CAUSE')}",
            f"VERDICT: {g.get('VERDICT')}",
            "",
            "## 155 CAUSES",
            "",
            f"N={c155.get('A2_155_N')} counts={c155.get('A2_155_CAUSE_COUNTS')} primary={c155.get('A2_155_PRIMARY_CAUSE')}.",
            str(c155.get("note") or ""),
            "",
            "## 1787 C3 MISSING HANDLING",
            "",
            str(miss.get("code_path") or ""),
            f"counts={miss.get('C3_MISSING_SCORE_HANDLING_COUNTS')} "
            f"live_C3_when_panel_OOF_missing={miss.get('LIVE_C3_SCORE_WHEN_PANEL_OOF_MISSING_N')}",
            "",
            "## Q1–Q7",
            "",
            f"Q1: {q.get('Q1')}",
            f"Q2: {q.get('Q2')}",
            f"Q3: {q.get('Q3')}",
            f"Q4: {q.get('Q4')}",
            f"Q5: {q.get('Q5')}",
            f"Q6: {q.get('Q6')}",
            f"Q7: {q.get('Q7')}",
            "",
            "## STOP",
            "",
            "No fix implemented. No C4. No execution-aware model. Runtime/C14 unchanged.",
            "Paper/OPVAL not operated. submit/cancel/live=0/0/0.",
            "",
        ]
    )
