"""Write report.json / report.md / audit.xlsx only."""
from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from research.am_entry_profit_improvement.publish import json_sanitize as _json_sanitize
from research.research_pause_and_data_requirements_redesign_v1 import ANALYSIS_ID
from research.research_pause_and_data_requirements_redesign_v1.isolation import OUT

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_ORDER = (
    "answers",
    "objective_alignment",
    "parent",
    "dataset_retirement",
    "burned_data",
    "current_information_boundary",
    "current_method_boundary",
    "meta_overfit",
    "missing_information",
    "future_method_requirements",
    "future_dev_protocol",
    "future_oos_protocol",
    "restart_gate",
    "decision",
    "safety",
)


def json_sanitize(obj: Any) -> Any:
    got = _json_sanitize(obj)
    if isinstance(got, float) and abs(got) == float("inf"):
        return "inf" if got > 0 else "-inf"
    if isinstance(got, dict):
        return {str(k): json_sanitize(v) for k, v in got.items()}
    if isinstance(got, list):
        return [json_sanitize(v) for v in got]
    return got


def _sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["empty"])
        return
    keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in keys:
                keys.append(k)
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for k in keys:
            v = r.get(k)
            if isinstance(v, (dict, list, tuple)):
                v = json.dumps(v, ensure_ascii=False, default=str)
            if isinstance(v, float) and v != v:
                v = None
            vals.append(v)
        ws.append(vals)
    for i, _k in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(_k)) + 2))


def _kv(obj: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"key": str(k), "value": v} for k, v in obj.items()]


def build_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        "answers": _kv(dict(report.get("answers") or {})),
        "objective_alignment": _kv(dict(report.get("objective_alignment") or {})),
        "parent": _kv(dict(report.get("parent") or {})),
        "dataset_retirement": _kv(dict(report.get("dataset_retirement") or {})),
        "burned_data": _kv(dict(report.get("burned_data") or {})),
        "current_information_boundary": _kv(dict(report.get("current_information_boundary") or {})),
        "current_method_boundary": _kv(dict(report.get("current_method_boundary") or {})),
        "meta_overfit": _kv(dict(report.get("meta_overfit") or {})),
        "missing_information": list(report.get("missing_information") or []),
        "future_method_requirements": _kv(dict(report.get("future_method_requirements") or {})),
        "future_dev_protocol": _kv(dict(report.get("future_dev_protocol") or {})),
        "future_oos_protocol": _kv(dict(report.get("future_oos_protocol") or {})),
        "restart_gate": _kv(dict(report.get("restart_gate") or {})),
        "decision": _kv(dict(report.get("decision") or {})),
        "safety": _kv(dict(report.get("safety") or {})),
    }


def build_markdown(report: dict[str, Any]) -> str:
    a = dict(report.get("answers") or {})
    d = dict(report.get("decision") or {})
    return "\n".join(
        [
            f"# {ANALYSIS_ID}",
            "",
            f"VERDICT: **{d.get('VERDICT')}**",
            f"NEXT: {d.get('NEXT')}",
            f"CASE: {d.get('CASE')}",
            f"STRATEGY_RESEARCH_STATUS: {d.get('STRATEGY_RESEARCH_STATUS')}",
            f"RESTART_ALLOWED: {d.get('RESTART_ALLOWED')}",
            "",
            str(d.get("INTERPRETATION") or ""),
            "",
            f"- parent pinned: {a.get('1_parent_verdict_pinned')}",
            f"- remaining eligible prior method N: {a.get('2_remaining_eligible_prior_method_N')}",
            f"- closed architecture family N: {a.get('3_closed_architecture_family_N')}",
            f"- LEGACY_DEV burned / selectable: {a.get('4_LEGACY_DEV_RESEARCH_BURNED')} / {a.get('5_legacy_DEV_usable_for_new_strategy_selection')}",
            f"- missing information categories: {a.get('17_missing_information_category_N')} {a.get('18_missing_information_category_IDs')}",
            f"- selected information category: {a.get('19_selected_recommended_information_category_ID')}",
            f"- NEW_DEV min days/weeks: {a.get('28_NEW_DEV_MIN_VALID_DAYS')} / {a.get('29_NEW_DEV_MIN_CALENDAR_WEEKS')}",
            f"- TRUE_OOS min days: {a.get('31_TRUE_OOS_MIN_VALID_DAYS')}",
            f"- future dates assigned: {a.get('32_future_dataset_dates_assigned')}",
            "",
            "Do not acquire data. Do not consume 20260907+. Do not reopen closed architectures.",
            "",
            "STOP.",
            "",
        ]
    )


def write_artifacts(report: dict[str, Any], sheets: dict[str, list[dict[str, Any]]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json_sanitize({k: v for k, v in report.items() if not str(k).startswith("_")})
    (OUT / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "report.md").write_text(str(report.get("_markdown") or build_markdown(report)), encoding="utf-8")
    wb = Workbook()
    first = True
    for name in SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        _sheet(ws, list(sheets.get(name) or []))
    xlsx = OUT / "audit.xlsx"
    wb.save(xlsx)
    assert xlsx.is_file()
